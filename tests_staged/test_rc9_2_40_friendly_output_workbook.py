"""The friendly output workbook: planner tabs first, every audit tab kept but hidden.

Business request (2026-10-04): make the outputs "more organized and more into
the points, avoid or hide unnecessary details". The published workbook opened
on ~45 visible tabs, most of them per-rule audit evidence and copies of the
input tabs.

What must hold, and is proven here on a real Voice schedule:

  * a planner sees eight tabs: Read Me First, Schedule, Break Plan (new: each
    person's shift and breaks per day), Break Schedule, FT Wise After Breaks,
    Coverage Before Breaks, Production Summary, Validation Log - plus
    No-Break Exceptions whenever it lists one;
  * no tab is removed: every audit and input tab is still in the workbook,
    hidden, so the evidence stays one right-click away;
  * the presentation changes no number: the independent validator passes the
    presented workbook with exactly the metrics it measures on the raw one;
  * the Break Plan is a view of the Break Schedule, not a second source;
  * the alternative exports (MAX_TARGET, MORE_CONSISTENT, ...) get the same
    presentation BEFORE they are validated, so what is validated is what ships.

Written to FAIL on the polisher before this change and pass after.
"""
from __future__ import annotations

import gzip
import importlib.util
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(ROOT / "engine" / "_tools"))
import l632_universal_scheduler as E  # noqa: E402

RUN = REPO / "fixtures" / "real_runs" / "language_hours"
BOOK = RUN / "Cricut_Voice_LANGUAGE_HOURS.xlsx"
CANDIDATE = RUN / "VOICE_FINAL_SHEET_CANDIDATE.json.gz"
VISIBLE = ["Read Me First", "Schedule", "Break Plan", "Break Schedule", "FT Wise After Breaks",
           "Coverage Before Breaks", "Production Summary", "Validation Log"]


def load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


POLISHER = load(ROOT / "engine" / "production" / "production_output_polisher.py", "polisher_c40")
IV = load(ROOT / "engine" / "tools" / "independent_validator.py", "iv_c40")
VALIDATOR = ROOT / "engine" / "tools" / "independent_validator.py"


class ThePresentedWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from openpyxl import load_workbook
        parsed = E.parse_input(BOOK)
        with gzip.open(CANDIDATE, "rt", encoding="utf-8") as fh:
            skeleton, solution = E.deserialize_break_candidate(parsed, json.load(fh))
        cls.tmp = Path(tempfile.mkdtemp())
        cls.raw = cls.tmp / "case_L6_3_2_3_MAX_TARGET_CANDIDATE.xlsx"
        E.write_output_workbook(BOOK, cls.raw, parsed, skeleton, solution, E.capacity_diagnostics(parsed),
                                [], {}, [], artifact_type="MAX_TARGET_CANDIDATE")
        cls.presented = cls.tmp / "presented" / cls.raw.name
        cls.presented.parent.mkdir()
        shutil.copy(cls.raw, cls.presented)
        cls.hidden = POLISHER.present_in_place(cls.presented, "MAX_TARGET_CANDIDATE", "ALTERNATIVE - test")
        cls.raw_wb, cls.wb = load_workbook(cls.raw), load_workbook(cls.presented)

    def test_a_planner_sees_the_planner_tabs_in_order(self):
        visible = [ws.title for ws in self.wb.worksheets if ws.sheet_state == "visible"]
        self.assertEqual(visible, VISIBLE)
        self.assertEqual(self.wb.active.title, "Read Me First")

    def test_no_tab_was_removed_only_hidden(self):
        raw = set(self.raw_wb.sheetnames)
        self.assertTrue(raw <= set(self.wb.sheetnames), raw - set(self.wb.sheetnames))
        self.assertEqual(set(self.wb.sheetnames) - raw, {"Break Plan"})
        self.assertEqual(self.hidden, len(self.wb.sheetnames) - len(VISIBLE))
        self.assertGreater(self.hidden, 30)

    def test_the_validator_measures_exactly_the_same_schedule(self):
        before = IV.validate(BOOK, self.raw, ROOT / "engine" / "_tools" / "l632_universal_scheduler.py")
        after = IV.validate(BOOK, self.presented, ROOT / "engine" / "_tools" / "l632_universal_scheduler.py")
        self.assertEqual(after["hard_fail_count"], 0, after["failures"][:3])
        self.assertEqual(after["metrics"], before["metrics"])

    def test_the_break_plan_is_a_view_of_the_break_schedule(self):
        plan, breaks = self.wb["Break Plan"], self.wb["Break Schedule"]
        heads = [c.value for c in plan[3]]
        self.assertEqual(heads[:2] + heads[2:9], ["Associate", "Language", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"])
        rows = [r for r in breaks.iter_rows(min_row=2, values_only=True) if r[0]]
        who, day = rows[0][0], rows[0][1]
        short = {"Break 1": "B1", "Break 2": "B2", "Lunch": "L"}
        expected = [f"{short.get(r[3], r[3])} {r[4]}" for r in rows if r[0] == who and r[1] == day]
        line = next(r for r in plan.iter_rows(min_row=4, values_only=True) if r[0] == who)
        cell = line[2 + ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"].index(day)]
        shift, plan_breaks = cell.split("\n")
        self.assertEqual(plan_breaks.split(" | "), expected)
        self.assertEqual(shift, rows[0][2])
        roster = [a.name for a in E.parse_input(BOOK).associates]
        self.assertEqual([r[0] for r in plan.iter_rows(min_row=4, values_only=True) if r[0]], roster)

    def test_no_break_exceptions_show_only_when_there_are_some(self):
        self.assertFalse(POLISHER._has_exceptions(self.wb))
        self.assertNotIn("No-Break Exceptions", [ws.title for ws in self.wb.worksheets if ws.sheet_state == "visible"])


class AlternativesArePresentedBeforeTheyAreValidated(unittest.TestCase):
    SOURCE = (ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py").read_text()

    def test_the_runner_presents_then_validates(self):
        block = self.SOURCE[self.SOURCE.index("def validate_alternative_exports("):]
        self.assertLess(block.index("presentation = present_alternative(path, role)"),
                        block.index("proc = subprocess.run(command"))
        self.assertIn("'presentation': presentation}", block)

    def test_a_presentation_failure_is_recorded_not_hidden(self):
        block = self.SOURCE[self.SOURCE.index("def present_alternative("):]
        block = block[:block.index("\ndef ")]
        self.assertIn("return f'FAILED: {type(exc).__name__}: {exc}'", block)
        self.assertIn("presentation failed", block)

    def test_the_published_schedule_uses_the_same_arrangement(self):
        source = (ROOT / "engine" / "production" / "production_output_polisher.py").read_text()
        clean = source[source.index("def clean_book("):source.index("def publish(")]
        # Re-pinned in Phase H (tests_staged/test_rc9_2_53_one_output_look.py):
        # the arrangement now takes the role, so the before-breaks and shortfall
        # schedules get their own visible tabs; the published schedule still goes
        # through the same _arrange_tabs.
        self.assertIn("_arrange_tabs(wb,role)", clean)


if __name__ == "__main__":
    unittest.main()
