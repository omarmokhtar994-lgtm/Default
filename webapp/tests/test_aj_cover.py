# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AJ item 2 (owner, 2026-10-15: "we might add someone from TM or suppot members not part of the roaster to
login to cover intervals they should count in interval wise as well so we need an option to add that explaining
he/she is support or tm and duration they will logjn to save an interval"). Answers: the name is typed; they count
as one agent for the minutes they cover; they pick a channel when the program has channels; posted to the group like
other changes; the export has a tick to include them or not."""
import html
import io
import re
import shutil
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from webapp.attendance import COVER_ROLES, DayBook
from webapp.exports import build
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_attendance import WED, THU
from webapp.tests.test_channel_day import _Day
from webapp.tests.test_overtime_side import _Overtime
from webapp.tests.test_runs import token
from webapp.tests.test_schedules import AFTER, INPUT, REPO

AR = "AE/AR B2B"
TM = "Team lead 01"  # a made-up name, typed as the RTA would


class _Book(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program=AR, week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))

    def cell(self, t, on=WED, **kw):
        return next(c for c in self.days.page(AR, on, **kw)["view"]["cells"] if c["t"] == t)

    def add(self, start="10:00", end="12:00", name=TM, role="Team manager", on=WED, **kw):
        return self.days.add_cover(AR, on, name, role, start, end, self.sara, **kw)


class TheCount(_Book):
    def test_roles(self):
        self.assertEqual(COVER_ROLES, ("Team manager", "Support"))

    def test_a_whole_interval_counts_one_more_on_the_floor(self):
        step = self.days.page(AR, WED)["view"]["interval"]
        before = self.cell(600)
        self.add("10:00", plain(600 + step))
        after = self.cell(600)
        self.assertEqual(after["now"], before["now"] + 1)
        self.assertEqual(after["low"], before["low"] + 1)
        self.assertEqual(after["outside"], 1)
        self.assertEqual(after["plan"], before["plan"])  # not in the plan
        if before["required"] > 0:
            self.assertAlmostEqual(after["pm"], before["pm"] + step / 60, places=2)

    def test_part_of_an_interval_counts_in_part(self):
        step = self.days.page(AR, WED)["view"]["interval"]
        before = self.cell(600)
        self.add("10:00", plain(600 + step // 2))
        after = self.cell(600)  # every 5 minutes of the first half has one more ("now" itself is rounded to 0.1)
        self.assertEqual([a - b for a, b in zip(after["slots"], before["slots"])],
                         [1] * (step // 10) + [0] * (step // 10))

    def test_past_midnight_counts_on_the_next_day(self):
        before, today = self.cell(0, THU), self.cell(23 * 60)
        self.add("23:00", "01:00")
        self.assertEqual(self.cell(0, THU)["now"], before["now"] + 1)
        self.assertEqual(self.cell(0, THU)["outside"], 1)
        self.assertEqual(self.cell(23 * 60)["now"], today["now"] + 1)

    def test_not_a_person_on_the_roster(self):
        before = self.days.page(AR, WED)["view"]
        self.add()
        after = self.days.page(AR, WED)["view"]
        self.assertEqual([l["name"] for l in after["lanes"]], [l["name"] for l in before["lanes"]])
        for key in ("planned", "present", "absent", "overtime"):
            self.assertEqual(after["tiles"][key], before["tiles"][key])
        self.assertEqual(after["tiles"]["cover"], 120)
        self.assertEqual(after["languages"], before["languages"])  # their language is not known

    def test_left_out_when_asked(self):
        before = self.cell(600)
        self.add()
        self.assertEqual(self.cell(600, cover=False)["now"], before["now"])

    def test_listed_with_who_kept_it(self):
        made = self.add(role="Support")
        self.assertEqual([(c["id"], c["associate"], c["role"], c["start"], c["end_min"], c["channel"])
                          for c in self.days.covers(AR, WED)], [(made, TM, "Support", 600, 720, "")])


class TheChecks(_Book):
    def refused(self, said, **kw):
        with self.assertRaises(ValueError) as got:
            self.add(**kw)
        self.assertEqual(str(got.exception), said)

    def test_each_refusal_says_why(self):
        self.refused("Type the name of the person covering.", name="  ")
        self.refused("Keep the name to 60 characters.", name="x" * 61)
        self.refused("Pick Team manager or Support.", role="Manager")
        self.refused("Give the start and end like 13:05.", start="10")
        self.refused("Cover goes in 5-minute steps.", start="10:03")
        self.refused("The end time must be after the start time.", start="10:00", end="10:00")
        self.refused("Cover is up to 12 hours.", start="08:00", end="20:05")
        self.refused("This program has no channels set up, so there is no channel to pick.", channel="P")
        self.assertIsNone(self.days.page(AR, date(2026, 11, 4)))
        self.refused(f"There is no schedule for {AR} in the week of 01 Nov.", on=date(2026, 11, 4))

    def test_one_name_cannot_cover_twice_at_once(self):
        self.add("10:00", "12:00")
        self.refused(f"{TM} already covers 10:00 to 12:00.", start="11:00", end="13:00")
        self.add("12:00", "13:00")  # back to back is fine
        self.add("11:00", "12:00", name="Support 02", role="Support")  # someone else is fine

    def test_a_dry_run_keeps_nothing(self):
        found = self.add(dry_run=True)
        self.assertEqual(found["what"], "Cover from outside the roster (Team manager) 10:00 to 12:00")
        self.assertEqual(self.days.covers(AR, WED), [])


class TheRecord(_Book):
    def test_logged_cancelled_and_undone(self):
        made = self.add()
        log = self.days.page(AR, WED)["log"]
        self.assertEqual((log[-1]["associate"], log[-1]["what"]),
                         (TM, "Cover from outside the roster (Team manager) 10:00 to 12:00"))
        self.days.cancel_cover(AR, WED, made, self.sara)
        self.assertEqual(self.days.covers(AR, WED), [])
        self.assertEqual(self.days.page(AR, WED)["log"][-1]["what"],
                         "Cancelled: cover from outside the roster (Team manager) 10:00 to 12:00")
        self.assertEqual(self.days.undo_last(AR, WED, self.sara),
                         f"Undone: {TM}, Cancelled: cover from outside the roster (Team manager) 10:00 to 12:00.")
        self.assertEqual(len(self.days.covers(AR, WED)), 1)
        self.assertEqual(self.days.page(AR, WED)["log"][-1]["what"], "Cover 10:00 to 12:00 back (undone)")

    def test_cancel_only_on_its_own_day(self):
        made = self.add()
        with self.assertRaises(ValueError):
            self.days.cancel_cover(AR, THU, made, self.sara)

    def test_old_days_are_cleaned_up(self):
        self.add()
        self.store.delete_day_before((WED + timedelta(days=1)).isoformat())
        self.assertEqual(self.days.covers(AR, WED), [])


class TheExport(_Book):
    def sheets(self, kinds):
        data, _, _ = build(self.store, self.days, WED, WED, kinds, program=AR)
        return load_workbook(io.BytesIO(data))

    def test_ticked_it_has_its_own_sheet_and_counts_in_the_summary(self):
        alone = self.sheets(["summary"])["Daily summary"]
        self.add("10:00", "12:00")
        wb = self.sheets(["summary", "cover"])
        rows = list(wb["Cover from outside the roster"].values)
        self.assertEqual(rows[0][:7], ("Shift date", "Program", "Name", "Role", "From", "To", "Minutes"))
        self.assertEqual(rows[1][:7], ("2026-10-14", AR, TM, "Team manager", "10:00", "12:00", 120))
        head = list(alone.values)[0]
        short = head.index("Hours short")
        with_cover = list(wb["Daily summary"].values)
        self.assertEqual(with_cover[0][-1], "Cover min")
        self.assertEqual(with_cover[1][-1], 120)
        before = float(list(alone.values)[1][short])
        self.assertLessEqual(float(with_cover[1][short]), before)

    def test_unticked_the_export_is_the_roster_alone(self):
        alone = list(self.sheets(["summary"])["Daily summary"].values)
        self.add("10:00", "12:00")
        again = list(self.sheets(["summary"])["Daily summary"].values)
        changes = alone[0].index("Changes recorded")  # the day's audit count: adding the cover is a change made
        self.assertEqual(again[1][changes], alone[1][changes] + 1)
        self.assertEqual([r[:changes] + r[changes + 1:] for r in again], [r[:changes] + r[changes + 1:] for r in alone])
        self.assertNotIn("Cover from outside the roster", self.sheets(["summary"]).sheetnames)


class TheChannels(_Day):
    def test_they_pick_a_channel_and_count_on_it(self):
        self.assertEqual(self.cell("Phone", 14 * 60)["low"], 0)
        with self.assertRaises(ValueError) as said:
            self.days.add_cover(self.key, date(2026, 10, 14), TM, "Support", "14:00", "15:00", self.omar)
        self.assertEqual(str(said.exception), "Pick the channel Team lead 01 takes: Phone or Chat.")  # the fixture's two
        with self.assertRaises(ValueError):
            self.days.add_cover(self.key, date(2026, 10, 14), TM, "Support", "14:00", "15:00", self.omar, channel="X")
        self.days.add_cover(self.key, date(2026, 10, 14), TM, "Support", "14:00", "15:00", self.omar, channel="P")
        self.assertEqual(self.cell("Phone", 14 * 60)["low"], 1)
        self.assertEqual(self.days.page(self.key, date(2026, 10, 14))["log"][-1]["what"],
                         "Cover from outside the roster (Support) 14:00 to 15:00, Phone")


class ThePage(_Overtime):
    def form(self, **data):
        return {"csrf_token": token(self.client), "program": self.key, "date": "2026-10-14", "view": "cover",
                "name": TM, "role": "Support", "from": "10:00", "to": "11:00", **data}

    def page(self):
        return html.unescape(self.client.get(f"/day?program={self.key.replace(' ', '+')}&date=2026-10-14&view=cover")
                             .get_data(as_text=True))

    def test_added_listed_and_deleted_from_the_overtime_and_vto_tab(self):
        body = self.page()
        self.assertIn("Cover from outside the roster", body)
        self.assertRegex(body, r'<form method="post" action="/day/cover"')
        got = self.client.post("/day/cover", data=self.form())
        self.assertEqual(got.status_code, 303)
        self.assertIn("view=cover", got.headers["Location"])
        body = self.page()
        self.assertIn(f"{TM} (Support) 10:00 to 11:00", body)
        made = re.search(r'action="/day/cover/cancel".*?name="id" value="(\d+)"', body, re.S).group(1)
        self.client.post("/day/cover/cancel", data={"csrf_token": token(self.client), "program": self.key,
                                                    "date": "2026-10-14", "id": made, "view": "cover"})
        self.assertNotIn(f"{TM} (Support) 10:00 to 11:00</b>", self.page())

    def test_a_refusal_is_said(self):
        got = self.client.post("/day/cover", data=self.form(name=""), follow_redirects=True)
        self.assertIn("Type the name of the person covering.", html.unescape(got.get_data(as_text=True)))


def plain(m):
    return f"{m // 60 % 24:02d}:{m % 60:02d}"


if __name__ == "__main__":
    unittest.main()


class TheBoard(_Book):
    def test_the_board_says_who_is_covering(self):
        from webapp.day import board
        self.add("10:00", "12:00", role="Support")
        rows = board(self.days.page(AR, WED)["view"], 0.9)
        row = next(r for r in rows if r["t"] == 600)
        self.assertIn({"name": TM, "text": "support from outside the roster, covering 10:00 to 12:00"}, row["events"])

    def test_the_24_hours_from_a_start_keep_one_stretch_past_midnight(self):
        from webapp.day import window_view
        self.add("23:00", "01:00")
        views = {d: (self.days.page(AR, WED + timedelta(days=d)) or {}).get("view") for d in (0, 1)}
        found = window_view(views, 12 * 60, WED)
        self.assertEqual([(o["name"], o["start"], o["end"]) for o in found["outside"]], [(TM, 1380, 1500)])


class TheLanguage(_Book):
    """Phase AJ item 4 (owner, 2026-10-15: "For the support tm cover should count in language as well so add language
    part"): a cover says the language the person takes, from the program's languages, and counts for it."""

    def lang(self, t, **kw):
        row = next(l for l in self.days.page(AR, WED, **kw)["view"]["languages"] if l["name"] == "English")
        return next(c for c in row["cells"] if c["t"] == t)

    def test_a_cover_in_a_language_counts_for_it(self):
        before = self.lang(600)["count"]
        self.add("10:00", "12:00", language="English")
        self.assertEqual(self.lang(600)["count"], before + 1)
        self.assertEqual(self.days.covers(AR, WED)[0]["language"], "English")
        self.assertEqual(self.days.page(AR, WED)["log"][-1]["what"],
                         "Cover from outside the roster (Team manager) 10:00 to 12:00, English")

    def test_without_a_language_it_counts_for_none(self):
        before = self.lang(600)["count"]
        self.add("10:00", "12:00")
        self.assertEqual(self.lang(600)["count"], before)

    def test_only_the_programs_languages(self):
        with self.assertRaises(ValueError) as said:
            self.add(language="Klingon")
        self.assertEqual(str(said.exception), f"Pick the language {TM} takes: English, or none.")

    def test_the_tab_offers_the_languages(self):
        self.assertEqual(self.days.cover_languages(self.days.page(AR, WED)), ["English"])


class TheChannelLanguage(_Day):
    def test_a_cover_counts_for_its_channel_and_language(self):
        rows = self.view()["channels"]["rows"]
        rule = next(r for r in rows if r.get("language"))
        cell = next(c for c in rule["cells"] if c["need"])
        names = {"P": "Phone", "C": "Chat", "E": "Email"}
        lo = cell["t"]
        self.days.add_cover(self.key, date(2026, 10, 14), TM, "Support", plain(lo), plain(lo + 60), self.omar,
                            channel=rule["letter"], language=rule["language"])
        after = next(c for r in self.view()["channels"]["rows"] if r["name"] == rule["name"] for c in r["cells"]
                     if c["t"] == lo)
        self.assertEqual(after["low"], cell["low"] + 1, names[rule["letter"]])


class TheUpgrade(unittest.TestCase):
    def test_a_database_from_the_first_phase_aj_package_gets_the_language_column(self):
        import sqlite3
        folder = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, folder, True)
        Store(folder / "scheduler.db")
        db = sqlite3.connect(folder / "scheduler.db")
        db.execute("drop table outside_cover")
        db.execute("create table outside_cover (id integer primary key autoincrement, program text not null, "
                   "shift_date text not null, associate text not null, role text not null, start integer not null, "
                   "end_min integer not null, channel text not null default '', user_id integer not null, "
                   "at real not null)")
        db.execute("insert into outside_cover (program, shift_date, associate, role, start, end_min, user_id, at) "
                   "values ('X', '2026-10-15', 'Team lead 01', 'Support', 600, 660, 1, 0)")
        db.commit()
        db.close()
        store = Store(folder / "scheduler.db")
        self.assertEqual(store.list_cover("X", ["2026-10-15"])[0]["language"], "")
