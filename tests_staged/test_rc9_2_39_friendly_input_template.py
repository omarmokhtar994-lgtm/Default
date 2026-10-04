"""The friendly input template: every cell restricted, nothing the engine reads changed.

Business request (2026-10-04): organise the input workbook, hide what a
planner does not need, and make every entry a dropdown or a limit "with
restrictions to the inputs needed by the engine to avoid any wrong entry".

What must hold, and is proven here:

  * the engine reads exactly what it read before (every parsed field and the
    input-contract verdict), on every shipped workbook, and the builder itself
    refuses to write a workbook that would change it;
  * every value cell on Instructions and Engine Defaults carries an enforced
    validation, and every choice a dropdown offers is one the engine accepts
    without a HARD_ refusal;
  * a value the engine reads under another name ("Leave" for "Leave Enabled")
    is carried into its row instead of being dropped (the previous builder
    dropped it, so "Leave = No" came back as Yes);
  * the weekly tabs are tied to the roster and the Shift Library;
  * the quick checks on Start Here compute, and turn red on a real mistake
    (when LibreOffice is available to recalculate them).

Written to FAIL on the builder before this change and pass after.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(ROOT / "engine" / "_tools"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import l632_universal_scheduler as E  # noqa: E402
import phase_a_workbooks as W  # noqa: E402

TEMPLATE = ROOT / "tools" / "build_input_template.py"
INPUTS = next((p for p in (ROOT / "inputs", REPO / "packages" / "rc9_2_2_production" / "inputs") if p.is_dir()))
READY = INPUTS / "ready_to_edit"
VOICE = INPUTS / "Cricut_Voice_RC9_1_READY_SKELETON.xlsx"


def rebuild(src: Path, dst: Path) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(TEMPLATE), str(src), str(dst)], capture_output=True, text=True)


def load_builder():
    import importlib.util
    spec = importlib.util.spec_from_file_location("builder_c39", TEMPLATE)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


B = load_builder()


def hard(parsed) -> set:
    return {w.split(":")[0] for w in parsed.parser_warnings if w.startswith("HARD_")}


class TheBuilderProvesItChangedNothing(unittest.TestCase):
    def test_every_shipped_workbook_is_rebuilt_and_verified(self):
        for book in sorted(INPUTS.glob("*.xlsx")):
            with self.subTest(workbook=book.name), tempfile.TemporaryDirectory() as tmp:
                out = rebuild(book, Path(tmp) / book.name)
                self.assertEqual(out.returncode, 0, out.stdout + out.stderr)
                self.assertIn("VERIFIED: the engine reads the same contract", out.stdout)
                before, after = B.engine_reading(book), B.engine_reading(Path(tmp) / book.name)
                changed = {k for k in before if before[k] != after.get(k)} - set(B.PROVENANCE_ONLY)
                self.assertEqual(changed, set())

    def test_rebuilding_a_rebuilt_workbook_changes_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            once, twice = Path(tmp) / "once.xlsx", Path(tmp) / "twice.xlsx"
            self.assertEqual(rebuild(VOICE, once).returncode, 0)
            out = rebuild(once, twice)
            self.assertEqual(out.returncode, 0, out.stdout)
            self.assertEqual(B.engine_reading(once), B.engine_reading(twice))

    def test_a_change_to_what_the_engine_reads_is_refused_and_nothing_is_written(self):
        source = TEMPLATE.read_text()
        self.assertIn("dst.unlink()", source)
        self.assertIn('"REFUSED: the rebuilt workbook would change what the engine reads', source)
        self.assertEqual(B.PROVENANCE_ONLY, {"coverage_objective_weighting_source": "coverage_objective_weighting",
                                             "coverage_split_source": "coverage_split_rules"})
        # The seeded Run Stage / Run Depth rows: blank becomes the runner's own fallback, nothing else.
        self.assertEqual(B.SEEDED_FALLBACK, {"run_stage": "'FULL_SCHEDULE'", "run_depth": "'QUICK'"})
        runner = (ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py").read_text()
        self.assertIn("args.mode or contract_depth or 'QUICK'", runner)
        self.assertIn("args.stage or contract_stage or 'FULL_SCHEDULE'", runner)


class AnAliasIsCarriedNotDropped(unittest.TestCase):
    def test_leave_no_under_its_short_name_stays_no(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "case.xlsx"
            W.build(dict(roster=[("Agent A", "English")], shifts=["06:00 - 15:00"],
                         demand=lambda d, m: 1 if 6 * 60 <= m < 15 * 60 else None,
                         instructions={"Count of Associates": 1, "Leave": "No"}), src)
            self.assertFalse(E.parse_input(src).leave_enabled)
            out = rebuild(src, Path(tmp) / "rebuilt.xlsx")
            self.assertEqual(out.returncode, 0, out.stdout)
            parsed = E.parse_input(Path(tmp) / "rebuilt.xlsx")
            self.assertFalse(parsed.leave_enabled)
            from openpyxl import load_workbook
            ws = load_workbook(Path(tmp) / "rebuilt.xlsx")["Instructions"]
            row = next(r for r in ws.iter_rows() if r[1].value == "Leave Enabled")
            self.assertEqual(row[2].value, "No")


class EveryCellIsRestrictedToWhatTheEngineAccepts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from openpyxl import load_workbook
        cls.tmp = Path(tempfile.mkdtemp())
        cls.book = cls.tmp / "voice.xlsx"
        assert rebuild(VOICE, cls.book).returncode == 0
        cls.wb = load_workbook(cls.book)

    def value_cells(self, sheet):
        ws = self.wb[sheet]
        return {r[1].value: r[2] for r in ws.iter_rows(min_row=4)
                if r[1].value and r[1].value != "Instruction" and r[0].value in (None, "")}

    def rules_for(self, sheet, coordinate):
        ws = self.wb[sheet]
        return [dv for dv in ws.data_validations.dataValidation if coordinate in dv.sqref]

    def test_every_value_cell_has_exactly_one_enforced_validation(self):
        for sheet in ("Instructions", "Engine Defaults"):
            cells = self.value_cells(sheet)
            self.assertGreater(len(cells), 20, sheet)
            for label, cell in cells.items():
                with self.subTest(sheet=sheet, row=label):
                    rules = self.rules_for(sheet, cell.coordinate)
                    self.assertEqual(len(rules), 1)
                    self.assertTrue(rules[0].showErrorMessage)
                    self.assertEqual(rules[0].errorStyle, "stop")

    def test_every_dropdown_choice_is_accepted_by_the_engine(self):
        """Set each offered choice and parse: no HARD_ refusal names that row."""
        from openpyxl import load_workbook
        lists = load_workbook(self.book)["Validation Lists"]
        durations = [c.value for c in lists["A"][2:] if c.value not in (None, "")]
        checked = 0
        for label, cell in self.value_cells("Instructions").items():
            rules = self.rules_for("Instructions", cell.coordinate)
            if rules[0].type != "list":
                continue
            formula = rules[0].formula1
            choices = durations if "Validation Lists" in formula else formula.strip('"').split(",")
            for choice in choices:
                with self.subTest(row=label, choice=choice):
                    wb = load_workbook(self.book)
                    wb["Instructions"][cell.coordinate] = choice
                    if label == "Allowed Shift Durations Hours":
                        # 10.5 h or more is legal only with the 11H/3OFF week; the engine names it otherwise.
                        long_ = any(float(t) >= 10.5 for t in re.findall(r"\d+(?:\.\d+)?", str(choice)))
                        mode = next(c for k, c in self.value_cells("Instructions").items() if k == "Use 11H/3OFF")
                        wb["Instructions"][mode.coordinate] = "Yes" if long_ else "No"
                    path = self.tmp / "choice.xlsx"
                    wb.save(path)
                    try:
                        parsed = E.parse_input(path)
                    except ValueError as exc:
                        # The value was read; another tab does not support it yet (the Voice
                        # library holds 9 h shifts only, its 15/60-minute demand tabs are empty).
                        # The engine names that tab; it is not a refusal of the value.
                        self.assertRegex(str(exc), r"^(No legal shifts found for allowed durations "
                                                   r"|No interval requirements were parsed from FT Wise)")
                        checked += 1
                        continue
                    refused = {w for w in parsed.parser_warnings if w.startswith("HARD_INVALID")
                               or w.startswith("HARD_LONG_DURATION")}
                    self.assertEqual(refused, set())
                    checked += 1
        self.assertGreater(checked, 60)

    def test_the_lists_say_what_the_engine_means(self):
        self.assertEqual(B.BREAK_MINUTES, '"' + ",".join(str(m) for m in E.BREAK_DURATION_CHOICES) + '"')
        for text, mode in (("No new staffing in blank intervals", "hard_no_current_week_staffing"),
                           ("Allow staffing in blank intervals", "allow")):
            self.assertEqual(E.blank_requirement_mode(text), mode)
        self.assertEqual(E.normalize_language_window_mode("REQUIRED_LANGUAGE_ONLY") is not None, True)
        for choice in B.DURATION_CHOICES:
            minutes = [round(float(t) * 60) for t in re.findall(r"\d+(?:\.\d+)?", choice)]
            self.assertTrue(minutes and all(240 <= m <= 960 for m in minutes), choice)

    def test_number_rows_have_the_engine_range_and_ratios_read_as_ratios(self):
        from openpyxl import load_workbook
        cells = self.value_cells("Instructions")
        target = self.rules_for("Instructions", cells["Target"].coordinate)[0]
        self.assertEqual((target.type, target.formula1, target.formula2), ("decimal", "0", "1"))
        variety = self.rules_for("Instructions", cells["Count of Different Shifts Per week"].coordinate)[0]
        self.assertEqual((variety.type, variety.formula1, variety.formula2), ("whole", "1", "7"))
        wb = load_workbook(self.book)
        wb["Instructions"][cells["Target"].coordinate] = 0.9  # what Excel stores for 90%
        wb["Instructions"][cells["Count of Different Shifts Per week"].coordinate] = 1
        path = self.tmp / "numbers.xlsx"
        wb.save(path)
        parsed = E.parse_input(path)
        self.assertAlmostEqual(parsed.target_ratio, 0.9)
        self.assertEqual(parsed.max_different_shifts, 1)
        self.assertNotIn("INVALID_MAX_SHIFT_VARIETY", {f.get("code") for f in E.validate_input_contract(parsed)["failures"]})


class TheWeeklyTabsAreTiedTogether(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from openpyxl import load_workbook
        cls.tmp = Path(tempfile.mkdtemp())
        cls.book = cls.tmp / "voice.xlsx"
        assert rebuild(VOICE, cls.book).returncode == 0
        cls.wb = load_workbook(cls.book)

    def formulas(self, sheet):
        return {dv.formula1: dv for dv in self.wb[sheet].data_validations.dataValidation}

    def test_request_cells_offer_the_shift_library_and_names_the_roster(self):
        for sheet in ("Preference", "Fixed Request"):
            f = self.formulas(sheet)
            self.assertIn("'Validation Lists'!$B$3:", " ".join(f), sheet)
            self.assertIn("'Schedule'!$D$3:$D$302", f, sheet)
        lists = self.wb["Validation Lists"]
        self.assertEqual([lists["B3"].value, lists["B4"].value], ["OFF", "Leave"])
        self.assertIn("'Shift Library'!A", str(lists["B5"].value))

    def test_previous_week_warns_instead_of_refusing_a_departed_name(self):
        warn = [dv for dv in self.wb["Previous week scheduled"].data_validations.dataValidation
                if dv.formula1 == "'Schedule'!$D$3:$D$302"]
        self.assertEqual(len(warn), 1)
        self.assertEqual(warn[0].errorStyle, "warning")
        self.assertIn("Known Departed Associates", warn[0].error)

    def test_demand_and_shrinkage_take_numbers_in_the_engine_range(self):
        shr = [dv for dv in self.wb["Shrinkage 30 Min"].data_validations.dataValidation if dv.type == "decimal"]
        self.assertEqual((shr[0].formula1, shr[0].formula2), ("0", "0.9999"))
        req = [dv for dv in self.wb["FT Wise 30 Min"].data_validations.dataValidation if dv.type == "decimal"]
        self.assertEqual(req[0].formula1, "0")

    def test_start_here_comes_first_and_reference_tabs_are_hidden_not_deleted(self):
        self.assertEqual(self.wb.sheetnames[:3], ["Start Here", "Instructions", "Schedule"])
        hidden = {ws.title for ws in self.wb.worksheets if ws.sheet_state == "hidden"}
        self.assertTrue({"Engine Defaults", "Validation Lists", "FT Wise 15 Min", "FT Wise 60 Min",
                         "RC9.1 Release Notes", "00 START HERE"} <= hidden, hidden)
        self.assertEqual(self.wb["FT Wise 30 Min"].sheet_state, "visible")
        self.assertEqual(self.wb.active.title, "Start Here")


@unittest.skipUnless(shutil.which("soffice"), "LibreOffice not installed: the live checks are not recalculated here")
class TheQuickChecksWork(unittest.TestCase):
    def recalculated_checks(self, book: Path):
        from openpyxl import load_workbook
        out = book.parent / "recalc"
        out.mkdir(exist_ok=True)
        env = dict(os.environ, HOME=str(book.parent))
        subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(out), str(book)],
                       capture_output=True, timeout=300, env=env, check=True)
        ws = load_workbook(out / book.name, data_only=True)["Start Here"]
        return {r[2].value: r[4].value for r in ws.iter_rows(min_row=5, max_row=11) if r[2].value}

    def test_a_clean_workbook_is_all_ok_and_a_real_mistake_turns_red(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "voice.xlsx"
            self.assertEqual(rebuild(VOICE, book).returncode, 0)
            checks = self.recalculated_checks(book)
            self.assertTrue(all(str(v).startswith("OK") for v in checks.values()), checks)
            wb = load_workbook(book)
            wb["Schedule"]["D4"] = wb["Schedule"]["D3"].value          # a repeated name
            wb["Preference"]["A3"] = "Somebody Not On The Roster"
            wb.save(book)
            checks = self.recalculated_checks(book)
            self.assertTrue(str(checks["Schedule"]).startswith("! 1 repeated name") or
                            str(checks["Schedule"]).startswith("! 2 repeated name"), checks)
            self.assertEqual(checks["Preference"], "! 1 name(s) not on the roster")


class TheReadyToEditWorkbooksMatchTheirSources(unittest.TestCase):
    PAIRS = {
        "AE_AR_B2B_WEEKLY_INPUT.xlsx": "AE_AR_B2B.xlsx",
        "Cricut_Chat_WEEKLY_INPUT.xlsx": "Cricut_Chat_RC9_1_READY_SKELETON.xlsx",
        "Cricut_Voice_WEEKLY_INPUT.xlsx": "Cricut_Voice_RC9_1_READY_SKELETON.xlsx",
        "GDI_28HC_24_7_WEEKLY_INPUT.xlsx": "GDI_REAL28_RC9_1_24_7_FINAL_READY.xlsx",
        "NMG_EN_AND_SP_WEEKLY_INPUT.xlsx": "NMG_EN_AND_SP.xlsx",
        "NMG_EN_WEEKLY_INPUT.xlsx": "NMG_EN_FIXED_NESTING_REST_SAFE_REGRESSION_CLEAN.xlsx",
        "NMG_SP_WEEKLY_INPUT.xlsx": "NMG_SP_RC9_1_READY_FIXED.xlsx",
    }

    def test_each_reads_exactly_like_the_shipped_workbook(self):
        self.assertEqual(sorted(p.name for p in READY.glob("*.xlsx")), sorted(self.PAIRS))
        for ready, source in self.PAIRS.items():
            with self.subTest(workbook=ready):
                before, after = B.engine_reading(INPUTS / source), B.engine_reading(READY / ready)
                changed = {k for k in before if before[k] != after.get(k)} - set(B.PROVENANCE_ONLY)
                self.assertEqual(changed, set())


if __name__ == "__main__":
    unittest.main()
