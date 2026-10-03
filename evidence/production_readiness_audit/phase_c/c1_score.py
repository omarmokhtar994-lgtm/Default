"""Score the C1 scenario runs against C1_RULE.txt rules R2-R6.

    python3 c1_score.py SCENARIO_OUT_DIR OUT.json

SCENARIO_OUT_DIR is the solver_scenarios.py output directory (runs/<id>/...)."""
import json, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(sys.argv[1])
R2_CASES = ["S04_ONE_IMPOSSIBLE_INTERVAL", "S05_LANGUAGE_SHORTAGE", "S09_LONE_NIGHT_BREAKS",
            "S11_OVERNIGHT_DAY_WINDOW", "S13_24x7_UNDERSTAFFED"]
R6_CASES = ["S10_FIXED_IN_BLANK_HOURS", "S14_DUPLICATE_PREFERENCE_ROW"]
ALLOWED_VALIDATOR = {"ZERO_STAFF_ACTIVE", "LANGUAGE_MINIMUM", "OPENING_MINIMUM", "COVERAGE_SPLIT",
                     "HARD_FLOOR", "NEXT_SUNDAY_CARRY_OUT"}
# The clean-room checker reports coverage minimums as metrics, not violations, except
# Coverage Split; every other violation it can report is a person rule.
ALLOWED_CLEAN_ROOM = {"coverage_split"}


def jload(p):
    try:
        return json.loads(Path(p).read_text())
    except (OSError, json.JSONDecodeError):
        return None


def shortfall_sheet(path):
    wb = load_workbook(path, read_only=True)
    try:
        if "Shortfalls" not in wb.sheetnames:
            return None
        rows = [r for r in wb["Shortfalls"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    finally:
        wb.close()
    return rows


out = {}
for sid in R2_CASES + R6_CASES:
    res = jload(ROOT / f"{sid}.result.json") or {}
    case = ROOT / "runs" / sid
    books = sorted(case.glob("*_HARD_RULE_SHORTFALL_SCHEDULE.xlsx"))
    row = {"runner_rc": res.get("runner_rc"), "engine_status": res.get("engine_status"),
           "published_final": res.get("published_final"), "shortfall_workbook": books[0].name if books else None}
    outcome = (case / "BUSINESS_OUTCOME.txt").read_text() if (case / "BUSINESS_OUTCOME.txt").exists() else ""
    if sid in R2_CASES:
        r2 = bool(books) or (res.get("runner_rc") == 0 and bool(res.get("published_final")))
        row["R2"] = "PASS" if r2 else "FAIL"
        if books:
            check = jload(case / "SHORTFALL_SCHEDULE_VALIDATION.json") or {}
            status = jload(case / "UNIVERSAL_RUN_STATUS.json") or {}
            sv = ((status.get("independent_validation") or {}).get("shortfall_schedule")) or {}
            cr = jload(case / "SHORTFALL_CLEAN_ROOM_CHECK.json") or {}
            types = sorted({f.get("type") for f in check.get("failures", [])})
            rows = shortfall_sheet(books[0]) or []
            row.update(validator_failure_types=types, runner_shortfall_check=sv.get("status"),
                       count_mismatches=sv.get("count_mismatches"), shortfall_rows=len(rows),
                       shortfall_families=sorted({str(r[0]) for r in rows}),
                       clean_room_violations_by_rule=cr.get("violations_by_rule"),
                       clean_room_engine_mismatches=cr.get("engine_mismatches"))
            r3 = (set(types) <= ALLOWED_VALIDATOR and sv.get("status") == "SHORTFALLS_CONFIRMED"
                  and not sv.get("count_mismatches") and cr != {}
                  and set((cr.get("violations_by_rule") or {})) <= ALLOWED_CLEAN_ROOM)
            row["R3"] = "PASS" if r3 else "FAIL"
            row["R5"] = ("PASS" if res.get("runner_rc") not in (0, None)
                         and "No schedule meets every hard rule" in outcome
                         and "Shortfalls (the attached schedule misses these minimums)" in outcome
                         else "FAIL")
            if sid == "S04_ONE_IMPOSSIBLE_INTERVAL":
                times = sorted({(str(r[2]), str(r[3])) for r in rows})
                row["shortfall_day_times"] = times
                row["R4"] = "PASS" if rows and all(d == "Sun" and t.startswith("03:") for d, t in times) else "FAIL"
    else:
        row["pre_solver_failures"] = res.get("pre_solver_failures")
        row["R6"] = "PASS" if (not books and res.get("engine_status") == "FAIL_PRE_SOLVER_CONTRACT"
                               and res.get("runner_rc") not in (0, None)) else "FAIL"
    out[sid] = row
verdict = {k: all(r.get(k, "PASS") == "PASS" for r in out.values()) for k in ("R2", "R3", "R4", "R5", "R6")}
json.dump({"cases": out, "rules_pass": verdict}, open(sys.argv[2], "w"), indent=1, default=str)
print(json.dumps(verdict))
for k, v in out.items():
    print(k, {x: v.get(x) for x in ("runner_rc", "engine_status", "shortfall_workbook", "R2", "R3", "R4", "R5", "R6", "shortfall_families", "shortfall_rows")})
