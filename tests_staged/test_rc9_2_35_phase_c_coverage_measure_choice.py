"""Phase C4 follow-up: each program chooses its coverage measure.

Some programs are accountable for service level (calls answered in time, a
volume-weighted measure); others for interval compliance (the share of
intervals at target). One global default cannot fit both, so the choice is
made per program, in the same two places as the Language Working Window:

  * the workbook: an Instructions row "Coverage Objective Weighting" with an
    enforced dropdown (Interval Count / Volume Weighted). The template builder
    adds it, seeded with Interval Count, the value an empty cell already
    means, so rebuilding a workbook changes no contract;
  * the run: an override on the engine, the production runner, the Colab
    runner and both notebooks ("workbook" = use the row).

The run override beats the workbook, which beats the default (Interval Count).
The measure used, and where it came from, is in the audit and on the first
lines of BUSINESS_OUTCOME.txt.

Written to FAIL on the engine before this change and pass after.
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(ROOT / "engine" / "_tools"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import l632_universal_scheduler as E  # noqa: E402
import phase_a_workbooks as W  # noqa: E402

PACKAGE = REPO / "packages" / "rc9_2_2_production"
TEMPLATE = REPO / "tools" / "build_input_template.py"

ONE_DAY = dict(
    roster=[("Agent A", "English")], shifts=["06:00 - 15:00", "15:00 - 00:00"],
    demand=lambda d, m: (0.5 if 6 * 60 <= m < 15 * 60 else 0.85 if 15 * 60 <= m < 21 * 60 else None) if d == 1 else None,
    instructions={"Count of Associates": 1, "Rest Gap Hours": 8},
)


def build(rows=None):
    spec = dict(ONE_DAY)
    spec["instructions"] = dict(ONE_DAY["instructions"], **(rows or {}))
    path = Path(tempfile.mkdtemp()) / "case.xlsx"
    W.build(spec, path)
    return path


class TheWorkbookOffersADropdown(unittest.TestCase):
    def test_the_template_lists_the_row_with_its_two_choices(self):
        source = TEMPLATE.read_text()
        self.assertIn('("Coverage Objective Weighting", \'"Interval Count,Volume Weighted"\')', source)

    def test_a_rebuilt_workbook_has_the_row_seeded_and_enforced(self):
        from openpyxl import load_workbook
        src = build()
        out = src.parent / "rebuilt.xlsx"
        subprocess.run([sys.executable, str(TEMPLATE), str(src), str(out)], check=True, capture_output=True)
        wb = load_workbook(out)
        ws = wb["Instructions"]
        cell = None
        for row in ws.iter_rows():
            for c in row:
                if str(c.value or "").strip() == "Coverage Objective Weighting":
                    cell = ws.cell(c.row, c.column + 1)
        self.assertIsNotNone(cell, "row missing from the rebuilt Instructions sheet")
        self.assertEqual(cell.value, "Interval Count")
        lists = [dv for dv in ws.data_validations.dataValidation if cell.coordinate in dv.sqref]
        self.assertEqual(len(lists), 1)
        self.assertEqual(lists[0].formula1, '"Interval Count,Volume Weighted"')
        self.assertTrue(lists[0].showErrorMessage)
        # Seeding the value an empty cell already means changes no contract.
        self.assertEqual(E.canonical_hash(E.canonical_contract_snapshot(E.parse_input(src))),
                         E.canonical_hash(E.canonical_contract_snapshot(E.parse_input(out))))


class TheRunCanOverrideTheWorkbook(unittest.TestCase):
    def test_the_engine_accepts_the_override(self):
        args = E.build_arg_parser().parse_args(["--coverage-objective-weighting", "VOLUME_WEIGHTED"])
        self.assertEqual(args.coverage_objective_weighting, "VOLUME_WEIGHTED")
        self.assertIsNone(E.build_arg_parser().parse_args([]).coverage_objective_weighting)

    def test_override_beats_workbook_beats_default(self):
        default = E.parse_input(build())
        stated = E.parse_input(build({"Coverage Objective Weighting": "Volume Weighted"}))
        self.assertEqual(E.apply_coverage_objective_override(default, None),
                         {"mode": "interval_count", "source": "default"})
        self.assertEqual(E.apply_coverage_objective_override(stated, None),
                         {"mode": "volume_weighted", "source": "workbook"})
        self.assertEqual(E.apply_coverage_objective_override(stated, "INTERVAL_COUNT"),
                         {"mode": "interval_count", "source": "run override"})
        self.assertEqual(stated.coverage_objective_weighting, "interval_count")
        self.assertEqual(E.volume_weight(stated, 700, 1, 18), 700)

    def test_the_override_is_in_the_contract_fingerprint(self):
        a = E.parse_input(build())
        b = E.parse_input(build())
        E.apply_coverage_objective_override(b, "VOLUME_WEIGHTED")
        self.assertNotEqual(E.canonical_hash(E.canonical_contract_snapshot(a)),
                            E.canonical_hash(E.canonical_contract_snapshot(b)))

    def test_an_unknown_override_is_refused(self):
        with self.assertRaises(SystemExit):
            E.build_arg_parser().parse_args(["--coverage-objective-weighting", "REVENUE"])

    def test_run_case_applies_and_records_it(self):
        source = (ROOT / "engine" / "_tools" / "l632_universal_scheduler.py").read_text()
        body = source[source.index("def run_case("):source.index("def run_case(") + 60000]
        self.assertIn("apply_coverage_objective_override(parsed, coverage_objective_weighting_override)", body)
        self.assertIn('"coverage_objective_weighting_override": coverage_objective_weighting_override', body)
        self.assertIn("coverage_objective_weighting_override=args.coverage_objective_weighting", source)


class EveryRunnerPassesItOn(unittest.TestCase):
    def test_the_production_runner(self):
        source = (ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py").read_text()
        self.assertIn("'--coverage-objective-weighting'", source)
        self.assertIn("command += ['--coverage-objective-weighting', args.coverage_objective_weighting]", source)

    def test_the_colab_runner(self):
        source = (PACKAGE / "runners" / "rc921_runner.py").read_text()
        self.assertIn('"--coverage-objective-weighting"', source)
        self.assertIn('command += ["--coverage-objective-weighting", args.coverage_objective_weighting]', source)

    def test_both_notebooks_offer_the_choice(self):
        for name in ("RC922_Colab_A_NO_DRIVE.ipynb", "RC922_Colab_B_WITH_DRIVE.ipynb"):
            with self.subTest(notebook=name):
                cells = json.loads((PACKAGE / "runners" / name).read_text())["cells"]
                run = next("".join(c["source"]) for c in cells if "#@title 5. Run" in "".join(c["source"]))
                self.assertIn('COVERAGE_MEASURE = "workbook" #@param ["workbook", "INTERVAL_COUNT", "VOLUME_WEIGHTED"]', run)
                self.assertIn('cmd += ["--coverage-objective-weighting", COVERAGE_MEASURE]', run)


class TheOutcomeSaysWhichMeasureWasUsed(unittest.TestCase):
    def test_the_first_lines_name_the_measure_and_its_source(self):
        audit = {"status": "PASS", "coverage_measure": {"mode": "volume_weighted", "source": "run override"}}
        outcome = E.build_business_outcome(audit, 0)
        self.assertEqual(outcome["coverage_measure"], {"mode": "volume_weighted", "source": "run override"})
        text = E.format_business_outcome(outcome)
        self.assertIn("Coverage measure: Volume Weighted (run override)", "\n".join(text.splitlines()[:6]))

    def test_the_runner_header_names_it_too(self):
        # BUSINESS_OUTCOME.txt, the file a planner reads, is written by the runner.
        sys.path.insert(0, str(ROOT / "engine"))
        import importlib.util
        spec = importlib.util.spec_from_file_location("runner_c4", ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py")
        runner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(runner)
        self.assertEqual(runner.coverage_measure_line({"coverage_measure": {"mode": "interval_count", "source": "workbook"}}),
                         "Coverage measure: Interval Count (workbook)\n")
        self.assertEqual(runner.coverage_measure_line({}), "")
        source = (ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py").read_text()
        header = source[source.index("rendered_text = ("):source.index("rendered_text = (") + 500]
        self.assertIn("coverage_measure_line(outcome)", header)


if __name__ == "__main__":
    unittest.main()
