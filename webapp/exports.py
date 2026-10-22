# © 2026 Omar Mokhtar. All rights reserved.
"""Exports: what was recorded on the website for a period, and who did it.

Shift-dated items (attendance, day activity, breaks, worked hours) follow the
shift's own date; logged actions (schedule changes, versions, runs, the other
actions) follow when they happened. All times are Egypt time. Excel holds an
"About this export" tab and a tab per item; CSV holds one item. Text typed by
people is always written as text (never a formula, in Excel or CSV).
"""
from __future__ import annotations

import csv
import io
import json
import time
from datetime import date, datetime, timedelta
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from .adherence import person_day
from .clock import EGYPT
from .day import ABSENT, AUX, CALLED_IN, EXTRA_BREAKS, MEASURES, pattern_breaks, planned
from .versions import DAYS, shift_span

MAX_DAYS = 400
KINDS = {"attendance": "Attendance", "activities": "Activities", "activity": "Day activity",
         "breaks": "Breaks planned vs taken",
         "changes": "Schedule changes and swaps", "versions": "Versions and in use", "runs": "Runs",
         "worked": "Worked hours and adherence", "summary": "Daily summary",
         "channels": "Channels planned and changed", "cover": "Cover from outside the roster",
         "leave": "Leave requests", "channel_hours": "Hours per channel",
         "absence": "Absenteeism and shrinkage", "incident_summary": "Incidents per associate",
         "incidents": "Incidents (missed breaks)", "sessions": "Booked sessions",
         "break_requests": "Break and coaching requests", "swaps": "Swap requests",
         "record": "Other actions"}
VERSION_EVENTS = {"version_created": "New version", "set_in_use": "Set in use"}
OTHER_EVENTS = {"run_uploaded": "Run uploaded", "run_started": "Run started", "run_stopped": "Run stopped",
                "run_resumed": "Run resumed", "roles_changed": "Roles changed", "team_manager_set": "Team manager set", "downloaded": "Downloaded", "exported": "Exported",
                "user_added": "Person added", "user_deleted": "Person deleted", "users_bulk_added": "People added from a file", "signed_in": "Signed in", "sign_in_failed": "Sign-in failed",
                "sign_in_held": "Sign-in held (too many tries)", "sign_in_released": "Sign-in hold lifted", "sessions_ended": "Other devices signed out", "users_exported": "People list downloaded", "user_disabled": "Person switched off", "user_enabled": "Person switched on",
                "password_reset": "Password reset", "run_details_changed": "Run details changed",
                "program_renamed": "Program renamed", "user_programs_changed": "Programs or role changed",
                "program_set_up": "Programs and LOBs", "program_deleted": "Program deleted",
                "lob_deleted": "LOB deleted", "with_list_changed": "Departments and people",
                "channels_changed": "Associate channels", "interval_target_changed": "Interval target",
                "channel_needs_added": "Channel needs added", "run_deleted": "Schedule deleted",
                "notify_saved": "Notifications changed", "notify_removed": "Group link removed",
                "notify_tested": "Test message sent", "shrinkage_mapped": "What counts as shrinkage",
                "operating_day_set": "Operating day start", "upload_refused": "Upload refused and deleted", "heavy_weeks_set": "Longest look changed"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RISKY = ("=", "+", "-", "@", "\t", "\r")


EXPORT_SECONDS = 300  # an export still working after this stops and says why (owner: it "keeps loading")


class ExportTooSlow(ValueError):
    """An export that ran past its time (said on the page with what to untick)."""


SLOW_SAID = (f"The export did not finish within {EXPORT_SECONDS // 60} minutes, so it was stopped. Pick a shorter "
             f"period, or untick {KINDS['worked']} and {KINDS['summary']} (the items that work out every day), then "
             "download again.")


class _Days:
    """The day book as one export sees it: each day worked out once (worked hours and the daily summary ask for the
    same days), and a check that the export is still within its time."""

    def __init__(self, days, deadline: Optional[float], cover: bool = True):
        self._days = days
        self._deadline = deadline
        self.cover = cover  # Phase AJ: people covering from outside the roster counted in the day's figures
        self._pages: Dict[Tuple[str, date, str], Any] = {}

    def __getattr__(self, name: str) -> Any:
        return getattr(self._days, name)

    def page(self, program: str, on: date, measure: Optional[str] = None, extra: Optional[Dict[str, Any]] = None):
        if extra:
            return self._days.page(program, on, measure, extra, cover=self.cover)
        key = (program, on, measure)
        if key not in self._pages:
            try:
                self._pages[key] = (self._days.page(program, on, measure, cover=self.cover), None)
            except ValueError as exc:  # said again to every item that asks for this day
                self._pages[key] = (None, exc)
        found, problem = self._pages[key]
        if problem is not None:
            raise problem
        return found

    def check(self) -> None:
        if self._deadline is not None and time.monotonic() > self._deadline:
            raise ExportTooSlow(SLOW_SAID)


def _ts(day: date) -> float:
    return datetime(day.year, day.month, day.day, tzinfo=EGYPT).timestamp()


def _when(ts: Optional[float]) -> str:
    return datetime.fromtimestamp(ts, EGYPT).strftime("%Y-%m-%d %H:%M") if ts else ""


def _hm(minute: Optional[int]) -> str:
    if minute is None:
        return ""
    minute %= 1440
    return f"{minute // 60:02d}:{minute % 60:02d}"


def _dates(start: date, end: date) -> Iterable[date]:
    day = start
    while day <= end:
        yield day
        day += timedelta(days=1)


def _day_index(day: date) -> int:
    return (day.weekday() + 1) % 7


class _Weeks:
    """The schedule in use (or the tool's) per program and week, read once."""

    def __init__(self, days):
        self.days = days
        self.seen: Dict[int, Dict[str, Any]] = {}
        self.dates: Dict[Tuple[str, date], Dict[str, Any]] = {}

    def week(self, program: str, day: date) -> Dict[str, Any]:
        """The week holding ``day`` (its start may be any weekday): looked up once per program and date,
        read once per version."""
        if (program, day) not in self.dates:
            row, _ = self.days.version(program, day)
            if row is not None and row["id"] not in self.seen:
                self.seen[row["id"]] = json.loads(row["week"] or "{}")
            self.dates[(program, day)] = self.seen[row["id"]] if row is not None else {}
        return self.dates[(program, day)]

    def person(self, program: str, day: date, name: str) -> Dict[str, Any]:
        return next((a for a in self.week(program, day).get("associates", []) if a["name"] == name), {})


def _minutes_off(status: str, shift: str, lo: Optional[int], hi: Optional[int]) -> Any:
    """Phase AV: how long a Late (shift start to arrival), a Left early (leaving to shift end) or a timed absence was."""
    span = shift_span(shift)
    if status == "Late" and span and hi is not None:
        return hi + (1440 if hi < span[0] else 0) - span[0]
    if status == "Left early" and span and lo is not None:
        return span[1] - (lo + (1440 if lo < span[0] else 0))
    if lo is not None and hi is not None:
        return hi - lo
    return ""


def _attendance(store, days, weeks, start, end, program, user_id, measure):
    yield ["Shift date", "Program", "Associate", "Slot", "Shift", "Status", "From", "To", "Minutes", "Billable",
           "With department", "With", "Why", "Recorded by", "Recorded at"]
    for r in store.attendance_between(start.isoformat(), end.isoformat(), program, user_id):
        day = date.fromisoformat(r["shift_date"])
        person = weeks.person(r["program"], day, r["associate"])
        billable = "Yes" if r["billable"] else ("No" if r["status"] in AUX else "")
        shift = (person.get("days") or [""] * 7)[_day_index(day)]
        yield [r["shift_date"], r["program"], r["associate"], person.get("slot", ""), shift, r["status"],
               _hm(r["from_min"]), _hm(r["to_min"]), _minutes_off(r["status"], shift, r["from_min"], r["to_min"]),
               billable, r.get("with_dept", ""), r.get("with_whom", ""), r.get("why", ""), r["by_name"], _when(r["at"])]


def _activities(store, days, weeks, start, end, program, user_id, measure):
    yield ["Shift date", "Program", "Associate", "Activity", "From", "To", "Minutes", "Billable", "With department",
           "With", "Why", "Note", "Recorded by", "Recorded at"]
    for r in store.activities_between(start.isoformat(), end.isoformat(), program, user_id):
        yield [r["shift_date"], r["program"], r["associate"], r["kind"], _hm(r["start"]), _hm(r["end_min"]),
               r["end_min"] - r["start"], "Yes" if r["billable"] else "No", r.get("with_dept", ""),
               r.get("with_whom", ""), r.get("why", ""), r["note"], r["by_name"], _when(r["at"])]


def _activity(store, days, weeks, start, end, program, user_id, measure):
    yield ["Shift date", "Program", "Associate", "What", "By", "At"]
    for r in store.day_log_between(start.isoformat(), end.isoformat(), program, user_id):
        yield [r["shift_date"], r["program"], r["associate"], r["what"], r["by_name"], _when(r["at"])]


def _programs_said(program) -> str:
    """The About tab's line: one program, a person's own, none, or all."""
    if program is None:
        return "All programs"
    if isinstance(program, str):
        return program
    return ", ".join(sorted(program)) or "None (no programs assigned)"


def _programs(days, program):
    """One program, several (a person's own), or every program with schedules."""
    if isinstance(program, (list, tuple, set, frozenset)):
        return sorted(program)
    return [program] if program else days.programs()


def _breaks(store, days, weeks, start, end, program, user_id, measure):
    yield ["Shift date", "Program", "Associate", "Slot", "Shift", "Break", "Minutes", "Planned", "Taken", "Moved by",
           "Moved at", "Status", "Cancelled because"]  # Phase AD: a break the RTA cancelled, and why
    for p in _programs(days, program):
        moved = {(r["shift_date"], r["associate"], r["idx"]): r for r in store.actual_breaks_between(
            start.isoformat(), end.isoformat(), p)}
        status = {(r["shift_date"], r["associate"]): r["status"] for r in store.attendance_between(
            start.isoformat(), end.isoformat(), p)}
        acts = store.activities_between(start.isoformat(), end.isoformat(), p)
        called = {(r["shift_date"], r["associate"]): r["note"] for r in acts if r["kind"] == CALLED_IN}
        added: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}  # breaks and lunches added on the day in RTA
        for r in acts:
            if r["kind"] in EXTRA_BREAKS:
                added.setdefault((r["shift_date"], r["associate"]), []).append(r)
        for day in _dates(start, end):
            days.check()
            week = weeks.week(p, day)
            for a in week.get("associates", []):
                state = status.get((day.isoformat(), a["name"]), "Present")
                shift, plan = a["days"][_day_index(day)], planned(week, _day_index(day), a["name"])
                label = called.get((day.isoformat(), a["name"]))
                if not plan and label:  # a day off cancelled: the plan's breaks for the shift they came in for
                    shift, plan = f"{label} (called in)", pattern_breaks(week, _day_index(day), label)
                for b in plan:
                    m = moved.get((day.isoformat(), a["name"], b["idx"]))
                    m = m if m and m["kind"] == b["kind"] else None  # a move kept against another plan is not this break
                    gone = bool(m and m.get("cancelled"))
                    taken = "" if state in ABSENT else ("Cancelled" if gone else _hm(m["start"] if m else b["start"]))
                    yield [day.isoformat(), p, a["name"], a.get("slot", ""), shift, b["kind"],
                           b["minutes"], _hm(b["start"]), taken, m["by_name"] if m else "", _when(m["at"]) if m else "",
                           state, m["why"] if gone else ""]
                for r in added.get((day.isoformat(), a["name"]), []):
                    yield [day.isoformat(), p, a["name"], a.get("slot", ""), shift, r["kind"], r["end_min"] - r["start"],
                           "", _hm(r["start"]), r["by_name"], _when(r["at"]), f"{state}, added on the day", ""]


def _changes(store, days, weeks, start, end, program, user_id, measure):
    yield ["Program", "Week", "Version", "Associate", "Day", "Old", "New", "Reason", "Severity", "Problems", "By", "At"]
    for r in store.changes_between(_ts(start), _ts(end + timedelta(days=1)), program, user_id):
        try:
            found = "; ".join(json.loads(r["problems"] or "[]"))
        except ValueError:
            found = r["problems"]
        severity = {"red": "Rule broken", "yellow": "Warning"}.get(r["severity"], "")
        yield [r["program"], r["week_start"], r["label"], r["associate"],
               "Whole week (slot swap)" if r["day"] == "Week" else r["day"], r["old"], r["new"], r["reason"], severity,
               found, r["by_name"], _when(r["at"])]


def _versions(store, days, weeks, start, end, program, user_id, measure):
    yield ["Program", "Week", "Version", "What", "Detail", "By", "At"]
    for r in store.list_events(_ts(start), _ts(end + timedelta(days=1)), program, user_id, list(VERSION_EVENTS)):
        yield [r["program"], r["week_start"], r["subject"], VERSION_EVENTS[r["kind"]], r["detail"], r["by_name"],
               _when(r["at"])]


def _runs(store, days, weeks, start, end, program, user_id, measure):
    yield ["Workbook", "Program", "Week", "Mode", "Status", "Uploaded by", "Uploaded at", "Finished at", "Run"]
    for r in store.runs_between(_ts(start), _ts(end + timedelta(days=1)), program, user_id):
        yield [r["workbook"], r["program"], r["week_start"], r["mode"], r["status"], r["by_name"], _when(r["created"]),
               _when(r["finished"]), r["id"]]


def _record(store, days, weeks, start, end, program, user_id, measure):
    yield ["What", "About", "Program", "Detail", "By", "At"]
    for r in store.list_events(_ts(start), _ts(end + timedelta(days=1)), program, user_id, list(OTHER_EVENTS)):
        yield [OTHER_EVENTS[r["kind"]], r["subject"], r["program"], r["detail"], r["by_name"], _when(r["at"])]


def _worked(store, days, weeks, start, end, program, user_id, measure):
    yield ["Shift date", "Program", "Associate", "Slot", "Shift", "Status", "Scheduled min", "Scheduled work min",
           "Worked min", "Late min", "Early min", "Absent min", "Aux billable min", "Aux non-billable min",
           "Breaks planned min", "Breaks taken min", "Out of schedule min", "VTO min", "Overtime min", "Adherence %",
           "Conformance %"]
    for p in _programs(days, program):
        for day in _dates(start, end):
            days.check()
            try:
                page = days.page(p, day, measure)
            except ValueError:  # the input workbook is gone: said on the day page, skipped here
                continue
            if page is None:
                continue
            for lane in page["view"]["lanes"]:
                if not any(s["offset"] == 0 for s in lane["segments"]):
                    continue
                r = person_day(page["view"], lane["name"])
                yield [day.isoformat(), p, lane["name"], lane.get("slot", ""), r["shift"], r["status"], r["scheduled"],
                       r["scheduled_work"], r["worked"], r["late"], r["early"], r["absent"], r["aux_billable"],
                       r["aux_unbillable"], r["breaks_planned"], r["breaks_taken"], r["out_of_schedule"], r["vto"],
                       r["overtime"], r["adherence"], r["conformance"]]


def _summary(store, days, weeks, start, end, program, user_id, measure):
    from .handover import note  # the handover note's numbers, one row per program and day
    yield ["Date", "Program", "Planned", "Present", "Absent", "Late or early", "In aux", "Breaks moved", "Overtime min",
           "VTO min", "Called in min", "Hours short", "Hours short in the plan", "Hours above", "Tightest (hours)",
           "Language gaps", "Adherence %", "Conformance %", "Changes recorded", "Interval target %",
           "Intervals at target", "Intervals with demand", "Cover min"]
    for p in _programs(days, program):
        for day in _dates(start, end):
            days.check()
            try:
                found = note(days, p, day, measure)
            except ValueError:
                continue
            if found is None:
                continue
            n = found["numbers"]
            yield [day.isoformat(), p, n["planned"], n["present"], n["absent"], n["late_early"], n["aux"], n["moved"],
                   n["overtime_minutes"], n["vto_minutes"], n["called_in_minutes"], n["short_hours"],
                   n["plan_short_hours"], n["over_hours"], n["tightest"], n["language_gaps"], n["adherence"],
                   n["conformance"], n["changes"], n["target"], n["at_target"], n["intervals"],
                   days.page(p, day, measure)["view"]["tiles"]["cover"] if days.cover else ""]


def _channels(store, days, weeks, start, end, program, user_id, measure):
    """Phase V: each planned channel block of the schedule in use, and each channel change made on the day."""
    from .channel_day import letter_at
    from .channels import CHANNELS
    from .versions import channel_blocks
    words = {**CHANNELS, "A": "All channels"}
    yield ["Shift date", "Program", "Associate", "Slot", "Shift", "Channel", "Start", "End", "Minutes", "Source",
           "Was", "By", "At"]
    for p in _programs(days, program):
        changes: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
        for r in store.channel_moves_between(start.isoformat(), end.isoformat(), p, user_id):
            changes.setdefault((r["shift_date"], r["associate"]), []).append(r)
        for day in _dates(start, end):
            days.check()
            week = weeks.week(p, day)
            for a in week.get("associates", []):
                blocks = channel_blocks(week, _day_index(day), a["name"])
                shift = a["days"][_day_index(day)]
                for b in blocks:
                    yield [day.isoformat(), p, a["name"], a.get("slot", ""), shift, words[b["channel"]],
                           _hm(b["start"]), _hm(b["end"]), b["end"] - b["start"], "Plan", "", "", ""]
                for r in changes.get((day.isoformat(), a["name"]), []):
                    was = letter_at([(b["start"], b["end"], b["channel"]) for b in blocks], [], r["start"], False)
                    yield [day.isoformat(), p, a["name"], a.get("slot", ""), shift, words[r["channel"]],
                           _hm(r["start"]), _hm(r["end_min"]), r["end_min"] - r["start"], "Changed on the day",
                           words.get(was, ""), r["by_name"], _when(r["at"])]


def _cover(store, days, weeks, start, end, program, user_id, measure):
    """Phase AJ: team managers and support members, not on the roster, who logged in to cover intervals."""
    from .channels import CHANNELS
    yield ["Shift date", "Program", "Name", "Role", "From", "To", "Minutes", "Channel", "Language", "Recorded by",
           "Recorded at"]
    for r in store.cover_between(start.isoformat(), end.isoformat(), program, user_id):
        yield [r["shift_date"], r["program"], r["associate"], r["role"], _hm(r["start"]), _hm(r["end_min"]),
               r["end_min"] - r["start"], CHANNELS.get(r["channel"], ""), r.get("language", ""), r["by_name"],
               _when(r["at"])]


def _leave(store, days, weeks, start, end, program, user_id, measure):
    """Phase AO: leave requests touching the period, whatever their status, so approved leave for a week that had no
    schedule yet can be taken from here for that week later."""
    yield ["Program", "Name", "Kind", "First day", "Last day", "Status", "Note", "Team manager", "RTA", "Outcome",
           "Reason"]
    words = {"waiting_tm": "Waiting for the team manager", "waiting_rta": "Waiting for the RTA",
             "applying": "Being approved", "approved": "Approved", "refused": "Refused", "cancelled": "Cancelled"}
    for r in store.leave_between(start.isoformat(), end.isoformat(), program):
        yield [r["program"], r["associate"], r["kind"], r["first_day"], r["last_day"], words.get(r["status"], r["status"]),
               r["note"], r["tm_name"], r["rta_name"], r["outcome"], r["reason"]]


def _units_of(days, program):
    from .programs import ProgramBook
    book = ProgramBook(days.store)
    out = {}
    for p in _programs(days, program):
        found = book.unit(p)
        out[p] = ((found["program"]["name"], (found["lob"] or {}).get("name") or found["program"]["name"])
                  if found else (p, p))
    return out


def _absence(store, days, weeks, start, end, program, user_id, measure):
    """Phase AV: absenteeism (the RTA's and client's view, and Sutherland's) and the three shrinkage figures, per LOB and day."""
    from . import insights
    yield ["Date", "Program", "LOB", "Shifts scheduled", "Unplanned leave", "Sick", "Planned leave",
           "Absenteeism % (RTA and client: unplanned + sick)", "Absenteeism % (Sutherland: unplanned)",
           "Shrinkage out of office %", "Shrinkage, the client's %", "Shrinkage in office %", "Adherence %", "Late", "Early",
           "Missed breaks"]
    facts = insights.collect(days, store, None, _units_of(days, program), start, end, measure, days.check)
    by: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
    for r in facts:
        by.setdefault((r["date"], r["program"], r["lob"]), []).append(r)
    for (day, p, lob), rows in sorted(by.items()):
        f, i = insights.figures(rows), insights.figures(rows, ("unplanned",))
        n = f["incidents"]
        yield [day, p, lob, f["shifts"], n["unplanned"]["n"], n["sick"]["n"], n["planned"]["n"], f["absenteeism"],
               i["absenteeism"], f["shrink_out"], f["shrink_client"], f["shrink_in"], f["adherence"], n["late"]["n"],
               n["early"]["n"], n["missed_break"]["n"]]


def _incident_summary(store, days, weeks, start, end, program, user_id, measure):
    """Phase AV: each associate's count and minutes for every incident type over the period."""
    from . import insights
    head = ["Associate", "Program", "LOB", "Shifts"]
    for _, label in insights.KINDS:
        head += [f"{label} (count)", f"{label} (min)"]
    yield head + ["Adherence %"]
    facts = insights.collect(days, store, None, _units_of(days, program), start, end, measure, days.check)
    for g in insights.grouped(facts, "associates"):
        row = [g["name"], g["program"], g["lob"], g["shifts"]]
        for k, _ in insights.KINDS:
            row += [g["incidents"][k]["n"], g["incidents"][k]["min"]]
        yield row + [g["adherence"]]


def _incidents(store, days, weeks, start, end, program, user_id, measure):
    """Phase AV: what was ticked as an incident when a break was moved (a missed break so far)."""
    yield ["Shift date", "Program", "Associate", "Slot", "Shift", "Incident", "Break", "Note", "Recorded by",
           "Recorded at"]
    for r in store.incidents_between(start.isoformat(), end.isoformat(), program, user_id):
        person = weeks.person(r["program"], date.fromisoformat(r["shift_date"]), r["associate"])
        shift = (person.get("days") or [""] * 7)[_day_index(date.fromisoformat(r["shift_date"]))]
        yield [r["shift_date"], r["program"], r["associate"], person.get("slot", ""), shift,
               "Missed break" if r["kind"] == "missed_break" else r["kind"], r["idx"] + 1, r["note"], r["by_name"],
               _when(r["at"])]


def _sessions(store, days, weeks, start, end, program, user_id, measure):
    """Phase AV: meetings, training and coaching booked for people together, one row per person."""
    yield ["Shift date", "Program", "Session", "Associate", "Activity", "From", "To", "Minutes", "Billable", "With",
           "Why", "People in it", "Recorded by", "Recorded at"]
    rows = [r for r in store.activities_between(start.isoformat(), end.isoformat(), program, user_id) if r.get("session")]
    size: Dict[Tuple[str, str], int] = {}
    for r in rows:
        size[(r["shift_date"], r["session"])] = size.get((r["shift_date"], r["session"]), 0) + 1
    for r in sorted(rows, key=lambda r: (r["shift_date"], r["program"], r["session"], r["start"], r["associate"])):
        yield [r["shift_date"], r["program"], r["session"], r["associate"], r["kind"], _hm(r["start"]), _hm(r["end_min"]),
               r["end_min"] - r["start"], "Yes" if r["billable"] else "No", r.get("with_whom", ""), r.get("why", ""),
               size[(r["shift_date"], r["session"])], r["by_name"], _when(r["at"])]


def _break_requests(store, days, weeks, start, end, program, user_id, measure):
    """Phase AV: break and coaching change requests, and how each was answered."""
    yield ["Shift date", "Program", "Kind", "People", "Asked by", "Asked at", "Detail", "Why", "Status", "Decided by",
           "Decided at", "Decision", "Outcome"]
    for r in store.requests_between(start.isoformat(), end.isoformat(), program):
        try:
            people = ", ".join(json.loads(r["names"] or "[]"))
        except ValueError:
            people = r["names"]
        yield [r["shift_date"], r["program"], r["kind"], people, r["requester_name"], _when(r["asked"]), r["detail"],
               r["why"], r["status"], r["decided_name"], _when(r["decided_at"]), r["decision"], r["outcome"]]


def _swaps(store, days, weeks, start, end, program, user_id, measure):
    """Phase AV: shift swap requests for days in the period, whatever their status."""
    yield ["Day", "Program", "Asked by", "Swap with", "Their shift", "Other's shift", "Note", "Status", "Asked at",
           "Decided by", "Decided at", "Reason", "Outcome"]
    for r in store.swaps_between(start.isoformat(), end.isoformat(), program):
        yield [r["day"], r["program"], r["asker"], r["other"], r["asker_cell"], r["other_cell"], r["note"], r["status"],
               _when(r["asked"]), r["decided_name"], _when(r["decided_at"]), r["reason"], r["outcome"]]


def _channel_hours(store, days, weeks, start, end, program, user_id, measure):
    """Phase AP (sample ap01, C): the hours each channel is staffed, per program and day, from the schedule in use."""
    from .channel_hours import CHANNEL_NAMES, day_hours
    yield ["Date", "Program", "Channel", "Hours"]
    for p in _programs(days, program):
        for day in _dates(start, end):
            days.check()
            try:
                page = days.page(p, day, measure)
            except ValueError:
                continue
            for letter, hours in sorted(day_hours(page["view"] if page else None).items()):
                yield [day.isoformat(), p, CHANNEL_NAMES.get(letter, letter), hours]


TABLES: Dict[str, Callable] = {"attendance": _attendance, "activities": _activities, "activity": _activity,
                               "breaks": _breaks,
                               "changes": _changes, "versions": _versions, "runs": _runs, "worked": _worked,
                               "summary": _summary, "channels": _channels, "cover": _cover, "leave": _leave, "channel_hours": _channel_hours,
                               "absence": _absence, "incident_summary": _incident_summary, "incidents": _incidents, "sessions": _sessions, "break_requests": _break_requests,
                               "swaps": _swaps, "record": _record}


def _text_cell(ws, value: Any) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # typed text stays text
    return cell


def _csv_safe(value: Any) -> Any:
    return "'" + value if isinstance(value, str) and value.startswith(RISKY) else value


def build(store, days, start: date, end: date, kinds: List[str], fmt: str = "xlsx", program: Optional[str] = None,
          user_id: Optional[int] = None, by: str = "", measure: str = "interval",
          deadline: Optional[float] = None) -> Tuple[bytes, str, str]:
    """The export as (bytes, file name, media type). Raises ValueError for a request it cannot answer, and
    ExportTooSlow (a ValueError) once ``deadline`` (time.monotonic()) has passed."""
    if end < start:
        raise ValueError("The period ends before it starts.")
    if (end - start).days + 1 > MAX_DAYS:
        raise ValueError(f"Pick at most {MAX_DAYS} days (records are kept 13 months).")
    if not kinds or any(k not in KINDS for k in kinds):
        raise ValueError("Pick what to export.")
    if fmt not in ("xlsx", "csv"):
        raise ValueError("Pick Excel or CSV.")
    if fmt == "csv" and len(kinds) != 1:
        raise ValueError("A CSV file holds one item: pick one, or choose Excel.")
    if measure not in MEASURES:
        raise ValueError("Pick a measure.")
    ordered = [k for k in KINDS if k in kinds]
    days = _Days(days, deadline, cover="cover" in kinds)  # Phase AJ: the owner's tick to include them or not
    weeks = _Weeks(days)
    who = (store.get_user(user_id) or {}).get("display_name", "") if user_id is not None else ""
    if fmt == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        for row in TABLES[ordered[0]](store, days, weeks, start, end, program, user_id, measure):
            writer.writerow([_csv_safe(v) for v in row])
        return (out.getvalue().encode("utf-8-sig"), f"Team_Scheduler_{ordered[0]}_{start}_to_{end}.csv", "text/csv")
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    about = wb.create_sheet("About this export")
    for key, value in (("Period", f"{start:%d %b %Y} to {end:%d %b %Y}"), ("Programs", _programs_said(program)),
                       ("Changes by", who or "Anyone"), ("Measure", MEASURES[measure]),
                       ("Items", ", ".join(KINDS[k] for k in ordered)), ("Downloaded by", by),
                       ("Downloaded at", datetime.now(EGYPT).strftime("%Y-%m-%d %H:%M")),
                       ("Times", "All times are Egypt time (Cairo clock: UTC+3 in summer, UTC+2 in winter)."),
                       ("Dates", "Attendance, day activity, breaks and worked hours follow the shift's own date (an "
                                 "overnight shift belongs to the day it started); changes, versions, runs and other "
                                 "actions follow when they happened."),
                       ("Changes by filter", "Applies to what people did (attendance, activity, changes, versions, "
                                             "runs, other actions), not to the breaks and worked-hours listings."),
                       ("Kept", "Records are kept 13 months; run files 30 to 60 days.")):
        head = WriteOnlyCell(about, value=key)
        head.font = bold
        about.append([head, _text_cell(about, value)])
    try:
        for kind in ordered:
            ws = wb.create_sheet(KINDS[kind][:31])
            rows = TABLES[kind](store, days, weeks, start, end, program, user_id, measure)
            header = next(rows)
            cells = []
            for value in header:
                cell = WriteOnlyCell(ws, value=value)
                cell.font = bold
                cells.append(cell)
            ws.append(cells)
            for row in rows:
                ws.append([_text_cell(ws, v) for v in row])
    except BaseException:  # a stopped export: finish the half-written tabs so nothing is left open
        for sheet in wb.worksheets:
            sheet.close()
        raise
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), f"Team_Scheduler_{start}_to_{end}.xlsx", XLSX
