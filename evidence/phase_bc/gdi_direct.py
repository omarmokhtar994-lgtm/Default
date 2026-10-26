# © 2026 Omar Mokhtar. All rights reserved.
"""Phase BC: re-plans the owner's channel plan file (channel_plan_GDI_170.xlsx) directly with the planner: the people, shifts, languages and breaks are
read back from the file, the needs and rules from the GDI input workbook's own channel tabs (weekday grids), and each setting is planned Sun to Sat with the
earlier days carried, breaks kept where the file has them. Reports the Chat and Email people next to what the file itself shows.
argv: plan file, input workbook, names (comma separated, Chat-and-Email only), then optional NAME=VALUE overrides of channel_plan constants."""
import json, sys, time
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from openpyxl import load_workbook
from webapp import channel_plan
from webapp.channels import read_channels
from webapp.channel_plan import plan_day

DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
RULES = {"segments": [[15, "Break 1"], [30, "Lunch"], [15, "Break 2"]], "by_length": [],
         "windows": {"break 1": {"earliest": None, "latest": None}, "lunch": {"earliest": None, "latest": None}, "break 2": {"earliest": None, "latest": None}},
         "edge_margin": None, "min_gap": 60, "preferred_gap": 150, "max_gap": 210, "max_concurrent": 99}

def minutes(text, after):
    h, m = (int(x) for x in str(text).split(":"))
    t = h * 60 + m
    while t < after:
        t += 1440
    return t

def read_plan(path):
    wb = load_workbook(path, data_only=True)
    week = {}
    owner = {}
    skills_h = {}
    for who, p, c, e, a in list(wb["Week"].iter_rows(min_row=2, values_only=True)):
        skills_h[who] = "".join(l for l, h in (("P", p), ("C", c), ("E", e)) if h) or "PCE"
    for d, day in enumerate(DAYS):
        rows = [r for r in wb[day].iter_rows(min_row=2, values_only=True) if r[0]]
        by = {}
        for who, lang, a, b, ch, mins in rows:
            by.setdefault(who, {"language": lang, "rows": []})["rows"].append((a, b, ch))
        people = []
        for who, info in by.items():
            t, spans = None, []
            for a, b, ch in info["rows"]:
                s = minutes(a, t if t is not None else 0) if t is not None else minutes(a, 0)
                e = minutes(b, s + 1)
                spans.append((s, e, ch)); t = e
            start, end = spans[0][0], spans[-1][1]
            breaks = [(ch, s, e - s) for s, e, ch in spans if ch in ("Break", "Lunch")]
            kinds = []
            n = 0
            for ch, s, mins_ in breaks:
                if ch == "Lunch":
                    kinds.append(("Lunch", s, mins_))
                else:
                    n += 1
                    kinds.append((f"Break {n}", s, mins_))
            people.append({"name": who, "language": info["language"], "start": start, "end": end, "breaks": kinds, "skills": "", "counts_for": {str(info["language"]).casefold()}})
            blocks = [(s, e, ch) for s, e, ch in spans if ch in ("Phone", "Chat", "Email", "All channels")]
            owner.setdefault(day, {})[who] = blocks
        week[day] = people
    return week, owner, skills_h

LETTER = {"Phone": "P", "Chat": "C", "Email": "E", "All channels": "A"}

def stats(by_day, names):
    out = {}
    for who in names:
        h = {"P": 0.0, "C": 0.0, "E": 0.0}; longest = []; runs_n = []
        for day, people in by_day.items():
            blocks = sorted(people.get(who, []))
            runs, last = [], None
            for a, b, c in blocks:
                c = LETTER.get(c, c)
                if c not in "PCE":
                    continue
                h[c] += (b - a) / 60
                if last == c:
                    runs[-1] += b - a
                else:
                    runs.append(b - a)
                last = c
            if runs:
                longest.append(max(runs) / 60); runs_n.append(len(runs))
        out[who] = {"chat_h": round(h["C"], 1), "email_h": round(h["E"], 1), "days": len(longest), "longest_h": round(sum(longest) / max(1, len(longest)), 2),
                    "stretches_a_day": round(sum(runs_n) / max(1, len(runs_n)), 2)}
    return out

def main():
    plan_file, book, names = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3].split(",")
    for item in sys.argv[4:]:
        k, v = item.split("="); setattr(channel_plan, k, int(v))
    week, owner, skills_h = read_plan(plan_file)
    setup = read_channels(book, 60)
    print("rules:", setup["rules"], "email_mode:", setup["email_mode"], flush=True)
    print("THE FILE      ", json.dumps(stats(owner, names)), flush=True)
    for label, flags in (("no tick", {}), ("fixed", {"consistent": True}), ("balanced", {"balanced": True})):
        conf = dict(setup, rules={**setup["rules"], "balanced": False, "consistent": False, **flags})
        carry, by_day, began = {}, {}, time.time()
        for d, day in enumerate(DAYS):
            people = week[day]
            if not people:
                continue
            for p in people:
                p["skills"] = "CE" if p["name"] in names else skills_h.get(p["name"], "PCE")
            plan = plan_day(people, conf, RULES, d, carry=carry, move_breaks=False)
            by_day[day] = plan["blocks"]
            for who, blocks in plan["blocks"].items():
                for a, b, c in blocks:
                    if c in "PCE":
                        carry.setdefault(who, {})[c] = carry.get(who, {}).get(c, 0) + b - a
        print(f"{label:14}", json.dumps(stats(by_day, names)), f"({time.time() - began:.0f}s)", flush=True)

if __name__ == "__main__":
    main()
