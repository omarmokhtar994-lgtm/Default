"""Phase C1 of the production-readiness audit (F-06, F-07).

Pre-registered rule: evidence/production_readiness_audit/phase_c/C1_RULE.txt.
Written to FAIL on the Phase B engine (commit 407a725) and pass after C1:

  - an elastic model: every coverage minimum may fall short by a slack that
    costs more than anything else, and the slacks are reported;
  - a shortfall pass that turns a no-schedule week into a non-releasable
    HARD_RULE_SHORTFALL_SCHEDULE listing every shortfall;
  - input errors never get one;
  - the next-Sunday floor is measured as coverage quality, not a hard failure.
"""
from __future__ import annotations

import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(ROOT / "engine" / "_tools"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import l632_universal_scheduler as E  # noqa: E402
import phase_a_workbooks as W  # noqa: E402

NAMES = [f"Agent {c}" for c in "ABCDEFGHIJKL"]
DAY_SHIFTS = [f"{h:02d}:00 - {(h + 9) % 24:02d}:00" for h in range(6, 12)]


def build(spec):
    tmp = Path(tempfile.mkdtemp())
    path = tmp / "case.xlsx"
    W.build(spec, path)
    return path, E.parse_input(path)


def load(path: Path, name: str):
    spec_ = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec_)
    sys.modules[name] = mod
    spec_.loader.exec_module(mod)
    return mod


# S04: 0.5 FTE on Sunday 03:00 that no legal day shift reaches.
S04 = dict(roster=[(n, "English") for n in NAMES[:10]], shifts=DAY_SHIFTS,
           demand=lambda d, m: 0.5 if (d == 0 and m == 180) else (3 if 8 * 60 <= m < 20 * 60 else None))
# S05: Spanish minimum 1 08:00-20:00 every day with only two Spanish agents.
S05 = dict(roster=[(n, "Spanish" if i < 2 else "English") for i, n in enumerate(NAMES[:10])],
           shifts=DAY_SHIFTS, demand=lambda d, m: 3 if 8 * 60 <= m < 20 * 60 else None,
           language_setup=[{"Language": "Spanish", "Coverage Start": "08:00", "Coverage End": "20:00",
                            "Active?": "Yes", "Coverage Group": "Spanish", "Minimum Per Interval": 1,
                            "Coverage Days": "All"}])
# One person, one Monday of demand: any break leaves a quarter empty.
LONE = dict(roster=[("Agent A", "English")], shifts=["08:00 - 17:00"],
            demand=lambda d, m: 1 if (d == 1 and 8 * 60 <= m < 17 * 60) else None,
            instructions={"Count of Associates": 1})


class C1ElasticStageOne(unittest.TestCase):
    def test_hard_config_has_an_elastic_switch_off_by_default(self):
        self.assertFalse(E.HardConfig().elastic)

    def test_an_unreachable_quarter_is_one_named_shortfall(self):
        _, parsed = build(S04)
        hard = E.build_skeleton(parsed, None, E.HardConfig(), 20, 2, io.StringIO())
        self.assertEqual(hard.cp_status, "INFEASIBLE")
        sk = E.build_skeleton(parsed, None, E.HardConfig(elastic=True), 20, 2, io.StringIO())
        self.assertIn(sk.cp_status, {"OPTIMAL", "FEASIBLE"})
        slacks = sk.diagnostics["elastic_slacks"]
        self.assertTrue(slacks)
        self.assertEqual({(s["family"], s["day"]) for s in slacks}, {("zero_active", "Sun")})
        self.assertTrue(all(s["time"].startswith("03:") for s in slacks), slacks)

    def test_a_language_shortage_is_named_as_language(self):
        _, parsed = build(S05)
        sk = E.build_skeleton(parsed, None, E.HardConfig(elastic=True), 30, 2, io.StringIO())
        self.assertIn(sk.cp_status, {"OPTIMAL", "FEASIBLE"})
        families = {s["family"] for s in sk.diagnostics["elastic_slacks"]}
        self.assertEqual(families, {"language"})

    def test_person_rules_stay_hard(self):
        _, parsed = build(S05)
        sk = E.build_skeleton(parsed, None, E.HardConfig(elastic=True), 30, 2, io.StringIO())
        for a, row in enumerate(sk.assignment):
            self.assertEqual(sum(v == "OFF" for v in row), 2, parsed.associates[a].name)


class C1ElasticStageTwo(unittest.TestCase):
    def test_a_lone_cover_takes_its_breaks_and_the_gap_is_named(self):
        _, parsed = build(LONE)
        sk = E.build_skeleton(parsed, None, E.HardConfig(), 20, 1, io.StringIO())
        self.assertIn(sk.cp_status, {"OPTIMAL", "FEASIBLE"})
        E.ensure_before_break_metrics(parsed, sk)  # as the pipeline does before Stage 2
        hard = E.solve_breaks(parsed, sk, 24, False, 20, 1, io.StringIO())
        self.assertNotIn(hard.cp_status, {"OPTIMAL", "FEASIBLE"})
        soft = E.solve_breaks(parsed, sk, 24, False, 20, 1, io.StringIO(), elastic=True)
        self.assertIn(soft.cp_status, {"OPTIMAL", "FEASIBLE"})
        self.assertFalse(soft.no_break_cells)
        slacks = soft.diagnostics["elastic_slacks"]
        self.assertEqual({s["family"] for s in slacks}, {"zero_coverage"})
        self.assertEqual(sum(s["shortfall"] for s in slacks), 4)  # 15 + 30 + 15 minutes


class F33BalanceTermsAreBuiltFromUncorruptedExpressions(unittest.TestCase):
    """Audit F-33, found while building C1. Stage 2 reused one stored
    ``headcount - breaks`` expression per quarter; OR-Tools 9.15 reduces
    ``k - (k - S)`` to the inner sum ``S`` and the following ``- allowed``
    modified it in place, so later balance terms read a wrong expression.

    The shortfall pass builds a fresh expression per use. Normal runs keep the
    measured objective: the pre-registered A/B of the fix did not pass
    (evidence/production_readiness_audit/phase_c/F33_AB_SCORE.json), so F-33
    stays open for them; their model is proven unchanged by
    f33_stage2_model_diff.py, not by a test here."""

    def test_the_shortfall_pass_uses_a_permitted_exception_instead_of_a_gap(self):
        # With the shared expression this one-person case had no solution at
        # all, even with a no-break exception permitted.
        spec = dict(LONE, instructions={"Count of Associates": 1,
                                        "Critical Coverage No-Break Exception Enabled": "Yes",
                                        "Critical Coverage No-Break Max Associate-Days": 1})
        _, parsed = build(spec)
        sk = E.build_skeleton(parsed, None, E.HardConfig(), 20, 1, io.StringIO())
        E.ensure_before_break_metrics(parsed, sk)
        sol = E.solve_breaks(parsed, sk, 24, True, 20, 1, io.StringIO(), exception_cap=1, elastic=True)
        self.assertIn(sol.cp_status, {"OPTIMAL", "FEASIBLE"})
        self.assertEqual(len(sol.no_break_cells), 1)
        self.assertEqual(sol.diagnostics["elastic_slacks"], [])

    def test_each_use_gets_a_fresh_expression(self):
        cp = E.import_cp_sat()
        model = cp.CpModel()
        bv = [model.NewBoolVar(f"b{i}") for i in range(3)]
        entry = (1, bv)
        first = E.after_break_raw_expression(entry)
        before = str(first)
        _ = 1 - E.after_break_raw_expression(entry) - 3
        _ = 1 - first - 3
        self.assertEqual(str(E.after_break_raw_expression(entry)), before)


class F34StageOneIsFundedForTheConfiguredSlice(unittest.TestCase):
    """Audit F-34, found while designing C3: the budget planner sized the Stage-1
    window for 45 s slices whatever slice the workbook configured."""

    def test_the_default_slice_is_unchanged(self):
        self.assertEqual(E.stage1_minimum_budget_seconds(3600, 15, 45.0), int(15 * 45 / 0.82))

    def test_a_deeper_slice_is_funded_up_to_the_cap(self):
        self.assertEqual(E.stage1_minimum_budget_seconds(3600, 15, 240.0), 1620)
        self.assertEqual(E.stage1_minimum_budget_seconds(3600, 3, 240.0), int(3 * 240 / 0.82))


class C1ShortfallPass(unittest.TestCase):
    def test_capacity_proofs_allow_the_pass_and_input_errors_do_not(self):
        self.assertTrue(E.shortfall_pass_allowed([{"code": "ZERO_ACTIVE_INTERVAL_PROVABLY_IMPOSSIBLE"}]))
        self.assertTrue(E.shortfall_pass_allowed([{"code": "LANGUAGE_WINDOW_MAX_CAPACITY_BELOW_MINIMUM"},
                                                  {"code": "SKILL_WINDOW_PAID_CAPACITY_PROVABLY_INSUFFICIENT"}]))
        self.assertFalse(E.shortfall_pass_allowed([{"code": "ZERO_ACTIVE_INTERVAL_PROVABLY_IMPOSSIBLE"},
                                                   {"code": "HARD_PREFERENCE_DUPLICATE_ASSOCIATE"}]))
        self.assertFalse(E.shortfall_pass_allowed([{"code": "FIXED_REQUEST_IN_BLANK_HOURS"}]))
        self.assertFalse(E.shortfall_pass_allowed([]))

    def test_the_pass_exports_a_listed_non_releasable_schedule(self):
        path, parsed = build(S04)
        out = path.parent / "case_L6_3_2_3_HARD_RULE_SHORTFALL_SCHEDULE.xlsx"
        # Phase H (owner-approved; evidence/phase_h/H4_DETERMINISTIC_BUDGET.md):
        # a solver-work budget instead of 40 s wall clock, so a slow or busy
        # machine (Colab: 3 of 10 failures) does the same search.
        # Re-pinned 10 -> 40 (2026-10-08): with 10 the refinement and break
        # solves stop short of their best answer about 1 run in 10 (4/40 on
        # x86; it failed the safety gate on the owner's ARM server), leaving
        # one extra uncovered half-hour (evidence/phase_i/I7_ROOT_CAUSE.md).
        # Chosen by the rule pre-registered in I7_SHORTFALL_BUDGET_RULE.txt:
        # 20 failed 1/20; 40 passed 20/20 unloaded and 10/10 under load.
        # The assertions below are unchanged.
        record = E.run_shortfall_pass(parsed, path, out, {}, E.capacity_diagnostics(parsed),
                                      time_limit=600, workers=2, log=io.StringIO(), random_seed=9000,
                                      deterministic_time=40.0)
        self.assertEqual(record["status"], "EXPORTED", record)
        self.assertTrue(out.exists())
        self.assertFalse(record["releasable"])
        # Shortfalls are measured on the schedule as published, after breaks.
        self.assertEqual({s["family"] for s in record["shortfalls"]}, {"zero_coverage"})
        self.assertTrue(all(s["day"] == "Sun" and s["time"].startswith("03:") for s in record["shortfalls"]),
                        record["shortfalls"])
        from openpyxl import load_workbook
        wb = load_workbook(out, read_only=True)
        self.assertIn("Shortfalls", wb.sheetnames)
        rows = [r for r in wb["Shortfalls"].iter_rows(min_row=2, values_only=True) if r and r[0]]
        self.assertEqual(len(rows), len(record["shortfalls"]))
        # The workbook names itself, so no reader can take it for a release.
        summary = {r[0]: r[1] for r in wb["Production Summary"].iter_rows(values_only=True) if r and r[0]}
        self.assertEqual(summary.get("Artifact Type"), "HARD_RULE_SHORTFALL_SCHEDULE")

        # The independent validator agrees: only the declared shortfall family fails.
        runner = load(ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py", "runner_phase_c")
        check = runner.validate_shortfall_schedule(path, out, None, out.parent / "SHORTFALL_VALIDATION.json")
        self.assertEqual(check["status"], "SHORTFALLS_CONFIRMED", check)
        self.assertEqual(check["validator_failure_types"], ["ZERO_STAFF_ACTIVE"])


class TheShortfallPassCanRunOnASolverWorkBudget(unittest.TestCase):
    """Phase H (evidence/phase_g/G2_SHORTFALL_REPRO.md): on a slow, busy
    2-core machine (Colab) the 40 s wall-clock pass did less search and the
    export test failed 3 of 10 times. An optional solver-work budget
    (CP-SAT max_deterministic_time) makes the amount of search independent of
    machine speed. Default None: production behaviour unchanged."""

    def test_the_limit_is_applied_only_when_set(self):
        cp_model = E.import_cp_sat()
        solver = cp_model.CpSolver()
        self.assertIsNone(E.configure_solver_limits(solver).get("max_deterministic_time"))
        old = E.SOLVER_DETERMINISTIC_TIME
        E.SOLVER_DETERMINISTIC_TIME = 7.5
        try:
            solver = cp_model.CpSolver()
            self.assertEqual(E.configure_solver_limits(solver)["max_deterministic_time"], 7.5)
            self.assertEqual(solver.parameters.max_deterministic_time, 7.5)
        finally:
            E.SOLVER_DETERMINISTIC_TIME = old

    def test_the_pass_splits_its_budget_and_restores_the_default(self):
        seen = []
        real_build, real_breaks = E.build_skeleton, E.solve_breaks

        def spy_build(*a, **kw):
            seen.append(("stage1", E.SOLVER_DETERMINISTIC_TIME))
            return real_build(*a, **kw)

        def spy_breaks(*a, **kw):
            seen.append(("stage2", E.SOLVER_DETERMINISTIC_TIME))
            return real_breaks(*a, **kw)
        path, parsed = build(S04)
        E.build_skeleton, E.solve_breaks = spy_build, spy_breaks
        try:
            E.run_shortfall_pass(parsed, path, path.parent / "case_SHORTFALL.xlsx", {}, E.capacity_diagnostics(parsed),
                                 time_limit=600, workers=2, log=io.StringIO(), random_seed=9000,
                                 deterministic_time=10.0)
        finally:
            E.build_skeleton, E.solve_breaks = real_build, real_breaks
        self.assertEqual(seen, [("stage1", 4.0), ("stage1", 3.0), ("stage2", 3.0)])
        self.assertIsNone(E.SOLVER_DETERMINISTIC_TIME)


class C1NextSundayFloorIsCoverageQuality(unittest.TestCase):
    def test_a_boundary_floor_miss_is_not_a_hard_failure(self):
        # Saturday-night carry-out covers next Sunday 00:00-05:00 with one
        # person against 3 FTE: the floor is missed, nobody is missing.
        spec = dict(roster=[("Agent A", "English"), ("Agent B", "English")],
                    shifts=["20:00 - 05:00", "08:00 - 17:00"],
                    demand=lambda d, m: 3 if m < 300 else (1 if 8 * 60 <= m < 17 * 60 else None))
        _, parsed = build(spec)
        idx = {s.label: s.index for s in parsed.shifts}
        rows = [["OFF", "OFF", "08:00 - 17:00", "08:00 - 17:00", "08:00 - 17:00", "08:00 - 17:00", "20:00 - 05:00"],
                ["OFF", "08:00 - 17:00", "08:00 - 17:00", "08:00 - 17:00", "08:00 - 17:00", "08:00 - 17:00", "OFF"]]
        sk = E.SkeletonSolution("t", "FEASIBLE", 0.0, 0.0, rows, [[idx.get(v) for v in r] for r in rows], {})
        m = E.calculate_metrics(parsed, sk, {}, [])
        self.assertGreater(m["week_boundary_floor_gap_count"], 0)
        self.assertEqual(m["week_boundary_zero_staffed_active_quarters"], 0)
        self.assertEqual(m["week_boundary_hard_failure_count"], 0)
        gate = E.production_quality_gate(parsed, m)
        codes = {row.get("code") for row in gate["failures"] + gate["warnings"] + gate["suppressed_by_disabled_gates"]}
        self.assertIn("NEXT_SUNDAY_FLOOR_GAPS", codes)


if __name__ == "__main__":
    unittest.main()
