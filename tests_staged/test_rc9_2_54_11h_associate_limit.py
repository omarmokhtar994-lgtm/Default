"""Phase H, task 2: "Max 11H/3OFF Associates" caps the long pattern per week.

Owner, 2026-10-07: "add an option to use 11/3 shift but with a count like 1
slot only or 2"; asked what the count means: "if i ask for 2 i expected 2
associates only with full shift 11 hours". 11/3 is the engine's existing
11H/3OFF pattern (11-hour shifts, three OFF days); each associate works either
it or the 9H/2OFF week. The new instruction caps how many associates may work
the 11H/3OFF pattern in the week. Blank = no limit (default, unchanged).

Pinned here:
  * default: no limit, contract fingerprint unchanged;
  * values: "2", "2.0", " 2 " read as 2; "two", "0", "-1" are named input errors;
  * the limit binds in the solve (a demand only 11h shifts reach);
  * set while Use 11H/3OFF is No: no effect;
  * more associates forced long than the limit: named before solving;
  * the independent validator fails a week over the limit;
  * Production Summary reports the limit and how many associates used it.
"""
import copy
import os
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(REPO / "tools"))
sys.path.insert(0, str(ROOT / "engine" / "tools"))
import build_synthetic_suite as B  # noqa: E402

E = B.load_engine(ROOT / "engine" / "_tools" / "l632_universal_scheduler.py")
import independent_validator as IV  # noqa: E402

TEMPLATE = REPO / "fixtures" / "SYNTHETIC_FIXTURE_FLOOR_NOT_80.xlsx"
KEY = "Max 11H/3OFF Associates"
LONG = ("08:00 - 19:00", "08:00", "19:00", 11, "11H", "synthetic")


def workbook(limit=None, use11="Yes", long_only=False, n=6):
    names, langs = B.english(n)
    settings = {"Short Break Count": 0, "Lunch Count": 0, "Use 11H/3OFF": use11,
                "Allowed Shift Durations Hours": "11" if long_only else "9, 11"}
    if limit is not None:
        settings[KEY] = limit
    case = dict(id="H2_11H", tier="test", associates=n, shift_starts=[] if long_only else [8 * 60],
                shift_minutes=540, shrinkage=0.0, interval_minutes=60, names=names, languages=langs,
                language_rules=B.ENGLISH_RULE, settings=settings, extra_shift_rows=[LONG])
    # 08:00-19:00: 17:00-19:00 is reachable only from the 11-hour shift.
    demand = [[(2.0 if 8 <= h < 19 else None) for h in range(24)] for _ in range(7)]
    out = Path(tempfile.mkdtemp()) / "H2_11H.xlsx"
    B.build_workbook(TEMPLATE, out, copy.deepcopy(case), E, demand)
    return out


def solve(p, seconds=30.0):
    prof = next(pr for pr in E.skeleton_profiles(["target90_restore_champion"])
                if pr["name"] == "target90_restore_champion")
    sk = E.build_skeleton(p, dict(prof), E.HardConfig(hard_floor=False, elastic=True), seconds, 2, None,
                          random_seed=9000)
    assert sk.cp_status in ("OPTIMAL", "FEASIBLE"), sk.cp_status
    return sk


def long_associates(p, sk):
    return {a for a, _, si in E.scheduled_cells(sk) if p.shifts[si].duration_min >= E.LONG_SHIFT_MIN_DURATION_MIN}


def hard_warnings(book):
    p = E.parse_input(book)
    return [w for w in getattr(p, "parser_warnings", []) if "11H_ASSOCIATE_LIMIT" in w]


class TheInstruction(unittest.TestCase):
    def test_default_is_no_limit_and_contract_unchanged(self):
        p = E.parse_input(workbook())
        self.assertIsNone(p.max_11h_associates)
        self.assertNotIn("max_11h_associates", E.input_contract_payload(p))
        self.assertEqual(E.input_contract_payload(E.parse_input(workbook(limit=2)))["max_11h_associates"], 2)

    def test_limit_values(self):
        for raw in ("2", "2.0", " 2 ", 2):
            with self.subTest(raw=raw):
                self.assertEqual(E.parse_input(workbook(limit=raw)).max_11h_associates, 2)
        for raw in ("two", "0", "-1", "1.5"):
            with self.subTest(raw=raw):
                self.assertTrue(any(w.startswith("HARD_INVALID_11H_ASSOCIATE_LIMIT") for w in hard_warnings(workbook(limit=raw))))


class TheLimitBindsInTheSolve(unittest.TestCase):
    def test_limit_caps_the_long_pattern(self):
        free = E.parse_input(workbook())
        capped = E.parse_input(workbook(limit=1))
        self.assertGreater(len(long_associates(free, solve(free))), 1)
        self.assertLessEqual(len(long_associates(capped, solve(capped))), 1)

    def test_limit_without_11h_mode_is_inert(self):
        p = E.parse_input(workbook(limit=1, use11="No"))
        self.assertEqual(long_associates(p, solve(p, 10.0)), set())
        rows = dict(E.long_mode_summary_rows(p, solve(p, 10.0)))
        self.assertEqual(rows, {})


class TheContractNamesAnImpossibleLimit(unittest.TestCase):
    def test_limit_below_forced_long_associates_is_named(self):
        p = E.parse_input(workbook(limit=2, long_only=True, n=6))
        failures = E.validate_input_contract(p)["failures"]
        hit = [f for f in failures if f.get("code") == "11H_ASSOCIATE_LIMIT_BELOW_FORCED"]
        self.assertTrue(hit, failures)
        self.assertIn("6", str(hit[0]))


class TheValidatorChecksTheLimit(unittest.TestCase):
    def test_validator_fails_a_week_over_the_limit(self):
        free_book, capped_book = workbook(), workbook(limit=1)
        free = E.parse_input(free_book)
        sk = solve(free)
        self.assertGreater(len(long_associates(free, sk)), 1)
        capped = E.parse_input(capped_book)
        patterns = E.generate_break_patterns(capped)
        solution = E.BreakSolution("test", sk.profile, "FEASIBLE", 0.0, 0.0, 115, False, {}, set(), patterns,
                                   {}, E.calculate_metrics(capped, sk, {}, patterns))
        out = capped_book.parent / "H2_11H_L6_3_2_3_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx"
        E.write_output_workbook(capped_book, out, capped, sk, solution, E.capacity_diagnostics(capped), [], {}, [])
        result = IV.validate(capped_book, out, ROOT / "engine" / "_tools" / "l632_universal_scheduler.py")
        self.assertIn("11H_ASSOCIATE_LIMIT_EXCEEDED", {f.get("type") for f in result["failures"]})


class TheOutputReportsIt(unittest.TestCase):
    def test_summary_rows(self):
        p = E.parse_input(workbook(limit=1))
        sk = solve(p)
        rows = dict(E.long_mode_summary_rows(p, sk))
        self.assertEqual(rows["Max 11H/3OFF Associates"], 1)
        self.assertEqual(rows["11H/3OFF Associates Used"], len(long_associates(p, sk)))
        free = E.parse_input(workbook())
        self.assertEqual(dict(E.long_mode_summary_rows(free, solve(free, 10.0)))["Max 11H/3OFF Associates"], "No limit")


class TheWorkbookAndTemplateCarryIt(unittest.TestCase):
    def test_the_output_workbook_reports_it(self):
        from openpyxl import load_workbook
        book = workbook(limit=1)
        p = E.parse_input(book)
        sk = solve(p, 10.0)
        patterns = E.generate_break_patterns(p)
        solution = E.BreakSolution("test", sk.profile, "FEASIBLE", 0.0, 0.0, 115, False, {}, set(), patterns,
                                   {}, E.calculate_metrics(p, sk, {}, patterns))
        out = book.parent / "H2_OUT.xlsx"
        E.write_output_workbook(book, out, p, sk, solution, E.capacity_diagnostics(p), [], {}, [])
        ws = load_workbook(out, read_only=True)["Production Summary"]
        rows = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r and r[0]}
        self.assertEqual(rows.get("Max 11H/3OFF Associates"), 1)
        self.assertEqual(rows.get("11H/3OFF Associates Used"), len(long_associates(p, sk)))

    def test_the_validator_cross_checks_the_raw_cell(self):
        book = workbook(limit=2)
        p = copy.deepcopy(E.parse_input(book))
        self.assertEqual(IV.independent_input_crosscheck(book, p)["mismatches"], [])
        p.max_11h_associates = None
        self.assertIn("max_11h_associates", {m["check"] for m in IV.independent_input_crosscheck(book, p)["mismatches"]})

    def test_the_input_template_offers_it(self):
        sys.path.insert(0, str(REPO / "tools"))
        import build_input_template as T
        rows = [name for _, section in T.SETUP_LAYOUT for name, _ in section]
        self.assertIn(KEY, rows)
        self.assertIn(KEY, T.HELP)


class TheReadyToEditWorkbooksOfferTheNewRows(unittest.TestCase):
    """The template builder had the Phase F and Phase H rows, but the seven
    ready-to-edit workbooks users fill in were never rebuilt, so neither
    option could be chosen in them."""

    def test_every_ready_workbook_has_both_rows(self):
        from openpyxl import load_workbook
        inputs = next(p for p in (REPO / "inputs", REPO / "packages" / "rc9_2_2_production" / "inputs") if p.is_dir())
        books = sorted((inputs / "ready_to_edit").glob("*.xlsx"))
        self.assertGreaterEqual(len(books), 7)
        for book in books:
            with self.subTest(book=book.name):
                ws = load_workbook(book, read_only=True)["Instructions"]
                labels = {str(c).strip() for row in ws.iter_rows(values_only=True, max_col=4) for c in row if c}
                self.assertIn("Allow Half-Hour Starts", labels)
                self.assertIn(KEY, labels)


if __name__ == "__main__":
    unittest.main()
