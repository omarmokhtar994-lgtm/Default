# © 2026 Omar Mokhtar. All rights reserved.
"""Phase O task 9: the shift handover note, and its row in the exports' daily summary."""
import io
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from webapp.attendance import DayBook
from webapp.day import day_view
from webapp.exports import build
from webapp.handover import note
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_schedules import AFTER, INPUT, REPO

WED = date(2026, 10, 14)


class TheHandover(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program="AE/AR B2B",
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))
        self.days.set_status("AE/AR B2B", WED, "Associate 012", "Sick", self.sara)
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Late", self.sara, end="13:00")
        self.days.move_break("AE/AR B2B", WED, "Associate 008", 1, "13:05", self.sara)
        self.days.add_activity("AE/AR B2B", WED, "Associate 028", "Overtime", "07:00", "09:00", self.sara)

    def test_note_lists_what_happened(self):
        n = note(self.days, "AE/AR B2B", WED)
        self.assertEqual([(p["name"], p["status"]) for p in n["absent"]], [("Associate 012", "Sick")])
        self.assertEqual([(p["name"], p["text"]) for p in n["late_early"]], [("Associate 001", "late, in at 13:00")])
        self.assertEqual([(m["name"], m["what"]) for m in n["moves"]], [("Associate 008", "Lunch moved 12:30 to 13:05")])
        self.assertEqual([(a["name"], a["text"]) for a in n["activities"]], [("Associate 028", "Overtime 07:00 to 09:00")])
        self.assertEqual(n["numbers"]["overtime_minutes"], 120)
        page = self.days.page("AE/AR B2B", WED)
        self.assertEqual(n["numbers"]["short_hours"], page["view"]["tiles"]["short_hours"])
        self.assertEqual(n["short"], [(c["t"], c["pm"]) for c in page["view"]["cells"] if c["pm"] is not None and c["pm"] < 0])

    def test_a_called_in_day_off_is_in_the_note(self):
        self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "12:00", "21:00", self.sara)
        n = note(self.days, "AE/AR B2B", WED)
        self.assertIn(("Associate 002", "Day off cancelled: called in 12:00 - 21:00"),
                      [(a["name"], a["text"]) for a in n["activities"]])
        self.assertEqual(n["numbers"]["called_in_minutes"], 9 * 60)
        data, *_ = build(self.store, self.days, WED, WED, ["summary"], by="Omar")
        rows = list(load_workbook(io.BytesIO(data))["Daily summary"].iter_rows(values_only=True))
        self.assertEqual(dict(zip(rows[0], rows[1]))["Called in min"], 9 * 60)

    def test_watch_tomorrow_from_the_plan(self):
        n = note(self.days, "AE/AR B2B", WED)
        page = self.days.page("AE/AR B2B", date(2026, 10, 15))
        expected = [(c["t"], c["plan_pm"]) for c in page["view"]["cells"] if c["plan_pm"] is not None and c["plan_pm"] < 0]
        self.assertEqual(n["tomorrow"]["short"], expected)
        self.assertEqual(n["tomorrow"]["date"], date(2026, 10, 15))

    def test_daily_summary_in_the_exports(self):
        data, *_ = build(self.store, self.days, WED, WED, ["summary"], by="Omar")
        ws = load_workbook(io.BytesIO(data))["Daily summary"]
        rows = list(ws.iter_rows(values_only=True))
        row = dict(zip(rows[0], rows[1]))
        n = note(self.days, "AE/AR B2B", WED)
        self.assertEqual((row["Program"], row["Absent"], row["Late or early"], row["Overtime min"], row["Breaks moved"],
                          row["Hours short"]),
                         ("AE/AR B2B", 1, 1, 120, 1, n["numbers"]["short_hours"]))


if __name__ == "__main__":
    unittest.main()
