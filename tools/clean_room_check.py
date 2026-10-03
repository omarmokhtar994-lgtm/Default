#!/usr/bin/env python3
"""Clean-room re-check of a finished schedule. Shares no code with the engine.

The shipped independent validator recomputes coverage from the output cells,
but it takes the input contract from the engine's own parse_input and calls
engine helpers for most rule definitions (rest, language rules, break caps).
A defect in one of those shared pieces is invisible to both. This script
re-reads the raw input workbook and the output workbook with openpyxl only,
re-implements every rule from its business meaning, and compares its numbers
with what the engine and the validator published.

    python3 tools/clean_room_check.py --input IN.xlsx --output OUT.xlsx \
        [--audit AUDIT.json] [--validation INDEPENDENT_VALIDATION.json] [--json-out R.json]

Scope (what it re-derives):
  * contract: target, floor, rest gap, strict/separate OFF, hard OFF, leave,
    fixed requests, 11H/3OFF, break contract, language rules and windows,
    previous-Saturday carry-in, demand and shrinkage grids;
  * schedule: OFF count and shape, rest (incl. Sat->Sun wrap and previous
    Saturday), leave / hard OFF / fixed honoured, max different shifts,
    language working window, break count/length/order/containment/gap;
  * coverage: before/after-break effective staffing per interval, target /
    floor / 80 / 90 / 100 hits, language minimum gaps, zero-staffed quarters.

Interpretation choices are written next to the code. Where this script and the
engine disagree, the disagreement is the finding; neither side is assumed right.
"""
from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from openpyxl import load_workbook

DAYS = ["sun", "mon", "tue", "wed", "thu", "fri", "sat"]
FULL = ["sunday", "monday", "tuesday", "wednesday", "thursday", "friday", "saturday"]
EPS = 1e-9


def n(v: Any) -> str:
    return re.sub(r"\s+", " ", str("" if v is None else v).strip()).lower()


def clock(v: Any) -> Optional[int]:
    """Minute of day from a cell: time object, Excel fraction, or 'HH:MM'."""
    if v is None or v == "":
        return None
    if hasattr(v, "hour") and hasattr(v, "minute"):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 0 <= v < 1:
            return int(round(v * 1440)) % 1440
        return None
    m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*", str(v))
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    return h * 60 + mi if h < 24 and mi < 60 else None


def shift_of(label: Any) -> Optional[Tuple[int, int]]:
    """(start minute, duration minutes) from 'HH:MM - HH:MM'."""
    found = re.findall(r"(\d{1,2}):(\d{2})", str(label or ""))
    if len(found) < 2:
        return None
    s = int(found[0][0]) * 60 + int(found[0][1])
    e = int(found[1][0]) * 60 + int(found[1][1])
    d = (e - s) % 1440 or 1440
    return s, d


def num(v: Any, pct: bool = False) -> Optional[float]:
    if v is None or v == "" or isinstance(v, bool):
        return None
    t = str(v).strip()
    try:
        if t.endswith("%"):
            return float(t[:-1]) / 100.0
        x = float(t)
    except ValueError:
        return None
    if pct and x > 1.5:
        x /= 100.0
    return x


YES = {"yes", "y", "true", "1", "enabled", "on", "hard", "required"}


def yes(v: Any, default: bool) -> bool:
    t = n(v)
    return default if not t else t in YES


def day_index(header: Any) -> Optional[int]:
    if hasattr(header, "weekday"):
        return (header.weekday() + 1) % 7
    h = n(header)
    for i, (s, f) in enumerate(zip(DAYS, FULL)):
        if h in (s, f):
            return i
    return None


# --------------------------------------------------------------------------- input
class Contract:
    pass


def instructions(wb) -> Dict[str, Any]:
    """Label -> value. Engine Defaults first, Instructions override; last row wins.

    Both layouts seen in delivered workbooks are read: label in column B with
    value in C (current template), and label in A with value in B (older).
    """
    out: Dict[str, Any] = {}
    for title in ("Engine Defaults", "Instructions"):
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        for r in range(1, ws.max_row + 1):
            a, b, c = (ws.cell(r, k).value for k in (1, 2, 3))
            if a not in (None, "") and b not in (None, ""):
                out[n(a)] = b
            if b not in (None, "") and c not in (None, ""):
                out[n(b)] = c
    return out


def get(im: Dict[str, Any], *names: str, default: Any = None) -> Any:
    for name in names:
        if n(name) in im:
            return im[n(name)]
    return default


def header_row(ws, must: Tuple[str, ...], rows: int = 30) -> Optional[int]:
    for r in range(1, min(ws.max_row, rows) + 1):
        cells = [n(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if all(any(m == x or m in x for x in cells) for m in must):
            return r
    return None


def day_columns(ws, r: int) -> Optional[List[int]]:
    cols: Dict[int, int] = {}
    for c in range(1, ws.max_column + 1):
        d = day_index(ws.cell(r, c).value)
        if d is not None and d not in cols:
            cols[d] = c
    return [cols[d] for d in range(7)] if len(cols) == 7 else None


def grid(wb, title: str, step: int) -> Dict[Tuple[int, int], Optional[float]]:
    ws = wb[title]
    r0 = next(r for r in range(1, 20) if day_columns(ws, r))
    cols = day_columns(ws, r0)
    out: Dict[Tuple[int, int], Optional[float]] = {}
    for r in range(r0 + 1, ws.max_row + 1):
        m = clock(ws.cell(r, 1).value)
        if m is None or m % step:
            continue
        for d, c in enumerate(cols):
            out[(d, m // step)] = num(ws.cell(r, c).value)
    return out


def read_contract(path: Path) -> Contract:
    wb = load_workbook(path, data_only=False)
    im = instructions(wb)
    k = Contract()
    k.im = im
    k.input_issues: List[Dict[str, Any]] = []
    k.step = int(num(get(im, "Interval Granularity Minutes", "Interval Minutes", "Interval")) or 60)
    req_title = str(get(im, "Requirements Source", "Requirement Source", default=f"FT Wise {k.step} Min"))
    shr_title = str(get(im, "Shrinkage Source", "Shrinkage Sheet", default=f"Shrinkage {k.step} Min"))
    req = grid(wb, next(t for t in wb.sheetnames if n(t) == n(req_title)), k.step)
    shr_sheet = next((t for t in wb.sheetnames if n(t) == n(shr_title)), None)
    shr = grid(wb, shr_sheet, k.step) if shr_sheet else {}
    k.per_day = 1440 // k.step
    k.req = [[req.get((d, i)) for i in range(k.per_day)] for d in range(7)]
    k.shr = [[(shr.get((d, i)) or 0.0) for i in range(k.per_day)] for d in range(7)]
    k.target = num(get(im, "Target", "Coverage Target", "Interval Target", default=0.9), pct=True)
    k.floor = num(get(im, "Minimum Per Interval", "Coverage Floor", "Minimum Coverage Percentage", default=0.8), pct=True)
    k.rest_h = num(get(im, "Difference Between Shifts", "Rest Gap Hours", "Minimum Rest Gap", default=12))
    k.strict_off = yes(get(im, "Strict 2 OFF", "Strict Two OFF", "Strict OFF Count"), True)
    k.separate_off = yes(get(im, "Separate OFF Days", "Separate Off Days", "Allow Separate OFF Days"), True)
    k.hard_off = yes(get(im, "Hard OFF Preferences", "Hard OFF", "OFF Preferences Hard"), True)
    k.leave = yes(get(im, "Leave", "Leave Days", "Leave Enabled"), True)
    k.fixed_on = yes(get(im, "Fixed Request Use", "Use Fixed Requests", "Fixed/Nesting Enabled", "Use Fixed/Nesting"), False)
    k.max_diff = int(num(get(im, "Count of Different Shifts Per week", "Max Different Shifts per Week", default=3)) or 3)
    k.long_min = 630
    shorts = int(num(get(im, "Short Break Count", "Count of Short Breaks", "Number of Short Breaks", default=2)))
    short_len = int(num(get(im, "Short Break Duration Minutes", "Short Break Duration", "Break Duration Minutes", default=15)))
    lunches = int(num(get(im, "Lunch Count", "Meal Count", "Count of Lunch Breaks", default=1)))
    lunch_len = int(num(get(im, "Lunch Duration Minutes", "Lunch Duration", "Meal Duration Minutes", default=30)))
    k.break_lengths = sorted([short_len] * shorts + [lunch_len] * lunches)
    hard_gap = num(get(im, "Break Absolute Minimum Gap Minutes", "Absolute Minimum Gap Between Breaks Minutes",
                       "Break Hard Minimum Gap Minutes", "Minimum Legal Break Separation Minutes"))
    k.break_min_gap = int(hard_gap) if hard_gap is not None else 60
    raw_mode = re.sub(r"[^a-z0-9]", "", n(get(im, "Language Working Window", "Language Working Hours",
                                             "Language Window Enforcement")))
    k.window_mode = ("OFF" if raw_mode in ("", "off", "no", "none", "disabled", "coverageonly", "minimumonly")
                     else "ALL_ROWS" if raw_mode.startswith("all") or raw_mode == "everyrow"
                     else "REQUIRED_LANGUAGE_ONLY" if "required" in raw_mode or "exclusive" in raw_mode
                     else "MINIMUM_ROWS")

    # Roster: Schedule sheet, located by header text.
    ws = wb["Schedule"]
    hr = header_row(ws, ("language",))
    heads = {n(ws.cell(hr, c).value): c for c in range(1, ws.max_column + 1)}
    name_c = next(c for h, c in heads.items() if h in ("sf name", "name", "associate name", "employee name"))
    lang_c = next(c for h, c in heads.items() if h == "language")
    k.names, k.lang = [], {}
    for r in range(hr + 1, ws.max_row + 1):
        nm = str(ws.cell(r, name_c).value or "").strip()
        if nm and n(nm) not in ("total", "staffed hc", "required hc"):
            k.names.append(nm)
            k.lang[n(nm)] = n(ws.cell(r, lang_c).value)

    def by_name_sheet(title: str) -> Dict[str, List[str]]:
        out: Dict[str, List[str]] = {}
        if title not in wb.sheetnames:
            return out
        s = wb[title]
        r0 = next((r for r in range(1, 20) if day_columns(s, r)), None)
        if r0 is None:
            raise ValueError(f"{title}: no Sun..Sat header row")
        cols = day_columns(s, r0)
        hc = {n(s.cell(r0, c).value): c for c in range(1, s.max_column + 1)}
        nc = next(c for h, c in hc.items() if "name" in h)
        ac = next((c for h, c in hc.items() if h.startswith("active")), None)
        for r in range(r0 + 1, s.max_row + 1):
            nm = n(s.cell(r, nc).value)
            if not nm or (ac and not yes(s.cell(r, ac).value, True)):
                continue
            if nm in out:
                # Two rows for one person: which one did the planner mean? Do not guess.
                k.input_issues.append({"rule": "duplicate_input_row", "sheet": title, "who": nm, "row": r})
            out[nm] = [str(s.cell(r, c).value or "").strip() for c in cols]
        return out

    k.pref = by_name_sheet("Preference")
    k.fixed = by_name_sheet("Fixed Request") if k.fixed_on else {}
    k.prev_sat: Dict[str, str] = {}
    if "Previous week scheduled" in wb.sheetnames:
        s = wb["Previous week scheduled"]
        hr2 = header_row(s, ("name",))
        hc = {n(s.cell(hr2, c).value): c for c in range(1, s.max_column + 1)}
        nc = next(c for h, c in hc.items() if "name" in h)
        sc = next(c for h, c in hc.items() if h.startswith("sat") or "saturday" in h)
        for r in range(hr2 + 1, s.max_row + 1):
            nm = n(s.cell(r, nc).value)
            if nm:
                k.prev_sat[nm] = str(s.cell(r, sc).value or "").strip()

    # Shift library: every 'HH:MM - HH:MM' label in the first column.
    k.library: Set[str] = set()
    s = wb["Shift Library"]
    for r in range(1, s.max_row + 1):
        if shift_of(s.cell(r, 1).value):
            k.library.add(n(s.cell(r, 1).value))

    # Language Setup -> minimum rules and working windows.
    k.rules: List[Dict[str, Any]] = []
    k.windows: Dict[str, List[Tuple[int, int, bool, Set[int]]]] = defaultdict(list)
    if "Language Setup" in wb.sheetnames:
        s = wb["Language Setup"]
        hr3 = header_row(s, ("language", "minimum"))
        hc = {n(s.cell(hr3, c).value): c for c in range(1, s.max_column + 1)}

        def col(*keys: str) -> Optional[int]:
            return next((c for h, c in hc.items() if any(key in h for key in keys)), None)
        c_l, c_s, c_e = col("language"), col("coverage start", "start time"), col("coverage end", "end time")
        c_can, c_act, c_grp = col("can cover"), col("active"), col("coverage group", "group")
        c_min, c_days = col("minimum"), col("coverage days", "active days")
        rows = []
        can: Dict[str, Set[str]] = defaultdict(set)
        for r in range(hr3 + 1, s.max_row + 1):
            lang = n(s.cell(r, c_l).value)
            if not lang:
                continue
            targets = {n(t) for t in re.split(r"[,;/|]+", str(s.cell(r, c_can).value or "")) if n(t)} if c_can else set()
            can[lang] |= targets | {lang}
            if c_act and not yes(s.cell(r, c_act).value, True):
                continue
            st, en = clock(s.cell(r, c_s).value), clock(s.cell(r, c_e).value)
            mn = num(s.cell(r, c_min).value) if c_min else 0
            days = parse_days(s.cell(r, c_days).value if c_days else "")
            grp = n(s.cell(r, c_grp).value) if c_grp else lang
            rows.append((lang, st, en, int(round(mn or 0)), days, grp or lang))
        for lang in k.lang.values():
            can[lang] |= {lang}
        grouped: Dict[Tuple, Set[str]] = defaultdict(set)
        for lang, st, en, mn, days, grp in rows:
            if st is None or en is None:
                continue
            if mn > 0:
                grouped[(grp, st, en, mn, tuple(sorted(days)))].add(lang)
            if st != en:
                k.windows[lang].append((st, en, mn > 0, days))
        for (grp, st, en, mn, days), required in grouped.items():
            eligible = {src for src, tg in can.items() if tg & required}
            k.rules.append({"group": grp, "start": st, "end": en, "minimum": mn,
                            "days": set(days), "eligible": eligible})
    return k


def parse_days(v: Any) -> Set[int]:
    t = re.sub(r"[^a-z0-9,\-;/ ]", "", n(v))
    squashed = t.replace(" ", "")
    if squashed in ("", "all", "alldays", "daily", "everyday", "7days"):
        return set(range(7))
    if squashed in ("weekday", "weekdays"):
        return {1, 2, 3, 4, 5}
    if squashed in ("weekend", "weekends"):
        return {0, 6}
    idx = {**{s: i for i, s in enumerate(DAYS)}, **{f: i for i, f in enumerate(FULL)},
           "tues": 2, "thur": 4, "thurs": 4}
    out: Set[int] = set()
    for part in re.split(r"[,;/]+", t):
        part = part.strip()
        if not part:
            continue
        ends = [p.strip() for p in re.split(r"\s*-\s*|\s+to\s+", part)]
        if len(ends) == 2 and ends[0] in idx and ends[1] in idx:
            d = idx[ends[0]]
            while True:
                out.add(d)
                if d == idx[ends[1]]:
                    break
                d = (d + 1) % 7
        elif part in idx:
            out.add(idx[part])
        else:
            raise ValueError(f"unreadable Coverage Days {v!r}")
    return out


def rule_in_force(rule: Dict[str, Any], day: int, minute: int) -> bool:
    """Business meaning of a day-specific window: it OPENS on each listed day.

    A Mon-Fri 18:00-05:00 window covers Mon 18:00 -> Tue 05:00 ... Fri 18:00 ->
    Sat 05:00. So the early-morning part belongs to the PREVIOUS day's opening;
    Monday 01:00 is in force only if Sunday is a listed day.
    """
    st, en, days = rule["start"], rule["end"], rule["days"]
    if st == en:
        return day in days
    if st < en:
        return day in days and st <= minute < en
    if minute >= st:
        return day in days
    if minute < en:
        return (day - 1) % 7 in days
    return False


# -------------------------------------------------------------------------- output
def read_output(path: Path) -> Tuple[Dict[str, List[str]], List[Dict[str, Any]], Set[Tuple[str, int]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Schedule"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, row in enumerate(rows[:20]) if any(n(v) in ("sf name", "name") for v in row))
    heads = rows[hi]
    name_i = next(i for i, v in enumerate(heads) if n(v) in ("sf name", "name"))
    day_i = {day_index(v): i for i, v in enumerate(heads) if day_index(v) is not None}
    sched: Dict[str, List[str]] = {}
    for row in rows[hi + 1:]:
        nm = n(row[name_i] if name_i < len(row) else None)
        if nm:
            sched[nm] = [str(row[day_i[d]] or "").strip() for d in range(7)]
    breaks: List[Dict[str, Any]] = []
    sheet = "Break Schedule Active" if "Break Schedule Active" in wb.sheetnames else "Break Schedule"
    brows = list(wb[sheet].iter_rows(values_only=True))
    bh = {n(v): i for i, v in enumerate(brows[0])}
    for row in brows[1:]:
        if not row or not row[bh["associate"]]:
            continue
        d = day_index(row[bh["day"]])
        if d is None:
            continue
        breaks.append({"name": n(row[bh["associate"]]), "day": d, "type": n(row[bh["break type"]]),
                       "start": clock(row[bh["start"]]), "len": int(row[bh["duration minutes"]] or 0),
                       "status": n(row[bh.get("status", 0)])})
    exceptions: Set[Tuple[str, int]] = set()
    if "No-Break Exceptions" in wb.sheetnames:
        erows = list(wb["No-Break Exceptions"].iter_rows(values_only=True))
        eh = {n(v): i for i, v in enumerate(erows[0])}
        for row in erows[1:]:
            nm = n(row[eh["associate"]]) if row else ""
            d = day_index(row[eh["day"]]) if row else None
            if nm and d is not None:
                exceptions.add((nm, d))
    wb.close()
    return sched, breaks, exceptions


# --------------------------------------------------------------------------- check
def check(k: Contract, sched: Dict[str, List[str]], breaks, exceptions) -> Dict[str, Any]:
    v: List[Dict[str, Any]] = list(getattr(k, "input_issues", []))
    kinds: Dict[str, List[str]] = {}
    shifts: Dict[str, List[Optional[Tuple[int, int]]]] = {}
    for nm in k.names:
        key = n(nm)
        row = sched.get(key)
        if row is None:
            v.append({"rule": "missing_associate", "who": nm})
            continue
        ks, ss = [], []
        for d, cell in enumerate(row):
            c = n(cell)
            if c == "off":
                ks.append("off"); ss.append(None)
            elif c in ("leave", "pto", "vacation"):
                ks.append("leave"); ss.append(None)
            elif shift_of(cell):
                if c not in k.library:
                    v.append({"rule": "shift_not_in_library", "who": nm, "day": DAYS[d], "value": cell})
                ks.append("shift"); ss.append(shift_of(cell))
            else:
                ks.append("blank"); ss.append(None)
                v.append({"rule": "blank_or_unknown_cell", "who": nm, "day": DAYS[d], "value": cell})
        kinds[key], shifts[key] = ks, ss
        pref = k.pref.get(key, [""] * 7)
        fix = k.fixed.get(key, [""] * 7)
        for d in range(7):
            p, f = n(pref[d]), n(fix[d])
            leave_words = ("leave", "annual leave", "vacation", "holiday", "sick", "sick leave", "pto")
            if (k.leave and p in leave_words) or f in leave_words:
                if ks[d] != "leave":
                    v.append({"rule": "leave_not_honoured", "who": nm, "day": DAYS[d], "got": row[d]})
            if (k.hard_off and p in ("off", "day off", "rest day")) or f == "off":
                if ks[d] != "off":
                    v.append({"rule": "hard_off_not_honoured", "who": nm, "day": DAYS[d], "got": row[d]})
            if shift_of(fix[d]) and n(row[d]) != f:
                v.append({"rule": "fixed_not_honoured", "who": nm, "day": DAYS[d], "want": fix[d], "got": row[d]})
        offs = [d for d in range(7) if ks[d] == "off"]
        long_mode = any(s and s[1] >= k.long_min for s in ss)
        if k.strict_off and len(offs) != (3 if long_mode else 2):
            v.append({"rule": "off_count", "who": nm, "offs": len(offs)})
        if not k.strict_off and len(offs) < 2:
            v.append({"rule": "off_minimum", "who": nm, "offs": len(offs)})
        if k.strict_off and not k.separate_off and not any((d + 1) % 7 in offs for d in offs):
            v.append({"rule": "off_not_consecutive", "who": nm})
        labels = {n(row[d]) for d in range(7) if ks[d] == "shift"}
        if len(labels) > k.max_diff:
            v.append({"rule": "max_different_shifts", "who": nm, "count": len(labels)})
        need = k.rest_h * 60
        for d in range(7):
            a, b = ss[d], ss[(d + 1) % 7]
            if a and b and 1440 + b[0] - (a[0] + a[1]) + EPS < need:
                v.append({"rule": "rest" if d < 6 else "rest_sat_to_sun_wrap", "who": nm, "from": DAYS[d]})
        prev = shift_of(k.prev_sat.get(key, ""))
        if prev and ss[0] and ss[0][0] - (prev[0] + prev[1] - 1440) + EPS < need:
            v.append({"rule": "rest_previous_saturday", "who": nm})
        if k.window_mode != "OFF":
            lang = k.lang.get(key, "")
            for d in range(7):
                if not ss[d]:
                    continue
                entries = [w for w in k.windows.get(lang, [])
                           if (k.window_mode != "MINIMUM_ROWS" or w[2])]
                todays = [w for w in entries if d in w[3]]
                if not todays:
                    continue  # no window opens this day: the engine reads that as unrestricted
                if not any((ss[d][0] - w[0]) % 1440 < ((w[1] - w[0]) % 1440 or 1440) for w in todays):
                    v.append({"rule": "language_working_window", "who": nm, "day": DAYS[d], "shift": row[d]})

    # Breaks
    bq: Dict[str, Set[int]] = defaultdict(set)
    per_cell: Dict[Tuple[str, int], List[Dict[str, Any]]] = defaultdict(list)
    for b in breaks:
        if "no break" in b["status"]:
            continue
        per_cell[(b["name"], b["day"])].append(b)
    for (key, d), rows in per_cell.items():
        s = shifts.get(key, [None] * 7)[d]
        if s is None:
            v.append({"rule": "break_on_non_shift_day", "who": key, "day": DAYS[d]})
            continue
        rels = []
        for b in rows:
            if b["start"] is None or b["start"] % 15 or b["len"] % 15:
                v.append({"rule": "break_not_on_quarter", "who": key, "day": DAYS[d]})
                continue
            rel = (b["start"] - s[0]) % 1440
            if rel + b["len"] > s[1]:
                v.append({"rule": "break_outside_shift", "who": key, "day": DAYS[d]})
            rels.append((rel, b["len"]))
            start_q = d * 96 + (s[0] + rel) // 15
            for q in range(start_q, start_q + b["len"] // 15):
                if q in bq[key]:
                    v.append({"rule": "break_overlap", "who": key, "day": DAYS[d]})
                bq[key].add(q)
        rels.sort()
        if sorted(x[1] for x in rels) != k.break_lengths:
            v.append({"rule": "break_set", "who": key, "day": DAYS[d], "got": sorted(x[1] for x in rels)})
        for (r1, l1), (r2, _l2) in zip(rels, rels[1:]):
            if r2 - (r1 + l1) < k.break_min_gap:
                v.append({"rule": "break_gap", "who": key, "day": DAYS[d], "gap": r2 - (r1 + l1)})
    for key, ss in shifts.items():
        for d in range(7):
            if ss[d] and (key, d) not in per_cell and (key, d) not in exceptions and k.break_lengths:
                v.append({"rule": "shift_without_breaks", "who": key, "day": DAYS[d]})

    # Coverage at quarter level: current-week shifts + previous-Saturday carry-in.
    work: Dict[int, List[str]] = defaultdict(list)
    carry: Dict[int, List[str]] = defaultdict(list)
    for key, ss in shifts.items():
        for d in range(7):
            if ss[d]:
                q0 = d * 96 + ss[d][0] // 15
                for q in range(q0, q0 + ss[d][1] // 15):
                    if q < 7 * 96:
                        work[q].append(key)
        prev = shift_of(k.prev_sat.get(key, ""))
        if prev:
            q0 = -96 + prev[0] // 15
            for q in range(q0, q0 + prev[1] // 15):
                if 0 <= q < 96:
                    carry[q].append(key)
    qpi = k.step // 15
    m = defaultdict(int)
    lang_gaps, zero = 0, 0
    for d in range(7):
        for i in range(k.per_day):
            r = k.req[d][i]
            if r is None or r <= 0:
                continue
            eff = 1.0 - float(k.shr[d][i] or 0.0)
            b_sum = a_sum = 0.0
            for q in range(qpi):
                slot = d * 96 + i * qpi + q
                present = work[slot] + carry[slot]
                after = [x for x in work[slot] if slot not in bq[x]] + carry[slot]
                b_sum += len(present) * eff
                a_sum += len(after) * eff
                zero += int(len(after) == 0)
                minute = i * k.step + q * 15
                for rule in k.rules:
                    if rule_in_force(rule, d, minute):
                        got = sum(1 for x in after if k.lang.get(x) in rule["eligible"])
                        lang_gaps += int(got < rule["minimum"])
            bp, ap = b_sum / qpi / r, a_sum / qpi / r
            m["active_intervals"] += 1
            for tag, pct in (("before", bp), ("after", ap)):
                m[f"{tag}_target"] += int(pct + EPS >= k.target)
                m[f"{tag}_floor"] += int(pct + EPS >= k.floor)
                m[f"{tag}100"] += int(pct + EPS >= 1.0)
                m[f"{tag}90"] += int(pct + EPS >= 0.9)
                m[f"{tag}80"] += int(pct + EPS >= 0.8)
            m["target_losses_from_breaks"] += int(bp + EPS >= k.target > ap + EPS)
            m["floor_losses_from_breaks"] += int(bp + EPS >= k.floor > ap + EPS)
    m["language_gap_count"] = lang_gaps
    m["zero_staffed_active_quarters"] = zero
    by_rule: Dict[str, int] = defaultdict(int)
    for row in v:
        by_rule[row["rule"]] += 1
    return {"metrics": dict(m), "violation_count": len(v), "violations_by_rule": dict(by_rule),
            "violations": v[:200], "language_window_mode": k.window_mode,
            "language_rules": [{**r, "days": sorted(r["days"]), "eligible": sorted(r["eligible"])} for r in k.rules]}


COMPARE = ("active_intervals", "before_target", "after_target", "before_floor", "after_floor",
           "before100", "before90", "before80", "after100", "after90", "after80",
           "target_losses_from_breaks", "floor_losses_from_breaks", "language_gap_count",
           "zero_staffed_active_quarters")


def compare(mine: Dict[str, Any], other: Dict[str, Any], aliases: Dict[str, str]) -> List[Dict[str, Any]]:
    out = []
    for key in COMPARE:
        theirs_key = aliases.get(key, key)
        if theirs_key not in other:
            continue
        if int(other[theirs_key]) != int(mine.get(key, 0)):
            out.append({"metric": key, "clean_room": int(mine.get(key, 0)), "published": int(other[theirs_key])})
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--output", required=True, type=Path)
    ap.add_argument("--audit", type=Path)
    ap.add_argument("--validation", type=Path)
    ap.add_argument("--json-out", type=Path)
    args = ap.parse_args()
    k = read_contract(args.input)
    sched, brk, exc = read_output(args.output)
    result = check(k, sched, brk, exc)
    if args.audit and args.audit.exists():
        audit = json.loads(args.audit.read_text())
        published = (audit.get("selected_candidate") or {}).get("metrics") or {}
        result["engine_mismatches"] = compare(result["metrics"], published, {"active_intervals": "active_intervals"})
    if args.validation and args.validation.exists():
        val = json.loads(args.validation.read_text())
        result["validator_mismatches"] = compare(result["metrics"], val.get("metrics") or {}, {})
        result["validator_status"] = val.get("status")
    text = json.dumps(result, indent=1, default=str)
    if args.json_out:
        args.json_out.write_text(text)
    print(json.dumps({key: result[key] for key in result if key not in ("violations", "language_rules")}, indent=1))
    return 0 if result["violation_count"] == 0 and not result.get("engine_mismatches") else 1


if __name__ == "__main__":
    raise SystemExit(main())
