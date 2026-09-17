#!/usr/bin/env python3
"""Build the purpose-built fixtures fix-plan W3 calls for.

The corpus measurement showed 4 of 12 contract dimensions have effectively no
variation, and that is *why* several findings stayed invisible:

  * 14 of 15 workbooks use floor_ratio 0.80, so `before_floor` and `before_80`
    are the same number and the ten `*_floor -> *_80` fallbacks cannot be seen.
  * 0 of 15 define a nesting group, so S6-1 and S6-2 live in code nothing runs.

Each fixture is DERIVED from a shipped workbook by changing one contract
dimension, and is written to a new file. No shipped asset is modified, and none
of these is a baseline -- they are test inputs, named so that cannot be
confused.
"""
import argparse
import importlib.util
import shutil
import sys
from pathlib import Path

import openpyxl


def instruction_row(ws, label, E):
    for r in range(1, ws.max_row + 1):
        for c in (1, 2):
            if E.norm(ws.cell(r, c).value) == E.norm(label):
                return r, c + 1
    return None, None


def set_instruction(ws, label, value, E):
    r, c = instruction_row(ws, label, E)
    if r is None:
        raise SystemExit("instruction %r not found" % label)
    ws.cell(r, c).value = value


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--engine", type=Path, required=True)
    ap.add_argument("--source", type=Path, required=True)
    ap.add_argument("--out-dir", type=Path, required=True)
    a = ap.parse_args()
    sys.path.insert(0, str(a.engine.parent))
    spec = importlib.util.spec_from_file_location("E", a.engine)
    E = importlib.util.module_from_spec(spec); sys.modules["E"] = E
    spec.loader.exec_module(E)
    a.out_dir.mkdir(parents=True, exist_ok=True)

    # ---- FX1: a floor that is not 80% ------------------------------------
    fx1 = a.out_dir / "SYNTHETIC_FIXTURE_FLOOR_NOT_80.xlsx"
    shutil.copy(a.source, fx1)
    wb = openpyxl.load_workbook(fx1)
    # 0.60 rather than a rounder 0.75: this contract's coverage is clustered
    # with nothing between 60% and 90%, so floors of 0.70-0.90 all select the
    # same 70 intervals as the 80% tier and the fixture would prove nothing.
    # At 0.60 the configured floor selects 82 where the 80% tier selects 70.
    set_instruction(wb["Instructions"], "Minimum Per Interval", 0.60, E)
    wb.save(fx1)
    p1 = E.parse_input(fx1)
    print("FX1 %s" % fx1.name)
    print("    floor_ratio=%s target_ratio=%s" % (p1.floor_ratio, p1.target_ratio))

    # ---- FX2: a nesting group whose members hold different leave ---------
    fx2 = a.out_dir / "SYNTHETIC_FIXTURE_NESTING_GROUP_MIXED_LEAVE.xlsx"
    shutil.copy(a.source, fx2)
    wb = openpyxl.load_workbook(fx2)
    set_instruction(wb["Instructions"], "Fixed Request Use", "Yes", E)

    base = E.parse_input(a.source)
    first, second = base.associates[0].name, base.associates[1].name

    fr = E._sheet_by_alias(wb, ["Fixed Request", "Fixed Requests", "Nesting", "Fixed/Nesting"])
    h = E._find_header_row(fr, ["name"])
    hdr = {E.norm(fr.cell(h, c).value): c for c in range(1, fr.max_column + 1)}
    c_active = next((c for k, c in hdr.items() if "active" in k), 1)
    c_name = next(c for k, c in hdr.items() if "name" in k)
    c_group = next(c for k, c in hdr.items() if "group" in k or "nest" in k)
    for offset, who in enumerate((first, second)):
        r = h + 1 + offset
        fr.cell(r, c_active).value = "Yes"
        fr.cell(r, c_name).value = who
        fr.cell(r, c_group).value = "GROUP_A"

    # one of the two carries approved leave, the other does not
    pref = E._sheet_by_alias(wb, ["Preference", "Prefrence", "Preferences"])
    ph = E._find_header_row(pref, ["name"], prefer_day_columns=True) \
        if "prefer_day_columns" in E._find_header_row.__code__.co_varnames \
        else E._find_header_row(pref, ["name"])
    pname = next((c for c in range(1, pref.max_column + 1)
                  if "name" in E.norm(pref.cell(ph, c).value)), 1)
    days = E._day_columns(pref, ph)
    placed = False
    for r in range(ph + 1, pref.max_row + 1):
        if E.norm(pref.cell(r, pname).value) == E.norm(first) and days:
            pref.cell(r, days[0]).value = "Leave"
            placed = True
            break
    wb.save(fx2)
    p2 = E.parse_input(fx2)
    groups = {}
    for assoc in p2.associates:
        if assoc.nesting_group:
            groups.setdefault(assoc.nesting_group, []).append(assoc.name)
    leave_holders = [assoc.name for assoc in p2.associates
                     if any(E.preference_kind(v) == "leave" for v in assoc.preferences)]
    print("FX2 %s" % fx2.name)
    print("    fixed_enabled=%s  nesting groups=%s" % (p2.fixed_enabled, groups))
    print("    leave placed=%s  associates holding leave=%s" % (placed, leave_holders))
    mixed = any(len({tuple(1 if E.preference_kind(v) == "leave" else 0
                           for v in next(x for x in p2.associates if x.name == n).preferences)
                     for n in names}) > 1 for names in groups.values())
    print("    group members hold DIFFERENT leave: %s" % mixed)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
