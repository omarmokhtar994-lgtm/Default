# © 2026 Omar Mokhtar. All rights reserved.
"""Phase L task 4: the clean input workbooks the website offers on its home page.

Owner, 2026-10-08: "we need source for clean input sheet to be download from
the homepage", then "C" (both a blank and an example). The example is the
synthetic SYNTH_M2 workbook (made-up names) rebuilt by the friendly template
builder, whose own check proves the engine reads it unchanged; the blank is
that example with people, demand and per-person rows cleared."""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(ROOT / "engine" / "_tools"))
sys.path.insert(0, str(ROOT / "tools"))
import build_input_template as B  # noqa: E402
import build_web_workbooks as W  # noqa: E402

BOOKS = ROOT / "webapp" / "workbooks"
EXAMPLE, BLANK = BOOKS / W.EXAMPLE_NAME, BOOKS / W.BLANK_NAME
SOURCE = ROOT / W.SOURCE


class TheWebWorkbooks(unittest.TestCase):
    def test_example_reads_like_its_synthetic_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            fresh = Path(tmp) / "fresh.xlsx"
            self.assertEqual(B.main(["build", str(SOURCE), str(fresh)]), 0)  # the builder's own engine check
            self.assertEqual(B.engine_reading(EXAMPLE), B.engine_reading(fresh))

    def test_example_passes_the_input_check(self):
        proc = subprocess.run([sys.executable, str(ROOT / "tools" / "check_input_workbook.py"), str(EXAMPLE)],
                              capture_output=True, text=True, timeout=600)
        self.assertEqual(proc.returncode, 0, proc.stdout[-2000:] + proc.stderr[-2000:])

    def test_example_names_are_made_up(self):
        names = [c.value for c in load_workbook(EXAMPLE)["Schedule"]["D"][2:] if c.value]
        self.assertTrue(names)
        self.assertTrue(all(str(n).startswith("Synth Agent") for n in names), names[:5])

    def test_blank_has_every_tab_and_no_people_or_demand(self):
        example, blank = load_workbook(EXAMPLE), load_workbook(BLANK)
        # Re-pinned in Phase Z (2026-10-10): the owner approved the channel tabs on the Home workbooks (sample 03,
        # "Panel + tabs"). A channel grid tab carries its workbook's own interval in its name (Phase V's approved
        # layout; a grid at another interval is refused once filled in): the Blank is set to 30 minutes and the
        # Example is 60, so those three names differ. Every other tab is the same, in the same order.
        grids = {f"{c} 60 Min": f"{c} 30 Min" for c in ("Chat", "Phone", "Email")}
        self.assertEqual(blank.sheetnames, [grids.get(n, n) for n in example.sheetnames])
        self.assertTrue(set(grids) <= set(example.sheetnames))
        for sheet, (first_row, first_col) in W.CLEARED.items():
            ws = blank[sheet]
            filled = [(c.coordinate, c.value) for row in ws.iter_rows(min_row=first_row, min_col=first_col)
                      for c in row if c.value not in (None, "")]
            self.assertEqual(filled, [], sheet)
        times = [blank["FT Wise 30 Min"].cell(r, 1).value for r in (3, 4)]
        self.assertEqual(times, ["00:00", "00:30"])  # the interval labels stay
        self.assertEqual([[c.value for c in r] for r in blank["Shift Library"].iter_rows()],
                         [[c.value for c in r] for r in example["Shift Library"].iter_rows()])
        self.assertEqual(len(blank["Schedule"].data_validations.dataValidation),
                         len(example["Schedule"].data_validations.dataValidation))

    def test_blank_settings_are_cleared_or_set_for_half_hours(self):
        ws = load_workbook(BLANK)["Instructions"]
        value = {r[1].value: r[2].value for r in ws.iter_rows() if r[1].value}
        self.assertIn(value["Program Name"], (None, ""))
        self.assertIn(value["Count of Associates"], (None, ""))
        self.assertEqual((str(value["Interval Minutes"]), value["Requirements Source"], value["Shrinkage Source"]),
                         ("30", "FT Wise 30 Min", "Shrinkage 30 Min"))


if __name__ == "__main__":
    unittest.main()
