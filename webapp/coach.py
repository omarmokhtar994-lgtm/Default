# © 2026 Omar Mokhtar. All rights reserved.
"""The shrinkage coach: what actually happened against the input's Shrinkage tab.

Per weekday and interval, over a period: the actual shrinkage (1 - people on the
floor / people paid on shift) averaged over the days someone recorded attendance,
activities or break moves for the program (a day nobody recorded would only echo
the plan). The suggestion is that average where at least ``min_days`` days were
recorded, else the input's own value (kept, and said so). The corrected tab keeps
the input's layout (title row, then Interval and Sun..Sat), ready to paste.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from .adherence import interval_shrinkage
from .versions import DAYS

MIN_DAYS = 3
EGYPT = timezone(timedelta(hours=3))


def actual_shrinkage(days, program: str, start: date, end: date, min_days: int = MIN_DAYS,
                     measure: str = "interval") -> Dict[str, Any]:
    seen: Dict[int, Dict[int, list]] = {d: {} for d in range(7)}
    planned: Dict[int, Dict[int, float]] = {d: {} for d in range(7)}
    recorded, interval = [], None
    day = start
    while day <= end:
        iso = day.isoformat()
        try:
            page = days.page(program, day, measure)
        except ValueError:
            page = None
        if page is not None:
            interval = page["view"]["interval"]
            weekday = page["day"]
            for minute, value in page["inputs"].get("planned_shrinkage", {}).get(weekday, {}).items():
                planned[weekday][minute] = value
            if days.store.list_day_log(program, iso):
                recorded.append(iso)
                for row in interval_shrinkage(page["view"], page["inputs"], weekday):
                    seen[weekday].setdefault(row["t"], []).append(row["actual"])
        day += timedelta(days=1)
    weekdays: Dict[int, Dict[int, Dict[str, Any]]] = {}
    for weekday in range(7):
        weekdays[weekday] = {}
        for minute in sorted(set(planned[weekday]) | set(seen[weekday])):
            values = seen[weekday].get(minute, [])
            actual = round(sum(values) / len(values), 3) if values else None
            kept = len(values) < min_days
            weekdays[weekday][minute] = {"actual": actual, "input": planned[weekday].get(minute, 0.0),
                                         "days": len(values), "kept_input": kept,
                                         "suggested": planned[weekday].get(minute, 0.0) if kept else actual}
    return {"program": program, "start": start, "end": end, "interval": interval or 30, "min_days": min_days,
            "recorded_days": recorded, "weekdays": weekdays}


def corrected_tab(found: Dict[str, Any], program: str) -> io.BytesIO:
    step = found["interval"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"Shrinkage {step} Min"
    ws["A1"] = f"Shrinkage {step} Min — from what happened ({program}, {found['start']:%d %b} to {found['end']:%d %b %Y})"
    ws["A1"].font = Font(bold=True)
    ws.append(["Interval"] + DAYS)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for minute in range(0, 1440, step):
        row = [f"{minute // 60:02d}:{minute % 60:02d}"]
        for weekday in range(7):
            entry = found["weekdays"][weekday].get(minute)
            row.append(entry["suggested"] if entry else 0.0)
        ws.append(row)
        for weekday in range(7):
            entry = found["weekdays"][weekday].get(minute)
            cell = ws.cell(ws.max_row, 2 + weekday)
            cell.number_format = "0.0%"
            if entry and entry["kept_input"]:
                cell.comment = Comment(f"Kept the input's value: {entry['days']} recorded day(s), "
                                       f"{found['min_days']} needed.", "Team Scheduler")
    how = wb.create_sheet("How this was made")
    for line in (f"Program: {program}", f"Period: {found['start']:%d %b %Y} to {found['end']:%d %b %Y}",
                 f"Days recorded on the day page: {len(found['recorded_days'])}",
                 "Actual shrinkage per interval = 1 - people on the floor / people paid on shift (absences, breaks "
                 "as taken, aux and VTO counted; billable aux counts as on the floor).",
                 f"A value replaces the input's where at least {found['min_days']} days of that weekday were "
                 "recorded; elsewhere the input's value is kept (the cell has a note).",
                 f"Made {datetime.now(EGYPT):%Y-%m-%d %H:%M} Egypt time. Paste the first tab over the input's "
                 f"Shrinkage {step} Min tab."):
        how.append([line])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
