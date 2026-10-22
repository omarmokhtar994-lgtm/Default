# © 2026 Omar Mokhtar. All rights reserved.
"""A look at every uploaded workbook before the site keeps it (Phase AV; owner, 2026-10-22).

Team Scheduler workbooks hold schedules, demand and shifts for people known by made-up or internal names. A file that is
not one of them (other tabs than the template's), or that carries what looks like card numbers or street
addresses, bank account numbers, macros, embedded objects or links to other files, is outsourced or unwanted data: it is
refused and deleted so nothing leaks through the website. Phone numbers are not looked for (hours and counts look like
them), and email addresses are not either (real rosters list work emails). What was found is named by kind and place, never by value, in the message and in the record.

An unreadable file is left to the checks that follow (they say what is wrong), so nothing that worked before breaks."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

MAX_CELLS = 400_000  # a scan stops here; the workbook check refuses an over-large file anyway
REQUIRED = ("Instructions", "Schedule", "Shift Library")  # in every real input workbook and in the template; the requirement
# tab is not asked for here, because the Instructions tab may name any tab as the requirement (the later check says so)
ACTIVE_PARTS = ("vbaproject", "xl/externallinks/", "xl/embeddings/", "xl/activex/", "xl/ctrlprops/")
CARD = re.compile(r"(?<![\w/\\.:-])(?:\d{4}[ -]?){3}\d{1,7}(?![\w/\\.:-])")  # 4-4-4-4 style, not inside a path or a code
STREET = re.compile(r"\b\d{1,5}\s+(?:[A-Z][a-z]+\s+){1,3}(?:Street|St|Road|Rd|Avenue|Ave|Boulevard|Blvd|Lane|Ln|Drive|Dr|Court|Ct|"
                    r"Square|Sq)\b\.?")
IBAN = re.compile(r"\b[A-Z]{2}\d{2}[A-Z0-9]{11,30}\b")
RISKY_FORMULA = re.compile(r"(?i)(?:\bcmd\b\s*\||\bdde\b|\bwebservice\s*\(|\bhyperlink\s*\(|\bimportxml\s*\(|\\\\\\\\|https?://|\[\d+\]!)")

LABELS = {"card": "a card number", "email": "an email address", "address": "a street address",
          "iban": "a bank account number", "formula": "a formula that reaches outside the file"}


def luhn(digits: str) -> bool:
    total, flip = 0, False
    for ch in reversed(digits):
        n = int(ch)
        if flip:
            n = n * 2 - 9 if n * 2 > 9 else n * 2
        total += n
        flip = not flip
    return total % 10 == 0


def kind_of(value: object) -> Optional[str]:
    """What sensitive thing a cell holds, or None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        return None  # a measurement, never a card
    text = str(value)
    if text.startswith("="):
        return "formula" if RISKY_FORMULA.search(text) else None
    if len(text) < 8:
        return None
    for m in CARD.finditer(text):
        digits = re.sub(r"\D", "", m.group())
        if 13 <= len(digits) <= 19 and luhn(digits) and len(set(digits)) > 1:
            return "card"
    if STREET.search(text):
        return "address"
    if IBAN.search(text.replace(" ", "")) and re.search(r"[A-Z]{2}\d{2}", text.replace(" ", "")) and len(text.replace(" ", "")) <= 40:
        return "iban"
    return None


def look(path: Path, structure: bool = True, personal: bool = True) -> Optional[Tuple[str, str]]:
    """None when the workbook may be kept, else ``(kind, message)``: kind is ``structure``, ``active`` or ``personal``.
    ``structure`` asks for the input workbook's tabs; ``personal`` looks for card numbers, addresses and the like."""
    try:
        with zipfile.ZipFile(path) as z:
            parts = [n.lower() for n in z.namelist()]
    except (zipfile.BadZipFile, OSError):
        return None  # not a workbook at all: the checks that follow say so
    bad = sorted({a.strip("/").split("/")[-1] or a for n in parts for a in ACTIVE_PARTS if a in n})
    if bad:
        return ("active", "This file carries macros, embedded objects or links to other files, which a Team Scheduler "
                "workbook never has. It could be outsourced or unsafe data, so it was refused and deleted from the "
                "server to avoid any leak.")
    handle = open(path, "rb")  # a file object, so any name works (an upload is saved as .upload)
    try:
        wb = load_workbook(handle, read_only=True, data_only=False)
    except Exception:  # unreadable by the library: left to the checks that follow
        handle.close()
        return None
    try:
        names = wb.sheetnames
        if structure:
            missing = [t for t in REQUIRED if t not in names]
            if missing:
                return ("structure", "This file does not match the Team Scheduler input workbook (missing tab"
                        + ("s" if len(missing) > 1 else "") + ": " + ", ".join(missing) + "). It looks like outsourced data "
                        "or a different kind of file, so it was refused and deleted from the server to avoid any data "
                        "leakage. Start from the workbook templates on the Home page.")
        if personal:
            found: Dict[str, List[str]] = {}
            seen = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for value in row:
                        seen += 1
                        k = kind_of(value)
                        if k:
                            found.setdefault(k, [])
                            if ws.title not in found[k]:
                                found[k].append(ws.title)
                    if seen > MAX_CELLS:
                        break
                if seen > MAX_CELLS:
                    break
            if found:
                what = "; ".join(f"{LABELS[k]} (tab {', '.join(tabs[:3])}{'…' if len(tabs) > 3 else ''})" for k, tabs in found.items())
                return ("personal", f"This file contains {what}. Team Scheduler workbooks hold no such data, so it could be "
                        "outsourced or personal data. It was refused and deleted from the server to avoid any leak; "
                        "remove those cells and upload again.")
    finally:
        wb.close()
        handle.close()
    return None
