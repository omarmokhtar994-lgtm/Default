# © 2026 Omar Mokhtar. All rights reserved.
"""Phase V task 7: channels on the RTA page (owner, 2026-10-09: "this some how needs to reflect in RTA view as well
with warnings when no channel coverage due to extra coaching etc considering the language as well").

Coverage counts the people really on each channel: breaks (as moved), absences, late, coaching, meetings and other
aux are taken out, and in all-channels times a person covers every channel they work. A warning names what took
the channel away (the person, the kind and, for an aux, its With and Why) and offers fixes: a channel change for
someone who can take it, moving the aux, moving a break. Booking an aux says the channel drop before it is kept.

The made-up Wednesday: Associate 001 (Arabic, 12:00 to 21:00) is the only one on Chat, Associate 031 (12:00 to
21:00, a break at 14:00) the only one on Phone, everyone else on Email. Chat and Phone each need one person from
14:00 to 16:00 (Chat with an Arabic speaker) and from 19:00 to 21:00, which is an all-channels time."""
import json
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from webapp.attendance import DayBook
from webapp.channel_people import ChannelPeople
from webapp.programs import ProgramBook
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_channels import add_channel_tabs
from webapp.tests.test_ready import make_ready
from webapp.tests.test_schedules import REPO
from webapp.versions import read_week, shift_span, write_breaks, write_channels

WED = date(2026, 10, 14)
ARABIC_SPARE = "Associate 013"  # on shift 11:00 to 20:00 on Wednesday, on Email


def need(windows):
    def grid(d, t):
        return 1 if d == 3 and any(lo <= t < hi for lo, hi in windows) else 0
    return grid


def set_languages(path, names):
    wb = load_workbook(path)
    ws = wb["Schedule"]
    head = next(r for r in range(1, 10) if "sun" in [str(c.value or "").strip().lower() for c in ws[r]])
    cols = {str(c.value or "").strip().lower(): c.column for c in ws[head] if c.value}
    for r in range(head + 1, ws.max_row + 1):
        if ws.cell(r, cols["sf name"]).value in names:
            ws.cell(r, cols["language"]).value = "Arabic"
    wb.save(path)


def channel_day(target, plan=True):
    """The ready workbook with channel tabs and Wednesday's channel plan and one break."""
    ready = make_ready(target.with_name("ready.xlsx"))
    set_languages(ready, {"Associate 001", ARABIC_SPARE})
    windows = [(14 * 60, 16 * 60), (19 * 60, 21 * 60)]
    book = add_channel_tabs(target.with_name("tabs.xlsx"), src=ready, phone=need(windows), chat=need(windows),
                            email_tab=False, hours=None, ft=None,
                            settings=[("Minimum block minutes", 60)],
                            languages=[("Chat", "Arabic", 1, "14:00", "16:00", "Wed", "Yes")],
                            blended=[("Wed", "19:00", "21:00", "Yes", "")])
    if not plan:
        shutil.copyfile(book, target)
        return target
    week = read_week(book)
    rows = {}
    for a in week["associates"]:
        span = shift_span(a["days"][3])
        if not span:
            continue
        if a["name"] == "Associate 001":
            blocks = [(span[0], span[1], "C")]
        elif a["name"] == "Associate 031":
            blocks = [(span[0], 840, "P"), (855, span[1], "P")]
        else:
            blocks = [(span[0], span[1], "E")]
        rows[(a["name"], "Wed")] = (a["days"][3], blocks)
    middle = target.with_name("with_breaks.xlsx")
    write_breaks(book, middle, {("Associate 031", "Wed"): ("12:00 - 21:00", [("Break 1", 840, 15)])})
    write_channels(middle, target, rows)
    return target


class _Day(unittest.TestCase):
    PLAN = True

    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.omar = store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        programs = ProgramBook(store)
        cls.program_id = programs.add_program("SAKS")
        cls.key = programs.add_lob(cls.program_id, "NMG Tier 2")
        store.add_run("cccccccccccc", cls.omar, "ready.xlsx", "READY", "DONE", program=cls.key, week_start="2026-10-11")
        upload = channel_day(cls.base / "upload.xlsx", plan=cls.PLAN)
        ScheduleBook(store, cls.base, REPO).ensure_ready(store.get_run("cccccccccccc"), upload)
        everyone = [a["name"] for a in read_week(upload)["associates"]]
        ChannelPeople(store).save(cls.program_id, {n: "E" for n in everyone if n not in
                                                    ("Associate 001", "Associate 031", ARABIC_SPARE)}, cls.omar)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))

    def view(self):
        return self.days.page(self.key, WED)["view"]

    def cell(self, row_name, t):
        rows = self.view()["channels"]["rows"]
        return next(c for r in rows if r["name"] == row_name for c in r["cells"] if c["t"] == t)

    def warnings(self):
        return self.view()["channels"]["warnings"]

    def heads(self):
        return [w["text"] for w in self.warnings()]


class TheChannelsOnTheDay(_Day):
    def test_as_planned_every_channel_and_language_is_covered(self):
        self.assertEqual((self.cell("Phone", 14 * 60)["low"], self.cell("Phone", 14 * 60)["need"]), (0, 1))
        self.assertEqual(self.cell("Chat", 15 * 60)["low"], 1)
        # 14:00 to 14:15 Associate 031 is on the break the schedule gave them: Phone has nobody then, as planned
        self.assertEqual(self.heads(), [])  # as planned is not a warning: the plan's own gap shows in the table
        self.assertEqual(self.view()["channels"]["now"], {})  # "right now" figures only on the day itself

    def test_coaching_on_the_only_arabic_chat_person_names_it_with_who_and_why_and_offers_fixes(self):
        self.days.add_activity(self.key, WED, "Associate 001", "Coaching", "15:00", "15:30", self.omar,
                               with_whom="Associate 021", with_dept="Team Leads", why="Quality follow-up")
        found = self.warnings()
        self.assertEqual([w["text"] for w in found],
                         ["15:00 to 16:00: nobody on Chat, and no Arabic speaker on Chat"])
        self.assertEqual(found[0]["causes"], ["Associate 001 is booked for Coaching, with Associate 021 (Team Leads): "
                                              "Quality follow-up"])
        fixes = [f["text"] for f in found[0]["fixes"]]
        self.assertIn(f"{ARABIC_SPARE}: Email → Chat from 15:00 to 16:00", fixes)
        self.assertIn("Move Associate 001's Coaching to 16:00", fixes)

    def test_the_channel_change_fix_clears_the_warning_and_is_logged(self):
        self.days.add_activity(self.key, WED, "Associate 001", "Coaching", "15:00", "15:30", self.omar)
        swap = next(f for f in self.warnings()[0]["fixes"] if f["kind"] == "swap")
        self.days.channel_move(self.key, WED, swap["who"], swap["start"], swap["end"], swap["channel"], self.omar)
        self.assertEqual(self.heads(), [])
        said = [r["what"] for r in self.store.list_day_log(self.key, WED.isoformat())]
        self.assertIn("Chat from 15:00 to 16:00 (was Email)", said)
        move = self.store.list_channel_moves(self.key, [WED.isoformat()])[0]
        self.days.cancel_channel_move(self.key, WED, move["id"], self.omar)
        self.assertEqual(len(self.heads()), 1)

    def test_moving_the_coaching_clears_the_warning(self):
        made = self.days.add_activity(self.key, WED, "Associate 001", "Coaching", "15:00", "15:30", self.omar,
                                      why="Quality follow-up")
        fix = next(f for f in self.warnings()[0]["fixes"] if f["kind"] == "activity")
        self.days.move_activity(self.key, WED, fix["id"], fix["start"], self.omar)
        self.assertEqual(self.heads(), [])
        moved = [a for a in self.store.list_activities(self.key, [WED.isoformat()]) if a["kind"] == "Coaching"]
        self.assertEqual([(a["start"], a["end_min"], a["why"]) for a in moved], [(960, 990, "Quality follow-up")])
        self.assertNotEqual(made, moved[0]["id"])  # re-made at the new time, with the same who and why

    def test_an_absence_is_named(self):
        self.days.set_status(self.key, WED, "Associate 031", "Sick", self.omar)
        found = self.warnings()
        self.assertEqual([w["text"] for w in found], ["14:00 to 16:00: nobody on Phone"])
        self.assertEqual(found[0]["causes"], ["Associate 031 is off (sick)"])

    def test_all_channels_time_with_a_break_and_a_meeting_leaves_nobody(self):
        self.days.add_activity(self.key, WED, "Associate 001", "Meeting", "20:00", "20:30", self.omar,
                               with_whom="Associate 030", with_dept="Workforce", why="Roster review")
        self.days.add_activity(self.key, WED, "Associate 031", "Break", "20:00", "20:15", self.omar)
        found = self.warnings()
        self.assertEqual([w["text"] for w in found], ["20:00 to 21:00 (all channels): nobody on Phone or Chat"])
        self.assertEqual(sorted(found[0]["causes"]), [
            "Associate 001 is booked for Meeting, with Associate 030 (Workforce): Roster review",
            "Associate 031 is on Break (added on the day)"])

    def test_a_moved_break_gives_its_time_to_the_block_before(self):
        self.days.move_break(self.key, WED, "Associate 031", 0, "16:00", self.omar)
        self.assertEqual(self.cell("Phone", 14 * 60)["low"], 1)  # 14:00 to 14:15 now on Phone, the block before
        self.assertEqual(self.heads(), [])

    def test_booking_says_the_channel_drop_and_times_that_keep_coverage(self):
        said = self.days.preview_item(self.key, WED, "Associate 001", "Coaching", "15:00", 30)
        self.assertEqual(said["level"], "bad")
        self.assertEqual(said["channels"], ["15:00 to 16:00: nobody on Chat, and no Arabic speaker on Chat"])
        self.assertEqual(said["better"][0], "16:00")


class TheChannelsTab(unittest.TestCase):
    """The RTA page's Channels tab: warnings with causes and fixes; a fix applied from the page."""

    @classmethod
    def setUpClass(cls):
        import io
        from webapp.tests.test_runs import make_app, sign_in, token, wait
        cls.dir = Path(tempfile.mkdtemp())
        upload = channel_day(cls.dir / "upload.xlsx")
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.omar = cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        programs = ProgramBook(cls.store)
        program_id = programs.add_program("SAKS")
        cls.key = programs.add_lob(program_id, "NMG Tier 2")
        cls.token = staticmethod(token)
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")
        cls.admin.post("/runs", data={"csrf_token": token(cls.admin), "kind": "ready", "mode": "QUICK",
                                      "program": cls.key, "week_start": "2026-10-11",
                                      "workbook": (io.BytesIO(upload.read_bytes()), "week.xlsx")},
                       content_type="multipart/form-data")
        wait(cls.store, cls.store.list_runs()[0]["id"], statuses=("DONE", "REJECTED", "FAILED"))
        everyone = [a["name"] for a in read_week(upload)["associates"]]
        ChannelPeople(cls.store).save(program_id, {n: "E" for n in everyone if n not in
                                                   ("Associate 001", "Associate 031", ARABIC_SPARE)}, cls.omar)
        cls.days = cls.app.extensions["days"]

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def page(self):
        import html
        url = f"/day?program={self.key.replace(' ', '+')}&date={WED.isoformat()}&view=channels"
        return html.unescape(self.admin.get(url).get_data(as_text=True))

    def test_the_tab_lists_the_warning_its_cause_and_fixes_and_a_fix_applies(self):
        self.days.add_activity(self.key, WED, "Associate 001", "Coaching", "15:00", "15:30", self.omar,
                               why="Quality follow-up")
        page = self.page()
        for words in ("1 channel warning", "15:00 to 16:00: nobody on Chat, and no Arabic speaker on Chat",
                      "Associate 001 is booked for Coaching: Quality follow-up",
                      f"{ARABIC_SPARE}: Email → Chat from 15:00 to 16:00", "Channels by interval"):
            self.assertIn(words, page)
        got = self.admin.post("/day/channel", data={"csrf_token": self.token(self.admin), "program": self.key,
                                                    "date": WED.isoformat(), "action": "swap",
                                                    "associate": ARABIC_SPARE, "from": "15:00", "to": "16:00",
                                                    "channel": "C", "view": "channels"})
        self.assertEqual(got.status_code, 303)
        page = self.page()
        self.assertIn(f"{ARABIC_SPARE}: Chat from 15:00 to 16:00.", page)
        self.assertNotIn("nobody on Chat", page)
        self.assertIn("Take back", page)


class TheDayWithoutChannels(_Day):
    PLAN = False

    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.omar = store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        programs = ProgramBook(store)
        cls.program_id = programs.add_program("SAKS")
        cls.key = programs.add_lob(cls.program_id, "NMG Tier 2")
        store.add_run("pppppppppppp", cls.omar, "ready.xlsx", "READY", "DONE", program=cls.key, week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure_ready(store.get_run("pppppppppppp"),
                                                         make_ready(cls.base / "upload.xlsx"))

    def test_the_day_and_the_booking_preview_are_as_before(self):
        self.assertIsNone(self.view()["channels"])
        said = self.days.preview_item(self.key, WED, "Associate 001", "Coaching", "15:00", 30)
        self.assertEqual(set(said), {"text", "level"})


if __name__ == "__main__":
    unittest.main()
