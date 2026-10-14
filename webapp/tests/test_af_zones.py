# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AF, time zones (A5, D5) and the operating day (D1, A4).

Owner (2026-10-14): times are optional per page, not fixed for a program. Every page shows Egypt time with a drop-down
to see another zone (SAKS NMG Tier 1 in EST, for example), and a day in the picked zone is that zone's day. At upload
the schedule's own zone is picked next to the program, the run mode and the week start, and it can be changed later.
The operating day runs from the program's start (16:00) to the next day's, as one day on the RTA.
Named zones keep summer time right: the offset is taken for the date shown, never a fixed number.
"""
import io
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from webapp.clock import HOME, ZONES, offset, pick_zone, shift, zone_label

WED = date(2026, 10, 14)
WINTER = date(2026, 11, 10)


class TheClock(unittest.TestCase):
    def test_zones(self):
        self.assertEqual(ZONES[0][0], HOME)
        self.assertEqual(HOME, "Africa/Cairo")
        self.assertIn("America/New_York", dict(ZONES))

    def test_offsets_follow_the_date(self):
        self.assertEqual(offset("Africa/Cairo", WED), 180)
        self.assertEqual(offset("Africa/Cairo", WINTER), 120)
        self.assertEqual(offset("America/New_York", WED), -240)
        self.assertEqual(offset("America/New_York", WINTER), -300)

    def test_shift_between_zones(self):
        self.assertEqual(shift(WED, "Africa/Cairo", "America/New_York"), -420)
        self.assertEqual(shift(WINTER, "Africa/Cairo", "America/New_York"), -420)
        self.assertEqual(shift(date(2026, 10, 30), "Africa/Cairo", "America/New_York"), -360)  # Cairo moved first
        self.assertEqual(shift(WED, "Africa/Cairo", "Africa/Cairo"), 0)

    def test_labels_and_picks(self):
        self.assertEqual(zone_label("America/New_York", WED), "New York (US Eastern), UTC-4")
        self.assertEqual(zone_label("Africa/Cairo", WINTER), "Cairo (Egypt), UTC+2")
        self.assertEqual(pick_zone("Mars/Base"), HOME)
        self.assertEqual(pick_zone(""), HOME)
        self.assertEqual(pick_zone("Europe/London"), "Europe/London")


class AtUpload(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from webapp.programs import ProgramBook
        from webapp.tests.test_channel_day import REPO, channel_day
        from webapp.tests.test_runs import make_app, sign_in, token
        cls.dir = Path(tempfile.mkdtemp())
        cls.upload = channel_day(cls.dir / "upload.xlsx")
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        cls.store.add_user("plan", "Planner", "Planner-pass-123", must_change=False)
        programs = ProgramBook(cls.store)
        cls.key = programs.add_lob(programs.add_program("SAKS"), "NMG Tier 2")
        cls.token = staticmethod(token)
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")
        cls.planner = sign_in(cls.app, "plan", "Planner-pass-123")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def send(self, **extra):
        before = {r["id"] for r in self.store.list_runs()}
        got = self.admin.post("/runs", data={"csrf_token": self.token(self.admin), "kind": "ready", "mode": "QUICK",
                                             "program": self.key, "week_start": "2026-10-11",
                                             "workbook": (io.BytesIO(self.upload.read_bytes()), "week.xlsx"), **extra},
                              content_type="multipart/form-data")
        made = [r for r in self.store.list_runs() if r["id"] not in before]
        return got, made

    def test_the_form_offers_the_zones(self):
        import html
        page = html.unescape(self.admin.get("/").get_data(as_text=True))
        self.assertIn('name="tz"', page)
        self.assertIn("The schedule's times are in", page)
        self.assertIn('<option value="Africa/Cairo" selected>', page)

    def test_the_zone_is_kept_with_the_run(self):
        _, made = self.send(tz="America/New_York")
        self.assertEqual(made[0]["tz"], "America/New_York")
        import html
        page = html.unescape(self.admin.get(f"/runs/{made[0]['id']}").get_data(as_text=True))
        self.assertIn("The schedule's times are in", page)
        self.assertIn('<option value="America/New_York" selected>', page)

    def test_a_zone_not_on_the_list_is_refused(self):
        got, made = self.send(tz="Mars/Base")
        self.assertEqual(made, [])
        import html
        self.assertIn("Pick the schedule's time zone from the list.",
                      html.unescape(self.admin.get(got.headers["Location"]).get_data(as_text=True)))

    def test_changed_later_by_the_uploader_or_an_admin(self):
        _, made = self.send(tz="Africa/Cairo")
        run_id = made[0]["id"]
        got = self.planner.post(f"/runs/{run_id}/zone", data={"csrf_token": self.token(self.planner),
                                                              "tz": "Europe/London"})
        self.assertEqual(got.status_code, 403)
        got = self.admin.post(f"/runs/{run_id}/zone", data={"csrf_token": self.token(self.admin),
                                                            "tz": "Europe/London"})
        self.assertEqual(got.status_code, 303)
        self.assertEqual(self.store.get_run(run_id)["tz"], "Europe/London")


class TheWeekGrid(unittest.TestCase):
    """The Week page's grid in another zone: each row moves by that date's shift; what leaves the week is dropped."""

    def test_cairo_rows_read_in_new_york(self):
        from webapp.week import in_zone
        rows = [[0, "00:00", 1, 1, 1, 1, 1, 0], [0, "07:00", 2, 2, 2, 2, 2, 0], [1, "06:00", 3, 3, 3, 3, 3, 0],
                [6, "23:00", 4, 4, 4, 4, 4, 0]]
        moved = in_zone(rows, date(2026, 10, 11), "Africa/Cairo", "America/New_York")
        self.assertEqual([(r[0], r[1], r[2]) for r in moved], [(0, "00:00", 2), (0, "23:00", 3), (6, "16:00", 4)])

    def test_the_same_zone_changes_nothing(self):
        from webapp.week import in_zone
        rows = [[0, "00:00", 1, 1, 1, 1, 1, 0]]
        self.assertEqual(in_zone(rows, date(2026, 10, 11), "Africa/Cairo", "Africa/Cairo"), rows)


class TheStitchedDay(unittest.TestCase):
    """One day-long window stitched from the schedule's days: the day in another zone, or the operating day."""

    @staticmethod
    def day(need, lanes=()):
        return {"interval": 60, "lanes": list(lanes),
                "cells": [{"t": t, "required": need, "now": need, "need": need, "pm": 0.0, "cls": "ok"}
                          for t in range(0, 1440, 60)]}

    def test_cells_come_from_both_days(self):
        from webapp.day import stitch
        found = stitch({0: self.day(1.0), 1: self.day(2.0)}, 16 * 60, 0)
        self.assertEqual(found["cells"][0]["label"], "16:00")
        self.assertEqual(found["cells"][0]["required"], 1.0)
        self.assertEqual(found["cells"][8]["label"], "00:00")
        self.assertEqual(found["cells"][8]["required"], 2.0)  # the next day's midnight

    def test_new_york_labels(self):
        from webapp.day import stitch
        found = stitch({0: self.day(1.0), 1: self.day(2.0)}, 7 * 60, -420)  # New York's day from Cairo's 07:00
        self.assertEqual([c["label"] for c in found["cells"][:2]], ["00:00", "01:00"])
        self.assertEqual(found["cells"][16]["required"], 1.0)  # Cairo's 23:00
        self.assertEqual(found["cells"][17]["required"], 2.0)  # Cairo's next 00:00

    def test_shifts_that_cross_the_window(self):
        from webapp.day import stitch
        evening = {"name": "Associate 07", "language": "English", "segments": [
            {"offset": 0, "start": 22 * 60, "end": 31 * 60, "status": "Present", "label": "22:00 - 07:00",
             "breaks": [{"start": 26 * 60, "minutes": 30, "kind": "Lunch"}]}]}
        found = stitch({0: self.day(1.0, [evening]), 1: self.day(1.0)}, 16 * 60, 0)
        lane = found["lanes"][0]
        self.assertEqual((lane["start"], lane["end"]), (6 * 60, 15 * 60))
        self.assertEqual(lane["label"], "22:00 to 07:00")
        self.assertEqual(lane["breaks"][0]["start"], 10 * 60)

    def test_a_missing_day_shows_no_figures(self):
        from webapp.day import stitch
        found = stitch({0: self.day(1.0), 1: None}, 16 * 60, 0)
        self.assertIsNone(found["cells"][10]["required"])


class TheStitchedRows(unittest.TestCase):
    """Phase AG (T8, left from Phase AF): the day in another zone has the language and channel rows too."""

    @staticmethod
    def day(gap_at, channel_cls):
        view = TheStitchedDay.day(1.0)
        view["languages"] = [{"name": "French", "cells": [{"t": t, "cls": "bad" if t == gap_at else "lok"}
                                                         for t in range(0, 1440, 60)]}]
        view["channels"] = {"rows": [{"name": "Chat", "letter": "C", "cells": [
            {"t": t, "need": 1.0, "have": 1.0, "cls": channel_cls} for t in range(0, 1440, 60)]}]}
        return view

    def test_languages_from_both_days(self):
        from webapp.day import stitch
        found = stitch({0: self.day(17 * 60, "ok"), 1: self.day(1 * 60, "bad")}, 16 * 60, 0)
        french = found["languages"][0]
        self.assertEqual(french["name"], "French")
        self.assertEqual([c["cls"] for c in french["cells"][:3]], ["lok", "bad", "lok"])  # 16:00, 17:00, 18:00
        self.assertEqual(french["cells"][9]["cls"], "bad")  # the next day's 01:00
        self.assertEqual(french["cells"][9]["label"], "01:00")

    def test_channels_from_both_days(self):
        from webapp.day import stitch
        found = stitch({0: self.day(0, "ok"), 1: self.day(0, "bad")}, 16 * 60, 0)
        chat = found["channels"][0]
        self.assertEqual((chat["name"], chat["letter"]), ("Chat", "C"))
        self.assertEqual([c["cls"] for c in chat["cells"][7:9]], ["ok", "bad"])  # 23:00 today, 00:00 the next day

    def test_a_day_without_them(self):
        from webapp.day import stitch
        found = stitch({0: TheStitchedDay.day(1.0), 1: None}, 16 * 60, 0)
        self.assertEqual((found["languages"], found["channels"]), ([], []))


class OnThePages(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from webapp.tests import test_channel_day
        test_channel_day.TheChannelsTab.setUpClass.__func__(cls)
        from webapp.tests.test_runs import sign_in
        cls.store.add_user("plan", "Planner", "Planner-pass-123", must_change=False)
        cls.planner = sign_in(cls.app, "plan", "Planner-pass-123")

    @classmethod
    def tearDownClass(cls):
        from webapp.tests import test_channel_day
        test_channel_day.TheChannelsTab.tearDownClass.__func__(cls)

    def get(self, url):
        import html
        return html.unescape(self.admin.get(url).get_data(as_text=True))

    def day_url(self, **extra):
        args = "".join(f"&{k}={v}" for k, v in extra.items())
        return f"/day?program={self.key.replace(' ', '+')}&date={WED.isoformat()}{args}"

    def test_the_rta_offers_the_zones_and_egypt_is_the_usual(self):
        page = self.get(self.day_url())
        self.assertIn('<select name="tz"', page)
        self.assertIn('<option value="Africa/Cairo" selected>', page)
        self.assertNotIn('aria-labelledby="zd-h"', page)  # the usual page, every tab and change as before

    def test_the_rta_in_new_york(self):
        page = self.get(self.day_url(tz="America/New_York"))
        self.assertIn('aria-labelledby="zd-h"', page)
        self.assertIn("Wednesday 14 Oct in New York (US Eastern), UTC-4", page)
        self.assertIn('<option value="America/New_York" selected>', page)
        self.assertIn("Read only", page)
        self.assertRegex(page, r'<th scope="row">00:00</th>')

    def test_new_york_has_the_language_and_channel_rows(self):
        # Phase AG (T8): every RTA view in another zone, read only: the languages and channels too
        page = self.get(self.day_url(tz="America/New_York"))
        rows = page[page.index("<h3>Languages and channels</h3>"):]
        self.assertRegex(rows, r'<th scope="row">English</th>')
        self.assertRegex(rows, r'<th scope="row">Phone</th>')
        self.assertRegex(rows, r'<td class="zd-c (lok|bad|ok|none|warn)" title="00:00')

    def test_the_operating_day(self):
        got = self.planner.post("/day/operating-day", data={"csrf_token": self.token(self.planner),
                                                            "program": self.key, "date": WED.isoformat(),
                                                            "start": "16:00"})
        self.assertEqual(got.status_code, 403)
        got = self.admin.post("/day/operating-day", data={"csrf_token": self.token(self.admin), "program": self.key,
                                                          "date": WED.isoformat(), "start": "16:00"})
        self.assertEqual(got.status_code, 303)
        page = self.get(self.day_url(window="operating"))
        self.assertIn("Operating day, Wednesday 16:00 to Thursday 16:00", page)
        self.assertRegex(page, r'<tr[^>]*><th scope="row">16:00</th>')

    def test_the_week_page_in_new_york(self):
        page = self.get(f"/week?program={self.key.replace(' ', '+')}&week=2026-10-11&tz=America/New_York")
        self.assertIn('<option value="America/New_York" selected>', page)
        self.assertIn("Times in New York (US Eastern)", page)

    def test_the_schedules_page_offers_the_zones(self):
        page = self.get(f"/schedules?program={self.key.replace(' ', '+')}")
        self.assertIn('<select name="tz"', page)


class TheWorkbookSaysWhenTheDayStarts(unittest.TestCase):
    def test_operating_day_start_in_instructions(self):
        from openpyxl import load_workbook
        from webapp.day import read_inputs
        from webapp.tests.test_channel_day import channel_day
        folder = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, folder, True)
        path = channel_day(folder / "u.xlsx")
        self.assertIsNone(read_inputs(path)["day_start"])
        wb = load_workbook(path)
        ws = next(wb[n] for n in wb.sheetnames if n.strip().lower() == "instructions")
        ws.append(["Operating day start", "16:00"])
        wb.save(path)
        self.assertEqual(read_inputs(path)["day_start"], 16 * 60)


if __name__ == "__main__":
    unittest.main()
