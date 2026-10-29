# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AU: the people an admin or supervisor manages can be exported to a sheet in the bulk upload's own columns, and
switched off or on several at a time. The same walls as one at a time: only people they manage, never themselves."""
import io
import re
import unittest

from openpyxl import load_workbook

from webapp.tests.test_ar_users import HEAD, PASSWORD, _Users
from webapp.tests.test_runs import sign_in, token


class ThePeopleSheet(_Users):
    def test_the_export_has_the_upload_columns_and_one_row_per_person_the_viewer_manages(self):
        got = self.admin.get("/admin/users/export.xlsx")
        self.assertEqual(got.status_code, 200)
        ws = load_workbook(io.BytesIO(got.data)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(rows[0][:len(HEAD) - 1]), HEAD[:-1])      # no password column: nobody's password is kept
        self.assertEqual(rows[0][len(HEAD) - 1], "Active")
        names = {r[0] for r in rows[1:]}
        self.assertIn("plan_a", names)
        self.assertIn("plan_b", names)

    def test_a_supervisor_gets_only_their_programs_people(self):
        sup = sign_in(self.app, "sup_a", PASSWORD)
        ws = load_workbook(io.BytesIO(sup.get("/admin/users/export.xlsx").data)).active
        names = {r[0] for r in list(ws.iter_rows(values_only=True))[1:]}
        self.assertIn("plan_a", names)
        self.assertNotIn("plan_b", names)
        self.assertNotIn("omar", names)

    def test_a_person_who_may_not_manage_people_gets_a_403(self):
        self.assertEqual(sign_in(self.app, "tm_a", PASSWORD).get("/admin/users/export.xlsx").status_code, 403)

    def test_a_cell_that_starts_like_a_formula_is_written_as_text(self):
        uid = self.store.add_user("formula_one", "=HYPERLINK(\"http://evil\")", PASSWORD, must_change=False)
        self.store.set_user_programs(uid, [self.alpha_id])
        ws = load_workbook(io.BytesIO(self.admin.get("/admin/users/export.xlsx").data)).active
        cell = next(r[1] for r in ws.iter_rows(values_only=True) if r[0] == "formula_one")
        self.assertFalse(str(cell).startswith("="))


class TheSeveralAtOnce(_Users):
    def make(self, *names):
        ids = []
        for n in names:
            uid = self.store.add_user(n, n.title(), PASSWORD, must_change=False)
            self.store.set_user_programs(uid, [self.alpha_id])
            ids.append(uid)
        return ids

    def post(self, client, **data):
        return client.post("/admin/users/switch", data={"csrf_token": token(client), **data})

    def test_several_are_switched_off_and_on_again(self):
        a, b = self.make("switch_a", "switch_b")
        self.post(self.admin, who=[str(a), str(b)], decision="off")
        self.assertEqual([self.store.get_user(i)["active"] for i in (a, b)], [0, 0])
        self.post(self.admin, who=[str(a), str(b)], decision="on")
        self.assertEqual([self.store.get_user(i)["active"] for i in (a, b)], [1, 1])

    def test_nobody_switches_themselves_off_and_a_supervisor_cannot_reach_another_programs_people(self):
        sup = sign_in(self.app, "sup_a", PASSWORD)
        me = next(u for u in self.store.list_users() if u["username"] == "sup_a")
        other = next(u for u in self.store.list_users() if u["username"] == "plan_b")
        mine = next(u for u in self.store.list_users() if u["username"] == "plan_a")
        self.post(sup, who=[str(me["id"]), str(other["id"]), str(mine["id"])], decision="off")
        self.assertEqual(self.store.get_user(me["id"])["active"], 1)
        self.assertEqual(self.store.get_user(other["id"])["active"], 1)
        self.assertEqual(self.store.get_user(mine["id"])["active"], 0)
        self.post(self.admin, who=[str(mine["id"])], decision="on")

    def test_the_answer_says_how_many_and_the_events_are_recorded(self):
        a, b = self.make("switch_c", "switch_d")
        got = self.admin.post("/admin/users/switch", data={"csrf_token": token(self.admin), "who": [str(a), str(b)],
                                                           "decision": "off"}, follow_redirects=True)
        self.assertIn("2 people switched off", got.get_data(as_text=True))
        kinds = [e["kind"] for e in self.store.list_events(0, 9e12, None, None)]
        self.assertGreaterEqual(kinds.count("user_disabled"), 2)

    def test_nothing_picked_says_so(self):
        got = self.admin.post("/admin/users/switch", data={"csrf_token": token(self.admin), "decision": "off"},
                              follow_redirects=True)
        self.assertIn("Tick the people first", got.get_data(as_text=True))

    def test_the_form_needs_the_token(self):
        self.assertEqual(self.admin.post("/admin/users/switch", data={"who": ["1"], "decision": "off"}).status_code, 400)


if __name__ == "__main__":
    unittest.main()


class ThePage(_Users):
    def test_the_people_page_offers_the_download_and_the_switch_form(self):
        page = self.admin.get("/admin/users").get_data(as_text=True)
        self.assertIn("/admin/users/export.xlsx", page)
        self.assertIn('data-filter="switch-list"', page)
        self.assertNotIn(f'name="who" value="{next(u for u in self.store.list_users() if u["username"] == "omar")["id"]}"', page.split('id="switch-list"')[1].split("</fieldset>")[0])
