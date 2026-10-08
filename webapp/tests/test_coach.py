# © 2026 Omar Mokhtar. All rights reserved.
"""Phase O task 8: the shrinkage coach.

Over a period, per weekday and interval: the actual shrinkage on days someone
recorded attendance or moves (1 - on the floor / paid on shift), against the
input's Shrinkage tab; the suggestion is the actual average where enough days
were recorded, else the input's value; the corrected tab keeps the input's
layout so it can be pasted into the next input workbook."""
import io
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from webapp.adherence import interval_shrinkage
from webapp.attendance import DayBook
from webapp.coach import actual_shrinkage, corrected_tab
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_schedules import AFTER, INPUT, REPO

WED, THU = date(2026, 10, 14), date(2026, 10, 15)
START, END = date(2026, 10, 11), date(2026, 10, 17)


class TheCoach(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program="AE/AR B2B",
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))
        for name in ("Associate 001", "Associate 012"):
            self.days.set_status("AE/AR B2B", WED, name, "Unplanned leave", self.sara)

    def test_only_recorded_days_count(self):
        found = actual_shrinkage(self.days, "AE/AR B2B", START, END, min_days=1)
        self.assertEqual(found["interval"], 60)
        self.assertEqual(found["recorded_days"], ["2026-10-14"])
        wed_noon = found["weekdays"][3][12 * 60]
        page = self.days.page("AE/AR B2B", WED)
        row = next(r for r in interval_shrinkage(page["view"], page["inputs"], 3) if r["t"] == 12 * 60)
        self.assertEqual((wed_noon["actual"], wed_noon["input"], wed_noon["days"]), (row["actual"], 0.17, 1))
        self.assertEqual(wed_noon["suggested"], row["actual"])
        self.assertEqual(found["weekdays"][4][12 * 60]["days"], 0)  # Thursday: nothing recorded
        self.assertEqual(found["weekdays"][4][12 * 60]["suggested"], 0.17)

    def test_too_few_days_keep_the_input(self):
        found = actual_shrinkage(self.days, "AE/AR B2B", START, END)  # three days needed by default
        cell = found["weekdays"][3][12 * 60]
        self.assertEqual((cell["days"], cell["suggested"], cell["kept_input"]), (1, 0.17, True))

    def test_corrected_tab_keeps_the_input_layout(self):
        found = actual_shrinkage(self.days, "AE/AR B2B", START, END, min_days=1)
        wb = load_workbook(io.BytesIO(corrected_tab(found, "AE/AR B2B").getvalue()))
        ws = wb["Shrinkage 60 Min"]
        self.assertEqual([c.value for c in ws[2]][:8], ["Interval", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"])
        self.assertEqual(ws.max_row, 2 + 24)
        self.assertEqual(ws.cell(3 + 12, 1).value, "12:00")
        self.assertAlmostEqual(ws.cell(3 + 12, 5).value, found["weekdays"][3][12 * 60]["suggested"])
        self.assertAlmostEqual(ws.cell(3 + 12, 6).value, 0.17)
        self.assertIn("How this was made", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
