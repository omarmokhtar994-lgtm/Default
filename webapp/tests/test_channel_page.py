# © 2026 Omar Mokhtar. All rights reserved.
"""Phase V task 6: the "Plan channels" page of a version (owner, 2026-10-09: "assign each associate for a while for
channel based on the requirements ... this can be done on the website only ... automatically plan same as for
breaks").

A grid per day of who is on what in 15-minute steps, coverage per channel and language, and the day's figures.
Suggest plans the day (channels, and breaks inside the break rules); a block can be changed by hand; Copy repeats a
day on days with the same shifts; nothing is kept until Save, which makes a new version checked like any other."""
import html
import io
import json
import re
import shutil
import tempfile
import unittest
from pathlib import Path

from webapp.channel_people import ChannelPeople
from webapp.programs import ProgramBook
from webapp.tests.test_channels import add_channel_tabs
from webapp.tests.test_ready import make_ready
from webapp.tests.test_runs import REPO, make_app, sign_in, token, wait
from webapp.versions import DAYS, shift_span

WED = 3


class _Page(unittest.TestCase):
    @classmethod
    def start(cls, channels=True):
        cls.dir = Path(tempfile.mkdtemp())
        ready = make_ready(cls.dir / "ready.xlsx")
        book = add_channel_tabs(cls.dir / "channels.xlsx", src=ready) if channels else ready
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        programs = ProgramBook(cls.store)
        cls.program = programs.add_program("SAKS")
        cls.key = programs.add_lob(cls.program, "NMG Tier 2")
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")
        cls.admin.post("/runs", data={"csrf_token": token(cls.admin), "kind": "ready", "mode": "QUICK",
                                      "program": cls.key, "week_start": "2026-10-11",
                                      "workbook": (io.BytesIO(book.read_bytes()), "week.xlsx")},
                       content_type="multipart/form-data")
        run_id = cls.store.list_runs()[0]["id"]
        wait(cls.store, run_id, statuses=("DONE", "REJECTED", "FAILED"))
        cls.run_id = run_id
        cls.version = cls.app.extensions["schedules"].versions(run_id)[0]
        cls.week = json.loads(cls.version["week"])
        cls.url = f"/schedules/{cls.version['id']}/channels"

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def page(self, day="Wed"):
        got = self.admin.get(f"{self.url}?day={day}")
        return got.status_code, html.unescape(got.get_data(as_text=True))

    def post(self, **data):
        return self.admin.post(self.url, data={"csrf_token": token(self.admin), "day": "Wed", **data})

    def working(self, d=WED):
        return [a for a in self.week["associates"] if shift_span(a["days"][d])]


class TheChannelPage(_Page):
    @classmethod
    def setUpClass(cls):
        cls.start()

    def test_the_page_before_any_plan(self):
        code, page = self.page("Thu")  # no test here plans Thursday
        self.assertEqual(code, 200)
        self.assertIn("SAKS, NMG Tier 2: plan channels", page)
        self.assertIn("Not planned yet: Suggest the day, or set blocks by hand.", page)
        for words in ("Suggest the day", "Move breaks if it helps", "Save as a new version", "Phone", "Chat", "Email",
                      "Break or lunch", self.working(4)[0]["name"]):
            self.assertIn(words, page)

    def test_suggest_fills_a_draft_and_nothing_is_kept_until_saved(self):
        self.assertEqual(self.post(action="suggest", move="1").status_code, 303)
        code, page = self.page()
        self.assertIn("Suggested channels for Wednesday", page)
        self.assertIn("Nothing is kept until you save.", page)
        self.assertRegex(page, r'class="c ch-[PCE]"')
        self.assertRegex(page, r"\d+(\.\d+)? h</b> short on Phone")
        draft = self.store.get_channel_draft(self.version["id"])
        self.assertTrue(draft["Wed"]["blocks"])
        self.assertEqual(json.loads(self.store.get_schedule(self.version["id"])["week"])["channels"], [])
        # a break placed by the suggestion cannot be given a channel
        name = self.working()[0]["name"]
        start = draft["Wed"]["breaks"][name][0]
        self.post(action="edit", who=name, start=f"{start // 60:02d}:{start % 60:02d}",
                  end=f"{(start + 15) // 60:02d}:{(start + 15) % 60:02d}", channel="P")
        self.assertIn(f"{name} is on a break then", self.page()[1])

    def test_a_block_set_by_hand_is_checked_then_kept_in_the_draft(self):
        person = self.working()[1]
        lo, hi = shift_span(person["days"][WED])
        hh = lambda m: f"{(m // 60) % 24:02d}:{m % 60:02d}"  # noqa: E731
        self.post(action="edit", who=person["name"], start="00:00", end="01:00", channel="P") if lo >= 120 else None
        if lo >= 120:
            self.assertIn(f"outside {person['name']}'s shift", self.page()[1])
        self.post(action="edit", who=person["name"], start=hh(lo + 10), end=hh(lo + 70), channel="P")
        self.assertIn("15-minute steps", self.page()[1])
        ChannelPeople(self.store).save(self.program, {person["name"]: "P"}, 1)
        self.post(action="edit", who=person["name"], start=hh(lo), end=hh(lo + 60), channel="C")
        self.assertIn(f"{person['name']} cannot work Chat (see Associate channels).", self.page()[1])
        ChannelPeople(self.store).save(self.program, {person["name"]: "PCE"}, 1)
        self.post(action="edit", who=person["name"], start=hh(lo), end=hh(lo + 60), channel="C")
        self.assertIn(f"{person['name']}: Chat from {hh(lo)} to {hh(lo + 60)}. Nothing is kept until you save.",
                      self.page()[1])
        blocks = self.store.get_channel_draft(self.version["id"])["Wed"]["blocks"][person["name"]]
        self.assertIn([lo, lo + 60, "C"], [list(b) for b in blocks])

    def test_all_channels_times_take_no_single_channel(self):
        evening = [a for a in self.week["associates"] if shift_span(a["days"][5])
                   and shift_span(a["days"][5])[0] <= 20 * 60 < shift_span(a["days"][5])[1]]
        if not evening:
            self.skipTest("nobody works Friday evening in this schedule")
        self.admin.post(self.url, data={"csrf_token": token(self.admin), "day": "Fri", "action": "edit",
                                        "who": evening[0]["name"], "start": "20:00", "end": "21:00", "channel": "P"})
        self.assertIn("is an all-channels time", self.page("Fri")[1])

    def test_the_schedules_page_links_to_it(self):
        page = html.unescape(self.admin.get(f"/runs/{self.run_id}/schedules").get_data(as_text=True))
        self.assertIn(f'href="{self.url}"', page)
        self.assertIn("Plan channels", page)


class TheChannelSave(_Page):
    @classmethod
    def setUpClass(cls):
        cls.start()
        cls.admin.post(cls.url, data={"csrf_token": token(cls.admin), "day": "Wed", "action": "suggest", "move": "1"})
        cls.saved = cls.admin.post(cls.url, data={"csrf_token": token(cls.admin), "day": "Wed", "action": "save",
                                                  "reason": "", "use": "1"})
        cls.new_id = int(re.search(r"/schedules/(\d+)/channels", cls.saved.headers["Location"]).group(1))

    def test_saving_makes_a_checked_version_with_the_plan_and_the_breaks(self):
        self.assertNotEqual(self.new_id, self.version["id"])
        row = self.store.get_schedule(self.new_id)
        week = json.loads(row["week"])
        self.assertTrue([c for c in week["channels"] if c["day"] == "Wed"])
        self.assertTrue([b for b in week["breaks"] if b["day"] == "Wed"])
        added = self.app.extensions["schedules"].view(self.new_id)["added"]
        self.assertEqual([p["text"] for p in added if p["severity"] == "red"], [])
        self.assertEqual(self.store.get_channel_draft(self.version["id"]), {})

    def test_the_change_log_says_channels_planned(self):
        changes = self.app.extensions["schedules"].view(self.new_id)["changes"]
        self.assertTrue(changes)
        self.assertEqual({c["reason"] for c in changes}, {"Channels planned"})
        self.assertTrue(any(re.search(r"(Phone|Chat|Email) \d\d:\d\d to \d\d:\d\d", c["new"]) for c in changes))


class TheProgramWithoutChannels(_Page):
    @classmethod
    def setUpClass(cls):
        cls.start(channels=False)

    def test_the_page_says_there_are_no_channel_tabs(self):
        code, page = self.page()
        self.assertEqual(code, 404)
        self.assertIn("This schedule's input workbook has no channel tabs", page)
        schedules = html.unescape(self.admin.get(f"/runs/{self.run_id}/schedules").get_data(as_text=True))
        self.assertNotIn("Plan channels", schedules)


if __name__ == "__main__":
    unittest.main()


class TheWholeWeek(_Page):
    """Phase AF, B1: one action plans every day of the week (Plan channels planned one day at a time)."""

    @classmethod
    def setUpClass(cls):
        cls.start()
        cls.done = cls.admin.post(cls.url, data={"csrf_token": token(cls.admin), "day": "Wed",
                                                 "action": "suggest_week"})

    def test_every_working_day_gets_a_suggestion(self):
        draft = self.store.get_channel_draft(self.version["id"])
        working = [DAYS[i] for i in range(7) if any(shift_span(a["days"][i]) for a in self.week["associates"])]
        self.assertTrue(working)
        self.assertEqual(sorted(draft), sorted(working))
        for day in working:
            self.assertTrue(draft[day]["blocks"], day)

    def test_the_page_says_so_and_offers_it(self):
        self.assertEqual(self.done.status_code, 303)
        said = self.page()[1]
        self.assertRegex(said, r"Suggested channels for \d+ days in [\d.]+ seconds\. Nothing is kept until you save\.")
        self.assertIn('value="suggest_week"', said)


class TheStretches(unittest.TestCase):
    """Phase AF (C1): Plan channels shows each person's day as stretches (one box per run), each opening the block
    form with its own times; the quarter-hour table stays, folded under "Show quarter hours"."""

    def test_the_grid_joins_quarters_into_stretches(self):
        from webapp.channel_page import grid
        person = {"name": "Associate 01", "start": 600, "end": 1140, "breaks": [("Break 1", None, 15)],
                  "language": "English", "label": "10:00 - 19:00", "skills": "PC"}
        plan = {"blocks": {"Associate 01": [(600, 720, "P"), (735, 900, "C")]}, "breaks": {"Associate 01": [720]}}
        row = grid([person], plan)["rows"][0]
        self.assertEqual(row["runs"], [(600, 720, "P"), (720, 735, "B"), (735, 900, "C"), (900, 1140, "")])


class TheStretchesOnThePage(_Page):
    @classmethod
    def setUpClass(cls):
        cls.start()
        cls.admin.post(cls.url, data={"csrf_token": token(cls.admin), "day": "Wed", "action": "suggest"})

    def test_each_stretch_opens_the_block_form_and_the_table_stays(self):
        status, page = self.page()
        self.assertEqual(status, 200)
        self.assertRegex(page, r'<a [^>]*data-cp-run data-who="[^"]+" data-from="\d\d:\d\d" data-to="\d\d:\d\d" '
                               r'data-letter="[PCE]"')
        self.assertRegex(page, r"<title>Associate \d+: (Phone|Chat|Email) \d\d:\d\d to \d\d:\d\d</title>")
        self.assertIn("<summary>Show quarter hours</summary>", page)
        self.assertIn('<table class="cp-grid"', page)


class TheChannelDownload(_Page):
    """Phase AF (C4): the channel plan as an Excel file in two views (owner, 2026-10-14: "the view shared is good;
    in addition a similar view same as the timeline"): per working day a table of stretches and a timeline, and a
    Week tab of hours per person and channel."""

    @classmethod
    def setUpClass(cls):
        cls.start()
        cls.admin.post(cls.url, data={"csrf_token": token(cls.admin), "day": "Wed", "action": "suggest"})
        cls.got = cls.admin.get(cls.url.replace("/channels", "/channels.xlsx"))

    def test_both_views_per_day_and_a_week_tab(self):
        from openpyxl import load_workbook
        self.assertEqual(self.got.status_code, 200)
        self.assertIn("spreadsheetml", self.got.headers["Content-Type"])
        wb = load_workbook(io.BytesIO(self.got.data))
        self.assertEqual(wb.sheetnames[0], "Read me")
        self.assertIn("Wed", wb.sheetnames)
        self.assertIn("Wed timeline", wb.sheetnames)
        self.assertEqual(wb.sheetnames[-1], "Week")
        table = list(wb["Wed"].iter_rows(values_only=True))
        self.assertEqual(table[0], ("Associate", "Language", "From", "To", "Channel", "Minutes"))
        self.assertTrue(any(r[4] in ("Phone", "Chat", "Email") for r in table[1:]))
        line = wb["Wed timeline"]
        letters = {c.value for row in line.iter_rows(min_row=2, min_col=2) for c in row if c.value}
        self.assertTrue(letters & {"P", "C", "E"}, letters)
        filled = next(c for row in line.iter_rows(min_row=2, min_col=2) for c in row if c.value == "P")
        self.assertEqual(filled.fill.fgColor.rgb[-6:], "3E6FD8")

    def test_plan_channels_offers_it(self):
        self.assertIn("Download the channel plan (Excel)", self.page()[1])
