# © 2026 Omar Mokhtar. All rights reserved.
"""Phase N task 1: schedule workbooks read, edited and checked.

Real run files already in the repository: a week with breaks (AE/AR B2B
slice, names "Associate 001"...) and a before-breaks week, each with the input
workbook it was built from. The checks are the independent validator's."""
import shutil
import tempfile
import unittest
from pathlib import Path

from webapp.versions import apply_change, marks, problems, read_week, swap_slots, validate

REPO = Path(__file__).resolve().parents[2]
AFTER_DIR = REPO / "fixtures" / "real_runs" / "week_boundary" / "B3_ARB2B_S30"
AFTER = AFTER_DIR / "production" / "B3_ARB2B_S30_L6_3_2_3_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx"
AFTER_INPUT = AFTER_DIR / "input_snapshot" / "AE_AR_B2B_SLICE30.xlsx"
BEFORE_DIR = REPO / "fixtures" / "real_runs" / "before_break" / "B1_BEFORE"
BEFORE = BEFORE_DIR / "B1_BEFORE_L6_3_2_3_BEST_BEFORE_BREAKS_SCHEDULE.xlsx"
BEFORE_INPUT = BEFORE_DIR / "input_snapshot" / "AE_FR_Choice.xlsx"


class TheWorkbooks(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def edited(self, name="Associate 001", day="Wed", value="21:00 - 06:00"):
        out = self.tmp / "edited.xlsx"
        apply_change(AFTER, out, name, day, value)
        return out

    def test_read_week(self):
        week = read_week(AFTER)
        first = week["associates"][0]
        self.assertEqual((first["name"], first["language"], first["days"][3]), ("Associate 001", "English", "12:00 - 21:00"))
        self.assertEqual(len(week["associates"]), 39)
        self.assertIn("21:00 - 06:00", week["shifts"])
        self.assertEqual(week["shifts"][-2:], ["OFF", "Leave"])
        self.assertEqual(week["stage"], "after")
        self.assertIn({"associate": "Associate 001", "day": "Wed", "kind": "Lunch", "start": "16:30", "minutes": 30},
                      week["breaks"])
        self.assertEqual(read_week(BEFORE)["stage"], "before")

    def test_apply_change_writes_both_tabs_and_drops_that_days_breaks(self):
        week = read_week(self.edited())
        self.assertEqual(week["associates"][0]["days"][3], "21:00 - 06:00")
        self.assertFalse([b for b in week["breaks"] if (b["associate"], b["day"]) == ("Associate 001", "Wed")])
        self.assertTrue([b for b in week["breaks"] if (b["associate"], b["day"]) == ("Associate 001", "Tue")])
        from openpyxl import load_workbook
        final = load_workbook(self.tmp / "edited.xlsx")["Final Schedule"]
        self.assertEqual(next(r for r in final.iter_rows(min_row=3, values_only=True) if r[3] == "Associate 001")[9],
                         "21:00 - 06:00")

    def test_slot_ids_are_read(self):
        people = read_week(AFTER)["associates"]
        self.assertEqual([(a["slot"], a["name"]) for a in people[:2]], [("1", "Associate 001"), ("2", "Associate 002")])

    def test_swap_slots_moves_the_people_not_the_shifts(self):
        before = read_week(AFTER)
        out = self.tmp / "swapped.xlsx"
        swap_slots(AFTER, out, "Associate 001", "Associate 002")
        after = read_week(out)
        old = {a["slot"]: a for a in before["associates"]}
        new = {a["slot"]: a for a in after["associates"]}
        # slot 1 keeps its shifts and now holds Associate 002 (with their Emp ID); slot 2 the reverse
        self.assertEqual((new["1"]["name"], new["1"]["emp_id"], new["1"]["days"]),
                         ("Associate 002", old["2"]["emp_id"], old["1"]["days"]))
        self.assertEqual((new["2"]["name"], new["2"]["emp_id"], new["2"]["days"]),
                         ("Associate 001", old["1"]["emp_id"], old["2"]["days"]))
        # the breaks stay with the shifts: slot 1's breaks now carry Associate 002's name
        mine = lambda week, name: sorted((b["day"], b["kind"], b["start"]) for b in week["breaks"] if b["associate"] == name)
        self.assertEqual(mine(after, "Associate 002"), mine(before, "Associate 001"))
        self.assertEqual(mine(after, "Associate 001"), mine(before, "Associate 002"))
        from openpyxl import load_workbook
        final = load_workbook(out)["Final Schedule"]
        self.assertEqual(next(r for r in final.iter_rows(min_row=3, values_only=True) if r[0] == 1)[3], "Associate 002")

    def test_swap_needs_two_different_people_on_the_schedule(self):
        with self.assertRaises(ValueError):
            swap_slots(AFTER, self.tmp / "x.xlsx", "Associate 001", "Associate 001")
        with self.assertRaises(ValueError):
            swap_slots(AFTER, self.tmp / "x.xlsx", "Associate 001", "Nobody")

    def test_unknown_associate_or_value_refused(self):
        with self.assertRaises(ValueError):
            apply_change(AFTER, self.tmp / "x.xlsx", "Nobody", "Wed", "OFF")
        with self.assertRaises(ValueError):
            apply_change(AFTER, self.tmp / "x.xlsx", "Associate 001", "Wed", "25:00 - 26:00")
        with self.assertRaises(ValueError):
            apply_change(AFTER, self.tmp / "x.xlsx", "Associate 001", "Someday", "OFF")

    def test_validate_finds_rest_variety_and_missing_breaks(self):
        result = validate(AFTER_INPUT, self.edited(), REPO)
        self.assertEqual(result["status"], "FAIL")
        found = {p["key"].split("|")[0]: p for p in problems(result, edited={("Associate 001", "Wed")})}
        self.assertEqual(found["REST_VIOLATION"]["severity"], "red")
        self.assertIn("rests only", found["REST_VIOLATION"]["text"])
        self.assertEqual(found["MAX_SHIFT_VARIETY_VIOLATION"]["severity"], "red")
        self.assertEqual(found["BREAK_SEGMENT_COUNT_OR_DURATION"]["severity"], "yellow")  # breaks still to plan
        self.assertEqual(len(result["intervals"]), 168)

    def test_the_unedited_week_passes(self):
        result = validate(AFTER_INPUT, AFTER, REPO)
        self.assertEqual((result["status"], [p for p in problems(result) if p["severity"] == "red"]), ("PASS", []))

    def test_validate_writes_nothing_beside_the_schedule(self):
        folder = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, folder, True)
        schedule = folder / "kept.xlsx"
        shutil.copyfile(AFTER, schedule)
        validate(AFTER_INPUT, schedule, REPO)
        self.assertEqual(sorted(f.name for f in folder.iterdir()), ["kept.xlsx"])

    def test_marks_cells_and_days(self):
        m = marks(problems(validate(AFTER_INPUT, self.edited(), REPO), edited={("Associate 001", "Wed")}))
        self.assertEqual(m["cells"]["Associate 001|Wed"], "red")
        self.assertEqual(m["cells"]["Associate 001|Thu"], "red")  # the rest runs into Thursday
        self.assertEqual(m["days"]["Wed"], "red")

    def test_before_breaks_version_has_no_break_problems(self):
        out = self.tmp / "before.xlsx"
        first = read_week(BEFORE)["associates"][0]
        day = next(d for d, v in zip(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"], first["days"]) if v != "OFF")
        apply_change(BEFORE, out, first["name"], day, "OFF")
        kinds = {p["key"].split("|")[0] for p in problems(validate(BEFORE_INPUT, out, REPO), edited={(first["name"], day)})}
        self.assertFalse({k for k in kinds if k.startswith("BREAK")})


if __name__ == "__main__":
    unittest.main()
