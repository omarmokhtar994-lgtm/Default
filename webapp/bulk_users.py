# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AR (owner, 2026-10-21: "i want to be able to upload bulk users"): the people sheet an admin or supervisor uploads
to add many people at once, and the sheet of temporary passwords they get back. The checks that decide whether a row may
be added are the same as for one person and live with the page (webapp/app.py); this module only reads and writes files."""
from __future__ import annotations

import csv
import io
import secrets
import zipfile
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

HEAD = ["Username", "Display name", "Role", "Programs", "Roster name", "Team manager", "Password"]
MAX_ROWS = 500
NOT_A_SHEET = ("Upload the people sheet: an Excel (.xlsx) or CSV file with the columns " + ", ".join(HEAD)
               + ". Download the template from this page.")
EXAMPLE = ["associate_01", "Associate 01", "associate", "Program A", "Associate 01", "team_manager_01", ""]
HELP = [
    ("Username", "Letters, numbers, dots, dashes and underscores. Lower case is used."),
    ("Role", "planner, associate, team_manager, supervisor or admin. A supervisor may add planners, associates and team managers."),
    ("Programs", "Program names, separated by semicolons. An admin may write ALL for every program (not for an associate)."),
    ("Roster name", "Associates only: their name as it is on the schedule."),
    ("Team manager", "Associates only, optional: the username of their team manager."),
    ("Password", "Optional. Leave it blank and one is made up for you and shown once."),
]


def new_password() -> str:
    """A temporary password (12 characters), made up for a person whose row or form gave none."""
    return secrets.token_urlsafe(9)


def _safe(value: Any) -> Any:
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(HEAD)
    ws.append(EXAMPLE)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFG", (20, 22, 14, 24, 22, 20, 18)):
        ws.column_dimensions[col].width = width
    notes = wb.create_sheet("How to fill it")
    notes.append(["Column", "What to write"])
    for row in HELP:
        notes.append(list(row))
    notes.append([])
    notes.append(["The example row is made up. Delete it before you upload."])
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["B"].width = 110
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def people_workbook(people: List[Dict[str, Any]]) -> bytes:
    """Phase AU: the people an admin or supervisor manages, in the upload's own columns (so a sheet can be edited and
    uploaded for new people), with Active last. No password column: nobody's password is kept in a readable form."""
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(HEAD[:-1] + ["Active"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for p in people:
        ws.append([_safe(p.get(k, "")) for k in ("username", "display_name", "role", "programs", "roster_name",
                                                  "team_manager")] + ["Yes" if p.get("active") else "No"])
    for col, width in zip("ABCDEFG", (20, 22, 14, 28, 22, 20, 8)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _rows(data: bytes, filename: str) -> List[List[Any]]:
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    except (zipfile.BadZipFile, InvalidFileException, KeyError, IndexError, OSError, ValueError) as exc:
        raise ValueError(NOT_A_SHEET) from exc


def read_people(data: bytes, filename: str) -> List[Dict[str, Any]]:
    """The rows of an uploaded people sheet as dicts ({"n": the row's number in the file, plus a key for each column}),
    blank rows left out. ValueError, in words, for a file that is not a people sheet or holds too many rows."""
    rows = _rows(data, filename)
    head = next((i for i, row in enumerate(rows) if any(_norm(c) == "username" for c in row)), None)
    if head is None:
        raise ValueError(NOT_A_SHEET)
    columns = {_norm(c): i for i, c in enumerate(rows[head]) if c not in (None, "")}
    keys = {"username": "username", "display name": "display_name", "role": "role", "programs": "programs",
            "roster name": "roster_name", "team manager": "team_manager", "password": "password"}
    out = []
    for offset, row in enumerate(rows[head + 1:], start=head + 2):
        values = {key: str(row[columns[name]]).strip() if name in columns and columns[name] < len(row)
                  and row[columns[name]] is not None else "" for name, key in keys.items()}
        if any(values.values()):
            out.append({"n": offset, **values})
    if len(out) > MAX_ROWS:
        raise ValueError(f"Add at most {MAX_ROWS} people at a time; this file has {len(out)}.")
    if not out:
        raise ValueError("The file has no people in it. Fill the template and upload it again.")
    return out


def passwords_workbook(rows: List[Dict[str, str]]) -> bytes:
    """Who got which temporary password ({"username", "display_name", "password"}), for the person who added them."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Temporary passwords"
    ws.append(["Username", "Temporary password", "Display name"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row["username"], row["password"], _safe(row["display_name"])])
    ws.append([])
    ws.append(["Each person chooses their own password at first sign-in. Share these once and delete this file."])
    for col, width in zip("ABC", (24, 22, 28)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
