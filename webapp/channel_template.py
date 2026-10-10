# © 2026 Omar Mokhtar. All rights reserved.
"""The channel tabs, laid out as the approved Phase V sample workbook (evidence/phase_v/samples/
channels_input_sample.xlsx): a read-me, Chat, Phone and Email per interval, Email Hours and Channel Setup.

Phase Z (owner, 2026-10-10: "i cannot find where i can add the required per channel"): the website hands these out
as a channel needs workbook made for one schedule's week, and the Blank and Example input workbooks carry them too.
``webapp.channels.read_channels`` reads what is filled in; the engine never reads these tabs."""
from pathlib import Path
from typing import Callable, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
TITLE = Font(bold=True, size=14)
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SETTINGS = [
    ("Minimum block minutes", 60, "The shortest time on a channel before switching (in whole intervals)."),
    ("Phone maximum continuous minutes", 120, "The longest time on Phone in a row; 0 means no limit."),
    ("Chat maximum continuous minutes", 0, "0 means no limit."),
    ("Email maximum continuous minutes", 0, "0 means no limit."),
    ("Priority when short", "Phone > Chat > Email",
     "Which channel is covered first when there are not enough people (pick an order)."),
    ("How strict is the order", "Strict", "Strict: a channel earlier in the order always comes first. Balanced: a "
     "long gap on a later channel can outweigh a short gap on an earlier one."),
    ("Spread channels fairly over the week", "Yes", "Each person gets a fair share of each channel they can work."),
]
ORDERS = ["Phone > Chat > Email", "Phone > Email > Chat", "Chat > Phone > Email", "Chat > Email > Phone",
          "Email > Phone > Chat", "Email > Chat > Phone"]
LANGUAGE_ROWS = 6  # empty rows offered under the language minimums
BLENDED_ROWS = 11  # and under the all-channels times


def _head(ws, row: int, labels: List[str]) -> None:
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row, c, label)
        cell.font, cell.fill = HEAD, HEAD_FILL


def _title(ws, title: str, note: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = note


def _read_me(wb, step: int) -> None:
    ws = wb.create_sheet("Channels - read me")
    ws.column_dimensions["A"].width = 120
    lines = [
        "Channels: what these tabs are for",
        f"The schedule is built exactly as today from FT Wise {step} Min; the engine does not read these tabs.",
        "The website then places each person on the floor on a channel, as it does for breaks.",
        "",
        f"Chat {step} Min, Phone {step} Min: people needed per interval, like FT Wise {step} Min. Chat + Phone + Email "
        "make up the requirement.",
        f"Email: either Email {step} Min (per interval, like Chat), or, when that tab is left out or empty, "
        "Email Hours",
        "   (hours per day, shared out automatically into the quieter times of the day).",
        "Channel Setup: the rotation rules, the language minimums per channel (with start and end times), and the",
        "   all-channels times, when 1 or 2 people cover every channel at once (no add-up warning there).",
        "Who can work which channel is set once on the website (Manage > Associate channels), not every week.",
        "",
        f"Check on upload: the website says where Chat + Phone + Email differ from FT Wise {step} Min, interval by "
        "interval.",
    ]
    for r, text in enumerate(lines, start=1):
        if text:
            ws.cell(r, 1, text)
    ws["A1"].font = TITLE


def _grid(wb, name: str, note: str, step: int, value: Optional[Callable[[int, int], float]], freeze: bool) -> None:
    ws = wb.create_sheet(name)
    _title(ws, name if value is not None else f"{name} (optional)", note)
    _head(ws, 4, ["Interval"] + DAYS)
    ws.column_dimensions["A"].width = 11
    for i, t in enumerate(range(0, 1440, step)):
        r = 5 + i
        ws.cell(r, 1, f"{t // 60:02d}:{t % 60:02d}")
        if value is not None:
            for d in range(7):
                ws.cell(r, 2 + d, value(d, t))
    if freeze:
        ws.freeze_panes = "B5"


def _email_hours(wb, languages: List[str]) -> None:
    ws = wb.create_sheet("Email Hours")
    _title(ws, "Email Hours", "Used when the Email tab per interval is empty: the hours of email work needed each "
           "day. The website shares them out inside the window, into the times with the most spare people, in "
           "blocks like any channel.")
    split = [f"Of which {lang}" for lang in languages] if len(languages) > 1 else []
    _head(ws, 4, ["Day", "Hours needed"] + split + ["Window start", "Window end"])
    for i, day in enumerate(DAYS):
        ws.cell(5 + i, 1, day)
        ws.cell(5 + i, 2, 0)
    for col, width in zip("ABCDE", (10, 14, 16, 14, 12)):
        ws.column_dimensions[col].width = width


def _setup(wb, step: int) -> None:
    ws = wb.create_sheet("Channel Setup")
    _title(ws, "Channel Setup", "The rotation rules, then the language minimums per channel. Leave a row inactive to "
           "switch it off without deleting it.")
    _head(ws, 4, ["Setting", "Value", "What it means"])
    for i, (name, value, means) in enumerate(SETTINGS):
        ws.cell(5 + i, 1, name)
        ws.cell(5 + i, 2, value)
        ws.cell(5 + i, 3, means)
    order, strict = DataValidation(type="list", formula1='"' + ",".join(ORDERS) + '"', allow_blank=False), \
        DataValidation(type="list", formula1='"Strict,Balanced"', allow_blank=False)
    order.add("B9")
    strict.add("B10")
    head = 5 + len(SETTINGS) + 2
    _head(ws, head, ["Channel", "Language", "Minimum per interval", "Coverage start", "Coverage end", "Coverage days",
                     "Active?"])
    channel = DataValidation(type="list", formula1='"Chat,Phone,Email"', allow_blank=True)
    channel.add(f"A{head + 1}:A{head + LANGUAGE_ROWS}")
    blended = head + LANGUAGE_ROWS + 2
    ws.cell(blended, 1, "All channels together").font = Font(bold=True)
    ws.cell(blended + 1, 1, f"Times when the 1 or 2 people on the floor cover every channel they can work at once. "
                            f"There the website does not add Chat + Phone + Email up against FT Wise {step} Min (no "
                            "false warning), plans no channel blocks, and on the day warns when a channel would have "
                            "nobody.")
    _head(ws, blended + 2, ["Days", "Start", "End", "Active?", "Note"])
    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    yes_no.add(f"G{head + 1}:G{head + LANGUAGE_ROWS}")
    yes_no.add(f"D{blended + 3}:D{blended + 2 + BLENDED_ROWS}")
    for dv in (order, strict, channel, yes_no):
        ws.add_data_validation(dv)
    for col, width in zip("ABCDEFG", (34, 20, 64, 16, 14, 26, 10)):
        ws.column_dimensions[col].width = width


def add_channel_tabs(wb, step: int, languages: List[str],
                     need: Optional[Callable[[str, int, int], float]] = None) -> None:
    """Add the six channel tabs to ``wb``, for a program on ``step``-minute intervals with these languages.
    ``need(letter, day, minute)`` fills Chat ("C") and Phone ("P"), else 0; Email per interval is left empty;
    Email Hours is 0 each day (with an "Of which" column per language when there are 2 or more); Channel Setup has
    the sample's settings and empty language and all-channels rows."""
    _read_me(wb, step)
    for letter, channel in (("C", "Chat"), ("P", "Phone")):
        _grid(wb, f"{channel} {step} Min", f"People needed on {channel} per interval, as on FT Wise {step} Min.",
              step, (lambda d, t, x=letter: need(x, d, t) if need else 0), freeze=True)
    _grid(wb, f"Email {step} Min", "Fill this tab to plan email per interval like Chat and Phone. Leave it empty (or "
          "delete it) to use Email Hours instead.", step, None, freeze=False)
    _email_hours(wb, languages)
    _setup(wb, step)


def channel_workbook(path: Path, step: int, languages: List[str]) -> Path:
    """A channel needs workbook on its own (the six tabs, nothing filled in), saved at ``path``."""
    wb = Workbook()
    wb.remove(wb.active)
    add_channel_tabs(wb, step, languages)
    wb.save(path)
    return Path(path)
