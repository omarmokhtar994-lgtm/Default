# © 2026 Omar Mokhtar. All rights reserved.
"""The day: the schedule in use plus attendance and actual breaks, interval by interval.

Demand, language setup, shifts and planned breaks come from the version's own
workbook; everything else is what people recorded for the day. The input's
shrinkage is not used: it stands for absences, breaks and aux, which the day
counts as they happen (owner, 2026-10-08). Counted per 5 minutes: a person is
on the floor when on shift, present, not on a break and not in aux. Two
measures: "Interval compliance" counts a billable aux (say, billable coaching)
as on the floor; "Service level" takes everyone in aux off the floor. Per
interval (the workbook's own length):

  needed       = ceil(required)
  on the floor = the average of the interval's 5-minute counts
  plus/minus   = (on the floor - required) x interval hours

The previous day's overnight shifts (and their breaks) count in the early hours.
"""
from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .versions import DAYS, shift_span

STEP = 5  # minutes: breaks move in 5-minute steps (owner, 2026-10-08)
PRESENT = "Present"
STATUSES = ["Present", "Unplanned leave", "Late", "Sick", "Left early", "Training", "Coaching", "Meeting",
            "System issue"]
ABSENT = {"Unplanned leave", "Sick"}  # the whole shift
LATE_EARLY = {"Late", "Left early"}
AUX = {"Training", "Coaching", "Meeting", "System issue"}  # billable or not, chosen when it is recorded
EXTRA_BREAKS = ("Break", "Lunch")  # added on the day in RTA (owner, 2026-10-08): always off the floor
TIMED = LATE_EARLY | AUX  # may carry a from/to time
MEASURES = {"interval": "Interval compliance", "sl": "Service level"}
KEEP = 8  # workbooks kept read
_CACHE: Dict[Tuple[str, float], Dict[str, Any]] = {}


class BreakRefused(ValueError):
    """A break the day cannot take (outside the shift, on another break, not a 5-minute step)."""


def _norm(v: Any) -> str:
    return " ".join(str(v or "").split()).casefold()


def _minute(v: Any) -> Optional[int]:
    if v is None or v == "":
        return None
    if hasattr(v, "hour"):
        return v.hour * 60 + v.minute
    m = re.search(r"(\d{1,2}):(\d{2})", str(v))
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


def _instruction(wb, names) -> Any:
    ws = next((wb[n] for n in wb.sheetnames if _norm(n) == "instructions"), None)
    for row in (ws.iter_rows(values_only=True) if ws is not None else []):
        cells = list(row)
        for i, v in enumerate(cells):
            if _norm(v) in names:
                return next((x for x in cells[i + 1:] if x not in (None, "")), None)
    return None


def _grid(wb, name: str) -> Dict[int, Dict[int, float]]:
    """{day index: {minute: value}} from a demand or shrinkage tab."""
    ws = next((wb[n] for n in wb.sheetnames if _norm(n) == _norm(name)), None)
    out: Dict[int, Dict[int, float]] = {d: {} for d in range(7)}
    if ws is None:
        return out
    header = next(r for r in range(1, 10) if _norm(ws.cell(r, 1).value) == "interval")
    cols = {DAYS.index(str(ws.cell(header, c).value).strip()): c for c in range(2, ws.max_column + 1)
            if str(ws.cell(header, c).value or "").strip() in DAYS}
    for r in range(header + 1, ws.max_row + 1):
        t = _minute(ws.cell(r, 1).value)
        if t is None:
            continue
        for d, c in cols.items():
            try:
                out[d][t] = float(ws.cell(r, c).value or 0)
            except (TypeError, ValueError):
                out[d][t] = 0.0
    return out


def _languages(wb) -> List[Dict[str, Any]]:
    ws = next((wb[n] for n in wb.sheetnames if _norm(n) in ("language setup", "skill setup", "skills setup")), None)
    if ws is None:
        return []
    header = next((r for r in range(1, 10) if _norm(ws.cell(r, 1).value) == "language"), None)
    if header is None:
        return []
    cols = {_norm(ws.cell(header, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(header, c).value}
    get = lambda r, k: ws.cell(r, cols[k]).value if k in cols else None  # noqa: E731
    rows = []
    for r in range(header + 1, ws.max_row + 1):
        name = str(get(r, "language") or "").strip()
        if not name or _norm(get(r, "active?")) in ("no", "n", "false"):
            continue
        covers = {_norm(x) for x in str(get(r, "can cover languages") or "").split(",") if x.strip()} | {_norm(name)}
        days_text = _norm(get(r, "coverage days"))
        days = list(range(7)) if days_text in ("", "all") else [DAYS.index(d) for d in DAYS if _norm(d) in days_text]
        try:
            minimum = int(float(get(r, "minimum per interval") or 0))
        except (TypeError, ValueError):
            minimum = 0
        start, end = _minute(get(r, "coverage start")) or 0, _minute(get(r, "coverage end")) or 0
        rows.append({"name": name, "covers": covers, "start": start, "end": end, "minimum": minimum, "days": days})
    return rows


def read_inputs(path: Path) -> Dict[str, Any]:
    """Demand, interval and language setup of a version's workbook (cached per file)."""
    key = (str(path), Path(path).stat().st_mtime)
    if key in _CACHE:
        return _CACHE[key]
    wb = load_workbook(path, data_only=True)
    step = int(float(_instruction(wb, ("interval minutes",)) or 30))
    demand = _grid(wb, str(_instruction(wb, ("requirements source",)) or f"FT Wise {step} Min"))
    previous = []
    ws = next((wb[n] for n in wb.sheetnames if _norm(n) == "previous week scheduled"), None)
    if ws is not None:
        for row in ws.iter_rows(min_row=3, values_only=True):
            if len(row) > 3 and row[1] and shift_span(row[3]):
                previous.append({"name": str(row[1]).strip(), "language": str(row[2] or "").strip(), "shift": str(row[3])})
    planned = _grid(wb, str(_instruction(wb, ("shrinkage source",)) or f"Shrinkage {step} Min"))
    planned = {d: {t: (v / 100 if v > 1 else v) for t, v in col.items()} for d, col in planned.items()}
    gaps = {}
    for key, names in (("gap_min", ("break absolute minimum gap minutes",)),
                       ("gap_max", ("break normal maximum gap minutes",))):
        try:
            gaps[key] = int(float(_instruction(wb, names)))
        except (TypeError, ValueError):
            gaps[key] = None  # not set for this program: no gap warning is made up
    # planned_shrinkage is only compared with what happened (adherence, the shrinkage coach); the
    # day's own figures never use it (owner: actuals replace it)
    found = {"interval": step if 1440 % step == 0 else 30, "required": demand, "planned_shrinkage": planned, **gaps,
             "languages": _languages(wb), "previous_saturday": previous}
    while len(_CACHE) >= KEEP:
        _CACHE.pop(next(iter(_CACHE)))
    _CACHE[key] = found
    return found


def planned(week: Dict[str, Any], d: int, name: str) -> List[Dict[str, Any]]:
    """A person's planned breaks on day ``d`` in time order: idx, kind, start (minutes from
    that day's midnight; after midnight reads +24h) and minutes."""
    span = next((shift_span(a["days"][d]) for a in week.get("associates", []) if a["name"] == name), None)
    if not span:
        return []
    found = []
    for b in week.get("breaks", []):
        m = _minute(b["start"])
        if b["associate"] == name and b["day"] == DAYS[d] and m is not None:
            found.append({"kind": b["kind"], "start": m + (1440 if m < span[0] else 0), "minutes": int(b["minutes"])})
    return [{"idx": i, **b} for i, b in enumerate(sorted(found, key=lambda b: b["start"]))]


SAME_WEEK = "same week"  # day_view's default for the day before: this week's previous column (the input's tab on Sunday)


def _yesterday(week: Dict[str, Any], day: int, before: Any) -> Optional[Tuple[Dict[str, Any], int]]:
    """Where the day before is read from: (week, column), or None for the input's Previous week tab."""
    if before == SAME_WEEK:
        return (week, day - 1) if day > 0 else None
    return before


def _segments(week: Dict[str, Any], inputs: Dict[str, Any], day: int, before: Any = SAME_WEEK) -> List[Dict[str, Any]]:
    """Shift pieces touching this day: the day's own shifts and the previous day's overnight ones (from
    ``before``: the (week, column) holding the day before, or None for the input's carry-in tab)."""
    out = []
    for offset in (-1, 0):
        source, d = (week, day) if offset == 0 else (_yesterday(week, day, before) or (None, None))
        if source is None:  # no schedule for the day before: overnight shifts from the Previous week scheduled tab
            for p in inputs.get("previous_saturday", []):
                span = shift_span(p["shift"])
                if span and span[1] > 1440:
                    out.append({"name": p["name"], "slot": "", "language": p["language"], "offset": -1,
                                "called_in": False, "label": p["shift"],
                                "start": span[0] - 1440, "end": span[1] - 1440, "planned": []})
            continue
        for a in source.get("associates", []):
            span = shift_span(a["days"][d])
            if not span:
                continue
            start, end = span[0] + 1440 * offset, span[1] + 1440 * offset
            if end <= 0 or start >= 1440:
                continue
            out.append({"name": a["name"], "slot": a.get("slot", ""), "language": a["language"], "offset": offset,
                        "called_in": False, "label": a["days"][d],
                        "start": start, "end": end,
                        "planned": [{**b, "start": b["start"] + 1440 * offset} for b in planned(source, d, a["name"])]})
    return out


CALLED_IN = "Called in"  # a day off cancelled: the person works a shift from the Shift Library


def pattern_breaks(week: Dict[str, Any], d: int, label: str) -> List[Dict[str, Any]]:
    """The breaks the plan gives this shift: copied from someone working the same shift that day,
    else on another day of the week; none when nobody works it."""
    span = shift_span(label)
    if not span:
        return []
    order = [d] + [x for x in range(7) if x != d]
    for day in order:
        for a in week.get("associates", []):
            if _norm(a["days"][day]) != _norm(label):
                continue
            mine = planned(week, day, a["name"])
            if mine:
                own = shift_span(a["days"][day])
                return [{**b, "start": span[0] + (b["start"] - own[0])} for b in mine]
    return []


def _called_in(week: Dict[str, Any], day: int, activities: Optional[Dict[Tuple[int, str], List[Dict[str, Any]]]],
               existing: set, before: Any = SAME_WEEK) -> List[Dict[str, Any]]:
    """Segments for people called in on a day off (this day, or yesterday's overnight shift)."""
    out = []
    for (offset, name), acts in (activities or {}).items():
        source, d = (week, day) if offset == 0 else (_yesterday(week, day, before) or (None, None))
        if source is None or (offset, name) in existing:
            continue
        people = {a["name"]: a for a in source.get("associates", [])}
        for a in acts:
            if a["kind"] != CALLED_IN:
                continue
            label = a.get("label") or f"{_hm(a['start'])} - {_hm(a['end'])}"
            start, end = a["start"] + 1440 * offset, a["end"] + 1440 * offset
            if end <= 0 or start >= 1440:
                continue
            person = people.get(name, {})
            out.append({"name": name, "slot": person.get("slot", ""), "language": person.get("language", ""),
                        "offset": offset, "called_in": True, "label": f"{label} (called in)", "start": start, "end": end,
                        "planned": [{**b, "start": b["start"] + 1440 * offset} for b in pattern_breaks(source, d, label)]})
    return out


def _away(seg: Dict[str, Any], mark: Optional[Dict[str, Any]]) -> Tuple[int, int]:
    """The (from, to) minutes this person is off the floor in this segment; (0, 0) when present."""
    if not mark or mark.get("status", PRESENT) == PRESENT:
        return (0, 0)
    status = mark["status"]
    shift = 1440 * seg["offset"]
    lo = mark.get("from")
    hi = mark.get("to")
    lo = seg["start"] if lo is None else lo + shift
    hi = seg["end"] if hi is None else hi + shift
    if status == "Late":
        lo = seg["start"]
    if status == "Left early":
        hi = seg["end"]
    if status not in TIMED:
        lo, hi = seg["start"], seg["end"]
    return (max(lo, seg["start"]), min(hi, seg["end"]))


def check_break(week: Dict[str, Any], day: int, offset: int, name: str, idx: int, start: int,
                actual: Dict[Tuple[int, str, int], int], inputs: Optional[Dict[str, Any]] = None,
                segment: Optional[Dict[str, Any]] = None) -> None:
    """Refuse a break that leaves the shift, overlaps another break, or is not on a 5-minute step.
    ``start`` is minutes from the shift's own day's midnight. ``segment`` (start, end, label,
    planned) is given for a shift that is not in the plan (a day off called in)."""
    if start % STEP:
        raise BreakRefused("Breaks move in 5-minute steps.")
    if segment is not None:
        seg = segment
    else:
        segs = [s for s in _segments(week, inputs or {"previous_saturday": []}, day + offset if offset == -1 else day)
                if s["name"] == name and s["offset"] == 0]
        if not segs:
            raise BreakRefused(f"{name} has no shift that day.")
        seg = segs[0]
    mine = {b["idx"]: b for b in seg["planned"]}
    if idx not in mine:
        raise BreakRefused("That break is not on the plan.")
    length = mine[idx]["minutes"]
    if start < seg["start"] or start + length > seg["end"]:
        raise BreakRefused(f"The break would fall outside {name}'s shift ({seg['label']}).")
    for other, b in mine.items():
        if other == idx:
            continue
        s = actual.get((offset, name, other), b["start"])
        if start < s + b["minutes"] and s < start + length:
            raise BreakRefused(f"The break would overlap {name}'s {b['kind']}.")


def day_view(week: Dict[str, Any], inputs: Dict[str, Any], day: int,
             attendance: Dict[Tuple[int, str], Dict[str, Any]],
             actual: Dict[Tuple[int, str, int], int], measure: str = "interval",
             activities: Optional[Dict[Tuple[int, str], List[Dict[str, Any]]]] = None,
             before: Any = SAME_WEEK) -> Dict[str, Any]:
    """The day's lanes, interval cells, language rows and tiles. ``before`` is where the day before is
    read from: (week, column), None for the input's carry-in tab, or this week's previous column. ``attendance``,
    ``actual`` and ``activities`` are keyed by (day offset, name[, break index]):
    0 for this day's shifts, -1 for the previous day's; times are minutes from that
    shift's own midnight. ``measure`` is "interval" (billable aux stays on the
    floor) or "sl". Activities: aux kinds (off the floor unless billable under
    interval compliance), "VTO" (off the floor) and "Overtime" (on the floor
    outside the shift; not in the plan)."""
    if measure not in MEASURES:
        raise ValueError(f"unknown measure {measure!r}")
    step = inputs["interval"]
    slots = 1440 // STEP
    now = [[] for _ in range(slots)]
    plan = [[] for _ in range(slots)]
    lanes: Dict[str, Dict[str, Any]] = {}
    pieces = []
    segments = _segments(week, inputs, day, before)
    segments += _called_in(week, day, activities, {(x["offset"], x["name"]) for x in segments}, before)
    for seg in segments:
        mark = attendance.get((seg["offset"], seg["name"])) or {}
        status = mark.get("status", PRESENT)
        billable = status in AUX and bool(mark.get("billable"))
        away = _away(seg, mark)
        counts_away = not (billable and measure == "interval")
        acts = [{**a, "start": a["start"] + 1440 * seg["offset"], "end": a["end"] + 1440 * seg["offset"]}
                for a in (activities or {}).get((seg["offset"], seg["name"]), [])]
        off = [(a["start"], a["end"]) for a in acts
               if a["kind"] == "VTO" or a["kind"] in EXTRA_BREAKS
               or (a["kind"] in AUX and not (a.get("billable") and measure == "interval"))]
        extra = [(a["start"], a["end"]) for a in acts if a["kind"] == "Overtime"]
        breaks = []
        for b in seg["planned"]:
            moved = actual.get((seg["offset"], seg["name"], b["idx"]))
            start = b["start"] if moved is None else moved + 1440 * seg["offset"]
            breaks.append({**b, "planned_start": b["start"], "start": start, "moved": moved is not None})
        for i in range(slots):
            t = i * STEP
            in_shift = seg["start"] <= t < seg["end"]
            if not in_shift and not any(a <= t < b for a, b in extra):
                continue
            if in_shift and not seg["called_in"] and not any(
                    b["planned_start"] <= t < b["planned_start"] + b["minutes"] for b in breaks):
                plan[i].append(seg)  # the plan: the schedule as made (a called-in day off is not in it)
            if status in ABSENT or any(b["start"] <= t < b["start"] + b["minutes"] for b in breaks):
                continue
            if (counts_away and away[0] <= t < away[1]) or any(a <= t < b for a, b in off):
                continue
            now[i].append(seg)
        pieces.append((seg, status, away, breaks, acts))
        lane = lanes.setdefault(seg["name"], {"name": seg["name"], "slot": seg["slot"], "language": seg["language"],
                                              "segments": []})
        lane["segments"].append({"offset": seg["offset"], "called_in": seg["called_in"], "label": seg["label"],
                                 "start": seg["start"],
                                 "end": seg["end"], "status": status, "billable": billable,
                                 "from": mark.get("from"), "to": mark.get("to"), "away": away, "breaks": breaks,
                                 "activities": acts, "with_whom": mark.get("with_whom", ""),
                                 "why": mark.get("why", ""), "with_dept": mark.get("with_dept", "")})

    def overlapping(lo: int, hi: int, t: int) -> bool:
        return lo < t + step and t < hi

    cells = []
    req = inputs["required"].get(day, {})
    for t in range(0, 1440, step):
        r = req.get(t, 0.0)
        idx = range(t // STEP, (t + step) // STEP)
        ticks = [len(now[i]) for i in idx]
        here = sum(ticks) / len(ticks)
        planned_here = sum(len(plan[i]) for i in idx) / len(idx)
        counts = {"absent": sum(1 for seg, st, away, br, ac in pieces if st in ABSENT
                                and overlapping(seg["start"], seg["end"], t)),
                  "aux": sum(1 for seg, st, away, br, ac in pieces
                             if (st in AUX and away[0] < away[1] and overlapping(away[0], away[1], t))
                             or any(a["kind"] in AUX and overlapping(a["start"], a["end"], t) for a in ac)),
                  "late_early": sum(1 for seg, st, away, br, ac in pieces if st in LATE_EARLY and away[0] < away[1]
                                    and overlapping(away[0], away[1], t)),
                  "breaks": sum(1 for seg, st, away, br, ac in pieces if st not in ABSENT
                                and (any(overlapping(b["start"], b["start"] + b["minutes"], t) for b in br)
                                     or any(a["kind"] in EXTRA_BREAKS and overlapping(a["start"], a["end"], t)
                                            for a in ac))),
                  "overtime": sum(1 for seg, st, away, br, ac in pieces
                                  if any(a["kind"] == "Overtime" and overlapping(a["start"], a["end"], t) for a in ac)),
                  "vto": sum(1 for seg, st, away, br, ac in pieces
                             if any(a["kind"] == "VTO" and overlapping(a["start"], a["end"], t) for a in ac))}
        cell = {"t": t, "required": r, "plan": round(planned_here, 1), "now": round(here, 1), "slots": ticks,
                "low": min(ticks), **counts}
        if r <= 0:
            cells.append({**cell, "need": 0, "pm": None, "plan_pm": None, "cls": "none"})
            continue
        pm = (here - r) * step / 60
        cls = "ok" if pm >= -1e-9 else ("warn" if here > r - 1 else "bad")
        cells.append({**cell, "need": math.ceil(r - 1e-9), "pm": round(pm, 2),
                      "plan_pm": round((planned_here - r) * step / 60, 2), "cls": cls})
    languages = []
    for rule in inputs["languages"]:
        eligible = {_norm(r["name"]) for r in inputs["languages"] if _norm(rule["name"]) in r["covers"]}
        row = []
        for t in range(0, 1440, step):
            if rule["start"] == rule["end"]:
                inside = True
            elif rule["start"] < rule["end"]:
                inside = rule["start"] <= t < rule["end"]
            else:
                inside = t >= rule["start"] or t < rule["end"]
            if not inside or day not in rule["days"]:
                row.append({"t": t, "count": None, "cls": "none"})
                continue
            count = min(sum(1 for seg in now[i] if _norm(seg["language"]) in eligible)
                        for i in range(t // STEP, (t + step) // STEP))
            cls = "bad" if count < rule["minimum"] else ("warn" if count == 0 else "lok")
            row.append({"t": t, "count": count, "cls": cls})
        languages.append({"name": rule["name"], "start": rule["start"], "end": rule["end"], "minimum": rule["minimum"],
                          "eligible": sorted(eligible), "cells": row,
                          "gaps": sum(1 for c in row if c["cls"] in ("bad", "warn"))})
    mine = [(seg, st) for seg, st, away, br, ac in pieces if seg["offset"] == 0]
    with_aux = sum(1 for seg, st, away, br, ac in pieces if seg["offset"] == 0
                   and (st in AUX or any(a["kind"] in AUX for a in ac)))
    absent = sum(1 for seg, st in mine if st in ABSENT)
    short = sum(-c["pm"] for c in cells if c["pm"] is not None and c["pm"] < 0)
    plan_short = sum(-c["plan_pm"] for c in cells if c["plan_pm"] is not None and c["plan_pm"] < 0)
    over = sum(c["pm"] for c in cells if c["pm"] is not None and c["pm"] > 0)
    tiles = {"planned": len(mine), "present": len(mine) - absent, "absent": absent,
             "late_early": sum(1 for seg, st in mine if st in LATE_EARLY),
             "aux": with_aux,
             "moved": sum(1 for seg, st, away, br, ac in pieces for b in br if b["moved"]),
             "overtime": sum(a["end"] - a["start"] for seg, st, away, br, ac in pieces for a in ac if a["kind"] == "Overtime"),
             "vto": sum(a["end"] - a["start"] for seg, st, away, br, ac in pieces for a in ac if a["kind"] == "VTO"),
             "called_in": sum(seg["end"] - seg["start"] for seg, st, away, br, ac in pieces
                              if seg["called_in"] and seg["offset"] == 0 and st not in ABSENT),
             "short_hours": round(short, 1), "plan_short_hours": round(plan_short, 1),
             "short_intervals": sum(1 for c in cells if c["pm"] is not None and c["pm"] < 0),
             "over_hours": round(over, 1), "language_gaps": sum(l["gaps"] for l in languages)}

    def first(lane):  # last night's people first, then by the day's start time
        own = [s["start"] for s in lane["segments"] if s["offset"] == 0]
        return (min(own) + 1440 if own else min(s["start"] for s in lane["segments"]), lane["name"])

    order = sorted(lanes.values(), key=first)
    return {"interval": step, "measure": measure, "lanes": order, "cells": cells, "languages": languages,
            "tiles": tiles}


def _hm(minute: int) -> str:
    minute %= 1440
    return f"{minute // 60:02d}:{minute % 60:02d}"


def advice(view: Dict[str, Any], inputs: Dict[str, Any], name: str, idx: int, start: int) -> Dict[str, Any]:
    """For moving break ``idx`` of ``name``'s shift that starts this day to ``start`` (minutes
    from this day's midnight): the gap warnings (the program's Break Absolute Minimum / Normal
    Maximum Gap Minutes, end of one break to start of the next), and up to three best starts
    (one per interval) that keep the breaks in order, off each other and within the gaps,
    ranked by the buffer of the tightest interval the break touches."""
    seg = next(s for l in view["lanes"] if l["name"] == name for s in l["segments"] if s["offset"] == 0)
    me = next(b for b in seg["breaks"] if b["idx"] == idx)
    others = [b for b in seg["breaks"] if b["idx"] != idx]
    lo_gap, hi_gap = inputs.get("gap_min"), inputs.get("gap_max")

    def gaps(m: int) -> List[Tuple[str, str, int]]:
        seq = sorted([(b["start"], b["minutes"], b["kind"]) for b in others] + [(m, me["minutes"], me["kind"])])
        return [(a[2], b[2], b[0] - (a[0] + a[1])) for a, b in zip(seq, seq[1:])]

    warnings = []
    for a, b, gap in gaps(start):
        if lo_gap is not None and gap < lo_gap:
            warnings.append(f"{a} ends and {b} starts {gap} minutes apart; this program's minimum gap "
                            f"between breaks is {lo_gap} minutes.")
        elif hi_gap is not None and gap > hi_gap:
            warnings.append(f"{a} ends and {b} starts {gap} minutes apart; this program's normal maximum "
                            f"gap between breaks is {hi_gap} minutes.")
    step = view["interval"]
    base = [x for c in view["cells"] for x in c["slots"]]
    required = {c["t"]: c["required"] for c in view["cells"]}
    counts_away = not (seg["billable"] and view["measure"] == "interval")

    def free(t: int) -> bool:  # on the floor at t, this break aside
        if not seg["start"] <= t < seg["end"] or (counts_away and seg["away"][0] <= t < seg["away"][1]):
            return False
        return not any(b["start"] <= t < b["start"] + b["minutes"] for b in others)

    def tightest(m: int) -> Optional[float]:
        slots = list(base)
        for t in range(me["start"], me["start"] + me["minutes"], STEP):
            if 0 <= t < 1440 and free(t):
                slots[t // STEP] += 1
        for t in range(m, m + me["minutes"], STEP):
            if 0 <= t < 1440 and free(t):
                slots[t // STEP] -= 1
        found = []
        for t0 in range(m - m % step, m + me["minutes"], step):
            r = required.get(t0, 0.0)
            if r > 0:
                here = slots[t0 // STEP:(t0 + step) // STEP]
                found.append((sum(here) / len(here) - r) * step / 60)
        return min(found) if found else None

    before = [b for b in others if b["idx"] < idx]
    after = [b for b in others if b["idx"] > idx]
    lo = max([b["start"] + b["minutes"] + (lo_gap or 0) for b in before] + [seg["start"]])
    hi = min([b["start"] - (lo_gap or 0) - me["minutes"] for b in after] + [seg["end"] - me["minutes"]])
    best: Dict[int, Tuple[float, int]] = {}
    for m in range(lo + (-lo) % STEP, min(hi, 1440 - me["minutes"]) + 1, STEP):  # after midnight: next day's page
        if hi_gap is not None and any(g > hi_gap for _, _, g in gaps(m)):
            continue
        worst = tightest(m)
        key = m - m % step
        if worst is not None and (key not in best or worst > best[key][0]):
            best[key] = (worst, m)
    ranked = sorted(best.values(), key=lambda x: (-x[0], x[1]))[:3]
    return {"warnings": warnings, "fits": [{"start": _hm(m), "end": _hm(m + me["minutes"]), "buffer": round(w, 2)}
                                           for w, m in ranked]}


def _short(kind: str) -> str:
    """Break 1 -> B1, Lunch -> L, anything else: its initials."""
    if kind.lower().startswith("break "):
        return "B" + kind.split()[-1]
    return "L" if kind.lower() == "lunch" else "".join(w[0] for w in kind.split()).upper()[:3]


def board(view: Dict[str, Any]) -> List[Dict[str, Any]]:
    """The interval board: one row per interval with demand or people on shift: buffer, status,
    the breaks starting in it (lunch and short breaks), late / early / aux starting in it, and
    the counts (unplanned leave, on break, aux)."""
    step = view["interval"]
    rows = []
    for c in view["cells"]:
        t = c["t"]
        if c["required"] <= 0 and not c["plan"] and not c["now"]:
            continue
        lunch, short, events = [], [], []
        for lane in view["lanes"]:
            for seg in lane["segments"]:
                if seg["status"] in ABSENT:
                    continue  # not on shift today: nothing to plot
                for b in seg["breaks"]:
                    if t <= b["start"] < t + step:
                        chip = {"name": lane["name"], "slot": lane["slot"], "offset": seg["offset"], "idx": b["idx"],
                                "kind": b["kind"], "short": _short(b["kind"]), "start": _hm(b["start"]),
                                "planned": _hm(b["planned_start"]), "moved": b["moved"], "minutes": b["minutes"]}
                        (lunch if b["kind"].lower() == "lunch" else short).append(chip)
                for a in seg.get("activities", []):  # added on the day: shown with the planned ones
                    if a["kind"] in EXTRA_BREAKS and t <= a["start"] < t + step:
                        chip = {"name": lane["name"], "slot": lane["slot"], "offset": seg["offset"], "idx": None,
                                "kind": a["kind"], "short": _short(a["kind"]), "start": _hm(a["start"]),
                                "planned": _hm(a["start"]), "moved": False, "minutes": a["end"] - a["start"],
                                "added": True, "id": a.get("id")}
                        (lunch if a["kind"] == "Lunch" else short).append(chip)
                lo, hi = seg["away"]
                st = seg["status"]
                if st == "Late" and t <= seg["start"] < t + step:
                    events.append({"name": lane["name"], "text": f"late login, in at {_hm(hi)}"})
                elif st == "Left early" and t <= lo < t + step:
                    events.append({"name": lane["name"], "text": f"early leave at {_hm(lo)}"})
                elif st in AUX and lo < hi and t <= lo < t + step:
                    events.append({"name": lane["name"], "text": f"{st} {_hm(lo)} to {_hm(hi)}, "
                                                                 f"{'billable' if seg['billable'] else 'non-billable'}"})
                for a in seg.get("activities", []):
                    if t <= a["start"] < t + step and a["kind"] not in EXTRA_BREAKS:
                        text = (f"called in on a day off, {_hm(a['start'])} to {_hm(a['end'])}" if a["kind"] == CALLED_IN else
                                f"overtime {_hm(a['start'])} to {_hm(a['end'])}" if a["kind"] == "Overtime" else
                                f"VTO {_hm(a['start'])} to {_hm(a['end'])}" if a["kind"] == "VTO" else
                                f"{a['kind'].lower()} {_hm(a['start'])} to {_hm(a['end'])}, "
                                f"{'billable' if a.get('billable') else 'non-billable'}")
                        events.append({"name": lane["name"], "text": text})
        if c["pm"] is None:
            state, label, why = "none", "No demand", ""
        else:
            room = c["low"] - c["need"]
            if room >= 1:
                state, label, why = "good", "Good buffer", f"room for {room} more on break"
            elif room >= 0:
                state, label, why = "tight", "Tight", "on target, no room for another break"
            else:
                more = c["need"] - c["low"]
                state, label, why = "short", "Short", f"{more} more {'person' if more == 1 else 'people'} needed at the lowest point"
        langs = [{"name": l["name"], "count": cell["count"], "cls": cell["cls"]}
                 for l in view["languages"] for cell in l["cells"] if cell["t"] == t and cell["count"] is not None]
        rows.append({**c, "end": _hm(t + step), "start": _hm(t), "state": state, "label": label, "why": why,
                     "cover": round(c["now"] / c["required"] * 100) if c["required"] > 0 else None,
                     "lunch": sorted(lunch, key=lambda x: x["start"]), "short": sorted(short, key=lambda x: x["start"]),
                     "events": events, "langs": langs})
    return rows


# ---------------------------------------------------------------- meetings, overtime and VTO (Phase O)
MEETING_MARGIN = 10  # minutes kept clear before and after a break
OVERTIME_MAX = 120  # minutes of overtime offered next to a shift


def _own(view: Dict[str, Any], name: str) -> Optional[Dict[str, Any]]:
    return next((s for l in view["lanes"] if l["name"] == name for s in l["segments"] if s["offset"] == 0), None)


def _busy(seg: Dict[str, Any], t: int, margin: int = 0) -> bool:
    """On a break (with a margin), away, or in another activity at minute ``t``."""
    if seg["status"] in ABSENT or not seg["start"] <= t < seg["end"]:
        return True
    if seg["away"][0] <= t < seg["away"][1]:
        return True
    if any(b["start"] - margin <= t < b["start"] + b["minutes"] + margin for b in seg["breaks"]):
        return True
    return any(a["start"] <= t < a["end"] for a in seg.get("activities", []) if a["kind"] not in ("Overtime", CALLED_IN))


def _tightest(view: Dict[str, Any], slots: List[float], lo: int, hi: int) -> Optional[float]:
    step = view["interval"]
    found = []
    for c in view["cells"]:
        if c["required"] > 0 and c["t"] < hi and lo < c["t"] + step:
            here = slots[c["t"] // STEP:(c["t"] + step) // STEP]
            found.append((sum(here) / len(here) - c["required"]) * step / 60)
    return min(found) if found else None


def meeting_slots(view: Dict[str, Any], names: List[str], minutes: int, earliest: int, latest: int,
                  billable: bool = False, top: int = 5) -> List[Dict[str, Any]]:
    """The best start times (one per interval) when every person is on shift, present, not on a break
    (10 minutes either side), away or in another activity, ranked by the floor's buffer at its
    tightest once they are off the queue (a billable session under interval compliance keeps them on)."""
    segs = {}
    for name in names:
        seg = _own(view, name)
        if seg is None:
            raise ValueError(f"{name} has no shift this day.")
        segs[name] = seg
    if minutes <= 0 or minutes % STEP:
        raise ValueError("Pick a length in 5-minute steps.")
    base = [x for c in view["cells"] for x in c["slots"]]
    off_floor = 0 if (billable and view["measure"] == "interval") else len(names)
    step = view["interval"]
    best: Dict[int, Dict[str, Any]] = {}
    for start in range(earliest + (-earliest) % STEP, latest - minutes + 1, STEP):
        if any(_busy(seg, t, MEETING_MARGIN) for seg in segs.values() for t in range(start, start + minutes, STEP)):
            continue
        trial = list(base)
        for t in range(start, start + minutes, STEP):
            if 0 <= t < 1440:
                trial[t // STEP] -= off_floor
        tight = _tightest(view, trial, start, start + minutes)
        score = tight if tight is not None else float("inf")
        key = start - start % step
        if key not in best or score > best[key]["score"]:
            best[key] = {"start": start, "end": start + minutes, "tightest": tight, "score": score}
    ranked = sorted(best.values(), key=lambda x: (-x["score"], x["start"]))[:top]
    return [{k: v for k, v in x.items() if k != "score"} for x in ranked]


def _either_side(week: Dict[str, Any], day: int, name: str, near: Optional[Dict[str, Any]]) -> Tuple[Any, Any]:
    """The person's working spans the day before and after: from ``near`` (worked out across the week's
    edges, with call-ins and overtime) when given, else from this week's schedule."""
    if near is not None:
        return near["prev"].get(name), near["next"].get(name)
    a = next(x for x in week.get("associates", []) if x["name"] == name)
    return (shift_span(a["days"][day - 1]) if day > 0 else None, shift_span(a["days"][day + 1]) if day < 6 else None)


def overtime_offers(view: Dict[str, Any], week: Dict[str, Any], day: int, rest_hours: float = 12,
                    after: int = 0, only: Optional[int] = None, near: Optional[Dict[str, Any]] = None
                    ) -> List[Dict[str, Any]]:
    """For each short interval (from ``after``), or only the interval at ``only``: people present that day
    whose shift ends just before it or starts just after it, offered to stay on or come in early (at most
    2 hours), keeping the rest gap to the previous and next day's work; the shortest offers first."""
    step = view["interval"]
    rest = int(rest_hours * 60)
    out = []
    for c in view["cells"]:
        if only is not None:
            if c["t"] != only:
                continue
        elif c["pm"] is None or c["pm"] >= 0 or c["t"] < after:
            continue
        t, offers = c["t"], []
        for a in week.get("associates", []):
            seg = _own(view, a["name"])
            if seg is None or seg["status"] != PRESENT or any(x["kind"] == "VTO" for x in seg.get("activities", [])):
                continue
            prv, nxt = _either_side(week, day, a["name"], near)
            if seg["end"] <= t and t + step - seg["end"] <= OVERTIME_MAX:
                lo, hi = seg["end"], t + step
            elif seg["start"] >= t + step and seg["start"] - t <= OVERTIME_MAX:
                lo, hi = t, seg["start"]
            else:
                continue
            if (nxt and nxt[0] + 1440 - hi < rest) or (prv and lo + 1440 - prv[1] < rest):
                continue
            if any(x["start"] < hi and lo < x["end"] for x in seg.get("activities", [])):
                continue
            side = "after" if lo == seg["end"] else "before"
            if side == "after":  # the longest the rules allow: 2 hours, the rest gap to tomorrow
                most = OVERTIME_MAX if not nxt else min(OVERTIME_MAX, nxt[0] + 1440 - rest - seg["end"])
            else:
                most = OVERTIME_MAX if not prv else min(OVERTIME_MAX, seg["start"] - (prv[1] - 1440 + rest))
            offers.append({"name": a["name"], "start": lo, "end": hi, "side": side, "max": most,
                           "text": f"stay {_hm(lo)} to {_hm(hi)}" if side == "after" else f"start {_hm(lo)} instead of {_hm(hi)}"})
        offers.sort(key=lambda o: (o["end"] - o["start"], o["name"]))
        out.append({"t": t, "buffer": c["pm"], "short": max(0, math.ceil(c["required"] - c["now"] - 1e-9)),
                    "offers": offers[:3] if only is None else offers[:6]})
    return out


def vto_offers(view: Dict[str, Any], after: int = 0, top: int = 6) -> List[Dict[str, Any]]:
    """Intervals comfortably above need (from ``after``): people present whose shift ends within 2 hours,
    offered to leave at the start of the interval, while every interval until they would have left keeps
    its need at every 5 minutes and every language keeps its minimum (and anyone it had)."""
    step = view["interval"]
    need = {c["t"]: c["need"] for c in view["cells"] if c["required"] > 0}
    base = [x for c in view["cells"] for x in c["slots"]]
    lang_of = {l["name"]: _norm(l["language"]) for l in view["lanes"]}
    out = []
    for c in view["cells"]:
        t = c["t"]
        if c["pm"] is None or c["low"] - c["need"] < 1 or t < after or len(out) >= top:
            continue
        candidates = sorted((seg["end"], l["name"]) for l in view["lanes"] for seg in l["segments"]
                            if seg["offset"] == 0 and seg["status"] == PRESENT and seg["start"] < t < seg["end"]
                            and seg["end"] - t <= OVERTIME_MAX and not seg.get("activities"))
        trial, gone, taken = list(base), {}, []
        for end, name in candidates:
            ok = all(trial[x // STEP] - 1 >= need.get(x - x % step, 0) for x in range(t, min(end, 1440), STEP))
            for lang in view["languages"]:
                if lang_of.get(name) not in lang["eligible"]:
                    continue
                for cell in lang["cells"]:
                    if cell["count"] is not None and t <= cell["t"] < end:
                        left = cell["count"] - gone.get(lang["name"], 0) - 1
                        if left < max(lang["minimum"], 1 if cell["count"] else 0):
                            ok = False
            if not ok:
                continue
            for x in range(t, min(end, 1440), STEP):
                trial[x // STEP] -= 1
            for lang in view["languages"]:
                if lang_of.get(name) in lang["eligible"]:
                    gone[lang["name"]] = gone.get(lang["name"], 0) + 1
            taken.append({"name": name, "start": t, "end": end, "text": f"leave at {_hm(t)} instead of {_hm(end)}"})
        if taken:
            out.append({"t": t, "buffer": c["pm"], "room": c["low"] - c["need"], "offers": taken})
    return out


# ---------------------------------------------------------------- the autopilot for the rest of the day's breaks
def replan(view: Dict[str, Any], inputs: Dict[str, Any], now: int = 0, rounds: int = 400) -> Dict[str, Any]:
    """Re-plot the breaks not yet started (people on shift and present, this day's shifts) to lift the
    tightest interval, then to cut the hours short; each break stays inside its shift, in its order,
    on a 5-minute step, clear of the person's other breaks, away time and activities, with the gaps to
    its neighbours within the program's rules, and no language drops below its minimum (or loses its
    last person). Greedy: the best single move each round, until nothing improves. Returns the moves
    and the day before and after (tightest buffer and hours short, as the day page shows them)."""
    step = view["interval"]
    per = step // STEP
    cells = view["cells"]
    req = [c["required"] for c in cells]
    slots = [x for c in cells for x in c["slots"]]
    sums = [sum(slots[k * per:(k + 1) * per]) for k in range(len(cells))]
    lo_gap, hi_gap = inputs.get("gap_min"), inputs.get("gap_max")
    interval_mode = view["measure"] == "interval"

    def pm(k: int, total: float) -> float:
        return round((total / per - req[k]) * step / 60, 2)

    def score(totals: List[float]) -> Tuple[float, float]:
        values = [pm(k, totals[k]) for k in range(len(cells)) if req[k] > 0]
        if not values:
            return (0.0, 0.0)
        return (min(values), -round(sum(-v for v in values if v < 0), 1))

    people = [(lane, seg) for lane in view["lanes"] for seg in lane["segments"]
              if seg["offset"] == 0 and seg["status"] not in ABSENT]
    starts = {(lane["name"], b["idx"]): b["start"] for lane, seg in people for b in seg["breaks"]}
    original = dict(starts)

    def free(seg: Dict[str, Any], t: int) -> bool:  # on the floor at t, breaks aside
        if not seg["start"] <= t < seg["end"]:
            return False
        if not (seg["billable"] and interval_mode) and seg["away"][0] <= t < seg["away"][1]:
            return False
        return not any(a["start"] <= t < a["end"] for a in seg.get("activities", [])
                       if a["kind"] == "VTO" or a["kind"] in EXTRA_BREAKS
                       or (a["kind"] in AUX and not (a.get("billable") and interval_mode)))

    # languages: per 5 minutes, people on the floor who count for each language (inside its hours)
    langs = []
    for lang in view["languages"]:
        inside = {c["t"] for c in lang["cells"] if c["count"] is not None}
        counts = [0] * len(slots)
        langs.append({"eligible": set(lang["eligible"]), "minimum": lang["minimum"], "inside": inside, "counts": counts})
    lang_of = {lane["name"]: _norm(lane["language"]) for lane in view["lanes"]}
    for lane in view["lanes"]:
        for seg in lane["segments"]:
            for i in range(len(slots)):
                t = i * STEP
                on = (seg["status"] not in ABSENT and free(seg, t)
                      and not any(b["start"] <= t < b["start"] + b["minutes"] for b in seg["breaks"]))
                if on:
                    for L in langs:
                        if lang_of[lane["name"]] in L["eligible"]:
                            L["counts"][i] += 1

    def window(start: int, minutes: int) -> range:
        return range(max(start, 0), min(start + minutes, 1440), STEP)

    def positions(lane, seg, b) -> List[int]:
        name = lane["name"]
        others = sorted(((starts[(name, o["idx"])], o) for o in seg["breaks"] if o["idx"] != b["idx"]), key=lambda x: x[0])
        before = [(s, o) for s, o in others if o["idx"] < b["idx"]]
        after = [(s, o) for s, o in others if o["idx"] > b["idx"]]
        lo = max(seg["start"], now)
        hi = seg["end"] - b["minutes"]
        if before:
            prev_end = before[-1][0] + before[-1][1]["minutes"]
            lo = max(lo, prev_end + (lo_gap or 0))
            if hi_gap is not None:
                hi = min(hi, prev_end + hi_gap)
        if after:
            nxt = after[0][0]
            hi = min(hi, nxt - (lo_gap or 0) - b["minutes"])
            if hi_gap is not None:
                lo = max(lo, nxt - hi_gap - b["minutes"])
        out = []
        for s in range(lo + (-lo) % STEP, hi + 1, STEP):
            if s == starts[(name, b["idx"])]:
                continue
            if not all(free(seg, t) for t in range(s, s + b["minutes"], STEP)):
                continue
            out.append(s)
        return out

    def delta(lane, seg, b, new: int) -> Dict[int, int]:
        old = starts[(lane["name"], b["idx"])]
        change: Dict[int, int] = {}
        for t in window(old, b["minutes"]):
            if free(seg, t) and not new <= t < new + b["minutes"]:
                change[t // STEP] = change.get(t // STEP, 0) + 1
        for t in window(new, b["minutes"]):
            if free(seg, t) and not old <= t < old + b["minutes"]:
                change[t // STEP] = change.get(t // STEP, 0) - 1
        return change

    def languages_hold(lane, change: Dict[int, int]) -> bool:
        mine = lang_of[lane["name"]]
        for L in langs:
            if mine not in L["eligible"]:
                continue
            for i, d in change.items():
                if d < 0 and (i * STEP - (i * STEP) % step) in L["inside"]:
                    if L["counts"][i] + d < max(L["minimum"], 1):
                        return False
        return True

    before_score = score(sums)
    current = before_score
    for _ in range(rounds):
        tight = {k for k in range(len(cells)) if req[k] > 0 and pm(k, sums[k]) < step / 60}  # under one person spare
        best = None
        for lane, seg in people:
            for b in seg["breaks"]:
                key = (lane["name"], b["idx"])
                start = starts[key]
                if start < now or not any((start - start % step) // step <= k <= (start + b["minutes"] - 1) // step
                                          for k in tight):
                    continue
                for new in positions(lane, seg, b):
                    change = delta(lane, seg, b, new)
                    if not change or not languages_hold(lane, change):
                        continue
                    totals = list(sums)
                    for i, d in change.items():
                        totals[i // per] += d
                    candidate = score(totals)
                    rank = (candidate, -abs(new - original[key]))
                    if candidate > current and (best is None or rank > best[0]):
                        best = (rank, lane, seg, b, new, change, totals)
        if best is None:
            break
        _, lane, seg, b, new, change, totals = best
        for i, d in change.items():
            for L in langs:
                if lang_of[lane["name"]] in L["eligible"]:
                    L["counts"][i] += d
        starts[(lane["name"], b["idx"])] = new
        sums = totals
        current = score(sums)
    moves = []
    for lane, seg in people:
        for b in seg["breaks"]:
            key = (lane["name"], b["idx"])
            if starts[key] != original[key]:
                moves.append({"name": lane["name"], "idx": b["idx"], "kind": b["kind"], "minutes": b["minutes"],
                              "from": original[key], "to": starts[key]})
    return {"moves": moves, "before": {"tightest": before_score[0], "short_hours": -before_score[1]},
            "after": {"tightest": current[0], "short_hours": -current[1]}}


def dayoff_offers(view: Dict[str, Any], week: Dict[str, Any], day: int, rest_hours: float, t: int,
                  top: int = 5, near: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """People off this day who could be called in on a Shift Library shift covering the interval at
    ``t``, keeping the rest gap to the day before and after; for each person the shift covering the most
    short intervals (then the one centred on ``t``); those covering most first."""
    step = view["interval"]
    rest = int(rest_hours * 60)
    short = {c["t"] for c in view["cells"] if c["pm"] is not None and c["pm"] < 0}
    working = {l["name"] for l in view["lanes"] for s in l["segments"] if s["offset"] == 0}
    out = []
    for a in week.get("associates", []):
        if _norm(a["days"][day]) != "off" or a["name"] in working:
            continue
        prv, nxt = _either_side(week, day, a["name"], near)
        best = None
        for label in week.get("shifts", []):
            span = shift_span(label)
            if not span or not (span[0] <= t and t + step <= span[1]):
                continue
            if rest and ((prv and span[0] + 1440 - prv[1] < rest) or (nxt and nxt[0] + 1440 - span[1] < rest)):
                continue
            covers = sum(1 for k in short if span[0] <= k and k + step <= span[1])
            rank = (covers, -abs((span[0] + span[1]) / 2 - (t + step / 2)))
            if best is None or rank > best[0]:
                best = (rank, label, span, covers)
        if best:
            _, label, span, covers = best
            out.append({"name": a["name"], "slot": a.get("slot", ""), "language": a["language"], "shift": label,
                        "start": span[0], "end": span[1], "covers": covers})
    out.sort(key=lambda o: (-o["covers"], o["name"]))
    return out[:top]
