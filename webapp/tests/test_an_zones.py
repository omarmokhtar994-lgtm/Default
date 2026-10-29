# © 2026 Omar Mokhtar. All rights reserved.
"""Dated export regressions split from coverage tests so each has its own dependencies."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from webapp.versions import zoned_copy, zoned_label


class ZonedExportRegressions(unittest.TestCase):
    def test_week_uses_each_days_offset(self):
        with tempfile.TemporaryDirectory() as root:
            src, dst = Path(root) / "source.xlsx", Path(root) / "out.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Schedule"
            ws.append(["SF Name", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"])
            ws.append(["Associate 001"] + ["08:00 - 17:00"] * 7)
            wb.save(src)
            zoned_copy(src, dst, -180, on=date(2026, 10, 25), from_zone="Africa/Cairo", to_zone="Europe/London")
            ws = load_workbook(dst)["Schedule"]
            self.assertEqual(ws.cell(2, 2).value, "05:00 - 14:00")
            self.assertEqual(ws.cell(2, 7).value, "06:00 - 15:00")

    def test_transition_uses_both_shift_endpoints(self):
        self.assertEqual(zoned_label("00:30 - 04:30", 420, on=date(2026, 11, 1),
                                     from_zone="America/New_York", to_zone="Africa/Cairo"), "06:30 - 11:30")

    def test_midnight_label_still_marks_original_day(self):
        self.assertEqual(zoned_label("08:00 - 17:00", -600, on=date(2026, 10, 11),
                                     from_zone="Africa/Cairo", to_zone="America/Los_Angeles"),
                         "22:00 - 07:00 (day before)")

    def test_nonexistent_source_time_is_not_silently_shifted(self):
        with self.assertRaisesRegex(ValueError, "does not exist"):
            zoned_label("02:30 - 11:30", 0, on=date(2026, 3, 8),
                        from_zone="America/New_York", to_zone="UTC")

if __name__ == "__main__":
    unittest.main()
