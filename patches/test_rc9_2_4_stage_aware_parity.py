#!/usr/bin/env python3
"""B-1/B-2: the release parity gate must know which stage produced the artifact.

Before this, a BEFORE_BREAKS_ONLY run published no canonical metric surface at
all. The parity gate read "did not publish" as "disagrees" and returned 41
mismatches (every field `engine: null`, `MISSING_CANONICAL_METRIC`), so every
skeleton-only run failed with rc=4 / FAIL_METRIC_PARITY on a schedule that had
nothing wrong with it.

The fix publishes the Stage-1 surface rather than exempting the stage from the
gate, so these tests are written from both sides: the gate must now pass a
correct skeleton run AND must still fail when the engine publishes nothing,
when the two sides disagree about the stage, or when a skeleton artifact claims
an after-break coverage figure that differs from its before-break twin.
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ENGINE = ROOT / "engine" / "_tools" / "l632_universal_scheduler.py"
RUNNER = ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py"
VALIDATOR = ROOT / "engine" / "tools" / "independent_validator.py"
REPORTER = ROOT / "engine" / "production" / "phase_c_quality_report.py"
sys.path.insert(0, str(ROOT / "engine" / "_tools"))

import canonical_metrics as CM  # noqa: E402


def load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def full_surface(**overrides):
    """A complete raw surface, so a test fails for its own reason only.

    Keys are written in the engine's own vocabulary (the first alias), not the
    canonical name. Two canonical fields - the overage concentration and
    consecutive-interval metrics - are named differently on both sides, so a
    dict keyed by canonical name would silently leave them unpopulated and
    every test would fail on the fixture rather than on its subject.
    """
    surface = {CM.ALIASES[name][0]: 1 for name in CM.PARITY_FIELDS}
    surface["max_concurrent_break_ratio_observed"] = 0.0
    for name in ("active_intervals",
                 "before_100", "before_90", "before_80",
                 "after_100", "after_90", "after_80",
                 "before_target", "after_target",
                 "before_floor", "after_floor"):
        surface[CM.ALIASES[name][0]] = 100
    for name, value in overrides.items():
        surface[CM.ALIASES.get(name, (name,))[0]] = value
    return surface


class StageLabelling(unittest.TestCase):
    def test_no_stage_argument_leaves_the_surface_exactly_as_before(self):
        surface = CM.canonicalize_metrics({"active_intervals": 5}, "engine")
        self.assertNotIn("stage", surface)
        self.assertNotIn("break_stage_executed", surface)
        self.assertEqual(surface["active_intervals"], 5)

    def test_before_breaks_stage_declares_that_no_break_stage_ran(self):
        surface = CM.canonicalize_metrics({}, "engine", stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(surface["stage"], "BEFORE_BREAKS_ONLY")
        self.assertIs(surface["break_stage_executed"], False)

    def test_full_schedule_stage_declares_that_the_break_stage_ran(self):
        surface = CM.canonicalize_metrics({}, "engine", stage="FULL_SCHEDULE")
        self.assertEqual(surface["stage"], "FULL_SCHEDULE")
        self.assertIs(surface["break_stage_executed"], True)

    def test_an_unknown_stage_is_named_unknown_and_never_guessed(self):
        surface = CM.canonicalize_metrics({}, "engine", stage="SOMETHING_ELSE")
        self.assertEqual(surface["stage"], "UNRECOGNIZED_STAGE")
        self.assertIsNone(surface["break_stage_executed"])

    def test_artifact_role_maps_to_the_stage_that_produced_it(self):
        self.assertEqual(
            CM.stage_for_artifact_role("BEST_BEFORE_BREAKS"), "BEFORE_BREAKS_ONLY")
        self.assertEqual(
            CM.stage_for_artifact_role("FINAL_AFTER_BREAKS"), "FULL_SCHEDULE")
        self.assertIsNone(CM.stage_for_artifact_role(None))
        self.assertIsNone(CM.stage_for_artifact_role("SOMETHING_ELSE"))

    def test_stage_is_recorded_on_the_published_surface_not_only_in_the_verdict(self):
        result = CM.compare_metric_surfaces(
            full_surface(), full_surface(),
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(result["engine"]["stage"], "BEFORE_BREAKS_ONLY")
        self.assertEqual(result["validator"]["stage"], "BEFORE_BREAKS_ONLY")


class ParityAtTheSkeletonStage(unittest.TestCase):
    def test_a_correct_skeleton_run_passes_the_gate(self):
        result = CM.compare_metric_surfaces(
            full_surface(), full_surface(),
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(result["status"], "PASS", result["mismatches"][:5])
        self.assertEqual(result["mismatch_count"], 0)
        self.assertTrue(result["before_after_identity_checked"])

    def test_every_canonical_field_is_still_compared_at_the_skeleton_stage(self):
        # The fix must not have narrowed the field set for skeleton runs.
        result = CM.compare_metric_surfaces(
            full_surface(), full_surface(),
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(list(result["compared_fields"]), list(CM.PARITY_FIELDS))

    def test_a_value_mismatch_still_fails_at_the_skeleton_stage(self):
        result = CM.compare_metric_surfaces(
            full_surface(), full_surface(active_intervals=99),
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(result["status"], "FAIL")
        self.assertIn("VALUE_MISMATCH", {m["reason"] for m in result["mismatches"]})

    def test_a_missing_field_still_fails_at_the_skeleton_stage(self):
        partial = full_surface()
        partial.pop(CM.ALIASES["language_gap_count"][0])
        result = CM.compare_metric_surfaces(
            partial, full_surface(),
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(result["status"], "FAIL")
        self.assertIn(
            "MISSING_CANONICAL_METRIC", {m["reason"] for m in result["mismatches"]})

    def test_the_two_sides_must_agree_on_which_stage_produced_the_artifact(self):
        result = CM.compare_metric_surfaces(
            full_surface(), full_surface(),
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="FULL_SCHEDULE")
        self.assertEqual(result["status"], "FAIL")
        reasons = {m["reason"] for m in result["mismatches"]}
        self.assertIn("STAGE_DISAGREEMENT", reasons)

    def test_a_side_that_declares_no_stage_blocks_and_is_named_apart(self):
        # Blocking either way, but "we disagree" and "you never said" are
        # different defects and must not be reported as the same one.
        missing_validator = CM.compare_metric_surfaces(
            full_surface(), full_surface(), engine_stage="FULL_SCHEDULE")
        self.assertEqual(missing_validator["status"], "FAIL")
        self.assertIn("VALIDATOR_DECLARED_NO_STAGE",
                      {m["reason"] for m in missing_validator["mismatches"]})
        missing_engine = CM.compare_metric_surfaces(
            full_surface(), full_surface(), validator_stage="FULL_SCHEDULE")
        self.assertEqual(missing_engine["status"], "FAIL")
        self.assertIn("ENGINE_DECLARED_NO_STAGE",
                      {m["reason"] for m in missing_engine["mismatches"]})

    def test_a_skeleton_artifact_may_not_claim_a_break_effect_it_cannot_have(self):
        # B-2: "after breaks" on an artifact with no breaks can only mean the
        # same schedule measured again. A different number is fabricated.
        fabricated = full_surface(after_target=93)
        result = CM.compare_metric_surfaces(
            full_surface(), fabricated,
            engine_stage="BEFORE_BREAKS_ONLY", validator_stage="BEFORE_BREAKS_ONLY")
        self.assertEqual(result["status"], "FAIL")
        offending = [m for m in result["mismatches"]
                     if m["reason"] == "AFTER_DIFFERS_FROM_BEFORE_WITHOUT_BREAK_STAGE"]
        self.assertTrue(offending, result["mismatches"])
        self.assertEqual(offending[0]["field"], "after_target")
        self.assertEqual(offending[0]["source"], "independent_validator")

    def test_break_caused_coverage_loss_is_legitimate_at_the_full_stage(self):
        # The same numbers that are a defect above are the normal outcome here.
        lost = full_surface(after_target=93, after_100=93, after_90=95, after_80=97,
                            after_floor=98)
        result = CM.compare_metric_surfaces(
            lost, lost, engine_stage="FULL_SCHEDULE", validator_stage="FULL_SCHEDULE")
        self.assertEqual(result["status"], "PASS", result["mismatches"][:5])
        self.assertFalse(result["before_after_identity_checked"])

    def test_omitting_both_stages_behaves_exactly_as_the_old_gate_did(self):
        self.assertEqual(
            CM.compare_metric_surfaces(full_surface(), full_surface())["status"], "PASS")
        self.assertEqual(
            CM.compare_metric_surfaces({"active_intervals": 1},
                                       {"active_intervals": 2})["status"], "FAIL")


class EngineSurfaceDiscovery(unittest.TestCase):
    def setUp(self):
        self.runner = load("rc924_runner", RUNNER)

    def test_a_skeleton_audit_is_read_as_the_skeleton_stage(self):
        surface = self.runner.engine_metric_surface({
            "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
            "stage_metric_surface": {
                "stage": "BEFORE_BREAKS_ONLY",
                "break_stage_executed": False,
                "after_metrics_basis": "NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE",
                "metrics": {"active_intervals": 112},
            },
        })
        self.assertEqual(surface["stage"], "BEFORE_BREAKS_ONLY")
        self.assertEqual(surface["metrics"], {"active_intervals": 112})
        self.assertIs(surface["break_stage_executed"], False)

    def test_a_full_audit_is_read_as_the_full_stage(self):
        surface = self.runner.engine_metric_surface({
            "artifact_state": "FINAL_VERIFIED",
            "stage_metric_surface": {
                "stage": "FULL_SCHEDULE",
                "break_stage_executed": True,
                "metrics": {"active_intervals": 112},
            },
            "selected_candidate": {"metrics": {"active_intervals": 112}},
        })
        self.assertEqual(surface["stage"], "FULL_SCHEDULE")
        self.assertIs(surface["break_stage_executed"], True)

    def test_an_archived_full_audit_without_the_new_key_is_still_readable(self):
        surface = self.runner.engine_metric_surface({
            "artifact_state": "FINAL_VERIFIED",
            "selected_candidate": {"metrics": {"active_intervals": 7}},
        })
        self.assertEqual(surface["stage"], "FULL_SCHEDULE")
        self.assertEqual(surface["metrics"], {"active_intervals": 7})

    def test_an_archived_skeleton_audit_reports_the_stage_and_no_metrics(self):
        # The pre-fix audits genuinely published nothing. The stage is still
        # recoverable, but the absence of metrics must stay visible.
        surface = self.runner.engine_metric_surface({
            "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
            "selected_before_break_skeleton": {"profile": "x"},
        })
        self.assertEqual(surface["stage"], "BEFORE_BREAKS_ONLY")
        self.assertEqual(surface["metrics"], {})

    def test_engine_selected_metrics_still_answers_for_both_stages(self):
        with tempfile.TemporaryDirectory() as tmp:
            audit = Path(tmp) / "audit.json"
            audit.write_text(json.dumps({
                "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
                "stage_metric_surface": {
                    "stage": "BEFORE_BREAKS_ONLY",
                    "metrics": {"active_intervals": 5},
                },
            }), encoding="utf-8")
            self.assertEqual(
                self.runner.engine_selected_metrics(audit), {"active_intervals": 5})


class ParityGateEndToEnd(unittest.TestCase):
    """Drive apply_metric_parity_gate exactly as a run does."""

    def setUp(self):
        self.runner = load("rc924_runner_gate", RUNNER)

    def run_gate(self, audit_obj, validation_obj):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        root = Path(tmp)
        audit = root / "audit.json"
        vjson = root / "INDEPENDENT_VALIDATION.json"
        vcsv = root / "INDEPENDENT_VALIDATION.csv"
        audit.write_text(json.dumps(audit_obj), encoding="utf-8")
        vjson.write_text(json.dumps(validation_obj), encoding="utf-8")
        rc, validation = self.runner.apply_metric_parity_gate(audit, vjson, vcsv)
        return rc, validation, vcsv

    def skeleton_audit(self, **over):
        metrics = full_surface()
        metrics.update(over)
        return {
            "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
            "status": "SKELETON_ONLY_COMPLETE",
            "stage_metric_surface": {
                "stage": "BEFORE_BREAKS_ONLY",
                "break_stage_executed": False,
                "after_metrics_basis": "NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE",
                "metrics": metrics,
            },
        }

    def skeleton_validation(self, **over):
        metrics = full_surface()
        metrics.update(over)
        return {
            "status": "PASS", "hard_fail_count": 0, "warning_count": 0,
            "artifact_role": "BEST_BEFORE_BREAKS",
            "break_stage_executed": False,
            "metrics": metrics,
        }

    def test_a_skeleton_run_that_agrees_now_returns_zero(self):
        # This is the exact case that used to fail with 41 mismatches.
        rc, validation, _ = self.run_gate(
            self.skeleton_audit(), self.skeleton_validation())
        self.assertEqual(validation["metric_parity"]["mismatch_count"], 0,
                         validation["metric_parity"]["mismatches"][:5])
        self.assertEqual(validation["status"], "PASS")
        self.assertEqual(rc, 0)

    def test_the_gate_records_which_stage_it_compared(self):
        _, validation, vcsv = self.run_gate(
            self.skeleton_audit(), self.skeleton_validation())
        parity = validation["metric_parity"]
        self.assertEqual(parity["engine_stage"], "BEFORE_BREAKS_ONLY")
        self.assertEqual(parity["validator_stage"], "BEFORE_BREAKS_ONLY")
        self.assertEqual(parity["validator_artifact_role"], "BEST_BEFORE_BREAKS")
        self.assertEqual(validation["canonical_metrics"]["stage"], "BEFORE_BREAKS_ONLY")
        self.assertIn("run_stage", vcsv.read_text(encoding="utf-8"))

    def test_an_engine_that_publishes_nothing_still_fails_the_gate(self):
        # The defect was a missing surface. The fix must not turn a missing
        # surface into a pass by comparing nothing against nothing.
        audit = {"artifact_state": "BEST_BEFORE_BREAKS_ONLY"}
        rc, validation, _ = self.run_gate(audit, self.skeleton_validation())
        self.assertEqual(rc, 2)
        self.assertEqual(validation["status"], "FAIL_METRIC_PARITY")
        reasons = {m["reason"] for m in validation["metric_parity"]["mismatches"]}
        self.assertIn("ENGINE_PUBLISHED_NO_METRIC_SURFACE", reasons)

    def test_a_stage_mix_up_between_engine_and_validator_fails_the_gate(self):
        # The runner asked for a skeleton but the validator read a final
        # workbook, or the reverse. Comparing them is meaningless.
        rc, validation, _ = self.run_gate(
            self.skeleton_audit(),
            dict(self.skeleton_validation(), artifact_role="FINAL_AFTER_BREAKS"))
        self.assertEqual(rc, 2)
        reasons = {m["reason"] for m in validation["metric_parity"]["mismatches"]}
        self.assertIn("STAGE_DISAGREEMENT", reasons)

    def test_a_disagreeing_skeleton_run_still_fails_the_gate(self):
        rc, validation, _ = self.run_gate(
            self.skeleton_audit(),
            self.skeleton_validation(language_gap_count=3))
        self.assertEqual(rc, 2)
        self.assertEqual(validation["status"], "FAIL_METRIC_PARITY")

    def test_a_full_run_still_passes_and_still_fails_for_the_right_reasons(self):
        audit = {
            "artifact_state": "FINAL_VERIFIED",
            "stage_metric_surface": {
                "stage": "FULL_SCHEDULE", "break_stage_executed": True,
                "metrics": full_surface(),
            },
            "selected_candidate": {"metrics": full_surface()},
        }
        validation = {
            "status": "PASS", "hard_fail_count": 0, "warning_count": 0,
            "artifact_role": "FINAL_AFTER_BREAKS", "metrics": full_surface(),
        }
        rc, result, _ = self.run_gate(audit, validation)
        self.assertEqual(rc, 0)
        self.assertEqual(result["status"], "PASS")
        rc2, result2, _ = self.run_gate(
            audit, dict(validation, metrics=full_surface(after_floor=3)))
        self.assertEqual(rc2, 2)
        self.assertEqual(result2["status"], "FAIL_METRIC_PARITY")

    def test_the_coverage_gate_cross_check_is_unchanged(self):
        audit = {
            "artifact_state": "FINAL_VERIFIED",
            "production_quality_gate": {"coverage_gate_status": "PASS"},
            "stage_metric_surface": {
                "stage": "FULL_SCHEDULE", "metrics": full_surface()},
        }
        validation = {
            "status": "PASS", "hard_fail_count": 0, "warning_count": 0,
            "artifact_role": "FINAL_AFTER_BREAKS",
            "coverage_quality_gate_status": "FAIL",
            "metrics": full_surface(),
        }
        rc, result, _ = self.run_gate(audit, validation)
        self.assertEqual(rc, 2)
        reasons = {m["reason"] for m in result["metric_parity"]["mismatches"]}
        self.assertIn("QUALITY_GATE_STATUS_MISMATCH", reasons)


class EnginePublishesTheStageSurface(unittest.TestCase):
    """Pin the engine call sites, because a silent removal reopens B-1."""

    def setUp(self):
        self.source = ENGINE.read_text(encoding="utf-8")

    def test_the_skeleton_branch_publishes_a_stage_metric_surface(self):
        branch = self.source[
            self.source.index('audit["selected_before_break_skeleton"] = explanation_rows[0]'):
            self.source.index('audit["status"] = "SKELETON_ONLY_COMPLETE"')]
        self.assertIn('audit["stage_metric_surface"]', branch)
        self.assertIn('"stage": "BEFORE_BREAKS_ONLY"', branch)
        self.assertIn('"break_stage_executed": False', branch)
        self.assertIn('NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE', branch)
        self.assertIn('compact_metric_surface(before_break_metrics)', branch)

    def test_the_full_branch_publishes_a_stage_metric_surface(self):
        self.assertIn('"stage": "FULL_SCHEDULE"', self.source)
        self.assertIn('"break_stage_executed": True', self.source)

    def test_both_stages_use_the_same_evaluator(self):
        # The skeleton surface is calculate_metrics run with no breaks, which
        # is the same function every after-break candidate is scored with.
        # Parity would be circular if the skeleton surface came from a second
        # implementation.
        self.assertIn(
            'skeleton.diagnostics["no_break_metrics"] = metrics', self.source)
        self.assertIn('metrics = calculate_metrics(', self.source)


class ValidatorStageDeclaration(unittest.TestCase):
    """B-2: the validator must say what produced its after_* figures."""

    def setUp(self):
        from openpyxl import Workbook
        self.Workbook = Workbook

    def summary_workbook(self, artifact_type):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        path = Path(tmp) / "book.xlsx"
        wb = self.Workbook()
        ws = wb.active
        ws.title = "Production Summary"
        ws.append(["Release", "RC9.2.2"])
        ws.append(["Artifact Type", artifact_type])
        wb.save(path)
        wb.close()
        return path

    def test_the_production_summary_is_a_second_signal_for_the_stage(self):
        validator = load("rc924_validator", VALIDATOR)
        self.assertEqual(
            validator.declared_artifact_type(
                self.summary_workbook("BEST BEFORE BREAKS SCHEDULE")),
            "BEST_BEFORE_BREAKS")
        self.assertEqual(
            validator.declared_artifact_type(
                self.summary_workbook("BEST FINAL AFTER BREAKS SCHEDULE")),
            "FINAL_AFTER_BREAKS")

    def test_a_workbook_that_declares_nothing_yields_no_second_signal(self):
        validator = load("rc924_validator", VALIDATOR)
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        path = Path(tmp) / "bare.xlsx"
        wb = self.Workbook()
        wb.active.title = "Final Schedule"
        wb.save(path)
        wb.close()
        self.assertIsNone(validator.declared_artifact_type(path))

    def test_a_conflicting_declaration_is_a_hard_failure_not_a_silent_choice(self):
        source = VALIDATOR.read_text(encoding="utf-8")
        self.assertIn("ARTIFACT_STAGE_DECLARATION_CONFLICT", source)
        self.assertIn("failures.append(stage_declaration_conflict)", source)

    def test_the_validator_publishes_the_basis_for_its_after_metrics(self):
        source = VALIDATOR.read_text(encoding="utf-8")
        self.assertIn('"break_stage_executed":artifact_role!="BEST_BEFORE_BREAKS"', source)
        self.assertIn("NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE", source)
        self.assertIn("BREAKS_PLACED_BY_STAGE_2", source)


class RealBeforeBreakArtifact(unittest.TestCase):
    """The same checks against a workbook the engine actually produced.

    A synthesized workbook proves the parser; only a real export proves that
    the engine writes the two signals this stage detection depends on.
    """

    SEARCH = (
        Path("/tmp/claude-0/-home-user-Default/57e8acb4-ab5e-5113-8a50-dec0489e4e6a"
             "/scratchpad/verify_b1"),
        Path("/tmp/claude-0/-home-user-Default/57e8acb4-ab5e-5113-8a50-dec0489e4e6a"
             "/scratchpad/loop"),
    )

    @classmethod
    def locate(cls):
        for root in cls.SEARCH:
            if not root.is_dir():
                continue
            for book in sorted(root.rglob("*_BEST_BEFORE_BREAKS_SCHEDULE.xlsx")):
                snapshot = sorted(book.parent.glob("input_snapshot/*.xlsx"))
                if snapshot:
                    return snapshot[0], book
        return None, None

    def setUp(self):
        self.input_path, self.output_path = self.locate()
        if self.output_path is None:
            self.skipTest("no before-break export present in this environment")
        self.validator = load("rc924_validator_real", VALIDATOR)

    def validate(self, output_path):
        return self.validator.validate(self.input_path, output_path, ENGINE)

    def test_a_real_before_break_export_is_recognized_as_the_skeleton_stage(self):
        result = self.validate(self.output_path)
        self.assertEqual(result["artifact_role"], "BEST_BEFORE_BREAKS")
        self.assertIs(result["break_stage_executed"], False)
        self.assertEqual(
            result["after_metrics_basis"], "NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE")
        self.assertEqual(result["canonical_metrics"]["stage"], "BEFORE_BREAKS_ONLY")

    def test_both_stage_signals_are_written_by_the_engine_and_agree(self):
        result = self.validate(self.output_path)
        signals = result["artifact_role_signals"]
        self.assertEqual(signals["break_sheet"], "BEST_BEFORE_BREAKS")
        self.assertEqual(signals["production_summary"], "BEST_BEFORE_BREAKS")
        self.assertTrue(signals["agree"])

    def test_its_after_metrics_equal_its_before_metrics(self):
        metrics = self.validate(self.output_path)["metrics"]
        for before, after in (("before100", "after100"), ("before90", "after90"),
                              ("before80", "after80"),
                              ("before_target", "after_target"),
                              ("before_floor", "after_floor")):
            self.assertEqual(metrics[before], metrics[after],
                             f"{after} moved without a break stage")

    def test_a_contradicted_stage_declaration_is_a_hard_failure(self):
        from openpyxl import load_workbook
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        tampered = Path(tmp) / "tampered.xlsx"
        shutil.copy(self.output_path, tampered)
        wb = load_workbook(tampered)
        sheet = wb["Production Summary"]
        for row in range(1, min(sheet.max_row, 40) + 1):
            if str(sheet.cell(row, 1).value or "").strip().casefold() == "artifact type":
                sheet.cell(row, 2).value = "BEST FINAL AFTER BREAKS SCHEDULE"
                break
        else:
            self.fail("real export carries no Artifact Type row to contradict")
        wb.save(tampered)
        wb.close()
        result = self.validate(tampered)
        self.assertEqual(result["status"], "FAIL")
        self.assertIn("ARTIFACT_STAGE_DECLARATION_CONFLICT",
                      {f.get("type") for f in result["failures"]})
        # The stricter reading wins, so the tampering cannot promote a skeleton
        # into something whose after_* figures read as a break outcome.
        self.assertEqual(result["artifact_role"], "BEST_BEFORE_BREAKS")
        self.assertIs(result["break_stage_executed"], False)


def skeleton_audit_with(metrics, **over):
    audit = {
        "version": "TEST",
        "status": "SKELETON_ONLY_COMPLETE",
        "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
        "functional_status": "PASS_BEFORE_BREAK_SKELETON_GENERATED",
        "pre_solver_contract_validation": {"status": "PASS", "failures": [], "warnings": []},
        "skeleton_only": {"enabled": True, "selected_profile": "target90_restore_champion"},
        "stage_metric_surface": {
            "stage": "BEFORE_BREAKS_ONLY",
            "break_stage_executed": False,
            "after_metrics_basis": "NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE",
            "metrics": metrics,
        },
    }
    audit.update(over)
    return audit


class PhaseCReadsTheStageThatRan(unittest.TestCase):
    """The quality report published zeros for every skeleton run."""

    def setUp(self):
        self.reporter = load("rc924_reporter", REPORTER)

    def test_a_skeleton_run_reports_the_coverage_it_actually_reached(self):
        report = self.reporter.build_report(
            skeleton_audit_with(full_surface(active_intervals=112, before_target=112,
                                             after_target=112, after_floor=112)))
        self.assertEqual(report["run_stage"], "BEFORE_BREAKS_ONLY")
        self.assertEqual(report["coverage"]["active_intervals"], 112)
        self.assertEqual(report["coverage"]["after_target"], 112)
        self.assertEqual(report["coverage"]["after_metrics_basis"],
                         "NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE")

    def test_a_skeleton_run_is_never_reported_as_releasable(self):
        report = self.reporter.build_report(skeleton_audit_with(full_surface()))
        self.assertIs(report["production_release_eligible"], False)

    def test_a_full_run_that_failed_stage_2_must_not_borrow_skeleton_metrics(self):
        # A FULL_SCHEDULE run that failed break placement still carries a
        # before-break skeleton. Reporting its coverage as this run's result
        # would present before-break numbers as an after-break outcome.
        stage, metrics = self.reporter._stage_and_metrics({
            "status": "FAIL_BREAKS_REQUIRE_EXPLICIT_EXCEPTION_FOR_TESTED_SKELETONS",
            "artifact_state": "FINAL_DIAGNOSTIC_NO_RELEASE_SCHEDULE",
            "selected_before_break_skeleton": {"metrics": full_surface()},
        })
        self.assertEqual(metrics, {})
        self.assertEqual(stage, "FULL_SCHEDULE")

    def test_a_run_that_stopped_at_stage_1_may_use_its_skeleton_metrics(self):
        stage, metrics = self.reporter._stage_and_metrics({
            "status": "SKELETON_ONLY_COMPLETE",
            "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
            "selected_before_break_skeleton": {"metrics": full_surface()},
        })
        self.assertEqual(stage, "BEFORE_BREAKS_ONLY")
        self.assertTrue(metrics)

    def test_a_full_verified_run_is_still_reported_as_releasable(self):
        report = self.reporter.build_report({
            "artifact_state": "FINAL_VERIFIED",
            "pre_solver_contract_validation": {"status": "PASS"},
            "selected_candidate": {"metrics": full_surface()},
            "stage_metric_surface": {"stage": "FULL_SCHEDULE", "metrics": full_surface()},
        })
        self.assertEqual(report["run_stage"], "FULL_SCHEDULE")
        self.assertIs(report["production_release_eligible"], True)
        self.assertEqual(report["coverage"]["after_metrics_basis"],
                         "BREAKS_PLACED_BY_STAGE_2")


class PhaseCExitContract(unittest.TestCase):
    """The release contract and the stage contract are different questions."""

    def run_reporter(self, audit, validation=None):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        root = Path(tmp)
        audit_path = root / "audit.json"
        audit_path.write_text(json.dumps(audit), encoding="utf-8")
        if validation is not None:
            (root / "INDEPENDENT_VALIDATION.json").write_text(
                json.dumps(validation), encoding="utf-8")
        import subprocess
        result = subprocess.run(
            [sys.executable, str(REPORTER), "--audit-json", str(audit_path),
             "--case-root", str(root), "--strict"],
            capture_output=True, text=True)
        report = json.loads((root / "PHASE_C_QUALITY_SUMMARY.json").read_text())
        return result.returncode, report

    def clean_validation(self):
        return {"status": "PASS", "hard_fail_count": 0,
                "metric_parity": {"status": "PASS"},
                "metrics": full_surface()}

    def test_a_validated_skeleton_run_no_longer_returns_a_blocking_code(self):
        # Every BEFORE_BREAKS_ONLY run used to return 2 here, because the
        # release contract asks for FINAL_VERIFIED - which this stage can never
        # be. That is what made the diagnostic stage unusable.
        rc, report = self.run_reporter(
            skeleton_audit_with(full_surface()), self.clean_validation())
        self.assertEqual(rc, 0)
        self.assertEqual(report["stage_gate"]["status"], "PASS")
        self.assertIs(report["production_release_eligible"], False)

    def test_a_skeleton_run_whose_validation_failed_still_blocks(self):
        rc, report = self.run_reporter(
            skeleton_audit_with(full_surface()),
            {"status": "FAIL", "hard_fail_count": 2, "metrics": full_surface()})
        self.assertEqual(rc, 2)
        self.assertIn("independent_validation_not_pass",
                      report["stage_gate"]["blocking_reasons"])

    def test_a_skeleton_run_whose_parity_failed_still_blocks(self):
        rc, report = self.run_reporter(
            skeleton_audit_with(full_surface()),
            dict(self.clean_validation(), metric_parity={"status": "FAIL"}))
        self.assertEqual(rc, 2)
        self.assertIn("metric_parity_not_pass", report["stage_gate"]["blocking_reasons"])

    def test_a_skeleton_run_with_a_failed_input_contract_still_blocks(self):
        audit = skeleton_audit_with(
            full_surface(),
            pre_solver_contract_validation={"status": "FAIL", "failures": [{"x": 1}]})
        rc, report = self.run_reporter(audit, self.clean_validation())
        self.assertEqual(rc, 2)
        self.assertIn("input_contract_not_pass", report["stage_gate"]["blocking_reasons"])

    def test_a_skeleton_run_that_exported_no_artifact_still_blocks(self):
        audit = skeleton_audit_with(full_surface(),
                                    artifact_state="NO_BEFORE_BREAK_SCHEDULE")
        rc, report = self.run_reporter(audit, self.clean_validation())
        self.assertEqual(rc, 2)
        self.assertIn("before_break_artifact_not_exported",
                      report["stage_gate"]["blocking_reasons"])

    def test_the_full_release_contract_is_unchanged(self):
        # A full run that is not FINAL_VERIFIED must still block, exactly as
        # before, and a clean one must still pass.
        blocked_rc, _ = self.run_reporter({
            "artifact_state": "FINAL_VERIFICATION_FAILED",
            "pre_solver_contract_validation": {"status": "PASS"},
            "selected_candidate": {"metrics": full_surface()},
        }, self.clean_validation())
        self.assertEqual(blocked_rc, 2)
        clean_rc, report = self.run_reporter({
            "artifact_state": "FINAL_VERIFIED",
            "pre_solver_contract_validation": {"status": "PASS"},
            "production_quality_gate": {"status": "PASS", "mode": "warn"},
            "selected_candidate": {
                "metrics": full_surface(after_severe_overage_count=0,
                                        after_extreme_overage_count=0,
                                        after_avoidable_overage_fte_sum=0.0),
                "output_validation": {"validation": {"hard_fail_count": 0}},
            },
        }, self.clean_validation())
        self.assertEqual(clean_rc, 0, report.get("phase_c_quality_status"))


class BusinessOutcomeForTheSkeletonStage(unittest.TestCase):
    """B-2: the run must not describe itself as something it is not."""

    def test_the_engine_gives_the_skeleton_stage_its_own_outcome(self):
        import l632_universal_scheduler as E
        outcome = E.build_business_outcome(
            skeleton_audit_with(full_surface(active_intervals=112, before_target=112)), 0)
        self.assertEqual(outcome["outcome_code"], "BEFORE_BREAK_SKELETON_GENERATED")
        self.assertIs(outcome["production_eligible"], False)
        self.assertNotIn("Final schedule generated successfully", outcome["headline"])
        self.assertIn("breaks are not assigned", outcome["headline"].casefold())
        self.assertEqual(outcome["best_proven"]["active_intervals"], 112)

    def test_a_full_run_outcome_is_unchanged(self):
        import l632_universal_scheduler as E
        outcome = E.build_business_outcome({
            "status": "PASS", "artifact_state": "FINAL_VERIFIED",
            "hard_valid_schedule_exists": True,
            "production_quality_gate": {"status": "PASS"},
            "selected_candidate": {"no_break_exception_count": 0},
        }, 0)
        self.assertEqual(outcome["outcome_code"], "FINAL_SCHEDULE_GENERATED")
        self.assertIs(outcome["production_eligible"], True)

    def reconcile(self, audit, validation):
        runner = load("rc924_runner_outcome", RUNNER)
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        root = Path(tmp)
        (root / "x.l6_3_2_3_solver_audit.json").write_text(
            json.dumps(audit), encoding="utf-8")
        (root / "BUSINESS_OUTCOME.json").write_text(json.dumps({
            "outcome_code": "FINAL_SCHEDULE_GENERATED",
            "headline": "Final schedule generated successfully",
            "production_eligible": True,
        }), encoding="utf-8")
        runner.reconcile_business_outcome_after_validation(
            root, validation, 0 if validation.get("status") == "PASS" else 4)
        return json.loads((root / "BUSINESS_OUTCOME.json").read_text())

    def test_a_validated_skeleton_run_is_not_reported_as_a_missing_schedule(self):
        # It used to read: "No final schedule workbook was produced, so
        # independent validation could not run" - on a run whose validation
        # had run and passed.
        outcome = self.reconcile(
            skeleton_audit_with(full_surface()),
            {"status": "PASS", "return_code": 0})
        self.assertEqual(outcome["outcome_code"], "BEFORE_BREAK_SKELETON_GENERATED")
        self.assertIs(outcome["production_eligible"], False)
        self.assertNotIn("validation could not run",
                         outcome["plain_language_summary"])
        self.assertIn("must not be used operationally",
                      outcome["plain_language_summary"])

    def test_a_skeleton_run_that_failed_validation_says_so(self):
        outcome = self.reconcile(
            skeleton_audit_with(full_surface()),
            {"status": "FAIL", "return_code": 2})
        self.assertEqual(outcome["outcome_code"], "BEFORE_BREAK_SKELETON_BLOCKED")
        self.assertIs(outcome["production_eligible"], False)

    def test_a_blocked_full_run_still_gets_its_existing_diagnosis(self):
        outcome = self.reconcile(
            {"status": "FAIL_PRE_SOLVER_CONTRACT",
             "artifact_state": "NO_SCHEDULE"},
            {"status": "FAIL", "return_code": 2})
        self.assertEqual(outcome["outcome_code"], "INPUT_OR_RESOURCE_CONTRACT_GAP")


class PreliminaryQualityGateWaitsForValidation(unittest.TestCase):
    """The first Phase C call runs before validation and must not outlive it.

    On a skeleton run the preliminary report is blocked (validation has not
    run yet), the runner adopted that blocking code, and the later clean
    report could never clear it - so a fully passing skeleton run still ended
    at rc=2.
    """

    def setUp(self):
        self.runner = load("rc924_runner_pending", RUNNER)

    def check(self, report, validation_exists=False):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        root = Path(tmp)
        (root / "PHASE_C_QUALITY_SUMMARY.json").write_text(
            json.dumps(report), encoding="utf-8")
        if validation_exists:
            (root / "INDEPENDENT_VALIDATION.json").write_text("{}", encoding="utf-8")
        return self.runner.quality_allows_validation(root, 2)

    def skeleton_report(self, **over):
        report = {
            "run_stage": "BEFORE_BREAKS_ONLY",
            "artifact_state": "BEST_BEFORE_BREAKS_ONLY",
            "contract": {"status": "PASS", "failure_count": 0},
            "safety": {"status": "NOT_VALIDATED", "hard_fail_count": 0},
            "production_quality_gate": {"status": "NOT_EVALUATED"},
        }
        report.update(over)
        return report

    def full_report(self, **over):
        report = {
            "run_stage": "FULL_SCHEDULE",
            "artifact_state": "FINAL_VERIFIED",
            "contract": {"status": "PASS", "failure_count": 0},
            "safety": {"status": "NOT_VALIDATED", "hard_fail_count": 0},
            "production_quality_gate": {"status": "WARN"},
        }
        report.update(over)
        return report

    def test_a_skeleton_run_may_proceed_to_its_validation(self):
        self.assertTrue(self.check(self.skeleton_report()))

    def test_a_skeleton_run_with_no_exported_artifact_may_not(self):
        self.assertFalse(self.check(
            self.skeleton_report(artifact_state="NO_BEFORE_BREAK_SCHEDULE")))

    def test_a_skeleton_run_with_a_failed_contract_may_not(self):
        self.assertFalse(self.check(self.skeleton_report(
            contract={"status": "FAIL", "failure_count": 3})))

    def test_a_skeleton_run_with_hard_failures_may_not(self):
        self.assertFalse(self.check(self.skeleton_report(
            safety={"status": "FAIL", "hard_fail_count": 1})))

    def test_a_full_run_is_unaffected(self):
        self.assertTrue(self.check(self.full_report()))
        self.assertFalse(self.check(
            self.full_report(artifact_state="FINAL_VERIFICATION_FAILED")))

    def test_the_skeleton_allowance_does_not_leak_into_the_full_path(self):
        # NOT_EVALUATED is honest for a stage that never runs the production
        # quality gate. A full run that reached Stage 2 without evaluating it
        # must still be treated as it was before.
        self.assertFalse(self.check(
            self.full_report(production_quality_gate={"status": "NOT_EVALUATED"})))

    def test_nothing_proceeds_once_validation_has_already_run(self):
        self.assertFalse(self.check(self.skeleton_report(), validation_exists=True))
        self.assertFalse(self.check(self.full_report(), validation_exists=True))


if __name__ == "__main__":
    unittest.main(verbosity=1)
