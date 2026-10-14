# © 2026 Omar Mokhtar. All rights reserved.
"""The channel plan as an Excel file (Phase AF, C4), in the two views the owner asked for: per working day a table
of stretches (as Plan channels lists them) and a timeline (a row per person, a column per quarter hour, each cell in
its channel's colour with its letter), then a Week tab of hours per person and channel."""
from __future__ import annotations

import io
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .channels import hm

NAMES = {"P": "Phone", "C": "Chat", "E": "Email", "A": "All channels", "B": "Break", "L": "Lunch"}
COLOURS = {"P": "3E6FD8", "C": "B4479A", "E": "4D7D8D", "A": "6B7A90", "B": "D0D5DD", "L": "D0D5DD"}
BOLD = Font(bold=True)


def channel_plan_workbook(label: str, days: List[Tuple[str, Dict[str, Any]]]) -> bytes:
    """``days``: (day name, Plan channels grid) for each working day, in week order."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in (f"Channel plan: {label}", "",
                 "Each day has two tabs: the day's stretches as a table (one row per stretch), and the day as a "
                 "timeline (one row per person, one column per quarter hour, in the channel's colour).",
                 "Week sums the hours per person and channel.", "",
                 "Phone, Chat, Email: the channel. All channels: times when the people on the floor cover every "
                 "channel they work. Break and Lunch: off the floor."):
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120
    totals: Dict[str, Dict[str, int]] = {}
    for day, grid in days:
        table = wb.create_sheet(day)
        table.append(["Associate", "Language", "From", "To", "Channel", "Minutes"])
        for cell in table[1]:
            cell.font = BOLD
        for row in grid["rows"]:
            for start, end, letter in row["runs"]:
                if not letter:
                    continue
                table.append([row["name"], row["language"], hm(start), hm(end), NAMES[letter], end - start])
                if letter in "PCEA":
                    mine = totals.setdefault(row["name"], {})
                    mine[letter] = mine.get(letter, 0) + end - start
        for col, width in zip("ABCDEF", (22, 14, 8, 8, 14, 9)):
            table.column_dimensions[col].width = width
        line = wb.create_sheet(f"{day} timeline")
        line.append(["Associate"] + [hm(t) if t % 60 == 0 else "" for t in grid["cols"]])
        line["A1"].font = BOLD
        for row in grid["rows"]:
            line.append([row["name"]] + [c["letter"] for c in row["cells"]])
            for c, cell in zip(row["cells"], line[line.max_row][1:]):
                if c["letter"]:
                    cell.fill = PatternFill("solid", fgColor=COLOURS[c["letter"]])
                    cell.font = Font(bold=True, color="FFFFFF" if c["letter"] in "PCEA" else "3A4250")
                    cell.alignment = Alignment(horizontal="center")
        line.column_dimensions["A"].width = 22
        for i in range(2, len(grid["cols"]) + 2):
            line.column_dimensions[get_column_letter(i)].width = 3.2
        line.freeze_panes = "B2"
    week = wb.create_sheet("Week")
    week.append(["Associate", "Phone h", "Chat h", "Email h", "All channels h"])
    for cell in week[1]:
        cell.font = BOLD
    for name in sorted(totals):
        week.append([name] + [round(totals[name].get(c, 0) / 60, 2) for c in "PCEA"])
    week.column_dimensions["A"].width = 22
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
