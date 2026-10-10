# © 2026 Omar Mokhtar. All rights reserved.
"""Phase V: the channel tabs of an input workbook (owner, 2026-10-09: "some programs handles different channels chat,
phone, email same associates should rotate on those 3 channels ... create 3 sheets in the input sheet we will assign
the need per interval for each channel ... this can be done on the website only no need to add it in the engine").

The engine never reads these tabs. The website does:
  Chat {n} Min, Phone {n} Min, Email {n} Min   people needed per interval, laid out like FT Wise {n} Min
  Email Hours                                   hours of email per day (used when Email {n} Min is empty or missing)
  Channel Setup                                 the rotation rules, the language minimums per channel, and the
                                                all-channels times (when 1 or 2 people cover every channel at once)

Everything is read strictly: a cell the website cannot use is refused with its tab and row, never read as 0.
``check_lines`` is what the upload says: where the channel tabs and FT Wise disagree (owner: "3 channels should be
covered by the requirements"), leaving out the all-channels times (owner: "to avoid invalid warning")."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .versions import DAYS

CHANNELS = {"P": "Phone", "C": "Chat", "E": "Email"}
LETTER = {name.casefold(): letter for letter, name in CHANNELS.items()}
QUARTER = 15  # minutes: the planning step (breaks start on quarter hours, as the validator checks)
GRID_TAB = re.compile(r"^(chat|phone|email) (\d+) min$")
# The engine's _discover_requirement_sheet: the instruction first, then these names, then the first Sun-Sat grid.
REQUIREMENT_INSTRUCTIONS = ("requirements source", "requirement source", "active requirements used")
REQUIREMENT_ALIASES = ("FT Wise 30 Min", "FT Wise 60 Min", "FT Wise", "Required HC", "Requirements", "Forecast")
SETTINGS = {"minimum block minutes": "min_block", "phone maximum continuous minutes": "max_P",
            "chat maximum continuous minutes": "max_C", "email maximum continuous minutes": "max_E",
            "priority when short": "order", "how strict is the order": "strict",
            "spread channels fairly over the week": "fair"}
YES, NO = ("yes", "y", "true"), ("no", "n", "false")
EPS = 1e-6
KEEP = 8
_CACHE: Dict[Tuple[Any, ...], Any] = {}  # read_channels by (path, mtime, step); has_channel_needs by ("needs", ...)


def _norm(value: Any) -> str:
    return " ".join(str(value if value is not None else "").split()).casefold()


def hm(minute: int) -> str:
    return f"{(minute // 60) % 24:02d}:{minute % 60:02d}"


def _n(value: float) -> str:
    return f"{value:g}"


def _time(value: Any, where: str) -> Optional[int]:
    """Minutes from midnight; None when blank; refused when it is not a time."""
    if value is None or str(value).strip() == "":
        return None
    if hasattr(value, "hour"):
        return value.hour * 60 + value.minute
    if isinstance(value, float) and 0 <= value < 1:  # an Excel time serial
        return int(round(value * 1440))
    found = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*", str(value))
    if found:
        h, m = int(found.group(1)), int(found.group(2))
        if m < 60 and (h < 24 or (h == 24 and m == 0)):
            return h * 60 + m
    raise ValueError(f"{where}: {value} is not a time (use HH:MM, for example 08:30).")


def _day(token: str) -> Optional[int]:
    token = token.strip().casefold()
    full = ("sunday", "monday", "tuesday", "wednesday", "thursday", "friday", "saturday")
    return next((i for i, name in enumerate(full) if len(token) >= 3 and name.startswith(token)), None)


def _days(value: Any, where: str) -> Set[int]:
    """"All" (or blank) for every day, or days like "Sun, Mon" or "Sun-Thu"; refused when a day is not one."""
    text = " ".join(str(value if value is not None else "").split())
    if text.casefold() in ("", "all", "every day", "daily"):
        return set(range(7))
    found: Set[int] = set()
    for part in re.split(r"[,;/]", text):
        if not part.strip():
            continue
        ends = [_day(x) for x in part.split("-")] if "-" in part else [_day(part)]
        if None in ends or len(ends) > 2:
            raise ValueError(f"{where}: {part.strip()} is not a day (use Sun to Sat, a range like Sun-Thu, or All).")
        if len(ends) == 2:
            d = ends[0]
            while True:
                found.add(d)
                if d == ends[1]:
                    break
                d = (d + 1) % 7
        else:
            found.add(ends[0])
    return found


def _days_text(days: Set[int]) -> str:
    return "All" if len(days) == 7 else ", ".join(DAYS[d] for d in sorted(days))


def _number(value: Any, where: str, whole: bool = False) -> float:
    if value is None or str(value).strip() == "":
        return 0.0
    if isinstance(value, bool):
        raise ValueError(f"{where}: {value} is not a number.")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{where}: {value} is not a number.") from None
    if number < 0:
        raise ValueError(f"{where}: {value} is below 0.")
    if whole and number != int(number):
        raise ValueError(f"{where}: {value} is not a whole number.")
    return number


def inside(days: Set[int], start: int, end: int, day: int, t: int) -> bool:
    """Whether minute ``t`` (0 to 1439) of ``day`` falls in a window on ``days``: start == end is the whole day,
    start > end runs past midnight into the next day."""
    if start == end:
        return day in days
    if start < end:
        return day in days and start <= t < end
    return (day in days and t >= start) or ((day - 1) % 7 in days and t < end)


def blended_at(setup: Optional[Dict[str, Any]], day: int, t: int) -> bool:
    """Whether ``t`` on ``day`` is an all-channels time (``t`` may run past 1440 for an overnight shift)."""
    if not setup:
        return False
    day, t = (day + t // 1440) % 7, t % 1440
    return any(inside(b["days"], b["start"], b["end"], day, t) for b in setup["blended"])


# ----------------------------------------------------------------------------------------------- reading the tabs
def _grid(ws, step: int, problems: List[str]) -> Dict[int, Dict[int, float]]:
    label = ws.title
    header = next((r for r in range(1, 11) if _norm(ws.cell(r, 1).value) == "interval"), None)
    cols = {}
    if header is not None:
        for c in range(2, ws.max_column + 1):
            d = _day(str(ws.cell(header, c).value or ""))
            if d is not None and d not in cols:
                cols[d] = c
    if header is None or len(cols) < 7:
        problems.append(f"{label}: needs a header row with Interval and Sun to Sat, as on FT Wise {step} Min.")
        return {}
    out: Dict[int, Dict[int, float]] = {d: {} for d in range(7)}
    for r in range(header + 1, ws.max_row + 1):
        where = f"{label}, row {r}"
        raw = ws.cell(r, 1).value
        values = {d: ws.cell(r, c).value for d, c in cols.items()}
        try:
            t = _time(raw, where)
        except ValueError as exc:
            problems.append(str(exc))
            continue
        if t is None:
            if any(v not in (None, "") for v in values.values()):
                problems.append(f"{where}: numbers without an interval time.")
            continue
        if t % step or t >= 1440:
            problems.append(f"{where}: {raw} is not on the {step}-minute grid.")
            continue
        for d, v in values.items():
            try:
                out[d][t] = _number(v, where)
            except ValueError:
                problems.append(f"{where}: {DAYS[d]} {hm(t)} reads {v!r}; a need is a number, 0 or more.")
    return out


def _email_hours(ws, problems: List[str]) -> Dict[int, Dict[str, Any]]:
    out = {d: {"hours": 0.0, "languages": {}, "start": 0, "end": 1440} for d in range(7)}
    header = next((r for r in range(1, 11) if _norm(ws.cell(r, 1).value) == "day"), None)
    cols = {_norm(ws.cell(header, c).value): c for c in range(1, ws.max_column + 1)
            if header and ws.cell(header, c).value not in (None, "")}
    hours_col = cols.get("hours needed", cols.get("hours"))
    if header is None or hours_col is None:
        problems.append("Email Hours: needs a header row with Day and Hours needed.")
        return out
    langs = {key[len("of which "):]: c for key, c in cols.items() if key.startswith("of which ")}
    names = {_norm(ws.cell(header, c).value)[len("of which "):]: str(ws.cell(header, c).value).strip()[9:].strip()
             for c in range(1, ws.max_column + 1) if _norm(ws.cell(header, c).value).startswith("of which ")}
    seen: Set[int] = set()
    for r in range(header + 1, ws.max_row + 1):
        where = f"Email Hours, row {r}"
        cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v in (None, "") for v in cells):
            continue
        day_text = str(ws.cell(r, cols["day"]).value or "").strip()
        d = _day(day_text) if day_text and "," not in day_text else None
        if d is None:
            problems.append(f"{where}: {day_text or 'the day'} is not a day (use one of Sun to Sat).")
            continue
        if d in seen:
            problems.append(f"{where}: {DAYS[d]} is given twice.")
            continue
        seen.add(d)
        try:
            hours = _number(ws.cell(r, hours_col).value, where)
            by = {}
            for key, c in langs.items():
                value = ws.cell(r, c).value
                if value in (None, ""):
                    continue
                part = _number(value, where)
                if part > hours + EPS:
                    raise ValueError(f"{where}: Of which {names[key]} is {_n(part)} h, more than the {_n(hours)} h "
                                     "needed.")
                if part:
                    by[names[key]] = part
            start = _time(ws.cell(r, cols["window start"]).value, where) if "window start" in cols else None
            end = _time(ws.cell(r, cols["window end"]).value, where) if "window end" in cols else None
            start, end = (0 if start is None else start), (1440 if end is None else end)
            if start >= end:
                raise ValueError(f"{where}: the window must start before it ends.")
        except ValueError as exc:
            problems.append(str(exc))
            continue
        out[d] = {"hours": hours, "languages": by, "start": start, "end": end}
    return out


def _setup(ws, step: int, problems: List[str], notes: List[str]) -> Tuple[Dict[str, Any], List, List]:
    found: Dict[str, Any] = {}
    languages: List[Dict[str, Any]] = []
    blended: List[Dict[str, Any]] = []
    section, cols = None, {}
    for r in range(1, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, max(ws.max_column, 7) + 1)]
        first = _norm(cells[0])
        filled = [i for i, v in enumerate(cells) if v not in (None, "")]
        if not filled:
            continue
        second = _norm(cells[1])
        if first == "setting" and second == "value":
            section = "settings"
            continue
        if first == "channel" and second == "language":
            section, cols = "languages", {_norm(v): i for i, v in enumerate(cells) if v not in (None, "")}
            continue
        if first == "days" and second == "start":
            section, cols = "blended", {_norm(v): i for i, v in enumerate(cells) if v not in (None, "")}
            continue
        where = f"Channel Setup, row {r}"
        try:
            if section == "settings":
                if first not in SETTINGS:
                    if filled == [0] and first not in LETTER:
                        section = None  # a title or a note
                        continue
                    raise ValueError(f"{where}: {cells[0]} is not a setting this tab knows ("
                                     + ", ".join(k[0].upper() + k[1:] for k in SETTINGS) + ").")
                key = SETTINGS[first]
                if key in found:
                    raise ValueError(f"{where}: {cells[0]} is set twice.")
                found[key] = (cells[1], where, str(cells[0]).strip())
            elif section == "languages":
                if filled == [0] and first not in LETTER:
                    section = None
                    continue
                get = lambda k: cells[cols[k]] if k in cols else None  # noqa: E731
                channel = LETTER.get(_norm(get("channel")))
                if channel is None:
                    raise ValueError(f"{where}: {get('channel')} is not a channel (Chat, Phone or Email).")
                language = " ".join(str(get("language") or "").split())
                if not language:
                    raise ValueError(f"{where}: give the language.")
                minimum = _number(get("minimum per interval"), where, whole=True)
                start = _time(get("coverage start"), where) or 0
                end = _time(get("coverage end"), where) or 0
                days = _days(get("coverage days"), where)
                languages.append({"channel": channel, "language": language, "minimum": int(minimum),
                                  "start": start % 1440, "end": end % 1440, "days": days,
                                  "active": _norm(get("active?")) not in NO})
            elif section == "blended":
                get = lambda k: cells[cols[k]] if k in cols else None  # noqa: E731
                if filled == [0] and _days_or_none(cells[0]) is None:
                    section = None
                    continue
                days = _days(get("days"), where)
                start, end = _time(get("start"), where), _time(get("end"), where)
                if start is None or end is None:
                    raise ValueError(f"{where}: give the start and the end of the all-channels time.")
                if _norm(get("active?")) not in NO:
                    blended.append({"days": days, "start": start % 1440, "end": end % 1440})
        except ValueError as exc:
            problems.append(str(exc))
    rules = {"min_block": max(QUARTER, step - step % QUARTER), "max_run": {"P": 0, "C": 0, "E": 0},
             "order": ["P", "C", "E"], "strict": True, "fair": True}
    said = {"min_block": f"a minimum block of {rules['min_block']} minutes", "max_P": "no Phone maximum",
            "max_C": "no Chat maximum", "max_E": "no Email maximum", "order": "Phone > Chat > Email",
            "strict": "Strict", "fair": "channels spread fairly"}
    missing = [said[k] for k in said if k not in found]
    if missing and found:
        notes.append("Channel Setup does not set everything, so these defaults are used: " + ", ".join(missing) + ".")
    for key, (value, where, label) in found.items():
        try:
            if key == "min_block" or key.startswith("max_"):
                minutes = _number(value, where, whole=True) if str(value if value is not None else "").strip() else -1
                if minutes < 0 or minutes % QUARTER or (key == "min_block" and minutes < QUARTER):
                    raise ValueError(f"{where}: {label} must be whole minutes in 15-minute steps"
                                     f"{' (15, 30, 45, 60 ...)' if key == 'min_block' else ' (0 means no limit)'};"
                                     f" it reads {value}.")
                if key == "min_block":
                    rules["min_block"] = int(minutes)
                else:
                    rules["max_run"][key[-1]] = int(minutes)
            elif key == "order":
                letters = [LETTER.get(_norm(x)) for x in re.split(r"[>,]", str(value or "")) if x.strip()]
                if sorted(x or "" for x in letters) != ["C", "E", "P"]:
                    raise ValueError(f"{where}: {label} must name each channel once, like Phone > Chat > Email; it "
                                     f"reads {value}.")
                rules["order"] = letters
            elif key == "strict":
                if _norm(value) not in ("strict", "balanced"):
                    raise ValueError(f"{where}: {label} must be Strict or Balanced; it reads {value}.")
                rules["strict"] = _norm(value) == "strict"
            elif key == "fair":
                if _norm(value) not in YES + NO:
                    raise ValueError(f"{where}: {label} must be Yes or No; it reads {value}.")
                rules["fair"] = _norm(value) in YES
        except ValueError as exc:
            problems.append(str(exc))
    return rules, languages, blended


def _days_or_none(value: Any) -> Optional[Set[int]]:
    try:
        return _days(value, "")
    except ValueError:
        return None


def has_channel_tabs(path: Path) -> bool:
    """Whether a workbook has any channel tab (a grid for any interval, Email Hours or Channel Setup). A file that
    cannot be opened as a workbook has none here: the run's own check reports it, as before Phase V."""
    try:
        wb = load_workbook(path, read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError):
        return False
    try:
        return any(GRID_TAB.match(_norm(n)) or _norm(n) in ("email hours", "channel setup") for n in wb.sheetnames)
    finally:
        wb.close()


def has_channel_needs(path: Path) -> bool:
    """Whether a workbook's channel tabs ask for anyone (Phase Z): a cell other than empty or 0 under the header of a
    Chat, Phone or Email tab (any interval) or of Email Hours, or a filled language or all-channels row on Channel
    Setup. Anything that is not a number counts, so it is read and refused by tab and row. The tabs as the Blank
    input workbook and the channel needs download hand them out ask for nobody: such a workbook is treated as one
    without channel tabs, as before Phase V. A file that cannot be opened as a workbook has none (see above). Read
    once per file version (the day page asks for each candidate break time and each exported day)."""
    try:
        stat = Path(path).stat()
    except OSError:
        return False
    key = ("needs", str(path), stat.st_mtime_ns, stat.st_size)
    if key not in _CACHE:
        _remember(key, _asks_for_someone(path))
    return bool(_CACHE[key])


def _asks_for_someone(path: Path) -> bool:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError):
        return False

    def asks(value: Any) -> bool:
        return value not in (None, "") and not (isinstance(value, (int, float)) and not isinstance(value, bool)
                                                and value == 0)
    try:
        for name in wb.sheetnames:
            key = _norm(name)
            rows = list(wb[name].iter_rows(values_only=True)) if (GRID_TAB.match(key) or key in (
                "email hours", "channel setup")) else []
            if GRID_TAB.match(key) or key == "email hours":
                head = next((i for i, r in enumerate(rows[:10]) if r and _norm(r[0]) in ("interval", "day")), None)
                if head is not None and any(asks(v) for r in rows[head + 1:] for v in r[1:]):
                    return True
            elif key == "channel setup":
                section = None
                for r in rows:
                    first, second = _norm(r[0] if r else None), _norm(r[1] if r and len(r) > 1 else None)
                    if (first, second) in (("channel", "language"), ("days", "start")):
                        section = first
                    elif section == "channel" and first in LETTER:
                        return True
                    elif section == "days" and any(asks(v) for v in r):
                        return True
        return False
    finally:
        wb.close()


def read_channels(path: Path, step: int) -> Optional[Dict[str, Any]]:
    """The channel tabs of an input workbook, or None when it has none. Refused (ValueError naming each tab and row,
    all of them at once) when a cell cannot be used. Read once per file version."""
    key = (str(path), Path(path).stat().st_mtime, step)
    if key in _CACHE:
        return _CACHE[key]
    wb = load_workbook(path, data_only=True)
    by_name = {_norm(n): n for n in wb.sheetnames}
    problems: List[str] = []
    notes: List[str] = []
    for norm_name, name in by_name.items():
        found = GRID_TAB.match(norm_name)
        if found and int(found.group(2)) != step:
            channel = found.group(1).capitalize()
            problems.append(f"{name} does not match the program's {step}-minute interval: name it {channel} {step} "
                            f"Min, with {step}-minute rows.")
    grids = {letter: wb[by_name[f"{CHANNELS[letter].casefold()} {step} min"]]
             for letter in "PCE" if f"{CHANNELS[letter].casefold()} {step} min" in by_name}
    hours_ws = wb[by_name["email hours"]] if "email hours" in by_name else None
    setup_ws = wb[by_name["channel setup"]] if "channel setup" in by_name else None
    if not grids and hours_ws is None and setup_ws is None and not problems:
        found_none = None
        _remember(key, found_none)
        return found_none
    need = {letter: _grid(ws, step, problems) for letter, ws in grids.items()}
    hours = _email_hours(hours_ws, problems) if hours_ws is not None else None
    if setup_ws is not None:
        rules, languages, blended = _setup(setup_ws, step, problems, notes)
    else:
        rules = {"min_block": max(QUARTER, step - step % QUARTER), "max_run": {"P": 0, "C": 0, "E": 0},
                 "order": ["P", "C", "E"], "strict": True, "fair": True}
        languages, blended = [], []
        notes.append(f"No Channel Setup tab: defaults used (a minimum block of {rules['min_block']} minutes, no "
                     "maximums, Phone > Chat > Email, Strict, no language minimums per channel, no all-channels "
                     "times).")
    if not grids and hours_ws is None and not problems:
        problems.append(f"Channel Setup: there is no channel tab to set up (Chat {step} Min, Phone {step} Min, "
                        f"Email {step} Min or Email Hours).")
    if problems:
        more = f" And {len(problems) - 12} more." if len(problems) > 12 else ""
        raise ValueError(" ".join(problems[:12]) + more)
    email_any = "E" in need and any(v for col in need["E"].values() for v in col.values())
    hours_any = hours is not None and any(h["hours"] for h in hours.values())
    if email_any:
        mode = "interval"
        if hours_ws is not None:
            notes.append(f"Email {step} Min has numbers, so Email Hours is not used.")
    else:
        need.pop("E", None)
        mode = "hours" if hours_any else "none"
    found = {"step": step, "need": need, "tabs": [ws.title for ws in grids.values()],
             "hours_tab": hours_ws.title if hours_ws is not None else "", "email_mode": mode,
             "email_hours": hours if mode == "hours" else {d: {"hours": 0.0, "languages": {}, "start": 0, "end": 1440}
                                                          for d in range(7)},
             "rules": rules, "languages": languages, "blended": blended, "notes": notes}
    _remember(key, found)
    return found


def _remember(key, value) -> None:
    while len(_CACHE) >= KEEP:
        _CACHE.pop(next(iter(_CACHE)))
    _CACHE[key] = value


# ----------------------------------------------------------------------------------------------- the requirement tab
def _requirement(path: Path) -> Tuple[Optional[str], str]:
    """(the tab the engine reads as its requirement, "instruction" or "alias"), or (None, "") when it would fall
    back to the first Sun-Sat grid, which with channel tabs could be a channel tab."""
    wb = load_workbook(path, read_only=True)
    try:
        by_name = {_norm(n): n for n in wb.sheetnames}
        ws = wb[by_name["instructions"]] if "instructions" in by_name else None
        for row in (ws.iter_rows(values_only=True) if ws is not None else []):
            cells = list(row)
            for i, v in enumerate(cells):
                if _norm(v) in REQUIREMENT_INSTRUCTIONS:
                    value = next((x for x in cells[i + 1:] if x not in (None, "")), None)
                    if value is not None and _norm(value) in by_name:
                        return by_name[_norm(value)], "instruction"
        for alias in REQUIREMENT_ALIASES:
            if _norm(alias) in by_name:
                return by_name[_norm(alias)], "alias"
        return None, ""
    finally:
        wb.close()


def requirement_tab(path: Path) -> Optional[str]:
    return _requirement(path)[0]


# ----------------------------------------------------------------------------------------------- the upload check
def _ranges(items: List[Tuple[int, int, tuple]], step: int) -> List[Tuple[int, int, int, tuple]]:
    """Consecutive intervals of one day with the same numbers, joined: (day, start, end, numbers)."""
    out: List[Tuple[int, int, int, tuple]] = []
    for d, t, numbers in items:
        if out and out[-1][0] == d and out[-1][2] == t and out[-1][3] == numbers:
            out[-1] = (d, out[-1][1], t + step, numbers)
        else:
            out.append((d, t, t + step, numbers))
    return out


def _when(d: int, lo: int, hi: int, step: int) -> str:
    return f"{DAYS[d]} {hm(lo)}" + (f" to {hm(hi)}" if hi - lo > step else "")


def check_lines(setup: Dict[str, Any], inputs: Dict[str, Any], path: Path) -> List[Tuple[str, str]]:
    """What the upload says about the channel tabs: (level, text), level "ok", "info" or "warn". Warnings never
    block anything and never change the schedule."""
    step = inputs["interval"]
    name, how = _requirement(path)
    ft = name or f"FT Wise {step} Min"
    lines: List[Tuple[str, str]] = []
    if name:
        why = "named in Requirements Source" if how == "instruction" else "an engine name for it"
        lines.append(("ok", f"{name} is the schedule's requirement ({why}), so a channel tab is never read as the "
                            "requirement."))
    else:
        lines.append(("warn", f"No requirement tab is named: add Requirements Source on the Instructions tab, or name "
                              f"the tab FT Wise {step} Min. Until then the engine may read a channel tab as the "
                              "requirement."))
    mode, need = setup["email_mode"], setup["need"]
    tabs = [t for t in setup["tabs"] if mode == "interval" or not t.casefold().startswith("email")]
    if mode == "hours":
        tabs.append(setup["hours_tab"])
    read = "Read " + (", ".join(tabs[:-1]) + " and " + tabs[-1] if len(tabs) > 1 else tabs[0]) + "."
    email_tab = next((t for t in setup["tabs"] if t.casefold().startswith("email")), "")
    if mode == "hours" and email_tab:
        read += f" {email_tab} is empty, so email follows Email Hours."
    elif mode == "none":
        read += " There is no email to plan."
    lines.append(("ok", read))
    lines += [("info", note) for note in setup["notes"]]
    letters = [c for c in ("C", "P", "E") if c in need]
    said = {"C": "Chat", "P": "Phone", "E": "Email"}
    req = inputs.get("required", {})
    counted = ok = in_blend = avoided = 0
    wrong: List[Tuple[int, int, tuple]] = []
    blend_wrong: List[Tuple[int, int, tuple]] = []
    left = {d: 0.0 for d in range(7)}
    for d in range(7):
        window = setup["email_hours"][d]
        for t in range(0, 1440, step):
            values = {c: need[c].get(d, {}).get(t, 0.0) for c in letters}
            f = float(req.get(d, {}).get(t, 0.0) or 0.0)
            total = sum(values.values())
            if not total and not f:
                continue
            off = total - f
            bad = abs(off) > EPS if mode == "interval" else off > EPS
            if blended_at(setup, d, t):
                in_blend += 1
                avoided += int(bad)
                top = max(letters, key=lambda c: values[c]) if letters else None
                if top and values[top] > f + EPS:
                    blend_wrong.append((d, t, (top, values[top], f)))
                continue
            counted += 1
            if bad:
                wrong.append((d, t, tuple(values[c] for c in letters) + (f,)))
            else:
                ok += 1
            if mode == "hours" and window["start"] <= t < window["end"]:
                left[d] += max(0.0, f - total) * step / 60
    joined = " + ".join(said[c] for c in letters)
    verb = "match" if mode == "interval" else "fit inside"
    if counted:
        lines.append(("ok", f"{joined} {verb} {ft} in {ok} of {counted} intervals."))
    for d, lo, hi, numbers in _ranges(wrong, step):
        parts = " + ".join(f"{said[c]} {_n(v)}" for c, v in zip(letters, numbers))
        lines.append(("warn", f"{_when(d, lo, hi, step)}: {parts} = {_n(sum(numbers[:-1]))}, but {ft} needs "
                              f"{_n(numbers[-1])}. The channel plan follows the channel tabs; the schedule keeps {ft}."))
    if setup["blended"]:
        desc = "; ".join(f"{_days_text(b['days'])} {hm(b['start'])} to {hm(b['end'])}" for b in setup["blended"])
        lines.append(("info", f"All-channels times (Channel Setup): {desc}. {in_blend} intervals there are not added "
                               f"up, so {avoided} warnings that would have been false are not raised. Checked "
                               f"instead: no single channel needs more people than {ft}."))
    for d, lo, hi, (c, v, f) in _ranges(blend_wrong, step):
        lines.append(("warn", f"{_when(d, lo, hi, step)} (all channels): {said[c]} alone needs {_n(v)}, but {ft} "
                              f"needs {_n(f)}."))
    if mode == "hours":
        days = [d for d in range(7) if setup["email_hours"][d]["hours"] > 0]
        fits = 0
        for d in days:
            hours = setup["email_hours"][d]["hours"]
            if left[d] + EPS < hours:
                lines.append(("warn", f"{DAYS[d]}: {ft} leaves {_n(round(left[d], 2))} h after Chat and Phone; Email "
                                      f"Hours asks for {_n(hours)} h. The other {_n(round(hours - left[d], 2))} h "
                                      "fits only where more people turn up than needed."))
            else:
                fits += 1
        lines.append(("ok", f"Email Hours fit inside {ft} on {fits} of {len(days)} days."))
    return lines
