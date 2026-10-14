# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AF, A3: Channel Setup is specific, roomy and made of drop-downs, and its language windows add up.

- 20 rows for language minimums (the blank had 6), each column a drop-down: channel, language (from Language Setup,
  the list the engine reads), minimum, times, days and active.
- Days read as in Language Setup: Weekdays, Weekends, ranges and single days. The sheet says Days is the day a window
  starts (Mon - Fri 16:00 to 07:00 does not cover Sunday night).
- A channel's language window sits inside that language's coverage in Language Setup; anything else is refused
  (owner's example: English covered 16:00 to 07:00 rules out a channel row of 15:00 to 08:00). A language with no
  Language Setup row is said in the notes.
"""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from webapp.channel_template import LANGUAGE_ROWS, add_channel_tabs as template_tabs, channel_workbook
from webapp.channels import read_channels
from webapp.tests.test_channels import add_channel_tabs


def header_row(ws, first):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == first)


def validation_for(ws, cell):
    return next((dv for dv in ws.data_validations.dataValidation if cell in dv.sqref), None)


def with_english(path, start, end, days="All"):
    """``path`` with Language Setup's English row covering ``start`` to ``end`` on ``days``."""
    wb = load_workbook(path)
    ws = wb["Language Setup"]
    head = next(r for r in range(1, 10) if str(ws.cell(r, 1).value or "").strip().lower() == "language")
    cols = {str(ws.cell(head, c).value).strip().lower(): c for c in range(1, ws.max_column + 1) if ws.cell(head, c).value}
    row = next(r for r in range(head + 1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").strip() == "English")
    ws.cell(row, cols["coverage start"]).value = start
    ws.cell(row, cols["coverage end"]).value = end
    ws.cell(row, cols["coverage days"]).value = days
    wb.save(path)
    return path


class TheSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.ws = load_workbook(channel_workbook(cls.dir / "needs.xlsx", 30, ["English", "Arabic"]))["Channel Setup"]
        cls.head = header_row(cls.ws, "Channel")

    def test_twenty_language_rows(self):
        self.assertEqual(LANGUAGE_ROWS, 20)
        self.assertEqual(header_row(self.ws, "All channels together"), self.head + 20 + 2)

    def test_every_column_is_a_drop_down(self):
        for row in (self.head + 1, self.head + 20):
            for col in "ABCDEFG":
                dv = validation_for(self.ws, f"{col}{row}")
                self.assertIsNotNone(dv, f"{col}{row}")
                if col != "C":
                    self.assertEqual(dv.type, "list", f"{col}{row}")

    def test_days_offer_weekdays_and_weekends(self):
        dv = validation_for(self.ws, f"F{self.head + 1}")
        source = self.ws.parent["Channel Setup"]
        first, last = dv.formula1.replace("$", "").split("!")[-1].split(":")
        listed = [c.value for row in source[first:last] for c in row]
        for word in ("All", "Weekdays", "Weekends", "Mon-Fri", "Sun-Thu", "Sat"):
            self.assertIn(word, listed)

    def test_the_sheet_says_days_is_the_day_a_window_starts(self):
        text = " ".join(str(c.value) for row in self.ws.iter_rows(max_row=self.head) for c in row if c.value)
        self.assertIn("the day a window starts", text)

    def test_languages_come_from_language_setup_in_an_input_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Language Setup"
        template_tabs(wb, 30, ["English"])
        dv = validation_for(wb["Channel Setup"], f"B{header_row(wb['Channel Setup'], 'Channel') + 1}")
        self.assertIn("'Language Setup'!", dv.formula1)


class WindowsAddUp(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())

    def book(self, name, row, start="16:00", end="07:00"):
        return with_english(add_channel_tabs(self.dir / name, languages=[row]), start, end)

    def test_a_window_outside_the_languages_coverage_is_refused(self):
        path = self.book("wide.xlsx", ("Phone", "English", 1, "15:00", "08:00", "All", "Yes"))
        with self.assertRaisesRegex(ValueError, r"Channel Setup, row \d+: Phone English 15:00 to 08:00 .*outside "
                                                r"English's coverage in Language Setup"):
            read_channels(path, 60)

    def test_a_window_inside_it_is_read(self):
        for name, row in (("same.xlsx", ("Phone", "English", 1, "16:00", "07:00", "All", "Yes")),
                          ("inner.xlsx", ("Chat", "English", 1, "18:00", "02:00", "Weekdays", "Yes"))):
            with self.subTest(name):
                self.assertTrue(read_channels(self.book(name, row), 60)["languages"])

    def test_a_language_without_a_language_setup_row_is_noted(self):
        path = self.book("arabic.xlsx", ("Phone", "Arabic", 1, "10:00", "20:00", "All", "Yes"))
        notes = read_channels(path, 60)["notes"]
        self.assertTrue(any("Arabic has no row in Language Setup" in n for n in notes), notes)


if __name__ == "__main__":
    unittest.main()
