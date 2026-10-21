# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AP (sample ap01, D approved: one workbook with a sheet for each associate; owner, 2026-10-21: a daily view
for each associate, the same as now, downloadable from the RTA after the changes during the day). The people are the
ones on the page the RTA shows: one LOB, or every LOB of the program when the page is merged; the times are in the
zone picked."""
import io
import unittest
from datetime import date

from openpyxl import load_workbook

from webapp.daily_views import sheet_name
from webapp.programs import ProgramBook
from webapp.tests import test_ap_merge_day as fx
from webapp.tests.test_runs import sign_in, token

WED = fx.WED


class TheSheetNames(unittest.TestCase):
    def test_names_are_clean_unique_and_short(self):
        taken = set()
        a = sheet_name("Associate 027", taken)
        b = sheet_name("Associate 027", taken)
        self.assertEqual((a, b), ("Associate 027", "Associate 027 (2)"))
        self.assertEqual(sheet_name("A/B:C*D?[E]", set()), "A B C D E")
        self.assertLessEqual(len(sheet_name("x" * 80, set())), 31)
        self.assertEqual(sheet_name("   ", set()), "Associate")


class TheDailyViews(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fx._TwoLobs.setUpClass.__func__(cls)
        cls.zone_of = lambda self, unit, on: fx.CAIRO if unit == self.one else fx.LONDON

    @classmethod
    def tearDownClass(cls):
        fx._TwoLobs.tearDownClass.__func__(cls)

    def download(self, client=None, **query):
        query = {"program": self.one, "date": WED.isoformat(), "tz": fx.CAIRO, **query}
        return (client or self.admin).get("/day/daily-views.xlsx", query_string=query)

    def book(self, **query):
        got = self.download(**query)
        self.assertEqual(got.status_code, 200, got.data[:200])
        return load_workbook(io.BytesIO(got.data))

    def cells(self, ws):
        return [[c.value for c in row] for row in ws.iter_rows()]

    def test_one_lob_gives_a_sheet_for_each_of_its_people_only(self):
        self.admin.get("/day", query_string={"program": self.one, "date": WED.isoformat(), "merge": ""})
        wb = self.book()
        self.assertEqual(wb.sheetnames[0], "Everyone")
        self.assertIn("Associate 027", wb.sheetnames)
        self.assertNotIn("Associate B027", wb.sheetnames)

    def test_a_merged_page_gives_every_lobs_people_in_the_zone_picked(self):
        self.admin.get("/day", query_string={"program": self.one, "date": WED.isoformat(), "merge": "1",
                                              "tz": fx.CAIRO})
        wb = self.book()
        self.assertIn("Associate 027", wb.sheetnames)
        self.assertIn("Associate B027", wb.sheetnames)
        everyone = {r[0]: r for r in self.cells(wb["Everyone"])[1:]}
        self.assertEqual(everyone["Associate 027"][1], "Tier 1")
        self.assertEqual(everyone["Associate B027"][1], "Tier 2")
        self.assertEqual(everyone["Associate 027"][2], "11:00 - 20:00")   # Cairo's own time
        self.assertEqual(everyone["Associate B027"][2], "13:00 - 22:00")  # London's 11:00 to 20:00, two hours ahead
        london = self.book(tz=fx.LONDON)
        self.assertEqual({r[0]: r for r in self.cells(london["Everyone"])[1:]}["Associate 027"][2], "09:00 - 18:00")

    def test_a_sheet_lists_the_day_and_what_the_rta_recorded(self):
        self.admin.get("/day", query_string={"program": self.one, "date": WED.isoformat(), "merge": "1",
                                              "tz": fx.CAIRO})
        self.admin.post("/day/activity", data={
            "csrf_token": token(self.admin), "program": self.one, "merge": "1", "date": WED.isoformat(),
            "tz": fx.CAIRO, "anchor": WED.isoformat(), "associate": "Associate B027", "kind": "Coaching",
            "from": "15:00", "to": "15:30", "billable": "0", "with_whom": "Associate 900", "why": "Made-up coaching"})
        text = "\n".join(" | ".join(str(v) for v in row if v is not None) for row in
                         self.cells(self.book()["Associate B027"]))
        self.assertIn("Coaching", text)
        self.assertIn("15:00", text)
        self.assertIn("Associate B027", text.splitlines()[0])

    def test_a_person_who_may_not_open_the_program_gets_nothing(self):
        other = ProgramBook(self.store).add_program("OTHER")
        uid = self.store.add_user("planner", "Planner", "Planner-pass-12", is_admin=False, must_change=False)
        self.store.set_user_programs(uid, [other])
        got = self.download(sign_in(self.app, "planner", "Planner-pass-12"))
        self.assertEqual(got.status_code, 403)

    def test_an_associate_account_is_not_allowed_the_endpoint(self):
        from webapp.auth import associate_allowed
        self.assertFalse(associate_allowed("day_daily_views"))

    def test_the_rta_page_offers_the_download(self):
        page = self.admin.get("/day", query_string={"program": self.one, "date": WED.isoformat()}).get_data(as_text=True)
        self.assertIn("/day/daily-views.xlsx", page)
        self.assertIn("Download daily views", page)


if __name__ == "__main__":
    unittest.main()
