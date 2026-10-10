# © 2026 Omar Mokhtar. All rights reserved.
"""Phase W task 6: the interval target (owner, 2026-10-09: "a drop for each week to choose the complaince target per
interval for example if 90% and on a daily basis view in RTA to have something showing count of intervals above 90%
or the choosen tatget / total intervals").

Each LOB and week has a target (100, 95, 90, 85 or 80 %; the workbook's own "Target" until someone picks one; admins
and supervisors pick). An interval is achieved when the people on the floor reach the target share of its demand;
the RTA counts the day's intervals with demand that reach it, now and in the plan, and shows each interval's share."""
import html
import io
import re
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path
from urllib.parse import urlencode

from openpyxl import load_workbook

from webapp.attendance import DayBook
from webapp.day import at_target
from webapp.exports import build
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_runs import make_app, sign_in, token
from webapp.tests.test_schedules import AFTER, INPUT, REPO

WED = date(2026, 10, 14)
PROGRAM = "AE/AR B2B"


def cell(required, now, plan):
    return {"t": 0, "required": required, "now": now, "plan": plan}


class TheCount(unittest.TestCase):
    def test_at_target_counts_against_required(self):
        cells = [cell(10, 9, 10), cell(10, 8.9, 9), cell(4.5, 4.05, 3), cell(0, 3, 3)]
        found = at_target(cells, 0.9)
        self.assertEqual((found["intervals"], found["now"], found["plan"]), (3, 2, 2))
        self.assertEqual([c["pct"] for c in found["cells"]], [90, 89, 90, None])
        self.assertEqual([c["ok"] for c in found["cells"]], [True, False, True, None])
        self.assertEqual(at_target(cells, 1.0)["now"], 0)

    def test_a_day_without_demand_says_so(self):
        found = at_target([cell(0, 2, 2), cell(0, 0, 0)], 0.9)
        self.assertEqual((found["intervals"], found["now"], found["plan"]), (0, 0, 0))
        from webapp.app import create_app  # noqa: F401  the tile is a template part
        app, *_ = make_app(start_worker=False)
        with app.test_request_context():
            from flask import render_template_string
            text = render_template_string('{% from "_target_tile.html" import target_tile %}{{ target_tile(t, "/week") }}',
                                          t=found)
        self.assertIn("No intervals with demand", text)
        self.assertNotIn("0 of 0", text)


class TheWeekTarget(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program=PROGRAM,
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))

    def test_the_week_target_defaults_to_the_workbook(self):
        self.assertEqual(self.days.target_for(PROGRAM, WED), (1.0, {"kind": "workbook"}))  # the test input: Target 1
        self.store.set_interval_target(PROGRAM, "2026-10-11", 90, self.sara)
        ratio, source = self.days.target_for(PROGRAM, WED)
        self.assertEqual((ratio, source["kind"], source["by"]), (0.9, "set", "Sara"))
        self.assertEqual(self.days.page(PROGRAM, WED)["target"]["target"], 0.9)

    def test_the_daily_summary_carries_the_target(self):
        self.store.set_interval_target(PROGRAM, "2026-10-11", 85, self.sara)
        data, *_ = build(self.store, self.days, WED, WED, kinds=["summary"], program=PROGRAM)
        rows = list(load_workbook(io.BytesIO(data))["Daily summary"].iter_rows(values_only=True))
        row = dict(zip(rows[0], rows[1]))
        found = at_target(self.days.page(PROGRAM, WED)["view"]["cells"], 0.85)
        self.assertEqual((row["Interval target %"], row["Intervals at target"], row["Intervals with demand"]),
                         (85, found["now"], found["intervals"]))

    def test_a_rename_moves_the_targets(self):
        self.store.set_interval_target(PROGRAM, "2026-10-11", 95, self.sara)
        self.store.rename_program(PROGRAM, "AE-AR B2B", [])
        self.assertEqual(self.store.get_interval_target("AE-AR B2B", "2026-10-11")["target"], 95)
        self.assertIsNone(self.store.get_interval_target(PROGRAM, "2026-10-11"))


class TheTargetPages(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app, cls.store, cls.base, _ = make_app(start_worker=False, VALIDATOR_ROOT=str(REPO))
        boss = cls.store.add_user("lead", "Lead", "Lead-pass-123", must_change=False)
        cls.store.update_user(boss, is_supervisor=1, all_programs=1)
        planner = cls.store.add_user("plan", "Plan", "Plan-pass-123", must_change=False)
        cls.store.update_user(planner, all_programs=1)
        cls.store.add_run("aaaaaaaaaaaa", boss, "AR_week.xlsx", "QUICK", "DONE", program=PROGRAM, week_start="2026-10-11")
        cls.app.extensions["schedules"].ensure(cls.store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)
        cls.lead = sign_in(cls.app, "lead", "Lead-pass-123")
        cls.planner = sign_in(cls.app, "plan", "Plan-pass-123")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def set(self, client, target):
        return client.post("/week/target", data={"csrf_token": token(client), "program": PROGRAM,
                                                 "week": "2026-10-11", "target": str(target)})

    def test_a_supervisor_sets_the_target_a_planner_cannot(self):
        self.assertEqual(self.set(self.planner, 85).status_code, 403)
        self.assertIsNone(self.store.get_interval_target(PROGRAM, "2026-10-11"))
        self.assertEqual(self.set(self.lead, 77).status_code, 400)  # not one of the choices
        self.assertEqual(self.set(self.lead, 90).status_code, 303)
        self.assertEqual(self.store.get_interval_target(PROGRAM, "2026-10-11")["target"], 90)
        week = html.unescape(self.lead.get(f"/week?{urlencode({'program': PROGRAM, 'week': '2026-10-11'})}")
                             .get_data(as_text=True))
        self.assertRegex(week, r'<option value="90" selected>90%</option>')
        self.assertIn("100% (the workbook's own target)", week)
        self.assertRegex(week, r"<small>Wed 14</small><b>\d+ of \d+</b>")

    def test_the_day_tile_and_row_use_the_week_target(self):
        self.set(self.lead, 90)
        page = html.unescape(self.lead.get(f"/day?{urlencode({'program': PROGRAM, 'date': WED.isoformat()})}")
                             .get_data(as_text=True))
        found = at_target(self.app.extensions["days"].page(PROGRAM, WED)["view"]["cells"], 0.9)
        share = round(100 * found["now"] / found["intervals"])
        # Re-pinned in Phase Z (2026-10-10): the owner approved sample 04, which moves this count from the summary
        # line into the "Achievement on ..." block above the tabs. The numbers measured here are the same; only the
        # markup that carries them changed (webapp.tests.test_achievement covers the block itself).
        self.assertIn(f'<p class="achv-num"><b>{found["now"]} of {found["intervals"]}</b> intervals at 90% or more</p>',
                      page)
        self.assertIn(f"<b>{share}%</b> of the day's intervals reach the target, with attendance as marked. The plan "
                      f"had {found['plan']} of {found['intervals']}.", page)
        self.assertIn('<span class="tl-lbl">Achieved (target 90%)</span>', page)
        self.assertEqual(len(re.findall(r'<span class="tc (?:attarget|below)">\d+%</span>', page)), found["intervals"])


if __name__ == "__main__":
    unittest.main()
