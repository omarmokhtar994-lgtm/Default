#!/usr/bin/env python3
"""Independent schedule validator for RC9.2 artifacts.

This script does not trust optimizer audit metrics.  It reads the input contract,
re-reads the exported schedule/break tables, and independently recomputes key
hard rules, interval coverage, language minima, break structure, rest, OFF shape,
and cyclic next-Sunday coverage.
"""
from __future__ import annotations
import argparse, csv, hashlib, importlib.util, json, math, re, sys, traceback
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from openpyxl import load_workbook

DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]

class ParsedAssignments(dict):
    """Schedule mapping with duplicate-row evidence preserved for validation."""
    def __init__(self):
        super().__init__()
        self.duplicate_rows: List[Dict[str, Any]] = []

def sha256_file(path: Path) -> str:
    digest=hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda:handle.read(1024*1024),b''): digest.update(chunk)
    return digest.hexdigest()

def norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip()).casefold()

def hhmm_to_min(v: Any) -> Optional[int]:
    if v is None or str(v).strip() == "": return None
    if hasattr(v, "hour") and hasattr(v, "minute"): return int(v.hour) * 60 + int(v.minute)
    text=str(v).strip()
    m=re.search(r"(?<!\d)(\d{1,2}):(\d{2})(?:\s*([AP]M))?", text, flags=re.I)
    if not m: return None
    h=int(m.group(1)); minute=int(m.group(2)); ap=(m.group(3) or "").upper()
    if ap:
        h%=12
        if ap=="PM": h+=12
    if h==24 and minute==0: h=0
    if not (0<=h<24 and 0<=minute<60): return None
    return h*60+minute

def load_engine(engine_path: Path):
    sys.path.insert(0, str(engine_path.parent))
    spec=importlib.util.spec_from_file_location("validator_contract_parser", engine_path)
    mod=importlib.util.module_from_spec(spec); sys.modules[spec.name]=mod; spec.loader.exec_module(mod)
    return mod

def find_header(ws, required=("sf name",)) -> int:
    for r in range(1, min(ws.max_row, 20)+1):
        vals={norm(ws.cell(r,c).value) for c in range(1,ws.max_column+1)}
        if any(req in vals for req in required): return r
    raise ValueError(f"Could not find header in {ws.title}")

def _is_day_header(header_text, short, full):
    """True when a normalized header names this weekday and nothing else.

    Mirrors `_is_day_header` in the engine so input parsing and output
    validation cannot drift apart. A suffixed day header is a DATE
    ("Sunday 12 Jul"); accepting any suffix would re-admit "Sunday Totals" and
    the "sunday shift" header that really appears in one delivered workbook.
    """
    if header_text == short or header_text == full:
        return True
    prefix = full + " "
    return header_text.startswith(prefix) and any(ch.isdigit() for ch in header_text[len(prefix):])


def parse_output_schedule(path: Path, expected_names: List[str]) -> Tuple[Dict[str,List[str]], Dict[str,str]]:
    wb=load_workbook(path,data_only=False,read_only=False)
    ws=wb["Schedule"] if "Schedule" in wb.sheetnames else wb["Final Schedule"]
    header=find_header(ws)
    headers={norm(ws.cell(header,c).value):c for c in range(1,ws.max_column+1)}
    name_col=next((c for h,c in headers.items() if h in {"sf name","associate","associate name","name"}),None)
    lang_col=next((c for h,c in headers.items() if h=="language"),None)
    # Day columns are not always labelled on the same row as the name column.
    # Exported schedules carry the day NAMES on the banner row and the calendar
    # DATES on the row holding "SF Name", so locking to the name row found no
    # day columns and the validator raised, aborting before it evaluated a
    # single rule.  Search the name row first, then its neighbours, accepting
    # short or full day names - the same convention the engine's own
    # _day_columns uses.
    day_cols=None
    for probe in (header, header-1, header+1, header-2):
        if probe<1 or probe>ws.max_row: continue
        probe_headers={norm(ws.cell(probe,c).value):c for c in range(1,ws.max_column+1)}
        candidate={}
        for d,full in zip(DAYS,("sunday","monday","tuesday","wednesday","thursday","friday","saturday")):
            # Exact short name, or a label that begins with the FULL day name
            # ("Sunday 12 Jul").  Deliberately NOT startswith(short name): that
            # binds "Monthly Total" to Mon, "Saturation"/"Satisfaction Score" to
            # Sat and "Month" to Mon, silently reading a non-day column as a
            # day's assignments.
            match=next((c for h,c in probe_headers.items()
                        if _is_day_header(h,norm(d),full)),None)
            candidate[d]=match
        if all(c is not None for c in candidate.values()):
            day_cols=candidate; break
    if name_col is None or day_cols is None:
        raise ValueError(
            "Output schedule is missing name/day columns "
            f"(name_col={name_col}, searched rows {max(1,header-2)}-{header+1} of sheet {ws.title!r})")
    assignments=ParsedAssignments(); languages={}
    for r in range(header+1,ws.max_row+1):
        name=str(ws.cell(r,name_col).value or "").strip()
        if not name: continue
        key=norm(name)
        if key in assignments:
            assignments.duplicate_rows.append({"associate":name,"row":r})
            continue
        assignments[key]=[str(ws.cell(r,day_cols[d]).value or "").strip() for d in DAYS]
        languages[key]=str(ws.cell(r,lang_col).value or "").strip() if lang_col else ""
    wb.close()
    return assignments,languages

def declared_artifact_type(path: Path) -> Optional[str]:
    """Read the workbook's own statement of what it is, or None if it makes none.

    The break sheet's "NOT ASSIGNED IN BEFORE-BREAK ARTIFACT" marker is one
    signal for the stage.  It is a single cell in a sheet a downstream tool
    could rewrite, so this reads a second, independently written signal from
    the Production Summary.  Two signals let the validator report a
    contradiction instead of quietly trusting whichever it happened to read.
    """
    wb=load_workbook(path,data_only=True,read_only=True)
    try:
        sheet=next((wb[s] for s in ["Production Summary"] if s in wb.sheetnames),None)
        if sheet is None: return None
        for row in sheet.iter_rows(min_row=1,max_row=40,max_col=2,values_only=True):
            if not row: continue
            label=norm(row[0] if len(row)>0 else None)
            if label!=norm("Artifact Type"): continue
            value=norm(row[1] if len(row)>1 else None)
            if not value: return None
            if "before break" in value: return "BEST_BEFORE_BREAKS"
            if "after break" in value or "final" in value: return "FINAL_AFTER_BREAKS"
            return "UNRECOGNIZED:"+value
        return None
    finally:
        wb.close()

def parse_breaks(path: Path) -> Tuple[List[Dict[str,Any]], bool]:
    wb=load_workbook(path,data_only=True,read_only=True)
    sheet=next((wb[s] for s in ["Break Schedule Active","Break Schedule"] if s in wb.sheetnames),None)
    out=[]; before_only=False; exception_keys=set()

    def locate_header(rows, required):
        """Find a table header without assuming that it is row one."""
        for row_index, row in enumerate(rows[:20]):
            headers={norm(v):i for i,v in enumerate(row)}
            if all(key in headers for key in required):
                return row_index, headers
        return None, {}

    if sheet is not None:
        rows=list(sheet.iter_rows(values_only=True))
        header_index, headers=locate_header(rows, ("associate","day","duration minutes"))
        # Column positions must come from the header row, never from a
        # positional guess.  A release validator must not silently read a
        # reordered or decorated sheet from the wrong columns.
        if header_index is None:
            wb.close()
            raise ValueError(
                f"Break sheet {sheet.title!r} is missing required column(s) "
                f"['associate', 'day', 'duration minutes']; found no compatible header")

        def cell(rr, key):
            idx=headers.get(key)
            return rr[idx] if idx is not None and idx < len(rr) else None

        for rr in rows[header_index+1:]:
            if not rr: continue
            status=str(cell(rr,"status") or "")
            if "not assigned in before-break" in status.casefold(): before_only=True
            name=str(cell(rr,"associate") or "").strip()
            day=str(cell(rr,"day") or "").strip()
            if not name or norm(day) not in {norm(d) for d in DAYS}: continue
            raw_duration=cell(rr,"duration minutes")
            try:
                duration=int(float(raw_duration or 0))
            except (TypeError, ValueError):
                # Record it as a bad row rather than aborting the whole
                # validation.  A duration of 0 is rejected downstream.
                duration=0
            exception = "no breaks scheduled" in status.casefold() or "no break exception" in status.casefold()
            if exception:
                exception_keys.add((norm(name), norm(day)))
            out.append({
                "associate":name,"day":day,"shift":str(cell(rr,"shift") or ""),
                "type":str(cell(rr,"break type") or ""),"start":cell(rr,"start"),
                "duration":duration,"status":status,"is_exception":exception,
            })

    # A no-break exception is an explicit, auditable business decision and is
    # exported on its own sheet.  The old validator read only the active-break
    # table, so an approved no-break Sunday was incorrectly reported as a
    # missing [15, 15, 30] break set.  Represent each exception as a synthetic
    # zero-break record; the existing exception policy path then exempts that
    # associate/day while still enforcing the configured exception limits.
    exception_sheet=wb["No-Break Exceptions"] if "No-Break Exceptions" in wb.sheetnames else None
    if exception_sheet is not None:
        exception_rows=list(exception_sheet.iter_rows(values_only=True))
        exception_header, exception_headers=locate_header(exception_rows, ("associate","day"))
        if exception_header is not None:
            def exception_cell(rr, key):
                idx=exception_headers.get(key)
                return rr[idx] if idx is not None and idx < len(rr) else None

            for rr in exception_rows[exception_header+1:]:
                if not rr: continue
                name=str(exception_cell(rr,"associate") or "").strip()
                day=str(exception_cell(rr,"day") or "").strip()
                if (not name or norm(name) in {"no exceptions", "none", "n/a"}
                        or norm(day) not in {norm(d) for d in DAYS}):
                    continue
                key=(norm(name), norm(day))
                if key in exception_keys: continue
                evidence=[]
                for column in ("review status", "mandatory comment", "exception proof", "permission source"):
                    value=exception_cell(rr, column)
                    if value is not None and str(value).strip():
                        evidence.append(str(value).strip())
                out.append({
                    "associate":name,
                    "day":day,
                    "shift":str(exception_cell(rr,"shift") or ""),
                    "type":"",
                    "start":None,
                    "duration":0,
                    "status":" | ".join(evidence) or "NO BREAK EXCEPTION",
                    "is_exception":True,
                })
                exception_keys.add(key)

    wb.close(); return out,before_only

def independent_employee_quality(parsed, assignments, shift_map):
    """Recompute the employee_quality gate from the OUTPUT schedule.

    Independent by construction: it reads the published assignment grid, not
    the engine's skeleton object, so an engine that mis-scores its own gate
    cannot hide it here.

    Mirrors employee_operational_quality(): late shifts start at or after 18:00
    or before 04:00; overnight shifts end past midnight; weekend is the Sunday
    and Saturday columns; an isolated OFF is an OFF with a shift either side,
    evaluated circularly across the week boundary.
    """
    late, overnight, weekend = [], [], []
    isolated_off = 0
    for assoc in parsed.associates:
        row = assignments.get(norm(assoc.name)) or [""] * 7
        kinds, shifts = [], []
        for value in row[:7]:
            key = norm(value)
            if key == "off":
                kinds.append("off"); shifts.append(None)
            elif key in {"leave", "pto", "vacation"}:
                kinds.append("leave"); shifts.append(None)
            elif key in shift_map:
                kinds.append("shift"); shifts.append(shift_map[key])
            else:
                kinds.append("blank"); shifts.append(None)
        while len(kinds) < 7:
            kinds.append("blank"); shifts.append(None)

        late.append(sum(1 for s in shifts
                        if s is not None and (s.start_min >= 18 * 60 or s.start_min < 4 * 60)))
        overnight.append(sum(1 for s in shifts if s is not None and s.end_abs_min > 1440))
        weekend.append(sum(1 for d in (0, 6) if shifts[d] is not None))
        isolated_off += sum(
            1 for d in range(7)
            if kinds[d] == "off" and kinds[(d - 1) % 7] == "shift" and kinds[(d + 1) % 7] == "shift"
        )

    def spread(values):
        return (max(values) - min(values)) if values else 0

    return {
        "late_shift_load_delta": spread(late),
        "overnight_load_delta": spread(overnight),
        "weekend_load_delta": spread(weekend),
        "isolated_offday_violation_count": isolated_off,
        "roster_size": len(parsed.associates),
    }

def max_gap_run(flags: List[Optional[bool]], intervals_per_day: int) -> int:
    best=0
    for d in range(7):
        run=0
        for flag in flags[d*intervals_per_day:(d+1)*intervals_per_day]:
            if flag is None: run=0
            elif flag: run+=1; best=max(best,run)
            else: run=0
    return best

def validate(input_path: Path, output_path: Path, engine_path: Path) -> Dict[str,Any]:
    eng=load_engine(engine_path)
    from canonical_metrics import canonicalize_metrics
    parsed=eng.parse_input(input_path)
    assignments,output_languages=parse_output_schedule(output_path,[a.name for a in parsed.associates])
    breaks,before_only=parse_breaks(output_path)
    # Two independent statements of the stage, so a disagreement is visible.
    declared_type=declared_artifact_type(output_path)
    break_sheet_role="BEST_BEFORE_BREAKS" if before_only else "FINAL_AFTER_BREAKS"
    artifact_role=break_sheet_role
    stage_declaration_conflict=None
    if declared_type is not None and declared_type!=break_sheet_role:
        stage_declaration_conflict={
            "type":"ARTIFACT_STAGE_DECLARATION_CONFLICT",
            "break_sheet_role":break_sheet_role,
            "production_summary_artifact_type":declared_type,
        }
        # A workbook that cannot say what stage it is cannot be released. The
        # stricter of the two readings is used for the metric labelling so the
        # conflict never lets after-break figures be presented as break
        # results, and the conflict itself is a hard failure below.
        artifact_role="BEST_BEFORE_BREAKS"
    shift_map={norm(s.label):s for s in parsed.shifts}
    failures=[]; warnings=[]
    # #38a: independently recompute the employee_quality gate. It is one of
    # three release gates the engine previously self-reported with no check,
    # and the only one observed returning a verdict other than PASS.
    employee_quality_independent = independent_employee_quality(parsed, assignments, shift_map)
    for _key, _cap_attr, _label in (
        ("late_shift_load_delta", "employee_max_late_shift_load_delta", "LATE_SHIFT"),
        ("overnight_load_delta", "employee_max_overnight_load_delta", "OVERNIGHT"),
        ("weekend_load_delta", "employee_max_weekend_load_delta", "WEEKEND"),
    ):
        _cap = getattr(parsed, _cap_attr, None)
        _observed = employee_quality_independent[_key]
        if isinstance(_cap, int) and _cap >= 0 and _observed > _cap:
            warnings.append({
                "type": f"EMPLOYEE_QUALITY_{_label}_LOAD_SPREAD",
                "observed_spread": _observed,
                "declared_cap": _cap,
                "note": "recomputed from the output schedule, independently of the engine gate",
            })

    if stage_declaration_conflict is not None:
        failures.append(stage_declaration_conflict)
    matrix=[]
    for a,assoc in enumerate(parsed.associates):
        row=assignments.get(norm(assoc.name))
        if row is None:
            failures.append({"type":"MISSING_ASSOCIATE","associate":assoc.name}); row=[""]*7
        matrix.append(row)
    extra=sorted(set(assignments)-{norm(a.name) for a in parsed.associates})
    if extra: failures.append({"type":"EXTRA_OUTPUT_ASSOCIATES","count":len(extra),"examples":extra[:10]})
    duplicate_rows=list(getattr(assignments,"duplicate_rows",[]))
    if duplicate_rows:
        failures.append({"type":"DUPLICATE_OUTPUT_ASSOCIATE","count":len(duplicate_rows),"examples":duplicate_rows[:20]})
    output_language_mismatches=[]
    for assoc in parsed.associates:
        actual=output_languages.get(norm(assoc.name),"")
        if actual and norm(actual)!=norm(assoc.language):
            output_language_mismatches.append({"associate":assoc.name,"expected":assoc.language,"actual":actual})
    if output_language_mismatches:
        failures.append({"type":"OUTPUT_LANGUAGE_MISMATCH","count":len(output_language_mismatches),"examples":output_language_mismatches[:20]})

    # Assignment legality, fixed/leave/OFF, OFF shape, shift variety and rest.
    for a,assoc in enumerate(parsed.associates):
        row=matrix[a]; off_days=[]; shift_days=[]
        for d,value in enumerate(row):
            kind="off" if norm(value)=="off" else "leave" if norm(value) in {"leave","pto","vacation"} else "shift" if norm(value) in shift_map else "blank"
            if kind=="blank": failures.append({"type":"UNKNOWN_OR_BLANK_ASSIGNMENT","associate":assoc.name,"day":DAYS[d],"value":value})
            if kind=="off": off_days.append(d)
            if kind=="shift": shift_days.append((d,shift_map[norm(value)]))
            pref_kind=eng.preference_kind(assoc.preferences[d] if d<len(assoc.preferences) else "")
            fixed_kind=eng.preference_kind(assoc.fixed_schedule[d] if d<len(assoc.fixed_schedule) else "")
            if (
                ((parsed.fixed_enabled and fixed_kind=="leave") or (parsed.leave_enabled and pref_kind=="leave"))
                and kind!="leave"
            ):
                failures.append({"type":"LEAVE_VIOLATION","associate":assoc.name,"day":DAYS[d],"actual":value})
            if (
                ((parsed.fixed_enabled and fixed_kind=="off") or (parsed.hard_off and pref_kind=="off"))
                and kind!="off"
            ):
                failures.append({"type":"HARD_OFF_VIOLATION","associate":assoc.name,"day":DAYS[d],"actual":value})
            if parsed.fixed_enabled and fixed_kind=="shift" and norm(value)!=norm(assoc.fixed_schedule[d]):
                failures.append({"type":"FIXED_SHIFT_VIOLATION","associate":assoc.name,"day":DAYS[d],"expected":assoc.fixed_schedule[d],"actual":value})
            if kind=="shift" and parsed.language_working_window_mode in {"ALL_ROWS", "MINIMUM_ROWS"}:
                entries = eng.associate_language_windows(parsed, assoc, day=d)
                if entries and not any(
                    eng.shift_within_language_window(shift_map[norm(value)], window)
                    for window in entries
                ):
                    failures.append({
                        "type":"LANGUAGE_WORKING_WINDOW",
                        "associate":assoc.name,
                        "language":assoc.language,
                        "day":DAYS[d],
                        "shift":value,
                        "window":", ".join(f"{eng.hhmm(start)}-{eng.hhmm(end)}" for start, end in entries),
                        "mode":parsed.language_working_window_mode,
                    })
                if parsed.language_working_window_mode=="REQUIRED_LANGUAGE_ONLY":
                    blocked = eng.shift_overlaps_required_language_for_noneligible(
                        parsed, assoc, shift_map[norm(value)], d
                    )
                    if blocked is not None:
                        failures.append({
                            "type":"REQUIRED_LANGUAGE_ONLY_VIOLATION",
                            "associate":assoc.name,
                            "language":assoc.language,
                            "required_language":sorted(blocked.required_languages),
                            "day":DAYS[d],
                            "shift":value,
                            "window":f"{eng.hhmm(blocked.start_min)}-{eng.hhmm(blocked.end_min)}",
                            "mode":parsed.language_working_window_mode,
                        })
        # Contract constant shared with the engine, not an independent judgement.
        # Independence means recomputing coverage from the exported workbook - it
        # does not mean re-guessing what the contract says a long shift is.  This
        # was a local 660 while the engine used 630, so a shift in [630, 660) was
        # long to the engine and short to the validator, producing a false
        # OFF_COUNT_VIOLATION that would block an otherwise valid release.
        long_mode=any(s.duration_min>=eng.LONG_SHIFT_MIN_DURATION_MIN for _,s in shift_days)
        expected_off=3 if long_mode else 2
        if parsed.strict_off and len(off_days)!=expected_off:
            failures.append({"type":"OFF_COUNT_VIOLATION","associate":assoc.name,"expected":expected_off,"actual":len(off_days),"off_days":[DAYS[d] for d in off_days]})
        if parsed.strict_off and not parsed.separate_off_days:
            consecutive=any(d in off_days and (d+1)%7 in off_days for d in range(7))
            if not consecutive: failures.append({"type":"CONSECUTIVE_OFF_VIOLATION","associate":assoc.name,"off_days":[DAYS[d] for d in off_days]})
        distinct=len({norm(s.label) for _,s in shift_days})
        if distinct>parsed.max_different_shifts:
            failures.append({"type":"MAX_SHIFT_VARIETY_VIOLATION","associate":assoc.name,"actual":distinct,"maximum":parsed.max_different_shifts})
        for d in range(6):
            left=shift_map.get(norm(row[d])); right=shift_map.get(norm(row[d+1]))
            if left and right and not eng.rest_compatible(left,right,parsed.rest_gap_hours):
                failures.append({"type":"REST_VIOLATION","associate":assoc.name,"from_day":DAYS[d],"to_day":DAYS[d+1],"from":left.label,"to":right.label})
        sat=shift_map.get(norm(row[6])); sun=shift_map.get(norm(row[0]))
        if sat and sun and not eng.rest_compatible(sat,sun,parsed.rest_gap_hours):
            failures.append({"type":"CYCLIC_REST_VIOLATION","associate":assoc.name,"from":"Sat","to":"Sun","from_shift":sat.label,"to_shift":sun.label})
        if sun and not eng.previous_saturday_compatible(assoc.previous_saturday,sun,parsed.rest_gap_hours):
            failures.append({"type":"PREVIOUS_SATURDAY_REST_VIOLATION","associate":assoc.name,"previous_saturday":assoc.previous_saturday,"sunday":sun.label})

    # Flexible nesting groups are equal-schedule constraints.  Exact fixed rows
    # deliberately have their nesting tag cleared by the parser because their
    # day values are already authoritative and may differ by design.
    nesting_groups=defaultdict(list)
    for a,assoc in enumerate(parsed.associates):
        if assoc.nesting_group:
            nesting_groups[norm(assoc.nesting_group)].append((assoc.name,[norm(value) for value in matrix[a]]))
    for group,members in nesting_groups.items():
        distinct={tuple(row) for _,row in members}
        if len(distinct)>1:
            failures.append({
                "type":"NESTING_GROUP_SCHEDULE_MISMATCH","group":group,
                "members":[{"associate":name,"schedule":row} for name,row in members[:20]],
            })

    # Build shift occurrences at quarter level.
    horizon=7*96
    before=[[] for _ in range(horizon)]
    # Keep current-week assignments separate from historical Saturday carry-in.
    # The blank-interval contract is specifically "no NEW staffing in the
    # current week"; counting prior-week spillover as current-week staffing
    # created false hard failures on valid nesting/rest-safe schedules.
    current_week_before=[[] for _ in range(horizon)]
    for a,assoc in enumerate(parsed.associates):
        for d,value in enumerate(matrix[a]):
            shift=shift_map.get(norm(value))
            if not shift: continue
            start=d*96+shift.start_min//15
            for q in range(start,start+shift.duration_q):
                if 0<=q<horizon:
                    before[q].append(a)
                    current_week_before[q].append(a)
        # Previous Saturday spill into current Sunday.
        #
        # Parse the label directly rather than requiring it to appear in THIS
        # week's shift library.  Carry-in is historical fact: the associate
        # worked that shift last week.  The library constrains what may be
        # ASSIGNED this week - it is filtered by the allowed durations and start
        # window - so demanding membership silently discarded real coverage.
        #
        # On NMG SP the library holds a single label while two associates carry
        # in "16:00 - 01:00" and "21:00 - 06:00"; both were dropped, and the
        # validator under-counted Sunday coverage against the engine by exactly
        # the observed margin.  The validator was also inconsistent with itself,
        # since it already parses this same label directly for the rest check
        # via eng.previous_saturday_compatible.
        pparts=eng.shift_parts(assoc.previous_saturday)
        if pparts:
            pstart,_,pdur=pparts
            start=-96+pstart//15
            for q in range(start,start+pdur//15):
                if 0<=q<96: before[q].append(a)

    break_qslots=defaultdict(set); break_qslots_all=defaultdict(set); break_fail=[]; break_spacing_rows=[]
    by_cell=defaultdict(list); exception_cells=set()
    for br in breaks:
        aidx=next((i for i,a in enumerate(parsed.associates) if norm(a.name)==norm(br['associate'])),None)
        didx=next((i for i,d in enumerate(DAYS) if norm(d)==norm(br['day'])),None)
        if aidx is None or didx is None: break_fail.append({"type":"UNKNOWN_BREAK_ASSOCIATE_OR_DAY","row":br}); continue
        if br.get('is_exception'):
            exception_cells.add((aidx,didx)); continue
        shift=shift_map.get(norm(matrix[aidx][didx])); start_clock=hhmm_to_min(br['start'])
        if shift is None or start_clock is None or br['duration']<=0:
            break_fail.append({"type":"INVALID_BREAK_ROW","row":br}); continue
        if start_clock % 15 or int(br['duration']) % 15:
            break_fail.append({"type":"BREAK_NOT_QUARTER_HOUR_ALIGNED","associate":parsed.associates[aidx].name,"day":DAYS[didx],"break":br})
        rel=start_clock-shift.start_min
        if rel<0: rel+=1440
        if rel<0 or rel+br['duration']>shift.duration_min:
            break_fail.append({"type":"BREAK_OUTSIDE_SHIFT","associate":parsed.associates[aidx].name,"day":DAYS[didx],"break":br})
        abs_start=didx*96+(shift.start_min+rel)//15
        duration_q=math.ceil(br['duration']/15)
        br['_rel_q']=rel//15; br['_duration_q']=duration_q; br['_normalized_type']=norm(br.get('type'))
        for q in range(abs_start,abs_start+duration_q):
            if q in break_qslots_all[aidx]: break_fail.append({"type":"OVERLAPPING_BREAK","associate":parsed.associates[aidx].name,"day":DAYS[didx],"qslot":q})
            break_qslots_all[aidx].add(q)
            if 0<=q<horizon:
                break_qslots[aidx].add(q)
        by_cell[aidx,didx].append(br)
    failures.extend(break_fail)
    if not before_only:
        expected_segments=[(q*15,label) for q,label in parsed.break_segments_q]
        for a,assoc in enumerate(parsed.associates):
            for d,value in enumerate(matrix[a]):
                if norm(value) not in shift_map: continue
                cell=by_cell.get((a,d),[])
                if (a,d) in exception_cells:
                    if cell:
                        failures.append({"type":"BREAKS_PRESENT_ON_EXCEPTION_CELL","associate":assoc.name,"day":DAYS[d]})
                    continue
                durations=sorted(int(r['duration']) for r in cell)
                expected=sorted(m for m,_ in expected_segments)
                if durations!=expected:
                    failures.append({"type":"BREAK_SEGMENT_COUNT_OR_DURATION","associate":assoc.name,"day":DAYS[d],"expected":expected,"actual":durations})
                    continue
                ordered=sorted(cell,key=lambda row:int(row['_rel_q']))
                actual_labels=[row['_normalized_type'] for row in ordered]
                expected_labels=[norm(label) for _,label in expected_segments]
                if actual_labels!=expected_labels:
                    failures.append({"type":"BREAK_ORDER_OR_TYPE","associate":assoc.name,"day":DAYS[d],"expected":expected_labels,"actual":actual_labels})
                shift=shift_map[norm(value)]
                for row in ordered:
                    rel_q=int(row['_rel_q']); length_q=int(row['_duration_q'])
                    margin=parsed.break_edge_margin_q
                    if margin is not None and (rel_q < int(margin) or rel_q+length_q > shift.duration_q-int(margin)):
                        failures.append({"type":"BREAK_EDGE_MARGIN","associate":assoc.name,"day":DAYS[d],"break_type":row.get('type'),"relative_start_q":rel_q,"edge_margin_q":margin})
                    rule=(parsed.break_window_rules_q or {}).get(norm(row.get('type')),{})
                    earliest=rule.get('earliest_q'); latest=rule.get('latest_q')
                    if earliest is not None and rel_q<int(earliest) or latest is not None and rel_q>int(latest):
                        failures.append({"type":"BREAK_WINDOW","associate":assoc.name,"day":DAYS[d],"break_type":row.get('type'),"relative_start_q":rel_q,"earliest_q":earliest,"latest_q":latest})
                for previous,current in zip(ordered,ordered[1:]):
                    gap=int(current['_rel_q'])-(int(previous['_rel_q'])+int(previous['_duration_q']))
                    if gap<int(parsed.break_min_gap_q):
                        failures.append({"type":"BREAK_MINIMUM_GAP","associate":assoc.name,"day":DAYS[d],"actual_gap_q":gap,"minimum_gap_q":parsed.break_min_gap_q})
                    if gap==int(parsed.break_preferred_gap_q): classification="PREFERRED"
                    elif gap<int(parsed.break_preferred_gap_q): classification="FLEXIBLE_COMPRESSION"
                    elif gap<=int(parsed.break_normal_max_gap_q): classification="FLEXIBLE_EXTENSION"
                    else: classification="EXTENDED_BEYOND_NORMAL_RANGE"
                    break_spacing_rows.append({
                        "associate":assoc.name,"day":DAYS[d],"from_segment":previous.get('type'),
                        "to_segment":current.get('type'),"actual_gap_minutes":gap*15,
                        "preferred_gap_minutes":int(parsed.break_preferred_gap_q)*15,
                        "normal_maximum_gap_minutes":int(parsed.break_normal_max_gap_q)*15,
                        "classification":classification,
                    })

        if exception_cells and (not parsed.allow_no_break_exceptions or len(exception_cells)>int(parsed.max_no_break_exceptions)):
            failures.append({"type":"NO_BREAK_EXCEPTION_POLICY","actual":len(exception_cells),"allowed":bool(parsed.allow_no_break_exceptions),"maximum":int(parsed.max_no_break_exceptions)})

    quality_gate_failures=[]; quality_gate_warnings=[]; quality_gate_suppressed=[]
    concurrency_violations=[]; max_concurrent=0; max_concurrent_ratio=0.0
    for qslot in range(horizon):
        staffed=len(before[qslot]); on_break=sum(qslot in slots for slots in break_qslots.values())
        max_concurrent=max(max_concurrent,on_break)
        max_concurrent_ratio=max(max_concurrent_ratio,on_break/max(1,staffed))
        allowed=eng.maximum_concurrent_breaks(parsed,staffed)
        if on_break>allowed:
            concurrency_violations.append({"day":DAYS[qslot//96],"time":eng.hhmm((qslot%96)*15),"staffed":staffed,"on_break":on_break,"maximum":allowed})
    if concurrency_violations:
        issue={"type":"BREAK_CONCURRENCY","count":len(concurrency_violations),"examples":concurrency_violations[:20]}
        break_mode=norm(parsed.break_concurrency_gate_mode)
        if break_mode=="fail": quality_gate_failures.append(issue)
        elif break_mode=="warn": quality_gate_warnings.append(issue); warnings.append(issue)
        else: quality_gate_suppressed.append({**issue,"gate":"break_concurrency","gate_mode":break_mode})
    extended_spacing=[row for row in break_spacing_rows if row['classification']=="EXTENDED_BEYOND_NORMAL_RANGE"]
    if extended_spacing:
        warnings.append({"type":"BREAK_SPACING_BEYOND_NORMAL_RANGE","count":len(extended_spacing),"examples":extended_spacing[:20]})

    after=[]
    for q,covers in enumerate(before):
        after.append([a for a in covers if q not in break_qslots[a]])

    # Coverage and language.
    interval_rows=[]; zero=[]; language_gaps=[]; opening_gaps=[]; floor_flags=[]; blank_staffed=[]
    before100=before90=before80=after100=after90=after80=before_target=after_target=before_floor=after_floor=0
    for d in range(7):
        openings=set(eng.opening_intervals_for_day(parsed,d))
        for i in range(parsed.intervals_per_day):
            if not parsed.active[d][i]:
                for q in range(parsed.qslots_per_interval):
                    slot=d*96+i*parsed.qslots_per_interval+q
                    if current_week_before[slot]:
                        blank_staffed.append({
                            "day":DAYS[d],
                            "time":eng.hhmm(i*parsed.interval_minutes+q*15),
                            "current_week_before_raw":len(current_week_before[slot]),
                            "before_raw":len(before[slot]),
                            "after_raw":len(after[slot]),
                            "prior_carry_in_raw":max(
                                0, len(before[slot])-len(current_week_before[slot])
                            ),
                        })
                floor_flags.append(None); continue
            req=float(parsed.requirements[d][i] or 0.0); eff=1-float(parsed.shrinkage[d][i]); qpi=parsed.qslots_per_interval
            bvals=[]; avals=[]
            for q in range(qpi):
                slot=d*96+i*qpi+q
                b=len(before[slot]); a=len(after[slot]); bvals.append(b); avals.append(a)
                minute=i*parsed.interval_minutes+q*15
                if a<=0: zero.append({"day":DAYS[d],"time":eng.hhmm(minute)})
                if parsed.opening_guard_enabled and i in openings and a<parsed.opening_minimum:
                    opening_gaps.append({"day":DAYS[d],"time":eng.hhmm(minute),"actual":a,"minimum":parsed.opening_minimum})
                for rule in eng.language_rules_at(parsed, minute, day=d):
                    if not rule.active or not rule.overlaps(minute,15): continue
                    eligible=sum(1 for ai in after[slot] if norm(parsed.associates[ai].language) in rule.eligible_languages)
                    if eligible<rule.minimum:
                        language_gaps.append({"day":DAYS[d],"time":eng.hhmm(minute),"group":rule.group,"minimum":rule.minimum,"actual":eligible})
            be=sum(bvals)*eff/qpi; ae=sum(avals)*eff/qpi
            bp=be/req if req>0 else 1.0; ap=ae/req if req>0 else 1.0
            before100+=bp>=1-1e-9; before90+=bp>=.9-1e-9; before80+=bp>=.8-1e-9
            after100+=ap>=1-1e-9; after90+=ap>=.9-1e-9; after80+=ap>=.8-1e-9
            before_target+=bp+1e-9>=parsed.target_ratio; after_target+=ap+1e-9>=parsed.target_ratio
            before_floor+=bp+1e-9>=parsed.floor_ratio; after_floor+=ap+1e-9>=parsed.floor_ratio
            floor_flags.append(ap+1e-9<parsed.floor_ratio)
            # Same ceil tolerance the engine applies.  Without it a requirement
            # that is integral in exact arithmetic but marginally above integral
            # in floating point rounds up by a whole associate, overstating
            # unavoidable staffing and understating avoidable overage - a silent
            # disagreement with the engine on the exact metric under review.
            unavoidable=math.ceil(req*parsed.target_ratio/max(eff,1e-9)-eng.OVERAGE_CEIL_TOLERANCE)*eff
            avoid=max(0.0,ae-unavoidable)
            soft_threshold=max(unavoidable,req*float(parsed.overage_soft_cap_ratio))
            severe_threshold=max(unavoidable,req*float(parsed.overage_severe_cap_ratio))
            extreme_threshold=max(unavoidable,req*float(parsed.overage_extreme_cap_ratio))
            interval_rows.append({
                "day":DAYS[d],"day_index":d,"interval_index":i,
                "interval":eng.hhmm(i*parsed.interval_minutes),"required":req,
                "before_effective":be,"after_effective":ae,"before_pct":bp,"after_pct":ap,
                "before_raw_min":min(bvals,default=0),"before_raw_max":max(bvals,default=0),
                "after_raw_min":min(avals,default=0),"after_raw_max":max(avals,default=0),
                "avoidable_overage_fte":avoid,
                "soft_overage":bool(ae>soft_threshold+1e-9),
                "severe_overage":bool(ae>severe_threshold+1e-9),
                "extreme_overage":bool(ae>extreme_threshold+1e-9),
            })
    active=len(interval_rows); severe_threshold=max(0,parsed.floor_ratio-.10)
    severe=sum(1 for r in interval_rows if r['after_pct']+1e-9<severe_threshold)
    maxrun=max_gap_run(floor_flags,parsed.intervals_per_day)
    overages=[r['avoidable_overage_fte'] for r in interval_rows]
    mean=sum(overages)/max(1,len(overages)); variance=sum((x-mean)**2 for x in overages)/max(1,len(overages))
    positive_overages=[value for value in overages if value>1e-9]
    top10_concentration=(sum(sorted(positive_overages,reverse=True)[:10])/sum(positive_overages)) if positive_overages else 0.0
    requirements_sorted=sorted(float(row['required']) for row in interval_rows)
    median_requirement=requirements_sorted[len(requirements_sorted)//2] if requirements_sorted else 0.0
    low_demand_overage=sum(float(row['avoidable_overage_fte']) for row in interval_rows if float(row['required'])<=median_requirement)
    maximum_overage_run=0
    for d in range(7):
        run=0
        for row in (r for r in interval_rows if int(r['day_index'])==d):
            if float(row['avoidable_overage_fte'])>1e-9:
                run+=1; maximum_overage_run=max(maximum_overage_run,run)
            else:
                run=0

    # Independent whole-week distribution gate.  Total overage alone is not
    # sufficient: avoidable excess concentrated in a few intervals, or abrupt
    # raw-HC jumps unsupported by demand, is also release debt.
    whole_week_cap_violations=[]; whole_week_imbalance=[]
    if parsed.whole_week_balance_enabled:
        for row in interval_rows:
            cap=eng.whole_week_raw_cap(parsed,int(row['day_index']),int(row['interval_index']))
            if int(row['after_raw_max'])>cap and float(row['avoidable_overage_fte'])>1e-9:
                whole_week_cap_violations.append({
                    "day":row['day'],"interval":row['interval'],"actual":row['after_raw_max'],
                    "cap":cap,"avoidable_overage_fte":row['avoidable_overage_fte'],
                })
        for d in range(7):
            for qslot in range(d*96,(d+1)*96-1):
                previous_interval=(qslot-d*96)//parsed.qslots_per_interval
                current_interval=(qslot+1-d*96)//parsed.qslots_per_interval
                if not (0<=previous_interval<parsed.intervals_per_day and 0<=current_interval<parsed.intervals_per_day): continue
                if not (parsed.active[d][previous_interval] and parsed.active[d][current_interval]): continue
                delta=abs(len(after[qslot+1])-len(after[qslot]))
                allowed=eng.whole_week_adjacent_raw_limit(parsed,d,previous_interval,current_interval)
                if delta>allowed:
                    whole_week_imbalance.append({
                        "day":DAYS[d],"from":eng.hhmm((qslot-d*96)*15),
                        "to":eng.hhmm((qslot+1-d*96)*15),"change":delta,"maximum":allowed,
                    })
    whole_week_issues=[]
    if len(whole_week_cap_violations)>int(parsed.whole_week_max_overage_cap_violations):
        whole_week_issues.append({"code":"WHOLE_WEEK_OVERAGE_CAP_EXCEEDED","actual":len(whole_week_cap_violations),"maximum":int(parsed.whole_week_max_overage_cap_violations),"examples":whole_week_cap_violations[:20]})
    if len(whole_week_imbalance)>int(parsed.whole_week_max_imbalance_violations):
        whole_week_issues.append({"code":"WHOLE_WEEK_ADJACENT_IMBALANCE_EXCEEDED","actual":len(whole_week_imbalance),"maximum":int(parsed.whole_week_max_imbalance_violations),"examples":whole_week_imbalance[:20]})
    if whole_week_issues:
        issue={"type":"WHOLE_WEEK_BALANCE","issues":whole_week_issues}
        whole_week_mode=norm(parsed.whole_week_gate_mode)
        if whole_week_mode=="fail": quality_gate_failures.append(issue)
        elif whole_week_mode=="warn": quality_gate_warnings.append(issue); warnings.append(issue)
        else: quality_gate_suppressed.append({**issue,"gate":"whole_week_balance","gate_mode":whole_week_mode})

    # Next-Sunday is reconstructed from Saturday spill plus the recurring
    # Sunday assignment, including the matching Saturday/Sunday break cells.
    # The engine deliberately protects only the Saturday-spill horizon. Once
    # a legal next-Sunday shift can start on its own, later Sunday coverage
    # belongs to the ordinary Sunday contract, not to the carry-out gate.
    # Iterating over all 24 Sunday intervals here duplicated current-week
    # floor gaps as false NEXT_SUNDAY_CARRY_OUT failures.
    protected_next_sunday_intervals = set(eng.next_sunday_interval_indices(parsed))
    # The solver constrains next-Sunday adjacency between neighbouring QUARTER
    # slots, and the engine metric measures it the same way. Comparing whole
    # intervals here (the interval maximum against the next interval's maximum)
    # checked a weaker rule under the same name: on a 60-minute workbook it saw
    # a quarter of the adjacent pairs, and an interval maximum hides a cliff
    # inside the hour. That made this a check of something the engine does not
    # enforce rather than an independent check of what it does.
    protected_next_sunday_qslots = {
        qslot for _interval, _quarter, qslot in eng.next_sunday_interval_quarters(parsed)
    }
    next_quarter_raw = []
    next_rows=[]; next_zero=[]; next_language=[]; next_opening=[]; next_blank=[]
    next_openings=set(eng.opening_intervals_for_day(parsed,0))
    for i in range(parsed.intervals_per_day):
        if i not in protected_next_sunday_intervals:
            continue
        active_row=bool(parsed.active[0][i])
        req=float(parsed.requirements[0][i] or 0); eff=1-float(parsed.shrinkage[0][i]); before_vals=[]; after_vals=[]
        for q in range(parsed.qslots_per_interval):
            minute=i*parsed.interval_minutes+q*15; pseudo=7*96+i*parsed.qslots_per_interval+q
            occurrences=[]
            for a,_assoc in enumerate(parsed.associates):
                sat=shift_map.get(norm(matrix[a][6])); sun=shift_map.get(norm(matrix[a][0]))
                if sat:
                    st=6*96+sat.start_min//15
                    if st<=pseudo<st+sat.duration_q: occurrences.append((a,pseudo))
                if sun:
                    st=7*96+sun.start_min//15
                    if st<=pseudo<st+sun.duration_q: occurrences.append((a,pseudo-7*96))
            remaining=[(a,source_q) for a,source_q in occurrences if source_q not in break_qslots_all[a]]
            before_raw=len(occurrences); after_raw=len(remaining)
            before_vals.append(before_raw); after_vals.append(after_raw)
            if pseudo in protected_next_sunday_qslots:
                next_quarter_raw.append((pseudo, i, after_raw))
            boundary_on_break = sum(
                source_q in break_qslots_all[a] for a, source_q in occurrences
            )
            max_concurrent = max(max_concurrent, boundary_on_break)
            boundary_ratio = boundary_on_break / before_raw if before_raw else 0.0
            max_concurrent_ratio = max(max_concurrent_ratio, boundary_ratio)
            boundary_allowed = eng.maximum_concurrent_breaks(parsed, before_raw)
            if boundary_on_break > boundary_allowed:
                concurrency_violations.append({
                    "day": "Next Sun", "time": eng.hhmm(minute),
                    "staffed": before_raw, "on_break": boundary_on_break,
                    "maximum": boundary_allowed, "scope": "NEXT_SUNDAY_CYCLIC",
                })
            if not active_row:
                if before_raw: next_blank.append({"time":eng.hhmm(minute),"before_raw":before_raw,"after_raw":after_raw})
                continue
            if after_raw<=0: next_zero.append({"time":eng.hhmm(minute)})
            if parsed.opening_guard_enabled and i in next_openings and after_raw<parsed.opening_minimum:
                next_opening.append({"time":eng.hhmm(minute),"actual":after_raw,"minimum":parsed.opening_minimum})
            for rule in eng.language_rules_at(parsed, minute, day=0):
                if not rule.active or not rule.overlaps(minute,15): continue
                eligible=sum(1 for a,_source in remaining if norm(parsed.associates[a].language) in rule.eligible_languages)
                if eligible<rule.minimum:
                    next_language.append({"time":eng.hhmm(minute),"group":rule.group,"minimum":rule.minimum,"actual":eligible})
        if active_row:
            be=sum(before_vals)*eff/max(1,len(before_vals)); ae=sum(after_vals)*eff/max(1,len(after_vals)); pct=ae/req if req>0 else 1
            next_rows.append({"interval_index":i,"interval":eng.hhmm(i*parsed.interval_minutes),"before_effective":be,"after_effective":ae,"pct":pct,"after_raw_max":max(after_vals,default=0)})

    failures.extend({"type":"ZERO_STAFF_ACTIVE","detail":row} for row in zero)
    failures.extend({"type":"LANGUAGE_MINIMUM","detail":row} for row in language_gaps)
    failures.extend({"type":"OPENING_MINIMUM","detail":row} for row in opening_gaps)
    if parsed.blank_requirement_mode=="hard_no_current_week_staffing" and blank_staffed:
        failures.append({"type":"BLANK_INTERVAL_STAFFING","count":len(blank_staffed),"examples":blank_staffed[:20]})
    if parsed.floor_mode=="hard":
        hard_ratio=float(parsed.hard_floor_ratio or parsed.floor_ratio)
        hard_gaps=[r for r in interval_rows if float(r['after_pct'])+1e-9<hard_ratio]
        if hard_gaps: failures.append({"type":"HARD_FLOOR","count":len(hard_gaps),"ratio":hard_ratio,"examples":hard_gaps[:20]})
    next_floor_gaps=[row for row in next_rows if float(row['pct'])+1e-9<float(parsed.floor_ratio)]
    if next_floor_gaps or next_zero or next_language or next_opening or (parsed.blank_requirement_mode=="hard_no_current_week_staffing" and next_blank):
        failures.append({"type":"NEXT_SUNDAY_CARRY_OUT","floor_gap_count":len(next_floor_gaps),"zero_count":len(next_zero),"language_gap_count":len(next_language),"opening_gap_count":len(next_opening),"blank_staffed_count":len(next_blank) if parsed.blank_requirement_mode=="hard_no_current_week_staffing" else 0})

    next_overage_cap_violations=[]; next_imbalance=[]; next_adjacent_deltas=[]
    for row in next_rows:
        i=int(row['interval_index'])
        cap=eng.next_sunday_raw_cap(parsed,i)
        if int(row.get('after_raw_max',0))>cap: next_overage_cap_violations.append({"interval":row['interval'],"actual":row.get('after_raw_max'),"cap":cap})
    ordered_quarter_raw=sorted(next_quarter_raw)
    for (previous_q,previous_i,previous_raw),(current_q,current_i,current_raw) in zip(
            ordered_quarter_raw, ordered_quarter_raw[1:]):
        # Only genuinely neighbouring quarters: a gap in the protected horizon
        # is not an adjacency, exactly as the solver treats it.
        if current_q != previous_q + 1: continue
        delta=abs(int(current_raw)-int(previous_raw))
        allowed=eng.next_sunday_adjacent_raw_limit(parsed,int(previous_i),int(current_i))
        # Every adjacency contributes to the observed maximum step, including
        # the ones that stay inside their limit.
        next_adjacent_deltas.append(delta)
        if delta>allowed:
            next_imbalance.append({
                "from":eng.hhmm(previous_i*parsed.interval_minutes
                                +(previous_q-eng.TOTAL_QSLOTS-previous_i*parsed.qslots_per_interval)*15),
                "to":eng.hhmm(current_i*parsed.interval_minutes
                              +(current_q-eng.TOTAL_QSLOTS-current_i*parsed.qslots_per_interval)*15),
                "change":delta,"maximum":allowed,
                "previous_raw":int(previous_raw),"current_raw":int(current_raw)})
    boundary_issue={"type":"NEXT_SUNDAY_BALANCE","overage_cap_violations":len(next_overage_cap_violations),"imbalance_violations":len(next_imbalance),"examples":(next_overage_cap_violations+next_imbalance)[:20]}
    if next_overage_cap_violations or next_imbalance:
        boundary_mode=norm(parsed.next_sunday_balance_gate_mode)
        if boundary_mode=="fail": quality_gate_failures.append(boundary_issue)
        elif boundary_mode=="warn": quality_gate_warnings.append(boundary_issue); warnings.append(boundary_issue)
        else: quality_gate_suppressed.append({**boundary_issue,"gate":"next_sunday_balance","gate_mode":boundary_mode})

    quality_limits=eng.coverage_quality_limits(parsed,active)
    quality_issues=[]
    if active-int(after_floor)>quality_limits['maximum_floor_gaps']: quality_issues.append({"code":"FLOOR_GAP_LIMIT","actual":active-int(after_floor),"maximum":quality_limits['maximum_floor_gaps']})
    if severe>quality_limits['maximum_severe_gaps']: quality_issues.append({"code":"SEVERE_GAP_LIMIT","actual":severe,"maximum":quality_limits['maximum_severe_gaps']})
    if maxrun>quality_limits['maximum_consecutive_floor_gaps']: quality_issues.append({"code":"FLOOR_RUN_LIMIT","actual":maxrun,"maximum":quality_limits['maximum_consecutive_floor_gaps']})
    if parsed.protected_after80_minimum is not None and int(after80)<int(parsed.protected_after80_minimum): quality_issues.append({"code":"PROTECTED_AFTER80_MINIMUM","actual":int(after80),"minimum":int(parsed.protected_after80_minimum)})
    if parsed.minimum_after_break_target_ratio is not None and active and int(after_target)/active+1e-12<float(parsed.minimum_after_break_target_ratio): quality_issues.append({"code":"AFTER_TARGET_RATIO_MINIMUM","actual":int(after_target)/active,"minimum":float(parsed.minimum_after_break_target_ratio)})
    quality_gate_mode=norm(parsed.quality_gate_mode)
    if quality_issues:
        coverage_quality_gate_status=("FAIL" if quality_gate_mode=="fail" else
                                      "WARN" if quality_gate_mode=="warn" else
                                      "NOT_ENFORCED")
    else:
        coverage_quality_gate_status="PASS"
    if quality_issues:
        issue={"type":"PRODUCTION_QUALITY_GATE","issues":quality_issues}
        if quality_gate_mode=="fail":
            # Quality debt is release-blocking, but it is not a hard-rule
            # defect in the workbook artifact. Keep it separate so a hard-valid
            # schedule can still be retained, independently checked, and
            # reported as review-only by the runner.
            quality_gate_failures.append(issue)
            quality_gate_status="FAIL"
        elif quality_gate_mode=="warn":
            quality_gate_warnings.append(issue); warnings.append(issue)
        else:
            quality_gate_suppressed.append({**issue,"gate":"coverage","gate_mode":quality_gate_mode})
    if quality_gate_failures:
        quality_gate_status="FAIL"
    elif quality_gate_warnings:
        quality_gate_status="WARN"
    elif quality_gate_suppressed:
        quality_gate_status="NOT_ENFORCED"
    else:
        quality_gate_status="PASS"
    metrics={
        "active_intervals":active,"before100":int(before100),"before90":int(before90),"before80":int(before80),
        "after100":int(after100),"after90":int(after90),"after80":int(after80),
        "before_target":int(before_target),"after_target":int(after_target),"before_floor":int(before_floor),"after_floor":int(after_floor),
        "floor_gaps":active-int(after_floor),"severe_floor_gaps":severe,"max_consecutive_floor_gaps":maxrun,
        "zero_staffed_active_quarters":len(zero),"language_gap_count":len(language_gaps),"opening_gap_count":len(opening_gaps),
        "avoidable_overage_fte_sum":sum(overages),"avoidable_overage_peak_fte":max(overages,default=0),
        "avoidable_overage_variance_fte":variance,"avoidable_overage_stddev_fte":math.sqrt(max(0,variance)),
        "avoidable_overage_positive_interval_count":len(positive_overages),
        "avoidable_overage_top10_concentration":top10_concentration,
        "avoidable_overage_max_consecutive_intervals":maximum_overage_run,
        "avoidable_overage_low_demand_fte_sum":low_demand_overage,
        "overage_distribution_median_requirement":median_requirement,
        "soft_overage_interval_count":sum(bool(r['soft_overage']) for r in interval_rows),
        "severe_overage_interval_count":sum(bool(r['severe_overage']) for r in interval_rows),
        "extreme_overage_interval_count":sum(bool(r['extreme_overage']) for r in interval_rows),
        "whole_week_overage_cap_violation_count":len(whole_week_cap_violations),
        "whole_week_imbalance_violation_count":len(whole_week_imbalance),
        "blank_staffed_quarters":len(blank_staffed),"no_break_exception_count":len(exception_cells),
        "break_spacing_observation_count":len(break_spacing_rows),
        "break_spacing_compressed_count":sum(row['classification']=="FLEXIBLE_COMPRESSION" for row in break_spacing_rows),
        "break_spacing_extended_count":sum(row['classification']=="FLEXIBLE_EXTENSION" for row in break_spacing_rows),
        "break_spacing_beyond_normal_count":len(extended_spacing),
        "max_concurrent_breaks_observed":max_concurrent,"max_concurrent_break_ratio_observed":max_concurrent_ratio,
        "break_concurrency_violation_count":len(concurrency_violations),
        "next_sunday_target_hits":sum(r['pct']+1e-9>=parsed.target_ratio for r in next_rows),
        "next_sunday_floor_hits":sum(r['pct']+1e-9>=parsed.floor_ratio for r in next_rows),
        "next_sunday_floor_gap_count":len(next_floor_gaps),"next_sunday_zero_staffed_quarters":len(next_zero),
        "next_sunday_language_gap_count":len(next_language),"next_sunday_opening_gap_count":len(next_opening),
        "next_sunday_blank_staffed_quarters":(
            len(next_blank)
            if parsed.blank_requirement_mode=="hard_no_current_week_staffing" else 0
        ),
        "next_sunday_overage_cap_violation_count":len(next_overage_cap_violations),
        "next_sunday_imbalance_violation_count":len(next_imbalance),
        # B-9: independently derived counterparts for six metrics that decide
        # hard validity or candidate selection.  Each is recomputed from the
        # parsed workbook rows above, never read from the optimizer audit.
        "target_losses_from_breaks":sum(
            1 for r in interval_rows
            if float(r['before_pct'])+1e-9>=parsed.target_ratio
            and float(r['after_pct'])+1e-9<parsed.target_ratio),
        "floor_losses_from_breaks":sum(
            1 for r in interval_rows
            if float(r['before_pct'])+1e-9>=parsed.floor_ratio
            and float(r['after_pct'])+1e-9<parsed.floor_ratio),
        "before_severe_floor_gap_count":sum(
            1 for r in interval_rows if float(r['before_pct'])+1e-9<severe_threshold),
        # Mirrors the engine exactly: no hard floor configured means the count
        # is 0, NOT a fall back to the soft floor.  The HARD_FLOOR failure
        # above deliberately falls back; this metric deliberately does not.
        "hard_floor_gap_count":(
            0 if parsed.hard_floor_ratio is None
            else sum(1 for r in interval_rows
                     if float(r['after_pct'])+1e-9<float(parsed.hard_floor_ratio))),
        "next_sunday_hard_failure_count":(
            len(next_floor_gaps)+len(next_zero)+len(next_language)+len(next_opening)
            +(len(next_blank)
              if parsed.blank_requirement_mode=="hard_no_current_week_staffing" else 0)),
        "next_sunday_max_adjacent_raw_change":max(next_adjacent_deltas,default=0),
        "next_sunday_max_coverage_ratio":round(
            max((float(r['pct']) for r in next_rows),default=0.0),8),
        "quality_gate_issue_count":len(quality_issues),
        "quality_gate_status":quality_gate_status,
        "coverage_quality_gate_status":coverage_quality_gate_status,
    }
    return {
        "schema_version":2,"validator":"RC9.2.2_INDEPENDENT_VALIDATOR_HARDENED","input":str(input_path),"output":str(output_path),
        "input_sha256":sha256_file(input_path),"output_sha256":sha256_file(output_path),"engine_sha256":sha256_file(engine_path),
        "artifact_role":artifact_role,
        "artifact_role_signals":{
            "break_sheet":break_sheet_role,
            "production_summary":declared_type,
            "agree":declared_type is None or declared_type==break_sheet_role,
        },
        # Stage 2 never ran on a before-break artifact, so its after_* metrics
        # are the same schedule measured a second time, not a break outcome.
        # Both facts are published rather than the numbers being dropped: a
        # consumer needs the values AND needs to know what produced them.
        "break_stage_executed":artifact_role!="BEST_BEFORE_BREAKS",
        "after_metrics_basis":(
            "NO_BREAKS_PLACED_AFTER_EQUALS_BEFORE" if artifact_role=="BEST_BEFORE_BREAKS"
            else "BREAKS_PLACED_BY_STAGE_2"),
        "status":"PASS" if not failures else "FAIL","hard_fail_count":len(failures),"warning_count":len(warnings),
        "quality_gate_status":quality_gate_status,
        "metrics":metrics,
        "employee_quality_independent":employee_quality_independent,
        "canonical_metrics":canonicalize_metrics(
            metrics, "independent_validator",
            stage="BEFORE_BREAKS_ONLY" if artifact_role=="BEST_BEFORE_BREAKS" else "FULL_SCHEDULE"),
        "failures":failures[:500],"quality_gate_failures":quality_gate_failures[:500],
        "quality_gate_suppressed":quality_gate_suppressed[:500],
        "warnings":warnings[:500],"interval_rows":interval_rows,"next_sunday_rows":next_rows,
        "evidence_note":"Metrics were recomputed from workbook schedule/break cells and input contract; optimizer audit values were not used."
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--input',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    p.add_argument('--engine',type=Path,default=Path(__file__).resolve().parents[1]/'_tools/l632_universal_scheduler.py')
    p.add_argument('--json-out',type=Path)
    p.add_argument('--csv-out',type=Path)
    args=p.parse_args()
    jout=args.json_out or args.output.with_name(args.output.stem+'_INDEPENDENT_VALIDATION.json')
    cout=args.csv_out or args.output.with_name(args.output.stem+'_INDEPENDENT_VALIDATION.csv')
    try:
        result=validate(args.input,args.output,args.engine)
        return_code=0 if result['status']=='PASS' else 2
    except Exception as exc:
        result={
            'schema_version':2,'validator':'RC9.2.2_INDEPENDENT_VALIDATOR_HARDENED',
            'input':str(args.input),'output':str(args.output),'status':'ERROR',
            'hard_fail_count':1,'warning_count':0,'metrics':{},
            'failures':[{'type':'VALIDATOR_EXCEPTION','exception':type(exc).__name__,'message':str(exc)}],
            'traceback':traceback.format_exc(limit=20),
            'input_sha256':sha256_file(args.input) if args.input.is_file() else None,
            'output_sha256':sha256_file(args.output) if args.output.is_file() else None,
            'engine_sha256':sha256_file(args.engine) if args.engine.is_file() else None,
        }
        return_code=3
    jout.parent.mkdir(parents=True,exist_ok=True); jout.write_text(json.dumps(result,indent=2,default=str),encoding='utf-8')
    cout.parent.mkdir(parents=True,exist_ok=True)
    with cout.open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f); w.writerow(['Metric','Value']);
        for k,v in result['metrics'].items(): w.writerow([k,v])
        w.writerow(['status',result['status']]); w.writerow(['hard_fail_count',result['hard_fail_count']]); w.writerow(['warning_count',result['warning_count']])
        for row in result.get('failures',[]): w.writerow(['failure',json.dumps(row,sort_keys=True,default=str)])
        for row in result.get('quality_gate_failures',[]): w.writerow(['quality_gate_failure',json.dumps(row,sort_keys=True,default=str)])
        for row in result.get('quality_gate_suppressed',[]): w.writerow(['quality_gate_suppressed',json.dumps(row,sort_keys=True,default=str)])
    print(json.dumps({k:result.get(k) for k in ['status','hard_fail_count','warning_count','artifact_role','metrics','failures']},indent=2))
    return return_code
if __name__=='__main__': raise SystemExit(main())
