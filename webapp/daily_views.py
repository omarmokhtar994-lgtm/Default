# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AP (sample ap01, D approved: one workbook, a sheet for each associate): the daily view of each person on the
page the RTA shows, taken after the day's changes, in the zone picked."""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional, Set

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .day import CHANNEL_WORD, my_day, my_items

RISKY = ("=", "+", "-", "@", "\t", "\r")
BAD_IN_SHEET_NAMES = re.compile(r"[\[\]:*?/\\]+")


def sheet_name(name: str, taken: Set[str]) -> str:
    """A worksheet name for ``name``: no characters Excel refuses, at most 31 long, different from the ones taken
    (compared without regard to case); the chosen name is added to ``taken``."""
    clean = " ".join(BAD_IN_SHEET_NAMES.sub(" ", name).split())[:31].strip() or "Associate"
    found, n = clean, 1
    while found.casefold() in taken:
        n += 1
        suffix = f" ({n})"
        found = clean[:31 - len(suffix)].rstrip() + suffix
    taken.add(found.casefold())
    return found


def _clock(minute: int, moved: int) -> str:
    minute = (minute + moved) % 1440
    return f"{minute // 60:02d}:{minute % 60:02d}"


def people_of(view: Dict[str, Any], moved: int, lob_names: Optional[Dict[str, str]] = None) -> List[Dict[str, Any]]:
    """Everyone whose shift starts in ``view`` with their day: the shift, the status, what they are on hour by hour
    (work or channel, breaks) and what the RTA recorded (aux, meetings, coaching, overtime, VTO). Times are the view's
    minutes shown ``moved`` minutes later."""
    out = []
    for lane in view["lanes"]:
        day = my_day(view, lane["name"])
        if day is None:
            continue
        seg = next(s for s in lane["segments"] if s["offset"] == 0)
        breaks = {(b["start"], b["kind"]) for b in seg["breaks"]}
        band = [(_clock(p["lo"], moved), _clock(p["hi"], moved),
                 "Work" if p["kind"] == "work" else p["what"] if p["kind"] == "break" else
                 CHANNEL_WORD.get(p["kind"], p["what"])) for p in day["band"]]
        recorded = [(_clock(i["lo"], moved), _clock(i["hi"], moved), i["what"]) for i in my_items(view, lane["name"])
                    if (i["lo"], i["what"]) not in breaks]
        out.append({"name": lane["name"], "lob": (lob_names or {}).get(lane.get("lob", ""), ""),
                    "shift": f"{_clock(day['start'], moved)} - {_clock(day['end'], moved)}", "status": day["status"],
                    "band": band, "recorded": recorded})
    return out


def _safe(value: Any) -> Any:
    return "'" + value if isinstance(value, str) and value.startswith(RISKY) else value


def workbook(people: List[Dict[str, Any]], heading: str) -> bytes:
    """The workbook: a sheet "Everyone" (name, LOB, shift, status) and one sheet for each person."""
    wb = Workbook()
    bold = Font(bold=True)
    first = wb.active
    first.title = "Everyone"
    first.append(["Name", "LOB", "Shift", "Status"])
    for person in people:
        first.append([_safe(person["name"]), _safe(person["lob"]), person["shift"], person["status"]])
    first.append([])
    first.append([heading])
    for cell in first[1]:
        cell.font = bold
    taken = {"everyone"}
    for person in people:
        ws = wb.create_sheet(sheet_name(person["name"], taken))
        ws.append([_safe(person["name"] + (f" ({person['lob']})" if person["lob"] else ""))])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([heading])
        ws.append(["Shift", person["shift"]])
        ws.append(["Status", person["status"]])
        ws.append([])
        ws.append(["From", "To", "Doing"])
        for cell in ws[ws.max_row]:
            cell.font = bold
        for row in person["band"]:
            ws.append(list(row))
        if person["recorded"]:
            ws.append([])
            ws.append(["Recorded on the RTA", "", ""])
            ws[f"A{ws.max_row}"].font = bold
            ws.append(["From", "To", "What"])
            for cell in ws[ws.max_row]:
                cell.font = bold
            for row in person["recorded"]:
                ws.append(list(row))
        for col, width in enumerate((22, 12, 34), 1):
            ws.column_dimensions[get_column_letter(col)].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
