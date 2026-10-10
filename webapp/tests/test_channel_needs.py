# © 2026 Omar Mokhtar. All rights reserved.
"""Phase Z (owner, 2026-10-10: "Associates channels found but i cannot find where i can add the required per channel
and where can it run"; approved sample 03, "Panel + tabs"): channel needs added to a schedule that already exists,
from a workbook made for that week, without a new schedule; every channel reader uses them."""
import html
import importlib.util
import io
import re
import sys
import shutil
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from webapp.channel_template import add_channel_tabs as needs_tabs, channel_workbook
from webapp.channels import check_lines, has_channel_needs, read_channels
from webapp.day import read_inputs
from webapp.programs import ProgramBook
from webapp.tests.test_channels import add_channel_tabs
from webapp.tests.test_ready import make_ready
from webapp.tests.test_runs import REPO, make_app, sign_in, token, wait

WEEK = "2026-10-11"
WED = "2026-10-14"
TABS = ["Channels - read me", "Chat 60 Min", "Phone 60 Min", "Email 60 Min", "Email Hours", "Channel Setup"]
WORKBOOKS = Path(__file__).resolve().parents[1] / "workbooks"
HOME = {"Scheduler_Input_Blank.xlsx": 30, "Scheduler_Input_Example.xlsx": 60}  # each workbook's own interval


def filled(target: Path, step: int = 60) -> Path:
    """A channel needs workbook with Chat 1 and Phone 2 from 08:00 to 22:00 every day."""
    wb = Workbook()
    wb.remove(wb.active)
    needs_tabs(wb, step, ["English"], need=lambda letter, d, t: (0 if not 8 * 60 <= t < 22 * 60 else
                                                                 {"C": 1, "P": 2}[letter]))
    wb.save(target)
    return target


class _Needs(unittest.TestCase):
    """SAKS: NMG Tier 2 has a plain upload that gets channel needs; NMG Tier 3 a plain upload that never does; NMG
    Tier 1 an upload with its own channel tabs that then gets channel needs too. All for the week of Sun 11 Oct."""

    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.ready = make_ready(cls.dir / "ready.xlsx")
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        programs = ProgramBook(cls.store)
        saks = programs.add_program("SAKS")
        cls.key = programs.add_lob(saks, "NMG Tier 2")
        cls.bare_key = programs.add_lob(saks, "NMG Tier 3")
        cls.tabbed_key = programs.add_lob(saks, "NMG Tier 1")
        cls.client = sign_in(cls.app, "omar", "Owner-pass-123")
        cls.book = cls.app.extensions["schedules"]
        cls.plain = cls.upload(cls.ready, "plain.xlsx", cls.key)
        cls.bare = cls.upload(cls.ready, "bare.xlsx", cls.bare_key)
        cls.tabbed = cls.upload(add_channel_tabs(cls.dir / "tabbed.xlsx", src=cls.ready), "tabbed.xlsx",
                                cls.tabbed_key)
        cls.empty_key = programs.add_lob(saks, "NMG Tier 4")
        cls.with_empty = cls.upload(cls.empty_tabs(cls.dir / "empty_tabs.xlsx"), "empty_tabs.xlsx", cls.empty_key)
        cls.runs_before = len(cls.store.list_runs())
        cls.added = cls.attach(cls.plain, filled(cls.dir / "needs.xlsx"))
        cls.added_tabbed = cls.attach(cls.tabbed, filled(cls.dir / "needs_t1.xlsx"))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    @classmethod
    def upload(cls, path, name, program):
        cls.client.post("/runs", data={"csrf_token": token(cls.client), "kind": "ready", "mode": "QUICK",
                                       "program": program, "week_start": WEEK,
                                       "workbook": (io.BytesIO(path.read_bytes()), name)},
                        content_type="multipart/form-data")
        run_id = cls.store.list_runs()[0]["id"]
        wait(cls.store, run_id, statuses=("DONE", "REJECTED", "FAILED"))
        return run_id

    @classmethod
    def empty_tabs(cls, target):
        """The ready workbook (60-minute) with the Blank's empty channel tabs, still at 30 minutes."""
        wb = load_workbook(cls.ready)
        needs_tabs(wb, 30, ["English"])
        wb.save(target)
        return target

    @classmethod
    def attach(cls, run_id, path):
        return cls.client.post(f"/runs/{run_id}/channel-needs",
                               data={"csrf_token": token(cls.client),
                                     "workbook": (io.BytesIO(path.read_bytes()), path.name)},
                               content_type="multipart/form-data", follow_redirects=True)

    def page(self, url):
        got = self.client.get(url)
        self.assertEqual(got.status_code, 200, url)
        return html.unescape(got.get_data(as_text=True))

    def panel(self, page):
        found = re.search(r'(?s)<section[^>]*id="channels".*?</section>', page)
        self.assertIsNotNone(found, "no Channels panel")
        return found.group(0)


class TheChannelNeeds(_Needs):
    def test_the_downloaded_workbook_reads_as_empty_needs(self):
        got = self.client.get(f"/runs/{self.bare}/channel-needs.xlsx")
        self.assertEqual(got.status_code, 200)
        self.assertIn("Channel_needs_SAKS_NMG_Tier_3_2026-10-11.xlsx", got.headers["Content-Disposition"])
        path = self.dir / "download.xlsx"
        path.write_bytes(got.data)
        self.assertEqual(load_workbook(path, read_only=True).sheetnames, TABS)
        setup = read_channels(path, 60)
        for letter in "CP":
            self.assertEqual({v for col in setup["need"][letter].values() for v in col.values()}, {0.0})
            self.assertEqual(len(setup["need"][letter][3]), 24)
        self.assertEqual(setup["email_mode"], "none")
        self.assertEqual((setup["languages"], setup["blended"], setup["notes"]), ([], [], []))

    def test_needs_uploaded_to_a_schedule_are_used_without_a_new_run(self):
        self.assertEqual(len(self.store.list_runs()), self.runs_before)
        self.assertTrue((self.book.root / self.plain / "channels.xlsx").is_file())
        said = html.unescape(self.added.get_data(as_text=True))
        self.assertIn("Channel needs added to plain.xlsx", said)
        version = self.book.versions(self.plain)[0]
        page = self.page(f"/runs/{self.plain}/schedules")
        self.assertIn(f'href="/schedules/{version["id"]}/channels">Plan channels</a>', page)
        self.assertEqual(self.client.get(f"/schedules/{version['id']}/channels").status_code, 200)

    def test_the_rta_counts_channels_from_the_attached_needs(self):
        page = self.page(f"/day?program={self.key.replace(' ', '+')}&date={WED}")
        self.assertRegex(page, r'<a href="/day\?[^"]*view=channels"[^>]*>Channels</a>')
        bare = self.page(f"/day?program={self.bare_key.replace(' ', '+')}&date={WED}")
        self.assertNotIn("view=channels", bare)

    def test_a_wrong_interval_file_is_refused_and_nothing_is_attached(self):
        said = html.unescape(self.attach(self.bare, filled(self.dir / "thirty.xlsx", step=30)).get_data(as_text=True))
        self.assertIn("These channel needs were not added:", said)
        self.assertIn("Chat 30 Min does not match the program's 60-minute interval", said)
        self.assertFalse((self.book.root / self.bare / "channels.xlsx").exists())
        self.assertEqual(list((self.book.root / self.bare).glob("channels*")), [])
        self.assertIsNone(self.book.channel_path(self.bare))

    def test_attached_needs_win_over_the_workbooks_own_tabs(self):
        self.assertEqual(self.book.channel_path(self.tabbed), self.book.root / self.tabbed / "channels.xlsx")
        today = datetime.now(timezone(timedelta(hours=3))).strftime("%a %d %b")
        panel = self.panel(self.page(f"/runs/{self.tabbed}/schedules"))
        self.assertIn(f"From channel needs added on {today}", panel)
        self.assertNotIn("Phone 5 = 7", panel)  # the workbook's own Phone tab is no longer read

    def test_the_panel_shows_three_steps_without_needs(self):
        panel = self.panel(self.page(f"/runs/{self.bare}/schedules"))
        self.assertIn("Channels: Chat, Phone and Email", panel)
        self.assertIn("This schedule has no channel needs yet, so its channels cannot be planned.", panel)
        for step in ("Download the channel needs workbook", "Fill in how many people each channel needs",
                     "Upload it here"):
            self.assertIn(step, panel)
        self.assertIn(f'href="/runs/{self.bare}/channel-needs.xlsx"', panel)
        self.assertIn(f'action="/runs/{self.bare}/channel-needs"', panel)
        self.assertIn("Made for SAKS, NMG Tier 3, week of 11 Oct: Chat 60 Min, Phone 60 Min and Email 60 Min", panel)
        self.assertNotIn("Plan channels", self.page(f"/runs/{self.bare}/schedules"))

    def test_associate_channels_points_to_the_schedules(self):
        page = self.page(f"/setup/channels?program={self.key.replace(' ', '+')}")
        self.assertIn("This page says who can work each channel. How many people each channel needs comes with each "
                      "week's schedule: open", page)
        self.assertRegex(page, r'open <a href="/schedules\?program=SAKS[^"]*">Schedules</a>, then Channels, to add the '
                               r'needs and plan the day\.')

    def test_empty_channel_tabs_at_another_interval_do_not_stop_an_upload(self):
        # The Blank input workbook now carries empty channel tabs at 30 minutes; a 60-minute week filled in it is
        # uploaded and treated as before Phase V (no channel needs), not refused for the tabs nobody filled in.
        self.assertEqual(self.store.get_run(self.with_empty)["workbook"], "empty_tabs.xlsx")  # not refused
        self.assertEqual(self.store.get_run(self.with_empty)["status"], "DONE")
        self.assertIsNone(self.book.channel_path(self.with_empty))
        panel = self.panel(self.page(f"/runs/{self.with_empty}/schedules"))
        self.assertIn("This schedule has no channel needs yet", panel)

    def test_needs_that_ask_for_nobody_are_not_added(self):
        empty = channel_workbook(self.dir / "empty_needs.xlsx", 60, ["English"])
        said = html.unescape(self.attach(self.with_empty, empty).get_data(as_text=True))
        self.assertIn("These channel needs were not added: This workbook asks for nobody on any channel", said)
        self.assertFalse((self.book.root / self.with_empty / "channels.xlsx").exists())

    def test_associate_channels_says_the_home_workbooks_have_the_tabs(self):
        page = self.page(f"/setup/channels?program={self.key.replace(' ', '+')}")
        self.assertIn("The Blank input workbook on Home has the Chat, Phone, Email, Email Hours and Channel Setup tabs "
                      "too.", page)


def _engine():
    """The engine module, loaded read-only (nothing in it is changed or called beyond reading a workbook)."""
    path = Path(__file__).resolve().parents[2] / "engine" / "_tools" / "l632_universal_scheduler.py"
    if "phase_z_engine" not in sys.modules:
        sys.path.insert(0, str(path.parent))
        spec = importlib.util.spec_from_file_location("phase_z_engine", path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = module
        spec.loader.exec_module(module)
    return sys.modules["phase_z_engine"]


class TheNeedsTest(unittest.TestCase):
    """has_channel_needs: channel tabs count only when they ask for someone; anything that is not a number counts,
    so it is read and refused by tab and row rather than ignored."""

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir, True)

    def test_empty_tabs_ask_for_nobody_and_filled_ones_do(self):
        self.assertFalse(has_channel_needs(channel_workbook(self.dir / "empty.xlsx", 30, ["English", "Arabic"])))
        self.assertTrue(has_channel_needs(filled(self.dir / "filled.xlsx")))
        self.assertFalse(has_channel_needs(make_ready(self.dir / "plain.xlsx")))

    def test_a_cell_that_is_not_a_number_counts_so_it_is_refused(self):
        path = channel_workbook(self.dir / "typo.xlsx", 60, ["English"])
        wb = load_workbook(path)
        wb["Chat 60 Min"]["D14"] = "two"
        wb.save(path)
        self.assertTrue(has_channel_needs(path))
        with self.assertRaisesRegex(ValueError, "Chat 60 Min, row 14"):
            read_channels(path, 60)

    def test_a_file_is_read_once_per_version(self):
        # Break advice asks for the day once per candidate time and exports once per day: the input workbook is
        # opened once per version of the file, as read_channels does, and again when the file changes.
        import os
        from unittest import mock
        from webapp import channels
        path = filled(self.dir / "needs.xlsx")
        with mock.patch.object(channels, "load_workbook", wraps=channels.load_workbook) as opened:
            self.assertTrue(has_channel_needs(path))
            self.assertTrue(has_channel_needs(path))
        self.assertEqual(opened.call_count, 1)
        before = path.stat().st_mtime_ns
        channel_workbook(path, 60, ["English"])  # the same file, now asking for nobody
        os.utime(path, ns=(before + 10 ** 9, before + 10 ** 9))
        self.assertFalse(has_channel_needs(path))

    def test_email_hours_or_a_language_minimum_alone_count(self):
        for sheet, cell, value in (("Email Hours", "B7", 4), ("Channel Setup", "A15", "Phone")):
            path = channel_workbook(self.dir / f"{sheet}.xlsx", 60, ["English"])
            wb = load_workbook(path)
            wb[sheet][cell] = value
            wb.save(path)
            self.assertTrue(has_channel_needs(path), sheet)


class TheHomeWorkbooks(unittest.TestCase):
    """Sample 03, last line: the Blank and Example workbooks on Home carry the channel tabs at their own interval,
    and the engine still reads FT Wise as the requirement."""

    def test_the_home_workbooks_have_the_channel_tabs(self):
        for name, step in HOME.items():
            names = load_workbook(WORKBOOKS / name, read_only=True).sheetnames
            for tab in ("Channels - read me", f"Chat {step} Min", f"Phone {step} Min", f"Email {step} Min",
                        "Email Hours", "Channel Setup"):
                self.assertIn(tab, names, name)

    def test_the_engine_still_reads_ft_wise(self):
        engine = _engine()
        for name, step in HOME.items():
            wb = load_workbook(WORKBOOKS / name)
            im = engine._instruction_map(engine._sheet_by_alias(wb, ["Engine Defaults", "Engine Default",
                                                                     "Scheduler Defaults"]))
            im.update(engine._instruction_map(engine._sheet_by_alias(wb, ["Instructions"])))
            self.assertEqual(engine._discover_requirement_sheet(wb, im).title, f"FT Wise {step} Min", name)

    def test_the_example_needs_match_ft_wise(self):
        path = WORKBOOKS / "Scheduler_Input_Example.xlsx"
        inputs = read_inputs(path)
        setup = read_channels(path, inputs["interval"])
        self.assertTrue(any(v for col in setup["need"]["P"].values() for v in col.values()))
        self.assertTrue(any(v for col in setup["need"]["C"].values() for v in col.values()))
        lines = check_lines(setup, inputs, path)
        self.assertEqual([text for level, text in lines if level == "warn"], [])


if __name__ == "__main__":
    unittest.main()
