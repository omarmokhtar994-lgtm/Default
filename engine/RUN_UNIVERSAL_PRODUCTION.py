#!/usr/bin/env python3
from __future__ import annotations
import argparse, atexit, csv, hashlib, json, os, platform, re, shutil, subprocess, sys, threading, zipfile
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / '_tools'))
from canonical_metrics import (
    canonicalize_metrics, compare_metric_surfaces, stage_for_artifact_role,
)
ENGINE = ROOT / "_tools" / "l632_universal_scheduler.py"
POLISHER = ROOT / "production" / "production_output_polisher.py"
QUALITY_REPORTER = ROOT / "production" / "phase_c_quality_report.py"
PACKAGER = ROOT / "production" / "package_phase_c_outputs.py"
VALIDATOR = ROOT / "tools" / "independent_validator.py"
# Shares no code with the engine; see its docstring. Lives beside engine/.
CLEAN_ROOM = ROOT.parent / "tools" / "clean_room_check.py"
def _engine_release() -> str:
    """Read the release identity from the engine itself.

    This was previously a hardcoded literal in this wrapper, and it had drifted:
    the wrapper stamped every run "L6.3.2.4-RC9.2-PROTECTED-BALANCE-RC1" while
    the engine it invoked was "L6.3.2.5-RC9.2.1-...".  Because this value is
    written into run identity and manifest output, RC9.2.1 runs were being
    recorded under the RC9.2 release name, which defeats the exact-engine-
    identity release gate.  Deriving it from the engine's own VERSION constant
    makes that class of drift impossible rather than merely corrected once.

    Fails loudly: a wrong-but-plausible release string is worse than no run.
    """
    text = ENGINE.read_text(encoding='utf-8', errors='replace')
    match = re.search(r'^VERSION\s*=\s*["\'](.+?)["\']', text, re.MULTILINE)
    if not match:
        raise SystemExit(
            f'Cannot determine engine release identity: no VERSION constant in {ENGINE}. '
            'Refusing to stamp runs with an assumed release name.'
        )
    return match.group(1)


RELEASE = _engine_release()

DEFAULT_SKELETON_PROFILES = (
    "target90_restore_champion,target90_restore_productive,release_gate_floor_satisfaction,"
    "target_floor_pareto_master,floor_gate_hunter_before,floor_gate_hunter_productive,"
    "aggregate_floor_binding,quality_convergence,daily_floor_balanced,break_safe_reserve,"
    "target_priority_balanced,floor_protected,coverage_rebalance,before_target_champion,"
    "protected_balance_polish"
)
DEFAULT_BREAK_OBJECTIVES = "target_priority,release_quality_guard,coverage_rebalance,quality_convergence,floor_protected,balanced,target_100"

_ACTIVE_CASE_LOCK: Path | None = None
_ACTIVE_HEARTBEAT: Path | None = None
_HEARTBEAT_STOP = threading.Event()
_HEARTBEAT_THREAD: threading.Thread | None = None


def _contract_run_settings(input_path):
    """Read Run Stage, Run Depth and the workbook's search controls, or blanks.

    Never fatal. If the workbook cannot be parsed here the engine will fail on
    it a moment later with a far better message, and a runner that dies while
    reading a convenience setting would hide that.
    """
    try:
        sys.path.insert(0, str(ENGINE.parent))
        import l632_universal_scheduler as engine_module
        parsed = engine_module.parse_input(Path(input_path))
        return parsed.run_stage, parsed.run_depth, set(parsed.search_controls)
    except Exception as exc:
        print(f'[run] could not read Run Stage/Run Depth from the workbook '
              f'({type(exc).__name__}); using defaults', flush=True)
        return None, None, set()


def engine_flags_for_run(
    args, mode_defaults, runner_parser, workbook_controls
) -> list:
    """Build the engine flags this runner supplies, minus the ones it must not.

    The runner passes a value for nearly every search control, taken from its
    own defaults or from the depth it picked. Those are the runner's opinion,
    not the scheduler's, and on the command line they are indistinguishable
    from a typed argument - so they used to outrank a workbook that had
    explicitly stated the parameter, which made the workbook route useless.

    A flag is therefore dropped when the workbook states that parameter, unless
    the operator typed it on this runner's own command line. Nothing changes
    for a workbook that states none of them.
    """
    keep_all = False
    typed: set = set()
    try:
        sys.path.insert(0, str(ENGINE.parent))
        import l632_universal_scheduler as engine_module
        typed = engine_module.explicitly_supplied_options(runner_parser, None)
    except (Exception, SystemExit):
        # Without a reliable read of what was typed, keep every flag: the
        # previous behavior. Silently letting the workbook win here could
        # change a run the operator thought they had pinned.
        keep_all = True
    flags = []
    for dest, parameter, flag, value in _runner_search_flag_values(args, mode_defaults):
        stated_by_workbook = parameter in workbook_controls
        typed_by_operator = dest is not None and dest in typed
        if stated_by_workbook and not typed_by_operator and not keep_all:
            continue
        flags.extend([flag, value])
    return flags


# (argparse dest on THIS runner or None, engine flag, engine parameter name)
_RUNNER_SEARCH_FLAGS = (
    ('num_workers', '--num-workers', 'workers'),
    ('pattern_widths', '--pattern-widths', 'pattern_widths'),
    ('repair_change_limits', '--repair-change-limits', 'repair_change_limits'),
    ('skeleton_profiles', '--skeleton-profiles', 'skeleton_profile_names'),
    ('break_objective_modes', '--break-objective-modes', 'break_objective_modes'),
    (None, '--min-after90-gain-per-after80-loss', 'min_after90_gain_per_after80_loss'),
    (None, '--max-after80-tradeoff-intervals', 'max_after80_tradeoff_intervals'),
    (None, '--adaptive-no-improvement-attempts', 'adaptive_no_improvement_attempts'),
    (None, '--primary-target-tolerance', 'primary_target_tolerance'),
    (None, '--exception-search-reserve-sec', 'exception_search_reserve_sec'),
    (None, '--post-break-repair-reserve-sec', 'post_break_repair_reserve_sec'),
    (None, '--target-lock-recovery-reserve-sec', 'target_lock_recovery_reserve_sec'),
    (None, '--finalization-reserve-sec', 'finalization_reserve_sec'),
    (None, '--safe-incumbent-reserve-sec', 'safe_incumbent_reserve_sec'),
    (None, '--conflict-refinement-reserve-sec', 'conflict_refinement_reserve_sec'),
    (None, '--coordinated-repair-reserve-sec', 'coordinated_repair_reserve_sec'),
    (None, '--coordinated-repair-cycles', 'coordinated_repair_cycles'),
    (None, '--joint-refinement-reserve-sec', 'joint_refinement_reserve_sec'),
    (None, '--joint-change-limits', 'joint_change_limits'),
    (None, '--joint-shift-options-per-cell', 'joint_shift_options_per_cell'),
    (None, '--joint-patterns-per-shift', 'joint_patterns_per_shift'),
    (None, '--adaptive-joint-attempts', 'adaptive_joint_attempts'),
    (None, '--adaptive-joint-no-improvement-limit', 'adaptive_joint_no_improvement_limit'),
    ('solver_random_seed', '--solver-random-seed', 'solver_random_seed'),
)


def _runner_search_flag_values(args, mode_defaults):
    """Yield (runner dest, engine parameter, flag, value) for each flag set here.

    The runner dest is None when this runner exposes no option of its own for
    that parameter - those values are purely its depth arithmetic, so the
    workbook always outranks them.
    """
    literal = {
        '--min-after90-gain-per-after80-loss': '1.0',
        '--max-after80-tradeoff-intervals': '1',
        '--adaptive-no-improvement-attempts': str(mode_defaults['adaptive']),
        '--primary-target-tolerance': '1',
        '--exception-search-reserve-sec': '180',
        '--post-break-repair-reserve-sec': str(mode_defaults['post']),
        '--target-lock-recovery-reserve-sec': str(mode_defaults['target']),
        '--finalization-reserve-sec': str(mode_defaults['final']),
        '--safe-incumbent-reserve-sec': str(mode_defaults['safe']),
        '--conflict-refinement-reserve-sec': '120',
        '--coordinated-repair-reserve-sec': '300',
        '--coordinated-repair-cycles': '1',
        '--joint-refinement-reserve-sec': str(mode_defaults['joint']),
        '--joint-change-limits': '8,16,24,36,54',
        '--joint-shift-options-per-cell': '10',
        '--joint-patterns-per-shift': '64',
        '--adaptive-joint-attempts': str(mode_defaults['joint_attempts']),
        '--adaptive-joint-no-improvement-limit': str(mode_defaults['joint_no_improve']),
    }
    from_args = {
        '--num-workers': lambda: str(args.num_workers),
        '--pattern-widths': lambda: args.pattern_widths,
        '--repair-change-limits': lambda: args.repair_change_limits,
        '--skeleton-profiles': lambda: args.skeleton_profiles,
        '--break-objective-modes': lambda: args.break_objective_modes,
        '--solver-random-seed': lambda: str(args.solver_random_seed),
    }
    for dest, flag, parameter in _RUNNER_SEARCH_FLAGS:
        value = literal[flag] if flag in literal else from_args[flag]()
        yield dest, parameter, flag, value


def safe_id(value: str) -> str:
    cleaned = ''.join(ch if ch.isalnum() or ch in '_.-' else '_' for ch in str(value or 'schedule'))
    return cleaned.strip('_') or 'schedule'


def _pid_is_alive(pid: int) -> bool:
    if pid <= 0:
        return False
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    except OSError:
        return False
    return True


def _heartbeat_loop(heartbeat_path: Path, schedule_id: str) -> None:
    while not _HEARTBEAT_STOP.wait(30.0):
        payload = {
            'schema_version': 1,
            'schedule_id': schedule_id,
            'pid': os.getpid(),
            'heartbeat_utc': datetime.now(timezone.utc).isoformat(),
        }
        temporary = heartbeat_path.with_name(f'.{heartbeat_path.name}.{os.getpid()}.tmp')
        try:
            temporary.write_text(json.dumps(payload, indent=2), encoding='utf-8')
            temporary.replace(heartbeat_path)
        except OSError:
            temporary.unlink(missing_ok=True)


def _write_initial_heartbeat(heartbeat_path: Path, schedule_id: str) -> None:
    """Publish liveness immediately; do not make operators wait for the first tick."""
    payload = {
        'schema_version': 1,
        'schedule_id': schedule_id,
        'pid': os.getpid(),
        'heartbeat_utc': datetime.now(timezone.utc).isoformat(),
    }
    temporary = heartbeat_path.with_name(f'.{heartbeat_path.name}.{os.getpid()}.tmp')
    try:
        temporary.write_text(json.dumps(payload, indent=2), encoding='utf-8')
        temporary.replace(heartbeat_path)
    except OSError:
        temporary.unlink(missing_ok=True)
        raise RuntimeError(f'Unable to publish run heartbeat: {heartbeat_path}')


HEARTBEAT_FRESH_SEC = 300


def _machine_identity() -> str:
    """This boot of this machine. A pid means nothing outside it."""
    try:
        return Path('/proc/sys/kernel/random/boot_id').read_text(encoding='utf-8').strip()
    except OSError:
        return f'host:{platform.node()}'


def _process_start_token(pid: int) -> "str | None":
    """Kernel start time of `pid` (Linux), so a reused pid is told apart."""
    try:
        stat = Path(f'/proc/{int(pid)}/stat').read_text(encoding='utf-8')
        return stat.rsplit(')', 1)[1].split()[19]
    except (OSError, IndexError, ValueError):
        return None


def _heartbeat_is_fresh(case_root: Path) -> bool:
    try:
        beat = json.loads((case_root / 'RUN_HEARTBEAT.json').read_text(encoding='utf-8'))
        last = datetime.fromisoformat(str(beat.get('heartbeat_utc')))
    except (OSError, ValueError, TypeError, json.JSONDecodeError):
        return False
    return (datetime.now(timezone.utc) - last).total_seconds() <= HEARTBEAT_FRESH_SEC


def _lock_owner_is_running(case_root: Path, existing: dict) -> bool:
    """Is the run that wrote this RUN_LOCK.json still running?

    A pid alone cannot answer it. On Colab the results live on Drive and
    outlive the VM: a run interrupted there leaves its lock behind, and in the
    next VM the same small pid numbers belong to other processes. The old
    check (is that pid alive?) then refused every seed at once with "Case is
    already running" - a whole portfolio ended in 0.4 s with no schedule.
    On the machine that wrote the lock, the pid must be alive with the same
    start time. Anywhere else (another VM sharing the Drive folder, or a lock
    written before this check existed), only a heartbeat newer than
    HEARTBEAT_FRESH_SEC shows a live run; the heartbeat is rewritten every 30 s.
    """
    owner = int(existing.get('pid') or 0)
    if existing.get('machine') and existing.get('machine') == _machine_identity():
        if not _pid_is_alive(owner):
            return False
        recorded, current = existing.get('pid_start'), _process_start_token(owner)
        return recorded is None or current is None or recorded == current
    return _heartbeat_is_fresh(case_root)


def acquire_case_lock(case_root: Path, schedule_id: str) -> None:
    """Prevent concurrent writers and make abandoned runs safely resumable."""
    global _ACTIVE_CASE_LOCK, _ACTIVE_HEARTBEAT, _HEARTBEAT_THREAD
    lock_path = case_root / 'RUN_LOCK.json'
    payload = {
        'schema_version': 2,
        'schedule_id': schedule_id,
        'pid': os.getpid(),
        'pid_start': _process_start_token(os.getpid()),
        'machine': _machine_identity(),
        'host': platform.node(),
        'created_utc': datetime.now(timezone.utc).isoformat(),
    }
    for _ in range(2):
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:
                existing = json.loads(lock_path.read_text(encoding='utf-8'))
            except (OSError, json.JSONDecodeError):
                existing = {}
            owner = int(existing.get('pid') or 0)
            if _lock_owner_is_running(case_root, existing):
                raise RuntimeError(
                    f'Case is already running: {case_root} (pid={owner}, host='
                    f'{existing.get("host") or "unknown"}). Use a different schedule-id or wait for '
                    'that run to finish. If no run is active, delete RUN_LOCK.json in that folder.'
                )
            lock_path.unlink(missing_ok=True)
            continue
        else:
            with os.fdopen(fd, 'w', encoding='utf-8') as handle:
                json.dump(payload, handle, indent=2)
            heartbeat_path = case_root / 'RUN_HEARTBEAT.json'
            _write_initial_heartbeat(heartbeat_path, schedule_id)
            _ACTIVE_CASE_LOCK = lock_path
            _ACTIVE_HEARTBEAT = heartbeat_path
            _HEARTBEAT_STOP.clear()
            _HEARTBEAT_THREAD = threading.Thread(
                target=_heartbeat_loop,
                args=(heartbeat_path, schedule_id),
                name='rc922-heartbeat',
                daemon=True,
            )
            _HEARTBEAT_THREAD.start()
            atexit.register(release_case_lock)
            return
    raise RuntimeError(f'Unable to acquire case lock: {lock_path}')


def release_case_lock() -> None:
    global _ACTIVE_CASE_LOCK, _ACTIVE_HEARTBEAT, _HEARTBEAT_THREAD
    _HEARTBEAT_STOP.set()
    if _HEARTBEAT_THREAD is not None and _HEARTBEAT_THREAD is not threading.current_thread():
        _HEARTBEAT_THREAD.join(timeout=2.0)
    for path in (_ACTIVE_HEARTBEAT, _ACTIVE_CASE_LOCK):
        if path is not None:
            path.unlink(missing_ok=True)
    _ACTIVE_CASE_LOCK = None
    _ACTIVE_HEARTBEAT = None
    _HEARTBEAT_THREAD = None


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def read_summary_metrics(case_root: Path) -> dict:
    summary_files = sorted(case_root.glob('*.l6_3_2_3_summary.csv'))
    if not summary_files:
        summary_files = sorted(case_root.glob('*_SKELETON_ONLY_SUMMARY.csv'))
    if not summary_files:
        return {}
    with summary_files[0].open(newline='', encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    return rows[0] if rows else {}


def hard_valid_artifact(audit_path: Path) -> bool:
    if not audit_path.exists():
        return False
    try:
        audit = json.loads(audit_path.read_text(encoding='utf-8'))
    except Exception:
        return False
    return audit.get('artifact_state') == 'FINAL_VERIFIED' and audit.get('hard_valid_schedule_exists') is True


def engine_selected_metrics(audit_path: Path) -> dict:
    """Read the exact candidate metrics used by the engine export decision."""
    try:
        audit = json.loads(audit_path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        return {}
    return engine_metric_surface(audit)['metrics']


def engine_metric_surface(audit: dict) -> dict:
    """Return the metrics the engine published, with the stage that produced them.

    A FULL_SCHEDULE run publishes ``selected_candidate``; a BEFORE_BREAKS_ONLY
    run stops before a final candidate exists and publishes the Stage-1
    surface instead. Both are read here so the parity gate always compares two
    real metric surfaces. ``stage`` is None only when the engine published
    neither, which is itself a gate failure rather than something to guess at.
    """
    stage_surface = audit.get('stage_metric_surface') or {}
    stage = stage_surface.get('stage')
    metrics = stage_surface.get('metrics') or {}
    if not metrics:
        # Older audits predate stage_metric_surface. Fall back to the keys they
        # do carry and derive the stage from artifact_state, so an archived run
        # is still checkable instead of failing as un-gateable.
        candidate = (audit.get('selected_candidate') or {}).get('metrics') or {}
        stopped_at_stage_1 = (
            str(audit.get('artifact_state') or '').upper() == 'BEST_BEFORE_BREAKS_ONLY'
            or str(audit.get('status') or '').upper() == 'SKELETON_ONLY_COMPLETE')
        if candidate:
            metrics = candidate
            stage = stage or 'FULL_SCHEDULE'
        elif stopped_at_stage_1:
            # A full run that reached Stage 2 and failed also carries a
            # before-break skeleton. Reading its metrics as that run's result
            # would report before-break coverage as an after-break outcome, so
            # only a run that actually stopped at Stage 1 may claim it.
            skeleton = (audit.get('selected_before_break_skeleton') or {}).get('metrics') or {}
            if skeleton:
                metrics = skeleton
                stage = stage or 'BEFORE_BREAKS_ONLY'
    if stage is None:
        state = str(audit.get('artifact_state') or '').upper()
        if state == 'BEST_BEFORE_BREAKS_ONLY':
            stage = 'BEFORE_BREAKS_ONLY'
        elif state.startswith('FINAL_'):
            stage = 'FULL_SCHEDULE'
    return {
        'stage': stage,
        'metrics': metrics,
        'break_stage_executed': stage_surface.get('break_stage_executed'),
        'after_metrics_basis': stage_surface.get('after_metrics_basis'),
    }


def apply_metric_parity_gate(
    audit_path: Path, validation_json: Path, validation_csv: Path
) -> tuple[int, dict]:
    """Attach canonical parity evidence and fail a mismatched release.

    The validator remains independent because it still recomputes coverage from
    workbook cells. This gate only compares the two published metric surfaces
    after normalization; it prevents naming or calculation drift from being
    silently released.
    """
    validation = json.loads(validation_json.read_text(encoding='utf-8'))
    audit = json.loads(audit_path.read_text(encoding='utf-8'))
    surface = engine_metric_surface(audit)
    # The validator names the artifact it actually read. Deriving the stage
    # from that rather than from the runner's intent means a run that asked
    # for BEFORE_BREAKS_ONLY and was handed a final workbook (or the reverse)
    # fails the gate instead of being compared against the wrong stage.
    validator_stage = stage_for_artifact_role(validation.get('artifact_role'))
    parity = compare_metric_surfaces(
        surface['metrics'], validation.get('metrics') or {},
        engine_stage=surface['stage'], validator_stage=validator_stage,
    )
    parity['engine_metrics_published'] = bool(surface['metrics'])
    parity['engine_after_metrics_basis'] = surface.get('after_metrics_basis')
    parity['validator_artifact_role'] = validation.get('artifact_role')
    if not surface['metrics']:
        # An unpublished surface is not parity evidence. Without this the gate
        # would compare {} against {} on a future path that publishes neither
        # and report PASS on no comparison at all.
        parity['status'] = 'FAIL'
        parity['mismatch_count'] += 1
        parity['mismatches'].append({
            'field': 'engine_metric_surface',
            'engine': None,
            'validator': 'PUBLISHED',
            'reason': 'ENGINE_PUBLISHED_NO_METRIC_SURFACE',
        })
    engine_coverage_gate = str(
        (audit.get('production_quality_gate') or {}).get('coverage_gate_status') or ''
    ).upper()
    validator_coverage_gate = str(
        validation.get('coverage_quality_gate_status') or ''
    ).upper()
    if engine_coverage_gate and validator_coverage_gate and engine_coverage_gate != validator_coverage_gate:
        parity['status'] = 'FAIL'
        parity['mismatch_count'] += 1
        parity['mismatches'].append({
            'field': 'coverage_quality_gate_status',
            'engine': engine_coverage_gate,
            'validator': validator_coverage_gate,
            'reason': 'QUALITY_GATE_STATUS_MISMATCH',
        })
    validation['canonical_metrics'] = canonicalize_metrics(
        validation.get('metrics') or {}, 'independent_validator',
        stage=validator_stage,
    )
    validation['metric_parity'] = parity
    if parity['status'] != 'PASS':
        validation['status'] = 'FAIL_METRIC_PARITY'
        validation['hard_fail_count'] = max(1, int(validation.get('hard_fail_count') or 0))
        validation.setdefault('failures', []).append({
            'type': 'METRIC_PARITY_MISMATCH',
            'mismatch_count': parity['mismatch_count'],
            'mismatches': parity['mismatches'][:100],
        })
    validation['return_code'] = 0 if validation['status'] == 'PASS' else 2
    validation_json.write_text(json.dumps(validation, indent=2, default=str), encoding='utf-8')
    with validation_csv.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(['Metric', 'Value'])
        for key, value in (validation.get('metrics') or {}).items():
            writer.writerow([key, value])
        writer.writerow(['status', validation['status']])
        writer.writerow(['hard_fail_count', validation.get('hard_fail_count', 0)])
        writer.writerow(['warning_count', validation.get('warning_count', 0)])
        writer.writerow(['metric_parity_status', parity['status']])
        writer.writerow(['metric_parity_mismatch_count', parity['mismatch_count']])
        writer.writerow(['run_stage', parity.get('engine_stage')])
        writer.writerow(['validator_artifact_role', validation.get('artifact_role')])
        for row in validation.get('failures', []):
            writer.writerow(['failure', json.dumps(row, sort_keys=True, default=str)])
    return int(validation['return_code']), validation


# Engine statuses that end without a schedule but carry the engine's own
# diagnosis of why; the reconciled outcome keeps it.
ENGINE_DIAGNOSED_NO_ARTIFACT = {'FAIL_HARD_CONTRACT_INFEASIBLE', 'FAIL_HARD_CONTRACT_UNKNOWN'}


def _blocking_reasons(case_root: Path, independent_validation: dict, runner_return_code: int) -> list:
    """Plain lines naming what stopped a generated schedule from release.

    Audit F-14: a blocked schedule used to show only the engine's quality
    findings, under "Main blockers", while the gate that actually blocked it
    (a parity mismatch, a validator rule, the clean-room check) was named
    nowhere in the text.
    """
    reasons = []
    status = str(independent_validation.get('status') or 'NOT_RUN')
    for row in ((independent_validation.get('metric_parity') or {}).get('mismatches') or [])[:10]:
        reasons.append(f"Engine and independent validator disagree on {row.get('field')}: "
                       f"engine {row.get('engine')}, validator {row.get('validator')}")
    validation_json = independent_validation.get('json')
    if status == 'FAIL' and validation_json:
        try:
            failures = json.loads(Path(validation_json).read_text(encoding='utf-8')).get('failures') or []
        except (OSError, json.JSONDecodeError):
            failures = []
        counts = {}
        for row in failures:
            counts[str(row.get('type'))] = counts.get(str(row.get('type')), 0) + 1
        for kind, count in sorted(counts.items()):
            reasons.append(f'Independent validation: {kind} x{count}')
    clean_room = independent_validation.get('clean_room') or {}
    if clean_room.get('status') == 'FAIL':
        for rule, count in sorted((clean_room.get('violations_by_rule') or {}).items()):
            reasons.append(f'Clean-room check: {rule} x{count}')
        for row in (clean_room.get('engine_mismatches') or [])[:10]:
            reasons.append(f"Clean-room check disagrees with the engine on {row.get('metric')}: "
                           f"clean room {row.get('clean_room')}, engine {row.get('published')}")
    if status.startswith('ERROR'):
        reasons.append(f'Independent validation did not complete ({status})')
    if status.startswith('SKIPPED'):
        reasons.append(f'Independent validation was skipped ({status})')
    quality_path = case_root / 'PHASE_C_QUALITY_SUMMARY.json'
    if quality_path.is_file():
        try:
            gate = json.loads(quality_path.read_text(encoding='utf-8')).get('production_quality_gate') or {}
        except (OSError, json.JSONDecodeError):
            gate = {}
        for row in gate.get('failures') or []:
            if isinstance(row, dict):
                reasons.append(f"Quality gate in FAIL mode: {row.get('code')}"
                               + (f" (actual {row.get('actual', row.get('count'))})" if row.get('actual', row.get('count')) is not None else ''))
    if not reasons and int(runner_return_code or 0) != 0:
        reasons.append(f'Run returned blocking code {runner_return_code}')
    deduped = []
    for line in reasons:
        if line not in deduped:
            deduped.append(line)
    return deduped


def _outcome_detail_text(outcome: dict, findings_label: "str | None" = None) -> str:
    """The findings, named requests and actions as readable lines.

    BUSINESS_OUTCOME.txt is the file a scheduler opens; it used to stop at the
    one-paragraph summary, so the rules and cells to fix were only in the JSON.
    """
    def finding(item) -> str:
        if not isinstance(item, dict):
            return str(item)[:300]
        if item.get('rule') and item.get('finding'):
            return f"{item['rule']}: {item['finding']}"
        named = ('code', 'failure_code', 'type', 'gate', 'detail', 'message', 'headline', 'summary')
        head = next((str(item[k]) for k in named[:4] if item.get(k)), '')
        text = next((str(item[k]) for k in named[4:] if item.get(k)), '')
        rest = [f'{k}={v}' for k, v in item.items()
                if k not in named and not isinstance(v, (dict, list)) and v not in (None, '')]
        line = ': '.join(x for x in (head, text) if x)
        if rest:
            line = f"{line} ({', '.join(rest)})" if line else ', '.join(rest)
        return line[:300]
    lines = []
    reasons = outcome.get('blocking_reasons') or []
    if reasons:
        lines += ['', 'Why it is blocked:'] + [f'- {r}' for r in reasons[:12]]
    shortfall_rows = outcome.get('shortfall_rows') or []
    if shortfall_rows:
        lines += ['', 'Shortfalls (the attached schedule misses these minimums):']
        for row in shortfall_rows[:20]:
            unit = ' (coverage units)' if str(row.get('unit')) != 'people' else ''
            group = f" [{row.get('rule')}]" if row.get('rule') else ''
            lines.append(f"- {row.get('day')} {row.get('time')} | {row.get('meaning')}{group} | "
                         f"required {row.get('required')}, short by {row.get('short_by')}{unit}")
        if len(shortfall_rows) > 20:
            lines.append(f'- ... and {len(shortfall_rows) - 20} more (the Shortfalls sheet lists them all)')
    check = (outcome.get('independent_validation') or {}).get('shortfall_schedule') or {}
    if check:
        lines += ['', f"Independent check of the shortfall schedule: {check.get('status')}"
                  + (f" (validator failure types: {', '.join(check.get('validator_failure_types') or [])})"
                     if check.get('validator_failure_types') else '')]
    findings = []
    for item in outcome.get('resource_findings') or []:
        text = finding(item)
        if text not in findings:
            findings.append(text)
    if findings:
        label = findings_label or ('Warnings:' if outcome.get('production_eligible') else 'Main blockers:')
        lines += ['', label] + [f'- {text}' for text in findings[:8]]
    examples = outcome.get('affected_examples') or []
    if examples:
        lines += ['', 'Affected examples:']
        for item in examples[:20]:
            if isinstance(item, dict):
                lines.append('- ' + ' | '.join(str(item.get(k) or '') for k in ('associate', 'day', 'shift', 'window')).strip(' |'))
        if len(examples) > 20:
            lines.append(f'- ... and {len(examples) - 20} more (BUSINESS_OUTCOME.json lists them all)')
    for alt in (outcome.get('independent_validation') or {}).get('alternative_exports') or []:
        if alt.get('role') != 'MORE_CONSISTENT_CANDIDATE':
            continue
        if alt.get('status') == 'PASS' and alt.get('coverage_no_worse'):
            before = ((alt.get('consistency') or {}).get('before') or {})
            after = ((alt.get('consistency') or {}).get('after') or {})
            lines += ['', 'A more consistent version of this schedule (same coverage, checked independently):',
                      f"- {Path(str(alt.get('workbook'))).name}",
                      f"- start-time movement {before.get('start_movement_hours')} h -> "
                      f"{after.get('start_movement_hours')} h; distinct start times "
                      f"{before.get('distinct_start_times')} -> {after.get('distinct_start_times')}",
                      '- The schedule above is unchanged; use either.']
    actions = outcome.get('recommended_actions') or []
    if actions:
        lines += ['', 'Required action:'] + [f'- {a}' for a in actions]
    return '\n'.join(lines) + ('\n' if lines else '')


def coverage_measure_line(outcome: dict) -> str:
    """Which coverage measure the schedule was optimised for, and where that came from."""
    measure = outcome.get('coverage_measure') or {}
    if not measure.get('mode'):
        return ''
    label = {'interval_count': 'Interval Count', 'volume_weighted': 'Volume Weighted'}.get(
        measure['mode'], measure['mode'])
    return f"Coverage measure: {label} ({measure.get('source') or 'default'})\n"


def reconcile_business_outcome_after_validation(
    case_root: Path, independent_validation: dict, runner_return_code: int
) -> None:
    """Make user-facing outcome files agree with the final release gates.

    The scheduler writes BUSINESS_OUTCOME before the exact polished workbook is
    independently checked. A hard-valid schedule with quality warnings could
    therefore leave ``production_eligible=true`` beside a later validator
    failure. The runner is the final authority for publication, so it must
    reconcile every business-facing copy before metrics and packaging are
    written.
    """
    outcome_path = case_root / 'BUSINESS_OUTCOME.json'
    debug_outcome_path = case_root / 'debug' / 'BUSINESS_OUTCOME.json'
    source_path = outcome_path if outcome_path.is_file() else debug_outcome_path
    if not source_path.is_file():
        return
    try:
        outcome = json.loads(source_path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        return

    validation_status = str(independent_validation.get('status') or 'NOT_RUN')
    validation_rc = independent_validation.get('return_code')
    blocked = validation_status != 'PASS' or int(runner_return_code or 0) != 0
    # The engine's own findings on a schedule it judged releasable are declared
    # quality debt (warnings), even when a later gate blocks the release.
    engine_said_releasable = bool(outcome.get('production_eligible'))
    outcome.pop('blocking_reasons', None)
    quality_gate_status = 'NOT_EVALUATED'
    quality_report_path = case_root / 'PHASE_C_QUALITY_SUMMARY.json'
    if quality_report_path.is_file():
        try:
            quality_report = json.loads(quality_report_path.read_text(encoding='utf-8'))
            quality_gate_status = str(
                (quality_report.get('production_quality_gate') or {}).get('status')
                or quality_gate_status
            )
        except (OSError, json.JSONDecodeError):
            pass
    validation_record = {
        'status': validation_status,
        'return_code': validation_rc,
        'workbook': independent_validation.get('workbook'),
        'json': independent_validation.get('json'),
        'csv': independent_validation.get('csv'),
        'quality_gate_status': independent_validation.get('quality_gate_status'),
        'coverage_quality_gate_status': independent_validation.get('coverage_quality_gate_status'),
        'metric_parity': independent_validation.get('metric_parity'),
        'clean_room': independent_validation.get('clean_room'),
        'alternative_exports': independent_validation.get('alternative_exports'),
        'shortfall_schedule': independent_validation.get('shortfall_schedule'),
    }
    outcome['independent_validation'] = validation_record
    audit_status = ''
    for audit_path in sorted(case_root.glob('*solver_audit.json')):
        try:
            audit_status = str(json.loads(audit_path.read_text(encoding='utf-8')).get('status') or '')
        except (OSError, json.JSONDecodeError):
            continue
        if audit_status:
            break
    # A declared output path in an audit is not proof that a workbook exists.
    # Reconciliation must distinguish a generated-but-blocked schedule from a
    # pre-solver contract failure, otherwise the latter is reported as if a
    # production schedule had been created.
    audit_paths = sorted(case_root.glob('*solver_audit.json'))
    schedule_generated = (
        any(hard_valid_artifact(path) for path in audit_paths)
        or any(case_root.rglob('*_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx'))
        # Preserve an explicit legacy outcome in isolated reconciliation calls
        # that have no audit yet.  A real pre-solver run has an audit status and
        # therefore cannot reach this compatibility fallback.
        or (not audit_status and str(outcome.get('outcome_code') or '').startswith('FINAL_SCHEDULE_GENERATED'))
    )
    # A BEFORE_BREAKS_ONLY run deliberately produces no final schedule, so
    # `schedule_generated` is false by design. Without this branch the run was
    # reported as "No final schedule workbook was produced, so independent
    # validation could not run" - on a run whose validation had in fact run and
    # passed. The stage keeps its own outcome and never claims releasability.
    if audit_status == 'SKELETON_ONLY_COMPLETE':
        outcome['production_eligible'] = False
        outcome['technical_return_code'] = int(runner_return_code or 0)
        outcome['outcome_code'] = 'BEFORE_BREAK_SKELETON_GENERATED'
        outcome['outcome_category'] = 'REVIEW_ONLY'
        outcome['headline'] = (
            'Before-break skeleton generated for review; breaks are not assigned')
        if blocked:
            outcome['outcome_code'] = 'BEFORE_BREAK_SKELETON_BLOCKED'
            outcome['outcome_category'] = 'ACTION_REQUIRED'
            outcome['headline'] = (
                'Before-break stage completed, but its release gates did not pass')
            outcome['technical_status'] = (
                'FAIL_INDEPENDENT_VALIDATION' if validation_status != 'PASS'
                else 'FAIL_STAGE_GATE')
            outcome['plain_language_summary'] = (
                'The before-break stage produced a skeleton, but '
                f'independent validation reported {validation_status} and the run '
                'returned a blocking code. Resolve the finding before relying on '
                'this artifact even for review.')
        else:
            outcome['technical_status'] = 'SKELETON_ONLY_COMPLETE'
            outcome['plain_language_summary'] = (
                'The run was asked for the before-break stage only, and it completed '
                'and passed independent validation. The workbook contains shift and '
                'OFF assignments with no breaks, so it is a review artifact and must '
                'not be used operationally. Run the full stage to produce a '
                'production schedule.')
        outcome['recommended_actions'] = [
            'Use this artifact to judge whether the roster is worth taking to break placement.',
            'Rerun with the full stage to produce a schedule that can be released.',
        ]

    elif audit_status == 'PASS_DIAGNOSTICS_ONLY' and not schedule_generated:
        # Audit F-14: a diagnostics-only run produces no schedule by design.
        # It is neither a success nor an "engine output problem".
        outcome['production_eligible'] = False
        outcome['technical_return_code'] = int(runner_return_code or 0)
        outcome['technical_status'] = 'PASS_DIAGNOSTICS_ONLY'
        outcome['outcome_code'] = 'DIAGNOSTICS_ONLY_COMPLETE'
        outcome['outcome_category'] = 'REVIEW_ONLY'
        outcome['headline'] = 'Diagnostics completed; no schedule was produced'
        outcome['plain_language_summary'] = (
            'This was a diagnostics-only run: the workbook was read and checked and the hard rules '
            'were tested for satisfiability, but no shifts or breaks were scheduled, so there was '
            'nothing to validate.')
        outcome['recommended_actions'] = ['Rerun at QUICK depth or deeper to produce a schedule.']

    elif blocked:
        outcome['production_eligible'] = False
        outcome['technical_return_code'] = int(runner_return_code or validation_rc or 4)
        if not schedule_generated:
            if audit_status == 'FAIL_PRE_SOLVER_CONTRACT':
                # Preserve the engine's truthful contract diagnosis and its
                # resource findings.  Only repair legacy/generic wording that
                # incorrectly claimed that a final schedule existed.
                prior_code = str(outcome.get('outcome_code') or '')
                if not prior_code or prior_code.startswith('FINAL_SCHEDULE_GENERATED'):
                    outcome['outcome_code'] = 'INPUT_OR_RESOURCE_CONTRACT_GAP'
                    outcome['outcome_category'] = 'ACTION_REQUIRED'
                    outcome['headline'] = 'Input or resource contract prevents schedule generation'
                    outcome['plain_language_summary'] = (
                        'The submitted workbook does not provide a legally possible staffing '
                        'assignment for every required interval. No final schedule was generated; '
                        'correct the input/resource contract and rerun.'
                    )
                outcome['technical_status'] = 'FAIL_PRE_SOLVER_CONTRACT'
            elif (validation_status == 'NOT_RUN'
                  and (audit_status in ENGINE_DIAGNOSED_NO_ARTIFACT or audit_status.startswith('FAIL'))
                  and not str(outcome.get('outcome_code') or '').startswith('FINAL_SCHEDULE_GENERATED')):
                # The engine said why there is no schedule (which hard rules
                # contradict, which break placement failed, and the requests
                # involved). Replacing that with "the engine output problem"
                # hid the reason from the person who has to fix the workbook.
                # Audit F-14: this used to cover only the two hard-contract
                # statuses, so every other no-schedule ending lost its reason.
                outcome['technical_status'] = audit_status
            else:
                outcome['outcome_code'] = 'NO_FINAL_SCHEDULE_GENERATED_VALIDATION_NOT_RUN'
                outcome['outcome_category'] = 'ACTION_REQUIRED'
                outcome['headline'] = 'No final schedule was generated; validation was not run'
                outcome['plain_language_summary'] = (
                    'No final schedule workbook was produced'
                    + (f' (engine status {audit_status})' if audit_status else '')
                    + ', so independent validation could not run. See the engine audit and log '
                    'for the cause, correct it and rerun.'
                )
                outcome['technical_status'] = (
                    'FAIL_INDEPENDENT_VALIDATION'
                    if validation_status != 'NOT_RUN' else 'FAIL_NO_FINAL_SCHEDULE'
                )
            outcome['recommended_actions'] = outcome.get('recommended_actions') or [
                'Review the contract or engine-output findings, correct the cause, and rerun.'
            ]
            # The generated-schedule branches below must not overwrite this
            # no-artifact diagnosis.
        if schedule_generated:
            outcome['technical_status'] = (
                'FAIL_INDEPENDENT_VALIDATION'
                if validation_status != 'PASS'
                else 'FAIL_RELEASE_GATE'
            )
        if schedule_generated and validation_status == 'PASS' and quality_gate_status == 'FAIL':
            outcome['outcome_code'] = 'FINAL_SCHEDULE_GENERATED_QUALITY_GATE_BLOCKED'
            outcome['outcome_category'] = 'QUALITY_BLOCKED'
            outcome['headline'] = 'Final schedule generated, but the production quality gate blocks release'
            outcome['plain_language_summary'] = (
                'The exact final workbook passed independent hard-rule validation, but the '
                'configured production quality gate still failed. The schedule is retained '
                'for review and cannot be treated as production-approved without remediation '
                'or explicit human approval.'
            )
        elif schedule_generated and validation_status == 'FAIL_METRIC_PARITY':
            outcome['outcome_code'] = 'FINAL_SCHEDULE_GENERATED_METRIC_PARITY_BLOCKED'
            outcome['outcome_category'] = 'QUALITY_BLOCKED'
            outcome['headline'] = 'Final schedule generated, but evaluator metric parity blocks release'
            outcome['plain_language_summary'] = (
                'The exact workbook passed the independent rule checks, but the engine and '
                'independent evaluator did not agree on one or more canonical metrics. '
                'The schedule is retained as evidence and blocked until the discrepancy is resolved.'
            )
        elif schedule_generated and validation_status == 'FAIL':
            outcome['outcome_code'] = 'FINAL_SCHEDULE_GENERATED_INDEPENDENT_VALIDATION_BLOCKED'
            outcome['outcome_category'] = 'QUALITY_BLOCKED'
            outcome['headline'] = 'Final schedule generated, but independent validation blocks release'
            outcome['plain_language_summary'] = (
                'The solver produced a schedule, but the exact polished workbook did not pass '
                'independent validation. It must not be used as a production schedule until '
                'the validation failure is resolved.'
            )
        elif schedule_generated and validation_status.startswith('ERROR'):
            outcome['outcome_code'] = 'FINAL_SCHEDULE_GENERATED_VALIDATION_ERROR_BLOCKED'
            outcome['outcome_category'] = 'QUALITY_BLOCKED'
            outcome['headline'] = 'Final schedule generated, but validation did not complete'
            outcome['plain_language_summary'] = (
                'The solver produced a schedule, but independent validation did not complete '
                'successfully. The schedule is blocked from production use.'
            )
        elif schedule_generated and validation_status.startswith('SKIPPED'):
            outcome['outcome_code'] = 'FINAL_SCHEDULE_GENERATED_VALIDATION_SKIPPED_BLOCKED'
            outcome['outcome_category'] = 'QUALITY_BLOCKED'
            outcome['headline'] = 'Final schedule generated, but validation was skipped'
            outcome['plain_language_summary'] = (
                'Independent validation was explicitly skipped. The schedule is blocked from '
                'production use.'
            )
        elif schedule_generated:
            outcome['outcome_code'] = 'FINAL_SCHEDULE_GENERATED_RELEASE_GATE_BLOCKED'
            outcome['outcome_category'] = 'QUALITY_BLOCKED'
            outcome['headline'] = 'Final schedule generated, but release gates block publication'
            outcome['plain_language_summary'] = (
                'The schedule was generated, but a final release gate did not pass. It is '
                'blocked from production use.'
            )
        if schedule_generated:
            outcome['recommended_actions'] = [
                'Review the independent validation and release-gate failures, correct the cause, and rerun validation.'
            ]
            outcome['blocking_reasons'] = _blocking_reasons(
                case_root, independent_validation, runner_return_code)

    rendered_text = (
        f"Outcome: {outcome.get('headline', '')}\n"
        f"Status: {outcome.get('technical_status', '')}\n"
        f"Production eligible: {bool(outcome.get('production_eligible'))}\n"
        f"Independent validation: {validation_status} (return code {validation_rc})\n"
        + coverage_measure_line(outcome) + "\n"
        f"{outcome.get('plain_language_summary', '')}\n"
    ) + _outcome_detail_text(outcome, 'Warnings:' if engine_said_releasable else None)
    for path in (outcome_path, debug_outcome_path):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(outcome, indent=2), encoding='utf-8')
        path.with_suffix('.txt').write_text(rendered_text, encoding='utf-8')

    for audit_path in sorted(case_root.glob('*solver_audit.json')):
        try:
            audit = json.loads(audit_path.read_text(encoding='utf-8'))
        except (OSError, json.JSONDecodeError):
            continue
        audit['business_outcome'] = outcome
        audit['independent_validation'] = validation_record
        if blocked:
            audit['production_eligible'] = False
        audit_path.write_text(json.dumps(audit, indent=2), encoding='utf-8')

    for summary_path in sorted(case_root.glob('*.l6_3_2_3_summary.csv')):
        try:
            with summary_path.open(newline='', encoding='utf-8') as handle:
                rows = list(csv.DictReader(handle))
            fieldnames = list(rows[0].keys()) if rows else []
            for name in (
                'business_outcome_code', 'business_outcome_category',
                'independent_validation_status', 'independent_validation_return_code',
                'independent_validation_json',
            ):
                if name not in fieldnames:
                    fieldnames.append(name)
            for row in rows:
                row['business_headline'] = str(outcome.get('headline') or '')
                row['business_message'] = str(outcome.get('plain_language_summary') or '')
                row['business_outcome_category'] = str(outcome.get('outcome_category') or '')
                row['business_outcome_code'] = str(outcome.get('outcome_code') or '')
                row['production_eligible'] = str(bool(outcome.get('production_eligible'))).upper()
                row['independent_validation_status'] = validation_status
                row['independent_validation_return_code'] = '' if validation_rc is None else str(validation_rc)
                row['independent_validation_json'] = str(independent_validation.get('json') or '')
            with summary_path.open('w', newline='', encoding='utf-8') as handle:
                writer = csv.DictWriter(handle, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)
        except (OSError, ValueError, csv.Error):
            continue


def run_clean_room_gate(input_path: Path, output_workbook: Path, audit_path: Path,
                        json_out: Path, validation_json: "Path | None" = None) -> dict:
    """Second, engine-independent check of the schedule about to be released.

    Audit F-19: the independent validator takes its contract from the engine's
    own parser and rule helpers, so a defect in one of those is invisible to
    both. The clean-room checker re-reads both workbooks with openpyxl alone.

      PASS        no rule violation and every compared metric equals the engine's
      FAIL        a violation, or a metric the engine published differently
                  (blocks release)
      NOT_CHECKED the checker could not read this workbook layout; recorded and
                  surfaced, not blocking, because the validator has passed it
    """
    result = {'status': 'NOT_CHECKED', 'json': str(json_out), 'checker': str(CLEAN_ROOM)}
    if not CLEAN_ROOM.exists():
        result['reason'] = 'clean-room checker not found'
        return result
    if json_out.exists():
        json_out.unlink()
    command = [sys.executable, '-u', str(CLEAN_ROOM), '--input', str(input_path),
               '--output', str(output_workbook), '--audit', str(audit_path), '--json-out', str(json_out)]
    if validation_json is not None:
        command += ['--validation', str(validation_json)]
    proc = subprocess.run(command, capture_output=True, text=True)
    result['return_code'] = proc.returncode
    try:
        report = json.loads(json_out.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        tail = (proc.stderr or proc.stdout or '').strip().splitlines()[-3:]
        result['reason'] = 'the checker did not finish reading this workbook: ' + ' | '.join(tail)
        return result
    result['violation_count'] = int(report.get('violation_count', 0) or 0)
    result['violations_by_rule'] = report.get('violations_by_rule') or {}
    result['engine_mismatches'] = report.get('engine_mismatches') or []
    result['validator_mismatches'] = report.get('validator_mismatches') or []
    result['input_issues'] = report.get('input_issues') or []
    blocking = result['violation_count'] > 0 or bool(result['engine_mismatches'])
    result['status'] = 'FAIL' if blocking else 'PASS'
    return result


def quality_allows_validation(case_root: Path, quality_rc: int) -> bool:
    """Allow exact-artifact validation before final quality approval.

    Phase C is first evaluated before the polisher and independent validator
    run.  That preliminary state must not block the validation it is waiting
    for. A quality-gate failure also must not stop schedule generation when a
    hard-valid final artifact exists: the artifact is validated and retained,
    while publication remains blocked. Real contract failures and hard-rule
    failures remain blocking.
    """
    if quality_rc == 0:
        return True
    report_path = case_root / 'PHASE_C_QUALITY_SUMMARY.json'
    try:
        report = json.loads(report_path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        return False
    contract = report.get('contract') or {}
    safety = report.get('safety') or {}
    gate = report.get('production_quality_gate') or {}
    safety_status = str(safety.get('status', '')).upper()
    hard_fail_count = int(safety.get('hard_fail_count') or 0)
    # Each stage has its own completed artifact. Requiring FINAL_VERIFIED of a
    # before-break run - which can never reach it - meant the preliminary
    # Phase C failure was never superseded, so the run kept the blocking code
    # even after its validation and parity gate both passed. The production
    # quality gate is a Stage-2 measurement, so NOT_EVALUATED is its honest
    # state here and only here.
    if str(report.get('run_stage') or '') == 'BEFORE_BREAKS_ONLY':
        required_artifact_state = 'BEST_BEFORE_BREAKS_ONLY'
        allowed_gate_states = {'PASS', 'WARN', 'FAIL', 'NOT_EVALUATED'}
    else:
        required_artifact_state = 'FINAL_VERIFIED'
        allowed_gate_states = {'PASS', 'WARN', 'FAIL'}
    return (
        report.get('artifact_state') == required_artifact_state
        and str(contract.get('status', '')).upper() in {'PASS', 'WARN'}
        and int(contract.get('failure_count') or 0) == 0
        and safety_status in {'NOT_VALIDATED', 'PASS', 'WARN'}
        and hard_fail_count == 0
        and str(gate.get('status', '')).upper() in allowed_gate_states
        and not (case_root / 'INDEPENDENT_VALIDATION.json').exists()
    )


def stage_input_snapshot(source: Path, case_root: Path, resume: bool) -> Path:
    """Create the immutable case-local contract used by every run phase."""
    target_dir = case_root / 'input_snapshot'
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / source.name
    source_hash = sha256_file(source)
    if target.exists():
        target_hash = sha256_file(target)
        if target_hash != source_hash:
            raise RuntimeError(
                f'Input snapshot conflict: {target} has SHA256 {target_hash}, '
                f'but the requested input has {source_hash}. Use a new schedule-id.'
            )
        return target
    shutil.copy2(source, target)
    if sha256_file(target) != source_hash:
        target.unlink(missing_ok=True)
        raise RuntimeError('Input snapshot verification failed after copy')
    return target


def finalize_production_manifest(case_root: Path, validation_json: Path, validation_workbook: Path) -> None:
    """Seal hard-valid artifacts while keeping quality/human approval explicit."""
    manifest_path = case_root / 'production' / 'PRODUCTION_ARTIFACT_MANIFEST.json'
    if not manifest_path.is_file():
        raise RuntimeError('Prepared production manifest is missing')
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    validation = json.loads(validation_json.read_text(encoding='utf-8'))
    if validation.get('status') != 'PASS' or int(validation.get('hard_fail_count', 1) or 0) != 0:
        raise RuntimeError('Independent validation did not pass cleanly')
    actual_output_hash = sha256_file(validation_workbook)
    if validation.get('output_sha256') != actual_output_hash:
        raise RuntimeError('Independent validation output hash does not match the polished workbook')
    final_record = (manifest.get('two_artifact_contract') or {}).get('BEST_FINAL_AFTER_BREAKS_SCHEDULE') or {}
    if final_record.get('sha256') != actual_output_hash:
        raise RuntimeError('Prepared manifest final-workbook hash does not match the validated workbook')
    quality_path = case_root / 'PHASE_C_QUALITY_SUMMARY.json'
    if not quality_path.is_file():
        raise RuntimeError('Final Phase C quality summary is missing')
    try:
        quality_summary = json.loads(quality_path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError('Final Phase C quality summary is unreadable') from exc
    quality_safety = (quality_summary.get('safety') or {}).get('status')
    if str(quality_safety or '').upper() != 'PASS':
        raise RuntimeError(
            f'Final Phase C quality summary safety status is not PASS: {quality_safety!r}'
        )
    manifest['independent_validation'] = {
        'status': 'PASS',
        'hard_fail_count': 0,
        'json': str(validation_json.relative_to(case_root)),
        'json_sha256': sha256_file(validation_json),
        'validated_workbook': str(validation_workbook.relative_to(case_root)),
        'validated_workbook_sha256': actual_output_hash,
        'input_sha256': validation.get('input_sha256'),
        'engine_sha256': validation.get('engine_sha256'),
        'metric_parity': validation.get('metric_parity'),
    }
    # Automated checks prove that the exact workbook is technically valid and
    # its hashes are sealed. They do not constitute human operational
    # approval, so keep the delivery available for review without overstating
    # it as production-ready.
    manifest['approval_status'] = 'AUTOMATED_HARD_GATES_PASSED_PENDING_HUMAN_APPROVAL'
    manifest['hard_gates_passed'] = True
    manifest['automated_hard_gates_passed'] = True
    manifest['production_ready'] = False
    manifest['quality_gate_status'] = (quality_summary.get('production_quality_gate') or {}).get('status', 'NOT_EVALUATED')
    manifest['release_disposition'] = (
        'REVIEW_ONLY_QUALITY_GATE_BLOCKED'
        if manifest['quality_gate_status'] == 'FAIL'
        else 'REVIEW_ONLY_HUMAN_APPROVAL_PENDING'
    )
    manifest['sealed_utc'] = datetime.now(timezone.utc).isoformat()
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding='utf-8')


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description='RC9 universal workbook-driven WFM production runner. No client/case branching.')
    p.add_argument('--input', type=Path, required=False, help='Any supported Universal WFM input workbook (.xlsx)')
    p.add_argument('--output-root', type=Path, default=ROOT / 'results')
    p.add_argument('--schedule-id', help='Optional display/job ID. Defaults to input filename stem.')
    p.add_argument('--mode', choices=['SMOKE','QUICK','DEEP','OVERNIGHT','FULL'], default=None,
                   help='Run depth. Omit to take Run Depth from the workbook, '
                        'falling back to QUICK. FULL is a legacy alias for OVERNIGHT.')
    p.add_argument('--stage', choices=['BEFORE_BREAKS_ONLY','FULL_SCHEDULE'], default=None,
                   help='Omit to take Run Stage from the workbook, falling back '
                        'to FULL_SCHEDULE. BEFORE_BREAKS_ONLY runs Stage 1 and '
                        'exports the before-break champion without placing breaks.')
    p.add_argument('--time-limit', type=int, help='Override total solver budget seconds')
    p.add_argument('--num-workers', type=int, default=max(1, min(8, os.cpu_count() or 4)))
    p.add_argument('--pattern-widths', default='24,44,60,115')
    p.add_argument('--repair-change-limits', default='2,4,6,8,12')
    p.add_argument('--skeleton-profiles', default=DEFAULT_SKELETON_PROFILES)
    p.add_argument('--break-objective-modes', default=DEFAULT_BREAK_OBJECTIVES)
    p.add_argument('--use-input-schedule-as-seed', action='store_true', default=True)
    p.add_argument('--disable-input-schedule-seed', action='store_true')
    p.add_argument('--coverage-objective-weighting', choices=['INTERVAL_COUNT', 'VOLUME_WEIGHTED'], default=None,
                   help="Override the workbook's Coverage Objective Weighting for this run. "
                        "INTERVAL_COUNT counts every interval at target once (interval compliance); "
                        "VOLUME_WEIGHTED weighs each interval by its requirement (service level).")
    p.add_argument('--allow-no-break-exceptions', action='store_true')
    p.add_argument('--disable-no-break-exceptions', action='store_true')
    p.add_argument('--max-no-break-exceptions', type=int)
    p.add_argument('--language-working-window',
                   choices=['OFF', 'MINIMUM_ROWS', 'ALL_ROWS', 'REQUIRED_LANGUAGE_ONLY'], default=None,
                   help="Override the workbook's Language Working Window setting. OFF keeps "
                        "Coverage Start/End as a coverage minimum only (default). MINIMUM_ROWS "
                        "also bounds working hours for language rows with a minimum >= 1. "
                        "ALL_ROWS bounds every active row. REQUIRED_LANGUAGE_ONLY prevents "
                        "associates who cannot cover an active required language from being "
                        "scheduled during that language's window.")
    p.add_argument('--allow-headcount-mismatch', action='store_true')
    p.add_argument('--acknowledge-departed', default='',
                   help='semicolon-separated names treated as departed (see Known Departed Associates)')
    p.add_argument('--diagnostics-only', action='store_true')
    p.add_argument('--skeleton-only', action='store_true', help='Run Stage 1 only and export ranked before-break skeletons.')
    p.add_argument('--export-top-skeletons', type=int, default=5)
    p.add_argument('--enable-bundled-regression-fallbacks', action='store_true', help='Off by default in production; regression assets must not steer new client workbooks unless explicitly enabled.')
    p.add_argument('--solver-random-seed', type=int, default=9000)
    p.add_argument('--overwrite', action='store_true', help='Explicitly replace artifacts for the same schedule-id. Off by default.')
    p.add_argument('--resume', action='store_true')
    p.add_argument('--skip-independent-validation', action='store_true', help='Development-only bypass. It always blocks production packaging and returns a non-production status.')
    p.add_argument('--enable-joint-refinement', action='store_true',
                   help='Force joint refinement on. Off by default at every depth: 474 '
                        'solver audits (QUICK and DEEP) with improved 0, ~90%% of peak memory, '
                        'and two DEEP runs killed at 13-14 GB inside it.')
    polish = p.add_mutually_exclusive_group()
    polish.add_argument('--shift-consistency-polish', dest='shift_consistency_polish', action='store_const',
                        const=True, default=None,
                        help='Re-deal same-day shifts between interchangeable associates after the schedule '
                             'is chosen, so each week is more uniform. No metric may get worse.')
    polish.add_argument('--no-shift-consistency-polish', dest='shift_consistency_polish', action='store_const',
                        const=False, help='Turn the shift consistency polish off.')
    p.add_argument('--enable-dnbs', action='store_true',
                   help='enable the day-neighbourhood break search (off by default since the 2026-09-28 audit, F-10)')
    p.add_argument('--enable-break-load-feedback', action='store_true',
                   help='Stage-2 -> Stage-1 break-load feedback in place of coordinated repair '
                        '(off by default until its registered A/B).')
    p.add_argument('--stage1-profile-rotation', default=None,
                   help='i/n: Stage-1 profile order for seed i of an n-seed portfolio '
                        '(set by RUN_PORTFOLIO.py; default unchanged).')
    p.add_argument('--enable-final-recovery-endgame', action='store_true',
                   help='Reopen the joint search when a run ends with no release candidate '
                        '(off by default: 9 runs, 0 candidates added).')
    p.add_argument('--selfcheck', action='store_true')
    return p


def selfcheck() -> int:
    errors = []
    for path in [ENGINE, POLISHER, QUALITY_REPORTER, PACKAGER, VALIDATOR]:
        if not path.exists():
            errors.append(f'MISSING {path}')
    for py in [ENGINE, POLISHER, QUALITY_REPORTER, PACKAGER, VALIDATOR, Path(__file__)]:
        if py.exists():
            rc = subprocess.call([sys.executable, '-m', 'py_compile', str(py)])
            if rc:
                errors.append(f'COMPILE_FAIL {py}')
    if ENGINE.exists():
        rc = subprocess.call([sys.executable, str(ENGINE), '--selfcheck'])
        if rc:
            errors.append('ENGINE_SELFCHECK_FAIL')
    if POLISHER.exists():
        rc = subprocess.call([sys.executable, str(POLISHER), '--selfcheck'])
        if rc:
            errors.append('POLISHER_SELFCHECK_FAIL')
    if errors:
        print(json.dumps({'status':'FAIL','release':RELEASE,'errors':errors}, indent=2))
        return 1
    print(json.dumps({'status':'PASS','release':RELEASE,'engine_sha256': sha256_file(ENGINE)}, indent=2))
    return 0


REQUIRED_RUNTIME_MODULES = ('ortools', 'openpyxl', 'numpy', 'scipy')


def missing_runtime_dependencies() -> list:
    """Modules the engine needs to behave as measured; empty when all import."""
    import importlib
    missing = []
    for name in REQUIRED_RUNTIME_MODULES:
        try:
            importlib.import_module(name)
        except Exception as exc:  # noqa: BLE001 - any import failure is a missing dependency
            missing.append(f'{name} ({type(exc).__name__})')
    return missing


def main() -> int:
    runner_parser = build_parser()
    args = runner_parser.parse_args()
    if args.selfcheck:
        return selfcheck()
    if args.input is None:
        raise SystemExit('--input is required unless --selfcheck is used')
    missing = missing_runtime_dependencies()
    if missing:
        # The engine's Stage-1 aggregate guide (scipy MILP) silently degrades to
        # UNAVAILABLE without scipy, and every schedule is then searched without
        # it. That is a different engine from the one production runs, so refuse
        # rather than produce a schedule nobody measured (audit F-20).
        print('[run] RUNTIME DEPENDENCY MISSING: ' + ', '.join(missing) + '. Install the pinned '
              'requirements (ortools==9.15.6755, openpyxl, pandas, numpy, scipy>=1.11).', flush=True)
        return 2
    original_input_path = args.input.resolve()
    if not original_input_path.exists():
        raise FileNotFoundError(original_input_path)
    if original_input_path.suffix.lower() != '.xlsx':
        raise ValueError('Input must be an .xlsx workbook')

    schedule_id = safe_id(args.schedule_id or original_input_path.stem)
    case_root = (args.output_root / schedule_id).resolve()
    case_root.mkdir(parents=True, exist_ok=True)
    acquire_case_lock(case_root, schedule_id)
    input_path = stage_input_snapshot(original_input_path, case_root, bool(args.resume))
    work_dir = case_root / 'debug'
    output = case_root / f'{schedule_id}_L6_3_2_3_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx'
    audit = case_root / f'{schedule_id}.l6_3_2_3_solver_audit.json'
    summary = case_root / f'{schedule_id}.l6_3_2_3_summary.csv'

    # Run stage and depth come from the workbook unless the command line says
    # otherwise, so a scheduler picks them from a dropdown in the business
    # contract instead of remembering flags. An explicit flag always wins.
    contract_stage, contract_depth, workbook_search_controls = _contract_run_settings(input_path)
    mode = args.mode or contract_depth or 'QUICK'
    if mode == 'FULL':
        mode = 'OVERNIGHT'
    stage = args.stage or contract_stage or 'FULL_SCHEDULE'

    all_mode_defaults = {
        # joint_enabled False at every depth. SMOKE/QUICK: 69 attempts across 7
        # workbooks produced improved: 0, while the phase accounted for ~90% of
        # peak RSS (8197 MB -> 820 MB when disabled, coverage unchanged).
        # DEEP/OVERNIGHT, measured since: 36 attempts in 4 DEEP runs, improved 0,
        # and two DEEP runs killed at 13.0 / 13.9 GB inside the phase; 474
        # solver audits in all, improved 0 in every one
        # (evidence/JOINT_REFINEMENT_REMOVED.md). Its time goes to the other
        # phases. --enable-joint-refinement turns it back on.
        'SMOKE': {'time_limit': 900, 'joint': 120, 'joint_enabled': False, 'safe': 120, 'post': 60, 'target': 60, 'final': 60, 'adaptive': 6, 'joint_attempts': 4, 'joint_no_improve': 2},
        'QUICK': {'time_limit': 3600, 'joint': 900, 'joint_enabled': False, 'safe': 180, 'post': 180, 'target': 180, 'final': 120, 'adaptive': 18, 'joint_attempts': 16, 'joint_no_improve': 6},
        # RC9.1 production defaults: Deep may use four hours and Overnight six.
        # The engine remains workbook-driven; longer time only expands universal search breadth/depth.
        'DEEP': {'time_limit': 14400, 'joint': 5400, 'joint_enabled': False, 'safe': 360, 'post': 600, 'target': 600, 'final': 240, 'adaptive': 48, 'joint_attempts': 48, 'joint_no_improve': 16},
        'OVERNIGHT': {'time_limit': 21600, 'joint': 8400, 'joint_enabled': False, 'safe': 480, 'post': 900, 'target': 900, 'final': 300, 'adaptive': 72, 'joint_attempts': 72, 'joint_no_improve': 22},
    }
    mode_defaults = all_mode_defaults[mode]
    time_limit = int(args.time_limit or mode_defaults['time_limit'])
    if time_limit < 60:
        print(f'[run] --time-limit {time_limit} is below the 60 s minimum the engine can plan (audit M-01)', flush=True)
        return 2
    diagnostics_only = bool(args.diagnostics_only or mode == 'SMOKE')
    # One source of truth. `--skeleton-only` already existed as a flag with full
    # downstream handling; the workbook's Run Stage is a second way to ask for
    # the same thing, so both feed one variable. Deriving the stage from the
    # flag as well keeps the two from disagreeing in the run record.
    skeleton_only = bool(args.skeleton_only) or stage == 'BEFORE_BREAKS_ONLY'
    if skeleton_only:
        stage = 'BEFORE_BREAKS_ONLY'
    # The engine renames its own output in skeleton-only mode; renaming it here
    # too produced a doubled ..._BEST_BEFORE_BREAKS_SCHEDULE_BEST_BEFORE_BREAKS
    # _SCHEDULE.xlsx, so the path stays as-is and the engine decides the name.
    print(f'[run] stage={stage} depth={mode} time_limit={time_limit}s '
          f'(stage source: {"cli" if args.stage else ("workbook" if contract_stage else "default")}; '
          f'depth source: {"cli" if args.mode else ("workbook" if contract_depth else "default")})',
          flush=True)

    command = [
        sys.executable, '-u', str(ENGINE),
        '--input', str(input_path),
        '--output', str(output),
        '--audit-json', str(audit),
        '--summary-csv', str(summary),
        '--work-dir', str(work_dir),
        '--time-limit', str(time_limit),
        '--max-final-before-target-loss', '6',
    ] + engine_flags_for_run(
        args, mode_defaults, runner_parser, workbook_search_controls)
    if workbook_search_controls:
        print(f'[run] workbook states {len(workbook_search_controls)} search control(s): '
              f'{", ".join(sorted(workbook_search_controls))}', flush=True)
    if not args.enable_bundled_regression_fallbacks:
        command.append('--disable-bundled-fallbacks')
    if not mode_defaults.get('joint_enabled', True) and not args.enable_joint_refinement:
        command.append('--disable-joint-refinement')
    if args.enable_final_recovery_endgame:
        command.append('--enable-final-recovery-endgame')
    if args.enable_break_load_feedback:
        command.append('--enable-break-load-feedback')
    if args.enable_dnbs:
        command.append('--enable-dnbs')
    if args.shift_consistency_polish is True:
        command.append('--shift-consistency-polish')
    elif args.shift_consistency_polish is False:
        command.append('--no-shift-consistency-polish')
    if args.stage1_profile_rotation:
        command += ['--stage1-profile-rotation', str(args.stage1_profile_rotation)]
    if args.use_input_schedule_as_seed and not args.disable_input_schedule_seed:
        command.append('--use-input-schedule-as-seed')
    if args.allow_no_break_exceptions:
        command.append('--allow-no-break-exceptions')
    if args.disable_no_break_exceptions:
        command.append('--disable-no-break-exceptions')
    if args.max_no_break_exceptions is not None:
        command += ['--max-no-break-exceptions', str(args.max_no_break_exceptions)]
    if args.language_working_window is not None:
        command += ['--language-working-window', args.language_working_window]
    if args.coverage_objective_weighting is not None:
        command += ['--coverage-objective-weighting', args.coverage_objective_weighting]
    if args.allow_headcount_mismatch:
        command.append('--allow-headcount-mismatch')
    if args.acknowledge_departed:
        command += ['--acknowledge-departed', args.acknowledge_departed]
    if diagnostics_only:
        command.append('--diagnostics-only')
    if skeleton_only:
        command += ['--skeleton-only', '--export-top-skeletons', str(max(0, args.export_top_skeletons))]
    if args.overwrite:
        command.append('--overwrite')
    if args.resume:
        command.append('--resume')

    identity = {
        'schema_version': 1,
        'release': RELEASE,
        'created_utc': datetime.now(timezone.utc).isoformat(),
        'schedule_id': schedule_id,
        'mode': mode, 'stage': stage,
        'input_path': str(input_path),
        'original_input_path': str(original_input_path),
        'input_snapshot': True,
        'input_sha256': sha256_file(input_path),
        'engine_sha256': sha256_file(ENGINE),
        'python_version': platform.python_version(),
        'case_branching': False,
        'rc9_1_search_recovery': True,
        'rc9_2_protected_balance': True,
        'skeleton_only': bool(skeleton_only),
        'stage2_strategy': 'breadth_first_expand_then_next_skeleton',
        'bundled_regression_fallbacks_enabled': bool(args.enable_bundled_regression_fallbacks),
        'command': command,
    }
    (case_root / 'UNIVERSAL_RUN_IDENTITY.json').write_text(json.dumps(identity, indent=2), encoding='utf-8')
    print('COMMAND:', ' '.join(command), flush=True)
    engine_rc = subprocess.call(command)
    rc = engine_rc

    # Complete the case-root identity from the engine's own run identity.
    #
    # The wrapper can only know the input and engine hashes before the run; the
    # contract hash, parameters hash and run id are derived by the engine while
    # parsing. They were therefore written only to work_dir/RUN_IDENTITY.json,
    # leaving UNIVERSAL_RUN_IDENTITY.json - the file at case root that a
    # reviewer or packager reads - carrying two of the four identity axes, with
    # contract_sha256 and run_id absent. The release gate requires identity by
    # input, contract and engine hash plus run metadata, so an artifact that
    # cannot state its own contract hash cannot satisfy it.
    engine_identity_path = work_dir / 'RUN_IDENTITY.json'
    if engine_identity_path.exists():
        try:
            engine_identity = json.loads(engine_identity_path.read_text(encoding='utf-8'))
        except (OSError, json.JSONDecodeError) as exc:
            identity['engine_identity_error'] = f'{type(exc).__name__}: {exc}'
        else:
            for key in ('contract_sha256', 'run_id', 'parameters_sha256', 'seed_sha256',
                        'canonical_coverage_evaluator', 'git_commit', 'git_describe', 'git_dirty'):
                if engine_identity.get(key) is not None:
                    identity[key] = engine_identity[key]
            # Cross-check rather than trust: if the engine resolved a different
            # input or engine hash than the wrapper stamped, the artifact is
            # self-inconsistent and that must be visible, not silently merged.
            for key in ('input_sha256', 'engine_sha256'):
                engine_value = engine_identity.get(key)
                if engine_value is not None and engine_value != identity.get(key):
                    identity.setdefault('identity_mismatches', {})[key] = {
                        'wrapper': identity.get(key), 'engine': engine_value,
                    }
        identity['engine_identity_source'] = str(engine_identity_path)
    else:
        identity['engine_identity_source'] = None
    (case_root / 'UNIVERSAL_RUN_IDENTITY.json').write_text(
        json.dumps(identity, indent=2), encoding='utf-8')

    qrc = 3
    initial_qrc = qrc
    quality_pending_validation = False
    if audit.exists():
        qrc = subprocess.call([sys.executable, '-u', str(QUALITY_REPORTER), '--audit-json', str(audit), '--case-root', str(case_root), '--strict'])
        initial_qrc = qrc
        quality_pending_validation = quality_allows_validation(case_root, qrc)
        if qrc != 0 and not quality_pending_validation and rc == 0:
            rc = qrc
    independent_validation = {
        'enabled': not bool(args.skip_independent_validation),
        'status': 'NOT_RUN',
        'return_code': None,
        'workbook': None,
        'json': None,
        'csv': None,
        'quality_gate_status': None,
        'coverage_quality_gate_status': None,
        'canonical_metrics': None,
        'metric_parity': None,
    }
    validation_workbook = None
    if not diagnostics_only and skeleton_only and engine_rc == 0:
        before_candidates = sorted(case_root.glob('*_BEST_BEFORE_BREAKS_SCHEDULE.xlsx'))
        validation_workbook = before_candidates[0] if before_candidates else None
    elif (not diagnostics_only and not skeleton_only
          and (engine_rc == 0 or quality_pending_validation)
          and (qrc == 0 or quality_pending_validation) and hard_valid_artifact(audit)):
        prc = subprocess.call([sys.executable, '-u', str(POLISHER), '--case-root', str(case_root), '--prepare-only'])
        if prc != 0:
            rc = prc
        else:
            polished = sorted((case_root / 'production').glob('*_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx'))
            validation_workbook = polished[0] if len(polished) == 1 else None

    if validation_workbook is not None and not args.skip_independent_validation:
        validation_json = case_root / 'INDEPENDENT_VALIDATION.json'
        validation_csv = case_root / 'INDEPENDENT_VALIDATION.csv'
        validator_command = [
            sys.executable, '-u', str(VALIDATOR),
            '--input', str(input_path), '--output', str(validation_workbook),
            '--json-out', str(validation_json), '--csv-out', str(validation_csv),
        ]
        if args.language_working_window is not None:
            # The engine enforced the run's override, so the validator must check it.
            validator_command += ['--language-working-window', args.language_working_window]
        vrc = subprocess.call(validator_command)
        # The validator exits 0 on PASS and 2 when it has evaluated the schedule
        # and found hard-rule violations.  Any other code means it did not
        # complete - a crash, a missing file, an unreadable sheet.  Collapsing
        # every non-zero code to 'FAIL' reported a broken tool as a broken
        # schedule, which is a materially different claim: one blocks release
        # for a real defect, the other for a bug in the checker.  Both still
        # block (an unvalidated schedule must never pass), but they are now
        # distinguishable.
        if vrc == 0:
            validation_status = 'PASS'
        elif vrc == 2:
            validation_status = 'FAIL'
        else:
            validation_status = 'ERROR_VALIDATOR_DID_NOT_COMPLETE'
        # Every non-zero validator result blocks release, including checker
        # errors that are not schedule defects.
        if vrc != 0:
            rc = 4
        validation = {}
        validation_read_error = None
        try:
            validation = json.loads(validation_json.read_text(encoding='utf-8'))
        except (OSError, json.JSONDecodeError) as exc:
            # A validator process can return nonzero before it writes its JSON
            # (for example, an import or filesystem failure).  Treat that as
            # an incomplete validation with explicit evidence rather than
            # dereferencing an undefined object or mislabelling it as a
            # schedule-rule failure.
            validation_read_error = f'{type(exc).__name__}: {exc}'
            vrc = 3
            rc = 4
            validation = {
                'status': 'ERROR_VALIDATOR_DID_NOT_COMPLETE',
                'return_code': vrc,
                'error': validation_read_error,
            }
            validation_status = 'ERROR_VALIDATOR_DID_NOT_COMPLETE'
            try:
                validation_json.write_text(json.dumps(validation, indent=2), encoding='utf-8')
            except OSError:
                pass
        independent_validation.update({
            'status': validation_status,
            'return_code': int(vrc),
            'workbook': str(validation_workbook),
            'json': str(validation_json),
            'csv': str(validation_csv),
            'quality_gate_status': validation.get('quality_gate_status'),
            'coverage_quality_gate_status': validation.get('coverage_quality_gate_status'),
            'error': validation_read_error,
        })
        independent_validation['canonical_metrics'] = validation.get('canonical_metrics')
        if vrc == 0:
            parity_rc, parity_validation = apply_metric_parity_gate(audit, validation_json, validation_csv)
            independent_validation['canonical_metrics'] = parity_validation.get('canonical_metrics')
            independent_validation['metric_parity'] = parity_validation.get('metric_parity')
            if parity_rc != 0:
                vrc = parity_rc
                validation_status = 'FAIL_METRIC_PARITY'
                independent_validation['status'] = validation_status
                independent_validation['return_code'] = int(vrc)
        if vrc == 0 and not skeleton_only:
            clean_room = run_clean_room_gate(input_path, validation_workbook, audit,
                                             case_root / 'CLEAN_ROOM_CHECK.json', validation_json)
            independent_validation['clean_room'] = clean_room
            print(f"[run] clean-room check: {clean_room['status']}"
                  + (f" ({clean_room.get('reason')})" if clean_room.get('reason') else ''), flush=True)
            if clean_room['status'] == 'FAIL':
                vrc = 2
                validation_status = 'FAIL_CLEAN_ROOM'
                independent_validation['status'] = validation_status
                independent_validation['return_code'] = int(vrc)
        if vrc != 0:
            rc = 4
        # Re-evaluate Phase C after the exact polished workbook has been
        # independently checked, whether that check passes or fails. This makes
        # the report reflect the final artifact instead of a pre-validation
        # placeholder and closes every quality-failure path.
        qrc = subprocess.call([
            sys.executable, '-u', str(QUALITY_REPORTER),
            '--audit-json', str(audit), '--case-root', str(case_root), '--strict',
        ])
        if qrc != 0 and rc == 0:
            rc = qrc
        if vrc == 0 and not skeleton_only:
            try:
                finalize_production_manifest(case_root, validation_json, validation_workbook)
            except Exception as exc:
                independent_validation['status'] = 'ERROR_RELEASE_SEAL_FAILED'
                independent_validation['seal_error'] = f'{type(exc).__name__}: {exc}'
                rc = 4
    elif validation_workbook is not None:
        independent_validation.update({'status': 'SKIPPED_BY_EXPLICIT_FLAG', 'workbook': str(validation_workbook)})
        if not skeleton_only:
            rc = 5
    elif not diagnostics_only and (engine_rc == 0 or quality_pending_validation):
        independent_validation['status'] = (
            'SKIPPED_QUALITY_GATE' if not quality_pending_validation and qrc != 0
            else 'FAIL_OUTPUT_NOT_FOUND'
        )
        # A successful engine exit without a validated workbook is never a
        # successful production run. Preserve an earlier nonzero gate code,
        # but fail explicitly when this is the first blocking condition.
        if rc == 0:
            rc = 5 if args.skip_independent_validation else 4

    shortfall_books = sorted(case_root.glob('*_HARD_RULE_SHORTFALL_SCHEDULE.xlsx'))
    if validation_workbook is None and shortfall_books and not args.skip_independent_validation:
        # Audit F-06: no releasable schedule, but the engine produced one that
        # meets every person rule and lists its missed minimums. Check it
        # independently; it never changes the (non-zero) return code.
        # Phase H: the same look as every other schedule, applied before the
        # validator reads it, so the workbook validated is the one shipped.
        shortfall_presentation = present_alternative(shortfall_books[0], 'HARD_RULE_SHORTFALL_SCHEDULE')
        check = validate_shortfall_schedule(input_path, shortfall_books[0], args.language_working_window,
                                            case_root / 'SHORTFALL_SCHEDULE_VALIDATION.json')
        check['clean_room'] = run_clean_room_gate(input_path, shortfall_books[0], audit,
                                                  case_root / 'SHORTFALL_CLEAN_ROOM_CHECK.json')
        check['presentation'] = shortfall_presentation
        independent_validation['shortfall_schedule'] = check
        print(f"[run] shortfall schedule (not releasable): {check.get('status')}", flush=True)
        if rc == 0:
            rc = 2
    if (validation_workbook is not None and not args.skip_independent_validation
            and not skeleton_only):
        alternatives = validate_alternative_exports(case_root, input_path, args.language_working_window)
        independent_validation['alternative_exports'] = alternatives
        for row in alternatives:
            print(f"[run] alternative export {row['role']}: {row['status']}", flush=True)
    reconcile_business_outcome_after_validation(case_root, independent_validation, rc)
    metrics = read_summary_metrics(case_root)
    run_status = {
        'schema_version': 1,
        'release': RELEASE,
        'schedule_id': schedule_id,
        'mode': mode, 'stage': stage,
        'return_code': rc,
        'engine_return_code': engine_rc,
        'initial_quality_report_return_code': initial_qrc,
        'quality_report_return_code': qrc,
        'hard_valid_artifact': hard_valid_artifact(audit),
        'metrics': metrics,
        'case_root': str(case_root),
        'independent_validation': independent_validation,
    }
    status_path = case_root / 'UNIVERSAL_RUN_STATUS.json'
    # Write the status before packaging so every package carries the exact gate
    # decision that authorized (or blocked) it.  Rewrite it if packaging itself
    # fails, preserving that final return code as well.
    status_path.write_text(json.dumps(run_status, indent=2, default=str), encoding='utf-8')
    if (
        not diagnostics_only and not skeleton_only
        and (engine_rc == 0 or quality_pending_validation)
        and (qrc == 0 or quality_pending_validation)
        and hard_valid_artifact(audit)
        and independent_validation.get('status') == 'PASS'
        and not args.skip_independent_validation
    ):
        pkg_rc = subprocess.call([sys.executable, '-u', str(PACKAGER), '--case-root', str(case_root)])
        if pkg_rc != 0:
            rc = pkg_rc
            run_status['return_code'] = rc
            run_status['packager_return_code'] = pkg_rc
            status_path.write_text(json.dumps(run_status, indent=2, default=str), encoding='utf-8')
    if not diagnostics_only and not skeleton_only:
        run_status['top_level_copies'] = publish_top_level_copies(case_root)
        status_path.write_text(json.dumps(run_status, indent=2, default=str), encoding='utf-8')
    print(json.dumps(run_status, indent=2, default=str), flush=True)
    return rc

TOP_LEVEL_SCHEDULES = ('*_BEST_BEFORE_BREAKS_SCHEDULE.xlsx', '*_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx')


def publish_top_level_copies(case_root: Path) -> list:
    """Phase H: the schedule files a user opens at the top of the run folder are
    the polished production files (byte-identical to what was validated and
    packaged), not the raw engine copies with every input tab showing. The raw
    copies move to debug/raw_engine_output/, where a resumed run's polisher reads
    them. A file without a production counterpart is left as it is."""
    replaced = []
    raw_dir = case_root / 'debug' / 'raw_engine_output'
    for pattern in TOP_LEVEL_SCHEDULES:
        for top in sorted(case_root.glob(pattern)):
            polished = case_root / 'production' / top.name
            if not polished.is_file():
                continue
            raw_dir.mkdir(parents=True, exist_ok=True)
            if not (raw_dir / top.name).exists():
                shutil.move(str(top), str(raw_dir / top.name))
            shutil.copy2(polished, top)
            replaced.append(top.name)
    return replaced


# Audit F-06. Each Shortfalls-sheet rule, and where the independent validator
# reports the same miss: (failure type, count field or None = one row each).
SHORTFALL_VALIDATOR_MAP = {
    'zero_coverage': ('ZERO_STAFF_ACTIVE', None),
    'language': ('LANGUAGE_MINIMUM', None),
    'opening': ('OPENING_MINIMUM', None),
    'coverage_split': ('COVERAGE_SPLIT', 'count'),
    'hard_floor': ('HARD_FLOOR', 'count'),
    'next_sunday_zero': ('NEXT_SUNDAY_CARRY_OUT', 'zero_count'),
    'next_sunday_language': ('NEXT_SUNDAY_CARRY_OUT', 'language_gap_count'),
    'next_sunday_opening': ('NEXT_SUNDAY_CARRY_OUT', 'opening_gap_count'),
}


def validate_shortfall_schedule(input_path: Path, workbook: Path, language_working_window: "str | None",
                                json_out: Path) -> dict:
    """Independent check of a HARD_RULE_SHORTFALL_SCHEDULE (audit F-06).

    SHORTFALLS_CONFIRMED means the validator's only failures are coverage
    minimums, and it counts exactly the misses the Shortfalls sheet lists.
    Anything else (a person-rule failure, a miss the sheet does not list, a
    listed miss the validator does not see) is INCONSISTENT. Either way the
    schedule is not releasable.
    """
    command = [sys.executable, '-u', str(VALIDATOR), '--input', str(input_path),
               '--output', str(workbook), '--json-out', str(json_out)]
    if language_working_window is not None:
        command += ['--language-working-window', language_working_window]
    proc = subprocess.run(command, capture_output=True, text=True)
    result = {'workbook': str(workbook), 'json': str(json_out), 'return_code': proc.returncode}
    try:
        report = json.loads(json_out.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        result['status'] = 'ERROR_VALIDATOR_DID_NOT_COMPLETE'
        return result
    failures = report.get('failures') or []
    allowed = {kind for kind, _field in SHORTFALL_VALIDATOR_MAP.values()}
    types = sorted({str(row.get('type')) for row in failures})
    result['validator_failure_types'] = types
    from openpyxl import load_workbook
    listed = {}
    wb = load_workbook(workbook, read_only=True)
    try:
        if 'Shortfalls' in wb.sheetnames:
            for row in wb['Shortfalls'].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    listed[str(row[0])] = listed.get(str(row[0]), 0) + 1
    finally:
        wb.close()
    result['listed'] = listed
    expected = {}
    for family, count in listed.items():
        if family in SHORTFALL_VALIDATOR_MAP:
            key = SHORTFALL_VALIDATOR_MAP[family]
            expected[key] = expected.get(key, 0) + count
    seen = {}
    for kind, field in set(SHORTFALL_VALIDATOR_MAP.values()):
        rows = [row for row in failures if row.get('type') == kind]
        seen[(kind, field)] = (len(rows) if field is None
                               else sum(int(row.get(field) or 0) for row in rows))
    mismatches = [
        {'validator_type': kind, 'field': field, 'validator': seen.get((kind, field), 0),
         'shortfalls_sheet': expected.get((kind, field), 0)}
        for kind, field in sorted(set(SHORTFALL_VALIDATOR_MAP.values()), key=str)
        if seen.get((kind, field), 0) != expected.get((kind, field), 0)
    ]
    outside = [t for t in types if t not in allowed]
    result['outside_declared_families'] = outside
    result['count_mismatches'] = mismatches
    result['status'] = 'SHORTFALLS_CONFIRMED' if not outside and not mismatches else 'INCONSISTENT'
    return result


ALTERNATIVE_EXPORT_ROLES = ('MAX_TARGET_CANDIDATE', 'MAX_FLOOR_CANDIDATE', 'BALANCED_CANDIDATE',
                            'SAFER_BALANCED_CANDIDATE', 'MORE_CONSISTENT_CANDIDATE')

# What "no worse" means for MORE_CONSISTENT_CANDIDATE, on the validator's own
# metrics: the first group may not fall, the second may not rise.
MORE_CONSISTENT_HIGHER_IS_BETTER = ('after_target', 'after_floor', 'after100', 'after90', 'after80')
MORE_CONSISTENT_LOWER_IS_BETTER = ('severe_floor_gaps', 'language_gap_count', 'zero_staffed_active_quarters',
                                   'coverage_split_gap_count', 'break_concurrency_violation_count')


def more_consistent_coverage_verdict(selected: dict, candidate: dict) -> tuple:
    """(approved, losses): the polished week is approved only if no coverage measure is worse.

    Both sides are the independent validator's numbers for the two workbooks,
    so the comparison does not trust the engine's own claim (audit F-35).
    """
    losses = [k for k in MORE_CONSISTENT_HIGHER_IS_BETTER
              if float(candidate.get(k) or 0) < float(selected.get(k) or 0)]
    losses += [k for k in MORE_CONSISTENT_LOWER_IS_BETTER
               if float(candidate.get(k) or 0) > float(selected.get(k) or 0)]
    return (not losses, losses)


ALTERNATIVE_USE = {
    'MAX_TARGET_CANDIDATE': 'ALTERNATIVE - most intervals at target; validated on its own',
    'MAX_FLOOR_CANDIDATE': 'ALTERNATIVE - fewest intervals below the floor; validated on its own',
    'BALANCED_CANDIDATE': 'ALTERNATIVE - balanced trade-off; validated on its own',
    'SAFER_BALANCED_CANDIDATE': 'ALTERNATIVE - safer balanced trade-off; validated on its own',
    'MORE_CONSISTENT_CANDIDATE': 'ALTERNATIVE - the same week with steadier start times; '
                                 'offered only if its coverage is no worse',
    'HARD_RULE_SHORTFALL_SCHEDULE': 'NOT RELEASABLE - meets every person rule; every missed coverage '
                                    'minimum is listed on its Shortfalls sheet',
}


def present_alternative(path: Path, role: str) -> str:
    """Same tabs, colours and Break Plan as the published schedule, applied
    BEFORE the validator runs, so the workbook validated is the workbook shipped.
    A presentation failure is recorded and the unpolished workbook is validated
    as written; it never blocks or hides the export."""
    import importlib.util
    try:
        spec = importlib.util.spec_from_file_location('production_output_polisher_alt', POLISHER)
        polisher = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(polisher)
        hidden = polisher.present_in_place(path, role, ALTERNATIVE_USE.get(role, 'ALTERNATIVE'))
    except Exception as exc:  # recorded in UNIVERSAL_RUN_STATUS.json and printed
        print(f"[run] alternative export {role}: presentation failed ({type(exc).__name__}: {exc}); "
              "validating the workbook as written", flush=True)
        return f'FAILED: {type(exc).__name__}: {exc}'
    return 'APPLIED' if hidden is not None else 'SKIPPED_NO_SUMMARY_SHEET'


def validate_alternative_exports(case_root: Path, input_path: Path,
                                 language_working_window: "str | None") -> list:
    """Run the independent validator on every alternative schedule the engine exported.

    Audit F-16: the MAX_TARGET / MAX_FLOOR / BALANCED workbooks sit beside the
    released schedule and a planner may pick one, but only the selected
    schedule was ever validated. Each now carries its own verdict. This does
    not change the release decision for the selected schedule.
    """
    results = []
    for path in sorted(case_root.glob('*_CANDIDATE.xlsx')):
        role = next((r for r in sorted(ALTERNATIVE_EXPORT_ROLES, key=len, reverse=True)
                     if path.stem.endswith('_' + r)), None)
        if role is None:
            continue
        presentation = present_alternative(path, role)
        json_out = case_root / f'INDEPENDENT_VALIDATION_{role}.json'
        command = [sys.executable, '-u', str(VALIDATOR), '--input', str(input_path),
                   '--output', str(path), '--json-out', str(json_out)]
        if language_working_window is not None:
            command += ['--language-working-window', language_working_window]
        proc = subprocess.run(command, capture_output=True, text=True)
        row = {'role': role, 'workbook': str(path), 'json': str(json_out), 'return_code': proc.returncode,
               'presentation': presentation}
        try:
            report = json.loads(json_out.read_text(encoding='utf-8'))
            row['status'] = 'PASS' if proc.returncode == 0 else str(report.get('status') or 'FAIL')
            row['hard_fail_count'] = report.get('hard_fail_count')
            row['failure_types'] = sorted({str(f.get('type')) for f in report.get('failures') or []})
            if role == 'MORE_CONSISTENT_CANDIDATE':
                _more_consistent_review(case_root, report, row)
        except (OSError, json.JSONDecodeError):
            row['status'] = 'ERROR_VALIDATOR_DID_NOT_COMPLETE'
        results.append(row)
    return results


def _more_consistent_review(case_root: Path, report: dict, row: dict) -> None:
    """Approve the polished week only if its coverage is no worse than the selected schedule's."""
    selected_json = case_root / 'INDEPENDENT_VALIDATION.json'
    try:
        selected = json.loads(selected_json.read_text(encoding='utf-8')).get('metrics') or {}
    except (OSError, json.JSONDecodeError):
        row['coverage_no_worse'] = False
        row['coverage_losses'] = ['selected schedule not validated']
    else:
        ok, losses = more_consistent_coverage_verdict(selected, report.get('metrics') or {})
        row['coverage_no_worse'] = ok
        row['coverage_losses'] = losses
    for audit_path in sorted(case_root.glob('*solver_audit.json')):
        try:
            polish = json.loads(audit_path.read_text(encoding='utf-8')).get('shift_consistency_polish') or {}
        except (OSError, json.JSONDecodeError):
            continue
        row['consistency'] = {'before': polish.get('before'), 'after': polish.get('after')}
    if row.get('status') == 'PASS' and not row.get('coverage_no_worse'):
        row['status'] = 'NOT_APPROVED_COVERAGE_WORSE'


if __name__ == '__main__':
    raise SystemExit(main())
