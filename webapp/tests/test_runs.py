# © 2026 Omar Mokhtar. All rights reserved.
"""Phase I task 2: upload, check, queue, run, score, download, expire.

The production runner, input check, gate and scorer are replaced by
webapp/tests/fakes.py so the website's own logic is tested in seconds; the
real runner is exercised end to end separately (evidence/phase_i)."""
import html
import io
import json
import re
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from webapp.app import create_app
from webapp.store import Store

FAKES = Path(__file__).with_name("fakes.py")
TOKEN = re.compile(r'name="csrf_token" value="([^"]+)"')
FINAL = {"DONE", "REVIEW", "FAILED", "REJECTED", "INTERRUPTED", "STOPPED", "EXPIRED"}


def make_app(start_worker=True, gate_fails=False, **extra):
    data = Path(tempfile.mkdtemp())
    package = data / "package"
    package.mkdir()
    (package / "MANIFEST.json").write_text('{"package": "test"}', encoding="utf-8")
    gate_marker = data / "gate_calls"
    if gate_fails:
        Path(str(gate_marker) + ".fail").write_text("x")
    fake = [sys.executable, str(FAKES)]
    config = {"DATA_DIR": str(data), "SECRET_KEY": "t", "TESTING": True, "HTTPS": False,
              "PACKAGE_ROOT": str(package), "START_WORKER": start_worker,
              "RUNNER_CMD": fake + ["runner"], "CHECK_CMD": fake + ["check"],
              "SCORE_CMD": fake + ["score"], "GATE_CMD": fake + ["gate", str(gate_marker)]}
    config.update(extra)
    app = create_app(config)
    store = Store(data / "scheduler.db")
    sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
    store.update_user(sara, all_programs=1)  # Phase Q: sara sees every program, as people from before programs do
    return app, store, data, gate_marker


def client_for(app):
    client = app.test_client()
    tok = TOKEN.search(client.get("/login").get_data(as_text=True)).group(1)
    client.post("/login", data={"username": "sara", "password": "Sara-pass-1", "csrf_token": tok})
    return client


def token(client):
    return TOKEN.search(client.get("/").get_data(as_text=True)).group(1)


def upload(client, name="week42.xlsx", data=b"PK\x03\x04 fake workbook", mode="QUICK", **fields):
    return client.post("/runs", data={"csrf_token": token(client), "mode": mode,
                                      "workbook": (io.BytesIO(data), name), **fields},
                       content_type="multipart/form-data", follow_redirects=False)


def run_id_of(response):
    return response.headers["Location"].rstrip("/").split("/")[-1]


def wait(store, run_id, statuses=FINAL, seconds=30):
    end = time.time() + seconds
    while time.time() < end:
        run = store.get_run(run_id)
        if run and run["status"] in statuses:
            return run
        time.sleep(0.1)
    raise AssertionError(f"run {run_id} still {store.get_run(run_id)['status']}")


class UploadAndCheck(unittest.TestCase):
    def test_upload_rejects_non_xlsx_and_oversize(self):
        app, store, *_ = make_app(start_worker=False)
        client = client_for(app)
        page = upload(client, name="notes.txt")
        self.assertIn("Upload an Excel workbook (.xlsx)", client.get("/").get_data(as_text=True)
                      if page.status_code == 302 else page.get_data(as_text=True))
        self.assertEqual(store.list_runs(), [])
        big = upload(client, data=b"x" * (26 * 1024 * 1024))
        self.assertEqual(big.status_code, 413)

    def test_check_rejection_is_shown_with_its_message(self):
        app, store, *_ = make_app(start_worker=False)
        client = client_for(app)
        run_id = run_id_of(upload(client, name="bad_week.xlsx"))
        run = store.get_run(run_id)
        self.assertEqual(run["status"], "REJECTED")
        # As the reader sees it: the page HTML-escapes the quotes.
        page = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("Schedule!C7: unknown associate 'Zed'", page)


class RunningAndResults(unittest.TestCase):
    def test_accepted_run_goes_through_to_done_with_zip_and_verdict(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client))
        run = wait(store, run_id)
        self.assertEqual((run["status"], run["exit_code"]), ("DONE", 0))
        self.assertIn("RELEASABLE", run["verdict"])
        download = client.get(f"/runs/{run_id}/download")
        self.assertEqual(download.status_code, 200)
        names = zipfile.ZipFile(io.BytesIO(download.data)).namelist()
        self.assertTrue(any(n.endswith("_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx") for n in names), names)
        schedule = client.get(f"/runs/{run_id}/schedule")
        self.assertEqual((schedule.status_code, schedule.data), (200, b"schedule"))

    def test_review_required_verdict_is_never_shown_as_approved(self):
        # The runner exits 0 for RELEASABLE and for REVIEW_REQUIRED (only
        # NOT_RELEASABLE fails it), so the run's label follows the verdict.
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REVIEW_week.xlsx"))
        run = wait(store, run_id, FINAL | {"REVIEW"})
        self.assertEqual((run["status"], run["exit_code"]), ("REVIEW", 0))
        page = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("Needs review", page)
        self.assertNotIn(">Approved<", page)
        self.assertIn("REVIEW_REQUIRED", page)
        self.assertEqual(client.get(f"/runs/{run_id}/schedule").status_code, 200)
        self.assertIn("Needs review", client.get("/").get_data(as_text=True))

    def test_readiness_check_says_ready_not_approved(self):
        app, store, *_ = make_app()
        client = client_for(app)
        # Re-pinned in Phase L: a readiness check is judged by the engine's own
        # outcome (DIAGNOSTICS_ONLY_COMPLETE), because the real runner exits 2
        # for every readiness check (its release scoring fails a run that builds
        # no schedule; owner: "smoke showing not approved regularly"). The fake
        # writes that outcome for READY workbooks; the word is the approved
        # sample's "Ready to run".
        run_id = run_id_of(upload(client, name="READY_week.xlsx", mode="SMOKE"))
        run = wait(store, run_id)
        self.assertEqual(run["status"], "DONE")
        page = client.get(f"/runs/{run_id}").get_data(as_text=True)
        self.assertIn(">Ready to run<", page)
        self.assertNotIn(">Approved<", page)

    def test_summary_saved_when_a_run_finishes(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REAL_week.xlsx"))
        wait(store, run_id)
        saved = json.loads((data / "runs" / run_id / "summary.json").read_text())
        self.assertEqual(saved["numbers"]["fully_covered"], 107)
        self.assertEqual(app.extensions["runs"].summary(run_id)["staffing"]["roster"], 7)
        plain = run_id_of(upload(client, name="week44.xlsx"))
        wait(store, plain)
        self.assertIsNone(app.extensions["runs"].summary(plain))

    def test_failed_runner_exit_code_is_shown(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="FAILS.xlsx"))
        run = wait(store, run_id)
        self.assertEqual(run["exit_code"], 2)
        self.assertIn("Not approved", client.get(f"/runs/{run_id}").get_data(as_text=True))

    def test_two_runs_queue_and_both_finish(self):
        app, store, *_ = make_app()
        client = client_for(app)
        first = run_id_of(upload(client, name="SLOW1.xlsx"))
        second = run_id_of(upload(client, name="week43.xlsx"))
        time.sleep(0.8)
        self.assertEqual(store.get_run(second)["status"], "QUEUED")
        self.assertEqual(wait(store, first)["status"], "DONE")
        self.assertEqual(wait(store, second)["status"], "DONE")

    def test_stop_a_running_run(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="SLOW2.xlsx"))
        wait(store, run_id, {"RUNNING"})
        client.post(f"/runs/{run_id}/stop", data={"csrf_token": token(client)})
        self.assertEqual(wait(store, run_id)["status"], "STOPPED")


class TheVerdictLabel(unittest.TestCase):
    def test_label_reads_the_verdict_word_not_its_note(self):
        from webapp.runs import outcome
        self.assertEqual(outcome("QUICK", "RELEASE VERDICT (run): RELEASABLE - gate 5: none configured")[0], "DONE")
        self.assertEqual(outcome("QUICK", "RELEASE VERDICT (run): REVIEW_REQUIRED - gate 5 behind")[0], "REVIEW")
        self.assertEqual(outcome("QUICK", "no verdict line")[0], "REVIEW")


class RestartAndResume(unittest.TestCase):
    def test_restart_marks_running_as_interrupted(self):
        app, store, data, _ = make_app(start_worker=False)
        store.add_run("abc123abc123", 1, "week.xlsx", "QUICK", "RUNNING")
        create_app(dict(app.config, START_WORKER=False))
        self.assertEqual(store.get_run("abc123abc123")["status"], "INTERRUPTED")

    def test_resume_passes_resume_flag(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client))
        wait(store, run_id)
        store.update_run(run_id, status="INTERRUPTED")
        client.post(f"/runs/{run_id}/resume", data={"csrf_token": token(client)})
        wait(store, run_id)
        argv = json.loads((data / "runs" / run_id / "results" / "argv.json").read_text())
        self.assertIn("--resume", argv["flags"])


class TheSafetyGate(unittest.TestCase):
    def test_gate_runs_once_per_package_and_failure_blocks_runs(self):
        app, store, data, marker = make_app()
        client = client_for(app)
        for name in ("a.xlsx", "b.xlsx"):
            wait(store, run_id_of(upload(client, name=name)))
        self.assertEqual(marker.read_text(), "1")
        app2, store2, data2, marker2 = make_app(gate_fails=True)
        client2 = client_for(app2)
        run = wait(store2, run_id_of(upload(client2)))
        self.assertEqual(run["status"], "FAILED")
        self.assertIn("safety gate", run["message"].lower())

    def test_skip_guards_only_with_pass_stamp(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client))
        wait(store, run_id)
        argv = json.loads((data / "runs" / run_id / "results" / "argv.json").read_text())
        self.assertIn("--skip-guards", argv["flags"])
        stamp = next((data / "runs" / "_gate").glob("*.json"))
        self.assertEqual(json.loads(stamp.read_text())["status"], "PASS")


class Retention(unittest.TestCase):
    def test_runs_older_than_30_days_expire(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client))
        wait(store, run_id)
        queue = app.extensions["runs"]
        self.assertEqual(queue.cleanup(now=time.time() + 29 * 86400), 0)
        self.assertEqual(queue.cleanup(now=time.time() + 31 * 86400), 1)
        self.assertEqual(store.get_run(run_id)["status"], "EXPIRED")
        self.assertFalse((data / "runs" / run_id).exists())
        self.assertEqual(client.get(f"/runs/{run_id}/download").status_code, 404)


class TheControlRoomPages(unittest.TestCase):
    """Phase J task 2 (server side): what the redesigned pages say."""

    def test_run_page_shows_the_week_wall_and_numbers(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REAL_week.xlsx"))
        wait(store, run_id)
        page = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertEqual(page.count('class="cell '), 2 * 126)  # after and before breaks, same places
        self.assertIn("107 of 126 half-hours fully covered", page)
        self.assertIn("Fri 21:30: 2 people on break at the same time", page)
        self.assertIn("about 73 spare productive hours", page)
        self.assertIn("Tue 20:30", page)

    def test_run_without_validation_says_why(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client))
        wait(store, run_id)
        page = client.get(f"/runs/{run_id}").get_data(as_text=True)
        self.assertNotIn('class="cell ', page)
        self.assertIn("No coverage data for this run", page)

    def test_technical_details_are_folded(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="FAILS.xlsx"))
        wait(store, run_id)
        page = client.get(f"/runs/{run_id}").get_data(as_text=True)
        self.assertRegex(page, r'<details class="tech">')
        self.assertIn("Not approved", page)

    def test_queue_position_and_start_estimate_are_shown(self):
        app, store, *_ = make_app()
        client = client_for(app)
        first = run_id_of(upload(client, name="SLOW_a.xlsx"))
        wait(store, first, {"RUNNING"})
        second = run_id_of(upload(client, name="week45.xlsx"))
        page = client.get(f"/runs/{second}").get_data(as_text=True)
        self.assertIn("1 run ahead of you", page)
        self.assertIn("Starts in about", page)
        self.assertIn("Egypt time", page)
        board = client.get("/").get_data(as_text=True)
        self.assertIn("1 waiting", board)
        wait(store, second)

    def test_dashboard_shows_the_latest_wall(self):
        app, store, *_ = make_app()
        client = client_for(app)
        wait(store, run_id_of(upload(client, name="REAL_week.xlsx")))
        board = html.unescape(client.get("/").get_data(as_text=True))
        self.assertIn("Latest schedule", board)
        self.assertEqual(board.count('class="cell '), 126)
        self.assertIn("85%", board)  # 107 of 126 fully covered

    def test_server_error_page_is_friendly(self):
        app, store, *_ = make_app(start_worker=False)

        def boom():
            raise RuntimeError("unexpected")
        app.add_url_rule("/boom", "boom", boom)
        app.testing = False
        app.config["PROPAGATE_EXCEPTIONS"] = False
        response = client_for(app).get("/boom")
        self.assertEqual(response.status_code, 500)
        self.assertIn("Something went wrong on our side. Your runs are safe. Tell your admin.",
                      html.unescape(response.get_data(as_text=True)))


class TheAdvancedOptions(unittest.TestCase):
    """Phase J task 6: the Colab notebook's settings, folded under the mode."""

    def argv(self, data, run_id):
        return " ".join(json.loads((data / "runs" / run_id / "results" / "argv.json").read_text())["argv"])

    def test_advanced_options_reach_the_runner(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, mode="OVERNIGHT", language_window="ALL_ROWS",
                                  coverage_measure="VOLUME_WEIGHTED", stage="BEFORE_BREAKS_ONLY"))
        wait(store, run_id)
        argv = self.argv(data, run_id)
        for needle in ("--mode OVERNIGHT", "--language-working-window ALL_ROWS",
                       "--coverage-objective-weighting VOLUME_WEIGHTED", "--stage BEFORE_BREAKS_ONLY"):
            self.assertIn(needle, argv)
        page = client.get(f"/runs/{run_id}").get_data(as_text=True)
        self.assertIn("All rows", page)

    def test_follow_the_workbook_adds_no_flags(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, language_window="workbook", coverage_measure="workbook"))
        wait(store, run_id)
        argv = self.argv(data, run_id)
        self.assertNotIn("--language-working-window", argv)
        self.assertNotIn("--coverage-objective-weighting", argv)
        self.assertIn("--stage FULL_SCHEDULE", argv)

    def test_unknown_option_values_are_refused(self):
        app, store, *_ = make_app(start_worker=False)
        client = client_for(app)
        upload(client, language_window="EVERYTHING")
        self.assertEqual(store.list_runs(), [])
        self.assertIn("Choose one of the listed options", client.get("/").get_data(as_text=True))

    def test_old_database_gains_the_options_column(self):
        import sqlite3
        data = Path(tempfile.mkdtemp())
        db = sqlite3.connect(data / "scheduler.db")
        db.executescript("""create table users (id integer primary key, username text unique not null,
            display_name text not null, password_hash text not null, is_admin integer not null default 0,
            active integer not null default 1, must_change integer not null default 1,
            failed integer not null default 0, locked_until real not null default 0, created real not null);
            create table runs (id text primary key, user_id integer not null, workbook text not null,
            mode text not null, status text not null, message text not null default '', exit_code integer,
            verdict text not null default '', created real not null, started real, finished real,
            resume integer not null default 0);
            insert into users values (1,'sara','Sara','x',0,1,0,0,0,0);
            insert into runs (id,user_id,workbook,mode,status,created) values ('abcabcabcabc',1,'w.xlsx','QUICK','DONE',0);""")
        db.commit()
        db.close()
        run = Store(data / "scheduler.db").get_run("abcabcabcabc")
        self.assertEqual(run["options"], "")


class TheProgramTags(unittest.TestCase):
    """Phase K task 2: program and week on every run; figures kept after the files expire."""

    def test_upload_records_program_and_week(self):
        app, store, *_ = make_app(start_worker=False)
        client = client_for(app)
        run_id = run_id_of(upload(client, program="NMG Spanish", week_start="2026-10-11"))
        run = store.get_run(run_id)
        self.assertEqual((run["program"], run["week_start"]), ("NMG Spanish", "2026-10-11"))

    def test_week_must_be_a_date(self):
        app, store, *_ = make_app(start_worker=False)
        client = client_for(app)
        upload(client, program="NMG", week_start="next week")
        self.assertEqual(store.list_runs(), [])

    def test_metrics_saved_when_a_run_finishes(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REAL_week.xlsx", program="NMG", week_start="2026-10-11"))
        wait(store, run_id)
        m = json.loads(store.get_run(run_id)["metrics"])
        self.assertEqual((m["associates"], m["fully_covered"], m["active"]), (7, 107, 126))
        self.assertEqual(m["outcome"], "DONE")  # kept with the figures: the status is lost on expiry

    def test_history_survives_file_expiry(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REAL_week.xlsx", program="NMG", week_start="2026-10-11"))
        wait(store, run_id)
        app.extensions["runs"].cleanup(now=time.time() + 31 * 86400)
        run = store.get_run(run_id)
        self.assertEqual(run["status"], "EXPIRED")
        self.assertEqual(json.loads(run["metrics"])["fully_covered"], 107)

    def test_runs_finished_before_the_update_join_the_history(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REAL_week.xlsx"))
        wait(store, run_id)
        store.update_run(run_id, metrics="")  # as a run finished before program history existed
        self.assertEqual(app.extensions["runs"].backfill(), 1)
        m = json.loads(store.get_run(run_id)["metrics"])
        self.assertEqual((m["fully_covered"], m["outcome"]), (107, "DONE"))
        self.assertEqual(app.extensions["runs"].backfill(), 0)  # once only

    def test_untagged_runs_can_be_tagged(self):
        app, store, *_ = make_app(start_worker=False)
        store.add_user("lina", "Lina", "Lina-pass-12", must_change=False)
        client = client_for(app)
        run_id = run_id_of(upload(client))
        other = app.test_client()
        tok = TOKEN.search(other.get("/login").get_data(as_text=True)).group(1)
        other.post("/login", data={"username": "lina", "password": "Lina-pass-12", "csrf_token": tok})
        refused = other.post(f"/runs/{run_id}/tag", data={"csrf_token": token(other), "program": "X",
                                                          "week_start": "2026-10-11"})
        self.assertEqual(refused.status_code, 403)
        from webapp.programs import ProgramBook
        ProgramBook(store).add_program("NMG")  # re-pinned (Phase R): a run's program is picked from the
        # programs set up under LOBs and defaults, never typed (owner, 2026-10-08: typing made stray programs)
        client.post(f"/runs/{run_id}/tag", data={"csrf_token": token(client), "program": "NMG",
                                                 "week_start": "2026-10-11"})
        run = store.get_run(run_id)
        self.assertEqual((run["program"], run["week_start"]), ("NMG", "2026-10-11"))

    def test_old_database_gains_analytics_columns(self):
        import sqlite3
        data = Path(tempfile.mkdtemp())
        db = sqlite3.connect(data / "scheduler.db")
        db.executescript("""create table users (id integer primary key, username text unique not null,
            display_name text not null, password_hash text not null, is_admin integer not null default 0,
            active integer not null default 1, must_change integer not null default 1,
            failed integer not null default 0, locked_until real not null default 0, created real not null);
            create table runs (id text primary key, user_id integer not null, workbook text not null,
            mode text not null, status text not null, message text not null default '', exit_code integer,
            verdict text not null default '', created real not null, started real, finished real,
            resume integer not null default 0, options text not null default '');
            insert into users values (1,'sara','Sara','x',0,1,0,0,0,0);
            insert into runs (id,user_id,workbook,mode,status,created) values ('abcabcabcabc',1,'w.xlsx','QUICK','DONE',0);""")
        db.commit()
        db.close()
        run = Store(data / "scheduler.db").get_run("abcabcabcabc")
        self.assertEqual((run["program"], run["week_start"], run["metrics"]), ("", "", ""))


def seed_week(store, program, week, outcome="DONE", by=1, n=[0], **changes):
    """A finished run with stored figures, as a run of the real Phase I workbook would leave."""
    from webapp.results import metrics, summarize
    m = metrics(summarize(Path(__file__).with_name("fixtures") / "real_run"))
    m.update(changes, outcome=outcome)
    n[0] += 1
    run_id = f"{n[0]:012x}"
    store.add_run(run_id, by, f"{program}_{week}.xlsx", "QUICK", outcome, program=program, week_start=week)
    store.update_run(run_id, metrics=json.dumps(m), started=time.time() - 3600, finished=time.time())
    return run_id


class TheAnalyticsPages(unittest.TestCase):
    """Phase K task 5: programs, one program's history, the team."""

    def setUp(self):
        self.app, self.store, *_ = make_app(start_worker=False)
        self.client = client_for(self.app)
        seed_week(self.store, "NMG", "2026-09-27", associates=8, fully_covered=100)
        seed_week(self.store, "NMG", "2026-10-04", associates=8, fully_covered=102)
        seed_week(self.store, "NMG", "2026-10-11")  # the real run: 7 associates, 107 of 126
        seed_week(self.store, "AE/AR <B2B>", "2026-10-11", outcome="REVIEW", associates=12)

    def page(self, url):
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200, url)
        return html.unescape(response.get_data(as_text=True))

    def test_programs_page_lists_programs_with_latest_figures(self):
        body = self.page("/programs")
        self.assertIn("NMG", body)
        self.assertIn("3 weeks", body)
        self.assertIn("85%", body)  # 107 of 126 after breaks, the latest NMG week
        self.assertRegex(body, r"7\s*<[^>]*>\s*down 1")  # associates with the change from last week

    def test_program_page_shows_history_insights_and_suggestions(self):
        body = self.page("/programs/NMG")
        self.assertIn("Associates: 7 this week, down 1 from 8 in the week of 04 Oct.", body)
        self.assertIn("Fully covered before breaks: 99%, so breaks cost 14 points this week.", body)
        self.assertIn("at 90% it is met in 96%", body)
        for chart in ("Associates per week", "Coverage before and after breaks", "Hours available and needed",
                      "What stops 100%", "If the interval target were lower", "Hours that keep coming up short"):
            self.assertIn(chart, body)
        self.assertEqual(body.count('<details class="tableview">'), 6)
        self.assertIn("2026-10-11", body)  # the week-by-week table

    def test_program_weeks_filter_limits_the_history(self):
        body = self.page("/programs/NMG?weeks=2")
        self.assertNotIn("27 Sep", body)
        self.assertIn("04 Oct", body)

    def test_program_names_are_escaped_and_url_safe(self):
        raw = self.client.get("/programs").get_data(as_text=True)
        self.assertNotIn("<B2B>", raw)
        self.assertIn("/programs/AE/AR%20%3CB2B%3E", raw)
        self.assertIn("12", self.page("/programs/AE/AR%20%3CB2B%3E"))
        self.assertEqual(self.client.get("/programs/Nothing").status_code, 404)

    def test_team_page(self):
        body = self.page("/team")
        self.assertIn("Sara", body)
        self.assertRegex(body, r"Sara</td>\s*<td>4</td>")

    def test_analytics_need_login(self):
        anonymous = self.app.test_client()
        for url in ("/programs", "/programs/NMG", "/team"):
            self.assertEqual(anonymous.get(url).status_code, 302, url)

    def test_upload_form_offers_program_and_week(self):
        body = self.page("/")
        self.assertIn('name="program"', body)
        self.assertIn('<option value="NMG">', body)
        # Re-pinned in Phase P (owner, 2026-10-08): the week is now a start-date dropdown (Sundays and Mondays).
        self.assertIn('<select name="week_start"', body)

    def test_run_page_lets_the_owner_tag_the_run(self):
        run_id = seed_week(self.store, "", "")
        self.assertIn(f'action="/runs/{run_id}/tag"', self.page(f"/runs/{run_id}"))


class TheThemes(unittest.TestCase):
    """Phase L: light or dark, the person's choice; the system's choice until they pick."""

    def test_theme_cookie_sets_the_page_theme(self):
        app, *_ = make_app(start_worker=False)
        client = client_for(app)
        self.assertNotIn("data-theme=", client.get("/").get_data(as_text=True).split(">", 2)[1])
        response = client.post("/theme", data={"csrf_token": token(client), "theme": "light", "next": "/programs"})
        self.assertEqual(response.headers["Location"], "/programs")
        self.assertIn("theme=light", response.headers["Set-Cookie"])
        self.assertIn('<html lang="en" data-theme="light">', client.get("/").get_data(as_text=True))
        client.post("/theme", data={"csrf_token": token(client), "theme": "system"})
        self.assertIn('<html lang="en">', client.get("/").get_data(as_text=True))

    def test_theme_toggle_needs_csrf(self):
        app, *_ = make_app(start_worker=False)
        client = client_for(app)
        self.assertEqual(client.post("/theme", data={"theme": "light"}).status_code, 400)

    def test_theme_only_returns_to_this_site(self):
        app, *_ = make_app(start_worker=False)
        client = client_for(app)
        response = client.post("/theme", data={"csrf_token": token(client), "theme": "dark",
                                                "next": "https://evil.example/"})
        self.assertEqual(response.headers["Location"], "/")

    def test_sign_in_page_follows_the_theme_too(self):
        app, *_ = make_app(start_worker=False)
        client = app.test_client()
        client.set_cookie("theme", "light")
        self.assertIn('data-theme="light"', client.get("/login").get_data(as_text=True))


class TheEngineOutcome(unittest.TestCase):
    """Phase L task 1: why a run stopped and what to change, on its own page."""

    def test_failed_run_page_shows_why_and_what_to_change(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="CONFLICT_week.xlsx"))
        run = wait(store, run_id)
        self.assertEqual(run["status"], "FAILED")
        self.assertEqual(run["message"], "No schedule satisfies all hard rules together")
        body = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("Can't be scheduled", body)
        self.assertIn("Why it stopped", body)
        self.assertIn("What to change in the input workbook", body)
        self.assertIn("widen that language's Coverage Start/End in Language Setup", body)
        self.assertIn("Associate 1 (English)", body)
        self.assertIn("Mon to Thu", body)
        self.assertNotIn("tools/check_input_workbook.py", body.split("Technical details")[0])
        self.assertNotIn('action="/runs/%s/resume"' % run_id, body)  # the same file would fail the same way
        self.assertIn("Upload the corrected workbook", body)

    def test_outcome_survives_file_expiry(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="CONFLICT_week.xlsx"))
        wait(store, run_id)
        app.extensions["runs"].cleanup(now=time.time() + 31 * 86400)
        self.assertEqual(json.loads(store.get_run(run_id)["engine_outcome"])["headline"],
                         "No schedule satisfies all hard rules together")
        self.assertIn("No schedule satisfies all hard rules together",
                      html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True)))

    def test_shortfall_schedule_offered_for_review(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="SHORTFALL_week.xlsx"))
        wait(store, run_id)
        body = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("Sun 03:00: Nobody on the floor after breaks, needs 1, short by 1", body)
        self.assertIn("Download the shortfall schedule (for review)", body)
        self.assertNotIn(">Download the schedule<", body)
        got = client.get(f"/runs/{run_id}/shortfall")
        self.assertEqual(got.data, b"shortfall")

    def test_crash_without_outcome_keeps_resume(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="FAILS_week.xlsx"))
        run = wait(store, run_id)
        self.assertEqual(run["status"], "FAILED")
        body = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("Not approved", body)
        self.assertIn("the engine wrote no outcome", body)
        self.assertIn(f'action="/runs/{run_id}/resume"', body)


class TheReadinessCheck(unittest.TestCase):
    """Phase L task 2: a readiness check says Ready to run or Not ready (never
    Not approved), and a ready one starts the real run from the same workbook."""

    def ready_run(self, **fields):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="READY_week.xlsx", mode="SMOKE", **fields))
        wait(store, run_id)
        return app, store, client, run_id

    def test_readiness_judged_by_engine_outcome(self):
        app, store, client, run_id = self.ready_run()
        run = store.get_run(run_id)
        self.assertEqual((run["status"], run["exit_code"]), ("DONE", 2))  # the runner's 2 is its release scoring
        self.assertIn(">Ready to run<", client.get(f"/runs/{run_id}").get_data(as_text=True))

    def test_readiness_not_ready_shows_reason(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="NOTREADY_week.xlsx", mode="SMOKE"))
        run = wait(store, run_id)
        self.assertEqual(run["status"], "FAILED")
        body = html.unescape(client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn(">Not ready<", body)
        self.assertIn("Why it is not ready", body)
        self.assertNotIn("/resume", body)

    def test_readiness_without_an_outcome_is_not_ready(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="plain_week.xlsx", mode="SMOKE"))
        self.assertEqual(wait(store, run_id)["status"], "FAILED")  # fail closed: no outcome, no "ready"

    def test_readiness_page_offers_start_not_resume(self):
        app, store, client, run_id = self.ready_run()
        body = client.get(f"/runs/{run_id}").get_data(as_text=True)
        self.assertIn(f'action="/runs/{run_id}/start"', body)
        for mode in ("QUICK", "DEEP", "OVERNIGHT"):
            self.assertIn(f'name="mode" value="{mode}"', body)
        self.assertIn("Start the run", body)
        self.assertNotIn("/resume", body)

    def test_old_readiness_checks_are_read_again_at_startup(self):
        """Checks finished under the old rule (runner exit code) read "Not approved";
        their engine outcome is read again when the website starts."""
        app, store, client, run_id = self.ready_run()
        store.update_run(run_id, status="FAILED", engine_outcome="",
                         message="Not approved: the runner exited with code 2.")  # as the old rule left it
        not_ready = run_id_of(upload(client, name="NOTREADY_old.xlsx", mode="SMOKE"))
        wait(store, not_ready)
        store.update_run(not_ready, engine_outcome="")
        self.assertEqual(app.extensions["runs"].backfill_outcomes(), 2)
        run = store.get_run(run_id)
        self.assertEqual(run["status"], "DONE")
        self.assertTrue(run["message"].startswith("Ready to run"))
        self.assertEqual(store.get_run(not_ready)["status"], "FAILED")
        self.assertEqual(json.loads(store.get_run(not_ready)["engine_outcome"])["code"], "HARD_RULE_COMBINATION_INFEASIBLE")
        self.assertEqual(app.extensions["runs"].backfill_outcomes(), 0)  # once only

    def test_start_from_readiness_queues_the_same_workbook(self):
        app, store, client, run_id = self.ready_run(program="NMG", week_start="2026-10-11", language_window="ALL_ROWS")
        response = client.post(f"/runs/{run_id}/start", data={"csrf_token": token(client), "mode": "QUICK"})
        new_id = run_id_of(response)
        self.assertNotEqual(new_id, run_id)
        new, old = store.get_run(new_id), store.get_run(run_id)
        self.assertEqual((new["mode"], new["program"], new["week_start"], new["options"], new["workbook"]),
                         ("QUICK", "NMG", "2026-10-11", old["options"], "READY_week.xlsx"))
        queue = app.extensions["runs"]
        self.assertEqual(queue.input_path(new_id).read_bytes(), queue.input_path(run_id).read_bytes())

    def test_start_takes_only_a_real_run_mode(self):
        app, store, client, run_id = self.ready_run()
        before = len(store.list_runs())
        client.post(f"/runs/{run_id}/start", data={"csrf_token": token(client), "mode": "SMOKE_OR_WHATEVER"})
        self.assertEqual(len(store.list_runs()), before)

    def test_start_needs_the_input(self):
        app, store, client, run_id = self.ready_run()
        import shutil
        shutil.rmtree(app.extensions["runs"].input_path(run_id).parent)
        before = len(store.list_runs())
        response = client.post(f"/runs/{run_id}/start", data={"csrf_token": token(client), "mode": "QUICK"},
                               follow_redirects=True)
        self.assertEqual(len(store.list_runs()), before)
        self.assertIn("The workbook of this check is no longer on the server", html.unescape(response.get_data(as_text=True)))

    def test_corrected_upload_link_prefills_program_and_week(self):
        app, *_ = make_app(start_worker=False)
        client = client_for(app)
        body = client.get("/?program=NMG%20Spanish&week=2026-10-11").get_data(as_text=True)
        # Re-pinned in Phase Q (owner, 2026-10-08): the program is picked from a Program and LOB dropdown now;
        # a prefilled program is its selected option.
        self.assertIn('<option value="NMG Spanish" selected>NMG Spanish</option>', body)
        # Re-pinned in Phase P (owner, 2026-10-08): the week is now a start-date dropdown (Sundays and Mondays),
        # so the prefilled date is the selected option rather than a date box's value.
        self.assertRegex(body, r'<option value="2026-10-11" data-day="0" selected>Sun 11 Oct 2026</option>')


class TheWayAround(unittest.TestCase):
    """Phase L task 3: Back and Home on every page but home."""

    def test_every_page_has_back_and_home(self):
        app, store, *_ = make_app(start_worker=False)
        store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        client = app.test_client()
        tok = TOKEN.search(client.get("/login").get_data(as_text=True)).group(1)
        client.post("/login", data={"username": "omar", "password": "Owner-pass-123", "csrf_token": tok})
        run_id = seed_week(store, "NMG", "2026-10-11")
        parents = {f"/runs/{run_id}": "/", "/programs/NMG": "/programs", "/programs": "/", "/team": "/",
                   "/admin/users": "/", "/account/password": "/"}
        for url, parent in parents.items():
            body = client.get(url).get_data(as_text=True)
            self.assertIn(f'<a href="{parent}" data-back>', body, url)
            self.assertIn('<a class="home" href="/">Home</a>', body, url)
        home = client.get("/").get_data(as_text=True)
        self.assertNotIn("data-back", home)
        # Re-pinned in Phase Q (owner, 2026-10-08): the menu moved to the left and marks the page open,
        # so on Home its first item carries that mark.
        self.assertIn('<a href="/" class="on" aria-current="page">Home</a>', home)  # the menu's first item


class TheCleanWorkbooks(unittest.TestCase):
    """Phase L task 4: a blank and an example input workbook, from the home page."""

    def test_only_the_two_workbooks_download(self):
        app, *_ = make_app(start_worker=False)
        client = client_for(app)
        home = client.get("/").get_data(as_text=True)
        for name in ("Scheduler_Input_Blank.xlsx", "Scheduler_Input_Example.xlsx"):
            self.assertIn(f'href="/workbooks/{name}"', home)
            got = client.get(f"/workbooks/{name}")
            self.assertEqual((got.status_code, got.data[:2]), (200, b"PK"), name)
            self.assertIn("attachment", got.headers["Content-Disposition"])
        for bad in ("other.xlsx", "..%2Fapp.py", "Scheduler_Input_Blank.xlsx.bak"):
            self.assertEqual(client.get(f"/workbooks/{bad}").status_code, 404, bad)
        self.assertEqual(app.test_client().get("/workbooks/Scheduler_Input_Blank.xlsx").status_code, 302)


class TheWeekPage(unittest.TestCase):
    """Phase M task 2: one run's week, or a program's week, interval by interval."""

    def finished(self, **fields):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, name="REAL_week.xlsx", **fields))
        wait(store, run_id)
        return app, store, client, run_id

    def page(self, client, url):
        response = client.get(url)
        self.assertEqual(response.status_code, 200, url)
        # re-pinned (Phase AC, review finding 9): a tile's "of N" is now a smaller part of the figure,
        # <b>107<small class="of"> of 126</small></b>, so it stays on one line; the figures checked are unchanged
        return re.sub(r'<small class="of">( of \d+)</small>', r"\1", html.unescape(response.get_data(as_text=True)))

    def test_week_view_for_a_run(self):
        app, store, client, run_id = self.finished()
        body = self.page(client, f"/runs/{run_id}/week")
        for text in ("Achieved every 30 minutes", "Where overtime is needed", "Extra hours available",
                     "107 of 126", "9.5 h", "24.5 h", "02:00 to 17:00"):
            self.assertIn(text, body)
        self.assertIn("125 of 126", self.page(client, f"/runs/{run_id}/week?side=before"))

    def test_week_view_by_program_and_week(self):
        app, store, client, run_id = self.finished(program="NMG", week_start="2026-10-11")
        body = self.page(client, "/week?program=NMG&week=2026-10-11")
        self.assertIn("NMG, week of 11 Oct", body)
        self.assertIn("107 of 126", body)
        self.assertIn('<option value="NMG" selected>', body)
        self.assertIn('<option value="2026-10-11" selected>', body)
        self.assertIn("107 of 126", self.page(client, "/week"))  # no choice yet: the latest program week

    def test_week_view_survives_file_expiry(self):
        app, store, client, run_id = self.finished()
        app.extensions["runs"].cleanup(now=time.time() + 31 * 86400)
        self.assertIn("107 of 126", self.page(client, f"/runs/{run_id}/week"))

    def test_older_runs_gain_intervals(self):
        app, store, client, run_id = self.finished()
        m = json.loads(store.get_run(run_id)["metrics"])
        self.assertIn("intervals", m)
        del m["intervals"]
        store.update_run(run_id, metrics=json.dumps(m))  # as a run finished before Phase M
        self.assertEqual(app.extensions["runs"].backfill(), 1)
        self.assertEqual(len(json.loads(store.get_run(run_id)["metrics"])["intervals"]), 126)  # intervals with demand

    def test_unknown_week_says_so(self):
        app, store, client, run_id = self.finished()
        self.assertIn("No finished run for this program and week", self.page(client, "/week?program=Nope&week=2026-10-11"))
        rejected = run_id_of(upload(client, name="bad.xlsx"))
        self.assertIn("No interval figures for this run", self.page(client, f"/runs/{rejected}/week"))

    def test_week_links_from_run_and_program_pages(self):
        app, store, client, run_id = self.finished(program="NMG", week_start="2026-10-11")
        self.assertIn(f'href="/runs/{run_id}/week"', self.page(client, f"/runs/{run_id}"))
        self.assertIn(f'href="/runs/{run_id}/week"', self.page(client, "/programs/NMG"))


REPO = Path(__file__).resolve().parents[2]
B3_INPUT = REPO / "fixtures" / "real_runs" / "week_boundary" / "B3_ARB2B_S30" / "input_snapshot" / "AE_AR_B2B_SLICE30.xlsx"


def versioned_run(app, store, client, **fields):
    """A finished run whose files are a real schedule week (fake runner, real workbooks)."""
    run_id = run_id_of(upload(client, name="VERSIONED_week.xlsx", data=B3_INPUT.read_bytes(), **fields))
    wait(store, run_id)
    end = time.time() + 60
    while time.time() < end and not app.extensions["schedules"].versions(run_id):
        time.sleep(0.2)
    return run_id


class TheKeptSchedules(unittest.TestCase):
    """Phase N task 2: a finished run's schedules are kept as versions."""

    def test_finished_run_keeps_its_schedules(self):
        app, store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        client = client_for(app)
        run_id = versioned_run(app, store, client, program="AE/AR B2B", week_start="2026-10-11")
        versions = app.extensions["schedules"].versions(run_id)
        self.assertEqual([(v["kind"], v["program"]) for v in versions], [("tool_after", "AE/AR B2B")])
        self.assertEqual(json.loads(versions[0]["checks"])["status"], "PASS")

    def test_a_finished_run_has_kept_its_schedules_already(self):
        # "finished" means the run's work is over: its versions are kept before the status says
        # so (a clean-up starting the moment a run reads finished must not race the copy)
        app, store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        client = client_for(app)
        run_id = run_id_of(upload(client, name="VERSIONED_week.xlsx", data=B3_INPUT.read_bytes()))
        wait(store, run_id)
        self.assertTrue(app.extensions["schedules"].versions(run_id))


class TheSchedulesPage(unittest.TestCase):
    """Phase N task 3: versions, edits with their warning, in use, downloads."""

    @classmethod
    def setUpClass(cls):
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("lina", "Lina", "Lina-pass-12", must_change=False)
        cls.client = client_for(cls.app)
        cls.run_id = versioned_run(cls.app, cls.store, cls.client, program="AE/AR B2B", week_start="2026-10-11")
        cls.tool = cls.app.extensions["schedules"].versions(cls.run_id)[0]["id"]

    def post(self, url, client=None, **data):
        client = client or self.client
        return client.post(url, data={"csrf_token": token(client), **data})

    def test_schedules_page_lists_versions_and_marks(self):
        body = html.unescape(self.client.get(f"/runs/{self.run_id}/schedules").get_data(as_text=True))
        self.assertIn("AE/AR B2B, week of 11 Oct: schedules", body)
        self.assertIn("Tool, after breaks", body)
        self.assertIn('data-name="Associate 001" data-day="Wed" data-value="12:00 - 21:00"', body)
        self.assertIn('<option value="21:00 - 06:00">', body)
        draft = self.app.extensions["schedules"].change(self.tool, 1, "Associate 001", "Wed", "21:00 - 06:00", "swap")
        body = html.unescape(self.client.get(f"/runs/{self.run_id}/schedules?v={draft}").get_data(as_text=True))
        self.assertRegex(body, r'class="ed changed sev-red"[^>]*data-name="Associate 001" data-day="Wed"')
        self.assertIn('class="dot red"', body)
        self.assertIn("rests only 6 hours", body)
        self.assertIn("swap", body)  # the change log

    def test_check_returns_new_problems_only(self):
        got = self.post(f"/schedules/{self.tool}/check", associate="Associate 001", day="Wed", value="21:00 - 06:00").get_json()
        self.assertEqual(got["severity"], "red")
        self.assertEqual(len(got["added"]), 3)
        self.assertTrue(any("rests only 6 hours" in p["text"] for p in got["added"]))
        same = self.post(f"/schedules/{self.tool}/check", associate="Associate 001", day="Wed", value="12:00 - 21:00").get_json()
        self.assertEqual((same["severity"], same["added"]), ("ok", []))

    def test_change_needs_a_reason_when_it_breaks_rules(self):
        refused = self.post(f"/schedules/{self.tool}/change", associate="Associate 004", day="Wed", value="21:00 - 06:00")
        self.assertEqual(refused.status_code, 400)
        self.assertIn("reason", refused.get_json()["error"])
        kept = self.post(f"/schedules/{self.tool}/change", associate="Associate 004", day="Wed", value="21:00 - 06:00",
                         reason="cover for a colleague").get_json()
        self.assertIn("/schedules?v=", kept["url"])
        self.assertNotEqual(kept["schedule_id"], self.tool)

    def test_slot_swap_checked_then_kept(self):
        body = html.unescape(self.client.get(f"/runs/{self.run_id}/schedules").get_data(as_text=True))
        self.assertIn("Slot 1 · English", body)
        got = self.post(f"/schedules/{self.tool}/swap-check", first="Associate 005", second="Associate 006").get_json()
        self.assertIn(got["severity"], ("ok", "yellow", "red"))
        refused = self.post(f"/schedules/{self.tool}/swap", first="Associate 005", second="Associate 005", reason="x")
        self.assertEqual(refused.status_code, 400)
        kept = self.post(f"/schedules/{self.tool}/swap", first="Associate 005", second="Associate 006",
                         reason="swap requested").get_json()
        week = json.loads(self.store.get_schedule(kept["schedule_id"])["week"])
        self.assertEqual([a["name"] for a in week["associates"] if a["slot"] in ("5", "6")],
                         ["Associate 006", "Associate 005"])
        page = html.unescape(self.client.get(kept["url"]).get_data(as_text=True))
        self.assertIn("Associate 005 moved from Slot 5 to Slot 6 (slot swap)", page)

    def test_in_use_only_for_owner_or_admin(self):
        other = self.app.test_client()
        tok = TOKEN.search(other.get("/login").get_data(as_text=True)).group(1)
        other.post("/login", data={"username": "lina", "password": "Lina-pass-12", "csrf_token": tok})
        self.assertEqual(self.post(f"/schedules/{self.tool}/in-use", client=other).status_code, 403)
        self.assertEqual(self.post(f"/schedules/{self.tool}/in-use").status_code, 302)
        self.assertEqual(self.app.extensions["schedules"].in_use("AE/AR B2B", "2026-10-11")["id"], self.tool)

    def test_download_a_version(self):
        tool = self.client.get(f"/schedules/{self.tool}/download")
        self.assertEqual((tool.status_code, tool.data[:2]), (200, b"PK"))
        draft = self.app.extensions["schedules"].change(self.tool, 1, "Associate 002", "Sun", "OFF", "leave")
        got = self.client.get(f"/schedules/{draft}/download")
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(got.data))
        self.assertEqual(book.sheetnames[0], "Version Notes")
        notes = [c for row in book["Version Notes"].iter_rows(values_only=True) for c in row if c]
        self.assertIn("Associate 002", " ".join(map(str, notes)))
        self.assertIn("attachment", got.headers["Content-Disposition"])

    def test_week_view_of_a_version(self):
        body = html.unescape(self.client.get(f"/schedules/{self.tool}/week").get_data(as_text=True))
        self.assertIn("Tool, after breaks", body)
        self.assertIn("Achieved every hour", body)  # this program's demand is hourly

    def test_links_to_the_schedules(self):
        self.assertIn(f'href="/runs/{self.run_id}/schedules"', self.client.get(f"/runs/{self.run_id}").get_data(as_text=True))


class TheDayPage(unittest.TestCase):
    """Phase N task 4: the day from the schedule in use, attendance and actual breaks."""

    @classmethod
    def setUpClass(cls):
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.client = client_for(cls.app)
        cls.run_id = versioned_run(cls.app, cls.store, cls.client, program="AE/AR B2B", week_start="2026-10-11")
        cls.url = "/day?program=AE/AR+B2B&date=2026-10-14"

    def post(self, url, **data):
        return self.client.post(url, data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                           "date": "2026-10-14", **data})

    def page(self):
        return html.unescape(self.client.get(self.url).get_data(as_text=True))

    def test_add_route_records_and_redirects_to_the_board_row(self):
        """Phase R task 4: "+ Add" posts here; back to the board at that interval with what was kept."""
        # re-pinned (Phase T): an aux is booked with who it is with and why (owner, 2026-10-09)
        got = self.post("/day/add", date="2026-10-15", associate="Associate 001", what="Coaching", **{"from": "15:00"},
                        minutes="30", billable="1", view="board", at="900", with_whom="Sara", why="Quality review")
        self.assertEqual(got.status_code, 303)
        self.assertIn("view=board", got.headers["Location"])
        self.assertTrue(got.headers["Location"].endswith("#row-900"))
        acts = self.store.list_activities("AE/AR B2B", ["2026-10-15"])
        self.assertEqual([(a["associate"], a["kind"], a["start"], a["end_min"], a["billable"]) for a in acts],
                         [("Associate 001", "Coaching", 900, 930, 1)])
        page = html.unescape(self.client.get(got.headers["Location"]).get_data(as_text=True))
        self.assertIn("Recorded: Associate 001, Coaching 15:00 to 15:30 (billable), with Sara: Quality review.", page)
        from datetime import date
        days = self.app.extensions["days"]
        days.cancel_activity("AE/AR B2B", date(2026, 10, 15), acts[0]["id"], 1)
        refused = self.post("/day/add", date="2026-10-15", associate="Associate 001", what="Late", **{"from": "22:00"},
                            view="board", at="900")
        page = html.unescape(self.client.get(refused.headers["Location"]).get_data(as_text=True))
        self.assertIn("outside Associate 001's shift", page)

    def test_wallboard_lists_breaks_added_on_the_day(self):
        from datetime import date
        days = self.app.extensions["days"]
        made = days.add_activity("AE/AR B2B", date(2026, 10, 15), "Associate 001", "Break", "15:00", "15:15", 1)
        try:
            body = html.unescape(self.client.get("/day/wallboard?program=AE/AR+B2B&date=2026-10-15&at=15:05")
                                 .get_data(as_text=True))
            on_break = body[body.index("On break now"):body.index("Next 30 minutes")]
            self.assertIn("<b>Associate 001</b> Break until 15:15", on_break)
        finally:
            days.cancel_activity("AE/AR B2B", date(2026, 10, 15), made, 1)

    def test_add_preview_says_the_effect_or_why_not(self):
        got = self.post("/day/add-preview", date="2026-10-15", associate="Associate 001", what="Break",
                        **{"from": "15:00"}, minutes="15")
        self.assertEqual(got.status_code, 200)
        self.assertIn("the floor at its tightest goes from", got.get_json()["text"])
        self.assertEqual(self.store.list_activities("AE/AR B2B", ["2026-10-15"]), [])
        got = self.post("/day/add-preview", date="2026-10-15", associate="Associate 001", what="Late", **{"from": "22:00"})
        self.assertEqual(got.status_code, 400)
        self.assertIn("outside", got.get_json()["error"])

    def test_undo_status_and_break(self):
        from datetime import date
        days, on = self.app.extensions["days"], date(2026, 10, 15)
        days.set_status("AE/AR B2B", on, "Associate 001", "Sick", 1)
        got = self.post("/day/undo", date="2026-10-15", associate="Associate 001", what="status", view="board")
        self.assertEqual(got.status_code, 303)
        self.assertEqual([r["status"] for r in self.store.list_attendance("AE/AR B2B", ["2026-10-15"])], ["Present"])
        seg = next(x for l in days.page("AE/AR B2B", on)["view"]["lanes"] if l["name"] == "Associate 001"
                   for x in l["segments"] if x["offset"] == 0)
        b = seg["breaks"][0]
        days.move_break("AE/AR B2B", on, "Associate 001", b["idx"], hm_(b["start"] + 15), 1)
        self.post("/day/undo", date="2026-10-15", associate="Associate 001", what="break", idx=str(b["idx"]))
        self.assertEqual(self.store.list_actual_breaks("AE/AR B2B", ["2026-10-15"]), [])

    def test_board_rows_offer_add(self):
        """Phase R task 5: every interval of the board has "+ Add"."""
        body = html.unescape(self.client.get(self.url + "&view=board").get_data(as_text=True))
        # Phase XY (owner-approved review finding 12): each + Add also names its interval for screen readers.
        self.assertRegex(body, r'<a class="rb-add" href="[^"]*add=600[^"]*#add">\+ Add<span class="sr-only"> to 10:00</span></a>')

    def test_add_dialog_lists_kinds_and_prefills_the_interval(self):
        body = html.unescape(self.client.get(self.url + "&view=board&add=600").get_data(as_text=True))
        dialog = body[body.index('<dialog id="add-dialog"'):]
        dialog = dialog[:dialog.index("</dialog>")]
        # re-pinned (Phase AC, review finding 10): the title read "Add to 10:00 to 11:00"; it now reads "Add for ..."
        self.assertIn("Add for 10:00 to", dialog)
        self.assertIn('action="/day/add"', dialog)
        for kind in ("Break", "Lunch", "Coaching", "Meeting", "Training", "System issue", "Unplanned leave", "Sick",
                     "Late", "Left early", "Overtime", "VTO"):
            self.assertIn(f'name="what" value="{kind}"', dialog)
        self.assertIn('name="from" value="10:00"', dialog)
        self.assertIn('<input type="hidden" name="at" value="600">', dialog)
        self.assertRegex(dialog, r'<option value="Associate \d+"')

    def test_add_dialog_calls_in_someone_off_on_typed_times(self):
        # owner, 2026-10-09: "cancel day off I'm not able to choose the shift start end manually or it's not
        # existing in RTA view"
        body = html.unescape(self.client.get(self.url + "&view=board&add=600").get_data(as_text=True))
        dialog = body[body.index('<dialog id="add-dialog"'):]
        dialog = dialog[:dialog.index("</dialog>")]
        self.assertRegex(dialog, r'<input type="radio" name="what" value="Day off cancelled">')
        off = re.search(r'<select name="associate" data-off[^>]*>(.*?)</select>', dialog, re.S)
        self.assertIsNotNone(off)
        listed = re.findall(r'<option value="(Associate \d+)">', off.group(1))
        self.assertTrue(listed)
        self.assertNotIn("Associate 001", listed)  # works that day (12:00 - 21:00)
        self.assertRegex(dialog, r'(?s)<select name="shift" data-shift-pick>.*?data-from="12:00" data-to="21:00"')
        self.assertRegex(dialog, r'<input type="time" name="to"')
        form = {"csrf_token": token(self.client), "program": "AE/AR B2B", "date": "2026-10-14",
                "what": "Day off cancelled", "from": "10:30", "to": "16:30"}
        who = next(n for n in listed  # the first one off whose rest gap allows 10:30 - 16:30
                   if self.client.post("/day/add-preview", data={**form, "associate": n}).status_code == 200)
        got = self.client.post("/day/add", data={**form, "associate": who, "view": "board", "at": "600",
                                                 "minutes": "15"}, follow_redirects=True)
        self.assertIn(f"Recorded: {who}, Day off cancelled: called in 10:30 - 16:30.",
                      html.unescape(got.get_data(as_text=True)))
        again = self.client.post("/day/add-preview", data={**form, "associate": who})
        self.assertEqual(again.status_code, 400)  # already called in: said, never recorded twice
        self.assertIn("already called in", again.get_json()["error"])

    def test_person_dialog_lists_records_with_delete(self):
        from datetime import date
        days, on = self.app.extensions["days"], date(2026, 10, 16)
        made = days.add_activity("AE/AR B2B", on, "Associate 001", "Coaching", "15:00", "15:30", 1, billable=True)
        days.set_status("AE/AR B2B", on, "Associate 001", "Late", 1, end="12:20")
        try:
            body = html.unescape(self.client.get("/day?program=AE/AR+B2B&date=2026-10-16&view=board&who=Associate+001")
                                 .get_data(as_text=True))
            dialog = body[body.index('<dialog id="person-dialog"'):]
            dialog = dialog[:dialog.index("</dialog>")]
            self.assertIn("<h2", dialog)
            self.assertIn("Associate 001", dialog)
            self.assertIn("Coaching 15:00 to 15:30, billable", dialog)
            self.assertRegex(dialog, rf'(?s)action="/day/activity/cancel".*?name="id" value="{made}".*?>Delete</button>')
            self.assertIn("Late, arrived 12:20", dialog)
            self.assertRegex(dialog, r'(?s)action="/day/undo".*?name="what" value="status".*?>Set back to present</button>')
            self.assertIn('select class="att', dialog)  # the attendance list, as on the timeline
            self.assertRegex(dialog, r'href="[^"]*add=[^"]*person=Associate(\+|%20)001[^"]*#add">\+ Add something</a>')
        finally:
            days.cancel_activity("AE/AR B2B", on, made, 1)
            days.set_status("AE/AR B2B", on, "Associate 001", "Present", 1)

    def test_recorded_list_says_delete(self):
        from datetime import date
        days, on = self.app.extensions["days"], date(2026, 10, 16)
        made = days.add_activity("AE/AR B2B", on, "Associate 001", "Overtime", "21:00", "22:00", 1)
        try:
            body = html.unescape(self.client.get("/day?program=AE/AR+B2B&date=2026-10-16&view=cover")
                                 .get_data(as_text=True))
            self.assertRegex(body, rf'(?s)name="id" value="{made}".*?class="link danger">Delete</button>')
            self.assertNotIn('class="link">Cancel</button>', body)
        finally:
            days.cancel_activity("AE/AR B2B", on, made, 1)

    def test_cover_panel_offers_overtime_lengths(self):
        """Phase R task 6: overtime offers come with a length to pick (15 minutes to the most the rules allow)."""
        body = html.unescape(self.client.get("/day?program=AE/AR+B2B&date=2026-10-16&view=board&cover=1260")
                             .get_data(as_text=True))
        panel = body[body.index('id="cover"'):]
        forms = [f for f in re.findall(r'(?s)<form[^>]*class="ot-offer"[^>]*>.*?</form>', panel) if "Associate 001" in f]
        self.assertEqual(len(forms), 1, "Associate 001 (12:00 - 21:00, off on Saturday) can stay on")
        options = re.findall(r'<option value="([0-9:|]+)"( selected)?>', forms[0])
        self.assertEqual(options[0][0], "21:00|21:15")
        self.assertEqual(options[-1][0], "21:00|23:00")
        self.assertIn(("21:00|22:00", " selected"), options)

    def test_record_overtime_with_span(self):
        got = self.post("/day/activity", date="2026-10-16", associate="Associate 001", kind="Overtime",
                        span="21:00|21:45", view="board")
        self.assertEqual(got.status_code, 303)
        acts = self.store.list_activities("AE/AR B2B", ["2026-10-16"])
        try:
            self.assertEqual([(a["kind"], a["start"], a["end_min"]) for a in acts], [("Overtime", 1260, 1305)])
        finally:
            from datetime import date
            for a in acts:
                self.app.extensions["days"].cancel_activity("AE/AR B2B", date(2026, 10, 16), a["id"], 1)

    def test_vto_offers_pick_from_and_to(self):
        body = html.unescape(self.client.get(self.url + "&view=cover").get_data(as_text=True))
        vto = body[body.index('id="vto-h"'):]
        form = re.search(r'(?s)<form[^>]*class="vto-offer"[^>]*>.*?</form>', vto)
        self.assertIsNotNone(form)
        self.assertRegex(form.group(0), r'<select name="from"')
        self.assertRegex(form.group(0), r'<select name="to"')
        self.assertRegex(form.group(0), r'>Add VTO</button>')

    def test_overview_and_rta_are_separate(self):
        body = html.unescape(self.client.get("/overview?program=AE/AR+B2B&date=2026-10-14").get_data(as_text=True))
        for words in ("On shift today", "Short of demand", "Next six hours", "Needs attention", "Open RTA",
                      "Handover note"):
            self.assertIn(words, body)
        self.assertIn("/day?program=AE/AR+B2B&date=2026-10-14", body)
        # Re-pinned in Phase AA (2026-10-10): the owner approved sample 03, so the menu's Schedules is a page of its
        # own (a week picker and the week's schedules side by side) that opens each schedule, instead of a redirect.
        got = self.client.get("/schedules?program=AE/AR+B2B")
        self.assertEqual(got.status_code, 200)
        self.assertRegex(got.get_data(as_text=True), r'href="/runs/[0-9a-f]{12}/schedules">Open</a>')

    def test_day_page_shows_lanes_rows_and_tiles(self):
        body = self.page()
        self.assertIn("AE/AR B2B, Wednesday 14 Oct: the day", body)
        self.assertIn("No version is marked in use for this week", body)
        self.assertIn('data-name="Associate 001" data-date="2026-10-14"', body)
        for row in ("Needed on the floor", "Planned on the floor", "On the floor now", "Plus / minus (hours)",
                    "English on the floor", "On shift today", "Short of demand"):
            self.assertIn(row, body)
        # Re-pinned in Phase Q (owner, 2026-10-08): the top bar's "Today" became the left menu's RTA, for the
        # program open.
        self.assertIn('href="/day?program=AE/AR+B2B&date=2026-10-14" class="on" aria-current="page">RTA</a>', body)

    def test_status_kept_or_refused_with_a_reason(self):
        refused = self.post("/day/attendance", associate="Associate 001", status="Late")
        self.assertEqual((refused.status_code, refused.get_json()["error"]), (400, "Give the time they arrived."))
        refused = self.post("/day/attendance", associate="Associate 001", status="Late", to="22:00")
        self.assertIn("outside Associate 001's shift", refused.get_json()["error"])
        self.assertEqual(self.post("/day/attendance", associate="Associate 001", status="Late", to="13:00").get_json(),
                         {"ok": True})
        body = self.page()
        self.assertRegex(body, r'<option value="Late" selected>Late</option>')
        self.assertIn("Late, arrived 13:00", body)  # the day's log

    def test_break_moved_or_refused(self):
        refused = self.post("/day/break", associate="Associate 008", idx="1", at="13:07")
        self.assertEqual(refused.status_code, 400)
        self.assertIn("5-minute steps", refused.get_json()["error"])
        self.assertEqual(self.post("/day/break", associate="Associate 008", idx="1", at="13:05").get_json(), {"ok": True})
        body = self.page()
        self.assertRegex(body, r'class="brk moved"[^>]*data-name="Associate 008"[^>]*data-idx="1"')
        self.assertIn("Lunch moved 12:30 to 13:05", body)
        self.assertEqual(self.post("/day/break", associate="Associate 008", idx="1", at="").get_json(), {"ok": True})
        self.assertIn("Lunch back to plan (12:30)", self.page())

    def test_bad_requests_answer_plainly(self):
        self.assertEqual(self.post("/day/break", associate="Associate 008", idx="x", at="13:05").status_code, 400)
        self.assertEqual(self.post("/day/attendance", associate="Nobody", status="Sick").status_code, 400)
        got = self.client.post("/day/attendance", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                        "date": "14-10-2026", "associate": "Associate 001",
                                                        "status": "Sick"})
        self.assertEqual((got.status_code, got.get_json()["error"]), (400, "Pick a day."))

    def test_measure_and_billable_aux(self):
        # re-pinned (Phase T): an aux is kept with who it is with and why (owner, 2026-10-09)
        self.assertEqual(self.post("/day/attendance", associate="Associate 012", status="Coaching", **{"from": "10:00"},
                                   to="11:00", billable="1", with_whom="Sara", why="Call quality review").get_json(),
                         {"ok": True})
        body = self.page()
        self.assertIn("Coaching (billable) 10:00 to 11:00", body)
        self.assertRegex(body, r'<option value="interval" selected>Interval compliance</option>')
        self.assertIn('<option value="sl">Service level</option>', body)
        sl = html.unescape(self.client.get(self.url + "&measure=sl").get_data(as_text=True))
        self.assertRegex(sl, r'<option value="sl" selected>Service level</option>')
        self.assertEqual(self.client.get(self.url + "&measure=weekly").status_code, 200)  # unknown: the default

    def test_slot_ids_and_the_interval_board(self):
        body = self.page()
        self.assertIn("Slot 1", body)
        board = html.unescape(self.client.get(self.url + "&view=board").get_data(as_text=True))
        for column in ("Interval", "Needed", "On the floor", "Buffer (h:mm)", "Status", "Lunch", "Short breaks",
                       "Late, early, aux", "Unplanned leave", "Languages"):
            self.assertIn(column, board)
        self.assertRegex(board, r'class="chip[^"]*"[^>]*data-name="Associate 008"[^>]*data-idx="1"')

    def test_a_broken_day_says_why(self):
        from unittest import mock
        with mock.patch.object(self.app.extensions["days"], "page", side_effect=ValueError("The input is missing.")):
            got = self.client.get(self.url)
        self.assertEqual(got.status_code, 200)
        self.assertIn("The input is missing.", got.get_data(as_text=True))

    def test_adherence_tab(self):
        self.post("/day/attendance", associate="Associate 021", status="Late", to="08:40")
        body = html.unescape(self.client.get(self.url + "&view=adherence").get_data(as_text=True))
        for words in ("Team adherence", "Team conformance", "Late and early leave", "By person (lowest first)",
                      "Actual shrinkage against your input"):
            self.assertIn(words, body)
        self.assertRegex(body, r"Associate 021</b></td><td>08:00 - 17:00</td><td>Late</td>")
        self.assertIn("Adherence", body)  # the tab

    def test_find_a_time_for_two_people_opens_no_person_dialog(self):
        # owner, 2026-10-09: "if i choosed more than 1 associate the pop up is showing ... 1 associate only": the
        # person dialog read the first of Find a time's people (both use "who") and covered the list of times
        url = self.url + "&view=meeting&who=Associate+001&who=Associate+012&minutes=30&from=13:00&to=17:00"
        body = html.unescape(self.client.get(url).get_data(as_text=True))
        self.assertNotIn('id="person-dialog"', body)
        self.assertEqual(len(re.findall(r'name="who" value="Associate 0(?:01|12)" checked', body)), 2)
        self.assertTrue(re.findall(r'name="start" value="(\d+)"', body))

    def test_every_aux_route_asks_who_and_why(self):
        # owner, 2026-10-09: "in case of any aux being placed like meeting coaching etc we need to specify with who
        # and why in a comment while reserving"; refused before anything is kept, on every way in
        def kept():
            return (len(self.store.list_activities("AE/AR B2B", ["2026-10-16"])),
                    len(self.store.list_attendance("AE/AR B2B", ["2026-10-16"])))
        def said(got):  # the message on the page the form returns to
            return html.unescape(self.client.get(got.headers["Location"]).get_data(as_text=True))

        before = kept()
        got = self.post("/day/add", date="2026-10-16", associate="Associate 001", what="Meeting", **{"from": "15:00"},
                        minutes="30", view="board", at="900")
        self.assertIn("Say who the meeting is with.", said(got))
        got = self.post("/day/activity", date="2026-10-16", associate="Associate 001", kind="Coaching",
                        **{"from": "15:00", "to": "15:30"}, with_whom="Sara")
        self.assertIn("Say why: a short reason for the coaching.", said(got))
        got = self.post("/day/attendance", date="2026-10-16", associate="Associate 001", status="Training",
                        billable="1", **{"from": "15:00", "to": "16:00"})
        self.assertEqual((got.status_code, got.get_json()["error"]), (400, "Say who the training is with."))
        got = self.post("/day/book", date="2026-10-16", who=["Associate 001"], start="900", minutes="30",
                        kind="Meeting", why="Process update")
        self.assertIn("Say who the meeting is with.", said(got))
        self.assertEqual(kept(), before)
        preview = self.post("/day/add-preview", date="2026-10-16", associate="Associate 001", what="Meeting",
                            **{"from": "15:00"}, minutes="30")
        self.assertEqual(preview.status_code, 200)  # the floor preview needs neither answer yet

    def test_an_aux_status_from_the_timeline_keeps_who_and_why(self):
        got = self.post("/day/attendance", date="2026-10-16", associate="Associate 012", status="System issue",
                        **{"from": "18:00", "to": "18:30"}, with_whom="IT ticket 4411", why="Headset not working")
        self.assertEqual(got.status_code, 200)
        try:
            row = next(r for r in self.store.list_attendance("AE/AR B2B", ["2026-10-16"]) if r["associate"] == "Associate 012")
            self.assertEqual((row["status"], row["with_whom"], row["why"]),
                             ("System issue", "IT ticket 4411", "Headset not working"))
            body = html.unescape(self.client.get("/day?program=AE/AR+B2B&date=2026-10-16&view=board&who=Associate+012")
                                 .get_data(as_text=True))
            dialog = body[body.index('<dialog id="person-dialog"'):]
            self.assertIn("with IT ticket 4411: Headset not working", dialog[:dialog.index("</dialog>")])
        finally:
            self.post("/day/undo", date="2026-10-16", associate="Associate 012", what="status")

    def test_the_forms_ask_who_and_why(self):
        body = html.unescape(self.client.get(self.url + "&view=board&add=600").get_data(as_text=True))
        dialog = body[body.index('<dialog id="add-dialog"'):]
        dialog = dialog[:dialog.index("</dialog>")]
        self.assertRegex(dialog, r'<input name="with_whom" list="team-names" maxlength="80"')
        self.assertRegex(dialog, r'<input name="why" maxlength="200"')
        self.assertRegex(body, r'(?s)<datalist id="team-names">.*?<option value="Sara">')  # the team, or type a name
        att = body[body.index('<dialog id="att-dialog"'):]
        self.assertRegex(att[:att.index("</dialog>")], r'name="with_whom".*name="why"')
        finder = html.unescape(self.client.get(self.url + "&view=meeting&who=Associate+001&minutes=30&from=13:00"
                                               "&to=17:00&with_whom=Ops+manager&why=Process+update")
                               .get_data(as_text=True))
        self.assertRegex(finder, r'<input name="with_whom" list="team-names" maxlength="80" required value="Ops manager"')
        self.assertRegex(finder, r'<input type="hidden" name="with_whom" value="Ops manager"><input type="hidden" '
                                 r'name="why" value="Process update">')  # carried to Book

    def test_find_a_time_and_book_it(self):
        url = (self.url + "&view=meeting&who=Associate+001&who=Associate+012&minutes=30&from=13:00&to=17:00"
               "&kind=Training&billable=1")
        body = html.unescape(self.client.get(url).get_data(as_text=True))
        self.assertIn("Find a time", body)
        starts = re.findall(r'name="start" value="(\d+)"', body)
        self.assertTrue(starts)
        got = self.client.post("/day/book", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                  "date": "2026-10-14", "who": ["Associate 001", "Associate 012"],
                                                  "start": starts[0], "minutes": "30", "kind": "Training", "billable": "1",
                                                  # re-pinned (Phase T): who it is with and why (owner, 2026-10-09)
                                                  "with_whom": "IT trainer", "why": "New CRM release"})
        self.assertEqual(got.status_code, 303)
        page = self.page()
        self.assertRegex(page, r'class="tl-act aux billable"')
        self.assertIn("Associate 012, Training", page)
        again = self.client.post("/day/book", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                    "date": "2026-10-14", "who": ["Associate 001"], "start": starts[0],
                                                    "minutes": "30", "kind": "Meeting", "with_whom": "Ops manager",
                                                    "why": "Process update"}, follow_redirects=True)
        self.assertIn("Not free then: Associate 001. Nobody was booked.", html.unescape(again.get_data(as_text=True)))

    def test_overtime_and_vto_tab(self):
        body = html.unescape(self.client.get(self.url + "&view=cover").get_data(as_text=True))
        for words in ("Short: offer overtime", "Above demand: offer voluntary time off", "Activities on this day"):
            self.assertIn(words, body)
        got = self.client.post("/day/activity", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                      "date": "2026-10-14", "associate": "Associate 028",
                                                      "kind": "Overtime", "from": "07:00", "to": "09:00", "view": "cover"})
        self.assertEqual(got.status_code, 303)
        self.assertIn("view=cover", got.headers["Location"])
        body = html.unescape(self.client.get(self.url + "&view=cover").get_data(as_text=True))
        self.assertIn("Associate 028, Overtime 07:00 to 09:00", body)
        made = re.search(r'Associate 028, Overtime 07:00 to 09:00</b>\s*<form[^>]*>(?:<input[^>]*>)*?<input type="hidden" '
                         r'name="id" value="(\d+)"', body).group(1)
        self.client.post("/day/activity/cancel", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                       "date": "2026-10-14", "id": made, "view": "cover"})
        self.assertIn("Cancelled: Overtime 07:00 to 09:00", self.page())
        refused = self.client.post("/day/activity", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                          "date": "2026-10-14", "associate": "Associate 028",
                                                          "kind": "Overtime", "from": "10:00", "to": "11:00",
                                                          "view": "cover"}, follow_redirects=True)
        self.assertIn("Overtime has to start when", html.unescape(refused.get_data(as_text=True)))

    def test_autopilot_preview_and_apply(self):
        body = html.unescape(self.client.get(self.url + "&view=replan").get_data(as_text=True))
        for words in ("Fix the rest of the day's breaks", "Tightest interval", "Hours short"):
            self.assertIn(words, body)
        got = self.client.post("/day/replan/apply", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                          "date": "2026-10-14", "move": ["Associate 021|1|775"]})
        self.assertEqual(got.status_code, 303)
        self.assertIn("Lunch moved 12:30 to 12:55 (autopilot)", self.page())

    def test_handover_note_page(self):
        body = html.unescape(self.client.get("/day/handover?program=AE/AR+B2B&date=2026-10-14").get_data(as_text=True))
        for words in ("Handover: AE/AR B2B, Wednesday 14 Oct", "In numbers", "Who was not on the floor",
                      "Breaks moved", "Watch tomorrow (Thursday 15 Oct)", "Print"):
            self.assertIn(words, body)
        self.assertIn('href="/day/handover?program=AE/AR+B2B&amp;date=2026-10-14"',
                      self.client.get(self.url).get_data(as_text=True))

    def test_wallboard(self):
        body = html.unescape(self.client.get("/day/wallboard?program=AE/AR+B2B&date=2026-10-14&at=12:40").get_data(as_text=True))
        self.assertIn("data-wallboard", body)
        for words in ("AE/AR B2B", "12:40", "Now", "12:00 to 13:00", "13:00 to 14:00", "15:00 to 16:00",
                      "On break now", "Next 30 minutes", "English"):
            self.assertIn(words, body)
        self.assertNotIn("16:00 to 17:00", body)  # now and the next three only
        self.assertIn("Associate 021", body)  # lunch 12:30 to 13:00: on break at 12:40
        self.assertEqual(self.app.test_client().get("/day/wallboard?program=AE/AR+B2B").status_code, 302)

    def test_add_cover_from_the_board(self):
        body = html.unescape(self.client.get(self.url + "&view=board&cover=840").get_data(as_text=True))
        for words in ("Cover 14:00 to 15:00", "Overtime", "Day off cancelled", "Add cover"):
            self.assertIn(words, body)
        offered = re.search(r"Associate 002: (\d\d:\d\d) - (\d\d:\d\d)", body)  # off on Wednesday
        self.assertIsNotNone(offered)
        lo, hi = offered.groups()
        self.assertTrue(lo <= "14:00" and "15:00" <= hi)
        got = self.client.post("/day/activity", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                      "date": "2026-10-14", "associate": "Associate 002",
                                                      "kind": "Called in", "from": lo, "to": hi,
                                                      "view": "board", "cover": "840"})
        self.assertEqual(got.status_code, 303)
        self.assertIn("cover=840", got.headers["Location"])
        page = self.page()
        self.assertIn(f"{lo} - {hi} (called in)", page)
        self.assertIn(f"Day off cancelled: called in {lo} - {hi}", page)

    def test_saturday_says_when_sunday_is_not_checked(self):
        note = "Next week's schedule is not here yet, so the rest gap to Sunday is not checked"
        self.assertNotIn(note, html.unescape(self.client.get(self.url + "&view=board&cover=840").get_data(as_text=True)))
        saturday = "/day?program=AE/AR+B2B&date=2026-10-17"
        self.assertIn(note, html.unescape(self.client.get(saturday + "&view=board&cover=840").get_data(as_text=True)))
        self.assertIn(note, html.unescape(self.client.get(saturday + "&view=cover").get_data(as_text=True)))
        got = self.client.post("/day/activity", data={"csrf_token": token(self.client), "program": "AE/AR B2B",
                                                      "date": "2026-10-17", "associate": "Associate 002",
                                                      "kind": "Overtime", "from": "18:00", "to": "19:00",
                                                      "view": "board", "cover": "1080"}, follow_redirects=True)
        body = html.unescape(got.get_data(as_text=True))
        self.assertIn("Recorded: Associate 002, overtime.", body)
        self.assertIn(note, body)

    def test_week_without_a_schedule_and_defaults(self):
        body = html.unescape(self.client.get("/day?program=AE/AR+B2B&date=2026-11-04").get_data(as_text=True))
        self.assertIn("No schedule for AE/AR B2B in the week of 01 Nov", body)
        self.assertEqual(self.client.get("/day").status_code, 200)
        self.assertEqual(self.client.get("/day?date=not-a-date").status_code, 200)


class TheRecord(unittest.TestCase):
    """Phase O task 1: runs, downloads and people admin are recorded with who and when."""

    def events(self, store):
        return [(e["kind"], e["by_name"], e["subject"]) for e in store.list_events(0, time.time() + 60)]

    def test_upload_stop_resume_recorded(self):
        app, store, *_ = make_app(start_worker=False)
        client = client_for(app)
        run_id = run_id_of(upload(client, name="week42.xlsx"))
        client.post(f"/runs/{run_id}/stop", data={"csrf_token": token(client)})
        client.post(f"/runs/{run_id}/resume", data={"csrf_token": token(client)})
        self.assertEqual(self.events(store), [("run_uploaded", "Sara", "week42.xlsx"), ("run_stopped", "Sara", "week42.xlsx"),
                                              ("run_resumed", "Sara", "week42.xlsx")])
        self.assertEqual({e["run_id"] for e in store.list_events(0, time.time() + 60)}, {run_id})

    def test_downloads_recorded(self):
        app, store, *_ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client))
        wait(store, run_id)
        self.assertEqual(client.get(f"/runs/{run_id}/download").status_code, 200)
        kind, by, subject = self.events(store)[-1]
        self.assertEqual((kind, by), ("downloaded", "Sara"))
        self.assertTrue(subject.endswith(".zip"))

    def test_people_admin_recorded_without_passwords(self):
        app, store, *_ = make_app(start_worker=False)
        store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        client = app.test_client()
        tok = TOKEN.search(client.get("/login").get_data(as_text=True)).group(1)
        client.post("/login", data={"username": "omar", "password": "Owner-pass-123", "csrf_token": tok})
        client.post("/admin/users", data={"csrf_token": token(client), "username": "lina", "display_name": "Lina",
                                          "password": "Lina-temp-pass-1"})
        lina = next(u for u in store.list_users() if u["username"] == "lina")["id"]
        for action in ("disable", "enable", "reset"):
            client.post(f"/admin/users/{lina}/{action}", data={"csrf_token": token(client)})
        self.assertEqual(self.events(store), [("user_added", "Omar", "lina"), ("user_disabled", "Omar", "lina"),
                                              ("user_enabled", "Omar", "lina"), ("password_reset", "Omar", "lina")])
        everything = " ".join(str(v) for e in store.list_events(0, time.time() + 60) for v in e.values())
        self.assertNotIn("Lina-temp-pass-1", everything)

    def test_run_files_kept_30_to_60_days(self):
        from webapp import serve
        base = {"SCHEDULER_DATA_DIR": "/var/lib/scheduler"}
        self.assertEqual(serve.config_from_env(base)["RUN_FILES_DAYS"], 30)
        self.assertEqual(serve.config_from_env({**base, "SCHEDULER_RUN_FILES_DAYS": "60"})["RUN_FILES_DAYS"], 60)
        for bad in ("29", "61", "forever"):
            with self.assertRaises(SystemExit):
                serve.config_from_env({**base, "SCHEDULER_RUN_FILES_DAYS": bad})
        app, store, data, _ = make_app(RUN_FILES_DAYS=45)
        client = client_for(app)
        run_id = run_id_of(upload(client))
        wait(store, run_id)
        queue = app.extensions["runs"]
        self.assertEqual(queue.cleanup(now=time.time() + 44 * 86400), 0)
        self.assertEqual(queue.cleanup(now=time.time() + 46 * 86400), 1)


def hm_(minute):
    return f"{(minute // 60) % 24:02d}:{minute % 60:02d}"


def sign_in(app, username, password):
    client = app.test_client()
    tok = TOKEN.search(client.get("/login").get_data(as_text=True)).group(1)
    client.post("/login", data={"username": username, "password": password, "csrf_token": tok})
    return client


class TheRunDetails(unittest.TestCase):
    """Phase P: an admin corrects any run's details and renames programs; a submitter fixes their own
    run's program and week. Versions move with the run; consequences are shown before saving."""

    def setUp(self):
        self.app, self.store, *_ = make_app(start_worker=False)
        self.omar = self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        self.lina = self.store.add_user("lina", "Lina", "Lina-pass-12", must_change=False)
        self.sara = sign_in(self.app, "sara", "Sara-pass-1")
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")
        self.run_id = run_id_of(upload(self.sara, program="AE/AR B2B", week_start="2026-10-11"))

    def post(self, client, url=None, **data):
        return client.post(url or f"/runs/{self.run_id}/tag", data={"csrf_token": token(client), **data})

    def test_admin_edits_every_detail(self):
        page = html.unescape(self.admin.get(f"/runs/{self.run_id}").get_data(as_text=True))
        # re-pinned (Phase R): the program is a select of the programs and LOBs set up, the run's own selected
        for words in ("Edit details", 'name="user_id"', 'name="workbook"', '<option value="AE/AR B2B" selected>'):
            self.assertIn(words, page)
        got = self.post(self.admin, program="AE/AR B2B", week_start="2026-10-11", user_id=str(self.lina),
                        workbook="AR week 42.xlsx", reason="Lina built it")
        self.assertEqual(got.status_code, 302)
        page = html.unescape(self.admin.get(f"/runs/{self.run_id}").get_data(as_text=True))
        self.assertIn("Details saved.", page)
        self.assertIn("<h1>AR week 42.xlsx</h1>", page)
        self.assertIn("started by Lina", page)
        from datetime import datetime
        from openpyxl import load_workbook
        from webapp.exports import build
        from webapp.schedules import EGYPT
        today = datetime.now(EGYPT).date()  # exports count days in Egypt time; the server clock is UTC (the test
        # read the UTC date and failed every day from 00:00 to 03:00 Egypt time)
        data, *_ = build(self.store, self.app.extensions["days"], today, today, ["record"], by="Omar")
        rows = list(load_workbook(io.BytesIO(data))["Other actions"].iter_rows(values_only=True))
        changed = [dict(zip(rows[0], r)) for r in rows[1:] if "Run details changed" in r]
        self.assertEqual(len(changed), 1)
        self.assertIn("Submitted by: Sara → Lina", " ".join(str(v) for v in changed[0].values()))

    def test_submitter_edits_program_and_week_only(self):
        from webapp.programs import ProgramBook
        ProgramBook(self.store).add_program("NMG")  # re-pinned (Phase R): programs are picked, never typed
        page = html.unescape(self.sara.get(f"/runs/{self.run_id}").get_data(as_text=True))
        self.assertNotIn('name="user_id"', page)
        self.assertEqual(self.post(self.sara, program="NMG", week_start="2026-10-18").status_code, 302)
        self.assertEqual((self.store.get_run(self.run_id)["program"], self.store.get_run(self.run_id)["week_start"]),
                         ("NMG", "2026-10-18"))
        self.assertEqual(self.post(self.sara, program="NMG", week_start="2026-10-18", workbook="x.xlsx").status_code, 403)
        self.assertEqual(self.post(self.sara, program="NMG", week_start="2026-10-18",
                                   user_id=str(self.lina)).status_code, 403)
        self.assertEqual(self.store.get_run(self.run_id)["workbook"], "week42.xlsx")

    def test_program_rename_page(self):
        # re-pinned (Phase R): a program's data moves into a program or LOB that is set up, picked from a
        # list; a typed new name is refused (owner, 2026-10-08: typing made stray programs). Display names
        # change under LOBs and defaults.
        from webapp.programs import ProgramBook
        book = ProgramBook(self.store)
        target = book.add_lob(book.add_program("AE"), "AR B2B")  # key "AE AR B2B"
        self.assertEqual(self.post(self.sara, "/programs/rename", old="AE/AR B2B", new=target).status_code, 403)
        page = html.unescape(self.post(self.admin, "/programs/rename", old="AE/AR B2B", new=target)
                             .get_data(as_text=True))
        for words in ("Move AE/AR B2B into AE, AR B2B", "1 run", "Give a reason"):
            self.assertIn(words, page)
        self.assertEqual(self.store.get_run(self.run_id)["program"], "AE/AR B2B")  # not yet
        got = self.post(self.admin, "/programs/rename", old="AE/AR B2B", new=target, confirm="1",
                        reason="into its LOB")
        self.assertEqual(got.status_code, 302)
        self.assertEqual(self.store.get_run(self.run_id)["program"], target)


class ThePickedProgram(unittest.TestCase):
    """Phase R task 1: a run's program and a program's move target are picked from the programs and LOBs
    set up under LOBs and defaults, never typed (owner, 2026-10-08: typing made stray programs)."""

    def setUp(self):
        from webapp.programs import ProgramBook
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        book = ProgramBook(self.store)
        saks = book.add_program("SAKS")
        book.add_lob(saks, "NMG Tier 1")
        self.tier2 = book.add_lob(saks, "NMG Tier 2")  # key "SAKS NMG Tier 2"
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")
        self.run_id = seed_week(self.store, self.tier2, "2026-10-11")

    def post(self, url, **data):
        return self.admin.post(url, data={"csrf_token": token(self.admin), **data}, follow_redirects=True)

    def test_edit_details_offers_program_and_lob_select(self):
        page = html.unescape(self.admin.get(f"/runs/{self.run_id}").get_data(as_text=True))
        self.assertIn('<select name="program"', page)
        self.assertIn('<optgroup label="SAKS">', page)
        self.assertRegex(page, r'<option value="SAKS NMG Tier 2" selected>SAKS, NMG Tier 2</option>')
        self.assertNotIn('<input name="program"', page)

    def test_typed_program_is_refused(self):
        page = html.unescape(self.post(f"/runs/{self.run_id}/tag", program="SAKS, NMG Tier 2",
                                       week_start="2026-10-11", reason="typed").get_data(as_text=True))
        self.assertIn("Pick a program and LOB set up under LOBs and defaults.", page)
        self.assertEqual(self.store.get_run(self.run_id)["program"], "SAKS NMG Tier 2")
        self.assertNotIn("SAKS, NMG Tier 2", [p["name"] for p in self.store.list_programs()])

    def test_picked_program_is_saved(self):
        self.post(f"/runs/{self.run_id}/tag", program="SAKS NMG Tier 1", week_start="2026-10-11", reason="moved")
        self.assertEqual(self.store.get_run(self.run_id)["program"], "SAKS NMG Tier 1")

    def test_program_page_moves_into_a_picked_unit(self):
        page = html.unescape(self.admin.get("/programs/SAKS%20NMG%20Tier%202").get_data(as_text=True))
        self.assertIn("Move everything into", page)
        self.assertRegex(page, r'(?s)<select name="new"[^>]*>.*<option value="SAKS NMG Tier 1">SAKS, NMG Tier 1</option>')
        self.assertNotIn('<input name="new"', page)
        page = html.unescape(self.post("/programs/rename", old="SAKS NMG Tier 2", new="SAKS Tier 9").get_data(as_text=True))
        self.assertIn("Pick a program and LOB set up under LOBs and defaults.", page)
        self.assertEqual(self.store.get_run(self.run_id)["program"], "SAKS NMG Tier 2")
        page = html.unescape(self.post("/programs/rename", old="SAKS NMG Tier 2", new="SAKS NMG Tier 1").get_data(as_text=True))
        self.assertIn("Move SAKS, NMG Tier 2 into SAKS, NMG Tier 1", page)


class ThePeoplePage(unittest.TestCase):
    """Phase Q: roles and programs. Admins manage everyone; supervisors the planners of their own programs
    (add, programs, reset password, switch on or off); planners nobody."""

    def setUp(self):
        from webapp.programs import ProgramBook
        self.app, self.store, *_ = make_app(start_worker=False)
        self.omar = self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        book = ProgramBook(self.store)
        self.ae, self.nmg = book.add_program("AE"), book.add_program("NMG")
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")

    def post(self, client, url, **data):
        return client.post(url, data={"csrf_token": token(client), **data})

    def person(self, username):
        return next(u for u in self.store.list_users() if u["username"] == username)

    def supervisor(self):
        sid = self.store.add_user("sami", "Sami", "Sami-pass-123", must_change=False)
        self.store.update_user(sid, is_supervisor=1)
        self.store.set_user_programs(sid, [self.ae])
        return sid, sign_in(self.app, "sami", "Sami-pass-123")

    def test_admin_adds_a_supervisor_with_programs(self):
        got = self.admin.post("/admin/users", data={"csrf_token": token(self.admin), "username": "rana",
                                                    "display_name": "Rana", "password": "Rana-temp-pass-1",
                                                    "role": "supervisor", "programs": [str(self.ae)]})
        self.assertEqual(got.status_code, 302)
        rana = self.person("rana")
        self.assertEqual((rana["is_supervisor"], rana["is_admin"], rana["all_programs"]), (1, 0, 0))
        self.assertEqual(self.store.user_program_ids(rana["id"]), [self.ae])
        page = html.unescape(self.admin.get("/admin/users").get_data(as_text=True))
        self.assertIn("Supervisor", page)
        self.assertRegex(page, r"(?s)Rana</b>.*?Supervisor.*?<td>AE</td>")  # name, role, programs on one row

    def test_supervisor_manages_planners_of_their_programs(self):
        sid, sami = self.supervisor()
        self.assertEqual(sami.get("/admin/users").status_code, 200)
        self.post(sami, "/admin/users", username="lina", display_name="Lina", password="Lina-temp-pass-1",
                  role="planner", programs=[str(self.ae)])
        lina = self.person("lina")
        self.assertEqual(self.store.user_program_ids(lina["id"]), [self.ae])
        page = html.unescape(self.post(sami, "/admin/users", username="hadi", display_name="Hadi",
                                       password="Hadi-temp-pass-1", role="planner",
                                       programs=[str(self.nmg)]).get_data(as_text=True))
        self.assertIn("You can only give your own programs.", page)
        self.assertFalse(any(u["username"] == "hadi" for u in self.store.list_users()))
        page = html.unescape(self.post(sami, "/admin/users", username="hadi", display_name="Hadi",
                                       password="Hadi-temp-pass-1", role="admin",
                                       programs=[str(self.ae)]).get_data(as_text=True))
        self.assertIn("Only an admin adds supervisors and admins.", page)
        got = self.post(sami, f"/admin/users/{lina['id']}/reset")
        self.assertEqual(got.status_code, 302)
        self.assertIn("Temporary password for lina", html.unescape(sami.get("/admin/users").get_data(as_text=True)))
        self.assertEqual(self.post(sami, f"/admin/users/{self.omar}/reset").status_code, 403)
        self.assertEqual(self.post(sami, f"/admin/users/{lina['id']}/programs", programs=[str(self.nmg)]).status_code,
                         403)
        self.assertEqual(self.store.user_program_ids(lina["id"]), [self.ae])
        self.assertEqual(self.post(sami, f"/admin/users/{lina['id']}/disable").status_code, 302)
        self.assertFalse(self.person("lina")["active"])

    def test_admin_changes_programs_and_role(self):
        lina = self.store.add_user("lina", "Lina", "Lina-pass-123", must_change=False)
        self.post(self.admin, f"/admin/users/{lina}/programs", programs=[str(self.ae), str(self.nmg)],
                  role="supervisor")
        self.assertEqual(self.store.user_program_ids(lina), [self.ae, self.nmg])
        self.assertEqual(self.store.get_user(lina)["is_supervisor"], 1)
        self.post(self.admin, f"/admin/users/{lina}/programs", all_programs="1", role="planner")
        self.assertEqual((self.store.get_user(lina)["all_programs"], self.store.get_user(lina)["is_supervisor"]), (1, 0))
        kinds = [(e["kind"], e["subject"]) for e in self.store.list_events(0, time.time() + 60)]
        self.assertIn(("user_programs_changed", "lina"), kinds)

    def test_planners_cannot_open_people(self):
        self.store.add_user("lina", "Lina", "Lina-pass-123", must_change=False)
        self.assertEqual(sign_in(self.app, "lina", "Lina-pass-123").get("/admin/users").status_code, 403)


class TheProgramSetup(unittest.TestCase):
    """Phase Q: an admin sets up programs, their LOBs and the defaults a new schedule starts from; the
    upload form offers each program's LOBs and fills in its defaults."""

    def setUp(self):
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        self.sara = client_for(self.app)
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")

    def post(self, **data):
        return self.admin.post("/setup/programs", data={"csrf_token": token(self.admin), **data})

    def program(self, name):
        from webapp.programs import ProgramBook
        return next(p for p in ProgramBook(self.store).tree() if p["name"] == name)

    def test_setup_page_offers_move_and_delete(self):
        """Phase R task 2: a stray program's data moves into its LOB, then it can be deleted; an empty LOB
        can be deleted; nothing with data under it can (owner, 2026-10-08: move first, then delete)."""
        from webapp.programs import ProgramBook
        book = ProgramBook(self.store)
        saks = book.add_program("SAKS")
        tier2 = book.add_lob(saks, "NMG Tier 2")
        book.add_lob(saks, "SAKS Tier 1")
        run_id_of(upload(self.sara, program="SAKS, NMG Tier 2", week_start="2026-10-18"))
        page = html.unescape(self.admin.get("/setup/programs").get_data(as_text=True))
        self.assertIn("Move all of it into", page)
        self.assertRegex(page, r'<button[^>]*value="delete_program"[^>]*disabled')  # the stray still has its run
        self.assertIn("still has 1 run", page)
        stray = self.program("SAKS, NMG Tier 2")
        tier1 = next(l for l in self.store.list_lobs() if l["key"] == "SAKS SAKS Tier 1")
        self.post(action="delete_lob", lob_id=str(tier1["id"]))
        self.assertNotIn("SAKS SAKS Tier 1", [l["key"] for l in self.store.list_lobs()])
        refused = self.admin.post("/setup/programs", data={"csrf_token": token(self.admin), "action": "delete_program",
                                                           "program_id": str(stray["id"])}, follow_redirects=True)
        self.assertIn("still has 1 run", html.unescape(refused.get_data(as_text=True)))
        self.admin.post("/programs/rename", data={"csrf_token": token(self.admin), "old": "SAKS, NMG Tier 2",
                                                  "new": tier2, "confirm": "1", "reason": "into its LOB"})
        self.post(action="delete_program", program_id=str(stray["id"]))
        self.assertNotIn("SAKS, NMG Tier 2", [p["name"] for p in ProgramBook(self.store).tree()])
        kinds = [e["kind"] for e in self.store.list_events(0, time.time() + 1)]
        self.assertIn("program_deleted", kinds)
        self.assertIn("lob_deleted", kinds)

    def test_a_stray_program_offers_its_lob_first(self):
        """Phase R: 'SAKS, NMG Tier 2' reads as SAKS's LOB NMG Tier 2, so that is the move offered first."""
        from webapp.programs import ProgramBook
        book = ProgramBook(self.store)
        book.add_program("AE")
        saks = book.add_program("SAKS")
        book.add_lob(saks, "NMG Tier 1")
        tier2 = book.add_lob(saks, "NMG Tier 2")
        run_id_of(upload(self.sara, program="SAKS, NMG Tier 2", week_start="2026-10-18"))
        page = html.unescape(self.admin.get("/setup/programs").get_data(as_text=True))
        self.assertIn(f'<option value="{tier2}" selected>SAKS, NMG Tier 2</option>', page)
        self.assertIn("SAKS, NMG Tier 2 looks like SAKS's LOB NMG Tier 2.", page)
        self.assertIn("SAKS, NMG Tier 2 still has 1 run: move it into a program or LOB first.", page)

    def test_set_up_a_program_and_use_it(self):
        self.assertEqual(self.sara.get("/setup/programs").status_code, 403)
        run_id = run_id_of(upload(self.sara, program="AE/AR B2B", week_start="2026-10-11"))
        self.assertEqual(self.post(action="add_program", name="AE").status_code, 302)
        ae = self.program("AE")["id"]
        self.post(action="add_lob", program_id=str(ae), name="IT")
        self.post(action="adopt", program_id=str(ae), key="AE/AR B2B", name="AR B2B")
        self.post(action="defaults", program_id=str(ae), start_day="1", run_mode="DEEP", stage="BEFORE_BREAKS_ONLY")
        found = self.program("AE")
        self.assertEqual([(l["name"], l["key"]) for l in found["lobs"]], [("AR B2B", "AE/AR B2B"), ("IT", "AE IT")])
        self.assertEqual((found["start_day"], found["run_mode"], found["options"]), (1, "DEEP", {"stage": "BEFORE_BREAKS_ONLY"}))
        page = html.unescape(self.admin.get("/setup/programs").get_data(as_text=True))
        for words in ("Programs and LOBs", "AR B2B", "AE IT"):
            self.assertIn(words, page)
        form = self.sara.get("/").get_data(as_text=True)  # raw: the defaults are JSON inside an attribute
        group = re.search(r'<optgroup label="AE" data-defaults="([^"]*)">(.*?)</optgroup>', form, re.S)
        self.assertIsNotNone(group)
        self.assertEqual(json.loads(html.unescape(group.group(1))), {"start_day": 1, "run_mode": "DEEP",
                                                                     "options": {"stage": "BEFORE_BREAKS_ONLY"}})
        # program and LOB, even closed; re-pinned (Phase S): the unit picked on the left (here the one Sara last
        # opened) is now preselected in the upload form (owner, 2026-10-09), so the option may carry "selected"
        self.assertRegex(group.group(2), r'<option value="AE/AR B2B"( selected)?>AE, AR B2B</option>')
        self.assertIn('<option value="AE IT">AE, IT</option>', group.group(2))
        other = run_id_of(upload(self.sara, program="AE IT", week_start="2026-10-12"))
        self.assertEqual(self.store.get_run(other)["program"], "AE IT")
        page = html.unescape(self.sara.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("AE, AR B2B", page)  # the run page names program and LOB

    def test_refusals_are_shown(self):
        self.post(action="add_program", name="AE")
        self.post(action="add_program", name="ae")
        self.assertIn("There is already a program called ae.",
                      html.unescape(self.admin.get("/setup/programs").get_data(as_text=True)))


class TheProgramAccess(unittest.TestCase):
    """Phase Q: people open only their programs' pages (403 otherwise); New schedule stays open to everyone
    for any program, and their own runs stay theirs."""

    def setUp(self):
        from webapp.programs import ProgramBook
        self.app, self.store, *_ = make_app(start_worker=False)
        self.sara = client_for(self.app)
        self.ae = run_id_of(upload(self.sara, name="AE_week.xlsx", program="AE/AR B2B", week_start="2026-10-11"))
        self.nmg = run_id_of(upload(self.sara, name="NMG_week.xlsx", program="NMG", week_start="2026-10-12"))
        book = ProgramBook(self.store)
        book.sync()
        nmg = next(p["id"] for p in book.tree() if p["name"] == "NMG")
        lina = self.store.add_user("lina", "Lina", "Lina-pass-123", must_change=False)
        self.store.set_user_programs(lina, [nmg])
        self.lina = sign_in(self.app, "lina", "Lina-pass-123")

    def test_other_programs_are_refused(self):
        for url in ("/day?program=AE/AR+B2B&date=2026-10-14", "/week?program=AE/AR+B2B&week=2026-10-11",
                    "/programs/AE/AR%20B2B", f"/runs/{self.ae}", f"/runs/{self.ae}/schedules",
                    f"/runs/{self.ae}/download", "/coach?program=AE/AR+B2B",
                    "/day/handover?program=AE/AR+B2B&date=2026-10-14", "/day/wallboard?program=AE/AR+B2B",
                    "/exports/download?program=AE/AR+B2B&from=2026-10-11&to=2026-10-17&kind=attendance"):
            self.assertEqual(self.lina.get(url).status_code, 403, url)
        got = self.lina.post("/day/attendance", data={"csrf_token": token(self.lina), "program": "AE/AR B2B",
                                                       "date": "2026-10-14", "associate": "Associate 001",
                                                       "status": "Sick"})
        self.assertEqual(got.status_code, 403)

    def test_own_programs_and_runs_open(self):
        self.assertEqual(self.lina.get("/day?program=NMG&date=2026-10-14").status_code, 200)
        self.assertEqual(self.lina.get(f"/runs/{self.nmg}").status_code, 200)
        home = html.unescape(self.lina.get("/").get_data(as_text=True))
        self.assertIn("NMG_week.xlsx", home)
        self.assertNotIn("AE_week.xlsx", home)
        exports = html.unescape(self.lina.get("/exports").get_data(as_text=True))
        self.assertNotIn('value="AE/AR B2B"', exports)
        mine = run_id_of(upload(self.lina, name="Lina_AE.xlsx", program="AE/AR B2B", week_start="2026-10-18"))
        self.assertEqual(self.lina.get(f"/runs/{mine}").status_code, 200)  # New schedule: any program, her run

    def test_all_programs_means_all_of_yours(self):
        self.store.add_activity(program="AE/AR B2B", shift_date="2026-10-14", associate="Associate 001",
                                kind="Training", start=600, end_min=660, user_id=1)
        self.store.add_activity(program="NMG", shift_date="2026-10-14", associate="Associate 002",
                                kind="Training", start=600, end_min=660, user_id=1)
        got = self.lina.get("/exports/download?from=2026-10-14&to=2026-10-14&kind=activities&format=csv")
        self.assertEqual(got.status_code, 200)
        body = got.get_data(as_text=True)
        self.assertIn("Associate 002", body)
        self.assertNotIn("Associate 001", body)
        from openpyxl import load_workbook
        got = self.lina.get("/exports/download?from=2026-10-14&to=2026-10-14&kind=activities")
        about = dict(r[:2] for r in load_workbook(io.BytesIO(got.data))["About this export"].iter_rows(values_only=True))
        self.assertEqual(about["Programs"], "NMG")  # all of hers


class TheNewLayout(unittest.TestCase):
    """Phase Q: a slim top bar (name, look, Settings, Sign out); every page in the left menu; Home starts
    from the person's programs."""

    def setUp(self):
        from webapp.programs import ProgramBook
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        seed_week(self.store, "NMG", "2026-10-11")
        seed_week(self.store, "AE/AR B2B", "2026-10-11")
        book = ProgramBook(self.store)
        book.sync()
        lina = self.store.add_user("lina", "Lina", "Lina-pass-123", must_change=False)
        self.store.set_user_programs(lina, [next(p["id"] for p in book.tree() if p["name"] == "NMG")])
        self.sara, self.lina = client_for(self.app), sign_in(self.app, "lina", "Lina-pass-123")
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")

    def part(self, body, tag, cls):
        found = re.search(rf'<{tag} class="{cls}"[^>]*>(.*?)</{tag}>', body, re.S)
        self.assertIsNotNone(found, f"{tag}.{cls}")
        return html.unescape(found.group(1))

    def test_top_bar_has_only_account_things(self):
        top = self.part(self.sara.get("/").get_data(as_text=True), "header", "top")
        self.assertIn("Sara", top)
        self.assertRegex(top, r"(Light|Dark) look")
        for words in (">Settings<", "Sign out"):
            self.assertIn(words, top)
        for word in ("Programs", "Weeks", "Exports", "Team", "People", "Today"):
            self.assertNotIn(f">{word}<", top)

    def test_left_menu_has_the_program_pages(self):
        side = self.part(self.sara.get("/day?program=NMG&date=2026-10-14").get_data(as_text=True), "aside", "leftnav")
        for words in ("New schedule", ">Home<", ">Overview<", ">RTA<", ">Weeks<", ">Schedules<", ">Analysis<",
                      ">Programs<", ">Exports<", ">Team<"):
            self.assertIn(words, side)
        self.assertIn('href="/overview?program=NMG&date=2026-10-14"', side)  # the menu keeps the day
        self.assertIn('aria-current="page">RTA</a>', side)
        self.assertIn('<option value="NMG" selected>NMG</option>', side)
        self.assertNotIn(">People<", side)  # planners manage nobody
        side = self.part(self.admin.get("/").get_data(as_text=True), "aside", "leftnav")
        for words in (">People<", ">LOBs and defaults<"):
            self.assertIn(words, side)

    def test_home_starts_from_my_programs(self):
        cards = self.part(self.lina.get("/").get_data(as_text=True), "section", "programs-home")
        self.assertIn("NMG", cards)
        self.assertNotIn("AE/AR B2B", cards)
        self.assertIn('href="/overview?program=NMG"', cards)
        cards = self.part(self.sara.get("/").get_data(as_text=True), "section", "programs-home")
        self.assertIn("AE/AR B2B", cards)


class TheHomeFollowsThePicker(unittest.TestCase):
    """Phase R task 3: Home shows the program picked on the left, with links to the person's other programs
    and to all of them; planners only ever see their own programs (owner, 2026-10-08)."""

    def setUp(self):
        from webapp.programs import ProgramBook
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        seed_week(self.store, "NMG", "2026-10-11")
        seed_week(self.store, "AE/AR B2B", "2026-10-11")
        book = ProgramBook(self.store)
        book.sync()
        lina = self.store.add_user("lina", "Lina", "Lina-pass-123", must_change=False)
        self.store.set_user_programs(lina, [next(p["id"] for p in book.tree() if p["name"] == "NMG")])
        self.lina = sign_in(self.app, "lina", "Lina-pass-123")
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")

    def home(self, client, url="/"):
        body = client.get(url).get_data(as_text=True)
        found = re.search(r'<section class="programs-home"[^>]*>(.*?)</section>', body, re.S)
        self.assertIsNotNone(found)
        return body, html.unescape(found.group(1))

    def test_home_shows_the_picked_program_only(self):
        _, cards = self.home(self.admin, "/?program=NMG")
        self.assertIn("<h2>NMG</h2>", cards)
        self.assertNotIn("<h2>AE/AR B2B</h2>", cards)
        self.assertRegex(cards, r'<a class="lob" href="/\?program=AE/AR\+B2B">AE/AR B2B</a>')
        self.assertIn('href="/?all=1">Show all my programs</a>', cards)

    def test_home_all_shows_every_program(self):
        _, cards = self.home(self.admin, "/?all=1")
        self.assertIn("<h2>NMG</h2>", cards)
        self.assertIn("<h2>AE/AR B2B</h2>", cards)

    def test_planner_home_never_lists_other_programs(self):
        for url in ("/", "/?all=1", "/?program=AE/AR+B2B"):
            _, cards = self.home(self.lina, url)
            self.assertIn("<h2>NMG</h2>", cards)
            self.assertNotIn("AE/AR B2B", cards)
            self.assertNotIn("Show all my programs", cards)  # one program: nothing else to show

    def upload_pick(self, client, url="/"):
        body = client.get(url).get_data(as_text=True)
        found = re.search(r'<select name="program" data-program-pick>(.*?)</select>', body, re.S)
        self.assertIsNotNone(found)
        return re.findall(r'<option value="([^"]*)" selected>', html.unescape(found.group(1)))

    def test_upload_form_takes_the_program_picked_on_the_left(self):
        # owner, 2026-10-09: "program while uploading should be selected automatically based on choosen program/lob"
        self.admin.get("/?program=NMG")  # picked on the left; the menu remembers it
        self.assertEqual(self.upload_pick(self.admin), ["NMG"])
        self.assertEqual(self.upload_pick(self.admin, "/?program=AE/AR+B2B"), ["AE/AR B2B"])
        self.admin.get("/day?program=AE/AR+B2B")  # picked on another page
        self.assertEqual(self.upload_pick(self.admin), ["AE/AR B2B"])
        self.assertEqual(self.upload_pick(self.admin, "/?all=1"), [""])  # all programs: the planner picks
        self.assertEqual(self.upload_pick(self.lina), ["NMG"])  # one program: always it

    def test_menu_picker_on_home_stays_on_home(self):
        body, _ = self.home(self.admin)
        self.assertIn('data-target="/?program={key}"', body)


class TheExportRange(unittest.TestCase):
    """Phase R (owner, 2026-10-09: "changes in breaks and auxs not reflecting in the exports"): the export's
    period reaches the latest day anything was recorded, so changes made for a schedule that starts after today
    are in it by default; "Next week" is one click."""

    def setUp(self):
        from datetime import datetime, timedelta
        from webapp.schedules import EGYPT
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")
        self.today = datetime.now(EGYPT).date()
        self.later = self.today + timedelta(days=5)
        self.store.set_attendance(program="NMG", shift_date=self.later.isoformat(), associate="Associate 001",
                                  status="Sick", from_min=None, to_min=None, billable=0, user_id=1)

    def test_the_period_reaches_the_latest_record(self):
        page = html.unescape(self.admin.get("/exports").get_data(as_text=True))
        self.assertIn(f'name="to" value="{self.later.isoformat()}"', page)

    def test_next_week_is_a_preset(self):
        from datetime import timedelta
        sunday = self.today - timedelta(days=(self.today.weekday() + 1) % 7) + timedelta(days=7)
        page = html.unescape(self.admin.get("/exports").get_data(as_text=True))
        self.assertIn(f'from={sunday.isoformat()}&to={(sunday + timedelta(days=6)).isoformat()}">Next week</a>', page)


class TheLobFilters(unittest.TestCase):
    """Phase Q: filters go by program, then LOB: a whole program means all of its LOBs (that the person may
    open); the programs list groups LOBs under their program."""

    def setUp(self):
        from webapp.programs import ProgramBook
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        seed_week(self.store, "AE/AR B2B", "2026-10-11")
        seed_week(self.store, "NMG", "2026-10-11")
        book = ProgramBook(self.store)
        book.sync()
        self.ae = book.add_program("AE")
        book.adopt("AE/AR B2B", self.ae, "AR B2B")
        book.add_lob(self.ae, "IT")
        for n, key in enumerate(("AE/AR B2B", "AE IT", "NMG"), start=1):
            self.store.add_activity(program=key, shift_date="2026-10-14", associate=f"Associate 00{n}",
                                    kind="Training", start=600, end_min=660, user_id=1)
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")
        nmg = next(p["id"] for p in book.tree() if p["name"] == "NMG")
        lina = self.store.add_user("lina", "Lina", "Lina-pass-123", must_change=False)
        self.store.set_user_programs(lina, [nmg])
        self.lina = sign_in(self.app, "lina", "Lina-pass-123")

    def csv(self, client, pick):
        got = client.get(f"/exports/download?from=2026-10-14&to=2026-10-14&kind=activities&format=csv&pick={pick}")
        self.assertEqual(got.status_code, 200)
        return got.get_data(as_text=True)

    def test_a_whole_program_or_one_lob(self):
        body = self.csv(self.admin, f"group:{self.ae}")
        self.assertIn("Associate 001", body)
        self.assertIn("Associate 002", body)
        self.assertNotIn("Associate 003", body)
        body = self.csv(self.admin, "key:AE IT")
        self.assertEqual(("Associate 001" in body, "Associate 002" in body), (False, True))
        page = html.unescape(self.admin.get("/exports").get_data(as_text=True))
        self.assertIn('<optgroup label="AE">', page)
        self.assertIn(f'<option value="group:{self.ae}">All of AE</option>', page)
        self.assertIn('<option value="key:AE/AR B2B">AE, AR B2B</option>', page)

    def test_others_programs_stay_closed(self):
        self.assertEqual(self.lina.get(f"/exports/download?from=2026-10-14&to=2026-10-14&kind=activities"
                                       f"&format=csv&pick=group:{self.ae}").status_code, 403)
        self.assertEqual(self.lina.get("/exports/download?from=2026-10-14&to=2026-10-14&kind=activities"
                                       "&format=csv&pick=key:AE IT").status_code, 403)
        self.assertNotIn('label="AE"', self.lina.get("/exports").get_data(as_text=True))

    def test_programs_list_groups_lobs(self):
        page = html.unescape(self.admin.get("/programs").get_data(as_text=True))
        self.assertRegex(page, r'(?s)<th[^>]*class="group"[^>]*>AE</th>.*AE, AR B2B')


class TheStartDate(unittest.TestCase):
    """Phase P: the schedule's first date is picked from Sundays and Mondays when uploading (some programs
    start on Monday), kept as picked, and changed by an admin from the run page."""

    def setUp(self):
        self.app, self.store, *_ = make_app(start_worker=False)
        self.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        self.client = client_for(self.app)
        self.admin = sign_in(self.app, "omar", "Owner-pass-123")

    def test_upload_keeps_the_chosen_start(self):
        body = html.unescape(self.client.get("/?program=NMG&week=2026-10-12").get_data(as_text=True))
        self.assertIn('<select name="week_start"', body)
        self.assertIn('<option value="2026-10-12" data-day="1" selected>Mon 12 Oct 2026</option>', body)
        self.assertIn('<option value="2026-10-11" data-day="0">Sun 11 Oct 2026</option>', body)
        run_id = run_id_of(upload(self.client, program="NMG", week_start="2026-10-12"))
        self.assertEqual(self.store.get_run(run_id)["week_start"], "2026-10-12")  # kept as picked, not moved to Sunday
        page = html.unescape(self.client.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn("starts Monday 12 Oct", page)

    def test_admin_moves_the_start_to_monday(self):
        run_id = run_id_of(upload(self.client, program="NMG", week_start="2026-10-11"))
        page = html.unescape(self.admin.get(f"/runs/{run_id}").get_data(as_text=True))
        self.assertIn('<option value="2026-10-11" data-day="0" selected>Sun 11 Oct 2026</option>', page)
        self.assertIn('<option value="2026-10-12" data-day="1">Mon 12 Oct 2026</option>', page)
        got = self.admin.post(f"/runs/{run_id}/tag", data={"csrf_token": token(self.admin), "program": "NMG",
                                                            "week_start": "2026-10-12", "reason": "starts Monday"})
        self.assertEqual(got.status_code, 302)
        self.assertEqual(self.store.get_run(run_id)["week_start"], "2026-10-12")
        self.assertIn("starts Monday 12 Oct", html.unescape(self.admin.get(f"/runs/{run_id}").get_data(as_text=True)))

    def test_a_date_that_is_not_a_date_is_refused(self):
        got = self.client.post("/runs", data={"csrf_token": token(self.client), "mode": "QUICK", "program": "NMG",
                                              "week_start": "next week",
                                              "workbook": (io.BytesIO(b"PK\x03\x04 x"), "w.xlsx")},
                               content_type="multipart/form-data", follow_redirects=True)
        self.assertIn("Pick the date the schedule starts.", html.unescape(got.get_data(as_text=True)))
        self.assertEqual(self.store.list_runs(), [])


class TheRunDetailsCheck(unittest.TestCase):
    """A change that leaves day records on a week without its schedule is shown first and needs a reason."""

    def test_check_before_saving(self):
        from datetime import date as day
        app, store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        sara = client_for(app)
        run_id = versioned_run(app, store, sara, program="AE/AR B2B", week_start="2026-10-11")
        app.extensions["days"].set_status("AE/AR B2B", day(2026, 10, 14), "Associate 001", "Sick",
                                          store.list_users()[0]["id"])
        post = lambda **data: sara.post(f"/runs/{run_id}/tag", data={  # noqa: E731
            "csrf_token": token(sara), "program": "AE/AR B2B", "week_start": "2026-10-18", **data})
        page = html.unescape(post().get_data(as_text=True))
        for words in ("Check before saving", "Schedule starts", "Sun 11 Oct 2026 → Sun 18 Oct 2026", "stay on their dates",
                      "no schedule", "rename the program instead"):
            self.assertIn(words, page)
        self.assertEqual(store.get_run(run_id)["week_start"], "2026-10-11")  # nothing yet
        page = html.unescape(post(confirm="1").get_data(as_text=True))
        self.assertIn("Give a reason.", page)
        self.assertEqual(store.get_run(run_id)["week_start"], "2026-10-11")
        self.assertEqual(post(confirm="1", reason="wrong week").status_code, 302)
        self.assertEqual(store.get_run(run_id)["week_start"], "2026-10-18")
        self.assertEqual({v["week_start"] for v in app.extensions["schedules"].versions(run_id)}, {"2026-10-18"})


class TheExportsPage(unittest.TestCase):
    """Phase O task 2: the Exports page and its downloads."""

    @classmethod
    def setUpClass(cls):
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.client = client_for(cls.app)
        cls.run_id = versioned_run(cls.app, cls.store, cls.client, program="AE/AR B2B", week_start="2026-10-11")
        cls.client.post("/day/attendance", data={"csrf_token": token(cls.client), "program": "AE/AR B2B",
                                                 "date": "2026-10-14", "associate": "Associate 001", "status": "Sick"})

    def test_page_lists_the_period(self):
        body = html.unescape(self.client.get("/exports?from=2026-10-01&to=2026-10-31").get_data(as_text=True))
        for words in ("Exports", "What to export", "In this period", "Attendance", "Schedule changes and swaps",
                      "Worked hours and adherence", 'name="kind" value="attendance" checked', 'href="/exports"'):
            self.assertIn(words, body)
        self.assertRegex(body, r"<td>Attendance</td><td class=\"n\">1</td><td>Sara, ")

    def test_download_is_recorded(self):
        got = self.client.get("/exports/download?from=2026-10-14&to=2026-10-14&kind=attendance&kind=activity"
                              "&format=xlsx&measure=interval")
        self.assertEqual(got.status_code, 200)
        self.assertIn("attachment; filename=Team_Scheduler_2026-10-14_to_2026-10-14.xlsx",
                      got.headers["Content-Disposition"])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(got.data))
        self.assertEqual(wb.sheetnames, ["About this export", "Attendance", "Day activity"])
        last = self.store.list_events(0, time.time() + 60, kinds=["exported"])[-1]
        self.assertEqual((last["by_name"], last["subject"]), ("Sara", "Team_Scheduler_2026-10-14_to_2026-10-14.xlsx"))

    def test_a_request_it_cannot_answer_says_why(self):
        got = self.client.get("/exports/download?from=2026-10-14&to=2026-10-01&kind=attendance&format=xlsx")
        self.assertEqual(got.status_code, 400)
        self.assertIn("The period ends before it starts.", got.get_data(as_text=True))
        got = self.client.get("/exports/download?from=2026-10-14&to=2026-10-14&kind=attendance&kind=runs&format=csv")
        self.assertIn("A CSV file holds one item", got.get_data(as_text=True))

    def test_shrinkage_coach(self):
        body = html.unescape(self.client.get("/coach?program=AE/AR+B2B&from=2026-10-11&to=2026-10-17").get_data(as_text=True))
        for words in ("Shrinkage coach", "Days recorded", "Download the corrected tab", "Wed"):
            self.assertIn(words, body)
        got = self.client.get("/coach/download?program=AE/AR+B2B&from=2026-10-11&to=2026-10-17")
        self.assertEqual(got.status_code, 200)
        self.assertIn("Shrinkage_60_Min_AE_AR_B2B_2026-10-11_to_2026-10-17.xlsx", got.headers["Content-Disposition"])
        last = self.store.list_events(0, time.time() + 60, kinds=["downloaded"])[-1]
        self.assertEqual(last["subject"], "Shrinkage_60_Min_AE_AR_B2B_2026-10-11_to_2026-10-17.xlsx")

    def test_signed_out_cannot_export(self):
        got = self.app.test_client().get("/exports/download?from=2026-10-14&to=2026-10-14&kind=attendance")
        self.assertEqual(got.status_code, 302)
        self.assertIn("/login", got.headers["Location"])


class Access(unittest.TestCase):
    def test_downloads_need_login(self):
        app, store, *_ = make_app()
        run_id = run_id_of(upload(client_for(app)))
        wait(store, run_id)
        anonymous = app.test_client()
        self.assertEqual(anonymous.get(f"/runs/{run_id}/download").status_code, 302)
        self.assertEqual(anonymous.get(f"/runs/{run_id}/schedule").status_code, 302)


class TheRunnerCommand(unittest.TestCase):
    def test_runner_command_matches_production_flags(self):
        app, store, data, _ = make_app()
        client = client_for(app)
        run_id = run_id_of(upload(client, mode="DEEP"))
        wait(store, run_id)
        argv = json.loads((data / "runs" / run_id / "results" / "argv.json").read_text())["argv"]
        joined = " ".join(argv)
        for needle in ("--package-root", "--results-root", "--mode DEEP", "--stage FULL_SCHEDULE",
                       "--num-workers", "--input", "--resume"):
            self.assertIn(needle, joined)
        self.assertNotIn("--seeds", joined)  # DEEP: the runner's own default (best of 4)
        quick = run_id_of(upload(client, name="q.xlsx", mode="QUICK"))
        wait(store, quick)
        quick_argv = json.loads((data / "runs" / quick / "results" / "argv.json").read_text())["argv"]
        self.assertIn("--seeds 2", " ".join(quick_argv))


if __name__ == "__main__":
    unittest.main()
