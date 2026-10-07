"""Phase H, task 1: every schedule workbook the user opens has the agreed look.

Owner, 2026-10-07: "make sure all output sheets are using the same agreed
visuals as before break and after break i think still using the old visuals".
Measured on a Phase F run folder:
  * the top-level *_BEST_BEFORE_BREAKS_SCHEDULE.xlsx and
    *_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx were the raw engine copies (every
    input tab visible: Start Here, Instructions, FT Wise 30 Min, ...); only the
    copies under production/ had the agreed look;
  * the polished before-breaks workbook's front page said "After Target" and
    "After Floor", and its visible tabs included an empty Break Schedule and
    "FT Wise After Breaks", although no break is placed in it.
The polisher now gives the before-breaks role its own front page and tabs
(misleading tabs hidden, never deleted), and the runner puts the validated
polished files at the top level, keeping the raw engine copies under
debug/raw_engine_output/.
"""
import importlib.util
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
RAW_BEFORE = (REPO / "fixtures" / "real_runs" / "before_break" / "B1_BEFORE"
              / "B1_BEFORE_L6_3_2_3_BEST_BEFORE_BREAKS_SCHEDULE.xlsx")
BEFORE_VISIBLE = ["Read Me First", "Schedule", "Coverage Before Breaks", "Production Summary", "Validation Log"]


def load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


POLISHER = load(ROOT / "engine" / "production" / "production_output_polisher.py", "polisher_h1")
RUNNER = load(ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py", "runner_h1")


def presented(role: str):
    from openpyxl import load_workbook
    tmp = Path(tempfile.mkdtemp()) / RAW_BEFORE.name
    shutil.copy(RAW_BEFORE, tmp)
    POLISHER.present_in_place(tmp, role, "REVIEW ONLY - test")
    return load_workbook(tmp)


def readme_text(wb) -> str:
    ws = wb["Read Me First"]
    return " | ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


class TheBeforeBreaksWorkbookSaysBefore(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = presented("BEST_BEFORE_BREAKS_SCHEDULE")

    def test_before_breaks_dashboard_says_before(self):
        text = readme_text(self.wb)
        self.assertIn("Before Target", text)
        self.assertIn("Before Floor", text)
        self.assertNotIn("After Target", text)
        self.assertNotIn("After Floor", text)

    def test_before_breaks_has_no_after_or_empty_break_sheets_in_view(self):
        visible = [ws.title for ws in self.wb.worksheets if ws.sheet_state == "visible"]
        self.assertEqual(visible, BEFORE_VISIBLE)
        # Hidden, not deleted: the evidence stays in the workbook.
        for name in ("FT Wise After Breaks", "Break Schedule"):
            self.assertIn(name, self.wb.sheetnames)
            self.assertEqual(self.wb[name].sheet_state, "hidden")


class TheAfterBreaksWorkbookIsUnchanged(unittest.TestCase):
    def test_after_breaks_unchanged(self):
        wb = presented("BEST_FINAL_AFTER_BREAKS_SCHEDULE")
        visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        self.assertIn("FT Wise After Breaks", visible)
        self.assertIn("Break Schedule", visible)
        self.assertIn("After Target", readme_text(wb))


def case_root(with_production=True, validation_status="PASS"):
    root = Path(tempfile.mkdtemp()) / "CASE"
    (root / "production").mkdir(parents=True)
    names = ("CASE_L6_3_2_3_BEST_BEFORE_BREAKS_SCHEDULE.xlsx", "CASE_L6_3_2_3_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx")
    for name in names:
        (root / name).write_bytes(b"raw engine copy " + name.encode())
        if with_production:
            (root / "production" / name).write_bytes(b"polished " + validation_status.encode() + name.encode())
    return root, names


class TheTopLevelFilesAreTheDeliverables(unittest.TestCase):
    def test_top_level_copies_are_the_validated_production_files(self):
        root, names = case_root()
        replaced = RUNNER.publish_top_level_copies(root)
        self.assertEqual(sorted(replaced), sorted(names))
        for name in names:
            self.assertEqual((root / name).read_bytes(), (root / "production" / name).read_bytes())
            self.assertEqual((root / "debug" / "raw_engine_output" / name).read_bytes(),
                             b"raw engine copy " + name.encode())

    def test_failed_validation_still_publishes_the_polished_file_with_its_status(self):
        # The polished file carries its own status on its front page; the raw
        # copy, which says nothing about validation, is not what a user opens.
        root, names = case_root(validation_status="FAIL")
        RUNNER.publish_top_level_copies(root)
        for name in names:
            self.assertTrue((root / name).read_bytes().startswith(b"polished FAIL"))

    def test_no_production_copy_leaves_the_top_level_file(self):
        root, names = case_root(with_production=False)
        self.assertEqual(RUNNER.publish_top_level_copies(root), [])
        for name in names:
            self.assertEqual((root / name).read_bytes(), b"raw engine copy " + name.encode())


class TheShortfallScheduleHasTheSameLook(unittest.TestCase):
    """The shortfall schedule (no week meets every hard rule; this one meets
    every person rule and lists each missed minimum) was the one schedule
    workbook still shipped with every input tab showing."""

    def test_shortfall_workbook_is_presented_with_its_shortfalls_in_view(self):
        from openpyxl import load_workbook
        tmp = Path(tempfile.mkdtemp()) / "CASE_L6_3_2_3_HARD_RULE_SHORTFALL_SCHEDULE.xlsx"
        shutil.copy(RAW_BEFORE, tmp)
        wb = load_workbook(tmp)
        wb.create_sheet("Shortfalls").append(["Rule", "Meaning", "Day", "Time"])
        wb.save(tmp)
        POLISHER.present_in_place(tmp, "HARD_RULE_SHORTFALL_SCHEDULE", "NOT RELEASABLE - test")
        wb = load_workbook(tmp)
        visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        self.assertEqual(visible[:3], ["Read Me First", "Schedule", "Shortfalls"])
        self.assertIn("NOT RELEASABLE", readme_text(wb))

    def test_the_runner_presents_it_before_validating_it(self):
        source = (ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py").read_text(encoding="utf-8")
        block = source[source.index("shortfall_books = sorted("):]
        self.assertLess(block.index("present_alternative(shortfall_books[0], 'HARD_RULE_SHORTFALL_SCHEDULE')"),
                        block.index("validate_shortfall_schedule("))


class AResumedRunPolishesTheEngineCopy(unittest.TestCase):
    """After publish_top_level_copies the top-level file is already polished;
    a resumed run must polish the raw engine copy again, not the polished one."""

    def test_the_polisher_prefers_the_raw_engine_copy(self):
        root, names = case_root()
        RUNNER.publish_top_level_copies(root)
        src = POLISHER.engine_copy(root, "*_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx")
        self.assertEqual(src, root / "debug" / "raw_engine_output" / names[1])

    # Final review (Phase H): an OVERWRITE rerun leaves the previous run's raw
    # copy under debug/raw_engine_output/ while the engine writes a fresh one at
    # the top level. The fresh one must be polished and kept, never the stale one.
    def test_a_rerun_polishes_the_fresh_engine_copy_not_the_stale_one(self):
        root, names = case_root()
        RUNNER.publish_top_level_copies(root)
        fresh = b"fresh engine copy of the rerun"
        (root / names[1]).write_bytes(fresh)
        self.assertEqual(POLISHER.engine_copy(root, "*_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx"), root / names[1])
        (root / "production" / names[1]).write_bytes(b"polished rerun")
        RUNNER.publish_top_level_copies(root)
        self.assertEqual((root / "debug" / "raw_engine_output" / names[1]).read_bytes(), fresh)
        self.assertEqual((root / names[1]).read_bytes(), b"polished rerun")

    def test_a_first_run_reads_the_top_level_file(self):
        root, names = case_root(with_production=False)
        self.assertEqual(POLISHER.engine_copy(root, "*_BEST_BEFORE_BREAKS_SCHEDULE.xlsx"), root / names[0])


if __name__ == "__main__":
    unittest.main()
