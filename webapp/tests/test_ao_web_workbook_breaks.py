# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AO (owner, 2026-10-21: an 11-hour shift should get four breaks, and "update the input sheet in the website in
instructions sheet to have that"). The engine already gives shifts of N hours or more their own break set from an
Instructions row ``Break Set For Shifts Of N Hours Or More`` (audit F-20, PRODUCTION_RUN_GUIDE); without the row every
shift gets the global set. The two workbooks the website offers now carry the row for 11 hours: 15, 30, 15, 15 (three
short breaks and a lunch), placed in the Breaks group, so shorter shifts keep their own set."""
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
WORKBOOKS = REPO / "webapp" / "workbooks"
ROW = "Break Set For Shifts Of 11 Hours Or More"
sys.path.insert(0, str(REPO / "engine" / "_tools"))


def instructions(name):
    ws = load_workbook(WORKBOOKS / name)["Instructions"]
    return [(r[1], r[2]) for r in ws.iter_rows(values_only=True) if r[1]]


class TheWebWorkbooks(unittest.TestCase):
    def test_both_carry_the_row_after_the_other_break_settings(self):
        for name in ("Scheduler_Input_Example.xlsx", "Scheduler_Input_Blank.xlsx"):
            with self.subTest(name):
                rows = instructions(name)
                labels = [k for k, _ in rows]
                self.assertIn(ROW, labels)
                self.assertEqual(dict(rows)[ROW], "15, 30, 15, 15")
                self.assertEqual(labels[labels.index(ROW) - 1], "Allow Back-to-Back Breaks")

    def test_the_engine_gives_an_11_hour_shift_four_breaks_and_a_9_hour_shift_its_own_set(self):
        import l632_universal_scheduler as E
        parsed = E.parse_input(WORKBOOKS / "Scheduler_Input_Example.xlsx")
        self.assertEqual(E.break_segments_for(parsed, 660),
                         ((1, "Break 1"), (2, "Lunch"), (1, "Break 2"), (1, "Break 3")))
        self.assertEqual(E.break_segments_for(parsed, 540), parsed.break_segments_q)
        self.assertEqual(len(parsed.break_segments_q), 3)
        self.assertFalse([w for w in parsed.parser_warnings if "BREAK" in str(w).upper()], parsed.parser_warnings)


if __name__ == "__main__":
    unittest.main()
