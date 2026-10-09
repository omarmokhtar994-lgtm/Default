# © 2026 Omar Mokhtar. All rights reserved.
"""Phase R task 7: a ready schedule (shifts already made, breaks still to plan) is uploaded without the engine.

The file is the input workbook with its Schedule tab filled: each person's shift (from the Shift Library) or OFF
for each day. It is checked in seconds, kept as the run's version 1 (kind "ready"), and its missing breaks show as
yellow ("planned next"), not red (owner, 2026-10-08: "upload a full schedule already with shifts and everything
without breaks and we can plot breaks manually")."""
import html
import io
import re
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from webapp.ready import check_ready
from webapp.tests.test_runs import REPO, make_app, run_id_of, sign_in, token, wait
from webapp.tests.test_schedules import AFTER, INPUT
from webapp.versions import read_week

DAYS = ["sun", "mon", "tue", "wed", "thu", "fri", "sat"]


def make_ready(target: Path) -> Path:
    """The B3 input workbook with its Schedule tab filled from the tool's own schedule for that week."""
    wb = load_workbook(INPUT)
    ws = wb["Schedule"]
    week = {a["name"]: a["days"] for a in read_week(AFTER)["associates"]}
    head = next(r for r in range(1, 10) if "sun" in [str(c.value or "").strip().lower() for c in ws[r]])
    cols = {str(c.value or "").strip().lower(): c.column for c in ws[head] if c.value}
    for r in range(head + 1, ws.max_row + 1):
        name = str(ws.cell(r, cols["sf name"]).value or "").strip()
        if name in week:
            for i, d in enumerate(DAYS):
                ws.cell(r, cols[d]).value = week[name][i]
    wb.save(target)
    return target


def schedule_cell(path: Path, name: str, day: str):
    wb = load_workbook(path)
    ws = wb["Schedule"]
    head = next(r for r in range(1, 10) if "sun" in [str(c.value or "").strip().lower() for c in ws[r]])
    cols = {str(c.value or "").strip().lower(): c.column for c in ws[head] if c.value}
    row = next(r for r in range(head + 1, ws.max_row + 1) if ws.cell(r, cols["sf name"]).value == name)
    return wb, ws.cell(row, cols[day]), row


class TheReadyCheck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.ready = make_ready(cls.dir / "ready.xlsx")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def changed(self, name, day, value):
        target = self.dir / f"changed_{name.replace(' ', '_')}_{day}.xlsx"
        wb, cell, row = schedule_cell(self.ready, name, day)
        cell.value = value
        wb.save(target)
        return target, row

    def test_a_filled_schedule_is_accepted(self):
        self.assertEqual(check_ready(self.ready), [])

    def test_ready_workbook_with_unknown_shift_is_rejected(self):
        bad, row = self.changed("Associate 004", "tue", "9-18")
        found = check_ready(bad)
        self.assertEqual(len(found), 1)
        self.assertIn(f"row {row}", found[0])
        self.assertIn("Associate 004", found[0])
        self.assertIn("Tue", found[0])
        self.assertIn("9-18", found[0])
        self.assertIn("Shift Library", found[0])

    def test_an_empty_day_is_rejected(self):
        bad, _ = self.changed("Associate 004", "wed", None)
        self.assertTrue(any("Associate 004" in p and "Wed" in p and "empty" in p for p in check_ready(bad)))

    def test_a_name_twice_is_rejected(self):
        wb, cell, row = schedule_cell(self.ready, "Associate 005", "sf name")
        cell.value = "Associate 004"
        target = self.dir / "twice.xlsx"
        wb.save(target)
        self.assertTrue(any("Associate 004" in p and "more than once" in p for p in check_ready(target)))

    def test_no_schedule_tab_is_said_plainly(self):
        wb = load_workbook(self.ready)
        del wb["Schedule"]
        target = self.dir / "no_tab.xlsx"
        wb.save(target)
        self.assertEqual(check_ready(target), ["There is no Schedule tab: use the input workbook and fill its "
                                               "Schedule tab with each person's shift or OFF for each day."])


class TheReadyUpload(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from webapp.programs import ProgramBook
        cls.dir = Path(tempfile.mkdtemp())
        cls.ready = make_ready(cls.dir / "ready.xlsx").read_bytes()
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        book = ProgramBook(cls.store)
        cls.key = book.add_lob(book.add_program("SAKS"), "NMG Tier 2")
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")
        cls.run_id = run_id_of(cls.upload(cls.ready))

    @classmethod
    def upload(cls, data, name="ready_week.xlsx"):
        return cls.admin.post("/runs", data={"csrf_token": token(cls.admin), "kind": "ready", "mode": "QUICK",
                                             "program": cls.key, "week_start": "2026-10-11",
                                             "workbook": (io.BytesIO(data), name)},
                              content_type="multipart/form-data")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def test_ready_upload_makes_a_version_without_the_engine(self):
        run = wait(self.store, self.run_id, statuses=("DONE", "REJECTED", "FAILED"))
        self.assertEqual((run["status"], run["mode"]), ("DONE", "READY"))
        self.assertIn("the engine did not run", run["message"])
        self.assertFalse(self.app.extensions["runs"].results_dir(self.run_id).exists())  # no engine output at all
        versions = self.app.extensions["schedules"].versions(self.run_id)
        self.assertEqual([(v["kind"], v["label"], v["program"], v["week_start"]) for v in versions],
                         [("ready", "Ready schedule (uploaded)", self.key, "2026-10-11")])

    def test_ready_missing_breaks_are_yellow_not_red(self):
        version = self.app.extensions["schedules"].versions(self.run_id)[0]
        found = self.app.extensions["schedules"].view(version["id"])["problems"]
        self.assertTrue(found)
        self.assertEqual({p["severity"] for p in found if "break" in p["text"].lower()}, {"yellow"})
        self.assertEqual([p for p in found if p["severity"] == "red"], [])

    def test_ready_run_page_and_rta_work(self):
        page = html.unescape(self.admin.get(f"/runs/{self.run_id}").get_data(as_text=True))
        self.assertIn("Ready schedule uploaded", page)
        self.assertIn("Plan breaks", page)
        self.assertNotIn("Download the results", page)
        day = html.unescape(self.admin.get(f"/day?program={self.key.replace(' ', '+')}&date=2026-10-14")
                            .get_data(as_text=True))
        self.assertIn("On shift today", day)
        self.assertIn("Associate 001", day)
        self.assertIn("so this is the uploaded ready schedule", day)  # not "the tool's own": the engine never ran
        self.assertNotIn("tool's own schedule", day)

    def test_ready_workbook_with_unknown_shift_is_rejected(self):
        target = self.dir / "bad.xlsx"
        target.write_bytes(self.ready)
        wb, cell, _ = schedule_cell(target, "Associate 004", "tue")
        cell.value = "9-18"
        wb.save(target)
        run_id = run_id_of(self.upload(target.read_bytes(), "bad.xlsx"))
        run = wait(self.store, run_id, statuses=("DONE", "REJECTED", "FAILED"))
        self.assertEqual(run["status"], "REJECTED")
        self.assertIn("9-18", run["message"])
        self.assertEqual(self.app.extensions["schedules"].versions(run_id), [])

    def test_schedules_page_shows_the_details_and_links_to_edit_them(self):
        # owner, 2026-10-09: "not able to change schedule name or start dates after upload": the Schedules page
        # (where the menu goes) said neither and had no way to the run page's editor
        page = html.unescape(self.admin.get(f"/runs/{self.run_id}/schedules").get_data(as_text=True))
        details = re.search(r'<p class="byline rundetails">(.*?)</p>', page, re.S)
        self.assertIsNotNone(details)
        for words in ("ready_week.xlsx", "SAKS, NMG Tier 2", "starts Sunday 11 Oct"):
            self.assertIn(words, details.group(1))
        self.assertIn(f'<a href="/runs/{self.run_id}?edit=1#edit">Edit the name, program or start date</a>',
                      details.group(1))
        run_page = self.admin.get(f"/runs/{self.run_id}?edit=1").get_data(as_text=True)
        self.assertRegex(run_page, r'<details class="advanced tag[^"]*" id="edit" open>')
        self.assertNotRegex(self.admin.get(f"/runs/{self.run_id}").get_data(as_text=True), r'id="edit" open>')

    def test_only_who_may_edit_gets_the_edit_link(self):
        from webapp.programs import ProgramBook
        saks = next(p["id"] for p in ProgramBook(self.store).tree() if p["name"] == "SAKS")
        if not any(u["username"] == "nour" for u in self.store.list_users()):
            nour = self.store.add_user("nour", "Nour", "Nour-pass-123", must_change=False)
            self.store.set_user_programs(nour, [saks])
        planner = sign_in(self.app, "nour", "Nour-pass-123")
        page = html.unescape(planner.get(f"/runs/{self.run_id}/schedules").get_data(as_text=True))
        self.assertIn('<p class="byline rundetails">', page)
        self.assertNotIn("?edit=1", page)

    def test_upload_form_offers_build_or_ready(self):
        page = html.unescape(self.admin.get("/").get_data(as_text=True))
        self.assertRegex(page, r'<input type="radio" name="kind" value="build" checked')
        self.assertRegex(page, r'<input type="radio" name="kind" value="ready"')
        self.assertIn("Upload a ready schedule", page)



class TheUnfiledRuns(unittest.TestCase):
    """Phase S: a schedule uploaded without a program is listed on Home, where it is filed under a program and LOB
    (owner, 2026-10-09: "in case program was not choosen in the upload im not able to located it and assign it to
    program")."""

    @classmethod
    def setUpClass(cls):
        from webapp.programs import ProgramBook
        cls.dir = Path(tempfile.mkdtemp())
        cls.ready = make_ready(cls.dir / "ready.xlsx").read_bytes()
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        book = ProgramBook(cls.store)
        saks = book.add_program("SAKS")
        cls.key = book.add_lob(saks, "NMG Tier 2")
        for name in ("lina", "nour"):
            user = cls.store.add_user(name, name.title(), f"{name.title()}-pass-123", must_change=False)
            cls.store.set_user_programs(user, [saks])
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")
        cls.lina = sign_in(cls.app, "lina", "Lina-pass-123")
        cls.nour = sign_in(cls.app, "nour", "Nour-pass-123")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def upload(self, client, name):
        got = client.post("/runs", data={"csrf_token": token(client), "kind": "ready", "program": "", "week_start": "",
                                         "workbook": (io.BytesIO(self.ready), name)},
                          content_type="multipart/form-data")
        run_id = run_id_of(got)
        self.assertEqual(wait(self.store, run_id, statuses=("DONE", "REJECTED", "FAILED"))["status"], "DONE")
        return run_id

    def unfiled(self, client):
        body = html.unescape(client.get("/").get_data(as_text=True))
        found = re.search(r'<section class="panel unfiled"[^>]*>(.*?)</section>', body, re.S)
        return found.group(1) if found else ""

    def test_a_run_without_a_program_is_listed_on_home_and_filed_there(self):
        run_id = self.upload(self.admin, "unfiled_admin.xlsx")
        panel = self.unfiled(self.admin)
        self.assertIn("Not filed under a program", panel)
        self.assertIn("unfiled_admin.xlsx", panel)
        self.assertIn(f'action="/runs/{run_id}/tag"', panel)
        self.assertIn(f'<option value="{self.key}">SAKS, NMG Tier 2</option>', panel)  # picked from the list
        got = self.admin.post(f"/runs/{run_id}/tag", data={"csrf_token": token(self.admin), "program": self.key,
                                                           "week_start": "2026-10-11", "back": "home"})
        self.assertIn(got.status_code, (302, 303))
        self.assertEqual(got.headers["Location"], "/")  # back to Home, where the next one waits
        self.assertEqual((self.store.get_run(run_id)["program"], self.store.get_run(run_id)["week_start"]),
                         (self.key, "2026-10-11"))
        self.assertEqual([v["program"] for v in self.app.extensions["schedules"].versions(run_id)], [self.key])
        self.assertNotIn("unfiled_admin.xlsx", self.unfiled(self.admin))

    def test_a_planner_files_their_own_and_nobody_elses(self):
        run_id = self.upload(self.lina, "unfiled_lina.xlsx")
        self.assertIn(f'action="/runs/{run_id}/tag"', self.unfiled(self.lina))
        self.assertNotIn("unfiled_lina.xlsx", self.unfiled(self.nour))  # not hers, and in no program she has
        self.assertIn("unfiled_lina.xlsx", self.unfiled(self.admin))  # an admin sees every one


if __name__ == "__main__":
    unittest.main()
