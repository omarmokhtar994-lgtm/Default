# © 2026 Omar Mokhtar. All rights reserved.
"""A ready schedule: shifts already made elsewhere, managed here without running the engine (Phase R).

The file is the input workbook (demand, Shift Library, rules) with its Schedule tab filled: each person's
shift from the Shift Library, or OFF / Leave, for each day. Breaks are planned afterwards on the website.
This check is the website's own and takes a second; the independent validator still checks the version."""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook

from .versions import DAYS, _header, _library, _name_col, _norm, shift_span

NO_TAB = ("There is no Schedule tab: use the input workbook and fill its Schedule tab with each person's shift or "
          "OFF for each day.")


def check_ready(path: Path) -> List[str]:
    """What stops this workbook from being used as a ready schedule, in plain words; empty when nothing does."""
    try:
        wb = load_workbook(path)
    except Exception as exc:  # a broken or non-Excel file: said, never a server error
        return [f"The file could not be read as an Excel workbook ({type(exc).__name__})."]
    if "Schedule" not in wb.sheetnames:
        return [NO_TAB]
    ws = wb["Schedule"]
    try:
        head, cols = _header(ws, ("sf name", "associate name"))
    except ValueError:
        return ["The Schedule tab has no name column (SF Name or Associate Name) in its first 20 rows."]
    day_cols = [cols.get(_norm(d)) for d in DAYS]
    if None in day_cols:
        return ["The Schedule tab needs a column for each day, Sun to Sat."]
    allowed = {_norm(x) for x in _library(wb)}  # the Shift Library's shifts, plus OFF and Leave
    found: List[str] = []
    if not any(shift_span(x) for x in _library(wb)):
        found.append("The Shift Library tab has no shifts: add the shifts this schedule uses.")
    first_row: Dict[str, int] = {}
    name_col = _name_col(cols)
    for r in range(head + 1, ws.max_row + 1):
        name = str(ws.cell(r, name_col).value or "").strip()
        if not name:
            continue
        if name in first_row:
            found.append(f"Schedule tab, rows {first_row[name]} and {r}: {name} is listed more than once.")
        first_row.setdefault(name, r)
        for day, c in zip(DAYS, day_cols):
            text = str(ws.cell(r, c).value or "").strip()
            if not text:
                found.append(f"Schedule tab, row {r} ({name}), {day}: the day is empty; put a shift, OFF or Leave.")
            elif _norm(text) not in allowed:
                found.append(f"Schedule tab, row {r} ({name}), {day}: “{text}” is not a shift in the Shift Library "
                             "(or OFF, Leave).")
    if not first_row:
        found.append("The Schedule tab has nobody in it: put each person's name and their shift for each day.")
    return found
