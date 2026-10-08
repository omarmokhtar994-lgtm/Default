# © 2026 Omar Mokhtar. All rights reserved.
"""Schedule workbooks: read a week, change one shift, check a version.

A schedule workbook from the engine carries the whole input (demand,
shrinkage, Shift Library, languages, rules) plus the schedule (Schedule and
Final Schedule tabs) and its breaks (Break Schedule tabs). A change writes the
new shift on both schedule tabs and removes that person's breaks for that day
(they belonged to the old shift). Checks are the independent validator's
(engine/tools/independent_validator.py), turned into plain problems with a
severity and the cells or days they concern.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import load_workbook

from .week import compact

DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
STATUSES = ["OFF", "Leave"]
SCHEDULE_TABS = ("Schedule", "Final Schedule")
BREAK_TABS = ("Break Schedule Active", "Break Schedule")
VALIDATOR = Path("engine") / "tools" / "independent_validator.py"
VALIDATE_TIMEOUT = 300
SETTINGS = {"rest_gap_hours": ("rest gap hours", "difference between shifts", "minimum rest gap"),
            "max_different_shifts": ("count of different shifts per week", "max different shifts per week")}


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def _hhmm(value: Any) -> str:
    if hasattr(value, "hour"):
        return f"{value.hour:02d}:{value.minute:02d}"
    m = re.search(r"(\d{1,2}):(\d{2})", str(value or ""))
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else str(value or "")


def _minutes(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def shift_span(label: str) -> Optional[Tuple[int, int]]:
    """(start, end) minutes of a shift label like "21:00 - 06:00"; end past midnight is +1440."""
    found = re.findall(r"(\d{1,2}):(\d{2})", str(label or ""))
    if len(found) != 2:
        return None
    start, end = (int(h) * 60 + int(m) for h, m in found)
    return start, end + (1440 if end <= start else 0)


def _header(ws, needles: Iterable[str]) -> Tuple[int, Dict[str, int]]:
    for r in range(1, min(ws.max_row, 20) + 1):
        cols = {_norm(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
        if any(n in cols for n in needles):
            return r, cols
    raise ValueError(f"{ws.title}: no header row with {', '.join(needles)}")


def _name_col(cols: Dict[str, int]) -> int:
    for key in ("sf name", "name", "associate name", "associate"):
        if key in cols:
            return cols[key]
    raise ValueError("no associate-name column")


def _library(wb) -> List[str]:
    ws = next((wb[n] for n in wb.sheetnames if _norm(n) == "shift library"), None)
    labels: List[str] = []
    if ws is not None:
        row, cols = _header(ws, ("allowed shift / status", "allowed shift"))
        col = cols.get("allowed shift / status") or cols.get("allowed shift")
        for r in range(row + 1, ws.max_row + 1):
            v = str(ws.cell(r, col).value or "").strip()
            if v and _norm(v) not in ("off", "leave", "planned") and shift_span(v):
                labels.append(v)
    return labels + STATUSES


def _settings(wb) -> Dict[str, Any]:
    ws = next((wb[n] for n in wb.sheetnames if _norm(n) == "instructions"), None)
    found: Dict[str, Any] = {}
    if ws is None:
        return found
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        for i, v in enumerate(cells):
            for key, names in SETTINGS.items():
                if key not in found and _norm(v) in names:
                    value = next((x for x in cells[i + 1:] if x not in (None, "")), None)
                    try:
                        found[key] = float(value)
                    except (TypeError, ValueError):
                        pass
    return found


def _breaks(wb) -> List[dict]:
    ws = next((wb[n] for n in BREAK_TABS if n in wb.sheetnames), None)
    out: List[dict] = []
    if ws is None:
        return out
    row, cols = _header(ws, ("associate",))
    get = lambda r, k: ws.cell(r, cols[k]).value if k in cols else None  # noqa: E731
    for r in range(row + 1, ws.max_row + 1):
        name = get(r, "associate")
        if not name:
            continue
        try:
            minutes = int(float(get(r, "duration minutes") or 0))
        except (TypeError, ValueError):
            minutes = 0
        out.append({"associate": str(name).strip(), "day": str(get(r, "day") or "").strip(),
                    "kind": str(get(r, "break type") or "").strip(), "start": _hhmm(get(r, "start")),
                    "minutes": minutes})
    return out


_REF = re.compile(r"^=\s*\$?([A-Z]{1,3})\$?(\d+)\s*([+-])\s*(\d+)\s*$")


def _slot(ws, r: int, c: int) -> str:
    """A slot number; a simple formula such as =A30+1 (common in Slot columns, and kept
    without a cached value when a file is saved by a script) is followed to its number.
    Anything else is shown as written."""
    from openpyxl.utils import column_index_from_string
    shown = str(ws.cell(r, c).value or "").strip()
    total, seen = 0.0, set()
    while (r, c) not in seen:
        seen.add((r, c))
        value = ws.cell(r, c).value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            number = float(value) + total
            return str(int(number)) if number.is_integer() else str(number)
        found = _REF.match(str(value or "").strip())
        if not found:
            break
        total += (1 if found.group(3) == "+" else -1) * float(found.group(4))
        r, c = int(found.group(2)), column_index_from_string(found.group(1))
    return shown


def read_week(path: Path) -> Dict[str, Any]:
    """The week in a schedule workbook: people, their seven days, breaks, the shifts on offer."""
    wb = load_workbook(path)
    ws = wb["Schedule"] if "Schedule" in wb.sheetnames else wb["Final Schedule"]
    row, cols = _header(ws, ("sf name", "associate name"))
    name_col = _name_col(cols)
    day_cols = [cols.get(_norm(d)) for d in DAYS]
    if None in day_cols:
        raise ValueError(f"{ws.title}: needs a column for each day, Sun to Sat")
    associates = []
    for r in range(row + 1, ws.max_row + 1):
        name = str(ws.cell(r, name_col).value or "").strip()
        if not name:
            continue
        cell = lambda key: str(ws.cell(r, cols[key]).value or "").strip() if key in cols else ""  # noqa: E731
        associates.append({"name": name, "slot": _slot(ws, r, cols["slot"]) if "slot" in cols else "",
                           "emp_id": cell("emp id"), "language": cell("language"),
                           "tl": cell("tl"), "days": [str(ws.cell(r, c).value or "").strip() for c in day_cols]})
    breaks = _breaks(wb)
    return {"associates": associates, "breaks": breaks, "shifts": _library(wb), "settings": _settings(wb),
            "stage": "after" if breaks else "before"}


PERSON = ("emp id", "email", "sf name", "associate name", "name", "tl", "language")  # who sits in a slot


def swap_slots(src: Path, dst: Path, first: str, second: str) -> None:
    """Two people swap slots: each slot keeps its shifts and breaks and takes the other person
    (Emp ID, Email, name, TL, language) on both schedule tabs; break rows follow the new names."""
    if _norm(first) == _norm(second):
        raise ValueError("Pick two different people to swap.")
    wb = load_workbook(src)
    found = 0
    for tab in SCHEDULE_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        row, cols = _header(ws, ("sf name", "associate name"))
        name_col = _name_col(cols)
        rows = {_norm(ws.cell(r, name_col).value): r for r in range(row + 1, ws.max_row + 1)}
        missing = [n for n in (first, second) if _norm(n) not in rows]
        if missing:
            raise ValueError(f"{missing[0]!r} is not on this schedule")
        a, b = rows[_norm(first)], rows[_norm(second)]
        for key in PERSON:
            if key in cols:
                c = cols[key]
                ws.cell(a, c).value, ws.cell(b, c).value = ws.cell(b, c).value, ws.cell(a, c).value
        found += 1
    if not found:
        raise ValueError("no schedule tab to swap on")
    for tab in BREAK_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        row, cols = _header(ws, ("associate",))
        c = cols["associate"]
        for r in range(row + 1, ws.max_row + 1):
            who = _norm(ws.cell(r, c).value)
            if who == _norm(first):
                ws.cell(r, c).value = second
            elif who == _norm(second):
                ws.cell(r, c).value = first
    wb.save(dst)


def apply_change(src: Path, dst: Path, name: str, day: str, value: str) -> None:
    """Write one shift (or OFF / Leave) for one person and day; drop that day's breaks."""
    if day not in DAYS:
        raise ValueError(f"unknown day {day!r}")
    wb = load_workbook(src)
    allowed = {_norm(s): s for s in _library(wb)}
    if _norm(value) not in allowed:
        raise ValueError(f"{value!r} is not in this program's Shift Library (or OFF / Leave)")
    value = allowed[_norm(value)]
    written, unchanged = 0, False
    for tab in SCHEDULE_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        row, cols = _header(ws, ("sf name", "associate name"))
        name_col, day_col = _name_col(cols), cols.get(_norm(day))
        for r in range(row + 1, ws.max_row + 1):
            if _norm(ws.cell(r, name_col).value) == _norm(name):
                unchanged = unchanged or _norm(ws.cell(r, day_col).value) == _norm(value)
                ws.cell(r, day_col).value = value
                written += 1
                break
    if not written:
        raise ValueError(f"{name!r} is not on this schedule")
    for tab in BREAK_TABS if not unchanged else ():  # the same shift keeps its breaks
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        row, cols = _header(ws, ("associate",))
        doomed = [r for r in range(row + 1, ws.max_row + 1)
                  if _norm(ws.cell(r, cols["associate"]).value) == _norm(name)
                  and _norm(ws.cell(r, cols["day"]).value) == _norm(day)]
        for r in reversed(doomed):
            ws.delete_rows(r)
    wb.save(dst)


def validate(input_path: Path, schedule_path: Path, package_root: Path,
             options: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    """The independent validator on one version; about 3 seconds."""
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "check.json"
        cmd = [sys.executable, str(Path(package_root) / VALIDATOR), "--input", str(input_path),
               "--output", str(schedule_path), "--json-out", str(out), "--csv-out", str(Path(tmp) / "check.csv")]
        if options and options.get("language_window"):
            cmd += ["--language-working-window", options["language_window"]]
        try:
            proc = subprocess.run(cmd, cwd=str(package_root), capture_output=True, text=True, timeout=VALIDATE_TIMEOUT)
            data = json.loads(out.read_text(encoding="utf-8"))
        except (OSError, ValueError, subprocess.TimeoutExpired) as exc:
            detail = getattr(locals().get("proc"), "stderr", "") or str(exc)
            return {"status": "ERROR", "failures": [{"type": "VALIDATOR_EXCEPTION", "detail": detail[-1500:]}],
                    "warnings": [], "metrics": {}, "intervals": []}
    return {"status": data.get("status"), "failures": data.get("failures") or [], "warnings": data.get("warnings") or [],
            "metrics": data.get("canonical_metrics") or {}, "intervals": compact(data.get("interval_rows") or [])}


def _rest_gap(left: str, right: str) -> Optional[float]:
    a, b = shift_span(left), shift_span(right)
    return None if not a or not b else round((1440 + b[0] - a[1]) / 60, 1)


def _num(v: Any) -> str:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return str(int(f)) if f.is_integer() else f"{f:g}"


def _humanise(code: str) -> str:
    words = code.replace("_", " ").lower()
    return words[:1].upper() + words[1:]


def _example(e: Dict[str, Any]) -> str:
    when = " ".join(str(e.get(k)) for k in ("day", "time", "interval") if e.get(k))
    return when or ", ".join(f"{k.replace('_', ' ')} {_num(v)}" for k, v in e.items() if not isinstance(v, (list, dict)))


def _describe(item: Dict[str, Any], settings: Dict[str, Any]) -> Tuple[str, List[Tuple[str, str]], List[str]]:
    """(text, cells, days) for one validator finding."""
    t, who = item.get("type", ""), item.get("associate", "")
    day, cells, days = item.get("day"), [], []
    if who and day in DAYS:
        cells = [(who, day)]
    if t == "REST_VIOLATION":
        gap = _rest_gap(item.get("from", ""), item.get("to", ""))
        rule = settings.get("rest_gap_hours")
        cells = [(who, item.get("from_day")), (who, item.get("to_day"))]
        return (f"{who} rests only {_num(gap)} hours between {item.get('from_day')} {item.get('from')} and "
                f"{item.get('to_day')} {item.get('to')}" + (f"; the rule is {_num(rule)}." if rule else "."), cells, days)
    if t == "CYCLIC_REST_VIOLATION":
        return (f"{who} has too little rest from Saturday {item.get('from_shift')} into the next Sunday "
                f"{item.get('to_shift')}.", [(who, "Sat"), (who, "Sun")], days)
    if t == "PREVIOUS_SATURDAY_REST_VIOLATION":
        return (f"{who} has too little rest from last Saturday ({item.get('previous_saturday')}) into Sunday "
                f"{item.get('sunday')}.", [(who, "Sun")], days)
    if t == "MAX_SHIFT_VARIETY_VIOLATION":
        return (f"{who} has {item.get('actual')} different shift times this week; the limit is {item.get('maximum')}.",
                [(who, "*")], days)
    if t == "OFF_COUNT_VIOLATION":
        offs = item.get("off_days") or []
        return (f"{who} has {item.get('actual')} days off ({', '.join(offs) or 'none'}); {item.get('expected')} are required.",
                [(who, "*")], days)
    if t == "CONSECUTIVE_OFF_VIOLATION":
        return f"{who}'s days off ({', '.join(item.get('off_days') or [])}) must be next to each other.", [(who, "*")], days
    if t == "LEAVE_VIOLATION":
        return f"{who} is on leave on {day} but has {item.get('actual')}.", cells, days
    if t == "HARD_OFF_VIOLATION":
        return f"{who} must be off on {day} (a firm day-off request) but has {item.get('actual')}.", cells, days
    if t == "FIXED_SHIFT_VIOLATION":
        return f"{who} has a fixed {item.get('expected')} on {day} but has {item.get('actual')}.", cells, days
    if t == "UNKNOWN_OR_BLANK_ASSIGNMENT":
        return f"{who} has no shift, or one the program does not offer, on {day}.", cells, days
    if t == "LANGUAGE_WORKING_WINDOW":
        return (f"{who} ({item.get('language')}) works {item.get('shift')} on {day}, outside that language's hours "
                f"({item.get('window')}).", cells, days)
    if t == "BREAK_SEGMENT_COUNT_OR_DURATION":
        if not item.get("actual"):
            return f"{who} has no breaks on {day} yet; \"Plan breaks\" adds them.", cells, days
        return (f"{who}'s breaks on {day} do not match the rule (expected {item.get('expected')} minutes, "
                f"has {item.get('actual')}).", cells, days)
    if t in ("ZERO_STAFF_ACTIVE", "LANGUAGE_MINIMUM", "OPENING_MINIMUM"):
        d = item.get("detail") or {}
        what = {"ZERO_STAFF_ACTIVE": "nobody on the floor",
                "LANGUAGE_MINIMUM": f"{d.get('actual')} for {d.get('group')}, the minimum is {d.get('minimum')}",
                "OPENING_MINIMUM": f"{d.get('actual')} at opening, the minimum is {d.get('minimum')}"}[t]
        return f"{d.get('day')} {d.get('time')}: {what}.", [], [d.get("day")]
    examples = item.get("examples") or (item.get("issues") or [{}])[0].get("examples") or []
    days = sorted({e.get("day") for e in examples if isinstance(e, dict) and e.get("day") in DAYS}, key=DAYS.index)
    if t == "BREAK_CONCURRENCY":
        e = examples[0] if examples else {}
        return (f"{item.get('count')} quarter-hours have more people on break than allowed, for example "
                f"{e.get('day')} {e.get('time')} ({e.get('on_break')} on break, the limit is {e.get('maximum')}).", [], days)
    if t == "HARD_FLOOR":
        return (f"{item.get('count')} intervals are below the hard floor, for example "
                f"{', '.join(_example(e) for e in examples[:3])}.", [], days)
    values = ", ".join(f"{k.replace('_', ' ')} {_num(v)}" for k, v in item.items()
                       if k not in ("type", "associate", "day", "note", "examples", "issues") and not isinstance(v, (list, dict)))
    lead = f"{who}, {day}: " if who and day else (f"{who}: " if who else "")
    tail = f" (for example {', '.join(_example(e) for e in examples[:3])})" if examples else ""
    return f"{lead}{_humanise(t)}" + (f" ({values})" if values else "") + tail + ".", cells, days


def _size(item: Dict[str, Any]) -> float:
    for key in ("count", "actual", "overage_cap_violations"):
        try:
            return float(item[key]) if not isinstance(item[key], (list, dict)) else float(len(item[key]))
        except (KeyError, TypeError, ValueError):
            continue
    return 1.0


def problems(result: Dict[str, Any], edited: Optional[Set[Tuple[str, str]]] = None,
             settings: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """The validator's findings in plain words: failures red, warnings yellow; missing
    breaks on a day whose shift was edited are yellow (they are planned next)."""
    edited, settings = edited or set(), settings or {}
    out: List[Dict[str, Any]] = []
    for kind, items in (("red", result.get("failures") or []), ("yellow", result.get("warnings") or [])):
        for item in items:
            if not isinstance(item, dict):
                continue
            text, cells, days = _describe(item, settings)
            severity = kind
            if item.get("type") == "BREAK_SEGMENT_COUNT_OR_DURATION" and not item.get("actual") \
                    and (item.get("associate"), item.get("day")) in edited:
                severity = "yellow"
            parts = [item.get("type", "")] + [str(item.get(k)) for k in ("associate", "day", "from_day", "to_day", "language")
                                              if item.get(k)]
            if item.get("detail"):
                parts += [str((item["detail"] or {}).get(k)) for k in ("day", "time", "group")]
            out.append({"key": "|".join(parts), "severity": severity, "text": text, "size": _size(item),
                        "cells": [(n, d) for n, d in cells if n and d], "days": [d for d in days if d]})
    return out


def added(before: List[Dict[str, Any]], after: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Problems in ``after`` that ``before`` did not have (or had fewer of)."""
    old = {p["key"]: p for p in before}
    return [p for p in after if p["key"] not in old or p["size"] > old[p["key"]]["size"]
            or (p["severity"] == "red" and old[p["key"]]["severity"] != "red")]


def marks(found: List[Dict[str, Any]], week: Optional[Dict[str, Any]] = None,
          edited: Optional[Set[Tuple[str, str]]] = None) -> Dict[str, Dict[str, str]]:
    """Worst severity per cell ("name|day") and per day. A week-level problem for a
    person ("*") marks the days of that person that were edited, else the person's row."""
    rank = {"yellow": 1, "red": 2}
    cells: Dict[str, str] = {}
    days: Dict[str, str] = {}

    def put(store: Dict[str, str], key: str, sev: str) -> None:
        if rank[sev] > rank.get(store.get(key, ""), 0):
            store[key] = sev

    for p in found:
        for name, day in p["cells"]:
            targets = [day] if day != "*" else [d for (n, d) in (edited or set()) if n == name] or ["*"]
            for d in targets:
                put(cells, f"{name}|{d}", p["severity"])
                if d in DAYS:
                    put(days, d, p["severity"])
        for d in p["days"]:
            put(days, d, p["severity"])
    return {"cells": cells, "days": days}


def with_notes(path: Path, view: Dict[str, Any]) -> "io.BytesIO":
    """The version's workbook with a first tab saying what changed, who, why, and what the checks found."""
    import io
    from datetime import datetime, timedelta, timezone

    from openpyxl.styles import Font

    egypt = timezone(timedelta(hours=3))
    wb = load_workbook(path)
    ws = wb.create_sheet("Version Notes", 0)
    row = view["row"]
    bold = Font(bold=True)
    lines = [
        ("Version notes", ""),
        ("Version", row["label"]),
        ("In use", "Yes" if row["in_use"] else "No"),
        ("Checked by the independent validator", f"{len(view['added'])} problem(s) added by the edits"),
        ("Note", "The Schedule and Break Schedule tabs hold this version. The audit tabs after them describe "
                 "the tool's original schedule."),
        ("", ""),
    ]
    for r, (a, b) in enumerate(lines, start=1):
        ws.cell(r, 1, a).font = bold
        ws.cell(r, 2, b)
    r = len(lines) + 1
    ws.cell(r, 1, "Problems the edits added").font = bold
    for p in view["added"] or [{"severity": "", "text": "None"}]:
        r += 1
        ws.cell(r, 1, {"red": "Rule broken", "yellow": "Warning"}.get(p["severity"], ""))
        ws.cell(r, 2, p["text"])
    r += 2
    for c, head in enumerate(("When (Egypt time)", "Who", "Associate", "Day", "From", "To", "Reason", "Severity"), 1):
        ws.cell(r, c, head).font = bold
    for change in view["changes"]:
        r += 1
        when = datetime.fromtimestamp(change["at"], egypt).strftime("%a %d %b %Y, %H:%M")
        for c, v in enumerate((when, change["by_name"], change["associate"], change["day"], change["old"],
                               change["new"], change["reason"], change["severity"]), 1):
            ws.cell(r, c, v)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
