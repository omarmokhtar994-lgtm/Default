# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AQ (owner, 2026-10-21: "go ahead with WHOLE_SHIFT as optional row"). The two workbooks the website offers carry
the Instructions row ``Language Working Window Shift Rule`` in the Language / Skill group, right after Language Working
Window, set to START (what the engine always did) with a START or WHOLE_SHIFT dropdown; the engine reads it."""
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
WORKBOOKS = REPO / "webapp" / "workbooks"
ROW = "Language Working Window Shift Rule"
sys.path.insert(0, str(REPO / "engine" / "_tools"))


class TheWebWorkbooks(unittest.TestCase):
    def test_both_carry_the_row_after_language_working_window_set_to_start(self):
        for name in ("Scheduler_Input_Example.xlsx", "Scheduler_Input_Blank.xlsx"):
            with self.subTest(name):
                ws = load_workbook(WORKBOOKS / name)["Instructions"]
                rows = [(r[1], r[2]) for r in ws.iter_rows(values_only=True) if r[1]]
                labels = [k for k, _ in rows]
                self.assertEqual(dict(rows)[ROW], "START")
                self.assertEqual(labels[labels.index(ROW) - 1], "Language Working Window")

    def test_the_dropdown_offers_start_and_whole_shift_only(self):
        ws = load_workbook(WORKBOOKS / "Scheduler_Input_Example.xlsx")["Instructions"]
        row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == ROW)
        found = [v.formula1 for v in ws.data_validations.dataValidation if f"C{row}" in str(v.sqref)]
        self.assertEqual(found, ['"START,WHOLE_SHIFT"'])

    def test_the_engine_reads_start_and_the_default_contract_is_unchanged(self):
        import l632_universal_scheduler as E
        parsed = E.parse_input(WORKBOOKS / "Scheduler_Input_Example.xlsx")
        self.assertEqual(parsed.language_window_shift_rule, "START")
        self.assertNotIn("language_window_shift_rule", E.canonical_contract_snapshot(parsed))
        self.assertFalse([w for w in parsed.parser_warnings if "LANGUAGE_WINDOW" in str(w).upper()],
                         parsed.parser_warnings)


if __name__ == "__main__":
    unittest.main()
