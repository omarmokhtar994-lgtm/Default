# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AF, A1: channel grids at 60 minutes, both ways (owner, 2026-10-12: "everything should be working both
ways").

A grid whose minutes differ from the program's interval was refused. Now a 15, 30 or 60-minute grid is read for any
of those intervals:
- a coarser grid is split: each interval of the program takes its hour's (or half-hour's) number of people;
- a finer grid is joined: each interval of the program takes the average of its rows, as FT Wise counts people per
  interval.
The upload check says which way it went. A grid of other minutes (45, say) is still refused by name.
"""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from webapp.channels import read_channels
from webapp.tests.test_channels import add_channel_tabs

MON = 1


def phone_halves(d, t):
    return 2 if t % 60 == 0 else 4  # 2 on the hour, 4 on the half hour: the hour's average is 3


class GridsBothWays(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())

    def test_an_hourly_grid_in_a_half_hour_program_is_split(self):
        path = add_channel_tabs(self.dir / "hourly.xlsx", step=60, ft=None)
        hourly = read_channels(path, 60)["need"]["P"][MON]
        found = read_channels(path, 30)
        self.assertEqual(found["need"]["P"][MON][10 * 60], hourly[10 * 60])
        self.assertEqual(found["need"]["P"][MON][10 * 60 + 30], hourly[10 * 60])
        self.assertTrue(any("Phone 60 Min" in n and "split" in n for n in found["notes"]), found["notes"])

    def test_a_half_hour_grid_in_an_hourly_program_is_averaged(self):
        path = add_channel_tabs(self.dir / "halves.xlsx", step=30, phone=phone_halves, ft=None)
        found = read_channels(path, 60)
        self.assertEqual(found["need"]["P"][MON][10 * 60], 3)
        self.assertNotIn(10 * 60 + 30, found["need"]["P"][MON])
        self.assertTrue(any("Phone 30 Min" in n and "average" in n for n in found["notes"]), found["notes"])

    def test_the_programs_own_grid_wins_over_another(self):
        path = add_channel_tabs(self.dir / "both.xlsx", step=60, ft=None)
        wb = load_workbook(path)
        extra = wb.copy_worksheet(wb["Phone 60 Min"])
        extra.title = "Phone 30 Min"
        wb.save(path)
        found = read_channels(path, 60)
        self.assertTrue(any("Phone 30 Min is not used" in n for n in found["notes"]), found["notes"])

    def test_other_minutes_are_still_refused_by_name(self):
        path = add_channel_tabs(self.dir / "odd.xlsx", step=60, ft=None)
        wb = load_workbook(path)
        wb["Chat 60 Min"].title = "Chat 45 Min"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, r"Chat 45 Min.*60-minute.*Chat 60 Min"):
            read_channels(path, 60)


if __name__ == "__main__":
    unittest.main()
