# © 2026 Omar Mokhtar. All rights reserved.
"""Phase O task 2: exports for any period: who did what and when.

Shift-dated items (attendance, day activity, breaks, worked hours) follow the
shift's own date; logged actions (schedule changes, versions, runs, the
record) follow when they happened, in Egypt time. Excel holds a tab per item
after an "About this export" tab; CSV holds one item. Text typed by people is
never written as a formula."""
import csv
import io
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from webapp.adherence import person_day
from webapp.attendance import DayBook
from webapp.exports import KINDS, build
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_schedules import AFTER, INPUT, REPO

WED, THU = date(2026, 10, 14), date(2026, 10, 15)


class TheExports(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        cls.lina = store.add_user("lina", "Lina", "Lina-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program="AE/AR B2B",
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.book = ScheduleBook(self.store, self.data, REPO)
        self.days = DayBook(self.store, self.book)
        self.tool = self.book.versions("aaaaaaaaaaaa")[0]
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Late", self.sara, end="13:00")
        self.days.set_status("AE/AR B2B", THU, "Associate 001", "Sick", self.lina)
        self.days.move_break("AE/AR B2B", WED, "Associate 008", 1, "13:05", self.lina)

    def export(self, start=WED, end=WED, kinds=None, fmt="xlsx", **kw):
        data, name, mime = build(self.store, self.days, start, end, kinds=kinds or list(KINDS), fmt=fmt,
                                 by="Omar", **kw)
        return data, name, mime

    def sheet(self, data, title):
        ws = load_workbook(io.BytesIO(data))[title]
        rows = list(ws.iter_rows(values_only=True))
        return [dict(zip(rows[0], r)) for r in rows[1:]]

    def test_one_tab_per_item_after_an_about_tab(self):
        data, name, mime = self.export(kinds=["attendance", "changes"])
        self.assertEqual(load_workbook(io.BytesIO(data)).sheetnames,
                         ["About this export", "Attendance", "Schedule changes and swaps"])
        self.assertEqual((name, mime[:20]), ("Team_Scheduler_2026-10-14_to_2026-10-14.xlsx", "application/vnd.open"))
        about = {r[0]: r[1] for r in load_workbook(io.BytesIO(data))["About this export"].iter_rows(values_only=True)}
        self.assertEqual((about["Period"], about["Downloaded by"]), ("14 Oct 2026 to 14 Oct 2026", "Omar"))
        self.assertIn("Egypt time", about["Times"])

    def test_period_follows_the_shift_date(self):
        data, *_ = self.export(kinds=["attendance"])
        rows = self.sheet(data, "Attendance")
        self.assertEqual([(r["Shift date"], r["Associate"], r["Status"], r["Slot"], r["Shift"], r["Recorded by"])
                          for r in rows], [("2026-10-14", "Associate 001", "Late", "1", "12:00 - 21:00", "Sara")])
        data, *_ = self.export(start=WED, end=THU, kinds=["attendance"])
        self.assertEqual(len(self.sheet(data, "Attendance")), 2)

    def test_who_filter(self):
        data, *_ = self.export(start=WED, end=THU, kinds=["activity"], user_id=self.lina)
        rows = self.sheet(data, "Day activity")
        self.assertEqual({r["By"] for r in rows}, {"Lina"})
        self.assertEqual(sorted(r["What"] for r in rows), ["Lunch moved 12:30 to 13:05", "Sick"])

    def test_breaks_planned_against_taken(self):
        data, *_ = self.export(kinds=["breaks"])
        rows = [r for r in self.sheet(data, "Breaks planned vs taken") if r["Associate"] == "Associate 008"]
        lunch = next(r for r in rows if r["Break"] == "Lunch")
        self.assertEqual((lunch["Planned"], lunch["Taken"], lunch["Moved by"]), ("12:30", "13:05", "Lina"))
        self.assertTrue(all(r["Taken"] == r["Planned"] for r in rows if r["Break"] != "Lunch"))

    def test_schedule_changes_and_versions(self):
        draft = self.book.change(self.tool["id"], self.lina, "Associate 001", "Wed", "10:00 - 19:00", "swap asked")
        self.book.set_in_use(draft, self.sara)
        today = date.today()
        data, *_ = self.export(start=today, end=today, kinds=["changes", "versions"])
        change = self.sheet(data, "Schedule changes and swaps")[0]
        self.assertEqual((change["Associate"], change["Day"], change["Old"], change["New"], change["Reason"], change["By"]),
                         ("Associate 001", "Wed", "12:00 - 21:00", "10:00 - 19:00", "swap asked", "Lina"))
        events = [(r["What"], r["By"]) for r in self.sheet(data, "Versions and in use")]
        self.assertEqual(events, [("New version", "Lina"), ("Set in use", "Sara")])

    def test_worked_hours_match_the_adherence_tab(self):
        data, *_ = self.export(kinds=["worked"])
        row = next(r for r in self.sheet(data, "Worked hours and adherence") if r["Associate"] == "Associate 001")
        page = self.days.page("AE/AR B2B", WED)
        expected = person_day(page["view"], "Associate 001")
        self.assertEqual((row["Late min"], row["Adherence %"], row["Conformance %"]),
                         (expected["late"], expected["adherence"], expected["conformance"]))

    def test_text_is_never_a_formula(self):
        self.book.change(self.tool["id"], self.lina, "Associate 002", "Sun", "OFF", '=HYPERLINK("http://x","click")')
        today = date.today()
        data, *_ = self.export(start=today, end=today, kinds=["changes"])
        ws = load_workbook(io.BytesIO(data))["Schedule changes and swaps"]
        cell = next(c for row in ws.iter_rows() for c in row if str(c.value).startswith("=HYPERLINK"))
        self.assertEqual(cell.data_type, "s")
        data, name, mime = self.export(start=today, end=today, kinds=["changes"], fmt="csv")
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows[0]["Reason"], """'=HYPERLINK("http://x","click")""")
        self.assertEqual((name, mime), ("Team_Scheduler_changes_{0}_to_{0}.csv".format(today.isoformat()), "text/csv"))

    def test_refusals(self):
        with self.assertRaises(ValueError):  # longer than 400 days
            self.export(start=date(2025, 1, 1), end=date(2026, 3, 1))
        with self.assertRaises(ValueError):  # the end before the start
            self.export(start=THU, end=WED)
        with self.assertRaises(ValueError):  # CSV holds one item
            self.export(kinds=["attendance", "activity"], fmt="csv")
        with self.assertRaises(ValueError):
            self.export(kinds=["everything"])


if __name__ == "__main__":
    unittest.main()
