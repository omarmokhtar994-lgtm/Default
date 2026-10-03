#!/usr/bin/env python3
"""Build small input workbooks from a Python spec, from scratch (no template).

Used by the production-readiness audit to probe the parser and the solver with
inputs designed to break them. Every workbook is written with openpyxl only;
nothing is copied from a shipped workbook.

spec keys (all optional except roster/demand):
  instructions: {label: value}         -> Instructions sheet (Section, Instruction, Value)
  roster: [(name, language)]           -> Schedule sheet (Slot, Emp ID, Email, SF Name, TL, Language, Sun..Sat)
  roster_rows: list of raw row lists   -> overrides roster layout completely (for layout probes)
  schedule_header: list                -> header row to use with roster_rows
  step: 15|30|60                       -> demand/shrinkage interval
  demand: fn(day, minute) -> value|None
  demand_rows: list of (time, [7 values]) -> raw demand rows (overrides demand)
  shrinkage: float or fn(day, minute)
  shifts: [labels]
  preferences: {name: [7 cells]}
  preference_header: list              -> override header (layout probes)
  fixed: {name: [7 cells]} ; fixed_header: list
  previous_saturday: {name: label}
  language_setup: list of dict(Language, Coverage Start, Coverage End, Can Cover Languages,
                  Active?, Coverage Group, Minimum Per Interval, Coverage Days)
  coverage_split: list of dict(Coverage Group, Start, End, Coverage Ratio, Exclusive?, Active?)
"""
from __future__ import annotations

from typing import Any, Dict

from openpyxl import Workbook

DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]


def hhmm(m: int) -> str:
    m %= 1440
    return f"{m // 60:02d}:{m % 60:02d}"


def build(spec: Dict[str, Any], path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws.append(["Section", "Instruction", "Value", "Notes"])
    step = int(spec.get("step", 60))
    base = {
        "Program Name": "AUDIT_CASE",
        "Count of Associates": len(spec.get("roster", [])) or None,
        "Interval Minutes": step,
        "Requirements Source": f"FT Wise {step} Min",
        "Shrinkage Source": f"Shrinkage {step} Min",
        "Target": 0.9,
        "Minimum Per Interval": 0.8,
        "Allowed Shift Durations Hours": 9,
        "Use 11H/3OFF": "No",
        "Strict OFF Count": "Yes",
        "Separate OFF Days": "Yes",
        "Rest Gap Hours": 12,
        "Count of Different Shifts Per week": 3,
        "Fixed Request Use": "No",
        "Hard OFF Preferences": "Yes",
        "Leave": "Yes",
        "Use Preferences": "Yes",
    }
    base.update(spec.get("instructions", {}))
    for key, value in base.items():
        if value is not None:
            ws.append(["Case", key, value, None])

    s = wb.create_sheet("Schedule")
    s.append(["Schedule"])
    if "roster_rows" in spec:
        s.append(spec["schedule_header"])
        for row in spec["roster_rows"]:
            s.append(row)
    else:
        s.append(["Slot", "Emp ID", "Email", "SF Name", "TL", "Language"] + DAYS + ["Notes"])
        for i, (name, lang) in enumerate(spec["roster"]):
            cells = spec.get("schedule_cells", {}).get(name, [None] * 7)
            s.append([i + 1, 100000 + i, f"a{i}@x.test", name, "TL One", lang] + list(cells) + [None])

    lib = wb.create_sheet("Shift Library")
    lib.append(["Allowed Shift / Status", "Start", "End", "Duration Hours"])
    for label in spec.get("shifts", [f"{h:02d}:00 - {(h + 9) % 24:02d}:00" for h in range(0, 24)]):
        lib.append([label, None, None, None])

    shr_value = spec.get("shrinkage", 0.0)
    for title, kind in ((f"FT Wise {step} Min", "demand"), (f"Shrinkage {step} Min", "shrink")):
        g = wb.create_sheet(title)
        g.append([title])
        g.append(["Interval"] + DAYS)
        if kind == "demand" and "demand_rows" in spec:
            for t, vals in spec["demand_rows"]:
                g.append([t] + list(vals))
            continue
        for m in range(0, 1440, step):
            if kind == "demand":
                vals = [spec["demand"](d, m) for d in range(7)]
            else:
                vals = [(shr_value(d, m) if callable(shr_value) else shr_value) for d in range(7)]
            g.append([hhmm(m)] + vals)

    p = wb.create_sheet("Preference")
    p.append(["Preference"])
    p.append(spec.get("preference_header", ["Associate Name", "Language"] + DAYS))
    for name, cells in spec.get("preferences", {}).items():
        p.append([name, None] + list(cells))
    for name, cells in spec.get("preference_rows", []):
        p.append([name, None] + list(cells))

    if "fixed" in spec or "fixed_rows" in spec:
        f = wb.create_sheet("Fixed Request")
        f.append(["Fixed"])
        f.append(spec.get("fixed_header", ["Active?", "Associate Name", "Nesting Group"] + DAYS))
        for name, cells in spec.get("fixed", {}).items():
            f.append(["Yes", name, None] + list(cells))
        for row in spec.get("fixed_rows", []):
            f.append(row)

    if "previous_saturday" in spec:
        v = wb.create_sheet("Previous week scheduled")
        v.append(["Previous"])
        v.append(["Slot", "Name", "Language", "Previous Saturday Shift"])
        for i, (name, label) in enumerate(spec["previous_saturday"].items()):
            v.append([i + 1, name, None, label])

    if "language_setup" in spec:
        cols = ["Language", "Coverage Start", "Coverage End", "Can Cover Languages", "Active?",
                "Coverage Group", "Minimum Per Interval", "Coverage Days"]
        ls = wb.create_sheet("Language Setup")
        ls.append(["Language Setup"])
        ls.append(cols)
        for row in spec["language_setup"]:
            ls.append([row.get(c) for c in cols])

    if "coverage_split" in spec:
        cols = ["Coverage Group", "Start", "End", "Coverage Ratio", "Exclusive?", "Active?"]
        cs = wb.create_sheet("Coverage Split")
        cs.append(cols)
        for row in spec["coverage_split"]:
            cs.append([row.get(c) for c in cols])
    wb.save(path)
