# © 2026 Omar Mokhtar. All rights reserved.
"""Phase V: Chat, Phone and Email planned per associate (owner, 2026-10-09).

Task 1: the channel tabs of the input workbook ("Chat {n} Min", "Phone {n} Min", "Email {n} Min", "Email Hours",
"Channel Setup"), read strictly: a cell the website cannot use is refused with its tab and row, never read as 0.
The check said on upload compares the channel tabs with FT Wise (owner: "3 channels should be covered by the
requirements"), except in all-channels times (owner: "during some intervals we might have 1-2 associates covering
all channels so i need an option to tell the same to the tool to avoid invalid warning")."""
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from webapp.channels import blended_at, check_lines, has_channel_tabs, read_channels, requirement_tab
from webapp.day import read_inputs
from webapp.tests.test_schedules import INPUT

DAYS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
SETTINGS = [("Minimum block minutes", 60), ("Phone maximum continuous minutes", 120),
            ("Chat maximum continuous minutes", 0), ("Email maximum continuous minutes", 0),
            ("Priority when short", "Phone > Chat > Email"), ("How strict is the order", "Strict"),
            ("Spread channels fairly over the week", "Yes")]
LANGUAGES = [("Phone", "Arabic", 1, "10:00", "20:00", "All", "Yes"),
             ("Chat", "Arabic", 1, "12:00", "20:00", "Sun, Mon, Tue, Wed, Thu", "Yes"),
             ("Email", "English", 2, "08:00", "22:00", "All", "No")]
BLENDED = [("Fri, Sat", "20:00", "22:00", "Yes", "Weekend evenings")]
HOURS = [(d, 10, {"Arabic": 3}, "08:00", "22:00") for d in DAYS[:5]] + [("Fri", 20, {}, "08:00", "22:00"),
                                                                        ("Sat", 0, {}, "", "")]


def ft_grid(d, t):
    """FT Wise used by the check tests: 6 people from 08:00 to 22:00, 2 in the weekend's all-channels evening."""
    if not 8 * 60 <= t < 22 * 60:
        return 0
    return 2 if d in (5, 6) and t >= 20 * 60 else 6


def phone_grid(d, t):
    if not 8 * 60 <= t < 22 * 60:
        return 0
    if d == 4 and t == 13 * 60:
        return 5  # Thursday 13:00: Chat + Phone above FT Wise
    if d == 6 and t == 21 * 60:
        return 3  # Saturday 21:00, all-channels time: Phone alone above FT Wise
    return 2 if d in (5, 6) and t >= 20 * 60 else 3


def chat_grid(d, t):
    if not 8 * 60 <= t < 22 * 60:
        return 0
    return 1 if d in (5, 6) and t >= 20 * 60 else 2


def _grid_tab(wb, name, value, step):
    ws = wb.create_sheet(name)
    ws["A1"] = name
    ws["A2"] = "People needed per interval."
    ws.append([])
    ws.append(["Interval"] + DAYS)
    for t in range(0, 1440, step):
        ws.append([f"{t // 60:02d}:{t % 60:02d}"] + [value(d, t) if value else None for d in range(7)])
    return ws


def add_channel_tabs(dst, *, step=60, phone=phone_grid, chat=chat_grid, email=None, email_tab=True, hours=HOURS,
                     settings=SETTINGS, languages=LANGUAGES, blended=BLENDED, ft=ft_grid, src=INPUT):
    """A copy of the B3 input workbook with the channel tabs laid out as in the approved sample workbook.
    ``None`` leaves a tab out; ``email_tab`` keeps an empty Email tab; ``ft`` rewrites FT Wise's grid."""
    wb = load_workbook(src)
    if ft is not None:
        ws = wb[f"FT Wise {step} Min"]
        head = next(r for r in range(1, 10) if str(ws.cell(r, 1).value or "").strip().lower() == "interval")
        cols = {str(ws.cell(head, c).value).strip(): c for c in range(2, ws.max_column + 1)
                if str(ws.cell(head, c).value or "").strip() in DAYS}
        for r in range(head + 1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v is None:
                continue
            t = v.hour * 60 + v.minute if hasattr(v, "hour") else int(str(v)[:2]) * 60 + int(str(v)[3:5])
            for day, c in cols.items():
                ws.cell(r, c).value = ft(DAYS.index(day), t)
    if phone is not None:
        _grid_tab(wb, f"Phone {step} Min", phone, step)
    if chat is not None:
        _grid_tab(wb, f"Chat {step} Min", chat, step)
    if email is not None or email_tab:
        _grid_tab(wb, f"Email {step} Min", email, step)
    if hours is not None:
        ws = wb.create_sheet("Email Hours")
        ws["A1"] = "Email Hours"
        ws.append(["Hours of email each day."])
        ws.append([])
        langs = sorted({lang for _, _, by, _, _ in hours for lang in by})
        ws.append(["Day", "Hours needed"] + [f"Of which {lang}" for lang in langs] + ["Window start", "Window end"])
        for day, h, by, lo, hi in hours:
            ws.append([day, h] + [by.get(lang) for lang in langs] + [lo, hi])
    if settings is not None or languages is not None or blended is not None:
        ws = wb.create_sheet("Channel Setup")
        ws["A1"] = "Channel Setup"
        ws.append(["The rotation rules, the language minimums per channel and the all-channels times."])
        ws.append([])
        ws.append(["Setting", "Value", "What it means"])
        for name, value in settings or []:
            ws.append([name, value, ""])
        ws.append([])
        ws.append(["Channel", "Language", "Minimum per interval", "Coverage start", "Coverage end", "Coverage days",
                   "Active?"])
        for row in languages or []:
            ws.append(list(row))
        ws.append([])
        ws.append(["All channels together"])
        ws.append(["Times when the 1 or 2 people on the floor cover every channel they can work at once."])
        ws.append(["Days", "Start", "End", "Active?", "Note"])
        for row in blended or []:
            ws.append(list(row))
    wb.save(dst)
    return dst


def unnamed(src, dst, rename=True):
    """A copy whose Instructions no longer name the requirement and, with ``rename``, whose FT Wise tabs carry no
    name the engine knows."""
    wb = load_workbook(src)
    for row in wb["Instructions"].iter_rows():
        for cell in row:
            if str(cell.value or "").strip().lower() == "requirements source":
                for other in row[cell.column:]:
                    other.value = None
    if rename:
        for n in (15, 30, 60):
            wb[f"FT Wise {n} Min"].title = f"Demand {n}"
    wb.save(dst)
    return dst


def setup_row(path, first_cell):
    """The Channel Setup row whose first cell reads ``first_cell`` (as the error must name it)."""
    ws = load_workbook(path)["Channel Setup"]
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == first_cell)


class TheChannelTabs(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir, True)

    def book(self, name="channels.xlsx", **kw):
        return add_channel_tabs(self.dir / name, **kw)

    def test_a_workbook_without_channel_tabs_has_none(self):
        self.assertIsNone(read_channels(INPUT, 60))

    def test_reads_needs_email_hours_rules_languages_and_all_channels_times(self):
        found = read_channels(self.book(), 60)
        self.assertEqual(found["need"]["P"][3][16 * 60], 3)
        self.assertEqual(found["need"]["C"][3][16 * 60], 2)
        self.assertEqual(found["need"]["P"][4][13 * 60], 5)
        self.assertNotIn("E", found["need"])  # the Email tab is empty: Email Hours is used
        self.assertEqual(found["email_mode"], "hours")
        self.assertEqual(found["email_hours"][3], {"hours": 10.0, "languages": {"Arabic": 3.0}, "start": 480,
                                                   "end": 1320})
        self.assertEqual(found["email_hours"][6], {"hours": 0.0, "languages": {}, "start": 0, "end": 1440})
        self.assertEqual(found["rules"], {"min_block": 60, "max_run": {"P": 120, "C": 0, "E": 0},
                                          "order": ["P", "C", "E"], "strict": True, "fair": True})
        self.assertEqual(found["languages"][1], {"channel": "C", "language": "Arabic", "minimum": 1, "start": 720,
                                                 "end": 1200, "days": {0, 1, 2, 3, 4}, "active": True})
        self.assertFalse(found["languages"][2]["active"])
        self.assertEqual(found["blended"], [{"days": {5, 6}, "start": 1200, "end": 1320}])

    def test_an_email_tab_with_numbers_is_used_and_email_hours_is_not(self):
        found = read_channels(self.book(email=lambda d, t: 1 if 9 * 60 <= t < 21 * 60 else 0), 60)
        self.assertEqual(found["email_mode"], "interval")
        self.assertEqual(found["need"]["E"][2][10 * 60], 1)
        self.assertTrue(any("Email Hours is not used" in n for n in found["notes"]))

    def test_without_channel_setup_the_defaults_are_used_and_said(self):
        found = read_channels(self.book(settings=None, languages=None, blended=None), 60)
        self.assertEqual(found["rules"], {"min_block": 60, "max_run": {"P": 0, "C": 0, "E": 0},
                                          "order": ["P", "C", "E"], "strict": True, "fair": True})
        self.assertEqual((found["languages"], found["blended"]), ([], []))
        self.assertTrue(any(n.startswith("No Channel Setup tab") for n in found["notes"]))

    def test_a_cell_it_cannot_use_is_refused_with_its_tab_and_row(self):
        bad_grid = lambda value: (lambda d, t: value if (d, t) == (2, 10 * 60) else phone_grid(d, t))  # noqa: E731
        cases = {
            "negative need": ({"phone": bad_grid(-1)}, r"Phone 60 Min, row 15: .*Tue 10:00"),
            "word in a grid": ({"chat": lambda d, t: "lots" if (d, t) == (2, 10 * 60) else 1},
                               r"Chat 60 Min, row 15: .*Tue 10:00"),
            "unknown channel": ({"languages": [("Fax", "Arabic", 1, "10:00", "20:00", "All", "Yes")]},
                                r"Channel Setup, row {Fax}: .*Fax"),
            "bad time": ({"languages": [("Phone", "Arabic", 1, "25:99", "20:00", "All", "Yes")]},
                         r"Channel Setup, row {Phone}: .*25:99"),
            "bad day": ({"languages": [("Phone", "Arabic", 1, "10:00", "20:00", "Funday", "Yes")]},
                        r"Channel Setup, row {Phone}: .*Funday"),
            "minimum not whole": ({"languages": [("Phone", "Arabic", 1.5, "10:00", "20:00", "All", "Yes")]},
                                  r"Channel Setup, row {Phone}: .*whole number"),
            "order not a permutation": ({"settings": [("Priority when short", "Phone > Phone > Email")]},
                                        r"Channel Setup, row {Priority when short}: .*each channel once"),
            "strict misspelt": ({"settings": [("How strict is the order", "Strictly")]},
                                r"Channel Setup, row {How strict is the order}: .*Strict or Balanced"),
            "block not a quarter": ({"settings": [("Minimum block minutes", 50)]},
                                    r"Channel Setup, row {Minimum block minutes}: .*15"),
            "unknown setting": ({"settings": [("Lunch rule", 3)]}, r"Channel Setup, row {Lunch rule}: .*Lunch rule"),
            "all-channels day": ({"blended": [("Someday", "20:00", "22:00", "Yes", "")]},
                                 r"Channel Setup, row {Someday}: .*Someday"),
            "email hours day": ({"hours": [("Someday", 4, {}, "08:00", "22:00")]}, r"Email Hours, row 5: .*Someday"),
            "of which above hours": ({"hours": [("Mon", 4, {"Arabic": 5}, "08:00", "22:00")]},
                                     r"Email Hours, row 5: .*Arabic"),
        }
        for name, (kw, pattern) in cases.items():
            with self.subTest(name):
                path = self.book(f"{name.replace(' ', '_')}.xlsx", **kw)
                for first in ("Fax", "Phone", "Priority when short", "How strict is the order",
                              "Minimum block minutes", "Lunch rule", "Someday"):
                    if "{" + first + "}" in pattern:
                        pattern = pattern.replace("{" + first + "}", str(setup_row(path, first)))
                with self.assertRaisesRegex(ValueError, pattern):
                    read_channels(path, 60)

    def test_a_tab_for_another_interval_is_refused_by_name(self):
        path = self.book()
        wb = load_workbook(path)
        wb["Chat 60 Min"].title = "Chat 30 Min"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, r"Chat 30 Min.*60-minute.*Chat 60 Min"):
            read_channels(path, 60)

    def test_a_file_that_is_not_a_workbook_is_left_to_the_run_check(self):
        # The upload only looks for channel tabs; a file openpyxl cannot open is reported by the run's own check,
        # as before Phase V (webapp.tests.test_runs uploads such files).
        fake = self.dir / "fake.xlsx"
        fake.write_bytes(b"PK\x03\x04 fake workbook")
        self.assertFalse(has_channel_tabs(fake))

    def test_all_channels_times_by_day_and_across_midnight(self):
        found = read_channels(self.book(blended=[("Fri, Sat", "20:00", "22:00", "Yes", ""),
                                                 ("Sat", "23:00", "02:00", "Yes", ""),
                                                 ("Mon", "06:00", "07:00", "No", "")]), 60)
        self.assertTrue(blended_at(found, 5, 21 * 60))
        self.assertFalse(blended_at(found, 5, 19 * 60))
        self.assertTrue(blended_at(found, 6, 23 * 60 + 30))
        self.assertTrue(blended_at(found, 0, 60))  # Sunday 01:00 is Saturday night's window
        self.assertFalse(blended_at(found, 6, 60))  # Saturday 01:00 is not (Friday has no night window)
        self.assertFalse(blended_at(found, 1, 6 * 60 + 30))  # an inactive row is not used


class TheRequirementTab(unittest.TestCase):
    """The engine reads the tab Requirements Source names, else one of its own names, else the first Sun to Sat
    grid it finds: with channel tabs that could be a channel tab, so the website insists on a named tab."""

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir, True)

    def test_by_instruction_by_engine_name_or_none(self):
        self.assertEqual(requirement_tab(INPUT), "FT Wise 60 Min")
        # Without the instruction the engine takes its first alias found, "FT Wise 30 Min", even in this 60-minute
        # workbook: the website reports the tab the engine would really read.
        self.assertEqual(requirement_tab(unnamed(INPUT, self.dir / "alias.xlsx", rename=False)), "FT Wise 30 Min")
        self.assertIsNone(requirement_tab(unnamed(INPUT, self.dir / "none.xlsx")))


class TheUploadCheck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.path = add_channel_tabs(cls.dir / "channels.xlsx")
        cls.lines = check_lines(read_channels(cls.path, 60), read_inputs(cls.path), cls.path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def said(self, level):
        return [text for lvl, text in self.lines if lvl == level]

    def test_the_requirement_and_the_tabs_read(self):
        self.assertIn("FT Wise 60 Min is the schedule's requirement (named in Requirements Source), so a channel "
                      "tab is never read as the requirement.", self.said("ok"))
        self.assertIn("Read Phone 60 Min, Chat 60 Min and Email Hours. Email 60 Min is empty, so email follows "
                      "Email Hours.", self.said("ok"))

    def test_chat_and_phone_against_ft_wise_outside_all_channels_times(self):
        # 7 days x 14 hours, less Friday's and Saturday's 2 all-channels hours each
        self.assertIn("Chat + Phone fit inside FT Wise 60 Min in 93 of 94 intervals.", self.said("ok"))
        self.assertIn("Thu 13:00: Chat 2 + Phone 5 = 7, but FT Wise 60 Min needs 6. The channel plan follows the "
                      "channel tabs; the schedule keeps FT Wise 60 Min.", self.said("warn"))
        self.assertFalse([t for t in self.said("warn") if t.startswith(("Fri 20:00", "Fri 21:00", "Sat 20:00"))])

    def test_all_channels_times_are_not_added_up_but_each_channel_is_checked(self):
        self.assertIn("All-channels times (Channel Setup): Fri, Sat 20:00 to 22:00. 4 intervals there are not added "
                      "up, so 4 warnings that would have been false are not raised. Checked instead: no single "
                      "channel needs more people than FT Wise 60 Min.", self.said("info"))
        self.assertIn("Sat 21:00 (all channels): Phone alone needs 3, but FT Wise 60 Min needs 2.", self.said("warn"))

    def test_email_hours_against_what_ft_wise_leaves_each_day(self):
        self.assertIn("Fri: FT Wise 60 Min leaves 12 h after Chat and Phone; Email Hours asks for 20 h. The other "
                      "8 h fits only where more people turn up than needed.", self.said("warn"))
        self.assertIn("Email Hours fit inside FT Wise 60 Min on 5 of 6 days.", self.said("ok"))

    def test_no_named_requirement_is_said_as_a_warning(self):
        lines = check_lines(read_channels(self.path, 60), read_inputs(self.path),
                            unnamed(self.path, self.dir / "unnamed.xlsx"))
        self.assertTrue(any(lvl == "warn" and text.startswith("No requirement tab is named") for lvl, text in lines))


if __name__ == "__main__":
    unittest.main()


class TheChannelUpload(unittest.TestCase):
    """Task 2: an upload with channel tabs is refused when the website cannot use them, or when the requirement tab
    is not named (the engine could then read a channel tab as its requirement); otherwise it is kept and its check
    is listed on the schedules page. A workbook without channel tabs is uploaded exactly as before."""

    @classmethod
    def setUpClass(cls):
        import html as _html
        from webapp.programs import ProgramBook
        from webapp.tests.test_ready import make_ready
        from webapp.tests.test_runs import REPO, make_app, sign_in, token
        cls.unescape, cls.token = staticmethod(_html.unescape), staticmethod(token)
        cls.dir = Path(tempfile.mkdtemp())
        cls.ready = make_ready(cls.dir / "ready.xlsx")
        cls.good = add_channel_tabs(cls.dir / "good.xlsx", src=cls.ready)
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        book = ProgramBook(cls.store)
        cls.key = book.add_lob(book.add_program("SAKS"), "NMG Tier 2")
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def upload(self, path, name="week.xlsx"):
        import io
        return self.admin.post("/runs", data={"csrf_token": self.token(self.admin), "kind": "ready", "mode": "QUICK",
                                              "program": self.key, "week_start": "2026-10-11",
                                              "workbook": (io.BytesIO(Path(path).read_bytes()), name)},
                               content_type="multipart/form-data", follow_redirects=True)

    def runs(self):
        return len(self.store.list_runs())

    def run_named(self, name):
        from webapp.tests.test_runs import wait
        run_id = next(r["id"] for r in self.store.list_runs() if r["workbook"] == name)
        wait(self.store, run_id, statuses=("DONE", "REJECTED", "FAILED"))
        return run_id

    def test_an_unnamed_requirement_is_refused_and_nothing_is_kept(self):
        before = self.runs()
        page = self.unescape(self.upload(unnamed(self.good, self.dir / "unnamed.xlsx")).get_data(as_text=True))
        self.assertIn("This workbook has channel tabs but does not name its requirement tab", page)
        self.assertIn("Requirements Source", page)
        self.assertEqual(self.runs(), before)
        self.assertFalse([p for p in (self.app.extensions["runs"].runs_root / "_incoming").glob("*")])

    def test_a_channel_setup_row_it_cannot_use_is_refused_naming_the_row(self):
        bad = add_channel_tabs(self.dir / "bad.xlsx", src=self.ready,
                               languages=[("Fax", "Arabic", 1, "10:00", "20:00", "All", "Yes")])
        page = self.unescape(self.upload(bad).get_data(as_text=True))
        self.assertIn(f"The channel tabs need fixing before this upload: Channel Setup, row {setup_row(bad, 'Fax')}: "
                      "Fax is not a channel (Chat, Phone or Email).", page)

    def test_a_good_upload_says_its_warnings_and_the_schedules_page_lists_them(self):
        page = self.unescape(self.upload(self.good, "good_week.xlsx").get_data(as_text=True))
        self.assertIn("Channel tabs: 3 warnings. They are listed on the schedules page and never change the "
                      "schedule.", page)
        run_id = self.run_named("good_week.xlsx")
        page = self.unescape(self.admin.get(f"/runs/{run_id}/schedules").get_data(as_text=True))
        self.assertIn("Channel tabs", page)
        self.assertIn("Thu 13:00: Chat 2 + Phone 5 = 7, but FT Wise 60 Min needs 6.", page)
        self.assertIn("Sat 21:00 (all channels): Phone alone needs 3, but FT Wise 60 Min needs 2.", page)

    def test_a_workbook_without_channel_tabs_shows_no_channel_panel(self):
        page = self.unescape(self.upload(self.ready, "plain_week.xlsx").get_data(as_text=True))
        self.assertNotIn("Channel tabs:", page)
        run_id = self.run_named("plain_week.xlsx")
        page = self.unescape(self.admin.get(f"/runs/{run_id}/schedules").get_data(as_text=True))
        self.assertNotIn("Channel tabs", page)


class TheChannelPlanTab(unittest.TestCase):
    """Task 4: a version keeps its channel plan on a "Channel Plan" tab (one row per block); the week read from it
    carries the blocks, and blocks that no longer fit the person's shift (after a shift edit or a swap) are named
    for re-planning instead of being counted."""

    @classmethod
    def setUpClass(cls):
        from webapp.tests.test_ready import make_ready
        cls.dir = Path(tempfile.mkdtemp())
        cls.ready = make_ready(cls.dir / "ready.xlsx")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def test_blocks_are_written_per_person_and_day_and_read_back(self):
        from webapp.versions import read_week, write_channels
        first, second = self.dir / "first.xlsx", self.dir / "second.xlsx"
        write_channels(self.ready, first, {("Associate 001", "Mon"): ("10:00 - 19:00", [(600, 720, "P"), (720, 840, "C")]),
                                           ("Associate 002", "Mon"): ("10:00 - 19:00", [(600, 1140, "A")])})
        write_channels(first, second, {("Associate 002", "Mon"): ("10:00 - 19:00", [(600, 900, "E")])})
        self.assertEqual(read_week(second)["channels"], [
            {"associate": "Associate 001", "day": "Mon", "start": "10:00", "end": "12:00", "channel": "P"},
            {"associate": "Associate 001", "day": "Mon", "start": "12:00", "end": "14:00", "channel": "C"},
            {"associate": "Associate 002", "day": "Mon", "start": "10:00", "end": "15:00", "channel": "E"}])
        ws = load_workbook(second)["Channel Plan"]
        self.assertEqual([c.value for c in ws[1]], ["Associate", "Day", "Shift", "Start", "End", "Channel"])
        self.assertEqual(ws.cell(2, 6).value, "Phone")
        self.assertEqual(read_week(self.ready)["channels"], [])

    def test_the_independent_validator_still_passes_a_version_with_the_tab(self):
        from webapp.tests.test_schedules import REPO
        from webapp.versions import validate, write_channels
        planned = self.dir / "planned.xlsx"
        write_channels(self.ready, planned, {("Associate 001", "Mon"): ("10:00 - 19:00", [(600, 1140, "P")])})
        before, after = validate(self.ready, self.ready, REPO), validate(planned, planned, REPO)
        self.assertEqual(after["status"], before["status"])
        self.assertEqual(sorted(f["type"] for f in after["failures"]), sorted(f["type"] for f in before["failures"]))

    def test_blocks_in_minutes_overnight_and_stale_ones_named(self):
        from webapp.versions import channel_blocks, stale_blocks
        week = {"associates": [{"name": "Associate 005", "days": ["OFF"] * 3 + ["22:00 - 07:00"] + ["OFF"] * 3},
                               {"name": "Associate 006", "days": ["OFF"] * 3 + ["12:00 - 21:00"] + ["OFF"] * 3},
                               {"name": "Associate 007", "days": ["OFF"] * 7}],
                "channels": [{"associate": "Associate 005", "day": "Wed", "start": "22:00", "end": "02:00", "channel": "P"},
                             {"associate": "Associate 005", "day": "Wed", "start": "02:00", "end": "07:00", "channel": "C"},
                             {"associate": "Associate 006", "day": "Wed", "start": "10:00", "end": "12:00", "channel": "P"},
                             {"associate": "Associate 006", "day": "Wed", "start": "12:00", "end": "16:00", "channel": "C"},
                             {"associate": "Associate 007", "day": "Wed", "start": "10:00", "end": "12:00", "channel": "E"}]}
        self.assertEqual(channel_blocks(week, 3, "Associate 005"),
                         [{"start": 1320, "end": 1560, "channel": "P"}, {"start": 1560, "end": 1860, "channel": "C"}])
        self.assertEqual(channel_blocks(week, 3, "Associate 006"), [{"start": 720, "end": 960, "channel": "C"}])
        self.assertEqual(channel_blocks(week, 3, "Associate 007"), [])
        self.assertEqual(stale_blocks(week, 3), [
            "Associate 006, Wed: part of the channel plan is outside the shift 12:00 - 21:00 (made for another shift); "
            "plan channels again for Wed.",
            "Associate 007, Wed: the channel plan has times, but the day is OFF; plan channels again for Wed."])
