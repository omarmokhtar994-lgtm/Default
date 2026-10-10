# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AD (owner, 2026-10-11: "I need option to delete break as well totally if we are going to cancel a break for
someone"; sample 03): a break can be cancelled with a reason; the person stays on the floor; the reason stays on the
site and in exports, never in the group post; Back to plan brings it back; Fix breaks treats the long gap a cancel
leaves as wanted. Made-up "Associate NN" data, Wednesday 14 Oct."""
import html
import io
import re
import unittest

from openpyxl import load_workbook

from webapp import day
from webapp.exports import build
from webapp.tests.test_notify_daybook import WED, _Tagged, hm
from webapp.tests.test_runs import token

WHY = "Queue is long, agreed with the team lead"


class _Cancel(_Tagged):
    def seg_of(self, name):
        lane = next(l for l in self.days.page(self.key, WED)["view"]["lanes"] if l["name"] == name)
        return next(s for s in lane["segments"] if s["offset"] == 0)

    def cell_at(self, minute):
        cells = self.days.page(self.key, WED)["view"]["cells"]
        return next(c for c in cells if c["t"] <= minute < c["t"] + 60)

    def cancelled(self, skip=()):
        name, idx, kind, start, _ = self.with_breaks(skip)
        self.addCleanup(self.days.move_break, self.key, WED, name, idx, None, self.omar)
        self.days.cancel_break(self.key, WED, name, idx, WHY, self.omar)
        return name, idx, kind, start


class TheCancel(_Cancel):
    def test_cancel_keeps_them_on_the_floor(self):
        name, idx, kind, start, _ = self.with_breaks()
        before = self.cell_at(start)["now"]
        self.addCleanup(self.days.move_break, self.key, WED, name, idx, None, self.omar)
        self.days.cancel_break(self.key, WED, name, idx, WHY, self.omar)
        seg = self.seg_of(name)
        self.assertNotIn(idx, [b["idx"] for b in seg["breaks"]])
        self.assertEqual([(b["idx"], b["kind"]) for b in seg["cancelled"]], [(idx, kind)])
        self.assertGreater(self.cell_at(start)["now"], before)
        self.assertEqual(self.logged()[0]["what"], f"{kind} cancelled ({hm(start)}): {WHY}")  # on the site, like an aux's why

    def test_cancel_needs_a_reason(self):
        name, idx, *_ = self.with_breaks()
        with self.assertRaises(ValueError) as said:
            self.days.cancel_break(self.key, WED, name, idx, "  ", self.omar)
        self.assertEqual(str(said.exception), "Say why the break is cancelled.")

    def test_started_break_cannot_be_cancelled(self):
        name, idx, kind, start, _ = self.with_breaks()
        with self.assertRaises(ValueError) as said:
            self.days.cancel_break(self.key, WED, name, idx, WHY, self.omar, now=start + 1)
        self.assertEqual(str(said.exception), "That break has started: it cannot be cancelled.")
        self.days.cancel_break(self.key, WED, name, idx, WHY, self.omar, now=start - 5)  # not started yet
        self.days.move_break(self.key, WED, name, idx, None, self.omar)

    def test_back_to_plan_brings_it_back(self):
        name, idx, kind, start = self.cancelled()
        self.days.move_break(self.key, WED, name, idx, None, self.omar)
        seg = self.seg_of(name)
        self.assertIn(idx, [b["idx"] for b in seg["breaks"]])
        self.assertEqual(seg["cancelled"], [])

    def test_group_post_has_no_reason(self):
        name, idx, kind, start = self.cancelled()
        item = self.items()[0]
        self.assertEqual((item["kind"], item["text"]), ("break", f"{kind} at {hm(start)} cancelled"))
        self.assertNotIn("Queue", item["text"])

    def test_breaks_export_says_cancelled(self):
        name, idx, kind, start = self.cancelled()
        data, *_ = build(self.store, self.days, WED, WED, ["breaks"], program=self.key, by="Omar")
        ws = load_workbook(io.BytesIO(data))["Breaks planned vs taken"]
        rows = list(ws.iter_rows(values_only=True))
        found = [dict(zip(rows[0], r)) for r in rows[1:]]
        row = next(r for r in found if r["Associate"] == name and r["Break"] == kind)
        self.assertEqual((row["Taken"], row["Cancelled because"]), ("Cancelled", WHY))

    def test_route_answers_like_a_move(self):
        name, idx, kind, start, _ = self.with_breaks()
        form = {"csrf_token": token(self.client), "program": self.key, "date": WED.isoformat(), "associate": name,
                "idx": str(idx), "why": ""}
        got = self.client.post("/day/break/cancel", data=form)
        self.assertEqual(got.status_code, 400)
        self.assertEqual(got.get_json()["error"], "Say why the break is cancelled.")
        self.addCleanup(self.days.move_break, self.key, WED, name, idx, None, self.omar)
        got = self.client.post("/day/break/cancel", data={**form, "why": WHY})
        self.assertEqual(got.get_json(), {"ok": True})


class ThePage(_Cancel):
    def page(self):
        got = self.client.get(f"/day?program={self.key.replace(' ', '+')}&date={WED.isoformat()}")
        return html.unescape(got.get_data(as_text=True))

    def test_dialog_has_cancel_and_close(self):
        dialog = re.search(r'(?s)<dialog id="break-dialog".*?</dialog>', self.page()).group(0)
        self.assertIn('<button class="danger" type="button" data-cancel-break>Cancel break</button>', dialog)
        self.assertRegex(dialog, r'<input[^>]*name="why"')
        self.assertIn(">Close</button>", dialog)
        self.assertNotIn(">Cancel</button>", dialog)  # "Cancel" never means two things (sample 03)

    def test_cancelled_break_on_the_timeline(self):
        name, idx, kind, start = self.cancelled()
        mark = re.search(rf'<rect class="brk cancelled"[^>]*data-name="{name}"[^>]*data-idx="{idx}"[^>]*>', self.page())
        self.assertIsNotNone(mark)
        self.assertIn('data-cancelled="1"', mark.group(0))


class TheGapRule(unittest.TestCase):  # pure: no fixture day needed
    def seg(self, cancelled):
        breaks = [{"idx": 0, "kind": "Break 1", "minutes": 15, "start": 600},
                  {"idx": 2, "kind": "Break 2", "minutes": 15, "start": 900}]
        lunch = {"idx": 1, "kind": "Lunch", "minutes": 30, "start": 750}
        return {"start": 480, "end": 1020, "breaks": breaks if cancelled else sorted(breaks + [lunch],
                                                                                      key=lambda b: b["idx"]),
                "cancelled": [lunch] if cancelled else []}

    def test_fix_breaks_moves_around_a_cancelled_break(self):
        seg = self.seg(cancelled=True)
        starts = {b["idx"]: b["start"] for b in seg["breaks"]}
        b2 = seg["breaks"][-1]
        self.assertEqual(day.break_window(seg, b2, starts, 0, 60, 210), (675, 1005))  # no 210-minute cap after B1
        seg = self.seg(cancelled=False)
        starts = {b["idx"]: b["start"] for b in seg["breaks"]}
        self.assertEqual(day.break_window(seg, seg["breaks"][-1], starts, 0, 60, 210), (840, 990))  # after lunch


if __name__ == "__main__":
    import unittest
    unittest.main()
