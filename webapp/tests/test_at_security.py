# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AT (owner, 2026-10-21: "all rules matching in all tabs; security high across every rule; nobody from another
program may see anything of it; no small change made by hand in the website's script may sign in as another account or
the like"). Everything a person can edit in the browser (the cookie, a hidden field, a form value, an address, the script)
is untrusted: the server decides from the account it loaded, never from what the page sent.

Four things are checked from outside, the way a person with a text editor would:
 1. the role rules: every route, for every role, against what the Roles matrix and the admin-only decorators say;
 2. the session: an edited, forged, replayed or old cookie signs nobody in;
 3. the forms: fields that name a role, a program, a person or a request never reach beyond the account's own rights;
 4. the answers: every response carries the security headers, escapes what people typed, and never leaves the site."""
import base64
import html
import io
import json
import os
import re
import stat
import tempfile
import unittest
import zlib
from pathlib import Path

from openpyxl import load_workbook

from webapp.auth import ASSOCIATE_PAGES
from webapp.exports import KINDS
from webapp.permissions import ENDPOINTS, allowed
from webapp.tests import test_ar_walls as fx
from webapp.tests.test_runs import make_app, sign_in, token

APP_SOURCE = (Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8")
ADMIN_ONLY = set(re.findall(r"@admin_required\n\s+def (\w+)\(", APP_SOURCE))
PASSWORD = fx.PASSWORD


class _Secured(fx._Walls):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from webapp.programs import ProgramBook  # noqa: F401
        store = cls.store
        admin = cls.admin_id
        # BRAVO's records in every table that lists people or what they did, so a leak has something to leak
        for program, name in ((cls.bravo, "Zorro Person One"), (cls.alpha, "Alpha Person One")):
            store.set_attendance(program=program, shift_date="2026-10-14", associate=name, status="Sick", from_min=None,
                                 to_min=None, billable=0, user_id=admin, at=1.0)
            store.add_activity(program=program, shift_date="2026-10-14", associate=name, kind="Coaching", start=600,
                               end_min=630, billable=0, note="", user_id=admin, at=1.0)
            store.add_day_log(program=program, shift_date="2026-10-14", associate=name, what=f"{name} noted",
                              user_id=admin, at=1.0)
            store.add_event(at=1.0, user_id=admin, kind="record", program=program, subject=f"{name} subject",
                            detail=f"{name} detail")
            store.add_leave(program=program, associate=name, user_id=admin, kind="Annual", first_day="2026-10-20",
                            last_day="2026-10-21", note=f"{name} note")

    def ids(self):
        return {"run_id": self.runs[self.alpha], "schedule_id": self.sched[self.alpha], "name": self.alpha,
                "user_id": self.users["plan_a"], "request_id": 1, "swap_id": 1}


# ---------------------------------------------------------------------------------------------- 1. the role rules
class TheRoleRules(_Secured):
    ROLES = ("plan_a", "tm_a", "sup_a")

    def expected_open(self, user, endpoint):
        """Whether this person's account may open ``endpoint``: None when the route decides inside (no claim made)."""
        if user["is_associate"]:
            return endpoint in ASSOCIATE_PAGES
        if user["is_admin"]:
            return True
        if endpoint in ADMIN_ONLY:
            return False
        if endpoint in ENDPOINTS:
            return allowed(self.store, user, ENDPOINTS[endpoint])
        return None

    def sweep(self, who, method):
        user = self.store.get_user(self.users[who])
        client = sign_in(self.app, who, PASSWORD)
        tok = token(client) if not user["is_associate"] else re.search(
            r'name="csrf_token" value="([^"]+)"', client.get("/me").get_data(as_text=True)).group(1)
        wrong, ids = [], self.ids()
        for rule in sorted(self.app.url_map.iter_rules(), key=lambda r: r.rule):
            if method not in rule.methods or rule.endpoint in ("static", "login", "logout", "set_theme",
                                                              "change_password"):
                continue
            claim = self.expected_open(user, rule.endpoint)
            if claim is None:
                continue
            args = {a: ids.get(a, "plan") for a in rule.arguments}
            with self.app.test_request_context():
                from flask import url_for
                url = url_for(rule.endpoint, **args, program=self.alpha) if method == "GET" else url_for(
                    rule.endpoint, **args)
            if method == "GET":
                got = client.get(url)
            else:
                got = client.post(url, data={**fx.FORM, "program": self.alpha, "csrf_token": tok})
            opened = got.status_code not in (403,) and not (user["is_associate"] and got.status_code in (302, 303)
                                                           and "/me" in got.headers.get("Location", "") + "/me")
            if user["is_associate"]:
                opened = got.status_code not in (403, 302, 303) or rule.endpoint in ASSOCIATE_PAGES
            if method == "POST" and claim:
                continue  # a form may be refused inside (only the uploader deletes a run): the claim here is "never more"
            if opened != claim:
                wrong.append((rule.rule, method, got.status_code, "opened" if opened else "refused",
                              "should open" if claim else "should be refused"))
        return wrong

    def test_every_get_route_follows_the_role_matrix_for_every_role(self):
        for who in self.ROLES + ("assoc_a",):
            with self.subTest(who):
                self.assertEqual(self.sweep(who, "GET"), [])

    def test_every_post_route_follows_the_role_matrix_for_every_role(self):
        for who in self.ROLES + ("assoc_a",):
            with self.subTest(who):
                self.assertEqual(self.sweep(who, "POST"), [])

    def test_the_admin_only_pages_are_the_ones_that_name_it(self):
        self.assertEqual(ADMIN_ONLY, {"roles_page", "notifications_page", "notifications_remove", "notifications_test",
                                      "program_rename", "program_setup", "admin_signins_release"})  # + Phase AU

    def test_an_admin_changing_a_role_changes_every_route_for_that_role_at_once(self):
        from webapp.permissions import matrix
        kept = self.store.role_permissions()
        grid = matrix(self.store)
        rows = {(role, key): value for role, keys in grid.items() for key, value in keys.items()}
        try:
            self.assertTrue(allowed(self.store, self.store.get_user(self.users["plan_a"]), "rta_change"))
            self.store.set_role_permissions({**rows, ("Planner", "rta_change"): False})
            self.assertFalse(allowed(self.store, self.store.get_user(self.users["plan_a"]), "rta_change"))
            self.assertEqual(self.sweep("plan_a", "POST"), [])
            self.assertEqual(self.sweep("plan_a", "GET"), [])
            client = sign_in(self.app, "plan_a", PASSWORD)
            got = client.post("/day/attendance", data={**fx.FORM, "program": self.alpha, "csrf_token": token(client)})
            self.assertEqual(got.status_code, 403)
        finally:
            self.store.set_role_permissions(kept)


class TheOffersMatchTheRole(_Secured):
    """Phase AU (owner: "all rules matching in all tabs"): a page offers only the links and forms its reader's role may
    use, so nobody meets a button that the server then refuses."""
    PAGES = ("/", "/overview", "/week", "/schedules", "/requests", "/exports", "/team", "/admin/users", "/programs",
             "/day", "/requests/coaching")

    def offered(self, who):
        client = sign_in(self.app, who, PASSWORD)
        user = self.store.get_user(self.users[who])
        adapter = self.app.url_map.bind("localhost")
        found = []
        pages = ("/me", "/me/breaks") if user["is_associate"] else self.PAGES
        for path in pages:
            sep = "&" if "?" in path else "?"
            got = client.get(path + sep + "program=" + self.alpha.replace(" ", "+"), follow_redirects=True)
            html_text = got.get_data(as_text=True)
            for kind, attr in (("GET", "href"), ("POST", "action")):
                for target in re.findall(rf'{attr}="(/[^"#]*)', html_text):
                    target = html.unescape(target).split("?")[0]
                    if target.startswith("/static") or target in ("/logout", "/theme", "/account/password") or (
                            target == "/" and user["is_associate"]):  # an associate's Home sends them to Today
                        continue
                    try:
                        endpoint, _ = adapter.match(target, method=kind)
                    except Exception:  # noqa: BLE001 (not a route of ours: a file, or a method the route lacks)
                        continue
                    found.append((path, kind, target, endpoint))
        return user, found

    def test_no_page_offers_what_the_role_is_refused(self):
        for who in ("plan_a", "tm_a", "sup_a", "assoc_a"):
            with self.subTest(who):
                user, found = self.offered(who)
                wrong = [(path, kind, target) for path, kind, target, endpoint in found
                         if TheRoleRules.expected_open(self, user, endpoint) is False]
                self.assertEqual(wrong, [])


# ---------------------------------------------------------------------------------------------- 2. the session
class TheSession(_Secured):
    def forged(self, payload, secret="not-the-sites-secret"):
        """A session cookie as the site would make it, signed with ``secret`` (a person without the site's own key)."""
        from flask.json.tag import TaggedJSONSerializer
        from itsdangerous import URLSafeTimedSerializer
        import hashlib
        return URLSafeTimedSerializer(secret, salt="cookie-session", serializer=TaggedJSONSerializer(),
                                      signer_kwargs={"key_derivation": "hmac", "digest_method": hashlib.sha1}).dumps(payload)

    def client_with(self, cookie):
        client = self.app.test_client()
        client.set_cookie("session", cookie)
        return client

    def test_a_cookie_signed_with_another_key_signs_nobody_in(self):
        admin = self.store.get_user_by_username("omar") if hasattr(self.store, "get_user_by_username") else {"id": self.admin_id}
        for payload in ({"user_id": admin["id"]}, {"user_id": admin["id"], "pw": "x" * 16, "csrf": "t"},
                        {"user_id": self.users["plan_a"]}):
            got = self.client_with(self.forged(payload)).get("/admin/users")
            self.assertEqual(got.status_code, 302, payload)
            self.assertIn("/login", got.headers["Location"])

    def test_a_cookie_edited_by_hand_is_refused(self):
        client = sign_in(self.app, "plan_a", PASSWORD)
        cookie = client.get_cookie("session").value
        body, rest = cookie.split(".", 1)
        raw = json.loads(zlib.decompress(base64.urlsafe_b64decode(body[1:] + "=" * (-len(body[1:]) % 4))) if body.startswith(".")
                         else base64.urlsafe_b64decode(body + "=" * (-len(body) % 4)))
        self.assertEqual(raw["user_id"], self.users["plan_a"])
        raw["user_id"] = self.admin_id  # what a person with a text editor would try
        edited = base64.urlsafe_b64encode(json.dumps(raw, separators=(",", ":")).encode()).decode().rstrip("=")
        got = self.client_with(f"{edited}.{rest}").get("/admin/users")
        self.assertEqual(got.status_code, 302)
        self.assertIn("/login", got.headers["Location"])
        # one character of the signature: from the middle, since the last character of a base64 string can carry bits
        # that are dropped when it is read (flipping it may leave the signature the same)
        at = cookie.rindex(".") + 1 + 5
        flipped = cookie[:at] + ("A" if cookie[at] != "A" else "B") + cookie[at + 1:]
        self.assertEqual(self.client_with(flipped).get("/admin/users").status_code, 302)

    def test_a_copy_of_a_cookie_stops_working_when_the_password_changes_the_account_is_switched_off_or_deleted(self):
        for how in ("password", "off", "deleted"):
            uid = self.store.add_user(f"cookie_{how}", f"Cookie {how}", PASSWORD, must_change=False)
            self.store.set_user_programs(uid, [self.alpha_id])
            client = sign_in(self.app, f"cookie_{how}", PASSWORD)
            copy = self.client_with(client.get_cookie("session").value)
            self.assertEqual(copy.get("/day").status_code in (200, 302, 403), True)
            self.assertNotIn("/login", copy.get("/day", query_string={"program": self.alpha}).headers.get("Location", ""))
            if how == "password":
                self.store.set_password(uid, "Another-pass-12345")
            elif how == "off":
                self.store.set_active(uid, False)
            else:
                self.store.delete_user(uid)
            got = copy.get("/day", query_string={"program": self.alpha})
            self.assertEqual(got.status_code, 302, how)
            self.assertIn("/login", got.headers["Location"], how)

    def test_sign_out_other_devices_ends_every_other_session_and_keeps_this_one(self):
        # Phase AU: a button on Settings; the password change did this already, this is for a lost phone
        here = sign_in(self.app, "plan_a", PASSWORD)
        there = sign_in(self.app, "plan_a", PASSWORD)
        self.assertEqual(there.get("/exports").status_code, 200)
        page = here.get("/account/password").get_data(as_text=True)
        self.assertIn("/account/sessions/end", page)
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)
        got = here.post("/account/sessions/end", data={"csrf_token": csrf})
        self.assertEqual(got.status_code, 302)
        self.assertEqual(here.get("/exports").status_code, 200)
        gone = there.get("/exports")
        self.assertEqual(gone.status_code, 302)
        self.assertIn("/login", gone.headers["Location"])
        self.assertEqual(sign_in(self.app, "plan_a", PASSWORD).get("/exports").status_code, 200)  # a new sign-in works

    def test_a_role_taken_away_takes_effect_on_the_next_request(self):
        uid = self.store.add_user("demoted", "Demoted", PASSWORD, must_change=False)
        self.store.update_user(uid, is_supervisor=1)
        self.store.set_user_programs(uid, [self.alpha_id])
        client = sign_in(self.app, "demoted", PASSWORD)
        self.assertEqual(client.get("/admin/users").status_code, 200)
        self.store.update_user(uid, is_supervisor=0)
        self.assertEqual(client.get("/admin/users").status_code, 403)

    def test_a_cookie_taken_before_signing_in_does_not_become_a_signed_in_one(self):
        client = self.app.test_client()
        client.get("/login")
        before = client.get_cookie("session")
        old = before.value if before else ""
        tok = re.search(r'name="csrf_token" value="([^"]+)"', client.get("/login").get_data(as_text=True)).group(1)
        client.post("/login", data={"username": "plan_a", "password": PASSWORD, "csrf_token": tok})
        self.assertEqual(self.client_with(old).get("/admin/users").status_code, 302) if old else None
        self.assertNotEqual(client.get_cookie("session").value, old)

    def test_the_cookie_is_http_only_same_site_and_secure_when_the_site_is_https(self):
        app, store, *_ = make_app(start_worker=False, HTTPS=True)
        store.add_user("hh", "HH", PASSWORD, must_change=False)
        client = app.test_client()
        tok = re.search(r'name="csrf_token" value="([^"]+)"', client.get("/login").get_data(as_text=True)).group(1)
        got = client.post("/login", data={"username": "hh", "password": PASSWORD, "csrf_token": tok})
        header = got.headers.get("Set-Cookie", "")
        for flag in ("HttpOnly", "SameSite=Lax", "Secure"):
            self.assertIn(flag, header)
        self.assertIn("Strict-Transport-Security", got.headers)

    def test_the_sites_own_secret_key_is_long_private_and_made_private_from_the_start(self):
        from webapp.app import _secret_key
        with tempfile.TemporaryDirectory() as folder:
            first = _secret_key(Path(folder))
            self.assertGreaterEqual(len(first), 48)
            self.assertEqual(stat.S_IMODE(os.stat(Path(folder) / "secret_key").st_mode), 0o600)
            self.assertEqual(_secret_key(Path(folder)), first)  # kept, not remade
        with tempfile.TemporaryDirectory() as other:
            self.assertNotEqual(_secret_key(Path(other)), first)


# ---------------------------------------------------------------------------------------------- 3. forms and addresses
class TheForms(_Secured):
    def post(self, who, url, **data):
        client = sign_in(self.app, who, PASSWORD)
        page = client.get("/", follow_redirects=True).get_data(as_text=True)
        tok = re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)
        return client.post(url, data={"csrf_token": tok, **data})

    def user(self, name):
        return self.store.get_user(self.users[name]) if name in self.users else next(
            u for u in self.store.list_users() if u["username"] == name)

    def test_a_supervisor_cannot_make_an_admin_or_give_every_program_by_editing_the_form(self):
        before = len(self.store.list_users())
        for fields in ({"role": "admin"}, {"role": "supervisor"}, {"is_admin": "on"}, {"all_programs": "1", "role": "planner"}):
            self.post("sup_a", "/admin/users", username="esc_new", password="Escalate-pass-1", display_name="E",
                      programs=[str(self.alpha_id)], **fields)
            made = next((u for u in self.store.list_users() if u["username"] == "esc_new"), None)
            if made is not None:  # only a plain planner of their own program may exist
                self.assertEqual((made["is_admin"], made["is_supervisor"], made["all_programs"]), (0, 0, 0), fields)
                self.assertEqual(self.store.user_program_ids(made["id"]), [self.alpha_id])
                self.store.delete_user(made["id"])
        self.assertEqual(len(self.store.list_users()), before)

    def test_nobody_raises_their_own_role_or_programs_through_the_people_routes(self):
        for who in ("plan_a", "tm_a", "sup_a"):
            me = self.user(who)
            for action in ("programs", "team", "reset", "enable", "disable", "delete"):
                self.post(who, f"/admin/users/{me['id']}/{action}", role="admin", all_programs="1",
                          programs=[str(self.alpha_id), str(self.bravo_id)], confirm="1", team_manager="")
            after = self.user(who)
            for field in ("is_admin", "is_supervisor", "is_team_manager", "all_programs", "is_associate", "active", "deleted"):
                self.assertEqual(after[field], me[field], (who, field))
            self.assertEqual(self.store.user_program_ids(me["id"]), [self.alpha_id])

    def test_a_supervisor_cannot_change_another_persons_role_or_give_them_a_program_they_do_not_hold(self):
        target = self.store.add_user("esc_target", "Esc Target", PASSWORD, must_change=False)
        self.store.set_user_programs(target, [self.alpha_id])
        before = self.store.get_user(target)
        self.post("sup_a", f"/admin/users/{target}/programs", role="admin", all_programs="1",
                  programs=[str(self.alpha_id), str(self.bravo_id)])
        after = self.store.get_user(target)
        self.assertEqual((after["is_admin"], after["is_supervisor"], after["all_programs"]),
                         (before["is_admin"], before["is_supervisor"], before["all_programs"]))
        self.assertEqual(self.store.user_program_ids(target), [self.alpha_id])

    def test_an_associate_asks_only_for_their_own_name_whatever_the_form_says(self):
        before = len(self.store.list_leave())
        self.post("assoc_a", "/me/leave", first="2026-11-01", last="2026-11-02", kind="Annual", note="x",
                  associate="Zorro Person One", program=self.bravo, user_id=str(self.users["plan_b"]))
        rows = self.store.list_leave()[before:]
        self.assertTrue(all(r["associate"] == "Alpha Person One" and r["program"] == self.alpha
                            and r["user_id"] == self.users["assoc_a"] for r in rows), rows)
        self.assertFalse([r for r in self.store.list_leave() if r["associate"] == "Zorro Person One" and r["note"] == "x"])

    def test_every_form_without_the_sites_token_is_refused_and_so_is_another_sessions_token(self):
        client = sign_in(self.app, "plan_a", PASSWORD)
        other = sign_in(self.app, "plan_b", PASSWORD)
        theirs = re.search(r'name="csrf_token" value="([^"]+)"', other.get("/").get_data(as_text=True)).group(1)
        for rule in self.app.url_map.iter_rules():
            if "POST" not in rule.methods or rule.endpoint == "static":
                continue
            args = {a: self.ids().get(a, "plan") for a in rule.arguments}
            with self.app.test_request_context():
                from flask import url_for
                url = url_for(rule.endpoint, **args)
            for data in ({}, {"csrf_token": ""}, {"csrf_token": theirs}):
                got = client.post(url, data={**fx.FORM, **data})
                self.assertEqual(got.status_code, 400, (rule.rule, data))

    def test_an_address_after_next_never_leaves_the_site(self):
        for target in ("//evil.example/x", "/\\evil.example", "https://evil.example", "/%09/evil.example", "javascript:alert(1)",
                       "///evil.example", "/\t/evil.example"):
            client = self.app.test_client()
            tok = re.search(r'name="csrf_token" value="([^"]+)"', client.get("/login").get_data(as_text=True)).group(1)
            got = client.post("/login", query_string={"next": target}, data={"username": "plan_a", "password": PASSWORD,
                                                                           "csrf_token": tok})
            where = got.headers.get("Location", "")
            self.assertTrue(where.startswith("/") and not where.startswith("//") and "\\" not in where
                            and not re.search(r"[\t\r\n]", where), (target, where))


# ---------------------------------------------------------------------------------------------- 4. the answers
class TheAnswers(_Secured):
    def test_every_answer_carries_the_security_headers(self):
        client = sign_in(self.app, "plan_a", PASSWORD)
        for url in ("/", "/day?program=" + self.alpha.replace(" ", "+"), "/requests", "/admin/users", "/nothing-here"):
            got = client.get(url)
            csp = got.headers["Content-Security-Policy"]
            for part in ("default-src 'self'", "frame-ancestors 'none'", "base-uri 'self'", "form-action 'self'",
                         "object-src 'none'"):
                self.assertIn(part, csp, url)
            self.assertNotIn("unsafe", csp)
            self.assertEqual(got.headers["X-Frame-Options"], "DENY")
            self.assertEqual(got.headers["X-Content-Type-Options"], "nosniff")
            self.assertEqual(got.headers["Referrer-Policy"], "same-origin")
            self.assertEqual(got.headers["Cache-Control"], "no-store", url)
            self.assertIn("camera=()", got.headers["Permissions-Policy"])
        self.assertEqual(self.app.test_client().get("/login").headers["Cache-Control"], "no-store")
        self.assertNotIn("no-store", client.get("/static/app.css").headers.get("Cache-Control", ""))

    def test_an_error_page_shows_no_trace_and_no_path(self):
        client = sign_in(self.app, "plan_a", PASSWORD)
        for url in ("/runs/zzzz/week", "/schedules/999999/download", "/programs/Nothing"):
            text = client.get(url).get_data(as_text=True)
            self.assertNotRegex(text, r"Traceback|File \"|/home/|/tmp/|sqlite|\.py\"")

    def test_what_people_type_is_escaped_wherever_it_is_shown(self):
        from webapp.contacts import ContactBook
        payloads = ('"><script>alert(11)</script>', "<img src=x onerror=alert(12)>", "'+alert(13)+'")
        uid = self.store.add_user("xss_user", payloads[0], PASSWORD, must_change=False)
        self.store.set_user_programs(uid, [self.alpha_id])
        ContactBook(self.store).add(self.alpha_id, "Dept " + payloads[1], "Contact " + payloads[0], self.admin_id)
        self.store.add_leave(program=self.alpha, associate="Alpha Person One", user_id=uid, kind="Annual",
                             first_day="2026-10-20", last_day="2026-10-21", note=payloads[1])
        self.store.add_change_request(kind="coaching", program=self.alpha, shift_date="2026-10-14",
                                      names=json.dumps(["Alpha Person One"]), requester_id=uid,
                                      requester_name=payloads[0], why=payloads[1],
                                      detail=json.dumps({"start": 600, "end": 630, "minutes": 30, "level": "high"}))
        raw = re.compile(r"<script>alert\(1[123]\)|<img src=x onerror=")
        admin = sign_in(self.app, "omar", "Owner-pass-123")
        values = {"run_id": self.runs[self.alpha], "schedule_id": self.sched[self.alpha], "name": self.alpha,
                  "user_id": uid, "request_id": 1, "swap_id": 1}
        for rule in self.app.url_map.iter_rules():
            if "GET" not in rule.methods or rule.endpoint == "static" or set(rule.arguments) - set(values):
                continue
            with self.app.test_request_context():
                from flask import url_for
                url = url_for(rule.endpoint, **{a: values[a] for a in rule.arguments}, program=self.alpha, q="x")
            got = admin.get(url)
            if got.headers.get("Content-Type", "").startswith(("text/html", "application/json")):
                self.assertNotRegex(got.get_data(as_text=True), raw, url)

    def test_no_template_turns_the_automatic_escaping_off(self):
        root = Path(__file__).resolve().parents[1]
        for path in list((root / "templates").glob("*.html")):
            text = path.read_text(encoding="utf-8")
            self.assertNotIn("|safe", text, path.name)
            self.assertNotIn("autoescape false", text, path.name)
        for path in root.glob("*.py"):  # the only Markup is the charts', whose every typed value goes through escape()
            if path.name != "charts.py":
                self.assertNotIn("Markup(", path.read_text(encoding="utf-8"), path.name)

    def test_the_sql_never_takes_a_column_or_table_name_from_a_request(self):
        root = Path(__file__).resolve().parents[1]
        text = (root / "app.py").read_text(encoding="utf-8")
        self.assertNotRegex(text, r"\*\*request\.(form|values|json)")
        self.assertNotRegex(text, r"\*\*dict\(request")


# ---------------------------------------------------------------------------------------------- exports by program
class TheExportsStayInTheirProgram(_Secured):
    def cells(self, data):
        wb = load_workbook(io.BytesIO(data))
        return " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c)

    def test_a_download_of_every_kind_holds_nothing_of_another_program(self):
        query = {"from": "2026-10-01", "to": "2026-10-31", "format": "xlsx", "kind": list(KINDS)}
        admin = sign_in(self.app, "omar", "Owner-pass-123")
        everything = self.cells(admin.get("/exports/download", query_string=query).data)
        self.assertIn("Zorro Person One", everything)  # the data is there for someone who has it
        self.assertIn("Alpha Person One", everything)
        for who in ("plan_a", "sup_a", "tm_a"):
            got = sign_in(self.app, who, PASSWORD).get("/exports/download", query_string=query)
            self.assertEqual(got.status_code, 200, who)
            text = self.cells(got.data)
            self.assertIn("Alpha Person One", text, who)
            self.assertNotRegex(text, r"(?i)zorro|bravo", who)

    def test_sign_in_history_is_for_people_who_see_everything_only(self):
        self.store.add_event(kind="sign_in_failed", subject="plan_b", detail="from 198.51.100.9: bad",
                             user_id=self.users["plan_b"])
        query = {"from": "2026-10-01", "to": "2026-10-31", "format": "xlsx", "kind": list(KINDS)}
        for who in ("plan_a", "sup_a"):
            got = sign_in(self.app, who, PASSWORD).get("/exports/download", query_string=query)
            self.assertNotIn("198.51.100.9", self.cells(got.data), who)
        admin = sign_in(self.app, "omar", "Owner-pass-123")
        self.assertIn("198.51.100.9", self.cells(admin.get("/exports/download", query_string=query).data))

    def test_asking_for_another_programs_download_by_any_route_is_refused(self):
        client = sign_in(self.app, "plan_a", PASSWORD)
        base = {"from": "2026-10-01", "to": "2026-10-31", "format": "xlsx", "kind": ["attendance", "record"]}
        for extra in ({"program": self.bravo}, {"pick": f"key:{self.bravo}"}, {"pick": f"group:{self.bravo_id}"},
                      ):
            self.assertEqual(client.get("/exports/download", query_string={**base, **extra}).status_code, 403, extra)

    def test_filtering_by_another_programs_person_shows_nothing_of_theirs(self):
        client = sign_in(self.app, "plan_a", PASSWORD)
        got = client.get("/exports/download", query_string={"from": "2026-10-01", "to": "2026-10-31", "format": "xlsx",
                                                           "kind": list(KINDS), "user": self.users["plan_b"]})
        self.assertEqual(got.status_code, 403)


if __name__ == "__main__":
    unittest.main()
