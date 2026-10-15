# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AH (owner, 2026-10-15; sample ah01 and ah02 approved): the Weeks page's Hours needed table. One row per
interval, a column per day, the hours needed (to 100% or, the RTA's choice, to the target with 10 minutes to spare),
coloured by the week's target; a filter for overtime or days off to cancel; as recorded on the RTA or as planned;
what to do per day with names; the same table in Excel. The Schedules page shows the RTA's records too. Made-up
"Associate NN" data."""
import io
import unittest
from datetime import date

from openpyxl import load_workbook

from webapp import needed
from webapp.day import STEP
from webapp.tests.test_notify_daybook import WED, _Tagged
from webapp.tests.test_runs import sign_in

SUN = date(2026, 10, 11)


def hour(t, ticks, required=5.0):
    return {"t": t, "required": required, "slots": list(ticks), "now": round(sum(ticks) / len(ticks), 1),
            "plan": required, "pm": (sum(ticks) / len(ticks) - required) if required else None}


def page(on, cells, target=0.9):
    return {"date": on, "view": {"interval": 60, "cells": cells}, "target": {"target": target}}


def day(on, cells, plan=None):
    return {"date": on, "page": page(on, cells), "plan": plan, "why": ""}


FULL = [5] * 12           # 300 person-minutes: 100%
FOUR = [4] * 12           # 240: 80%, below 90%
FIVE_OVER = [5] * 7 + [4] * 5  # 275: 92%, 5 minutes over the target


def week(cells_by_day, plans=None):
    plans = plans or {}
    return [day(date(2026, 10, 11 + i), cells_by_day.get(i, [hour(600, FULL), hour(660, FULL), hour(720, FULL)]),
                plans.get(i)) for i in range(7)]


class TheTable(unittest.TestCase):
    def test_hours_to_100_and_to_the_target_round_up_to_the_quarter(self):
        days = week({4: [hour(600, FOUR), hour(660, FIVE_OVER), hour(720, FULL)]})
        full = needed.week_table(days, "full", "all")
        cells = {r["time"]: r["cells"][4] for r in full["rows"] if "cells" in r}
        self.assertEqual(cells["10:00"]["hours"], 1.0)   # 60 person-minutes short of 100%
        self.assertEqual(cells["11:00"]["hours"], 0.5)   # 25 minutes, up to the quarter hour
        self.assertEqual((cells["10:00"]["cls"], cells["11:00"]["cls"]), ("bad", "warn"))
        self.assertTrue(cells["11:00"]["tight"])
        target = needed.week_table(days, "target", "all")
        cells = {r["time"]: r["cells"][4] for r in target["rows"] if "cells" in r}
        self.assertEqual(cells["10:00"]["hours"], 0.75)  # 270 + 10 = 280 against 240: 40 minutes
        self.assertEqual(cells["11:00"]["hours"], 0.25)  # 5 minutes more for 10 over the target
        self.assertEqual(target["totals"][4]["hours"], 1.0)
        self.assertEqual((full["totals"][4]["bad"], full["totals"][4]["warn"], full["totals"][4]["tight"]), (1, 1, 1))

    def test_rows_needing_nothing_fold(self):
        days = week({2: [hour(600, FULL), hour(660, FOUR), hour(720, FULL)]})
        rows = needed.week_table(days, "full", "all")["rows"]
        self.assertEqual(rows[0], {"fold": "00:00 to 11:00"})
        self.assertEqual(rows[1]["time"], "11:00")
        self.assertEqual(rows[2], {"fold": "12:00 to 24:00"})

    def test_fix_tags_and_filters(self):
        plan = {"proposals": [{"text": "Cancel 1 day off, 10:00 to 19:00", "start": 600, "end": 1140, "people": 1,
                               "intervals": 6, "covers": "10:00 to 12:00",
                               "candidates": [{"name": "Associate 07", "bends": [], "note": ""}]}],
                "overtime": [{"text": "11:00 to 12:00: 1 person", "start": 660, "end": 720, "people": 1,
                              "offers": [{"name": "Associate 03", "text": "stay 10:00 to 12:00"}]}]}
        days = week({5: [hour(600, FOUR), hour(660, [3] * 12), hour(720, FULL)]}, {5: plan})
        table = needed.week_table(days, "full", "all")
        fix = {r["time"]: r["cells"][5]["fix"] for r in table["rows"] if "cells" in r}
        self.assertEqual(fix, {"10:00": "Day off", "11:00": "Day off + OT"})
        self.assertEqual(table["counts"], {"all": 2, "ot": 1, "dayoff": 2})
        only_ot = needed.week_table(days, "full", "ot")
        shown = [r["time"] for r in only_ot["rows"] if "cells" in r]
        self.assertEqual(shown, ["11:00"])
        card = table["cards"][0]
        self.assertEqual(card["date"], date(2026, 10, 16))
        self.assertIn("cancel 1 day off", card["verdict"].lower())
        self.assertEqual(card["proposals"][0]["candidates"][0]["name"], "Associate 07")
        self.assertEqual(card["overtime"][0]["names"], ["Associate 03 (stay 10:00 to 12:00)"])

    def test_a_day_without_a_schedule_says_why(self):
        days = week({})
        days[3] = {"date": date(2026, 10, 14), "page": None, "plan": None, "why": "The input workbook is missing."}
        table = needed.week_table(days, "full", "all")
        self.assertEqual(table["days"][3]["why"], "The input workbook is missing.")
        self.assertTrue(all(r.get("cells", [None] * 7)[3] is None for r in table["rows"]))

    def test_bad_choices_fall_back(self):
        self.assertEqual(needed.pick("nonsense", needed.AIMS), "full")
        self.assertEqual(needed.pick("target", needed.AIMS), "target")
        self.assertEqual(needed.pick(None, needed.BASES), "recorded")
        self.assertEqual(needed.pick("<b>", needed.SHOWS), "all")


class TheWeekPage(_Tagged):
    URL = "/week?program={key}&week=2026-10-11"

    def html(self, **args):
        extra = "".join(f"&{k}={v}" for k, v in args.items())
        return self.client.get(self.URL.format(key=self.key) + extra).get_data(as_text=True)

    def test_the_panel_shows_the_week(self):
        html = self.html()
        self.assertIn('id="needed"', html)
        self.assertIn("Hours needed", html)
        for label in ("Sun 11", "Mon 12", "Sat 17"):
            self.assertIn(label, html)
        self.assertIn("As recorded on the RTA", html)
        self.assertIn("Download this table", html)
        self.assertNotIn('style="', html)

    def test_recorded_counts_a_sick_mark_and_planned_does_not(self):
        page = self.days.page(self.key, WED)
        name = next(l["name"] for l in page["view"]["lanes"]
                    if any(s["offset"] == 0 and s["status"] == "Present" for s in l["segments"]))
        self.days.set_status(self.key, WED, name, "Sick", self.omar)
        try:
            recorded = self.days.page(self.key, WED)
            planned = self.days.page(self.key, WED, recorded=False)
            self.assertLess(sum(sum(c["slots"]) for c in recorded["view"]["cells"]),
                            sum(sum(c["slots"]) for c in planned["view"]["cells"]))
            self.assertEqual(planned["view"]["tiles"]["absent"], 0)
            self.assertIn("As planned", self.html(basis="planned"))
        finally:
            self.days.undo_last(self.key, WED, self.omar)

    def test_the_target_aim_needs_no_more_than_100(self):
        days = [self.days.page(self.key, date(2026, 10, 11 + i)) for i in range(7)]
        rows = [{"date": p["date"], "page": p, "plan": None, "why": ""} for p in days]
        full, target = needed.week_table(rows, "full", "all"), needed.week_table(rows, "target", "all")
        for a, b in zip(full["totals"], target["totals"]):
            self.assertLessEqual(b["hours"], a["hours"])

    def test_download_is_the_same_table(self):
        got = self.client.get(f"/week/hours-needed.xlsx?program={self.key}&week=2026-10-11&aim=full")
        self.assertEqual(got.status_code, 200)
        wb = load_workbook(io.BytesIO(got.data))
        self.assertEqual(wb.sheetnames, ["Hours needed", "What to do"])
        ws = wb["Hours needed"]
        self.assertEqual(ws.cell(2, 1).value, "Interval")
        self.assertEqual(ws.cell(2, 2).value, "Sun 11 Oct")
        fills = {ws.cell(r, c).fill.fgColor.rgb[-6:] for r in range(3, ws.max_row + 1) for c in range(2, 9)
                 if ws.cell(r, c).fill.fill_type == "solid"}
        self.assertTrue(fills <= {"D5F2EE", "FBE9C9", "F8D3D0"}, fills)

    def test_download_checks_access(self):
        self.store.add_user("noone", "No One", "Owner-pass-123", must_change=False)
        other = sign_in(self.app, "noone", "Owner-pass-123")
        got = other.get(f"/week/hours-needed.xlsx?program={self.key}&week=2026-10-11")
        self.assertEqual(got.status_code, 403)

    def test_schedules_page_shows_the_rtas_records(self):
        run = next(r for r in self.store.list_runs() if r.get("program") == self.key)
        html = self.client.get(f"/runs/{run['id']}/schedules").get_data(as_text=True)
        self.assertIn("As recorded on the RTA", html)


if __name__ == "__main__":
    unittest.main()
