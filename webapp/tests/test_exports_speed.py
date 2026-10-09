# © 2026 Omar Mokhtar. All rights reserved.
"""Phase W task 1: exports that finish (owner, 2026-10-09: "Exporting as well lags alot and sometimes it just keeps
loading with no actual export").

Measured: a month's export read the input workbook again for every LOB-day (the cache never kept anything: a loop in
read_inputs reused the name of its cache key), and worked hours and the daily summary each worked out the same days.
Each workbook is now read once, each day once per export, an export that runs past its time stops and says what to
untick, and the download answers a script in JSON so the page can say "Preparing", "Ready" or what went wrong."""
import json
import shutil
import tempfile
import time
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from webapp import day as day_module
from webapp.attendance import DayBook
from webapp.exports import KINDS, ExportTooSlow, build
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.tests.test_schedules import AFTER, INPUT, REPO

WED, THU = date(2026, 10, 14), date(2026, 10, 15)


def with_target(dst: Path, value) -> Path:
    """A copy of the test input whose Instructions "Target" reads ``value`` (None: the row is emptied)."""
    shutil.copyfile(INPUT, dst)
    wb = load_workbook(dst)
    ws = next(wb[n] for n in wb.sheetnames if n.lower() == "instructions")
    for row in ws.iter_rows():
        for i, cell in enumerate(row[:-1]):
            if isinstance(cell.value, str) and cell.value.strip().lower() == "target":
                row[i + 1].value = value
    wb.save(dst)
    return dst


class TheWorkbookRead(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir, True)

    def test_a_workbook_is_read_once(self):
        path = self.dir / "input.xlsx"
        shutil.copyfile(INPUT, path)
        with mock.patch.object(day_module, "load_workbook", wraps=day_module.load_workbook) as loads:
            first = day_module.read_inputs(path)
            second = day_module.read_inputs(path)
        self.assertEqual(loads.call_count, 1)
        self.assertIs(first, second)

    def test_the_workbook_target_is_read(self):
        self.assertEqual(day_module.read_inputs(with_target(self.dir / "a.xlsx", 1))["target"], 1.0)
        self.assertEqual(day_module.read_inputs(with_target(self.dir / "b.xlsx", 0.85))["target"], 0.85)
        self.assertEqual(day_module.read_inputs(with_target(self.dir / "c.xlsx", 95))["target"], 0.95)
        self.assertEqual(day_module.read_inputs(with_target(self.dir / "d.xlsx", None))["target"], 0.9)


class TheExportWork(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program="AE/AR B2B",
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Late", self.sara, end="13:00")

    def test_worked_and_summary_share_each_day(self):
        asked = []
        real = DayBook.page

        def counting(days, program, on, measure="interval", extra=None):
            asked.append((program, on, measure))
            return real(days, program, on, measure, extra)

        with mock.patch.object(DayBook, "page", counting):
            build(self.store, self.days, WED, THU, kinds=["worked", "summary"], program="AE/AR B2B")
        # two days asked by both items, plus the summary's look at the day after the last (Friday): each once
        self.assertEqual(sorted(set(asked)), sorted(asked))
        self.assertEqual(len(asked), 3)

    def test_an_export_past_its_deadline_stops_and_says_what_to_untick(self):
        with self.assertRaises(ExportTooSlow) as said:
            build(self.store, self.days, WED, THU, kinds=list(KINDS), program="AE/AR B2B",
                  deadline=time.monotonic() - 1)
        self.assertIn(KINDS["worked"], str(said.exception))
        self.assertIn(KINDS["summary"], str(said.exception))
        self.assertIsInstance(said.exception, ValueError)  # said on the page like any other refusal


class TheDownloadForAScript(unittest.TestCase):
    def test_the_download_answers_a_script_in_json(self):
        from webapp.tests.test_runs import make_app, sign_in
        app, store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        client = sign_in(app, "omar", "Owner-pass-123")
        got = client.get("/exports/download?from=2026-10-14&to=2026-10-01&kind=attendance&format=xlsx",
                         headers={"X-Requested-With": "fetch"})
        self.assertEqual(got.status_code, 400)
        self.assertEqual(got.mimetype, "application/json")
        self.assertEqual(json.loads(got.get_data(as_text=True)), {"error": "The period ends before it starts."})
        page = client.get("/exports/download?from=2026-10-14&to=2026-10-01&kind=attendance&format=xlsx")
        self.assertEqual((page.status_code, page.mimetype), (400, "text/html"))  # without a script: the page says it


if __name__ == "__main__":
    unittest.main()
