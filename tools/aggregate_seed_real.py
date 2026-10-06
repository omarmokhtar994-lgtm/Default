#!/usr/bin/env python3
"""Phase D3 stage 2: a seed week for a real program workbook, from the workbook alone.

Stage 1 (tools/aggregate_seed.py) covers single-language workbooks without
leave or fixed requests. Real programs have leave, hard OFF, fixed shifts,
language minimums, carry-in and per-associate rules, so this model keeps one
decision per associate, day and shift (no tour enumeration) and applies the
engine's own Stage-1 hard rules to it:

  * day status: exactly one of a shift, OFF or pinned leave;
  * leave (approved / fixed), hard OFF (preference or fixed), fixed shifts;
  * the weekly OFF rule (add_weekly_off_rule semantics: exactly 2 OFF where
    room allows, a cyclic adjacent pair unless Separate OFF Days);
  * rest gap between consecutive days, Saturday -> Sunday included, and
    after last Saturday's carry-in;
  * at most the workbook's number of different shifts;
  * nesting groups share OFF/leave/shift when fixed requests are on;
  * no new staffing in blank intervals when the workbook forbids it.

Breaks are counted per language class (associates with the same eligibility
for every language rule with a minimum), using the engine's legal break
patterns and its concurrent-break cap. Coverage is the engine metric
(shrinkage, target, carry-in). Language minimums are counted on on-floor
eligible heads plus eligible carry-in. Objective: shortfall first (language
shortfall weighted 100x), then intervals at target hinted with phase 1.

The week is written into the Schedule sheet, which the engine reads as a seed
(a hint, never a constraint). Changes no engine file. 11H/3OFF is refused.

    python3 tools/aggregate_seed_real.py --engine engine/_tools/l632_universal_scheduler.py \
        --out-dir DIR WORKBOOK.xlsx [...]
"""
from __future__ import annotations

import argparse
import copy
import json
import math
import shutil
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_synthetic_suite as S  # noqa: E402
from aggregate_seed import add_break_cap  # noqa: E402

SCALE = 10_000
DAY_COLS = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]


def candidates(E, p, max_shifts: int, guide_seconds: float) -> Tuple[List[int], Dict[str, Any]]:
    q = copy.copy(p)
    q.aggregate_guide_mixed_durations = True
    g = E.aggregate_pattern_mix_guidance(q, time_limit_sec=guide_seconds)
    totals: Dict[int, int] = {}
    for key, v in (g.get("shift_count_targets") or {}).items():
        s = int(key.split(":")[1])
        totals[s] = totals.get(s, 0) + int(v)
    chosen = sorted(totals, key=lambda s: (-totals[s], s))[:max_shifts]
    by_label = {E.norm(sh.label): sh.index for sh in p.shifts}
    if p.fixed_enabled:
        for a in p.associates:
            for v in a.fixed_schedule:
                if E.preference_kind(v) == "shift" and E.norm(v) in by_label:
                    chosen.append(by_label[E.norm(v)])
    if not chosen:  # no guide counts: every shift (real libraries hold 1-24)
        chosen = [sh.index for sh in p.shifts]
    return sorted(set(chosen)), {"guide_status": g.get("status"), "guide_shift_totals": totals}


def build(E, wb: Path, out_dir: Path, max_shifts: int, time_limit: float, workers: int,
          all_shifts: bool = False, fix_week: Dict[str, List[str]] = None) -> Dict[str, Any]:
    """Per-associate variant. Diagnostics: all_shifts uses the whole shift library as
    candidates; fix_week (name -> 7 day cells) pins x/off to a given week so the model
    scores that week under its own metric (e.g. the engine's result)."""
    from ortools.sat.python import cp_model
    t0 = time.time()
    p = E.parse_input(wb)
    if p.use_11h_3off:
        return {"workbook": wb.name, "status": "REFUSED", "reason": "11H/3OFF out of scope"}
    cand, ginfo = candidates(E, p, max_shifts, min(45.0, time_limit / 4))
    if all_shifts:
        cand = [sh.index for sh in p.shifts]
    shifts = [p.shifts[i] for i in cand]
    A, D, J = len(p.associates), 7, len(shifts)
    qpi = p.qslots_per_interval
    m = cp_model.CpModel()
    x = {(a, d, j): m.NewBoolVar(f"x{a}_{d}_{j}") for a in range(A) for d in range(D) for j in range(J)}
    off = {(a, d): m.NewBoolVar(f"o{a}_{d}") for a in range(A) for d in range(D)}
    blank_hard = getattr(p, "blank_requirement_mode", "allow") != "allow"
    nest: Dict[str, List[int]] = {}
    for a, assoc in enumerate(p.associates):
        leave_days = 0
        for d in range(D):
            pref = assoc.preferences[d] if d < len(assoc.preferences) else ""
            fixed = assoc.fixed_schedule[d] if d < len(assoc.fixed_schedule) else ""
            kind, fkind = E.preference_kind(pref), E.preference_kind(fixed)
            is_leave = (p.fixed_enabled and fkind == "leave") or (p.leave_enabled and kind == "leave")
            if is_leave:
                leave_days += 1
                m.Add(off[a, d] == 0)
                for j in range(J):
                    m.Add(x[a, d, j] == 0)
                continue
            m.Add(sum(x[a, d, j] for j in range(J)) + off[a, d] == 1)
            if (p.fixed_enabled and fkind == "off") or (p.hard_off and kind == "off"):
                m.Add(off[a, d] == 1)
            if p.fixed_enabled and fkind == "shift":
                js = [j for j, sh in enumerate(shifts) if E.norm(sh.label) == E.norm(fixed)]
                if js:
                    m.Add(x[a, d, js[0]] == 1)
            for j, sh in enumerate(shifts):
                if d == 0 and not E.previous_saturday_compatible(assoc.previous_saturday, sh, p.rest_gap_hours):
                    m.Add(x[a, d, j] == 0)
                if blank_hard:
                    base = d * 96 + sh.start_min // 15
                    for o in range(sh.duration_q):
                        t = base + o
                        if t < 7 * 96 and not p.active[t // 96][(t % 96) // qpi]:
                            m.Add(x[a, d, j] == 0)
                            break
        room = 7 - leave_days
        offs = [off[a, d] for d in range(D)]
        if p.strict_off:
            m.Add(sum(offs) == (2 if room >= 3 else max(0, room)))
            if not p.separate_off_days and room >= 2:
                pairs = []
                for d in range(D):
                    pr = m.NewBoolVar(f"pr{a}_{d}")
                    m.Add(pr <= offs[d]); m.Add(pr <= offs[(d + 1) % 7]); m.Add(pr >= offs[d] + offs[(d + 1) % 7] - 1)
                    pairs.append(pr)
                m.Add(sum(pairs) >= 1)
        else:
            m.Add(sum(offs) >= min(2, max(0, room)))
        ys = []
        for j in range(J):
            y = m.NewBoolVar(f"y{a}_{j}")
            for d in range(D):
                m.Add(y >= x[a, d, j])
            ys.append(y)
        m.Add(sum(ys) <= int(p.max_different_shifts))
        for d in range(D):
            for j1, s1 in enumerate(shifts):
                for j2, s2 in enumerate(shifts):
                    if not E.rest_compatible(s1, s2, p.rest_gap_hours):
                        m.Add(x[a, d, j1] + x[a, (d + 1) % 7, j2] <= 1)
        if p.fixed_enabled and assoc.nesting_group:
            nest.setdefault(E.norm(assoc.nesting_group), []).append(a)
    pinned = 0
    for a, assoc in enumerate(p.associates):
        cells = (fix_week or {}).get(E.norm(assoc.name))
        if not cells:
            continue
        for d in range(D):
            v = str(cells[d] or "").strip()
            js = [j for j, sh in enumerate(shifts) if E.norm(sh.label) == E.norm(v)]
            if js:
                m.Add(x[a, d, js[0]] == 1); pinned += 1
            elif E.preference_kind(v) == "off":
                m.Add(off[a, d] == 1); pinned += 1
            elif E.preference_kind(v) == "shift":
                raise ValueError(f"fix_week: {assoc.name} day {d} shift {v!r} is not a candidate")
    for members in nest.values():
        for b in members[1:]:
            for d in range(D):
                m.Add(off[b, d] == off[members[0], d])
                for j in range(J):
                    m.Add(x[b, d, j] == x[members[0], d, j])
    # Language classes: eligibility for every rule with a minimum.
    min_rules = [r for r in p.language_rules if int(r.minimum or 0) > 0]
    cls_of = [tuple(bool(E.language_eligible(r, assoc)) for r in min_rules) for assoc in p.associates]
    classes = sorted(set(cls_of))
    patterns = E.generate_break_patterns(p)
    onfloor = {c: [[] for _ in range(7 * 96)] for c in classes}
    onbreak_total = [[] for _ in range(7 * 96)]
    zlist = []
    for c in classes:
        members = [a for a in range(A) if cls_of[a] == c]
        for d in range(D):
            for j, sh in enumerate(shifts):
                cnt = sum(x[a, d, j] for a in members)
                legal = [pt for pt in patterns if pt.duration_q == sh.duration_q] or [None]
                zs = []
                for pt in legal:
                    z = m.NewIntVar(0, len(members), f"z{classes.index(c)}_{d}_{j}_{pt.index if pt else 'nb'}")
                    zs.append(z); zlist.append(z)
                    broken = set(pt.broken_offsets) if pt else set()
                    base = d * 96 + sh.start_min // 15
                    for o in range(sh.duration_q):
                        t = base + o
                        if t >= 7 * 96:
                            break
                        if o in broken:
                            onbreak_total[t].append(z)
                        else:
                            onfloor[c][t].append(z)
                m.Add(sum(zs) == cnt)
    prior = [len(E.prior_covering_associates(p, q)) for q in range(7 * 96)]
    ratio, absolute = float(p.break_max_concurrent_ratio), int(p.break_max_concurrent_absolute)
    for t in range(7 * 96):
        if onbreak_total[t]:
            b = sum(onbreak_total[t])
            f = sum(sum(onfloor[c][t]) for c in classes)
            add_break_cap(m, b, b + f + prior[t], ratio, absolute)
    rows = []
    for d in range(D):
        for i in range(p.intervals_per_day):
            if not p.active[d][i]:
                continue
            need = p.target_ratio * float(p.requirements[d][i] or 0.0) * qpi
            eff = 1.0 - float(p.shrinkage[d][i])
            if need <= 1e-9 or eff <= 0:
                continue
            coef = int(round(eff * SCALE))
            qs = [d * 96 + i * qpi + k for k in range(qpi)]
            lhs = coef * sum(sum(onfloor[c][q]) for c in classes for q in qs) + coef * sum(prior[q] for q in qs)
            rows.append((lhs, int(math.ceil(need * SCALE - 1e-6))))
    lang_short = []
    for t in range(7 * 96):
        d, minute = t // 96, (t % 96) * 15
        if not p.active[d][(t % 96) // qpi]:
            continue
        for ri, r in enumerate(min_rules):
            if r not in E.language_rules_at(p, minute, day=d):
                continue
            elig = sum(sum(onfloor[c][t]) for c in classes if c[ri])
            have = len(E.prior_covering_associates(p, t, r))
            ls = m.NewIntVar(0, int(r.minimum), f"ls{t}_{ri}")
            m.Add(ls >= int(r.minimum) - have - elig)
            lang_short.append(ls)
    short = []
    for k, (lhs, rhs) in enumerate(rows):
        s_ = m.NewIntVar(0, rhs, f"s{k}")
        m.Add(s_ >= rhs - lhs)
        short.append(s_)
    lang_w = SCALE * 100
    m.Minimize(lang_w * sum(lang_short) + sum(short))
    sv = cp_model.CpSolver()
    sv.parameters.max_time_in_seconds = time_limit * 0.5
    sv.parameters.num_workers = workers
    sv.parameters.random_seed = 1
    st1 = sv.Solve(m)
    info = {"workbook": wb.name, "associates": A, "pinned_cells": pinned, "candidate_shifts": [sh.label for sh in shifts],
            "language_classes": len(classes), "phase1": {"status": sv.StatusName(st1), "seconds": round(sv.WallTime(), 1)}, **ginfo}
    if st1 not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        info["status"] = "NO_SEED"
        return info
    p1_zero = [sv.Value(s_) == 0 for s_ in short]
    info["phase1"].update(shortfall_zero_intervals=sum(p1_zero), language_shortfall=sum(sv.Value(v) for v in lang_short))
    p1_x = {k: sv.Value(v) for k, v in x.items()}
    for v in list(x.values()) + list(off.values()) + zlist:
        m.AddHint(v, sv.Value(v))
    for v in lang_short:
        m.Add(v == 0) if info["phase1"]["language_shortfall"] == 0 else None
    m.ClearObjective()
    hits = []
    for k, (lhs, rhs) in enumerate(rows):
        h = m.NewBoolVar(f"h{k}")
        m.Add(lhs >= rhs).OnlyEnforceIf(h)
        m.AddHint(h, 1 if p1_zero[k] else 0)
        hits.append(h)
    m.Maximize(sum(hits) - 1000 * sum(lang_short))
    sv2 = cp_model.CpSolver()
    sv2.parameters.max_time_in_seconds = time_limit * 0.5
    sv2.parameters.num_workers = workers
    sv2.parameters.random_seed = 1
    st2 = sv2.Solve(m)
    p2_ok = st2 in (cp_model.OPTIMAL, cp_model.FEASIBLE)
    p2_hits = sum(sv2.Value(h) for h in hits) if p2_ok else -1
    use1 = info["phase1"]["shortfall_zero_intervals"] >= p2_hits
    chosen = p1_x if use1 else {k: sv2.Value(v) for k, v in x.items()}
    info.update(status="SEED", seed_from_phase=1 if use1 else 2,
                seed_model_hits=info["phase1"]["shortfall_zero_intervals"] if use1 else p2_hits,
                phase2={"status": sv2.StatusName(st2), "hits": p2_hits}, active_intervals=sum(
                    1 for d in range(D) for i in range(p.intervals_per_day) if p.active[d][i]))
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (wb.stem + "_AGGSEED.xlsx")
    shutil.copy(wb, out)
    import openpyxl
    book = openpyxl.load_workbook(out)
    ws = book["Schedule"]
    hdr = next(r for r in range(1, 6) if any(str(ws.cell(r, c).value or "").strip() == "Sun" for c in range(1, 30)))
    col = {str(ws.cell(hdr, c).value or "").strip(): c for c in range(1, 40)}
    name_col = col.get("SF Name") or col.get("Name")
    row_of = {E.norm(str(ws.cell(r, name_col).value)): r for r in range(hdr + 1, ws.max_row + 1) if ws.cell(r, name_col).value}
    for a, assoc in enumerate(p.associates):
        r = row_of.get(E.norm(assoc.name))
        if r is None:
            continue
        for d, dn in enumerate(DAY_COLS):
            if any(chosen[a, d, j] for j in range(J)):
                ws.cell(r, col[dn]).value = shifts[next(j for j in range(J) if chosen[a, d, j])].label
            else:
                pref = assoc.preferences[d] if d < len(assoc.preferences) else ""
                fixed = assoc.fixed_schedule[d] if d < len(assoc.fixed_schedule) else ""
                leave = (p.fixed_enabled and E.preference_kind(fixed) == "leave") or (p.leave_enabled and E.preference_kind(pref) == "leave")
                ws.cell(r, col[dn]).value = "Leave" if leave else "OFF"
    book.save(out)
    seed = E.load_seed_skeleton(E.parse_input(out), out, "input_schedule_seed")
    info.update(seed_workbook=out.name, seed_valid_cells=seed.diagnostics.get("valid_seed_cells") if seed else 0,
                seed_invalid_cells=seed.diagnostics.get("invalid_seed_cells") if seed else None,
                total_seconds=round(time.time() - t0, 1))
    return info


def group_tours(E, p, assoc, shifts, var_cap: int):
    """Legal weekly tours for one associate's constraints, over the candidate shifts.

    Each tour is a 7-tuple of a shift index, "O" (OFF) or "L" (leave), under
    the same rules as the per-associate model above. Shift variety is capped at
    min(the workbook's limit, var_cap) to keep the enumeration small; a seed
    that is stricter than the rule is still legal.
    """
    import itertools
    status = []
    leave_days = 0
    for d in range(7):
        pref = assoc.preferences[d] if d < len(assoc.preferences) else ""
        fixed = assoc.fixed_schedule[d] if d < len(assoc.fixed_schedule) else ""
        kind, fkind = E.preference_kind(pref), E.preference_kind(fixed)
        if (p.fixed_enabled and fkind == "leave") or (p.leave_enabled and kind == "leave"):
            status.append("L"); leave_days += 1
        elif (p.fixed_enabled and fkind == "off") or (p.hard_off and kind == "off"):
            status.append("O")
        elif p.fixed_enabled and fkind == "shift":
            js = [j for j, sh in enumerate(shifts) if E.norm(sh.label) == E.norm(fixed)]
            status.append(("F", js[0]) if js else "?")
        else:
            status.append(None)
    if "?" in status:
        return []
    room = 7 - leave_days
    need_off = (2 if room >= 3 else max(0, room)) if p.strict_off else None
    cap = max(1, min(int(p.max_different_shifts), var_cap))
    free = [d for d in range(7) if status[d] is None]
    out = []
    for off_n in ([need_off] if need_off is not None else range(min(2, max(0, room)), room + 1)):
        forced_off = [d for d in range(7) if status[d] == "O"]
        extra = off_n - len(forced_off)
        if extra < 0 or extra > len(free):
            continue
        for offs in itertools.combinations(free, extra):
            off_set = set(offs) | set(forced_off)
            if p.strict_off and not p.separate_off_days and room >= 2 and off_n >= 2:
                if not any(d in off_set and (d + 1) % 7 in off_set for d in range(7)):
                    continue
            work = [d for d in free if d not in off_set]
            fixed_js = {status[d][1] for d in range(7) if isinstance(status[d], tuple)}
            for combo in itertools.product(range(len(shifts)), repeat=len(work)):
                if len(set(combo) | fixed_js) > cap:
                    continue
                row = list(status)
                for d in range(7):
                    if d in off_set:
                        row[d] = "O"
                    elif isinstance(row[d], tuple):
                        row[d] = row[d][1]
                for d, j in zip(work, combo):
                    row[d] = j
                ok = True
                if isinstance(row[0], int) and not E.previous_saturday_compatible(assoc.previous_saturday, shifts[row[0]], p.rest_gap_hours):
                    ok = False
                for d in range(7):
                    a_, b_ = row[d], row[(d + 1) % 7]
                    if ok and isinstance(a_, int) and isinstance(b_, int) and not E.rest_compatible(shifts[a_], shifts[b_], p.rest_gap_hours):
                        ok = False
                if ok and getattr(p, "blank_requirement_mode", "allow") != "allow":
                    for d in range(7):
                        if isinstance(row[d], int):
                            sh = shifts[row[d]]
                            base = d * 96 + sh.start_min // 15
                            if any(t < 7 * 96 and not p.active[t // 96][(t % 96) // p.qslots_per_interval]
                                   for t in range(base, base + sh.duration_q)):
                                ok = False
                                break
                if ok:
                    out.append(tuple(row))
    return out


def build_grouped(E, wb: Path, out_dir: Path, max_shifts: int, time_limit: float, workers: int, var_cap: int) -> Dict[str, Any]:
    """Grouped variant: identical associates share integer tour counts (no symmetry)."""
    from ortools.sat.python import cp_model
    t0 = time.time()
    p = E.parse_input(wb)
    if p.use_11h_3off:
        return {"workbook": wb.name, "status": "REFUSED", "reason": "11H/3OFF out of scope"}
    if p.fixed_enabled and any(a.nesting_group for a in p.associates):
        return {"workbook": wb.name, "status": "REFUSED", "reason": "nesting groups not supported by the grouped variant"}
    cand, ginfo = candidates(E, p, max_shifts, min(45.0, time_limit / 4))
    shifts = [p.shifts[i] for i in cand]
    qpi = p.qslots_per_interval
    min_rules = [r for r in p.language_rules if int(r.minimum or 0) > 0]

    def signature(a):
        assoc = p.associates[a]
        days = []
        for d in range(7):
            pref = assoc.preferences[d] if d < len(assoc.preferences) else ""
            fixed = assoc.fixed_schedule[d] if d < len(assoc.fixed_schedule) else ""
            days.append((E.preference_kind(pref), E.norm(pref) if E.preference_kind(pref) in ("leave", "off") else "",
                         E.preference_kind(fixed), E.norm(fixed)))
        return (tuple(bool(E.language_eligible(r, assoc)) for r in min_rules), tuple(days),
                E.norm(assoc.previous_saturday or ""))
    groups: Dict[Any, List[int]] = {}
    for a in range(len(p.associates)):
        groups.setdefault(signature(a), []).append(a)
    keys = list(groups)
    tours = {k: group_tours(E, p, p.associates[groups[k][0]], shifts, var_cap) for k in keys}
    empty = [k for k in keys if not tours[k]]
    if empty:
        return {"workbook": wb.name, "status": "NO_SEED", "reason": "a group has no legal tour over the candidate shifts",
                "groups": len(keys), **ginfo}
    m = cp_model.CpModel()
    tv = {}
    for gi, k in enumerate(keys):
        vs = [m.NewIntVar(0, len(groups[k]), f"t{gi}_{ti}") for ti in range(len(tours[k]))]
        m.Add(sum(vs) == len(groups[k]))
        tv[k] = vs
    classes = sorted({k[0] for k in keys})
    patterns = E.generate_break_patterns(p)
    onfloor = {c: [[] for _ in range(7 * 96)] for c in classes}
    onbreak_total = [[] for _ in range(7 * 96)]
    zlist = []
    for c in classes:
        for d in range(7):
            for j, sh in enumerate(shifts):
                users = [v for k in keys if k[0] == c for ti, v in enumerate(tv[k]) if tours[k][ti][d] == j]
                if not users:
                    continue
                legal = [pt for pt in patterns if pt.duration_q == sh.duration_q] or [None]
                zs = []
                for pt in legal:
                    z = m.NewIntVar(0, len(p.associates), f"z{classes.index(c)}_{d}_{j}_{pt.index if pt else 'nb'}")
                    zs.append(z); zlist.append(z)
                    broken = set(pt.broken_offsets) if pt else set()
                    base = d * 96 + sh.start_min // 15
                    for o in range(sh.duration_q):
                        t = base + o
                        if t >= 7 * 96:
                            break
                        (onbreak_total[t] if o in broken else onfloor[c][t]).append(z)
                m.Add(sum(zs) == sum(users))
    prior = [len(E.prior_covering_associates(p, q)) for q in range(7 * 96)]
    ratio, absolute = float(p.break_max_concurrent_ratio), int(p.break_max_concurrent_absolute)
    cap_slack = []  # soft in the seed: the engine re-places every break in Stage 2
    for t in range(7 * 96):
        if onbreak_total[t]:
            b = sum(onbreak_total[t]); f = sum(sum(onfloor[c][t]) for c in classes)
            sl = m.NewIntVar(0, len(p.associates), f"cs{t}")
            add_break_cap(m, b, b + f + prior[t], ratio, absolute, slack=sl)
            cap_slack.append(sl)
    rows = []
    for d in range(7):
        for i in range(p.intervals_per_day):
            if not p.active[d][i]:
                continue
            need = p.target_ratio * float(p.requirements[d][i] or 0.0) * qpi
            eff = 1.0 - float(p.shrinkage[d][i])
            if need <= 1e-9 or eff <= 0:
                continue
            coef = int(round(eff * SCALE))
            qs = [d * 96 + i * qpi + kq for kq in range(qpi)]
            lhs = coef * sum(sum(onfloor[c][q]) for c in classes for q in qs) + coef * sum(prior[q] for q in qs)
            rows.append((lhs, int(math.ceil(need * SCALE - 1e-6))))
    lang_short = []
    for t in range(7 * 96):
        d, minute = t // 96, (t % 96) * 15
        if not p.active[d][(t % 96) // qpi]:
            continue
        for ri, r in enumerate(min_rules):
            if r not in E.language_rules_at(p, minute, day=d):
                continue
            elig = sum(sum(onfloor[c][t]) for c in classes if c[ri])
            ls = m.NewIntVar(0, int(r.minimum), f"ls{t}_{ri}")
            m.Add(ls >= int(r.minimum) - len(E.prior_covering_associates(p, t, r)) - elig)
            lang_short.append(ls)
    short = []
    for kk, (lhs, rhs) in enumerate(rows):
        s_ = m.NewIntVar(0, rhs, f"s{kk}")
        m.Add(s_ >= rhs - lhs)
        short.append(s_)
    m.Minimize(SCALE * 100 * sum(lang_short) + SCALE * 10 * sum(cap_slack) + sum(short))
    sv = cp_model.CpSolver()
    sv.parameters.max_time_in_seconds = time_limit * 0.5
    sv.parameters.num_workers = workers
    sv.parameters.random_seed = 1
    st1 = sv.Solve(m)
    info = {"workbook": wb.name, "variant": "grouped", "associates": len(p.associates), "groups": len(keys),
            "tour_vars": sum(len(v) for v in tv.values()), "candidate_shifts": [sh.label for sh in shifts],
            "phase1": {"status": sv.StatusName(st1), "seconds": round(sv.WallTime(), 1)}, **ginfo}
    if st1 not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        info["status"] = "NO_SEED"
        return info
    p1_zero = [sv.Value(s_) == 0 for s_ in short]
    lang1 = sum(sv.Value(v) for v in lang_short)
    info["phase1"].update(shortfall_zero_intervals=sum(p1_zero), language_shortfall=lang1,
                          break_cap_excess=sum(sv.Value(v) for v in cap_slack))
    p1_counts = {k: [sv.Value(v) for v in tv[k]] for k in keys}
    for k in keys:
        for v in tv[k]:
            m.AddHint(v, sv.Value(v))
    for z in zlist:
        m.AddHint(z, sv.Value(z))
    if lang1 == 0:
        for v in lang_short:
            m.Add(v == 0)
    m.ClearObjective()
    hits = []
    for kk, (lhs, rhs) in enumerate(rows):
        h = m.NewBoolVar(f"h{kk}")
        m.Add(lhs >= rhs).OnlyEnforceIf(h)
        m.AddHint(h, 1 if p1_zero[kk] else 0)
        hits.append(h)
    m.Maximize(sum(hits) - 1000 * sum(lang_short) - 100 * sum(cap_slack))
    sv2 = cp_model.CpSolver()
    sv2.parameters.max_time_in_seconds = time_limit * 0.5
    sv2.parameters.num_workers = workers
    sv2.parameters.random_seed = 1
    st2 = sv2.Solve(m)
    p2_ok = st2 in (cp_model.OPTIMAL, cp_model.FEASIBLE)
    p2_hits = sum(sv2.Value(h) for h in hits) if p2_ok else -1
    use1 = sum(p1_zero) >= p2_hits
    counts = p1_counts if use1 else {k: [sv2.Value(v) for v in tv[k]] for k in keys}
    info.update(status="SEED", seed_from_phase=1 if use1 else 2, seed_model_hits=sum(p1_zero) if use1 else p2_hits,
                phase2={"status": sv2.StatusName(st2), "hits": p2_hits, "bound": int(math.floor(sv2.BestObjectiveBound() + 1e-6)) if p2_ok else None},
                active_intervals=sum(1 for d in range(7) for i in range(p.intervals_per_day) if p.active[d][i]))
    week: Dict[int, Tuple] = {}
    for k in keys:
        pool = []
        for ti, c in enumerate(counts[k]):
            pool.extend([tours[k][ti]] * c)
        for a, tour in zip(groups[k], pool):
            week[a] = tour
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (wb.stem + "_AGGSEED.xlsx")
    shutil.copy(wb, out)
    import openpyxl
    book = openpyxl.load_workbook(out)
    ws = book["Schedule"]
    hdr = next(r for r in range(1, 6) if any(str(ws.cell(r, c).value or "").strip() == "Sun" for c in range(1, 30)))
    col = {str(ws.cell(hdr, c).value or "").strip(): c for c in range(1, 40)}
    name_col = col.get("SF Name") or col.get("Name")
    row_of = {E.norm(str(ws.cell(r, name_col).value)): r for r in range(hdr + 1, ws.max_row + 1) if ws.cell(r, name_col).value}
    for a, tour in week.items():
        r = row_of.get(E.norm(p.associates[a].name))
        if r is None:
            continue
        for d, dn in enumerate(DAY_COLS):
            v = tour[d]
            ws.cell(r, col[dn]).value = shifts[v].label if isinstance(v, int) else ("Leave" if v == "L" else "OFF")
    book.save(out)
    seed = E.load_seed_skeleton(E.parse_input(out), out, "input_schedule_seed")
    info.update(seed_workbook=out.name, seed_valid_cells=seed.diagnostics.get("valid_seed_cells") if seed else 0,
                seed_invalid_cells=seed.diagnostics.get("invalid_seed_cells") if seed else None,
                total_seconds=round(time.time() - t0, 1))
    return info


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--engine", type=Path, required=True)
    ap.add_argument("--out-dir", type=Path, required=True)
    ap.add_argument("--max-shifts", type=int, default=12)
    ap.add_argument("--time-limit", type=float, default=300)
    ap.add_argument("--workers", type=int, default=2)
    ap.add_argument("--variant", choices=("grouped", "per-associate"), default="grouped")
    ap.add_argument("--variety-cap", type=int, default=2)
    ap.add_argument("--all-shifts", action="store_true", help="per-associate: every library shift is a candidate")
    ap.add_argument("--fix-week-from", type=Path, help="per-associate diagnostic: pin the week in this workbook's "
                    "'Final Schedule' (or 'Schedule') sheet and score it under the seed model's metric")
    ap.add_argument("workbooks", nargs="+", type=Path)
    a = ap.parse_args()
    E = S.load_engine(a.engine)
    fix = None
    if a.fix_week_from:
        import openpyxl
        book = openpyxl.load_workbook(a.fix_week_from, read_only=True)
        ws = book["Final Schedule"] if "Final Schedule" in book.sheetnames else book["Schedule"]
        rows = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows) if "Sun" in [str(v or "").strip() for v in r])
        hdr = [str(v or "").strip() for v in rows[h]]
        nc, dc = hdr.index("SF Name") if "SF Name" in hdr else hdr.index("Name"), [hdr.index(dn) for dn in DAY_COLS]
        fix = {E.norm(str(r[nc])): [r[c] for c in dc] for r in rows[h + 1:] if r[nc]}
    report = {}
    for wb in a.workbooks:
        row = (build_grouped(E, wb, a.out_dir, a.max_shifts, a.time_limit, a.workers, a.variety_cap)
               if a.variant == "grouped" else build(E, wb, a.out_dir, a.max_shifts, a.time_limit, a.workers,
                                                  a.all_shifts, fix))
        report[wb.name] = row
        print(json.dumps(row, default=str), flush=True)
        a.out_dir.mkdir(parents=True, exist_ok=True)
        (a.out_dir / "AGGREGATE_SEED_REAL_REPORT.json").write_text(json.dumps(report, indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
