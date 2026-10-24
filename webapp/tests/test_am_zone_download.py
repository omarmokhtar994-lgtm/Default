# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AM item 2 (owner, 2026-10-16: "When changing timezone in schedules tab it's working fine now but not
downloadable with the new timezone"). With "Times in" on another zone, Download gives a copy in that zone's clock,
each shift as the page shows it (marked when it starts the day before or the next day) and each break moved the same;
the workbook kept stays as it is."""
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from webapp.tests import test_overtime_side as ots
from webapp.versions import BREAK_TABS, SCHEDULE_TABS, _header, _name_col, _norm


def cells(data):
    wb = load_workbook(io.BytesIO(data))
    ws = wb[next(t for t in SCHEDULE_TABS if t in wb.sheetnames)]
    row, cols = _header(ws, ("sf name", "associate name"))
    shifts = {ws.cell(r, _name_col(cols)).value: ws.cell(r, cols["wed"]).value for r in range(row + 1, ws.max_row + 1)}
    starts = []
    for tab in BREAK_TABS:
        if tab in wb.sheetnames:
            bs = wb[tab]
            brow, bcols = _header(bs, ("associate",))
            starts = [str(bs.cell(r, bcols["start"]).value) for r in range(brow + 1, bs.max_row + 1)
                      if _norm(bs.cell(r, bcols["associate"]).value) == _norm("Associate 008")
                      and _norm(bs.cell(r, bcols["day"]).value) == "wed"]
            break
    return shifts, starts


class TheDownload(ots._Overtime):
    def version(self):
        run = self.store.list_runs()[0]["id"]
        return run, self.app.extensions["schedules"].versions(run)[0]

    def get(self, tz=""):
        _, v = self.version()
        # AN followup: consume the download and close its file before fixture cleanup.
        return self.client.get(f"/schedules/{v['id']}/download" + (f"?tz={tz}" if tz else ""), buffered=True)

    def test_the_link_carries_the_zone(self):
        run, v = self.version()
        body = self.client.get(f"/runs/{run}/schedules?tz=Europe/London").get_data(as_text=True)
        self.assertIn(f'href="/schedules/{v["id"]}/download?tz=Europe/London"', body)
        self.assertIn(f'href="/schedules/{v["id"]}/download"', self.client.get(f"/runs/{run}/schedules").get_data(as_text=True))

    def test_a_copy_in_that_zones_clock(self):
        own, moved = self.get(), self.get("Europe/London")  # Cairo is 2 hours ahead of London on 11 October 2026
        (o_shifts, _), (m_shifts, _) = cells(own.data), cells(moved.data)
        self.assertEqual((o_shifts["Associate 008"], m_shifts["Associate 008"]), ("08:00 - 17:00", "06:00 - 15:00"))
        self.assertTrue(moved.headers["Content-Disposition"].endswith("_London_time.xlsx"))

    def test_across_midnight_it_says_so(self):
        shifts, _ = cells(self.get("America/Los_Angeles").data)
        self.assertEqual(shifts["Associate 008"], "22:00 - 07:00 (day before)")

    def test_the_kept_workbook_is_unchanged(self):
        _, v = self.version()
        kept = self.app.extensions["schedules"].path(v["id"]).read_bytes()
        self.get("Europe/London")
        self.assertEqual(self.app.extensions["schedules"].path(v["id"]).read_bytes(), kept)
        self.assertEqual(self.get().data, kept)


class TheBreaks(unittest.TestCase):
    """A week with breaks (the real B3 AE/AR week, "Associate NN" names): each break moves with its shift."""

    def test_breaks_move_the_same_minutes(self):
        from webapp.tests.test_versions import AFTER
        from webapp.versions import read_week, zoned_copy
        out = Path(tempfile.mkdtemp()) / "zoned.xlsx"
        zoned_copy(AFTER, out, -120)
        week = read_week(out)
        self.assertIn({"associate": "Associate 001", "day": "Wed", "kind": "Lunch", "start": "14:30", "minutes": 30},
                      week["breaks"])
        self.assertEqual(week["associates"][0]["days"][3], "10:00 - 19:00")
