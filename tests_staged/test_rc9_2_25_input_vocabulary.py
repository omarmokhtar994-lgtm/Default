"""Input vocabulary fails closed without rejecting ordinary values (F-02, C-2).

F-02: the Preference/Fixed vocabulary rejected ordinary spellings ("Public
Holiday", "Casual Leave", "OFF (approved)", "N/A", "-") as hard contract
failures, so a workbook that ran under RC5 produced no schedule at all.
C-2: a yes/no instruction holding anything else ("Enable", "Yes.") read as No,
which could switch Leave or Hard OFF off for the whole roster.

Contract smoke over all 42 repository workbooks: 0 contract results and 0
canonical contract hashes changed (evidence/independent_audit_2026_09_28).
"""
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("CANDIDATE_ENGINE", REPO)).resolve()
sys.path.insert(0, str(ROOT / "engine" / "_tools"))
import l632_universal_scheduler as E  # noqa: E402

INPUTS = next(p for p in (ROOT / "inputs", REPO / "packages" / "rc9_2_2_production" / "inputs") if p.exists())
BOOK = INPUTS / "NMG_SP_RC9_1_READY_FIXED.xlsx"


class PreferenceVocabulary(unittest.TestCase):
    def test_ordinary_leave_and_off_spellings_are_recognised(self):
        cases = {
            "leave": ["Public Holiday", "Bank Holiday", "PH", "Casual Leave", "CL", "Medical Leave",
                      "Maternity Leave", "Vacation Leave", "annual-leave", "Leave (approved)", "Unpaid", "LOA"],
            "off": ["Comp Off", "OFF (approved)", "Requested Off", "Off Request", "R/D", "D/O"],
            "blank": ["N/A", "-", "TBD"],
        }
        for kind, values in cases.items():
            for value in values:
                with self.subTest(value=value):
                    self.assertEqual(E.preference_kind(value), kind)

    def test_ambiguous_values_still_fail_closed(self):
        for value in ("Training", "WFH", "Half Day", "9-6", "Morning", "Sick Off", "Off Duty"):
            with self.subTest(value=value):
                self.assertEqual(E.preference_kind(value), "other")

    def test_values_recognised_before_keep_their_meaning(self):
        for value in sorted(E.OFF_WORDS):
            self.assertEqual(E.preference_kind(value), "off", value)
        for value in sorted(E.LEAVE_WORDS):
            self.assertEqual(E.preference_kind(value), "leave", value)
        for value in ("09:00-18:00", "07:00 - 16:00", "22:00-07:00"):
            self.assertEqual(E.preference_kind(value), "shift", value)


def _edit(book: Path, edits):
    from openpyxl import load_workbook
    wb = load_workbook(book)
    for edit in edits:
        edit(wb)
    wb.save(book)


def _set_instruction(key: str, value):
    def edit(wb):
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and E.norm(cell.value) == E.norm(key):
                        ws.cell(cell.row, cell.column + 1, value)
                        return
        raise AssertionError(f"instruction {key!r} not found")
    return edit


def _set_first_preference(value):
    def edit(wb):
        ws = next(wb[n] for n in wb.sheetnames if E.norm(n) in {"preference", "prefrence", "preferences"})
        header = E._find_header_row(ws, ["name"], prefer_day_columns=True)
        headers = {E.norm(ws.cell(header, c).value): c for c in range(1, ws.max_column + 1)}
        name_col = next((c for h, c in headers.items() if "name" in h), 1)
        day_cols = E._day_columns(ws, header)
        first_associate = E.parse_input(BOOK).associates[0].name
        row = header + 1
        ws.cell(row, name_col, first_associate)
        ws.cell(row, day_cols[0], value)
    return edit


def _add_mapping(rows):
    def edit(wb):
        ws = wb.create_sheet("Preference Code Mapping")
        ws.append(["Value", "Meaning"])
        for row in rows:
            ws.append(list(row))
    return edit


class ContractBehaviour(unittest.TestCase):
    def contract(self, *edits):
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / BOOK.name
            shutil.copy(BOOK, book)
            _edit(book, edits)
            parsed = E.parse_input(book)
            codes = {f.get("code") for f in E.validate_input_contract(parsed).get("failures", [])}
            return parsed, codes

    def test_the_unedited_workbook_passes(self):
        _, codes = self.contract()
        self.assertFalse(codes)

    def test_an_unknown_switch_value_fails_the_contract(self):
        _, codes = self.contract(_set_instruction("Leave Enabled", "Enable"))
        self.assertIn("HARD_INVALID_INSTRUCTION_BOOLEAN", codes)

    def test_yes_and_no_still_pass(self):
        for value in ("Yes", "No"):
            _, codes = self.contract(_set_instruction("Leave Enabled", value))
            self.assertNotIn("HARD_INVALID_INSTRUCTION_BOOLEAN", codes, value)

    def test_an_ordinary_leave_spelling_no_longer_blocks_the_run(self):
        parsed, codes = self.contract(_set_first_preference("Public Holiday"))
        self.assertNotIn("UNRECOGNISED_PREFERENCE_VALUE", codes)

    def test_an_unknown_code_still_blocks_the_run(self):
        _, codes = self.contract(_set_first_preference("TRN"))
        self.assertIn("UNRECOGNISED_PREFERENCE_VALUE", codes)

    def test_a_mapping_sheet_turns_a_site_code_into_leave(self):
        parsed, codes = self.contract(_set_first_preference("TRN"), _add_mapping([("TRN", "Leave")]))
        self.assertNotIn("UNRECOGNISED_PREFERENCE_VALUE", codes)
        self.assertTrue(any("Leave" in a.preferences for a in parsed.associates))
        self.assertTrue(any(w.startswith("PREFERENCE_MAPPING_APPLIED") for w in parsed.parser_warnings))

    def test_a_mapping_with_an_unknown_meaning_fails(self):
        _, codes = self.contract(_add_mapping([("TRN", "Training")]))
        self.assertIn("HARD_INVALID_PREFERENCE_MAPPING", codes)


class BaselineWorkbookDepartedOverride(unittest.TestCase):
    """AE_AR_B2B is baseline-protected; its manifest names 5 departed associates (F-21)."""

    def test_the_manifest_names_clear_exactly_the_unknown_previous_week_rows(self):
        import json
        manifest = next(p for p in (ROOT / "SCENARIOS.json",
                                    REPO / "packages" / "rc9_2_2_production" / "SCENARIOS.json") if p.exists())
        row = next(r for r in json.loads(manifest.read_text())["scenarios"] if r["scenario_id"] == "AE_AR_B2B")
        book = INPUTS / row["input"]
        codes = lambda parsed: {f.get("code") for f in E.validate_input_contract(parsed).get("failures", [])}
        self.assertIn("HARD_PREVIOUS_SATURDAY_UNKNOWN_ASSOCIATE", codes(E.parse_input(book)))
        parsed = E.parse_input(book, acknowledged_departed_override=row["acknowledged_departed"])
        self.assertFalse(codes(parsed))
        self.assertEqual(sum(w.startswith("DEPARTED_ASSOCIATE_ROW_IGNORED") for w in parsed.parser_warnings),
                         len(row["acknowledged_departed"]))
        # acknowledging one person must not silence another
        partial = E.parse_input(book, acknowledged_departed_override=row["acknowledged_departed"][:-1])
        self.assertIn("HARD_PREVIOUS_SATURDAY_UNKNOWN_ASSOCIATE", codes(partial))


class TheTemplateCarriesTheDepartedRow(unittest.TestCase):
    """The rebuilt template has a visible 'Known Departed Associates' row, and
    filling it is the workbook's own way to clear the unknown-associate failure."""

    def test_listing_the_names_in_the_template_row_clears_exactly_those(self):
        import json
        import subprocess
        from openpyxl import load_workbook
        manifest = next(p for p in (ROOT / "SCENARIOS.json",
                                    REPO / "packages" / "rc9_2_2_production" / "SCENARIOS.json") if p.exists())
        row = next(r for r in json.loads(manifest.read_text())["scenarios"] if r["scenario_id"] == "AE_AR_B2B")
        names = row["acknowledged_departed"]
        codes = lambda book: {f.get("code") for f in E.validate_input_contract(E.parse_input(book)).get("failures", [])}
        with tempfile.TemporaryDirectory() as tmp:
            rebuilt = Path(tmp) / "rebuilt.xlsx"
            subprocess.run([sys.executable, str(ROOT / "tools" / "build_input_template.py"),
                            str(INPUTS / row["input"]), str(rebuilt)], check=True, capture_output=True)
            for listed, expect_clean in ((names[:-1], False), (names, True)):
                wb = load_workbook(rebuilt)
                ws = wb["Instructions"]
                cell = next(c for r in ws.iter_rows() for c in r
                            if isinstance(c.value, str) and c.value.strip() == "Known Departed Associates")
                ws.cell(cell.row, cell.column + 1, "; ".join(listed))
                book = Path(tmp) / f"filled_{len(listed)}.xlsx"
                wb.save(book)
                self.assertEqual("HARD_PREVIOUS_SATURDAY_UNKNOWN_ASSOCIATE" not in codes(book), expect_clean, len(listed))


class TheWorkbookPreCheckSaysWhatTheRunWillDo(unittest.TestCase):
    """tools/check_input_workbook.py: exit 0 exactly when the run would accept it."""

    def check(self, *args):
        import subprocess
        return subprocess.run([sys.executable, str(ROOT / "tools" / "check_input_workbook.py"), *map(str, args)],
                              capture_output=True, text=True, timeout=300)

    def test_accepted_refused_and_acknowledged(self):
        ok = self.check(BOOK)
        self.assertEqual(ok.returncode, 0, ok.stdout)
        self.assertIn("ACCEPTED", ok.stdout)
        refused = self.check(INPUTS / "AE_AR_B2B.xlsx")
        self.assertEqual(refused.returncode, 1)
        self.assertIn("Known Departed Associates", refused.stdout)
        import json
        manifest = next(p for p in (ROOT / "SCENARIOS.json",
                                    REPO / "packages" / "rc9_2_2_production" / "SCENARIOS.json") if p.exists())
        row = next(r for r in json.loads(manifest.read_text())["scenarios"] if r["scenario_id"] == "AE_AR_B2B")
        acknowledged = self.check(INPUTS / "AE_AR_B2B.xlsx", "--acknowledge-departed", "; ".join(row["acknowledged_departed"]))
        self.assertEqual(acknowledged.returncode, 0, acknowledged.stdout)

    def test_a_typed_yes_with_a_dot_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / BOOK.name
            shutil.copy(BOOK, book)
            _edit(book, [_set_instruction("Leave Enabled", "Yes.")])
            out = self.check(book)
            self.assertEqual(out.returncode, 1)
            self.assertIn("HARD_INVALID_INSTRUCTION_BOOLEAN", out.stdout)


VOICE = INPUTS / "Cricut_Voice_RC9_1_READY_SKELETON.xlsx"


class ThePreCheckShowsTheLanguageHours(unittest.TestCase):
    """Language Setup's Coverage Start/End limit who works when only if the
    Instructions row "Language Working Window" is set. A Voice run reported
    English associates in International hours and the other way round: the
    window was OFF, and Cricut Voice's roster has English-labelled associates
    with fixed 03:00-05:00 starts, so switching it on is refused as infeasible
    without naming why. The pre-check now says both before the run."""

    check = TheWorkbookPreCheckSaysWhatTheRunWillDo.check

    def test_off_is_shown_with_a_warning_and_the_conflicts_it_would_hit(self):
        out = self.check(VOICE)
        self.assertEqual(out.returncode, 0, out.stdout)
        self.assertIn("Language Working Window = OFF", out.stdout)
        self.assertIn("these hours do NOT limit", out.stdout)
        self.assertIn("if you set ALL_ROWS, 20 fixed request(s)", out.stdout)
        self.assertIn("Jhonny Mascarenhas (English) Sun: fixed 05:00 - 14:00, English hours 16:00-03:00", out.stdout)

    def test_an_enforced_window_that_contradicts_a_fixed_request_is_refused(self):
        import subprocess
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "voice_template.xlsx"
            subprocess.run([sys.executable, str(ROOT / "tools" / "build_input_template.py"), str(VOICE), str(book)],
                           check=True, capture_output=True)
            _edit(book, [_set_instruction("Language Working Window", "ALL_ROWS")])
            out = self.check(book)
            self.assertEqual(out.returncode, 1, out.stdout)
            self.assertIn("Language Working Window = ALL_ROWS", out.stdout)
            self.assertIn("FAIL: 20 fixed request(s) start outside", out.stdout)
            self.assertNotIn("these hours do NOT limit", out.stdout)

    def test_per_day_rows_are_shown_per_day(self):
        import datetime
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / VOICE.name
            shutil.copy(VOICE, book)
            wb = load_workbook(book)
            ws = wb["Language Setup"]
            ws.cell(2, 10, "Coverage Days")
            ws.cell(3, 10, "Mon-Fri")
            for c in range(1, 10):
                ws.cell(5, c, ws.cell(3, c).value)
            ws.cell(5, 2, datetime.time(17, 0))
            ws.cell(5, 3, datetime.time(4, 0))
            ws.cell(5, 10, "Sat,Sun")
            wb.save(book)
            out = self.check(book)
            self.assertIn("english: Sun 17:00-04:00, Mon 16:00-03:00, Tue 16:00-03:00, Wed 16:00-03:00, "
                          "Thu 16:00-03:00, Fri 16:00-03:00, Sat 17:00-04:00", out.stdout)
            # Sunday's own window is 17:00-04:00, so a fixed 05:00 start is still outside it.
            self.assertIn("Jhonny Mascarenhas (English) Sun: fixed 05:00 - 14:00, English hours 17:00-04:00",
                          out.stdout)


class TheRunNamesTheConflictingRules(unittest.TestCase):
    """A hard-rule contradiction used to print each constraint-isolation row as
    a raw solver dump under "Main blockers". It now names the rule families
    that conflict and, for an enforced language window, each fixed request
    that starts outside its person's hours."""

    ISOLATION = [
        {"relaxed_family": "language", "status": "FEASIBLE", "interpretation": "x",
         "diagnostics": {"solver_telemetry": {"num_conflicts": 0}}},
        {"relaxed_family": "fixed", "status": "FEASIBLE", "diagnostics": {"variables": 1}},
        {"relaxed_family": "rest", "status": "INFEASIBLE", "diagnostics": {"variables": 1}},
    ]

    def test_the_engine_lists_fixed_requests_outside_an_enforced_window(self):
        parsed = E.parse_input(VOICE)
        self.assertEqual(E.fixed_requests_outside_language_windows(parsed), [],
                         "with the window OFF nothing contradicts")
        parsed.language_working_window_mode = "ALL_ROWS"
        rows = E.fixed_requests_outside_language_windows(parsed)
        self.assertEqual(len(rows), 20)
        self.assertIn({"associate": "Jhonny Mascarenhas", "language": "English", "day": "Sun",
                       "shift": "05:00 - 14:00", "window": "16:00-03:00"}, rows)

    def test_the_outcome_text_names_rules_and_requests_without_solver_dumps(self):
        audit = {"status": "FAIL_HARD_CONTRACT_INFEASIBLE", "constraint_isolation": self.ISOLATION,
                 "hard_conflict_examples": [{"associate": "A B", "language": "English", "day": "Sun",
                                             "shift": "05:00 - 14:00", "window": "16:00-03:00"}]}
        text = E.format_business_outcome(E.build_business_outcome(audit, 2))
        self.assertNotIn("diagnostics", text)
        self.assertNotIn("{", text)
        self.assertIn("fixed requests (Fixed Request sheet): relaxing this rule alone makes a schedule possible", text)
        self.assertIn("language rules", text)
        self.assertNotIn("minimum rest between shifts", text)
        self.assertIn("A B (English) | Sun | fixed 05:00 - 14:00 | language hours 16:00-03:00", text)
        self.assertIn("1 fixed request(s) start outside", text)

    def test_no_single_family_and_no_isolation_are_said_plainly(self):
        combo = E.build_business_outcome({"status": "FAIL_HARD_CONTRACT_INFEASIBLE",
                                          "constraint_isolation": self.ISOLATION[2:]}, 2)
        self.assertIn("No single rule family explains it", combo["plain_language_summary"])
        none = E.build_business_outcome({"status": "FAIL_HARD_CONTRACT_INFEASIBLE"}, 2)
        self.assertIn("not enough time left", none["plain_language_summary"])

    def test_the_results_folder_keeps_the_engine_diagnosis(self):
        """The wrapper used to replace it with "the engine output problem" and
        wrote only the summary paragraph to BUSINESS_OUTCOME.txt."""
        import importlib.util
        import json
        spec = importlib.util.spec_from_file_location(
            "wrapper_hard_conflict", ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py")
        wrapper = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = wrapper
        spec.loader.exec_module(wrapper)
        audit = {"status": "FAIL_HARD_CONTRACT_INFEASIBLE", "constraint_isolation": self.ISOLATION,
                 "hard_conflict_examples": [{"associate": "A B", "language": "English", "day": "Sun",
                                             "shift": "05:00 - 14:00", "window": "16:00-03:00"}]}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "BUSINESS_OUTCOME.json").write_text(json.dumps(E.build_business_outcome(audit, 2)))
            (root / "CASE.l6_3_2_3_solver_audit.json").write_text(json.dumps({"status": audit["status"]}))
            wrapper.reconcile_business_outcome_after_validation(root, {"status": "NOT_RUN", "return_code": None}, 2)
            outcome = json.loads((root / "BUSINESS_OUTCOME.json").read_text())
            text = (root / "BUSINESS_OUTCOME.txt").read_text()
        self.assertEqual(outcome["outcome_code"], "HARD_RULE_COMBINATION_INFEASIBLE")
        self.assertEqual(outcome["technical_status"], "FAIL_HARD_CONTRACT_INFEASIBLE")
        self.assertFalse(outcome["production_eligible"])
        self.assertNotIn("engine output problem", text)
        self.assertIn("No schedule satisfies all hard rules together", text)
        self.assertIn("fixed requests (Fixed Request sheet): relaxing this rule alone", text)
        self.assertIn("A B (English) | Sun | fixed 05:00 - 14:00 | language hours 16:00-03:00", text)
        self.assertIn("Required action:", text)


LANG_RUN = REPO / "fixtures" / "real_runs" / "language_hours"


class TheValidatorChecksLanguageHoursInEveryEnforcedMode(unittest.TestCase):
    """The validator checked shift starts against language hours only under
    ALL_ROWS and MINIMUM_ROWS, and its REQUIRED_LANGUAGE_ONLY check sat inside
    that branch, so it never ran. The engine enforces the hours in all three
    modes. Fixture: the enforced Voice run of 2026-10-03 (evidence/language_hours)."""

    INPUT = LANG_RUN / "Cricut_Voice_LANGUAGE_HOURS.xlsx"
    OUTPUT = LANG_RUN / "VOICE_FINAL_SHEET_L6_3_2_3_BEST_FINAL_AFTER_BREAKS_SCHEDULE.xlsx"

    def validate(self, book, output):
        sys.path.insert(0, str(ROOT / "engine" / "tools"))
        import independent_validator as V
        result = V.validate(Path(book), Path(output), ROOT / "engine" / "_tools" / "l632_universal_scheduler.py")
        return [f for f in result.get("failures", []) if "LANGUAGE" in str(f.get("type"))]

    def _move_international_to_evening(self, tmp):
        from openpyxl import load_workbook
        out = Path(tmp) / self.OUTPUT.name
        shutil.copy(self.OUTPUT, out)
        wb = load_workbook(out)
        ws = wb["Schedule"]
        header = next(r for r in range(1, 10) if "Language" in [c.value for c in ws[r]])
        cols = {c.value: c.column for c in ws[header]}
        for r in range(header + 1, ws.max_row + 1):
            if ws.cell(r, cols["Language"]).value == "International":
                for day in ("Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"):
                    if ":" in str(ws.cell(r, cols[day]).value or ""):
                        ws.cell(r, cols[day], "18:00 - 03:00")
                        wb.save(out)
                        return out, ws.cell(r, cols["SF Name"]).value, day
        raise AssertionError("no International shift in the fixture")

    def _with_mode(self, tmp, mode, international_minimum=None):
        book = Path(tmp) / f"in_{mode}.xlsx"
        shutil.copy(self.INPUT, book)
        edits = [_set_instruction("Language Working Window", mode)]
        if international_minimum is not None:
            def set_minimum(wb):
                ws = wb["Language Setup"]
                cols = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
                for r in range(3, ws.max_row + 1):
                    if ws.cell(r, cols["Language"]).value == "International":
                        ws.cell(r, cols["Minimum Per Interval"], international_minimum)
            edits.append(set_minimum)
        _edit(book, edits)
        return book

    def test_the_enforced_run_passes_and_a_moved_shift_is_caught_under_all_rows(self):
        self.assertEqual(self.validate(self.INPUT, self.OUTPUT), [])
        with tempfile.TemporaryDirectory() as tmp:
            out, name, day = self._move_international_to_evening(tmp)
            failures = self.validate(self.INPUT, out)
        self.assertIn(("LANGUAGE_WORKING_WINDOW", name, day),
                      {(f["type"], f["associate"], f["day"]) for f in failures})

    def test_required_language_only_also_checks_the_hours(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = self._with_mode(tmp, "REQUIRED_LANGUAGE_ONLY")
            self.assertEqual(self.validate(book, self.OUTPUT), [])
            out, name, day = self._move_international_to_evening(tmp)
            failures = self.validate(book, out)
        self.assertIn(("LANGUAGE_WORKING_WINDOW", name, day),
                      {(f["type"], f["associate"], f["day"]) for f in failures})

    def test_required_language_only_exclusivity_is_reachable(self):
        # International minimum 1 in 00:00-16:00: English associates cannot
        # cover International, so English shifts reaching into those hours are
        # violations under REQUIRED_LANGUAGE_ONLY and only under it.
        with tempfile.TemporaryDirectory() as tmp:
            required = self._with_mode(tmp, "REQUIRED_LANGUAGE_ONLY", international_minimum=1)
            all_rows = self._with_mode(tmp, "ALL_ROWS", international_minimum=1)
            types_required = {f["type"] for f in self.validate(required, self.OUTPUT)}
            types_all_rows = {f["type"] for f in self.validate(all_rows, self.OUTPUT)}
        self.assertIn("REQUIRED_LANGUAGE_ONLY_VIOLATION", types_required)
        self.assertNotIn("REQUIRED_LANGUAGE_ONLY_VIOLATION", types_all_rows)


class AStaleRunLockNeverBlocksANewRun(unittest.TestCase):
    """A Colab portfolio ended in 0.4 s with no schedule: every seed refused
    with "Case is already running" because RUN_LOCK.json files left on Drive by
    an interrupted run named pids that, in the new VM, belonged to other live
    processes. Reproduced with pid 1 in the lock (2026-10-03)."""

    def setUp(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location(
            "wrapper_run_lock", ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py")
        self.w = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = self.w
        spec.loader.exec_module(self.w)
        self.tmp = tempfile.TemporaryDirectory()
        self.case = Path(self.tmp.name)

    def tearDown(self):
        self.w.release_case_lock()
        self.tmp.cleanup()

    def _lock(self, **fields):
        import json
        (self.case / "RUN_LOCK.json").write_text(json.dumps({"schema_version": 1, "pid": 1, **fields}))

    def _heartbeat(self, seconds_ago):
        import json
        from datetime import datetime, timedelta, timezone
        when = datetime.now(timezone.utc) - timedelta(seconds=seconds_ago)
        (self.case / "RUN_HEARTBEAT.json").write_text(json.dumps({"pid": 1, "heartbeat_utc": when.isoformat()}))

    def test_a_lock_from_another_session_with_a_live_pid_is_cleared(self):
        self._lock()  # written before this check existed: no machine identity
        self.w.acquire_case_lock(self.case, "X")
        self.assertTrue((self.case / "RUN_LOCK.json").exists())

    def test_a_lock_from_another_machine_is_cleared_once_its_heartbeat_stops(self):
        self._lock(machine="another-vm")
        self._heartbeat(seconds_ago=3600)
        self.w.acquire_case_lock(self.case, "X")

    def test_a_live_run_elsewhere_still_holds_the_lock(self):
        self._lock(machine="another-vm")
        self._heartbeat(seconds_ago=10)
        with self.assertRaisesRegex(RuntimeError, "already running"):
            self.w.acquire_case_lock(self.case, "X")

    def test_a_live_run_on_this_machine_still_holds_the_lock(self):
        import os
        self._lock(pid=os.getpid(), machine=self.w._machine_identity(),
                   pid_start=self.w._process_start_token(os.getpid()))
        with self.assertRaisesRegex(RuntimeError, "already running"):
            self.w.acquire_case_lock(self.case, "X")

    def test_a_reused_pid_on_this_machine_is_cleared(self):
        self._lock(pid=1, machine=self.w._machine_identity(), pid_start="not-the-start-of-pid-1")
        self.w.acquire_case_lock(self.case, "X")


class OnlyTheRowTheEngineReadsIsChecked(unittest.TestCase):
    def test_a_later_alias_with_prose_is_ignored(self):
        warnings = []
        im = {E.norm("Use Preferences"): "No", E.norm("Preferences"): "Separate OFF From Preference Only"}
        E._validate_boolean_instructions(im, warnings)
        self.assertEqual(warnings, [])
        E._validate_boolean_instructions({E.norm("Preferences"): "Separate OFF"}, warnings)
        self.assertEqual(len(warnings), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
