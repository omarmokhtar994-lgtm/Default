# © 2026 Omar Mokhtar. All rights reserved.
"""Phase N task 4: attendance and actual breaks kept per program and shift date.

The day is read from the version in use for its week (the tool's own version,
with a note, when none is). Statuses and break moves are checked before they
are kept, and every one is logged with who and when."""
import shutil
import tempfile
import time
import unittest
from datetime import date
from pathlib import Path

from webapp.attendance import DayBook, aux_details, hm
from webapp.day import board
from webapp.day import BreakRefused
from webapp.schedules import ScheduleBook
from webapp.store import Store
from webapp.versions import shift_span
from webapp.tests.test_schedules import AFTER, INPUT, REPO

WED = date(2026, 10, 14)  # the run's week starts Sunday 11 October
SAT = date(2026, 10, 17)
SUN, MON, TUE, THU = date(2026, 10, 11), date(2026, 10, 12), date(2026, 10, 13), date(2026, 10, 15)


class TheAttendance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        """One run with its tool version (the validator runs once); each test gets a copy."""
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program="AE/AR B2B",
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.book = ScheduleBook(self.store, self.data, REPO)
        self.tool = self.book.versions("aaaaaaaaaaaa")[0]
        self.days = DayBook(self.store, self.book)

    def lane(self, page, name):
        return next(l for l in page["view"]["lanes"] if l["name"] == name)

    def test_tool_version_with_a_note_until_one_is_in_use(self):
        page = self.days.page("AE/AR B2B", WED)
        self.assertEqual(page["version"]["id"], self.tool["id"])
        self.assertIn("No version is marked in use", page["note"])
        draft = self.book.change(self.tool["id"], self.sara, "Associate 001", "Wed", "10:00 - 19:00", "swap")
        self.book.set_in_use(draft, self.sara)
        page = self.days.page("AE/AR B2B", WED)
        self.assertEqual((page["version"]["id"], page["note"]), (draft, ""))
        self.assertEqual(self.lane(page, "Associate 001")["segments"][0]["label"], "10:00 - 19:00")
        self.assertEqual(self.days.programs(), ["AE/AR B2B"])
        self.assertIsNone(self.days.page("AE/AR B2B", date(2026, 11, 4)))  # no schedule that week

    def test_status_kept_and_logged(self):
        before = self.days.page("AE/AR B2B", WED)["view"]["tiles"]
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Unplanned leave", self.sara)
        page = self.days.page("AE/AR B2B", WED)
        self.assertEqual(page["view"]["tiles"]["absent"], before["absent"] + 1)
        self.assertEqual(self.lane(page, "Associate 001")["segments"][0]["status"], "Unplanned leave")
        self.assertEqual([(e["by_name"], e["associate"], e["what"]) for e in page["log"]],
                         [("Sara", "Associate 001", "Unplanned leave")])
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Present", self.sara)
        self.assertEqual(self.days.page("AE/AR B2B", WED)["view"]["tiles"]["absent"], before["absent"])

    def test_timed_statuses_need_times_inside_the_shift(self):
        with self.assertRaises(ValueError):  # a late arrival needs the time they arrived
            self.days.set_status("AE/AR B2B", WED, "Associate 001", "Late", self.sara)
        with self.assertRaises(ValueError):  # Wednesday is 12:00 - 21:00
            self.days.set_status("AE/AR B2B", WED, "Associate 001", "Late", self.sara, end="22:00")
        with self.assertRaises(ValueError):
            self.days.set_status("AE/AR B2B", WED, "Associate 001", "Meeting", self.sara, start="15:00", end="14:00")
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Late", self.sara, end="13:00")
        seg = self.lane(self.days.page("AE/AR B2B", WED), "Associate 001")["segments"][0]
        self.assertEqual((seg["status"], seg["away"]), ("Late", (12 * 60, 13 * 60)))
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Meeting", self.sara, start="15:00", end="16:00")
        self.assertEqual(self.lane(self.days.page("AE/AR B2B", WED), "Associate 001")["segments"][0]["away"],
                         (15 * 60, 16 * 60))

    def test_aux_is_billable_or_not_and_the_measure_decides(self):
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Coaching", self.sara, start="15:00", end="16:00",
                             billable=True)
        page = self.days.page("AE/AR B2B", WED)
        seg = self.lane(page, "Associate 001")["segments"][0]
        self.assertEqual((seg["status"], seg["billable"]), ("Coaching", True))
        self.assertEqual(page["log"][-1]["what"], "Coaching (billable) 15:00 to 16:00")
        cell = lambda measure: next(c for c in self.days.page("AE/AR B2B", WED, measure)["view"]["cells"]
                                    if c["t"] == 15 * 60)
        self.assertEqual(cell("interval")["now"], cell("interval")["plan"])  # billable: still on the floor
        self.assertEqual(cell("sl")["now"], cell("sl")["plan"] - 1)  # any aux takes them off the queue
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Coaching", self.sara, start="15:00", end="16:00")
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"][-1]["what"], "Coaching (non-billable) 15:00 to 16:00")
        self.assertEqual(cell("interval")["now"], cell("interval")["plan"] - 1)
        with self.assertRaises(ValueError):
            self.days.page("AE/AR B2B", WED, "weekly")

    def test_break_advice_warns_on_gaps_and_ranks_the_best_times(self):
        # Wednesday 12:00 - 21:00: Break 1 13:45-14:00, Lunch 16:30, Break 2 19:30; gaps 60 to 210 minutes
        advice = self.days.break_advice("AE/AR B2B", WED, "Associate 001", 1, "14:30")
        self.assertEqual(advice["warnings"], ["Break 1 ends and Lunch starts 30 minutes apart; this program's "
                                              "minimum gap between breaks is 60 minutes.",
                                              "Lunch ends and Break 2 starts 270 minutes apart; this program's "
                                              "normal maximum gap between breaks is 210 minutes."])
        fits = advice["fits"]
        self.assertTrue(0 < len(fits) <= 3)
        for fit in fits:  # Lunch between Break 1 and Break 2, within the 60 to 210 minute gaps
            self.assertTrue("15:30" <= fit["start"] <= "17:30", fit)
        self.assertEqual([f["buffer"] for f in fits], sorted((f["buffer"] for f in fits), reverse=True))
        self.assertEqual(self.days.break_advice("AE/AR B2B", WED, "Associate 001", 1, "16:30")["warnings"], [])

    def test_a_missing_input_workbook_is_said_plainly(self):
        self.book.input_path("aaaaaaaaaaaa").unlink()
        with self.assertRaises(ValueError) as said:
            self.days.page("AE/AR B2B", WED)
        self.assertIn("input workbook kept with Tool, after breaks is missing", str(said.exception))

    def acts(self, name):
        seg = self.lane(self.days.page("AE/AR B2B", WED), name)["segments"][0]
        return [(a["kind"], a["start"], a["end"], a["billable"]) for a in seg["activities"]]

    def test_activity_added_logged_and_cancelled(self):
        made = self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Meeting", "15:00", "15:30", self.sara)
        self.assertEqual(self.acts("Associate 001"), [("Meeting", 900, 930, False)])
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"][-1]["what"], "Meeting 15:00 to 15:30 (non-billable)")
        self.days.cancel_activity("AE/AR B2B", WED, made, self.sara)
        self.assertEqual(self.acts("Associate 001"), [])
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"][-1]["what"],
                         "Cancelled: Meeting 15:00 to 15:30 (non-billable)")

    def test_activity_rules(self):
        with self.assertRaises(ValueError):  # Wednesday is 12:00 - 21:00
            self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Training", "11:00", "12:00", self.sara)
        self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Training", "15:00", "16:00", self.sara, billable=True)
        with self.assertRaises(ValueError):  # on top of the training
            self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Coaching", "15:30", "16:30", self.sara)
        with self.assertRaises(ValueError):  # overtime must touch the shift
            self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Overtime", "22:00", "23:00", self.sara)
        with self.assertRaises(ValueError):  # at most 2 hours
            self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Overtime", "21:00", "23:30", self.sara)
        with self.assertRaises(ValueError):  # Tuesday 14:00 - 23:00: in at 09:00 leaves 10 hours' rest
            self.days.add_activity("AE/AR B2B", WED, "Associate 013", "Overtime", "09:00", "11:00", self.sara)
        with self.assertRaises(ValueError):  # Thursday 05:00 - 14:00: staying to 19:00 leaves 10 hours' rest
            self.days.add_activity("AE/AR B2B", WED, "Associate 008", "Overtime", "17:00", "19:00", self.sara)
        with self.assertRaises(ValueError):
            self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Lunch party", "15:00", "16:00", self.sara)
        self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Overtime", "21:00", "22:00", self.sara)
        self.days.add_activity("AE/AR B2B", WED, "Associate 001", "VTO", "20:00", "21:00", self.sara)
        self.assertEqual([a[0] for a in self.acts("Associate 001")], ["Training", "VTO", "Overtime"])

    def test_booking_a_meeting_is_all_or_nothing(self):
        slot = self.days.meeting_slots("AE/AR B2B", WED, ["Associate 001", "Associate 013"], 30, "13:00", "19:00")[0]
        self.days.book_session("AE/AR B2B", WED, ["Associate 001", "Associate 013"], slot["start"], 30, "Training", self.sara,
                       billable=True)
        self.assertEqual(len(self.acts("Associate 001")), 1)
        self.assertEqual(len(self.acts("Associate 013")), 1)
        with self.assertRaises(ValueError):  # Associate 013 is now busy then: nobody is booked
            self.days.book_session("AE/AR B2B", WED, ["Associate 012", "Associate 013"], slot["start"], 30, "Meeting", self.sara)
        self.assertEqual(self.acts("Associate 012"), [])

    def test_a_day_off_called_in_works_like_any_shift(self):
        # Associate 002: Tuesday 09:00 - 18:00, off on Wednesday
        self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "12:00", "21:00", self.sara)
        page = self.days.page("AE/AR B2B", WED)
        seg = self.lane(page, "Associate 002")["segments"][0]
        self.assertEqual((seg["label"], seg["called_in"], [b["kind"] for b in seg["breaks"]]),
                         ("12:00 - 21:00 (called in)", True, ["Break 1", "Lunch", "Break 2"]))
        self.assertEqual(page["log"][-1]["what"], "Day off cancelled: called in 12:00 - 21:00")
        self.days.move_break("AE/AR B2B", WED, "Associate 002", 1, "17:00", self.sara)  # its lunch, planned 16:30
        self.days.set_status("AE/AR B2B", WED, "Associate 002", "Late", self.sara, end="12:30")
        seg = self.lane(self.days.page("AE/AR B2B", WED), "Associate 002")["segments"][0]
        self.assertEqual((seg["breaks"][1]["start"], seg["breaks"][1]["moved"], seg["status"]), (17 * 60, True, "Late"))
        self.assertEqual(self.days.page("AE/AR B2B", WED)["stale"], 0)
        self.assertIn("gap", " ".join(self.days.break_advice("AE/AR B2B", WED, "Associate 002", 1, "14:15")["warnings"]))

    def test_calling_in_keeps_the_rules(self):
        with self.assertRaises(ValueError):  # works that day: overtime instead
            self.days.add_activity("AE/AR B2B", WED, "Associate 001", "Called in", "12:00", "21:00", self.sara)
        with self.assertRaises(ValueError):  # not a Shift Library shift, and typed times go in 15-minute steps
            self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "12:05", "21:05", self.sara)
        with self.assertRaises(ValueError):  # Tuesday 23:00 - 08:00 leaves 1 hour's rest before 09:00
            self.days.add_activity("AE/AR B2B", WED, "Associate 004", "Called in", "09:00", "18:00", self.sara)
        self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "12:00", "21:00", self.sara)
        with self.assertRaises(ValueError):  # once a day
            self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "09:00", "18:00", self.sara)

    def test_the_rest_gap_reaches_across_the_week_and_what_was_added(self):
        # Sunday: no schedule is kept for the week before, so the previous Saturday comes from the input
        with self.assertRaises(ValueError) as refused:  # Associate 001: Saturday 16:00 - 01:00
            self.days.add_activity("AE/AR B2B", SUN, "Associate 001", "Called in", "09:00", "18:00", self.sara)
        self.assertIn("would rest 8 hours after the previous shift", str(refused.exception))
        self.days.add_activity("AE/AR B2B", SUN, "Associate 001", "Called in", "13:00", "22:00", self.sara)
        with self.assertRaises(ValueError) as refused:  # Associate 010: Saturday 18:00 - 03:00, Sunday from 16:00
            self.days.add_activity("AE/AR B2B", SUN, "Associate 010", "Overtime", "14:00", "16:00", self.sara)
        self.assertIn("would rest 11 hours after the previous shift", str(refused.exception))
        # a call-in is a shift for the days either side: Associate 002 is off Wednesday and Thursday
        self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "15:00", "00:00", self.sara)
        with self.assertRaises(ValueError) as refused:
            self.days.add_activity("AE/AR B2B", THU, "Associate 002", "Called in", "09:00", "18:00", self.sara)
        self.assertIn("would rest 9 hours after the previous shift", str(refused.exception))
        # and overtime lengthens the shift: Associate 003 works Monday 09:00 - 18:00, off on Tuesday
        self.days.add_activity("AE/AR B2B", MON, "Associate 003", "Overtime", "18:00", "20:00", self.sara)
        with self.assertRaises(ValueError) as refused:
            self.days.add_activity("AE/AR B2B", TUE, "Associate 003", "Called in", "07:00", "16:00", self.sara)
        self.assertIn("would rest 11 hours after the previous shift", str(refused.exception))
        self.days.add_activity("AE/AR B2B", TUE, "Associate 003", "Called in", "08:00", "17:00", self.sara)
        # the other way: Associate 018 (Monday 09:00 - 18:00, off Tuesday) called in at 06:00 has exactly 12 hours
        self.days.add_activity("AE/AR B2B", TUE, "Associate 018", "Called in", "06:00", "15:00", self.sara)
        with self.assertRaises(ValueError) as refused:
            self.days.add_activity("AE/AR B2B", MON, "Associate 018", "Overtime", "18:00", "19:00", self.sara)
        self.assertIn("would rest 11 hours before the next shift", str(refused.exception))

    def test_board_offers_keep_the_rest_gap_across_the_week(self):
        page = self.days.page("AE/AR B2B", SUN)
        cover = self.days.cover_offers(page, 9 * 60)
        offered = [o["name"] for o in cover["dayoff"]]
        self.assertTrue(offered)
        # Associate 001 (Saturday to 01:00) and 012 (to 02:00) cannot start by 09:00 with 12 hours' rest
        self.assertFalse({"Associate 001", "Associate 012"} & set(offered))
        cover = self.days.cover_offers(page, 14 * 60)
        self.assertNotIn("Associate 010", [o["name"] for o in cover["overtime"]])  # 14:00 leaves 11 hours

    def test_cancelling_a_call_in_leaves_nothing_behind(self):
        called = self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Called in", "12:00", "21:00", self.sara)
        training = self.days.add_activity("AE/AR B2B", WED, "Associate 002", "Training", "15:00", "16:00", self.sara)
        self.days.set_status("AE/AR B2B", WED, "Associate 002", "Late", self.sara, end="12:30")
        self.days.move_break("AE/AR B2B", WED, "Associate 002", 1, "17:00", self.sara)
        with self.assertRaises(ValueError) as refused:  # its records would be left on a day off
            self.days.cancel_activity("AE/AR B2B", WED, called, self.sara)
        self.assertEqual(str(refused.exception), "Associate 002 still has a training 15:00 to 16:00, the status Late "
                         "and a moved Lunch on this shift: cancel or set those back first.")
        self.days.cancel_activity("AE/AR B2B", WED, training, self.sara)
        self.days.set_status("AE/AR B2B", WED, "Associate 002", "Present", self.sara)
        self.days.move_break("AE/AR B2B", WED, "Associate 002", 1, None, self.sara)
        self.days.cancel_activity("AE/AR B2B", WED, called, self.sara)
        page = self.days.page("AE/AR B2B", WED)
        self.assertEqual(page["log"][-1]["what"], "Call-in cancelled: back to the day off (was 12:00 - 21:00)")
        self.assertNotIn("Associate 002", [l["name"] for l in page["view"]["lanes"]])  # off again: no lane
        self.assertEqual(page["view"]["tiles"]["called_in"], 0)

    def test_unknown_status_or_person_refused(self):
        with self.assertRaises(ValueError):
            self.days.set_status("AE/AR B2B", WED, "Associate 001", "Holiday", self.sara)
        with self.assertRaises(ValueError):
            self.days.set_status("AE/AR B2B", WED, "Nobody", "Sick", self.sara)
        with self.assertRaises(ValueError):  # Associate 001 is off on Saturday
            self.days.set_status("AE/AR B2B", SAT, "Associate 001", "Sick", self.sara)
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"], [])

    def test_break_moved_then_back_to_plan(self):
        # Wednesday 12:00 - 21:00: Break 1 13:45, Lunch 16:30, Break 2 19:30
        self.days.move_break("AE/AR B2B", WED, "Associate 001", 1, "17:05", self.sara)
        page = self.days.page("AE/AR B2B", WED)
        lunch = self.lane(page, "Associate 001")["segments"][0]["breaks"][1]
        self.assertEqual((lunch["kind"], lunch["start"], lunch["planned_start"], lunch["moved"]),
                         ("Lunch", 17 * 60 + 5, 16 * 60 + 30, True))
        self.assertEqual(page["log"][-1]["what"], "Lunch moved 16:30 to 17:05")
        self.days.move_break("AE/AR B2B", WED, "Associate 001", 1, None, self.sara)
        page = self.days.page("AE/AR B2B", WED)
        lunch = self.lane(page, "Associate 001")["segments"][0]["breaks"][1]
        self.assertEqual((lunch["start"], lunch["moved"]), (16 * 60 + 30, False))
        self.assertEqual(page["log"][-1]["what"], "Lunch back to plan (16:30)")

    def test_refused_break_is_not_kept(self):
        for at in ("17:07", "20:50", "19:30", "bad"):  # not a 5-minute step, past 21:00, on Break 2, not a time
            with self.assertRaises((BreakRefused, ValueError)):
                self.days.move_break("AE/AR B2B", WED, "Associate 001", 1, at, self.sara)
        self.assertFalse(any(b["moved"] for b in self.lane(self.days.page("AE/AR B2B", WED),
                                                           "Associate 001")["segments"][0]["breaks"]))
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"], [])

    def test_a_move_that_no_longer_matches_the_plan_is_reported_not_applied(self):
        self.store.set_actual_break(program="AE/AR B2B", shift_date=WED.isoformat(), associate="Associate 001",
                                    idx=1, kind="Break 9", start=16 * 60, user_id=self.sara)
        page = self.days.page("AE/AR B2B", WED)
        self.assertEqual(page["stale"], 1)
        self.assertFalse(self.lane(page, "Associate 001")["segments"][0]["breaks"][1]["moved"])

    def test_overnight_status_belongs_to_the_shift_date(self):
        tue = date(2026, 10, 13)
        person = next(a for a in self.book.view(self.tool["id"])["week"]["associates"]
                      if (shift_span(a["days"][2]) or (0, 0))[1] > 1440)  # Tuesday night into Wednesday
        self.days.set_status("AE/AR B2B", tue, person["name"], "Sick", self.sara)
        carry = [s for s in self.lane(self.days.page("AE/AR B2B", WED), person["name"])["segments"] if s["offset"] == -1]
        self.assertEqual(carry[0]["status"], "Sick")

    def test_day_records_deleted_after_13_months(self):
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Sick", self.sara)
        self.days.move_break("AE/AR B2B", WED, "Associate 001", 1, "17:05", self.sara)
        self.assertEqual(self.days.cleanup(now=time.time() + 300 * 86400), 0)
        # the status, the break move and their two log lines
        self.assertEqual(self.days.cleanup(now=time.mktime((2027, 11, 20, 0, 0, 0, 0, 0, -1))), 4)
        self.assertEqual(self.store.list_attendance("AE/AR B2B", [WED.isoformat()]), [])


class TheAddDialog(unittest.TestCase):
    """Phase R task 4: RTA adds anything in an interval: a break or lunch, an aux, unplanned leave, sick,
    late, left early, overtime or VTO (any stretch of the shift), with the same checks as before and a
    preview that records nothing (owner, 2026-10-08)."""

    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        store.add_run("aaaaaaaaaaaa", cls.sara, "AR_week.xlsx", "QUICK", "DONE", program="AE/AR B2B",
                      week_start="2026-10-11")
        ScheduleBook(store, cls.base, REPO).ensure(store.get_run("aaaaaaaaaaaa"), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.days = DayBook(self.store, ScheduleBook(self.store, self.data, REPO))

    def seg(self, name="Associate 001", on=WED):
        page = self.days.page("AE/AR B2B", on)
        return next(s for l in page["view"]["lanes"] if l["name"] == name for s in l["segments"] if s["offset"] == 0)

    def free_start(self, minutes, name="Associate 001"):
        """The first time in the shift, an hour in, with no break or activity for ``minutes``."""
        seg = self.seg(name)
        busy = [(b["start"], b["start"] + b["minutes"]) for b in seg["breaks"]]
        busy += [(a["start"], a["end"]) for a in seg["activities"]]
        for t in range(seg["start"] + 60, seg["end"] - minutes, 5):
            if all(t + minutes <= lo or hi <= t for lo, hi in busy):
                return t
        raise AssertionError("no free time")

    def on_floor(self, t):
        cell = next(c for c in self.days.page("AE/AR B2B", WED)["view"]["cells"] if c["t"] <= t < c["t"] + 30)
        return cell["slots"][(t - cell["t"]) // 5]

    def test_add_break_in_an_interval_takes_them_off_the_floor(self):
        t = self.free_start(15)
        before = self.on_floor(t)
        got = self.days.add_item("AE/AR B2B", WED, "Associate 001", "Break", hm(t), 15, self.sara)
        self.assertEqual(got["text"], f"Break {hm(t)} to {hm(t + 15)}")
        self.assertEqual(self.on_floor(t), before - 1)
        rows = board(self.days.page("AE/AR B2B", WED)["view"])
        chip = next(c for r in rows for c in r["short"] if c["name"] == "Associate 001" and c.get("added"))
        self.assertEqual((chip["kind"], chip["start"]), ("Break", hm(t)))

    def test_add_unplanned_leave_from_the_dialog(self):
        tiles = self.days.page("AE/AR B2B", WED)["view"]["tiles"]
        got = self.days.add_item("AE/AR B2B", WED, "Associate 001", "Unplanned leave", "", 0, self.sara)
        self.assertEqual(got["text"], "Unplanned leave")
        self.assertEqual(self.days.page("AE/AR B2B", WED)["view"]["tiles"]["present"], tiles["present"] - 1)

    def test_add_late_outside_shift_is_refused(self):
        with self.assertRaises(ValueError) as said:  # Wednesday is 12:00 - 21:00
            self.days.add_item("AE/AR B2B", WED, "Associate 001", "Late", "22:00", 0, self.sara)
        self.assertIn("outside Associate 001's shift (12:00 to 21:00)", str(said.exception))
        with self.assertRaises(ValueError) as said:
            self.days.add_item("AE/AR B2B", WED, "Associate 001", "Late", "", 0, self.sara)
        self.assertIn("Give the time they arrived.", str(said.exception))
        self.assertEqual(self.store.list_attendance("AE/AR B2B", [WED.isoformat()]), [])

    def test_add_overtime_with_a_chosen_length(self):
        self.days.add_item("AE/AR B2B", WED, "Associate 001", "Overtime", "21:00", 45, self.sara)
        acts = self.store.list_activities("AE/AR B2B", [WED.isoformat()])
        self.assertEqual([(a["kind"], a["start"], a["end_min"]) for a in acts], [("Overtime", 1260, 1305)])

    def test_day_off_cancelled_on_typed_times(self):
        # owner, 2026-10-09: "cancel day off I'm not able to choose the shift start end manually or it's not
        # existing in RTA view": "+ Add" calls in someone off that day, on typed times (Associate 002 is off Wed)
        got = self.days.add_item("AE/AR B2B", WED, "Associate 002", "Day off cancelled", "10:30", 0, self.sara,
                                 end="16:30")
        self.assertEqual(got["text"], "Day off cancelled: called in 10:30 - 16:30")
        seg = self.seg("Associate 002")
        self.assertEqual((seg["label"], seg["start"], seg["end"], seg["breaks"]),
                         ("10:30 - 16:30 (called in)", 630, 990, []))  # nobody works it: breaks are added by hand

    def test_day_off_cancelled_on_a_library_shift_takes_its_breaks(self):
        self.days.add_item("AE/AR B2B", WED, "Associate 002", "Day off cancelled", "12:00", 0, self.sara, end="21:00")
        self.assertEqual([b["kind"] for b in self.seg("Associate 002")["breaks"]], ["Break 1", "Lunch", "Break 2"])

    def test_day_off_cancelled_overnight_on_typed_times(self):
        self.days.add_item("AE/AR B2B", WED, "Associate 002", "Day off cancelled", "22:00", 0, self.sara, end="04:00")
        seg = self.seg("Associate 002")
        self.assertEqual((seg["start"], seg["end"]), (22 * 60, 28 * 60))

    def test_typed_call_in_times_are_checked(self):
        for start, end, said in (("10:10", "16:30", "15-minute steps"), ("10:00", "10:45", "at least 1 hour"),
                                 ("06:00", "22:00", "9 hours, the longest shift in the Shift Library"),
                                 ("10", "16:30", "like 08:00")):
            with self.assertRaises(ValueError) as refused:
                self.days.add_item("AE/AR B2B", WED, "Associate 002", "Day off cancelled", start, 0, self.sara,
                                   end=end)
            self.assertIn(said, str(refused.exception))
        with self.assertRaises(ValueError) as refused:  # works that day: overtime instead
            self.days.add_item("AE/AR B2B", WED, "Associate 001", "Day off cancelled", "10:00", 0, self.sara,
                               end="16:00")
        self.assertIn("not off", str(refused.exception))
        self.assertIsNone(self.days.called_in_on("AE/AR B2B", WED, "Associate 002"))

    def test_day_off_cancelled_preview_records_nothing(self):
        found = self.days.preview_item("AE/AR B2B", WED, "Associate 002", "Day off cancelled", "12:00", 0,
                                       end="21:00")
        self.assertIn("Day off cancelled: called in 12:00 - 21:00", found["text"])
        self.assertIsNone(self.days.called_in_on("AE/AR B2B", WED, "Associate 002"))

    def test_people_off_today_are_listed_for_the_dialog(self):
        off = self.days.off_today("AE/AR B2B", WED)
        self.assertIn("Associate 002", [name for name, _ in off])
        self.assertNotIn("Associate 001", [name for name, _ in off])
        self.days.add_item("AE/AR B2B", WED, "Associate 002", "Day off cancelled", "12:00", 0, self.sara, end="21:00")
        self.assertNotIn("Associate 002", [name for name, _ in self.days.off_today("AE/AR B2B", WED)])

    def test_vto_any_stretch_inside_the_shift(self):
        got = self.days.add_item("AE/AR B2B", WED, "Associate 001", "VTO", "14:00", 60, self.sara)
        self.assertEqual(got["text"], "VTO 14:00 to 15:00")
        with self.assertRaises(ValueError):  # before the shift
            self.days.add_item("AE/AR B2B", WED, "Associate 001", "VTO", "06:00", 60, self.sara)

    def test_added_break_cannot_overlap_a_planned_break(self):
        planned = self.seg()["breaks"][0]
        with self.assertRaises(ValueError) as said:
            self.days.add_item("AE/AR B2B", WED, "Associate 001", "Lunch", hm(planned["start"]), 30, self.sara)
        self.assertIn(f"already on {planned['kind'].lower()}", str(said.exception))

    def test_lengths_go_in_five_minute_steps(self):
        with self.assertRaises(ValueError) as said:
            self.days.add_item("AE/AR B2B", WED, "Associate 001", "Coaching", "15:00", 0, self.sara)
        self.assertIn("Pick a length", str(said.exception))

    def test_preview_changes_nothing_and_says_the_effect(self):
        t = self.free_start(15)
        found = self.days.preview_item("AE/AR B2B", WED, "Associate 001", "Break", hm(t), 15, False, "interval")
        self.assertIn(f"{hm(t)} to {hm(t + 15)}", found["text"])
        self.assertIn(found["level"], ("ok", "warn", "bad"))
        self.assertEqual(self.store.list_activities("AE/AR B2B", [WED.isoformat()]), [])
        with self.assertRaises(ValueError):  # refused the same way the real thing is
            self.days.preview_item("AE/AR B2B", WED, "Associate 001", "Late", "22:00", 0, False, "interval")



class TheAuxDetails(unittest.TestCase):
    """Phase T: every aux says who it is with and why, kept with the record and in the day log (owner, 2026-10-09:
    "in case of any aux being placed like meeting coaching etc we need to specify with who and why in a comment
    while reserving and to reflect in the export report")."""

    setUpClass = classmethod(TheAddDialog.setUpClass.__func__)
    tearDownClass = classmethod(TheAddDialog.tearDownClass.__func__)
    setUp = TheAddDialog.setUp
    seg = TheAddDialog.seg
    free_start = TheAddDialog.free_start

    def test_both_answers_are_required_for_every_aux(self):
        for kind in ("Coaching", "Meeting", "Training", "System issue"):
            with self.assertRaises(ValueError) as said:
                aux_details(kind, "", "Monthly review")
            self.assertEqual(str(said.exception), f"Say who the {kind.lower()} is with.")
            with self.assertRaises(ValueError) as said:
                aux_details(kind, "Sara", "  ")
            self.assertEqual(str(said.exception), f"Say why: a short reason for the {kind.lower()}.")
        self.assertEqual(aux_details("Coaching", "  Sara   Ali ", " Monthly\n review "), ("Sara Ali", "Monthly review"))
        self.assertEqual(aux_details("Overtime", "", ""), ("", ""))  # not an aux: nothing is asked

    def test_long_answers_are_refused_not_cut(self):
        with self.assertRaises(ValueError) as said:
            aux_details("Meeting", "x" * 81, "Why")
        self.assertIn("80 characters", str(said.exception))
        with self.assertRaises(ValueError) as said:
            aux_details("Meeting", "Sara", "y" * 201)
        self.assertIn("200 characters", str(said.exception))

    def test_an_activity_keeps_who_and_why_and_logs_them(self):
        t = self.free_start(30)
        self.days.add_item("AE/AR B2B", WED, "Associate 001", "Coaching", hm(t), 30, self.sara, billable=True,
                           with_whom="Sara (team leader)", why="Monthly quality review")
        row = next(r for r in self.store.list_activities("AE/AR B2B", [WED.isoformat()]) if r["kind"] == "Coaching")
        self.assertEqual((row["with_whom"], row["why"]), ("Sara (team leader)", "Monthly quality review"))
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"][-1]["what"],
                         f"Coaching {hm(t)} to {hm(t + 30)} (billable), with Sara (team leader): Monthly quality review")
        act = next(a for a in self.seg()["activities"] if a["kind"] == "Coaching")
        self.assertEqual((act["with_whom"], act["why"]), ("Sara (team leader)", "Monthly quality review"))

    def test_an_aux_status_keeps_who_and_why_and_logs_them(self):
        self.days.set_status("AE/AR B2B", WED, "Associate 001", "Training", self.sara, start="13:00", end="14:00",
                             with_whom="IT trainer", why="New CRM release")
        row = self.store.list_attendance("AE/AR B2B", [WED.isoformat()])[0]
        self.assertEqual((row["with_whom"], row["why"]), ("IT trainer", "New CRM release"))
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"][-1]["what"],
                         "Training (non-billable) 13:00 to 14:00, with IT trainer: New CRM release")
        seg = self.seg()
        self.assertEqual((seg["with_whom"], seg["why"]), ("IT trainer", "New CRM release"))

    def test_a_booked_session_keeps_who_and_why_for_everyone(self):
        names = ["Associate 001", "Associate 012"]
        slots = self.days.meeting_slots("AE/AR B2B", WED, names, 30, "13:00", "17:00")
        self.days.book_session("AE/AR B2B", WED, names, slots[0]["start"], 30, "Meeting", self.sara,
                               with_whom="Ops manager", why="Process update")
        rows = [r for r in self.store.list_activities("AE/AR B2B", [WED.isoformat()]) if r["kind"] == "Meeting"]
        self.assertEqual({(r["associate"], r["with_whom"], r["why"]) for r in rows},
                         {(n, "Ops manager", "Process update") for n in names})

    def test_the_department_is_kept_and_logged(self):
        # Phase U (owner, 2026-10-09): who an aux is with is a department and a person from the program's lists
        t = self.free_start(30)
        self.days.add_item("AE/AR B2B", WED, "Associate 001", "Coaching", hm(t), 30, self.sara,
                           with_dept="Quality", with_whom="Lina", why="Monthly quality review")
        row = next(r for r in self.store.list_activities("AE/AR B2B", [WED.isoformat()]) if r["kind"] == "Coaching")
        self.assertEqual((row["with_dept"], row["with_whom"]), ("Quality", "Lina"))
        self.assertEqual(self.days.page("AE/AR B2B", WED)["log"][-1]["what"],
                         f"Coaching {hm(t)} to {hm(t + 30)} (non-billable), with Lina (Quality): Monthly quality review")
        self.assertEqual(next(a for a in self.seg()["activities"] if a["kind"] == "Coaching")["with_dept"], "Quality")
        self.days.set_status("AE/AR B2B", WED, "Associate 012", "Meeting", self.sara, start="13:00", end="13:30",
                             with_dept="Workforce", with_whom="Sara", why="Schedule review")
        row = next(r for r in self.store.list_attendance("AE/AR B2B", [WED.isoformat()]) if r["associate"] == "Associate 012")
        self.assertEqual((row["with_dept"], row["with_whom"]), ("Workforce", "Sara"))
        self.assertEqual(self.seg("Associate 012")["with_dept"], "Workforce")
        names = ["Associate 013", "Associate 027"]
        slots = self.days.meeting_slots("AE/AR B2B", WED, names, 30, "13:00", "18:00")
        self.days.book_session("AE/AR B2B", WED, names, slots[0]["start"], 30, "Training", self.sara,
                               with_dept="Training", with_whom="IT trainer", why="New CRM release")
        self.assertEqual({r["with_dept"] for r in self.store.list_activities("AE/AR B2B", [WED.isoformat()])
                          if r["kind"] == "Training"}, {"Training"})

    def test_a_database_from_before_gets_the_columns(self):
        import sqlite3
        old = self.data / "old.db"
        with sqlite3.connect(old) as db:  # the two tables as they were before Phase T
            db.execute("create table attendance (program text not null, shift_date text not null, associate text "
                       "not null, status text not null, from_min integer, to_min integer, billable integer not null "
                       "default 0, user_id integer not null, at real not null, primary key (program, shift_date, "
                       "associate))")
            db.execute("create table activities (id integer primary key autoincrement, program text not null, "
                       "shift_date text not null, associate text not null, kind text not null, start integer not "
                       "null, end_min integer not null, billable integer not null default 0, note text not null "
                       "default '', user_id integer not null, at real not null)")
            db.execute("insert into activities (program, shift_date, associate, kind, start, end_min, user_id, at) "
                       "values ('P', '2026-10-14', 'A', 'Meeting', 600, 630, 1, 0)")
        store = Store(old)
        with sqlite3.connect(old) as db:
            for table in ("attendance", "activities"):
                self.assertTrue({"with_whom", "why"} <= {r[1] for r in db.execute(f"pragma table_info({table})")})
        self.assertEqual([(r["with_whom"], r["why"]) for r in store.list_activities("P", ["2026-10-14"])], [("", "")])


class TheStartDay(unittest.TestCase):
    """Phase P: a schedule's first date is chosen when it is uploaded (Sunday or Monday). A date reads the
    column of its weekday from the run whose seven days hold it; last night's overnight shifts come from the
    day before, in the same run or the run before (the input's carry-in tab when there is none)."""

    @classmethod
    def setUpClass(cls):
        cls.base = Path(tempfile.mkdtemp())
        store = Store(cls.base / "scheduler.db")
        cls.sara = store.add_user("sara", "Sara", "Sara-pass-1", must_change=False)
        book = ScheduleBook(store, cls.base, REPO)
        for run_id, program, start in (("cccccccccccc", "MON", "2026-10-12"), ("dddddddddddd", "MON", "2026-10-19"),
                                       ("eeeeeeeeeeee", "MIX", "2026-10-11"), ("ffffffffffff", "MIX", "2026-10-12")):
            store.add_run(run_id, cls.sara, "AR_week.xlsx", "QUICK", "DONE", program=program, week_start=start)
            book.ensure(store.get_run(run_id), INPUT, AFTER, None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.base, True)

    def setUp(self):
        self.data = Path(tempfile.mkdtemp()) / "data"
        self.addCleanup(shutil.rmtree, self.data.parent, True)
        shutil.copytree(self.base, self.data)
        self.store = Store(self.data / "scheduler.db")
        self.book = ScheduleBook(self.store, self.data, REPO)
        self.days = DayBook(self.store, self.book)

    def run_of(self, program, on):
        row, _ = self.days.version(program, on)
        return row["run_id"] if row else None

    def carried(self, program, on, name):
        lane = next((l for l in self.days.page(program, on)["view"]["lanes"] if l["name"] == name), None)
        return [(s["start"], s["end"]) for s in (lane or {}).get("segments", []) if s["offset"] == -1]

    def test_a_monday_start_covers_monday_to_sunday(self):
        for d in range(12, 19):
            self.assertEqual(self.run_of("MON", date(2026, 10, d)), "cccccccccccc", d)
        self.assertIsNone(self.run_of("MON", date(2026, 10, 11)))
        self.assertEqual(self.run_of("MON", date(2026, 10, 19)), "dddddddddddd")
        self.assertIsNone(self.run_of("MON", date(2026, 10, 26)))
        page = self.days.page("MON", date(2026, 10, 18))
        self.assertEqual((page["day"], page["week_start"]), (0, "2026-10-12"))  # Sunday's column, the run's start

    def test_the_last_day_takes_saturday_night_from_the_same_run(self):
        # Associate 004: Saturday 20:00 - 05:00 runs into Sunday 18 October, the run's last day
        self.assertEqual(self.carried("MON", date(2026, 10, 18), "Associate 004"), [(20 * 60 - 1440, 5 * 60)])

    def test_the_first_day_takes_last_night_from_the_run_before(self):
        # Monday 19: Sunday 18 (the run before, Associate 004 23:00 - 08:00); not the input's carry-in tab
        self.assertEqual(self.carried("MON", date(2026, 10, 19), "Associate 004"), [(23 * 60 - 1440, 8 * 60)])
        self.assertEqual(self.carried("MON", date(2026, 10, 19), "Associate 001"), [])

    def test_the_first_day_without_a_run_before_uses_the_input_tab(self):
        # Monday 12 has no run for Sunday 11: the tab's Associate 001 16:00 - 01:00; never this run's own Sunday
        self.assertEqual(self.carried("MON", date(2026, 10, 12), "Associate 001"), [(16 * 60 - 1440, 60)])
        self.assertEqual(self.carried("MON", date(2026, 10, 12), "Associate 004"), [])

    def test_overlapping_runs_latest_start_unless_in_use(self):
        self.assertEqual(self.run_of("MIX", date(2026, 10, 11)), "eeeeeeeeeeee")
        self.assertEqual(self.run_of("MIX", date(2026, 10, 14)), "ffffffffffff")  # the later start counts
        self.book.set_in_use(self.book.versions("eeeeeeeeeeee")[0]["id"], self.sara)
        self.assertEqual(self.run_of("MIX", date(2026, 10, 14)), "eeeeeeeeeeee")  # unless the other is in use

    def test_tomorrow_unchecked_on_the_last_day(self):
        self.assertTrue(self.days.next_week_unknown("MON", date(2026, 10, 25)))
        self.assertFalse(self.days.next_week_unknown("MON", date(2026, 10, 24)))
        self.assertFalse(self.days.next_week_unknown("MON", date(2026, 10, 18)))  # Monday 19 is covered

    def test_breaks_export_follows_a_monday_start(self):
        import io
        from openpyxl import load_workbook
        from webapp.exports import build
        data, *_ = build(self.store, self.days, date(2026, 10, 11), date(2026, 10, 12), ["breaks"], program="MON",
                         by="Omar")
        rows = list(load_workbook(io.BytesIO(data))["Breaks planned vs taken"].iter_rows(values_only=True))
        self.assertEqual({r[0] for r in rows[1:]}, {"2026-10-12"})

    def test_exports_look_each_day_up_once(self):
        import io
        from webapp.exports import build
        for name in ("Associate 001", "Associate 006", "Associate 008"):  # all work on Wednesday
            self.days.set_status("MON", date(2026, 10, 14), name, "Sick", self.sara)
        asked = []
        version = self.days.version
        self.days.version = lambda program, on: (asked.append((program, on)), version(program, on))[1]
        build(self.store, self.days, date(2026, 10, 14), date(2026, 10, 14), ["attendance"], program="MON", by="Omar")
        self.assertEqual(asked, [("MON", date(2026, 10, 14))])  # three records, one lookup


if __name__ == "__main__":
    unittest.main()
