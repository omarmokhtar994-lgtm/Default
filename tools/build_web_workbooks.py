#!/usr/bin/env python3
# © 2026 Omar Mokhtar. All rights reserved.
"""Build the two clean input workbooks the website offers on its home page.

    python3 tools/build_web_workbooks.py [OUT_DIR]      (default webapp/workbooks)

Example: the synthetic SYNTH_M2 workbook (made-up "Synth Agent" names, a
roster that schedules) rebuilt by tools/build_input_template.py, whose own
check proves the engine reads it exactly as it reads the source; nothing is
written if it would not.

Blank: the example with people, demand, shrinkage, preferences, fixed
requests, last week's shifts and languages cleared, the program name and
headcount emptied, and the 30-minute demand tabs selected (the interval the
team's programs use). Every tab, dropdown, the Start Here checklist and the
Shift Library stay.

Both then get the channel tabs (Phase Z, webapp/channel_template.py) at their
own interval: empty on the Blank; on the Example, Phone is 60% of FT Wise
rounded up and Chat the rest, so the website's channel check matches FT Wise.
The engine never reads these tabs: nothing is written unless it reads exactly
the same from each workbook with the tabs as without them.
"""
from __future__ import annotations

import math
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_input_template as B  # noqa: E402

sys.path.insert(0, str(ROOT))
from webapp.channel_template import add_channel_tabs  # noqa: E402
from webapp.day import read_inputs  # noqa: E402

SOURCE = "fixtures/synthetic_suite/SYNTH_M2_MULTI_START_WITH_BREAKS.xlsx"
EXAMPLE_NAME = "Scheduler_Input_Example.xlsx"
BLANK_NAME = "Scheduler_Input_Blank.xlsx"
# sheet: (first data row, first data column); everything from there on is cleared
CLEARED = {
    "Schedule": (3, 1), "Previous week scheduled": (3, 1), "Preference": (3, 1), "Fixed Request": (3, 1),
    "Language Setup": (3, 1), "Coverage Split": (6, 1),
    "FT Wise 15 Min": (3, 2), "FT Wise 30 Min": (3, 2), "FT Wise 60 Min": (3, 2),
    "Shrinkage 15 Min": (3, 2), "Shrinkage 30 Min": (3, 2), "Shrinkage 60 Min": (3, 2),
}
# the header the clearing relies on: (row, column, text)
HEADERS = {
    "Schedule": (2, 1, "Slot"), "Previous week scheduled": (2, 1, "Slot"), "Preference": (2, 1, "Associate Name"),
    "Fixed Request": (2, 1, "Active?"), "Language Setup": (2, 1, "Language"), "Coverage Split": (4, 1, "Coverage Group"),
    **{f"{kind} {n} Min": (2, 1, "Interval") for kind in ("FT Wise", "Shrinkage") for n in (15, 30, 60)},
}
AFTER_TAB = "Previous week scheduled"  # the channel tabs follow the last weekly input tab
INPUT_TAB_COLOUR = "ED7D31"
PHONE_SHARE = 0.6
BLANK_SETTINGS = {"Program Name": None, "Count of Associates": None, "Interval Minutes": "30",
                  "Requirements Source": "FT Wise 30 Min", "Shrinkage Source": "Shrinkage 30 Min"}


# Phase AO (owner, 2026-10-21: an 11-hour shift gets four breaks, in the input sheet the website offers): an Instructions
# row the engine already reads (audit F-20). A shift of 11 hours or more gets these breaks in this order (30 minutes or
# more is a lunch); a shorter shift keeps the Short Break and Lunch counts. The row sits in the Breaks group.
LONG_BREAKS_ROW = "Break Set For Shifts Of 11 Hours Or More"
LONG_BREAKS_VALUE = "15, 30, 15, 15"
LONG_BREAKS_HELP = ("Shifts of 11 hours or more get these breaks in this order, in minutes (30 or more is a lunch): "
                    "15, 30, 15, 15 = two short breaks around a lunch, then one more. Shorter shifts keep the counts above.")


def offer_long_shift_breaks(source: Path, kept: Path) -> Path:
    """``source`` with the row added (so the engine reads the built template exactly as it reads its source), and the
    template builder told to lay the row out in the Breaks group. Written to ``kept``."""
    wb = load_workbook(source)
    ws = wb["Instructions"]
    if any(str(r[1] or "").strip() == LONG_BREAKS_ROW for r in ws.iter_rows(values_only=True)):
        raise SystemExit(f"{source.name} already has {LONG_BREAKS_ROW!r}; nothing written")
    row = ws.max_row + 1
    ws.cell(row, 2, LONG_BREAKS_ROW)
    ws.cell(row, 3, LONG_BREAKS_VALUE)
    wb.save(kept)
    if not any(key == LONG_BREAKS_ROW for _name, items in B.SETUP_LAYOUT for key, _c in items):  # once per process
        B.SETUP_LAYOUT = [(name, items + [(LONG_BREAKS_ROW, None)]) if name == "Breaks" else (name, items)
                          for name, items in B.SETUP_LAYOUT]
    B.HELP[LONG_BREAKS_ROW] = LONG_BREAKS_HELP
    B.NUMBER_RULES[LONG_BREAKS_ROW] = ("textLength", None, 40, None)
    return kept


def blank_from(example: Path, blank: Path) -> None:
    wb = load_workbook(example)
    for sheet, (row, col, text) in HEADERS.items():
        found = wb[sheet].cell(row, col).value
        if found != text:
            raise SystemExit(f"{sheet}: expected the header {text!r} at row {row}, found {found!r}; nothing written")
    for sheet, (first_row, first_col) in CLEARED.items():
        for cells in wb[sheet].iter_rows(min_row=first_row, min_col=first_col):
            for cell in cells:
                if cell.value is not None and not str(cell.value).startswith("="):
                    cell.value = None
    settings = wb["Instructions"]
    seen = set()
    for cells in settings.iter_rows():
        name = cells[1].value
        if name in BLANK_SETTINGS:
            old = cells[2].value
            new = BLANK_SETTINGS[name]
            cells[2].value = int(new) if isinstance(old, int) and new is not None else new
            seen.add(name)
    missing = set(BLANK_SETTINGS) - seen
    if missing:
        raise SystemExit(f"Instructions rows not found: {sorted(missing)}; nothing written")
    wb.save(blank)


def _ft_wise(wb, step: int) -> dict:
    """FT Wise at ``step`` minutes as {(day, minute): people}; empty cells are 0."""
    ws = wb[f"FT Wise {step} Min"]
    head = next(r for r in range(1, 11) if str(ws.cell(r, 1).value or "").strip().lower() == "interval")
    days = {c - 2: c for c in range(2, 9)}
    grid = {}
    for r in range(head + 1, ws.max_row + 1):
        text = str(ws.cell(r, 1).value or "").strip()
        if len(text) < 5 or text[2] != ":":
            continue
        t = int(text[:2]) * 60 + int(text[3:5])
        for d, c in days.items():
            grid[(d, t)] = float(ws.cell(r, c).value or 0)
    return grid


def engine_view(path: Path) -> dict:
    """What the engine reads from a workbook (build_input_template's check). A workbook it cannot parse (the Blank
    has no people) gives the refusal and the requirement tab it picks, compared the same way."""
    try:
        return B.engine_reading(path)
    except ValueError as exc:
        import l632_universal_scheduler as E  # on the path once B.engine_reading has run
        wb = load_workbook(path)
        im = E._instruction_map(E._sheet_by_alias(wb, ["Engine Defaults", "Engine Default", "Scheduler Defaults"]))
        im.update(E._instruction_map(E._sheet_by_alias(wb, ["Instructions"])))
        return {"parse refused": str(exc), "requirement tab": E._discover_requirement_sheet(wb, im).title}


def with_channel_tabs(path: Path, filled: bool) -> str:
    """Add the channel tabs to a built workbook, after its weekly input tabs. Refused (the workbook is put back as
    it was) when the engine would read anything differently."""
    inputs = read_inputs(path)
    step = inputs["interval"]
    before = engine_view(path)
    keep = path.with_suffix(".before-channels.xlsx")
    shutil.copyfile(path, keep)
    wb = load_workbook(path)
    ft = _ft_wise(wb, step) if filled else {}

    def need(letter: str, d: int, t: int) -> float:
        people = ft.get((d, t), 0)
        phone = math.ceil(PHONE_SHARE * people)
        return phone if letter == "P" else people - phone
    count = len(wb.sheetnames)
    add_channel_tabs(wb, step, [lang["name"] for lang in inputs["languages"]], need if filled else None)
    added = wb.worksheets[count:]
    at = wb.sheetnames.index(AFTER_TAB) + 1
    for i, ws in enumerate(added):
        wb.move_sheet(ws, offset=at + i - wb.sheetnames.index(ws.title))
        ws.sheet_properties.tabColor = INPUT_TAB_COLOUR
    wb.save(path)
    after = engine_view(path)
    changed = sorted(k for k in set(before) | set(after) if before.get(k) != after.get(k))
    if changed:
        shutil.move(keep, path)
        raise SystemExit(f"{path.name}: the channel tabs would change what the engine reads ({', '.join(changed)}); "
                         "the workbook is left without them")
    keep.unlink()
    return f"{path.name}: channel tabs at {step} minutes{' filled from FT Wise' if filled else ''}; the engine reads " \
           f"the same ({len(before)} fields compared)"


def main(argv: list) -> int:
    out = Path(argv[1]) if len(argv) > 1 else ROOT / "webapp" / "workbooks"
    out.mkdir(parents=True, exist_ok=True)
    example = out / EXAMPLE_NAME
    with tempfile.TemporaryDirectory() as tmp:
        source = offer_long_shift_breaks(ROOT / SOURCE, Path(tmp) / "source_with_long_breaks.xlsx")
        if B.main(["build", str(source), str(example)]) != 0:
            return 2
    blank_from(example, out / BLANK_NAME)
    print(f"{BLANK_NAME}: every tab and dropdown, no people or demand, 30-minute tabs selected")
    print(with_channel_tabs(example, filled=True))
    print(with_channel_tabs(out / BLANK_NAME, filled=False))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
