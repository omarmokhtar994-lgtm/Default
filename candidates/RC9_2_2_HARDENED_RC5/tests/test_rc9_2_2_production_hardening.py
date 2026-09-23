#!/usr/bin/env python3
"""Fast regression guards for the RC9.2.2 consolidated hardening pass."""
from __future__ import annotations

import importlib.util
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree
from pathlib import Path
from types import SimpleNamespace

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ENGINE_PATH = ROOT / "engine" / "_tools" / "l632_universal_scheduler.py"
VALIDATOR_PATH = ROOT / "engine" / "tools" / "independent_validator.py"
RUNNER_PATH = ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py"
sys.path.insert(0, str(ENGINE_PATH.parent))
import l632_universal_scheduler as E  # noqa: E402


def load_validator():
    spec = importlib.util.spec_from_file_location("rc922_validator", VALIDATOR_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def pair(profile: str, values: dict):
    skeleton = E.SkeletonSolution(profile, "FEASIBLE", 0.0, 0.0, [], [], {})
    solution = E.BreakSolution(
        profile=f"{profile}_breaks", skeleton_profile=profile, cp_status="FEASIBLE",
        elapsed_sec=0.0, objective=0.0, pattern_width=115, exception_mode=False,
        selected_pattern={(0, 0): 1}, no_break_cells=set(), patterns=[],
        diagnostics={"objective_mode": "target_priority"}, metrics=values,
    )
    return skeleton, solution


class ProtectedTierTradeoffGuard(unittest.TestCase):
    def test_three_target_hits_cannot_buy_nine_protected_hits(self):
        parsed = SimpleNamespace(target_ratio=0.90, floor_ratio=0.75)
        common = {
            "active_intervals": 252, "after_100": 140, "hard_floor_gap_count": 0,
            "week_boundary_hard_failure_count": 0, "week_boundary_after_target": 20,
            "week_boundary_after_floor": 24, "severe_floor_gap_count": 0,
            "max_consecutive_floor_gaps": 0, "floor_deficit_max": 0.01,
            "floor_deficit_sum": 0.10, "before_target": 160,
            "after_avoidable_overage_fte_sum": 20.0, "after_extreme_overage_count": 0,
            "after_severe_overage_count": 0, "target_deficit_sum": 0.0,
        }
        safe = pair("safe", {**common, "after_target": 150, "after_90": 150,
                              "after_80": 212, "after_floor": 220})
        aggressive = pair("aggressive", {**common, "after_target": 153, "after_90": 153,
                                          "after_80": 203, "after_floor": 211})
        selected, rows = E.target_priority_tradeoff_select(parsed, [safe, aggressive])
        self.assertEqual(selected[0].profile, "safe")
        aggressive_row = next(row for row in rows if row["skeleton_profile"] == "aggressive")
        self.assertFalse(aggressive_row["accepted_by_tradeoff_guard"])
        self.assertEqual(aggressive_row["tradeoff_cap"], 1)
        self.assertTrue(aggressive_row["tradeoff_blockers"])

    def test_global_dominance_cannot_revert_to_target_only_on_three_candidates(self):
        """A candidate best on one safety axis must not erase the safe anchor."""
        parsed = SimpleNamespace(target_ratio=0.90, floor_ratio=0.75)
        common = {
            "active_intervals": 252, "after_100": 140, "hard_floor_gap_count": 0,
            "week_boundary_hard_failure_count": 0, "week_boundary_after_target": 20,
            "week_boundary_after_floor": 24, "floor_deficit_sum": 0.10,
            "target_deficit_sum": 0.0, "after_avoidable_overage_fte_sum": 20.0,
            "after_extreme_overage_count": 0, "after_severe_overage_count": 0,
        }
        safe = pair("safe", {**common, "after_target": 150, "after_90": 150,
                              "after_80": 212, "after_floor": 220,
                              "severe_floor_gap_count": 2,
                              "max_consecutive_floor_gaps": 2,
                              "floor_deficit_max": 0.10})
        aggressive = pair("aggressive", {**common, "after_target": 153, "after_90": 153,
                                          "after_80": 203, "after_floor": 211,
                                          "severe_floor_gap_count": 9,
                                          "max_consecutive_floor_gaps": 6,
                                          "floor_deficit_max": 0.40})
        odd = pair("odd", {**common, "after_target": 150, "after_90": 150,
                           "after_80": 210, "after_floor": 218,
                           "severe_floor_gap_count": 0,
                           "max_consecutive_floor_gaps": 8,
                           "floor_deficit_max": 0.50})
        selected, rows = E.target_priority_tradeoff_select(
            parsed, [safe, aggressive, odd]
        )
        self.assertEqual(selected[0].profile, "safe")
        self.assertFalse(any(row["protected_safe_fallback_used"] for row in rows))
        aggressive_row = next(row for row in rows if row["skeleton_profile"] == "aggressive")
        self.assertFalse(aggressive_row["accepted_by_tradeoff_guard"])


class LanguageContractHardening(unittest.TestCase):
    def test_required_language_only_blocks_nonqualified_spillover(self):
        rule = E.LanguageRule("UK", "UK", 7 * 60, 2 * 60, 1, True,
                              {"uk"}, {"uk"})
        parsed = SimpleNamespace(language_working_window_mode="REQUIRED_LANGUAGE_ONLY",
                                 language_rules=[rule])
        english = E.Associate(0, 1, "", "", "", "English", "", "english")
        uk_bilingual = E.Associate(1, 2, "", "", "", "UK bilingual", "", "uk")
        shift = E.Shift(0, "08:00-17:00", 8 * 60, 17 * 60, 540)
        self.assertIsNotNone(E.shift_overlaps_required_language_for_noneligible(
            parsed, english, shift, 0))
        self.assertIsNone(E.shift_overlaps_required_language_for_noneligible(
            parsed, uk_bilingual, shift, 0))

    def test_absent_language_setup_has_stable_three_value_contract(self):
        wb = openpyxl.Workbook()
        values = E._parse_language_rules(wb, [])
        self.assertEqual(len(values), 3)
        self.assertEqual(values, ([], {}, {}))

    def test_non_aligned_window_protects_every_overlapping_quarter(self):
        rule = E.LanguageRule("English", "English", 9 * 60 + 10, 17 * 60 + 5,
                              1, True, {"english"}, {"english"})
        self.assertTrue(rule.overlaps(9 * 60, 15))
        self.assertTrue(rule.overlaps(17 * 60, 15))
        self.assertFalse(rule.overlaps(17 * 60 + 15, 15))

    def test_invalid_coverage_days_fails_closed(self):
        self.assertEqual(E._parse_language_days("Sun-Thu"), {0, 1, 2, 3, 4})
        self.assertEqual(E._parse_language_days("Sunday to Thursday"), {0, 1, 2, 3, 4})
        with self.assertRaises(ValueError):
            E._parse_language_days("garbage")
        warnings = []
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Language Setup"
        ws.append(["Language", "Coverage Start", "Coverage End", "Can Cover Languages",
                   "Active?", "Coverage Group", "Minimum Per Interval", "Coverage Days"])
        ws.append(["International", "00:00", "20:00", "International", "Yes", "International", 1, "not a day"])
        rules, capabilities, windows = E._parse_language_rules(
            wb, [SimpleNamespace(language="International")], warnings
        )
        self.assertTrue(any(item.startswith("HARD_INVALID_LANGUAGE_COVERAGE_DAYS:") for item in warnings))

    def test_multiple_windows_for_one_language_are_retained(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Language Setup"
        ws.append(["Language", "Coverage Start", "Coverage End", "Can Cover Languages",
                   "Active?", "Coverage Group", "Minimum Per Interval"])
        ws.append(["Spanish", "06:00", "12:00", "Spanish", "Yes", "Spanish", 1])
        ws.append(["Spanish", "18:00", "23:00", "Spanish", "Yes", "Spanish", 1])
        rules, _, windows = E._parse_language_rules(
            wb, [SimpleNamespace(language="Spanish")]
        )
        self.assertEqual(len(rules), 2)
        self.assertEqual(windows["spanish"], [(360, 720, True), (1080, 1380, True)])
        parsed = SimpleNamespace(language_working_window_mode="ALL_ROWS", language_windows=windows)
        self.assertEqual(
            E.associate_language_windows(parsed, SimpleNamespace(language="Spanish")),
            [(360, 720), (1080, 1380)],
        )

    def _mutated_nmg(self, mutate):
        source = ROOT / "inputs" / "RC9_2_2_NMG_EN_PRODUCTION_INPUT.xlsx"
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        path = Path(handle.name)
        wb = openpyxl.load_workbook(source)
        mutate(wb)
        wb.save(path)
        return path

    def test_invalid_language_time_is_a_hard_contract_failure(self):
        def mutate(wb):
            ws = wb["Language Setup"]
            header = E._find_header_row(ws, ["language", "minimum"])
            ws.cell(header + 1, 2).value = "09:99"

        path = self._mutated_nmg(mutate)
        try:
            parsed = E.parse_input(path)
            result = E.validate_input_contract(parsed)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any(row["code"] == "HARD_INVALID_LANGUAGE_TIME" for row in result["failures"]))
        finally:
            path.unlink(missing_ok=True)

    def test_unknown_language_active_value_is_a_hard_contract_failure(self):
        def mutate(wb):
            ws = wb["Language Setup"]
            header = E._find_header_row(ws, ["language", "minimum"])
            headers = {E.norm(ws.cell(header, c).value): c for c in range(1, ws.max_column + 1)}
            active_col = next(c for h, c in headers.items() if "active" in h or "enabled" in h)
            ws.cell(header + 1, active_col).value = "Maybe"

        path = self._mutated_nmg(mutate)
        try:
            parsed = E.parse_input(path)
            result = E.validate_input_contract(parsed)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any(row["code"] == "HARD_INVALID_LANGUAGE_ACTIVE" for row in result["failures"]))
        finally:
            path.unlink(missing_ok=True)

    def test_duplicate_roster_identity_is_a_hard_contract_failure(self):
        def mutate(wb):
            ws = wb["Schedule"]
            header = E._find_header_row(ws, ["name", "language"])
            name_col = next(
                c for c in range(1, ws.max_column + 1)
                if "name" in E.norm(ws.cell(header, c).value)
            )
            ws.cell(header + 2, name_col).value = ws.cell(header + 1, name_col).value

        path = self._mutated_nmg(mutate)
        try:
            parsed = E.parse_input(path)
            result = E.validate_input_contract(parsed)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any(row["code"] == "HARD_DUPLICATE_ROSTER_NAME" for row in result["failures"]))
        finally:
            path.unlink(missing_ok=True)

    def test_malformed_requirement_value_is_not_coerced_to_zero(self):
        def mutate(wb):
            ws = E._discover_requirement_sheet(wb, E._instruction_map(wb["Instructions"]))
            header = E._find_header_row(ws, ["sun", "mon"])
            day_col = E._day_columns(ws, header)[0]
            ws.cell(header + 1, day_col).value = "not-a-number"

        path = self._mutated_nmg(mutate)
        try:
            parsed = E.parse_input(path)
            result = E.validate_input_contract(parsed)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any(row["code"] == "HARD_INVALID_REQUIREMENT_VALUE" for row in result["failures"]))
        finally:
            path.unlink(missing_ok=True)

    def test_invalid_clock_tokens_are_not_normalized_with_modulo(self):
        self.assertIsNone(E.minute_of_day("09:99"))
        self.assertIsNone(E.minute_of_day("25:30"))
        self.assertIsNone(E.shift_parts("25:30 - 10:30"))

    def test_malformed_numeric_instruction_is_a_hard_contract_failure(self):
        def mutate(wb):
            ws = wb["Instructions"]
            for row in ws.iter_rows():
                for cell in row:
                    if E.norm(cell.value) == E.norm("Target"):
                        ws.cell(cell.row, cell.column + 1).value = "ninety percent"
                        return
            raise AssertionError("Target instruction row not found")

        path = self._mutated_nmg(mutate)
        try:
            parsed = E.parse_input(path)
            result = E.validate_input_contract(parsed)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any(row["code"] == "HARD_INVALID_INSTRUCTION_NUMBER" for row in result["failures"]))
        finally:
            path.unlink(missing_ok=True)


class LongShiftContractHardening(unittest.TestCase):
    def test_11h_enabled_keeps_long_duration(self):
        durations, enabled = E._parse_duration_set({
            E.norm("Use 11H/3OFF"): "Yes",
            E.norm("Allowed Shift Durations Hours"): "9, 11",
        })
        self.assertTrue(enabled)
        self.assertEqual(durations, {540, 660})

    def test_11h_prohibited_removes_long_duration(self):
        durations, enabled = E._parse_duration_set({
            E.norm("Use 11H/3OFF"): "No",
            E.norm("Allowed Shift Durations Hours"): "9, 11",
        })
        self.assertFalse(enabled)
        self.assertEqual(durations, {540})


class OutputIdentityAndReleaseChain(unittest.TestCase):
    def test_validator_reads_no_break_exception_sheet(self):
        validator = load_validator()
        wb = openpyxl.Workbook()
        active = wb.active
        active.title = "Break Schedule Active"
        active.append(["Associate", "Day", "Shift", "Break Type", "Start", "Duration Minutes", "Status"])
        active.append(["A", "Sun", "09:00 - 18:00", "Break 1", "11:00", 15, "Assigned"])
        exceptions = wb.create_sheet("No-Break Exceptions")
        exceptions.append([
            "Associate", "Day", "Shift", "Language/Skill", "Blocking Rules / Windows",
            "Exception Proof", "Mandatory Comment", "Permission Source", "Review Status",
        ])
        exceptions.append([
            "B", "Sun", "10:00 - 19:00", "UK", "critical window", "INFEASIBLE",
            "NO BREAKS SCHEDULED - critical coverage conflict", "Instructions", "MANUAL REVIEW REQUIRED",
        ])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            path = Path(handle.name)
        try:
            wb.save(path)
            breaks, _ = validator.parse_breaks(path)
            exception = next(row for row in breaks if row["associate"] == "B")
            self.assertTrue(exception["is_exception"])
            self.assertEqual(exception["duration"], 0)
            self.assertIn("MANUAL REVIEW REQUIRED", exception["status"])
        finally:
            path.unlink(missing_ok=True)

    def test_duplicate_output_associate_is_preserved_as_evidence(self):
        validator = load_validator()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Schedule"
        ws.append(["SF Name", "Language", *E.DAY_NAMES])
        ws.append(["A", "English", *(["OFF"] * 7)])
        ws.append(["A", "English", *(["OFF"] * 7)])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            path = Path(handle.name)
        try:
            wb.save(path)
            assignments, _ = validator.parse_output_schedule(path, ["A"])
            self.assertEqual(len(assignments), 1)
            self.assertEqual(len(assignments.duplicate_rows), 1)
        finally:
            path.unlink(missing_ok=True)

    def test_validator_crash_writes_machine_readable_error_and_returns_three(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            jout = root / "validation.json"
            cout = root / "validation.csv"
            result = subprocess.run([
                sys.executable, str(VALIDATOR_PATH), "--input", str(root / "missing.xlsx"),
                "--output", str(root / "also_missing.xlsx"), "--json-out", str(jout),
                "--csv-out", str(cout),
            ], capture_output=True, text=True, check=False)
            self.assertEqual(result.returncode, 3)
            self.assertEqual(json.loads(jout.read_text())["status"], "ERROR")
            self.assertTrue(cout.is_file())

    def test_runner_records_status_before_invoking_packager(self):
        source = (ROOT / "engine" / "RUN_UNIVERSAL_PRODUCTION.py").read_text()
        status_write = source.index("status_path.write_text")
        package_call = source.index("str(PACKAGER)", status_write)
        self.assertLess(status_write, package_call)

    def test_polisher_never_writes_a_zip(self):
        source = (ROOT / "engine" / "production" / "production_output_polisher.py").read_text()
        self.assertNotIn("ZIP_DEFLATED", source)
        self.assertNotIn("archive.write(", source)
        self.assertNotIn("production_zip", source)

    def test_production_output_includes_the_current_control_center(self):
        source = (ROOT / "engine" / "production" / "production_output_polisher.py").read_text()
        self.assertIn("OUTPUT_STYLE_VERSION='RC9.2.2-OUTPUT-UX-RC1'", source)
        self.assertIn("Read Me First", source)
        self.assertIn("Schedule Control Center", source)

    def test_runner_terminates_the_full_engine_process_group(self):
        source = (ROOT / "runners" / "rc921_runner.py").read_text()
        self.assertIn("start_new_session=(os.name == \"posix\")", source)
        self.assertIn("os.killpg(os.getpgid(proc.pid)", source)
        self.assertIn("except KeyboardInterrupt", source)

    def test_wrapper_uses_a_case_lock_and_heartbeat(self):
        source = RUNNER_PATH.read_text()
        self.assertIn("os.O_CREAT | os.O_EXCL", source)
        self.assertIn("RUN_LOCK.json", source)
        self.assertIn("RUN_HEARTBEAT.json", source)
        self.assertIn("Case is already running", source)

    def test_case_lock_blocks_live_owner_and_reclaims_dead_owner(self):
        spec = importlib.util.spec_from_file_location("rc922_case_lock", RUNNER_PATH)
        runner = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = runner
        spec.loader.exec_module(runner)
        with tempfile.TemporaryDirectory() as tmp:
            case_root = Path(tmp)
            runner.acquire_case_lock(case_root, "CASE")
            self.assertTrue((case_root / "RUN_HEARTBEAT.json").is_file())
            with self.assertRaises(RuntimeError):
                runner.acquire_case_lock(case_root, "CASE")
            runner.release_case_lock()
            (case_root / "RUN_LOCK.json").write_text(
                json.dumps({"pid": 999999999, "schedule_id": "OLD"}), encoding="utf-8")
            runner.acquire_case_lock(case_root, "CASE")
            self.assertEqual(json.loads((case_root / "RUN_LOCK.json").read_text())["pid"], os.getpid())
            runner.release_case_lock()
            self.assertFalse((case_root / "RUN_LOCK.json").exists())

    def test_phase_c_zip_is_published_only_after_atomic_verification(self):
        source = (ROOT / "engine" / "production" / "package_phase_c_outputs.py").read_text()
        self.assertIn("temporary.replace(target)", source)
        self.assertIn("temporary.unlink(missing_ok=True)", source)
        self.assertIn("COMPONENT_MANIFEST.json", source)
        self.assertIn("completion_status", source)

    def test_prevalidation_not_validated_state_allows_exact_output_validation(self):
        spec = importlib.util.spec_from_file_location("rc922_runner_core", RUNNER_PATH)
        runner = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = runner
        spec.loader.exec_module(runner)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "PHASE_C_QUALITY_SUMMARY.json").write_text(json.dumps({
                "artifact_state": "FINAL_VERIFIED",
                "contract": {"status": "WARN", "failure_count": 0},
                "safety": {"status": "NOT_VALIDATED"},
                "production_quality_gate": {"status": "WARN"},
            }), encoding="utf-8")
            self.assertTrue(runner.quality_allows_validation(root, 2))

    def test_quality_gate_failure_still_allows_exact_artifact_validation(self):
        spec = importlib.util.spec_from_file_location("rc922_runner_core_fail", RUNNER_PATH)
        runner = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = runner
        spec.loader.exec_module(runner)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "PHASE_C_QUALITY_SUMMARY.json").write_text(json.dumps({
                "artifact_state": "FINAL_VERIFIED",
                "contract": {"status": "PASS", "failure_count": 0},
                "safety": {"status": "NOT_VALIDATED", "hard_fail_count": 0},
                "production_quality_gate": {"status": "FAIL"},
            }), encoding="utf-8")
            self.assertTrue(runner.quality_allows_validation(root, 2))

    def test_validator_limits_next_sunday_to_engine_spill_horizon(self):
        source = VALIDATOR_PATH.read_text()
        self.assertIn(
            "protected_next_sunday_intervals = set(eng.next_sunday_interval_indices(parsed))",
            source,
        )
        self.assertIn("if i not in protected_next_sunday_intervals:", source)

    def test_next_sunday_language_repair_uses_key_day_not_stale_qslot(self):
        source = ENGINE_PATH.read_text()
        self.assertIn("rule_day = 0 if day_scope >= 7 else day_scope", source)
        self.assertNotIn(
            "day=(0 if qslot >= TOTAL_QSLOTS else day_scope)",
            source,
        )

    def test_missing_validated_output_has_nonzero_runner_code(self):
        source = RUNNER_PATH.read_text()
        self.assertIn("rc = 5 if args.skip_independent_validation else 4", source)

    def test_template_exposes_language_window_and_day_controls(self):
        source = (ROOT / "tools" / "build_input_template.py").read_text()
        self.assertIn('(\"Language Working Window\", LANGUAGE_WINDOW_CHOICES)', source)
        self.assertIn('"Coverage Days"', source)
        self.assertIn('showErrorMessage = True', source)

    def test_rc4_colab_notebook_uses_the_canonical_rc922_runner(self):
        notebook = ROOT / "runners" / "RC922_Colab_B_WITH_DRIVE.ipynb"
        self.assertTrue(notebook.is_file())
        source = notebook.read_text(encoding="utf-8")
        self.assertIn("rc922_runner.py", source)
        self.assertIn("/content/drive/MyDrive/RC922_RC5", source)
        self.assertNotIn("rc921_runner.py", source)

        no_drive = ROOT / "runners" / "RC922_Colab_A_NO_DRIVE.ipynb"
        self.assertTrue(no_drive.is_file())
        no_drive_source = no_drive.read_text(encoding="utf-8")
        self.assertIn("RC9_2_2_FIX_VALIDATION_RC5.zip", no_drive_source)
        self.assertIn("rc922_runner.py", no_drive_source)
        self.assertIn("start_new_session", no_drive_source)
        self.assertIn("KeyboardInterrupt", no_drive_source)
        self.assertIn("RESUME = True", no_drive_source)
        self.assertIn("--resume", no_drive_source)

    def test_rc4_helper_runners_import_the_canonical_entry_point(self):
        for name in ("run_deep.py", "run_standard_regression.py", "run_targeted_regression.py"):
            source = (ROOT / "runners" / name).read_text(encoding="utf-8")
            with self.subTest(runner=name):
                self.assertIn("from rc922_runner import main", source)
                self.assertNotIn("from rc921_runner import main", source)

    def test_package_builder_excludes_legacy_colab_notebooks(self):
        source = (ROOT / "tools" / "build_production_package.py").read_text(encoding="utf-8")
        self.assertIn('name.startswith("RC921_")', source)
        self.assertIn('name.endswith(".ipynb")', source)

    def test_failed_independent_validation_forces_business_outcome_blocked(self):
        spec = importlib.util.spec_from_file_location("rc922_runner_reconcile", RUNNER_PATH)
        runner = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = runner
        spec.loader.exec_module(runner)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "BUSINESS_OUTCOME.json").write_text(json.dumps({
                "technical_status": "PASS_WITH_QUALITY_WARNINGS",
                "technical_return_code": 0,
                "production_eligible": True,
                "outcome_code": "FINAL_SCHEDULE_GENERATED_WITH_DECLARED_QUALITY_DEBT",
                "outcome_category": "CONDITIONAL_SUCCESS",
                "headline": "Hard-valid final schedule generated",
                "plain_language_summary": "Warnings remain.",
            }), encoding="utf-8")
            summary = root / "JOB.l6_3_2_3_summary.csv"
            summary.write_text(
                "business_headline,business_message,production_eligible\n"
                "old,old,TRUE\n", encoding="utf-8"
            )
            runner.reconcile_business_outcome_after_validation(root, {
                "status": "FAIL", "return_code": 2, "json": str(root / "INDEPENDENT_VALIDATION.json")
            }, 4)
            outcome = json.loads((root / "BUSINESS_OUTCOME.json").read_text())
            self.assertFalse(outcome["production_eligible"])
            self.assertEqual(outcome["outcome_category"], "QUALITY_BLOCKED")
            self.assertEqual(outcome["independent_validation"]["status"], "FAIL")
            self.assertIn("must not be used", outcome["plain_language_summary"])
            self.assertIn("QUALITY_BLOCKED", summary.read_text())

    def test_pre_solver_failure_never_claims_a_final_schedule(self):
        spec = importlib.util.spec_from_file_location("rc922_runner_reconcile_contract", RUNNER_PATH)
        runner = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = runner
        spec.loader.exec_module(runner)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "BUSINESS_OUTCOME.json").write_text(json.dumps({
                "technical_status": "FAIL_PRE_SOLVER_CONTRACT",
                "technical_return_code": 2,
                "production_eligible": False,
                "outcome_code": "INPUT_OR_RESOURCE_CONTRACT_GAP",
                "outcome_category": "ACTION_REQUIRED",
                "headline": "Input or resource contract prevents schedule generation",
                "plain_language_summary": "The input is impossible.",
            }), encoding="utf-8")
            (root / "CASE_l6_3_2_3_solver_audit.json").write_text(json.dumps({
                "status": "FAIL_PRE_SOLVER_CONTRACT",
                "artifact_state": "NO_ARTIFACT",
            }), encoding="utf-8")
            runner.reconcile_business_outcome_after_validation(root, {
                "status": "NOT_RUN", "return_code": None,
            }, 2)
            outcome = json.loads((root / "BUSINESS_OUTCOME.json").read_text())
            self.assertFalse(outcome["production_eligible"])
            self.assertEqual(outcome["technical_status"], "FAIL_PRE_SOLVER_CONTRACT")
            self.assertNotIn("Final schedule generated", outcome["headline"])
            self.assertEqual(outcome["outcome_code"], "INPUT_OR_RESOURCE_CONTRACT_GAP")

    def test_production_package_carries_rc91_comparator(self):
        source = (ROOT / "tools" / "build_production_package.py").read_text(encoding="utf-8")
        production_start = source.index("def populate_production")
        validation_start = source.index("def populate_validation")
        production_source = source[production_start:validation_start]
        self.assertIn('"evidence/RC9_1_BASELINE.json"', production_source)
        self.assertIn('"evidence/RC9_1_BASELINE_PROVENANCE.txt"', production_source)

    def test_wrapper_reads_validation_json_before_quality_fields(self):
        source = RUNNER_PATH.read_text(encoding="utf-8")
        read_validation = source.index("validation = json.loads(validation_json.read_text")
        quality_field = source.index("'quality_gate_status': validation.get", read_validation)
        self.assertLess(read_validation, quality_field)

    def test_validator_separates_current_week_from_saturday_carry_in(self):
        source = VALIDATOR_PATH.read_text(encoding="utf-8")
        self.assertIn("current_week_before=[[] for _ in range(horizon)]", source)
        self.assertIn("if current_week_before[slot]:", source)
        self.assertIn("prior_carry_in_raw", source)

    def test_stage2_budget_protects_anchor_and_adaptive_search(self):
        source = ENGINE_PATH.read_text(encoding="utf-8")
        self.assertIn(
            "stage2_guard_reserve_sec = bounded_stage2_guard_reserve_seconds(",
            source,
        )
        self.assertIn("stage2_adaptive_reserve_sec = min(", source)
        recovery_block = source[source.index("breakability_recovery_deadline = min("):]
        recovery_block = recovery_block[:recovery_block.index("joint_refinement_deadline")]
        self.assertIn("stage2_guard_reserve_sec", recovery_block)
        self.assertIn("stage2_adaptive_reserve_sec", recovery_block)

    def test_stage2_reserve_is_large_enough_for_measured_break_depth(self):
        guard = E.bounded_stage2_guard_reserve_seconds(3600, 904)
        adaptive = min(E.BREAK_MIN_MEANINGFUL_SLICE_SEC, 904 - guard)
        self.assertGreaterEqual(guard, E.BREAK_MIN_MEANINGFUL_SLICE_SEC)
        self.assertGreaterEqual(adaptive, E.BREAK_MIN_MEANINGFUL_SLICE_SEC)
        self.assertLessEqual(guard + adaptive, 904)
        self.assertEqual(guard, 240)

    def test_stage2_anchor_keeps_unknown_high_coverage_skeleton_eligible(self):
        unknown = E.SkeletonSolution(
            "coverage_champion_unknown", "FEASIBLE", 0.0, 0.0, [], [],
            {"minimum_exception_count": None},
        )
        proven_zero = E.SkeletonSolution(
            "lower_coverage_zero_exception", "FEASIBLE", 0.0, 0.0, [], [],
            {"minimum_exception_count": 0},
        )
        diagnostic_probe = E.SkeletonSolution(
            "hard_feasibility_probe", "FEASIBLE", 0.0, 0.0, [], [],
            {"minimum_exception_count": 0},
        )
        candidates = E.coverage_first_stage2_anchor_candidates(
            [proven_zero, unknown, diagnostic_probe]
        )
        self.assertEqual(
            [s.profile for s in candidates],
            ["lower_coverage_zero_exception", "coverage_champion_unknown"],
        )

    def test_stage2_anchor_defers_only_proven_positive_minimum(self):
        proven_positive = E.SkeletonSolution(
            "proven_positive", "FEASIBLE", 0.0, 0.0, [], [],
            {"minimum_exception_count": 1},
        )
        unknown = E.SkeletonSolution(
            "unknown", "FEASIBLE", 0.0, 0.0, [], [],
            {"minimum_exception_count": None},
        )
        candidates = E.coverage_first_stage2_anchor_candidates(
            [proven_positive, unknown]
        )
        self.assertEqual([s.profile for s in candidates], ["unknown"])


class FixedRequestAndNestingGuards(unittest.TestCase):
    def test_fixed_off_is_independent_of_preference_hardness(self):
        source = ENGINE_PATH.read_text()
        self.assertGreaterEqual(source.count('(parsed.fixed_enabled and fixed_kind == "off")'), 2)
        self.assertIn('(hard.fixed and parsed.fixed_enabled and fixed_kind == "off")', source)
        self.assertNotIn('parsed.hard_off and (pref_kind == "off" or fixed_kind == "off")', source)

    def test_independent_validator_checks_flexible_nesting(self):
        source = VALIDATOR_PATH.read_text()
        self.assertIn("NESTING_GROUP_SCHEDULE_MISMATCH", source)


class ProductionWorkbookContract(unittest.TestCase):
    def test_supported_nmg_workbook_parses_and_preflights(self):
        workbook = ROOT / "inputs" / "RC9_2_2_NMG_EN_PRODUCTION_INPUT.xlsx"
        parsed = E.parse_input(workbook)
        preflight = E.validate_input_contract(parsed)
        self.assertEqual(preflight["status"], "PASS", preflight)
        self.assertEqual(len(parsed.associates), 42)
        self.assertEqual(len(parsed.shifts), 10)
        self.assertEqual(parsed.run_stage, "FULL_SCHEDULE")
        self.assertEqual(parsed.run_depth, "QUICK")
        self.assertEqual(parsed.quality_gate_mode, "fail")
        self.assertEqual(parsed.break_concurrency_gate_mode, "fail")
        self.assertEqual(parsed.next_sunday_balance_gate_mode, "fail")

    def test_validator_exports_distribution_evidence(self):
        source = VALIDATOR_PATH.read_text()
        for key in (
            "avoidable_overage_top10_concentration",
            "avoidable_overage_max_consecutive_intervals",
            "whole_week_overage_cap_violation_count",
            "whole_week_imbalance_violation_count",
            "break_spacing_beyond_normal_count",
        ):
            self.assertIn(key, source)

    def test_all_shipped_workbooks_have_enforced_validation_alerts(self):
        namespace = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        workbooks = sorted((ROOT / 'inputs').glob('*.xlsx'))
        self.assertEqual(len(workbooks), 8)
        for workbook in workbooks:
            with self.subTest(workbook=workbook.name), ZipFile(workbook) as archive:
                validations = []
                formula_errors = []
                for member in archive.namelist():
                    if not (member.startswith('xl/worksheets/sheet') and member.endswith('.xml')):
                        continue
                    root = ElementTree.fromstring(archive.read(member))
                    validations.extend(root.findall('.//x:dataValidation', namespace))
                    formula_errors.extend(
                        node.text or '' for node in root.findall('.//x:f', namespace)
                        if any(token in (node.text or '') for token in ('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A'))
                    )
                self.assertTrue(validations)
                self.assertTrue(all(node.attrib.get('showErrorMessage') == '1' for node in validations))
                self.assertFalse(formula_errors)

    def test_scenario_manifest_matches_every_shipped_workbook(self):
        manifest = json.loads((ROOT / 'SCENARIOS.json').read_text(encoding='utf-8'))
        for row in manifest['scenarios']:
            with self.subTest(scenario=row['scenario_id']):
                digest = hashlib.sha256((ROOT / 'inputs' / row['input']).read_bytes()).hexdigest()
                self.assertEqual(digest, row['input_sha256'])


if __name__ == "__main__":
    unittest.main()
