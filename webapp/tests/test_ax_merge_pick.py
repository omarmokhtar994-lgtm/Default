# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AX (owner, 2026-10-22: tick the LOBs to see together, two out of five, not all of them): the RTA's merged
day takes the LOBs ticked in the View pick. The LOB the page is opened on is always one of them. The made-up program
has three LOBs: Tier 1 (the fixture week), Tier 2 ("Associate B..." people) and Tier 3 ("Associate C..." people),
all on Cairo time so only the choice is under test."""
import html
import io
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from webapp.programs import ProgramBook
from webapp.tests.test_ap_merge_day import CAIRO, renamed
from webapp.tests.test_ready import make_ready
from webapp.tests.test_runs import make_app, sign_in, token, wait
from webapp.tests.test_schedules import REPO

import os

# Phase AY: pages default to the zone of the computer (the pc_tz cookie); the browser here is in Egypt, as the expected times are
os.environ["TZ"] = "Africa/Cairo"

WED = date(2026, 10, 14)


def renamed_c(path: Path) -> Path:
    wb = load_workbook(renamed(path))
    ws = wb["Schedule"]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("Associate B"):
                cell.value = cell.value.replace("Associate B", "Associate C")
    wb.save(path)
    return path


class ThePickOfLobs(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        files = (make_ready(cls.dir / "one.xlsx"), renamed(make_ready(cls.dir / "two.xlsx")),
                 renamed_c(make_ready(cls.dir / "three.xlsx")))
        cls.app, cls.store, *_ = make_app(VALIDATOR_ROOT=str(REPO))
        cls.store.add_user("omar", "Omar", "Owner-pass-123", is_admin=True, must_change=False)
        programs = ProgramBook(cls.store)
        pid = programs.add_program("SAKS")
        cls.one, cls.two, cls.three = (programs.add_lob(pid, n) for n in ("Tier 1", "Tier 2", "Tier 3"))
        cls.admin = sign_in(cls.app, "omar", "Owner-pass-123")
        for key, path in zip((cls.one, cls.two, cls.three), files):
            cls.admin.post("/runs", data={"csrf_token": token(cls.admin), "kind": "ready", "mode": "QUICK",
                                          "program": key, "week_start": "2026-10-11",
                                          "workbook": (io.BytesIO(path.read_bytes()), "week.xlsx")},
                           content_type="multipart/form-data")
            wait(cls.store, cls.store.list_runs()[0]["id"], statuses=("DONE", "REJECTED", "FAILED"))
        for run in cls.store.list_runs():
            cls.store.update_run(run["id"], tz=CAIRO)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)

    def setUp(self):
        self.admin.get("/day", query_string={"program": self.one, "date": WED.isoformat(), "merge": ""})

    def text(self, **query):
        query = {"program": self.one, "date": WED.isoformat(), "tz": CAIRO, **query}
        return html.unescape(self.admin.get("/day", query_string=query).get_data(as_text=True))

    def test_two_lobs_ticked_show_those_two_and_not_the_third(self):
        page = self.text(merge="1", pick="1", lobs=[self.two])
        self.assertIn("Associate 027", page)
        self.assertIn("Associate B027", page)
        self.assertNotIn("Associate C027", page)

    def test_the_lob_the_page_is_opened_on_is_always_shown(self):
        page = self.text(merge="1", pick="1", lobs=[self.two, self.three], program=self.one)
        self.assertIn("Associate 027", page)

    def test_the_page_says_which_lobs_it_shows_together(self):
        page = self.text(merge="1", pick="1", lobs=[self.two])
        self.assertIn("Tier 1 + Tier 2 together.", page)
        self.assertNotIn("Every LOB of the program together", page)

    def test_every_lob_ticked_shows_all_three(self):
        page = self.text(merge="1", pick="1", lobs=[self.two, self.three])
        for who in ("Associate 027", "Associate B027", "Associate C027"):
            self.assertIn(who, page)

    def test_nothing_else_ticked_shows_the_one_lob(self):
        page = self.text(merge="1", pick="1")
        self.assertNotIn("Associate B027", page)
        self.assertNotIn("Every LOB of the program together", page)

    def test_the_choice_is_kept_when_a_link_only_says_merge(self):
        self.text(merge="1", pick="1", lobs=[self.three])
        page = self.text(merge="1")
        self.assertIn("Associate C027", page)
        self.assertNotIn("Associate B027", page)

    def test_a_name_that_is_not_a_lob_of_the_program_is_ignored(self):
        page = self.text(merge="1", pick="1", lobs=[self.two, "Elsewhere LOB"])
        self.assertIn("Associate B027", page)
        self.assertNotIn('value="Elsewhere LOB"', page)
        self.assertNotIn("Associate C027", page)

    def test_the_pick_lists_each_lob_with_a_tick_box_ticked_as_chosen(self):
        page = self.text(merge="1", pick="1", lobs=[self.three])
        for lob in (self.two, self.three):
            self.assertIn(f'name="lobs" value="{lob}"', page)
        self.assertRegex(page, rf'name="lobs" value="{self.three}"\s+checked')
        self.assertNotRegex(page, rf'name="lobs" value="{self.two}"\s+checked')

    def test_the_week_page_follows_the_same_choice(self):
        self.text(merge="1", pick="1", lobs=[self.two])
        got = html.unescape(self.admin.get("/week", query_string={
            "program": self.one, "week": "2026-10-11", "tz": CAIRO, "merge": "1"}).get_data(as_text=True))
        self.assertIn("<title>SAKS, Tier 1 + Tier 2, week of", got)


class ThePickInTheBrowser(ThePickOfLobs):
    """The tick boxes work as a person uses them: open the pick, tick a LOB, press Show together."""
    SCREENS = Path(__file__).resolve().parents[2] / "evidence" / "phase_ax" / "screens"

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import threading
        from playwright.sync_api import sync_playwright
        from werkzeug.serving import make_server
        cls.server = make_server("127.0.0.1", 0, cls.app, threaded=True)
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        cls.pw = sync_playwright().start()
        cls.browser = cls.pw.chromium.launch(headless=True, executable_path="/opt/pw-browsers/chromium")
        cls.SCREENS.mkdir(parents=True, exist_ok=True)

    @classmethod
    def tearDownClass(cls):
        cls.browser.close()
        cls.pw.stop()
        cls.server.shutdown()
        super().tearDownClass()

    def test_ticking_one_other_lob_and_pressing_show_together_shows_two_of_three(self):
        page = self.browser.new_page(viewport={"width": 1280, "height": 860})
        self.addCleanup(page.close)
        base = f"http://127.0.0.1:{self.server.server_port}"
        page.goto(base + "/login")
        page.get_by_label("Username").fill("omar")
        page.get_by_label("Password").fill("Owner-pass-123")
        page.get_by_role("button", name="Sign in").click()
        page.wait_for_url(base + "/")
        page.goto(f"{base}/day?program={self.one}&date={WED.isoformat()}&tz={CAIRO}")
        page.get_by_text("View: This LOB").click()
        page.get_by_role("checkbox", name="Tier 3").check()
        page.screenshot(path=str(self.SCREENS / "pick_open.png"))
        page.get_by_role("button", name="Show together").click()
        page.wait_for_load_state()
        body = page.content()
        self.assertIn("Associate C027", body)
        self.assertNotIn("Associate B027", body)
        self.assertIn("View: Tier 1 + Tier 3", body)
        page.wait_for_timeout(700)  # the page-to-page crossfade
        page.screenshot(path=str(self.SCREENS / "two_of_three.png"))


if __name__ == "__main__":
    unittest.main()
