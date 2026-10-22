# © 2026 Omar Mokhtar. All rights reserved.
"""Phase AV (owner, 2026-10-22: "for late and early leave in export it has to have a duration, also I think we are
missing a lot of reports now that can be added based on all updates done"). Late and Left early rows say how long;
incidents, booked sessions, break and coaching requests and swap requests have their own tabs."""
import io
import unittest

from openpyxl import load_workbook

from webapp.tests import test_exports as _ex
from webapp.tests.test_exports import WED

KEY = "AE/AR B2B"


class TheNewExports(_ex.TheExports):
    def rows(self, kind, name, **kw):
        data, *_ = self.export(kinds=[kind], **kw)
        return self.sheet(data, name)

    def test_late_says_how_long_and_left_early_too(self):
        self.days.set_status(KEY, WED, "Associate 021", "Left early", self.sara, start="15:00")
        got = {r["Associate"]: r for r in self.rows("attendance", "Attendance")}
        self.assertEqual((got["Associate 001"]["Status"], got["Associate 001"]["Minutes"]), ("Late", 60))  # 12:00 shift, in at 13:00
        self.assertEqual((got["Associate 021"]["Status"], got["Associate 021"]["Minutes"]), ("Left early", 120))  # out 15:00, shift to 17:00

    def test_a_missed_break_is_an_incident_row(self):
        self.days.move_break(KEY, WED, "Associate 008", 1, "13:10", self.lina, missed=True)
        got = self.rows("incidents", "Incidents (missed breaks)")
        self.assertEqual([(r["Associate"], r["Incident"], r["Break"], r["Recorded by"]) for r in got],
                         [("Associate 008", "Missed break", 2, "Lina")])

    def test_a_booking_is_a_session_row_per_person(self):
        names = ["Associate 021", "Associate 037"]
        start = next(s["start"] for s in self.days.meeting_slots(KEY, WED, names, 30, "09:00", "16:00"))
        self.days.book_session(KEY, WED, names, start, 30, "Coaching", self.sara, with_whom="Team lead", why="Weekly")
        got = self.rows("sessions", "Booked sessions")
        self.assertEqual(sorted((r["Associate"], r["Minutes"], r["People in it"]) for r in got),
                         [(n, 30, 2) for n in names])
        self.assertEqual(len({r["Session"] for r in got}), 1)

    def test_the_request_tabs_open_even_when_empty(self):
        for kind in ("break_requests", "swaps"):
            data, *_ = self.export(kinds=[kind])
            self.assertTrue(load_workbook(io.BytesIO(data)).sheetnames[1])


if __name__ == "__main__":
    unittest.main()
